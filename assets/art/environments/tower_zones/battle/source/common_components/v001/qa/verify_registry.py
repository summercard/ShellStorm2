import openpyxl,copy,hashlib,json
from pathlib import Path

ROOT=Path('/Users/summercards/ShellStorm2')
OUT=ROOT/'assets/art/environments/tower_zones/battle/source/common_components/v001'
a=openpyxl.load_workbook(OUT/'qa/registry_before.xlsx')
b=openpyxl.load_workbook(ROOT/'assets/registry/ShellStorm2_美术资产台账_v001.xlsx')
changed=[]
for sa in a:
    sb=b[sa.title]
    for row in sa.iter_rows(max_row=89):
        for ca in row:
            if ca.value!=sb[ca.coordinate].value:changed.append((sa.title,ca.coordinate))
s=b['3D-场景通用']
styles=[]
for col in range(1,17):
    for prop in ['font','fill','alignment','border','number_format','protection']:
        if copy.copy(getattr(s.cell(89,col),prop))!=copy.copy(getattr(s.cell(90,col),prop)):styles.append((col,prop))
meta=json.loads((OUT/'qa/registry_update.json').read_text())
sha=hashlib.sha256((OUT/'战局区块_通用组件库_v001.blend').read_bytes()).hexdigest()
passed=not changed and not styles and s['A90'].value=='ENV-BATTLE-L01-COMMON-COMPONENT-LIBRARY' and sha==meta['source_sha256']
report={'passed':passed,'preexisting_cell_changes':changed,'new_row':90,'style_differences_from_template':styles,'source_sha256':sha}
(OUT/'qa/registry_validation.json').write_text(json.dumps(report,ensure_ascii=False,indent=2))
print(json.dumps(report,ensure_ascii=False));raise SystemExit(0 if passed else 1)
