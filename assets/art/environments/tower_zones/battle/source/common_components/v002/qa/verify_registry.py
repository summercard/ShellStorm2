import openpyxl,copy,hashlib,json
from pathlib import Path
ROOT=Path('/Users/summercards/ShellStorm2');OUT=ROOT/'assets/art/environments/tower_zones/battle/source/common_components/v002'
a=openpyxl.load_workbook(OUT/'qa/registry_before.xlsx');b=openpyxl.load_workbook(ROOT/'assets/registry/ShellStorm2_美术资产台账_v001.xlsx')
changes=[]
for sa in a:
 sb=b[sa.title]
 for row in sa:
  for ca in row:
   if ca.value!=sb[ca.coordinate].value:changes.append((sa.title,ca.coordinate))
expected={('3D-场景通用',c+'90') for c in ['E','F','O','P']}
meta=json.loads((OUT/'qa/registry_update.json').read_text());sha=hashlib.sha256((OUT/'战局区块_通用组件库_v002.blend').read_bytes()).hexdigest()
styles=[];sa=a['3D-场景通用'];sb=b['3D-场景通用']
for col in range(1,17):
 for prop in ['font','fill','alignment','border','number_format','protection']:
  if copy.copy(getattr(sa.cell(90,col),prop))!=copy.copy(getattr(sb.cell(90,col),prop)):styles.append((col,prop))
report={'passed':set(changes)==expected and not styles and sha==meta['source_sha256'],'changes':changes,'style_differences':styles,'source_sha256':sha}
(OUT/'qa/registry_validation.json').write_text(json.dumps(report,ensure_ascii=False,indent=2));print(json.dumps(report,ensure_ascii=False));raise SystemExit(0 if report['passed'] else 1)
