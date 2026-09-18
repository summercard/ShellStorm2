"""Read-only cell/formula/style comparison against this task's ledger backup."""
from pathlib import Path
from copy import copy
import hashlib
import json
import openpyxl

ROOT = Path('/Users/summercards/ShellStorm2')
OUT = ROOT / 'assets/art/environments/tower_zones/battle/source/room_instances/main_room_02/v003'
a = openpyxl.load_workbook(OUT / 'qa/ledger_before.xlsx')
b = openpyxl.load_workbook(ROOT / 'assets/registry/ShellStorm2_美术资产台账_v001.xlsx')
changes, styles = [], []
assert a.sheetnames == b.sheetnames
for sa in a:
    sb = b[sa.title]
    assert (sa.max_row, sa.max_column) == (sb.max_row, sb.max_column)
    for row in sa:
        for ca in row:
            cb = sb[ca.coordinate]
            if ca.value != cb.value:
                changes.append((sa.title, ca.coordinate))
            for prop in ['font', 'fill', 'alignment', 'border', 'number_format', 'protection']:
                if copy(getattr(ca, prop)) != copy(getattr(cb, prop)):
                    styles.append((sa.title, ca.coordinate, prop))
expected = {('3D-场景通用', col + '73') for col in ['B', 'E', 'F', 'N', 'O', 'P']}
update = json.loads((OUT / 'qa/ledger_update.json').read_text())
digest = hashlib.sha256((OUT / 'env_battle_l01_main_02_data_room_layout_source_v003.blend').read_bytes()).hexdigest()
passed = set(changes) == expected and not styles and digest == update['source_sha256']
report = dict(passed=passed, changes=changes, style_differences=styles, source_sha256=digest)
(OUT / 'qa/ledger_validation.json').write_text(json.dumps(report, ensure_ascii=False, indent=2))
print(json.dumps(report, ensure_ascii=False))
raise SystemExit(0 if passed else 1)
