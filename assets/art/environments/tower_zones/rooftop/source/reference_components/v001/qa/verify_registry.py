import openpyxl,json,hashlib,copy
from pathlib import Path
R=Path('/Users/summercards/ShellStorm2');O=Path(__file__).resolve().parents[1]
a=openpyxl.load_workbook(O/'qa/registry_before.xlsx');b=openpyxl.load_workbook(R/'assets/registry/ShellStorm2_美术资产台账_v001.xlsx')
unexpected=[];styles=[]
for sa in a:
    sb=b[sa.title]
    for row in sa:
        for c in row:
            if (sa.title=='资产主表' and c.row==429) or (sa.title=='3D-场景通用' and 97<=c.row<=134):continue
            target=sb[c.coordinate]
            if c.value!=target.value:unexpected.append([sa.title,c.coordinate,str(c.value),str(target.value)])
            if c.value is not None:
                for prop in ['font','fill','alignment','border','number_format','protection']:
                    if copy.copy(getattr(c,prop))!=copy.copy(getattr(target,prop)):styles.append([sa.title,c.coordinate,prop])
cat=json.loads((O/'component_packages_v001/catalog.json').read_text());ids=[b['3D-场景通用'].cell(i,1).value for i in range(97,135)]
sha=hashlib.sha256((O/'天台区块_参考组件库_v001.blend').read_bytes()).hexdigest()
checks={'existing_values_preserved':not unexpected,'existing_styles_preserved':not styles,'all_package_ids_registered':ids[1:]==[p['package_id'] for p in cat],'master_hash_matches':b['资产主表']['T429'].value==sha,'parent_rows_agree':ids[0]==b['资产主表']['A429'].value,'canonical_sheet_count_preserved':a.sheetnames==b.sheetnames}
report={'passed':all(checks.values()),'checks':checks,'unexpected_changes':unexpected,'style_changes':styles,'source_sha256':sha,'package_rows':37}
(O/'qa/registry_validation.json').write_text(json.dumps(report,ensure_ascii=False,indent=2));print(json.dumps(report,ensure_ascii=False));raise SystemExit(0 if report['passed'] else 1)
