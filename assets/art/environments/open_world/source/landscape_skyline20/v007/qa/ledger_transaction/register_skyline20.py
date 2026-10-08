import argparse
import ctypes
import hashlib
import json
import re
import shutil
import subprocess
import sys
import traceback
from collections import Counter
from copy import copy
from datetime import date, datetime
from pathlib import Path
from zipfile import ZipFile
from xml.etree import ElementTree as ET

ROOT = Path(__file__).resolve().parents[9]
WORK = Path(__file__).resolve().parent
SOURCE = WORK.parent.parent
sys.dont_write_bytecode = True
sys.path.insert(0, str(ROOT / 'scripts'))
sys.path.insert(0, str(ROOT / 'tools/asset_pipeline'))
from openpyxl import load_workbook
from ledger_registry import LedgerIndex
from split_asset_ledger import (
    HEADER_ROW, FIRST_DATA_ROW, COLUMN_COUNT, CONTENT_COLUMNS, DERIVED_COLUMNS,
    read_source_rows, _row_digest, sheet_digest, col_digest, rescope_asset_sheet,
)
from check_asset_registry import ACCEPTED_STATUSES

PAGE = '3D-场景通用'
STATUS = 'validated'
NOTE = '技术通过、美术待审；reference_art_approved=false；runtime_integrated=false；未正式入库/未正式接入；无碰撞，不制作底层室内。'
NAMES = {
    'floor_facade_module': '重复楼层外立面模块',
    'facade_hvac_unit': '外立面空调单元',
    'roof_guardrail_kit': '屋顶护栏与结构套件',
    'roof_service_hut': '屋顶服务小屋',
    'roof_hvac_unit': '屋顶空调单元',
    'roof_satellite_dish': '屋顶卫星天线',
    'roof_truss_tower': '屋顶桁架通信塔',
    'facade_billboard': '外立面广告牌',
}


def sha(path):
    h = hashlib.sha256()
    with Path(path).open('rb') as f:
        for chunk in iter(lambda: f.read(1024 * 1024), b''):
            h.update(chunk)
    return h.hexdigest()


def read_json(path):
    return json.loads(Path(path).read_text(encoding='utf-8'))


def write_json(path, value):
    Path(path).write_text(json.dumps(value, ensure_ascii=False, indent=2, default=str) + '\n', encoding='utf-8')


def relative(path):
    return Path(path).relative_to(ROOT).as_posix()


# 在现有门禁返回时收集完整问题数组，不修改门禁源码、不依赖截断的前40条。
RUNNER = '''import sys, runpy, json, pathlib
sys.dont_write_bytecode=True
import openpyxl
mapping=json.loads(sys.argv[1]); full_path=pathlib.Path(sys.argv[2]); script=sys.argv[3]; args=sys.argv[4:]
original=openpyxl.load_workbook
def load(path,*a,**kw):
    key=str(pathlib.Path(path).resolve())
    return original(mapping.get(key,path),*a,**kw)
openpyxl.load_workbook=load
captured={}
def profile(frame,event,arg):
    if event=='return' and frame.f_code.co_name in ('check_registry','main'):
        loc=frame.f_locals
        if frame.f_code.co_name=='check_registry' and 'issues' in loc:
            captured['issues']=dict(loc['issues'])
        if frame.f_code.co_name=='main' and 'failures' in loc:
            captured['failures']=loc['failures']
    return profile
sys.setprofile(profile); sys.argv=[script]+args
try:
    runpy.run_path(script,run_name='__main__')
finally:
    sys.setprofile(None)
    full_path.write_text(json.dumps(captured,ensure_ascii=False,indent=2,default=str)+'\\n',encoding='utf-8')
'''


def gates(tag, mapping=None):
    specs = {
        'structure': ('scripts/check_asset_registry.py', ['--project-root', str(ROOT), '--scope', 'structure']),
        'full_scenes': ('scripts/check_asset_registry.py', ['--project-root', str(ROOT), '--scope', 'full', '--ledger', 'scenes']),
        'baseline': ('tools/asset_pipeline/verify_ledger_split.py', ['--project-root', str(ROOT)]),
        'naming': ('scripts/check_asset_runtime_naming.py', ['--json']),
        'refs': ('scripts/check_ledger_refs.py', ['--project-root', str(ROOT)]),
    }
    results = {}
    for name, (script, args) in specs.items():
        output = WORK / f'{tag}_{name}.json'
        full = WORK / f'{tag}_{name}_complete.json'
        if name != 'naming':
            args += ['--json-output', str(output)]
        command = [sys.executable, '-I', '-X', 'utf8', '-c', RUNNER,
                   json.dumps(mapping or {}), str(full), str(ROOT / script), *args]
        process = subprocess.run(command, cwd=WORK, capture_output=True, text=True, encoding='utf-8')
        (WORK / f'{tag}_{name}.log').write_text(process.stdout + '\nSTDERR:\n' + process.stderr, encoding='utf-8')
        if name == 'naming':
            output.write_text(process.stdout, encoding='utf-8')
        data = read_json(output) if output.is_file() else {'execution_error': process.stderr}
        complete = read_json(full) if full.is_file() else {}
        if name in ('structure', 'full_scenes'):
            errors = [{'kind': k, **v} for k, values in complete.get('issues', {}).items() for v in values]
            assert len(errors) == sum(data.get('issue_counts', {}).values()), '完整问题捕获数量不符'
        elif name == 'baseline':
            errors = complete.get('failures', [])
            assert len(errors) == data.get('failure_count'), '完整基线问题捕获数量不符'
        elif name == 'refs':
            errors = [{'kind': 'pending', 'path': p} for p in data.get('pending', [])]
        else:
            errors = [] if process.returncode == 0 else [{'kind': 'naming_debt', 'data': data}]
        results[name] = {'exit': process.returncode, 'errors': errors, 'error_count': len(errors),
                         'result': relative(output), 'log': relative(WORK / f'{tag}_{name}.log')}
        print(f'{tag}/{name}: exit={process.returncode} errors={len(errors)}', flush=True)
    write_json(WORK / f'{tag}_gates.json', results)
    return results


def compare(before, after):
    result = {}
    for key in before:
        def signature(error):
            value = dict(error)
            if value.get('kind') == 'row_content_mutated' and isinstance(value.get('row'), list):
                value['row'] = [v for i, v in enumerate(value['row'], 1) if i not in DERIVED_COLUMNS]
            return json.dumps(value, ensure_ascii=False, sort_keys=True)
        b = {signature(e) for e in before[key]['errors']}
        a = {signature(e) for e in after[key]['errors']}
        result[key] = {'before_exit': before[key]['exit'], 'after_exit': after[key]['exit'],
                       'before_errors': len(b), 'after_errors': len(a),
                       'new': [json.loads(x) for x in sorted(a - b)],
                       'resolved': [json.loads(x) for x in sorted(b - a)], 'unchanged': len(a & b)}
    return result


def assert_no_new(diff):
    assert not any(v['new'] for v in diff.values()), '本次候选新增门禁错误，禁止提交'


def facts():
    index = LedgerIndex.load(ROOT)
    domain = index.domain_for_key('scenes')
    catalog = read_json(SOURCE / 'component_catalog.json')
    qa = read_json(SOURCE / 'qa/qa_report.json')
    runtime_path = ROOT / 'assets/art/environments/open_world/runtime/landscape_skyline20/asset_manifest.json'
    runtime = read_json(runtime_path)
    assert STATUS in ACCEPTED_STATUSES
    assert qa['reference_art_approved'] is False and qa['runtime_integrated'] is False
    assert runtime['asset_id'] == catalog['asset_id'] == qa['asset_id']
    assert qa['version'] == catalog['version'] == runtime['version'] == 'v007'
    assert qa['godot_runtime_probe'] == {'fail': 0, 'pass': 1430, 'status': 'passed'}
    assert len(catalog['components']) == 8
    hashes = {}
    for key, hash_key in [('source_blend', 'source_sha256'), ('optimized_blend', 'optimized_sha256')]:
        p = ROOT / catalog[key]
        hashes[relative(p)] = sha(p)
        assert hashes[relative(p)] == catalog[hash_key]
    assert hashes[catalog['source_blend']] == qa['source_hash']
    assert hashes[catalog['optimized_blend']] == qa['optimized_hash']
    export = read_json(SOURCE / 'export/v007/export_manifest.json')
    assert export['source'] == catalog['source_blend'] and export['optimized'] == catalog['optimized_blend']
    assert export['source_sha256'] == catalog['source_sha256'] and export['optimized_sha256'] == catalog['optimized_sha256']
    records = [{'asset_id': catalog['asset_id'], 'slug': 'landscape_root', 'name': 'SKYLINE20二十层景观建筑',
                'prefab': runtime['prefab'], 'glb': None,
                'size': '主体50×50m；20层；层距3.8m；包络53.23×50.93×86.16m',
                'spec': '主体50×50m；20层×3.8m；包络53.23×50.93×86.16m；8组件/37实例；7720三角；屋顶4572三角(59.22%)',
                'instances': 37, 'triangles': 7720}]
    for x in catalog['components']:
        mp = ROOT / Path(x['prefab']).parent / 'asset_manifest.json'
        manifest = read_json(mp)
        assert manifest['asset_id'] == manifest['component_id'] == x['asset_id'] == x['component_id']
        assert manifest['parent_asset_id'] == catalog['asset_id'] and manifest['version'] == 'v007'
        assert manifest['glb'] == x['glb'] and manifest['prefab'] == x['prefab']
        assert manifest['collision'] == 'none_visual_only'
        hashes[x['glb']] = sha(ROOT / x['glb'])
        assert hashes[x['glb']] == manifest['glb_sha256'] == x['glb_sha256']
        hashes[relative(mp)] = sha(mp)
        size = '×'.join(f'{float(v):.4f}'.rstrip('0').rstrip('.') for v in x['bbox_m']) + 'm（Blender XYZ）'
        records.append({'asset_id': x['asset_id'], 'slug': x['slug'], 'name': 'SKYLINE20' + NAMES[x['slug']],
                        'prefab': x['prefab'], 'glb': x['glb'], 'size': size,
                        'spec': f"包络{size}；单组件{x['triangles_after']}三角；{len(x['instance_transforms'])}实例；bottom_center；公共色盘",
                        'instances': len(x['instance_transforms']), 'triangles': x['triangles_after']})
    for x in records:
        p = ROOT / x['prefab']
        text = p.read_text(encoding='utf-8')
        assert '[gd_scene' in text and 'type="Node3D"' in text
        assert not any(v in text for v in ('CollisionShape3D', 'StaticBody3D', 'script = ExtResource'))
        hashes[x['prefab']] = sha(p)
        x['prefab_sha256'] = hashes[x['prefab']]
    assert sum(x['instances'] * x['triangles'] for x in records[1:]) == 7720
    assert sum(x['instances'] for x in records[1:]) == 37
    root_text = (ROOT/runtime['prefab']).read_text(encoding='utf-8')
    refs = re.findall(r'path="res://([^"]+)"', root_text)
    assert set(refs) == {x['prefab'] for x in records[1:]}
    assert len(re.findall(r'parent="\." instance=', root_text)) == 37
    for p in [SOURCE / 'component_catalog.json', SOURCE / 'qa/qa_report.json', runtime_path,
              SOURCE / 'export/v007/export_manifest.json']:
        hashes[relative(p)] = sha(p)
    return index, domain, catalog, records, hashes


def prepare():
    index, domain, catalog, records, hashes = facts()
    ids = {x['asset_id'] for x in records}
    seen = {}
    for d in index.domains:
        wb = load_workbook(d.path)
        hashes[d.relative_path] = sha(d.path)
        for row, values in read_source_rows(wb['资产主表']):
            if values[0] in ids:
                assert values[0] not in seen and d.key == 'scenes', 'ID重复或已属于其他域'
                seen[values[0]] = row
    baseline_path = ROOT / 'assets/registry/ledger_split_baseline.json'
    for p in [index.master_path, baseline_path, ROOT / 'assets/registry/ledger_index.json']:
        hashes[relative(p)] = sha(p)
    baseline = read_json(baseline_path)
    ws = load_workbook(domain.path)[PAGE]
    assert baseline['sheet_digests'][PAGE] == sheet_digest(ws), '授权Prefab页存在历史指纹漂移，禁止刷新洗白'
    before = gates('before')
    assert all(sha(ROOT / p) == h for p, h in hashes.items()), '只读核实期间文件发生并发变化，请重新准备'
    backups = {}
    for label, p in [('scenes', domain.path), ('master', index.master_path), ('baseline', baseline_path)]:
        dest = WORK / (label + '_before' + p.suffix)
        assert not dest.exists(), '备份已存在，禁止覆盖'
        shutil.copy2(p, dest)
        assert sha(dest) == hashes[relative(p)]
        backups[label] = relative(dest)
    write_json(WORK / 'prepared.json', {'root': str(ROOT), 'asset_id': catalog['asset_id'], 'status': STATUS,
                                        'records': records, 'hashes': hashes, 'backups': backups,
                                        'existing_rows': seen, 'prepared_at': datetime.now().isoformat()})
    print('准备完成：哈希核实、备份与before门禁已保存。', flush=True)


def logical_cells(wb):
    return {(ws.title, c.coordinate): c.value for ws in wb for c in ws._cells.values() if c.value is not None}


def workbook_diff(old_path, new_path, output):
    old, new = load_workbook(old_path), load_workbook(new_path)
    a, b = logical_cells(old), logical_cells(new)
    changes = [{'sheet': s, 'cell': c, 'before': a.get((s,c)), 'after': b.get((s,c))}
               for s, c in sorted(a.keys() | b.keys()) if a.get((s,c)) != b.get((s,c))]
    write_json(output, changes)
    return changes


def changelog(wb, kind, text):
    ws = wb['域变更日志']
    last = max(c.row for c in ws._cells.values() if c.column == 1 and c.value is not None)
    version = re.fullmatch(r'v(\d+)\.(\d+)\.(\d+)', ws.cell(last, 1).value)
    assert version
    a, b, c = map(int, version.groups())
    values = [f'v{a}.{b}.{c+1}', date.today().isoformat(), kind, 'Skyline20 v007', text, NOTE, 'WorkBuddy']
    for col, value in enumerate(values, 1):
        ws.cell(last + 1, col)._style = copy(ws.cell(last, col)._style)
        ws.cell(last + 1, col).value = value
    ws.row_dimensions[last + 1].height = ws.row_dimensions[last].height


def check_hashes(hashes):
    changes = [p for p, h in hashes.items() if sha(ROOT / p) != h]
    assert not changes, '检测到并发变化，停止提交：' + str(changes)


def commit_candidates(candidates, hashes, stage):
    # Windows独占句柄覆盖写：锁内复核SHA，全部锁取得前绝不改动；失败时仅在锁内恢复。
    from ctypes import wintypes
    kernel = ctypes.WinDLL('kernel32', use_last_error=True)
    kernel.CreateFileW.argtypes = [wintypes.LPCWSTR, wintypes.DWORD, wintypes.DWORD, ctypes.c_void_p,
                                   wintypes.DWORD, wintypes.DWORD, wintypes.HANDLE]
    kernel.CreateFileW.restype = wintypes.HANDLE
    kernel.ReadFile.argtypes = [wintypes.HANDLE, ctypes.c_void_p, wintypes.DWORD, ctypes.POINTER(wintypes.DWORD), ctypes.c_void_p]
    kernel.WriteFile.argtypes = kernel.ReadFile.argtypes
    kernel.SetFilePointerEx.argtypes = [wintypes.HANDLE, ctypes.c_longlong, ctypes.c_void_p, wintypes.DWORD]
    kernel.SetEndOfFile.argtypes = [wintypes.HANDLE]
    kernel.FlushFileBuffers.argtypes = [wintypes.HANDLE]
    kernel.CloseHandle.argtypes = [wintypes.HANDLE]
    handles, originals = {}, {}
    def write(handle, content):
        assert kernel.SetFilePointerEx(handle, 0, None, 0)
        written = wintypes.DWORD()
        buffer = ctypes.create_string_buffer(content)
        assert kernel.WriteFile(handle, buffer, len(content), ctypes.byref(written), None)
        assert written.value == len(content)
        assert kernel.SetEndOfFile(handle) and kernel.FlushFileBuffers(handle)
    check_hashes(hashes)
    try:
        for p in candidates:
            handle = kernel.CreateFileW(str(ROOT / p), 0xC0000000, 0, None, 3, 0x80, None)
            if handle == ctypes.c_void_p(-1).value:
                raise OSError(ctypes.get_last_error(), '目标账本被其他程序持有，停止提交', p)
            handles[p] = handle
            parts = []
            while True:
                buffer = ctypes.create_string_buffer(1024 * 1024)
                count = wintypes.DWORD()
                assert kernel.ReadFile(handle, buffer, len(buffer), ctypes.byref(count), None)
                if count.value == 0:
                    break
                parts.append(buffer.raw[:count.value])
            originals[p] = b''.join(parts)
            assert hashlib.sha256(originals[p]).hexdigest() == hashes[p], '锁内哈希不一致'
        check_hashes({p: h for p, h in hashes.items() if p not in handles})
        payloads = {p: Path(candidate).read_bytes() for p, candidate in candidates.items()}
        changed = []
        try:
            for p, data in payloads.items():
                changed.append(p)
                write(handles[p], data)
        except BaseException:
            for p in changed:
                write(handles[p], originals[p])
            raise
    finally:
        for handle in handles.values():
            kernel.CloseHandle(handle)
    for p, candidate in candidates.items():
        assert sha(ROOT / p) == sha(candidate)
        hashes[p] = sha(ROOT / p)
    write_json(WORK / f'{stage}_commit.json', {'committed': list(candidates), 'hashes': hashes})


def refresh_aggregate(baseline, index, scene_candidate):
    rows = []
    for d in index.domains:
        wb = load_workbook(scene_candidate if d.key == 'scenes' else d.path)
        rows.extend(read_source_rows(wb['资产主表']))
    baseline['asset_count'] = len(baseline['assets'])
    baseline['column_digests'] = {str(c): col_digest(rows, c) for c in CONTENT_COLUMNS}
    baseline['category_counts'] = dict(sorted(Counter(str(v[2]) for _, v in rows).items()))


def apply():
    assert not (WORK/'asset_commit.json').exists(), '事务已有提交记录，不得自动重复或覆盖；先检查登记报告'
    prepared = read_json(WORK / 'prepared.json')
    hashes = dict(prepared['hashes'])
    check_hashes(hashes)
    index, domain, catalog, records, asset_hashes = facts()
    baseline_path = ROOT / 'assets/registry/ledger_split_baseline.json'
    baseline_rel = relative(baseline_path)
    stage1 = WORK / 'scenes_asset_candidate.xlsx'
    master1 = WORK / 'master_asset_candidate.xlsx'
    baseline1 = WORK / 'baseline_asset_candidate.json'
    wb = load_workbook(domain.path)
    ws = wb['资产主表']
    old_rows = read_source_rows(ws)
    old_last = max(r for r, _ in old_rows)
    assert old_last == HEADER_ROW + len(old_rows), '现有主表有空行，不自动迁移历史行'
    row_map = {v[0]: r for r, v in old_rows}
    parent = catalog['asset_id']
    added = []
    for x in records:
        aid = x['asset_id']
        row = row_map.get(aid)
        if row is None:
            row = old_last + 1 + len(added)
            added.append(aid)
            for col in range(1, COLUMN_COUNT + 1):
                ws.cell(row, col)._style = copy(ws.cell(old_last, col)._style)
            ws.row_dimensions[row].height = ws.row_dimensions[old_last].height
        x['asset_row'] = row
        values = [aid, x['name'], '场景', 'environment_kit_3d' if aid == parent else 'environment_module_3d',
                  'open_world_landscape_skyline20', x['slug'], None if aid == parent else parent,
                  'Top3D / Godot Y-up', 'default', '开放世界景观；独立组件与本资产装配；未用于正式房间',
                  STATUS, 'P1', 'v007', x['spec'], x['prefab'],
                  relative(SOURCE / 'qa/qa_report.json') + '; ' + relative(SOURCE / 'component_catalog.json'),
                  'SKYLINE20;二十层景观建筑;' + x['slug'], None, None, x['prefab_sha256'], 'WorkBuddy',
                  date.today().isoformat(), '用户提供参考图；项目内原创建模；无第三方模型下载',
                  aid.removeprefix('ENV-OPENWORLD-'),
                  NOTE + ' Godot探针1430 pass/0 fail；source=' + catalog['source_blend'] + '; source_sha256=' + catalog['source_sha256']
                  + '; optimized=' + catalog['optimized_blend'] + '; optimized_sha256=' + catalog['optimized_sha256']
                  + (('; glb=' + x['glb'] + '; glb_sha256=' + asset_hashes[x['glb']]) if x['glb'] else '; root引用8个组件Prefab，无整屋GLB。')]
        assert len(values) == COLUMN_COUNT
        for col, value in enumerate(values, 1):
            ws.cell(row, col).value = value
    last = old_last + len(added)
    rescope_asset_sheet(ws, len(old_rows) + len(added))
    for cell in wb['总览']._cells.values():
        if isinstance(cell.value, str) and cell.value.startswith('=') and '资产主表' in cell.value and cell.row != 13:
            cell.value = re.sub(r'(?<=\$)' + str(old_last) + r'\b', str(last), cell.value)
    changelog(wb, '登记技术验证阶段资产', '父资产+8组件；validated；源/优化源/GLB与9个Prefab哈希实测；美术待审，未正式接入。')
    wb.save(stage1)
    reread = load_workbook(stage1)
    original_wb = load_workbook(domain.path)
    for sheet in domain.sheet_scope:
        assert sheet_digest(reread[sheet]) == sheet_digest(original_wb[sheet]), '主表事务不得改专表'
    authorized_ids = {x['asset_id'] for x in records}
    original_content = {v[0]: _row_digest(v) for _, v in old_rows if v[0] not in authorized_ids}
    candidate_content = {v[0]: _row_digest(v) for _, v in read_source_rows(reread['资产主表']) if v[0] not in authorized_ids}
    assert original_content == candidate_content, '非授权历史资产内容改变'
    with ZipFile(stage1) as z:
        tree = ET.fromstring(z.read('xl/worksheets/sheet3.xml'))
        ns = {'m': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
        actual = {n.attrib['sqref'] for n in tree.findall('.//m:dataValidation', ns)}
        assert actual == {f'C{FIRST_DATA_ROW}:C{last}', f'K{FIRST_DATA_ROW}:K{last}', f'L{FIRST_DATA_ROW}:L{last}'}
        cf_ranges = {n.attrib['sqref'] for n in tree.findall('.//m:conditionalFormatting', ns)}
        assert cf_ranges == {f'S{FIRST_DATA_ROW}:S{last}', f'K{FIRST_DATA_ROW}:K{last}'}
        assert reread['资产主表'].auto_filter.ref == f'A{HEADER_ROW}:X{last}'
    base = read_json(baseline_path)
    prior_baseline = read_json(baseline_path)
    prior_assets = dict(base['assets'])
    ids = {x['asset_id'] for x in records}
    for aid in ids:
        assert aid not in prior_assets or aid in row_map, '基线中已有但账本丢失的ID不能被本事务静默复活'
    for r, v in read_source_rows(reread['资产主表']):
        if v[0] in ids:
            base['assets'][v[0]] = {'v': _row_digest(v), 'c': v[2], 'd': 'scenes'}
    assert all(base['assets'][aid] == v for aid, v in prior_assets.items() if aid not in ids)
    refresh_aggregate(base, index, stage1)
    assert all(base[k] == prior_baseline[k] for k in prior_baseline if k not in ('assets','asset_count','column_digests','category_counts'))
    write_json(WORK/'baseline_asset_diff.json', {
        'authorized_assets': {aid:{'before':prior_assets.get(aid),'after':base['assets'][aid]} for aid in sorted(ids)},
        'asset_count': [prior_baseline['asset_count'],base['asset_count']],
        'column_digests': {k:[prior_baseline['column_digests'].get(k),v] for k,v in base['column_digests'].items() if prior_baseline['column_digests'].get(k)!=v},
        'category_counts': [prior_baseline['category_counts'],base['category_counts']],
        'unrelated_asset_fingerprints_unchanged': True, 'all_sheet_fingerprints_unchanged': True,
    })
    write_json(baseline1, base)
    master = load_workbook(index.master_path)
    for row in master['分账本索引'].iter_rows():
        if row[1].value == domain.name:
            master['分账本索引'].cell(row[0].row, 7).value = len(old_rows) + len(added)
    for row in master['总览'].iter_rows():
        if row[0].value == '场景':
            master['总览'].cell(row[0].row, 4).value = sum(v[2] == '场景' for _, v in read_source_rows(reread['资产主表']))
    master.save(master1)
    diffs1 = workbook_diff(domain.path, stage1, WORK / 'asset_transaction_diff.json')
    for change in diffs1:
        if change['sheet'] == '资产主表':
            coord = change['cell']
            row = int(re.search(r'\d+', coord).group())
            assert row in {x['asset_row'] for x in records} or coord.startswith(('R', 'S'))
        else:
            assert change['sheet'] in ('总览', '域变更日志')
    master_changes1 = workbook_diff(index.master_path, master1, WORK / 'master_asset_diff.json')
    allowed_master1 = set()
    for row in master['分账本索引'].iter_rows():
        if row[1].value == domain.name:
            allowed_master1.add(('分账本索引',f'G{row[0].row}'))
    for row in master['总览'].iter_rows():
        if row[0].value == '场景':
            allowed_master1.add(('总览',f'D{row[0].row}'))
    assert all((c['sheet'],c['cell']) in allowed_master1 for c in master_changes1)
    mapping1 = {str(domain.path.resolve()): str(stage1), str(index.master_path.resolve()): str(master1),
                str(baseline_path.resolve()): str(baseline1)}
    # 基线由Path读取而非load_workbook；候选门禁在runner中额外重定向该唯一JSON路径。
    candidate1 = gates_with_baseline('candidate_asset', mapping1, baseline1)
    diff1 = compare(read_json(WORK / 'before_gates.json'), candidate1)
    write_json(WORK / 'candidate_asset_gate_diff.json', diff1)
    assert_no_new(diff1)
    commit_candidates({domain.relative_path: stage1, index.master_relative_path: master1, baseline_rel: baseline1}, hashes, 'asset')
    middle = gates('after_asset')
    assert_no_new(compare(read_json(WORK / 'before_gates.json'), middle))
    write_json(WORK / 'asset_rows_committed.json', records)
    write_json(SOURCE / 'qa/ledger_registration.json', {
        'asset_id': parent, 'version': 'v007', 'status': STATUS,
        'registration_status': 'ASSET_ROWS_COMMITTED_PREFAB_TRANSACTION_PENDING',
        'reference_art_approved': False, 'runtime_integrated': False, 'note': NOTE,
        'paths': {'ledger': domain.relative_path, 'baseline': baseline_rel, 'master': index.master_relative_path},
        'rows': records, 'exit': {k:v['exit'] for k,v in middle.items()},
        'errors': {k:v['errors'] for k,v in middle.items()}, 'backups': prepared['backups'],
        'diff': compare(read_json(WORK/'before_gates.json'), middle),
    })
    stage2 = WORK / 'scenes_prefab_candidate.xlsx'
    master2 = WORK / 'master_prefab_candidate.xlsx'
    baseline2 = WORK / 'baseline_prefab_candidate.json'
    check_hashes(hashes)
    for label, p in [('scenes', domain.path), ('master', index.master_path), ('baseline', baseline_path)]:
        shutil.copy2(p, WORK / (label + '_before_prefab' + p.suffix))
    check_hashes(hashes)
    wb = load_workbook(domain.path)
    page = wb[PAGE]
    assert [page.cell(4,c).value for c in range(1,17)] == ['AssetID','中文名','Prefab路径','GLB模型路径','Blender源文件','功能说明','功能脚本路径','碰撞开关','碰撞归属','碰撞方式','标准尺寸','原点与朝向','使用位置','制作状态','版本','备注']
    previous = read_json(baseline_path)
    existing_page_digests = {sheet: sheet_digest(wb[sheet]) for sheet in domain.sheet_scope}
    assert previous['sheet_digests'][PAGE] == sheet_digest(page)
    pmap = {page.cell(r,1).value: r for r in range(5,page.max_row+1) if page.cell(r,1).value}
    end = max(pmap.values(), default=4)
    for x in records:
        row = pmap.get(x['asset_id'])
        if row is None:
            end += 1
            row = end
            for col in range(1,17):
                page.cell(row,col)._style = copy(page.cell(5,col)._style)
            page.row_dimensions[row].height = page.row_dimensions[5].height
        x['prefab_row'] = row
        values = [x['asset_id'], x['name'], x['prefab'], x['glb'], catalog['source_blend'], x['spec'], None,
                  '无','无','无',x['size'],'bottom_center；Godot Y-up/-Z；Blender Z-up/-Y',
                  '开放世界景观候选；未接入正式房间',STATUS,'v007',NOTE + ' prefab_sha256=' + x['prefab_sha256']]
        for col,value in enumerate(values,1):
            page.cell(row,col).value=value
    # 只修正授权Prefab页统计；其他页历史双等号/错误汇总保留为历史问题。
    for row in wb['总览'].iter_rows():
        if row[0].value == PAGE:
            for col,letter in [(2,'A'),(3,'D'),(4,'G'),(5,'H')]:
                wb['总览'].cell(row[0].row,col).value = f"=COUNTA('{PAGE}'!$A$5:$A${end})" if col==2 else f'=COUNTIF(\'{PAGE}\'!${letter}$5:${letter}${end},"' + ('开' if col==5 else '<>') + '")'
    changelog(wb,'独立登记PackedScene专表','9个实际存在Prefab（root+8组件）；专表独立事务；validated；无碰撞/功能脚本；美术待审、未正式接入。')
    wb.save(stage2)
    after_wb = load_workbook(stage2)
    for sheet, digest in existing_page_digests.items():
        if sheet != PAGE:
            assert sheet_digest(after_wb[sheet]) == digest, 'Prefab事务不得改其他专表'
    diffs2 = workbook_diff(domain.path,stage2,WORK/'prefab_transaction_diff.json')
    authorized_rows = {x['prefab_row'] for x in records}
    for change in diffs2:
        if change['sheet'] == PAGE:
            assert int(re.search(r'\d+',change['cell']).group()) in authorized_rows
        else:
            assert change['sheet'] in ('总览','域变更日志')
    assert read_source_rows(after_wb['资产主表']) == read_source_rows(load_workbook(domain.path)['资产主表'])
    base = read_json(baseline_path)
    base['sheet_digests'][PAGE] = sheet_digest(after_wb[PAGE])
    assert base['assets'] == previous['assets']
    assert all(base[k] == previous[k] for k in previous if k != 'sheet_digests')
    assert all(base['sheet_digests'][k] == v for k,v in previous['sheet_digests'].items() if k != PAGE)
    write_json(WORK/'baseline_prefab_diff.json', {'sheet':PAGE,'before':previous['sheet_digests'][PAGE],
               'after':base['sheet_digests'][PAGE],'all_asset_fingerprints_unchanged':True,'other_sheet_fingerprints_unchanged':True})
    write_json(baseline2,base)
    master = load_workbook(index.master_path)
    page_rows = [r for r in range(5,end+1) if page.cell(r,1).value]
    snapshot = [len(page_rows), sum(bool(page.cell(r,4).value) for r in page_rows),
                sum(bool(page.cell(r,7).value) for r in page_rows),sum(page.cell(r,8).value=='开' for r in page_rows)]
    for row in master['3D Prefab总控'].iter_rows():
        if row[1].value == PAGE:
            for col,value in zip((3,4,5,6),snapshot):
                master['3D Prefab总控'].cell(row[0].row,col).value=value
    master.save(master2)
    master_changes2 = workbook_diff(index.master_path,master2,WORK/'master_prefab_diff.json')
    allowed_master2 = set()
    for row in master['3D Prefab总控'].iter_rows():
        if row[1].value == PAGE:
            allowed_master2.update(('3D Prefab总控',f'{letter}{row[0].row}') for letter in ('C','D','E','F'))
    assert all((c['sheet'],c['cell']) in allowed_master2 for c in master_changes2)
    mapping2 = {str(domain.path.resolve()):str(stage2),str(index.master_path.resolve()):str(master2),str(baseline_path.resolve()):str(baseline2)}
    candidate2=gates_with_baseline('candidate_prefab',mapping2,baseline2)
    diff2=compare(middle,candidate2)
    write_json(WORK/'candidate_prefab_gate_diff.json',diff2)
    assert_no_new(diff2)
    commit_candidates({domain.relative_path:stage2,index.master_relative_path:master2,baseline_rel:baseline2},hashes,'prefab')
    after=gates('after')
    final_diff=compare(read_json(WORK/'before_gates.json'),after)
    assert_no_new(final_diff)
    check_hashes(hashes)
    report={'asset_id':parent,'version':'v007','status':STATUS,'registration_status':'REGISTERED_TECHNICAL_VALIDATION_ART_PENDING',
            'reference_art_approved':False,'runtime_integrated':False,'note':NOTE,
            'paths':{'ledger':domain.relative_path,'prefab_sheet':PAGE,'master':index.master_relative_path,'baseline':baseline_rel,
                     'source':catalog['source_blend'],'optimized':catalog['optimized_blend'],'script':relative(Path(__file__))},
            'rows':records,'exit':{k:v['exit'] for k,v in after.items()},'errors':{k:v['errors'] for k,v in after.items()},
            'before':read_json(WORK/'before_gates.json'),'after':after,'diff':final_diff,
            'backups':{**prepared['backups'],'prefab':{label:relative(WORK/(label+'_before_prefab'+p.suffix)) for label,p in [('scenes',domain.path),('master',index.master_path),('baseline',baseline_path)]}},
            'frozen_hashes':prepared['hashes'],'after_hashes':hashes,'asset_hashes':asset_hashes,
            'logical_diff':{'asset_cells':len(diffs1),'prefab_cells':len(diffs2),'master_asset_cells':len(master_changes1),'master_prefab_cells':len(master_changes2)},
            'preservation':{'historical_asset_fingerprints_unchanged':True,'unauthorized_sheet_fingerprints_unchanged':True,
                            'historical_asset_content_unchanged_except_derived_formulas':True,'source_glb_prefab_geometry_unchanged':True},
            'remaining':'美术参考图审批未完成；runtime_integrated=false；全项目门禁既有红项原样记录，不宣称全绿。',
            'historical_overview_debt': {
                'scenes': {c.coordinate:c.value for c in load_workbook(domain.path)['总览']._cells.values()
                           if isinstance(c.value,str) and c.value.startswith('==')},
                'master_prefab_total': [load_workbook(index.master_path)['3D Prefab总控'].cell(14,c).value for c in range(3,7)],
                'note':'历史双等号及非本资产汇总债务不在本次资产指纹更新范围；授权Prefab页统计已覆盖新行。',
            }}
    report['baseline_change_reports'] = [relative(WORK/'baseline_asset_diff.json'),relative(WORK/'baseline_prefab_diff.json')]
    report['workbook_change_reports'] = [relative(WORK/p) for p in ('asset_transaction_diff.json','prefab_transaction_diff.json','master_asset_diff.json','master_prefab_diff.json')]
    report['script_sha256'] = sha(Path(__file__))
    write_json(SOURCE/'qa/ledger_registration.json',report)
    print('登记完成；追溯报告已保存。',flush=True)


def gates_with_baseline(tag,mapping,baseline):
    global RUNNER
    original=RUNNER
    RUNNER=RUNNER.replace('original=openpyxl.load_workbook',
        'original_read=pathlib.Path.read_text\n'
        'def read(self,*a,**kw):\n'
        '    key=str(self.resolve())\n'
        '    target=mapping.get(key)\n'
        '    return original_read(pathlib.Path(target) if target else self,*a,**kw)\n'
        'pathlib.Path.read_text=read\noriginal=openpyxl.load_workbook')
    try:
        return gates(tag,mapping)
    finally:
        RUNNER=original


if __name__=='__main__':
    parser=argparse.ArgumentParser(description='Skyline20 v007真实阶段登记，独立主表与Prefab事务。')
    parser.add_argument('mode',choices=['prepare','apply'])
    args=parser.parse_args()
    try:
        prepare() if args.mode=='prepare' else apply()
    except BaseException as exc:
        write_json(WORK/'registration_error.json',{'status':'BLOCKED','error':str(exc),'traceback':traceback.format_exc()})
        raise
