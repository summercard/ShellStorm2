#!/usr/bin/env python3
"""把 8 个弹种专属弹体特效「制成」结果回填进特效账本，并登记《3D-特效》分页。

一次性、幂等拒绝重跑的登记脚本（前置断言必须全部成立，否则直接退出 1）。

改动面（仅 ShellStorm2_特效账本_v001.xlsx；主目录跨域契约不变，故不动）
--------------------------------------------------------------------
《资产主表》 r27..r34（8 条弹种专属特效行）：
  * K 列 制作状态  待制作 → 已完成
  * N 列 规格      待定 → 实际程序几何参数（逐弹种）
  * O 列 文件路径  待定 → assets/art/vfx/combat_3d/vfx_bullet_<type>_root_top3d.tscn
  * T 列 SHA-256   None → 该 Prefab 的 sha256 十六进制
  * U 列 负责人   待分配 → 摩斯拉
  * V 列 更新时间  → 2026-10-13 00:00:00
  * W 列 来源/许可 待制作/原创程序几何 → 原创程序几何
  * Y 列 备注      重写为「已制成」说明（原备注写的是制成后应做的动作）
  * r26 的 Y 列修掉一处悬空占位符「§…」→「§1.3/§3.5」（顺手修正，属观察到的笔误）

《3D-特效》分页 追加 1 个组头 + 8 条数据行：
  * 组头 r25「▼ 01 战斗反馈（续）· 子弹特效（FX01-06 系列）」
  * r26..r33 编号 FX01-08..FX01-15（承接 FX01-06 通用弹体这一「系列」的序号）
  * 每行登记实际存在的 PackedScene（该页铁律：一行 = 一个实际存在的 Godot PackedScene）
  * r2 说明行把「共 12 条」更新为实测条数，并补一条 2026-10-13 进度

《域变更日志》 追加 v0.1.15。

设计依据：
  内容库《武器》S 列（每弹种战斗表现原文）、04 §6.2 子弹模块必填 visual_id、
  14.6 §1.3（《资产主表》为唯一登记源）/§1.1（独立 Prefab）/§2.2（combat_3d 目录）。
"""

from __future__ import annotations

import hashlib
import sys
from copy import copy
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path("/Users/summercards/Documents/ShellStorm2")
sys.path.insert(0, str(ROOT / "scripts"))
sys.path.insert(0, str(ROOT / "tools" / "asset_pipeline"))

from ledger_registry import LedgerIndex  # noqa: E402
from split_asset_ledger import (  # noqa: E402
    ASSET_SHEET,
    FIRST_DATA_ROW,
    HEADER_ROW,
    _row_digest,
    dedupe_key_formula,
    dedupe_result_formula,
    sheet_digest,
)

VFX_PATH = ROOT / "assets/registry/ledgers/ShellStorm2_特效账本_v001.xlsx"
VFX_PAGE = "3D-特效"
VFX_DIR = "assets/art/vfx/combat_3d"
DATE = "2026-10-13"
TS = f"{DATE} 00:00:00"
ROW_HEIGHT = 15.0

# r27..r34 的 8 条（顺序与《资产主表》一致）
FIRST_BULLET_ROW = 27
LAST_BULLET_ROW = 34
NEW_LAST = LAST_BULLET_ROW  # 资产主表末数据行不变（只改内容，不增行）
PAGE_LAST_DATA_ROW = 24     # 《3D-特效》改动前末数据行
FX_FIRST = 8                # FX01-08 起

GROUP_HEADER = "▼ 01 战斗反馈（续）· 子弹特效（FX01-06 系列）"
PARENT_ID = "VFX-BULLET-VISUAL-3D"

# (slug, PascalCase, 中文名, module_cn, 内容库《武器》S 行号, S 列设计原文, 功能说明, 标准尺寸)
BULLETS = [
    ("standard", "Standard", "标准子弹弹体与尾迹特效", "标准子弹模块", 17,
     "使用清晰中性弹体、尾迹和命中；不附加误导性的元素效果",
     "中性基准弹体：细胶囊弹芯 + 柔和单锥尾迹，无任何元素附加物，不误导观感。",
     "弹芯胶囊 r=0.045/h=0.20；外壳 r=0.075/h=0.18；单锥尾迹 底0.055/长0.46；灯 energy1.15/range1.75m"),
    ("sticky", "Sticky", "黏黏弹附着与倒计时特效", "黏黏弹模块", 18,
     "命中后弹体留在表面并显示倒计时/状态，结束时反馈与结算同步",
     "短粗黏稠弹体 + 滴落小珠尾迹（drip 动画），强调「附着」质感。",
     "弹芯胶囊 r=0.050/h=0.16；外壳 r=0.086/h=0.16；锥尾迹 底0.062/长0.34 + 3 滴落珠；灯 energy1.20/range1.80m"),
    ("bounce", "Bounce", "回旋镖弹反弹折线与计数特效", "回旋镖弹模块", 13,
     "反弹点有折线闪光与计数递减，飞行方向变化清楚可读",
     "圆润弹体 + 三段左右错位的折线尾迹，方向变化一眼可读。",
     "弹芯胶囊 r=0.048/h=0.14；外壳 r=0.072/h=0.13；三段折线锥 底0.052/0.040/0.030；灯 energy1.20/range1.80m"),
    ("piercing", "Piercing", "穿甲弹贯穿线与层数特效", "穿甲弹模块", 16,
     "穿透命中显示连续贯穿线与剩余层数，不在首个目标处假性消失",
     "细长针形弹体 + 细长双段贯穿尾迹，笔直的贯穿感。",
     "弹芯胶囊 r=0.028/h=0.30；外壳 r=0.044/h=0.28；两段针形尾迹 底0.036/0.022；灯 energy1.35/range1.95m"),
    ("explosive", "Explosive", "爆炸弹弹体与爆炸时点特效", "爆炸弹模块", 14,
     "弹体、尾迹、爆炸半径与伤害时点同步，低特效下保留轮廓圈",
     "粗壮弹体 + 带飞散火星的锥尾迹（flicker 动画），体积感与弹重匹配。",
     "弹芯胶囊 r=0.055/h=0.22；外壳 r=0.094/h=0.20；锥尾迹 底0.078/长0.40 + 3 火星条；灯 energy1.65/range2.40m"),
    ("homing", "Homing", "蜂卵弹锁定框与弯曲尾迹特效", "蜂卵弹模块", 15,
     "锁定框、弯曲尾迹和转向动作一致；目标丢失有解除提示",
     "椭圆蜂卵弹体 + 向两侧摆动的弯曲尾迹（wave 动画），读得出「在追」。",
     "弹芯胶囊 r=0.050/h=0.18；外壳 r=0.082/h=0.16；三段摆动尾迹；灯 energy1.25/range1.90m"),
    ("blackhole", "Blackhole", "黑洞弹核心与吸入流线特效", "黑洞弹模块", 12,
     "黑洞核心、吸入流线和范围边界清晰；命中/结束有收束反馈",
     "暗核球形 + 紫光环 + 外扩喇叭尾迹（spin 旋转），读得出「在吸」。",
     "弹芯球 r=0.058；外壳球 r=0.098；外扩喇叭三段 底0.040/0.068/0.096；灯 energy1.45/range2.20m"),
    ("balloon", "Balloon", "气球弹膨胀弹体与命中体积特效", "气球弹模块", 11,
     "膨胀弹体、较重尾迹与命中体积一致，不能只改伤害数值",
     "大球体弹体 + 粗短厚重尾迹，并支持 apply_growth 继续膨胀。",
     "弹芯球 r=0.070；外壳球 r=0.118；粗短锥尾迹 底0.104/长0.36；灯 energy1.30/range2.05m"),
]

ORIGIN_AXIS = ("原点=宿主原点（local position 保持 0，随 Projectile3D 变换）；"
               "朝向=弹体沿 local -Z，拖尾向 local +Z（身后）")


def require(cond: bool, msg: str) -> None:
    if not cond:
        raise SystemExit(f"前置断言失败: {msg}")


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def read_rows(ws) -> list[tuple[int, tuple]]:
    rows = []
    for row in range(FIRST_DATA_ROW, ws.max_row + 1):
        if str(ws.cell(row=row, column=1).value or "").strip():
            rows.append((row, tuple(ws.cell(row=row, column=c).value for c in range(1, 26))))
    return rows


def page_data_rows(ws) -> list[int]:
    return [
        r for r in range(4, ws.max_row + 1)
        if str(ws.cell(row=r, column=1).value or "").startswith("FX")
    ]


def main() -> int:
    index = LedgerIndex.load(ROOT)
    domain = index.domain_for_key("vfx")
    require(domain.categories == ("特效", "子弹特效"), f"vfx 大类不是两项: {domain.categories}")

    # 逐弹种 Prefab / 脚本存在性 + 哈希
    facts = []
    for slug, pascal, cn, module_cn, src_row, s_text, g_text, spec in BULLETS:
        prefab = ROOT / VFX_DIR / f"vfx_bullet_{slug}_root_top3d.tscn"
        script = ROOT / "src/vfx" / f"VfxBullet{pascal}3D.gd"
        require(prefab.is_file(), f"缺少 Prefab: {prefab}")
        require(script.is_file(), f"缺少脚本: {script}")
        text = prefab.read_text(encoding="utf-8")
        require(f'metadata/asset_id = "VFX-BULLET-{slug.upper()}-3D"' in text,
                f"{prefab.name} 的 asset_id 不匹配")
        require('metadata/asset_version = "v001"' in text, f"{prefab.name} 的 asset_version 不是 v001")
        facts.append((slug, pascal, cn, module_cn, src_row, s_text, g_text, spec,
                      f"{VFX_DIR}/{prefab.name}", sha256(prefab)))

    wb = load_workbook(VFX_PATH)

    # ------------------------------------------------------------ 前置断言
    require(wb.sheetnames == ["总览", "账本说明", ASSET_SHEET, VFX_PAGE, "域变更日志"],
            f"页签不符: {wb.sheetnames}")
    main_ws = wb[ASSET_SHEET]
    rows_before = read_rows(main_ws)
    require([r for r, _ in rows_before] == list(range(FIRST_DATA_ROW, NEW_LAST + 1)),
            "资产主表数据行不连续或末行不是 34")
    for (r, v), fact in zip(rows_before[-8:], facts):
        require(str(v[0]) == f"VFX-BULLET-{fact[0].upper()}-3D", f"r{r} AssetID 不符: {v[0]}")
        require(str(v[10]).strip() == "待制作", f"r{r} K 列不是待制作: {v[10]!r}")
        require(str(v[14]).strip() == "待定", f"r{r} O 列不是待定: {v[14]!r}")
        require(v[19] in (None, ""), f"r{r} T 列本应为空: {v[19]!r}")

    page_ws = wb[VFX_PAGE]
    page_rows_before = page_data_rows(page_ws)
    require(page_rows_before[-1] == PAGE_LAST_DATA_ROW,
            f"《3D-特效》末数据行不是 {PAGE_LAST_DATA_ROW}: {page_rows_before[-1]}")
    require("FX01-08" not in {str(page_ws.cell(row=r, column=1).value) for r in page_rows_before},
            "FX01-08 已存在，拒绝重跑")
    page_before = sheet_digest(page_ws)
    note_before = str(page_ws.cell(row=2, column=1).value or "")

    # ------------------------------------------------------- 1. 资产主表回填
    style_done = copy(main_ws.cell(row=26, column=1)._style)   # r26 是「已完成」行，形态一致
    for (r, v), fact in zip(rows_before[-8:], facts):
        slug, pascal, cn, module_cn, src_row, s_text, g_text, spec, path, sha = fact
        main_ws.cell(row=r, column=11).value = "已完成"          # K 制作状态
        main_ws.cell(row=r, column=14).value = spec              # N 规格
        main_ws.cell(row=r, column=15).value = path              # O 文件路径
        main_ws.cell(row=r, column=20).value = sha               # T SHA-256
        main_ws.cell(row=r, column=21).value = "摩斯拉"           # U 负责人
        main_ws.cell(row=r, column=22).value = TS                # V 更新时间
        main_ws.cell(row=r, column=23).value = "原创程序几何"      # W 来源/许可
        main_ws.cell(row=r, column=25).value = (
            f"已制成并实装：Prefab={path} + 脚本=src/vfx/VfxBullet{pascal}3D.gd；"
            f"根节点=VfxBullet{pascal}3D(Node3D)；《3D-特效》分页编号见该页（FX01-06 系列）。"
            f"变体父体 {PARENT_ID}（复用其弹芯/外壳/拖尾几何分层，仅按弹种换外观与行为段）。"
            "调用方 Projectile3D._build_visual() 经 BULLET_VISUAL_REGISTRY 按 visual_id 路由，"
            "作为子节点持有、生命周期随 ProjectilePool3D，刻意不进 VfxPool3D；纯视觉无碰撞体。"
            "验收：verify_bullet_vfx_profiles（8 profile 可区分 + AssetID 路由 + 尺寸线性 + "
            "端到端开火 + 无碰撞 + 反向对照）。"
            f"设计依据：内容库《武器》S{src_row}「{s_text}」。"
        )
        for col in range(1, 26):
            if main_ws.cell(row=r, column=col)._style is None:
                main_ws.cell(row=r, column=col)._style = copy(style_done)
        main_ws.row_dimensions[r].height = main_ws.row_dimensions[26].height or ROW_HEIGHT

    # r26 备注：修掉悬空占位符「§…」
    y26 = str(main_ws.cell(row=26, column=25).value or "")
    if "§…" in y26:
        main_ws.cell(row=26, column=25).value = y26.replace("§…", "§1.3/§3.5")

    # ------------------------------------------------------- 2. 3D-特效 分页
    hdr_style = [copy(page_ws.cell(row=7, column=c)._style) for c in range(1, 18)]     # 组头
    data_style = [copy(page_ws.cell(row=13, column=c)._style) for c in range(1, 18)]   # 数据行

    head_r = PAGE_LAST_DATA_ROW + 1
    for c in range(1, 18):
        cell = page_ws.cell(row=head_r, column=c)
        cell.value = GROUP_HEADER if c == 1 else None
        cell._style = copy(hdr_style[c - 1])
    page_ws.row_dimensions[head_r].height = page_ws.row_dimensions[7].height or 66.0

    for offset, fact in enumerate(facts):
        slug, pascal, cn, module_cn, src_row, s_text, g_text, spec, path, sha = fact
        r = head_r + 1 + offset
        vals = [None] * 17
        vals[0] = f"FX01-{FX_FIRST + offset:02d}"
        vals[1] = f"VFX-BULLET-{slug.upper()}-3D"
        vals[2] = cn
        vals[3] = path
        vals[4] = "无（程序几何）"
        vals[5] = "无（程序几何）"
        vals[6] = g_text
        vals[7] = f"src/vfx/VfxBullet{pascal}3D.gd"
        vals[8] = "无"
        vals[9] = "无"
        vals[10] = "无"
        vals[11] = spec
        vals[12] = ORIGIN_AXIS
        vals[13] = f"玩家/敌人 {module_cn}（mod_bullet_{slug}）"
        vals[14] = "已完成"
        vals[15] = "v001"
        vals[16] = (
            f"已实装；Prefab={path} + 脚本=src/vfx/VfxBullet{pascal}3D.gd；"
            f"根节点=VfxBullet{pascal}3D(Node3D)；变体父体 {PARENT_ID}；"
            "调用方 Projectile3D._build_visual()（BULLET_VISUAL_REGISTRY 路由）；"
            "无碰撞体；不进 VfxPool3D；验收：verify_bullet_vfx_profiles。"
        )
        for c in range(1, 18):
            cell = page_ws.cell(row=r, column=c)
            cell.value = vals[c - 1]
            cell._style = copy(data_style[c - 1])
        page_ws.row_dimensions[r].height = page_ws.row_dimensions[13].height or ROW_HEIGHT

    # 说明行：把「共 12 条」更新为实测条数 + 补 2026-10-13 进度
    total_entries = len(page_data_rows(page_ws))
    note = note_before
    if "共 12 条" in note:
        note = note.replace("共 12 条", f"共 {total_entries} 条")
    note = note.rstrip() + (
        f"。 【{DATE} 进度】新增 8 个逐弹种专属弹体 Prefab（FX01-08..FX01-15，"
        "VFX-BULLET-{STANDARD,STICKY,BOUNCE,PIERCING,EXPLOSIVE,HOMING,BLACKHOLE,BALLOON}-3D），"
        "为 FX01-06 通用弹体的弹种专属变体：共享基类 VfxBulletBody3D（弹芯/外壳/尾迹/跟随灯分层），"
        "由 Projectile3D 依 visual_id 路由、随 ProjectilePool3D 回收，不进 VfxPool；"
        "专项验收 verify_bullet_vfx_profiles。"
    )
    page_ws.cell(row=2, column=1).value = note

    # ------------------------------------------------------- 3. 域变更日志
    log = wb["域变更日志"]
    log_row = log.max_row + 1
    while log_row > 6 and str(log.cell(row=log_row - 1, column=1).value or "").strip() == "":
        log_row -= 1
    log_style = copy(log.cell(row=log_row - 1, column=1)._style)
    require(str(log.cell(row=log_row - 1, column=1).value) == "v0.1.14",
            f"域变更日志最后一条不是 v0.1.14: {log.cell(row=log_row - 1, column=1).value!r}")
    entry = [
        "v0.1.15", DATE, "条目更新（待制作 → 已完成）+ 新增 Prefab", "特效 · 子弹特效",
        "把 8 个弹种专属弹体特效（VFX-BULLET-{STANDARD,STICKY,BOUNCE,PIERCING,EXPLOSIVE,"
        "HOMING,BLACKHOLE,BALLOON}-3D）从「待制作」推进到「已完成」："
        "《3D-特效》分页登记 8 个实际存在的 PackedScene（FX01-08..FX01-15，"
        "为 FX01-06 通用弹体 VFX-BULLET-VISUAL-3D 的弹种专属变体）；"
        "《资产主表》逐行回填 O 列路径、T 列 SHA-256，K 列推进「已完成」并同步 N 规格/U 负责人/V 更新时间/W 来源。"
        "实现：共享基类 src/vfx/VfxBulletBody3D.gd 提供弹芯（胶囊/球）/外壳（弹色半透明）/尾迹（锥段/小珠/火星条）/跟随灯四层，"
        "8 个子类只覆写 _bullet_profile() 声明弹体形态、配色、尾迹构图与灯参数；"
        "由 Projectile3D 依 visual_id 经 BULLET_VISUAL_REGISTRY 路由为子节点持有，随 ProjectilePool3D 回收，"
        "刻意不进 VfxPool3D（避免弹体池与特效池双重 ownership）。专项验收 verify_bullet_vfx_profiles。",
        "稳定 AssetID 全部保持不变（r27..r34 的 AssetID/子类/逻辑ID 逐格不变）；"
        "新增 8 个 Prefab 路径与哈希；O 列由「待定」变为去版本化运行路径（无 _vNNN），版本仍锁在 Prefab 根 meta asset_version=v001 与账本 M 列；"
        "纯视觉无碰撞体、不含玩法行为，运行时调用方式（AssetID/visual_id 路由）为向后兼容增强。",
        "CodeBuddy",
    ]
    for i, value in enumerate(entry, start=1):
        cell = log.cell(row=log_row, column=i)
        cell.value = value
        cell._style = copy(log_style)
    log.row_dimensions[log_row].height = 46

    # ------------------------------------------------- 回读自检（内存，通过才落盘）
    cws = wb[ASSET_SHEET]
    crows = {r: v for r, v in read_rows(cws)}
    expected_end = FIRST_DATA_ROW - 1 + len(crows)
    for r in crows:
        ls = str(cws.cell(row=r, column=19).value)
        require(ls == dedupe_result_formula(r, expected_end), f"S{r} 形状错: {ls}")
        require(str(cws.cell(row=r, column=18).value) == dedupe_key_formula(r), f"R{r} 形状错")
    for (r, _v), fact in zip(rows_before[-8:], facts):
        require(str(cws.cell(row=r, column=11).value) == "已完成", f"回读 r{r} K 错")
        require(str(cws.cell(row=r, column=15).value) == fact[8], f"回读 r{r} O 错")
        require(str(cws.cell(row=r, column=20).value) == fact[9], f"回读 r{r} T 错")
    cpage = wb[VFX_PAGE]
    crow_list = page_data_rows(cpage)
    require(len(crow_list) == len(page_rows_before) + 8, "回读 3D-特效 行数不对")
    require(str(cpage.cell(row=PAGE_LAST_DATA_ROW + 1, column=1).value) == GROUP_HEADER, "组头缺失")
    for offset, fact in enumerate(facts):
        r = PAGE_LAST_DATA_ROW + 2 + offset
        require(str(cpage.cell(row=r, column=2).value) == f"VFX-BULLET-{fact[0].upper()}-3D",
                f"回读 3D-特效 r{r} AssetID 错")
    clast = wb["域变更日志"]
    require(str(clast.cell(row=clast.max_row, column=1).value) == "v0.1.15", "回读日志末行错")

    # --------------------------------------------------------------- 落盘
    wb.save(VFX_PATH)

    # 落盘后跨进程回读
    final = load_workbook(VFX_PATH)
    fws = final[ASSET_SHEET]
    frows = {r: v for r, v in read_rows(fws)}
    require(set(frows) == set(crows), "落盘后资产主表行集合不一致")
    for r in frows:
        require(str(fws.cell(row=r, column=20).value) == str(cws.cell(row=r, column=20).value),
                f"落盘后 r{r} T 列不一致")
    fpage = final[VFX_PAGE]
    require(str(fpage.cell(row=PAGE_LAST_DATA_ROW + 1, column=1).value) == GROUP_HEADER, "落盘后组头丢失")
    require(len(page_data_rows(fpage)) == len(page_rows_before) + 8, "落盘后 3D-特效 行数不对")
    flog = final["域变更日志"]
    require(str(flog.cell(row=flog.max_row, column=1).value) == "v0.1.15", "落盘后日志末行错")

    print("APPLY_OK main_rows", len(frows), "page_rows", len(page_data_rows(fpage)),
          "page_last", page_data_rows(fpage)[-1], "log_last", flog.max_row)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
