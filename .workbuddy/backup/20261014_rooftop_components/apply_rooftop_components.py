# -*- coding: utf-8 -*-
"""在《场景账本》资产主表新增 9 条 100F 天台独立组件（设计阶段）。

纪律（见 skill `shellstorm2-asset-ledger-row-authoring`）：
- 只追加到末行，绝不 insert_rows（会破坏公式引用）
- R/S 派生列用权威函数 dedupe_key_formula / dedupe_result_formula 逐行重写
- 总览写死区间 $<旧末行> → $<新末行>
- DataValidation sqref 扩到新末行
- 域变更日志追加一行
- 其余（无损基线）由独立步骤补录

用法：python add_rooftop_components.py [--apply]
不带 --apply 只做 dry-run 校验。
"""
from __future__ import annotations

import argparse
import re
import sys
from copy import copy
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "tools" / "asset_pipeline"))
sys.path.insert(0, str(ROOT / "scripts"))

from split_asset_ledger import (  # noqa: E402
    ASSET_SHEET,
    COLUMN_COUNT,
    FIRST_DATA_ROW,
    HEADER_ROW,
    dedupe_key_formula,
    dedupe_result_formula,
)

LEDGER = ROOT / "assets/registry/ledgers/ShellStorm2_场景账本_v001.xlsx"
TEMPLATE_ROW = 345          # ENV-OPENWORLD-TOWER03-MAIN-STRUCTURE（场景独立组件行，同形态）
EXPECTED_OLD_LAST = 928     # 改前末行
EXPECTED_S6 = dedupe_result_formula(6, EXPECTED_OLD_LAST)
UPDATED_AT_SERIAL = 46309   # 2026-10-14

COMMON = {
    "C": "场景",
    "D": "environment_module_3d",
    "E": "rooftop_decor",
    "G": None,
    "H": "Blender Z-up / -Y",
    "I": "default",
    "J": "100F天台专用；场景独立组件（仅登记，未制作）",
    "K": "design_only",
    "L": "P1",
    "M": "v001",
    "O": None,
    "P": "docs/v0.1/design/rooftop_component_library.md",
    "T": None,
    "U": "Codex",
    "V": UPDATED_AT_SERIAL,
    "W": "用户设计",
    "X": "—",
    "Y": "2026-10-14 按用户要求登记为 100F 天台场景独立组件（设计阶段，未制作）：本次仅入账本，"
         "不建 Blender 源 / GLB / Prefab，不含碰撞与玩法语义。",
}

NEW_ROWS = [
    ("ENV-ROOFTOP-FLOWERBED-BROKEN", "天台破损花圃", "broken_flowerbed",
     "约6.0×1.5m（长6m、宽1.5m）；破损花圃：残破花池+塌落土石，落地件，底面Z=0",
     "天台破损花圃;破损花池;花圃;3D;Prefab;场景独立组件;broken_flowerbed;天台;100F"),
    ("ENV-ROOFTOP-SANDBAG-BARRIER", "天台防御沙包墙", "sandbag_barrier",
     "约7.0×2.0m（长7m、宽2m）；防御沙包掩体：堆叠沙包，落地件，底面Z=0",
     "天台防御沙包;沙包墙;沙包掩体;3D;Prefab;场景独立组件;sandbag_barrier;天台;100F"),
    ("ENV-ROOFTOP-RUBBLE-PILE-A", "天台石块堆 A", "rubble_pile_a",
     "高约1.0~1.3m；石块堆3变种之A，落地件，底面Z=0",
     "天台石块堆;石块堆A;瓦砾堆;3D;Prefab;场景独立组件;rubble_pile_a;天台;100F"),
    ("ENV-ROOFTOP-RUBBLE-PILE-B", "天台石块堆 B", "rubble_pile_b",
     "高约1.0~1.3m；石块堆3变种之B，落地件，底面Z=0",
     "天台石块堆;石块堆B;瓦砾堆;3D;Prefab;场景独立组件;rubble_pile_b;天台;100F"),
    ("ENV-ROOFTOP-RUBBLE-PILE-C", "天台石块堆 C", "rubble_pile_c",
     "高约1.0~1.3m；石块堆3变种之C，落地件，底面Z=0",
     "天台石块堆;石块堆C;瓦砾堆;3D;Prefab;场景独立组件;rubble_pile_c;天台;100F"),
    ("ENV-ROOFTOP-DEBRIS-PILE-LARGE", "天台大型杂物堆", "debris_pile_large",
     "约12.0×10.0m（12m×10m）；大型杂物堆，落地件，底面Z=0",
     "天台大型杂物堆;杂物堆;废墟堆;3D;Prefab;场景独立组件;debris_pile_large;天台;100F"),
    ("ENV-ROOFTOP-BROKEN-CRATE-A", "天台破败木箱 A", "broken_crate_a",
     "破败木箱3套变种之A；落地件，底面Z=0",
     "天台破败木箱;破木箱A;木箱;3D;Prefab;场景独立组件;broken_crate_a;天台;100F"),
    ("ENV-ROOFTOP-BROKEN-CRATE-B", "天台破败木箱 B", "broken_crate_b",
     "破败木箱3套变种之B；落地件，底面Z=0",
     "天台破败木箱;破木箱B;木箱;3D;Prefab;场景独立组件;broken_crate_b;天台;100F"),
    ("ENV-ROOFTOP-BROKEN-CRATE-C", "天台破败木箱 C", "broken_crate_c",
     "破败木箱3套变种之C；落地件，底面Z=0",
     "天台破败木箱;破木箱C;木箱;3D;Prefab;场景独立组件;broken_crate_c;天台;100F"),
]

COL_MAP = {"A": 1, "B": 2, "C": 3, "D": 4, "E": 5, "F": 6, "G": 7, "H": 8, "I": 9, "J": 10,
           "K": 11, "L": 12, "M": 13, "N": 14, "O": 15, "P": 16, "Q": 17, "T": 20, "U": 21,
           "V": 22, "W": 23, "X": 24, "Y": 25}


def build_values(index: int) -> list:
    asset_id, name_zh, slug, spec, keywords = NEW_ROWS[index]
    vals: list = [None] * COLUMN_COUNT
    for key, val in COMMON.items():
        vals[COL_MAP[key] - 1] = val
    vals[0] = asset_id
    vals[1] = name_zh
    vals[5] = slug          # F 组件槽
    vals[13] = spec         # N 规格
    vals[16] = keywords     # Q 关键词
    return vals


def check_preconditions(wb) -> None:
    assert ASSET_SHEET in wb.sheetnames, "缺少《资产主表》"
    ws = wb[ASSET_SHEET]
    assert ws.cell(row=HEADER_ROW, column=1).value == "AssetID", "表头行不是 5"
    assert ws.cell(row=HEADER_ROW, column=25).value == "备注", "第 25 列不是「备注」"
    last = EXPECTED_OLD_LAST
    assert ws.cell(row=last, column=1).value == "ENV-EXPEDITION-L99-BOSS-ROOM-02", \
        f"末行 {last} 不是预期资产：{ws.cell(row=last, column=1).value}"
    assert not ws.cell(row=last + 1, column=1).value, f"{last+1} 行非空，末行假设失效"
    actual_s6 = ws.cell(row=6, column=19).value
    assert actual_s6 == EXPECTED_S6, f"S6 公式形态不符：{actual_s6!r}"
    assert ws.cell(row=6, column=18).value == dedupe_key_formula(6), "R6 公式形态不符"
    existing = {ws.cell(row=r, column=1).value for r in range(FIRST_DATA_ROW, last + 1)}
    for asset_id, *_ in NEW_ROWS:
        assert asset_id not in existing, f"AssetID 已存在：{asset_id}"
    ov = wb["总览"]
    for ref in ("A6", "C6", "E6", "G6", "B10", "C10", "B11", "C11", "B12", "C12", "B13", "C13"):
        assert f"${last}" in str(ov[ref].value or ""), f"总览 {ref} 区间不含 ${last}"


def patch_overview(ov, old_last: int, new_last: int) -> int:
    n = 0
    for row in ov.iter_rows():
        for cell in row:
            v = cell.value
            if isinstance(v, str) and "$" + str(old_last) in v:
                cell.value = re.sub(r"\$%d\b" % old_last, f"${new_last}", v)
                n += 1
    return n


def patch_validations(ws, new_last: int) -> list:
    rebuilt = []
    old = list(ws.data_validations.dataValidation)
    ws.data_validations.dataValidation = []
    for dv in old:
        ranges = []
        for rng in dv.sqref.ranges:
            text = str(rng)
            m = re.fullmatch(r"([A-Z]+)(\d+):([A-Z]+)(\d+)", text)
            if m and int(m.group(4)) < new_last:
                text = f"{m.group(1)}{m.group(2)}:{m.group(3)}{new_last}"
            ranges.append(text)
        ndv = DataValidation(
            type=dv.type, formula1=dv.formula1, formula2=dv.formula2,
            allow_blank=dv.allow_blank, showDropDown=dv.showDropDown,
            showErrorMessage=dv.showErrorMessage, showInputMessage=dv.showInputMessage,
            errorStyle=dv.errorStyle, operator=dv.operator, imeMode=dv.imeMode,
        )
        for text in ranges:
            ndv.add(text)
        ws.add_data_validation(ndv)
        rebuilt.append((dv.type, ranges))
    return rebuilt


def append_change_log(wb, new_last: int, added: list[str]) -> int:
    log = wb["域变更日志"]
    last = max(r for r in range(1, log.max_row + 1) if log.cell(row=r, column=1).value)
    text = (
        "新增 100F 天台场景独立组件 9 条（ENV-ROOFTOP-FLOWERBED-BROKEN 破损花圃 6×1.5m / "
        "ENV-ROOFTOP-SANDBAG-BARRIER 防御沙包 7×2m / ENV-ROOFTOP-RUBBLE-PILE-A·B·C 石块堆 3 变种 高1.0~1.3m / "
        "ENV-ROOFTOP-DEBRIS-PILE-LARGE 大型杂物堆 12×10m / ENV-ROOFTOP-BROKEN-CRATE-A·B·C 破败木箱 3 套变种）。"
        "本次只登记设计意图与尺寸，制作状态=design_only，不建 Blender 源 / GLB / Prefab，无碰撞与玩法语义。"
    )
    compat = ("AssetID 全部为新增，不改既有行；不动 3D-场景通用 / 3D-设施 / 3D-其他 三张 Prefab 分页"
              "（无 Prefab 不加行）；无损基线同步补录 9 个指纹并重算列摘要与分类计数。")
    values = [f"v0.1.36", "2026-10-14", "新增条目", "关卡场景 / 100F天台", text, compat, "Codex"]
    new_row = last + 1
    for col in range(1, 8):
        cell = log.cell(row=new_row, column=col)
        cell._style = copy(log.cell(row=last, column=col)._style)
        cell.value = values[col - 1]
    log.row_dimensions[new_row].height = log.row_dimensions[last].height
    return new_row


def main() -> int:
    ap = argparse.ArgumentParser()
    ap.add_argument("--apply", action="store_true")
    args = ap.parse_args()

    wb = load_workbook(LEDGER)
    check_preconditions(wb)
    ws = wb[ASSET_SHEET]
    ov = wb["总览"]

    old_last = EXPECTED_OLD_LAST
    new_last = old_last + len(NEW_ROWS)

    for i in range(len(NEW_ROWS)):
        r = old_last + 1 + i
        values = build_values(i)
        for c in range(1, COLUMN_COUNT + 1):
            cell = ws.cell(row=r, column=c)
            cell._style = copy(ws.cell(row=TEMPLATE_ROW, column=c)._style)
            cell.value = values[c - 1]
        # 不加显式行高：与最近的追加（927/928）一致，交给 Excel 自动行高

    # R/S 逐行重写（含既有行：区间必须覆盖新末行）
    rewritten = 0
    for r in range(FIRST_DATA_ROW, new_last + 1):
        if not ws.cell(row=r, column=1).value:
            continue
        ws.cell(row=r, column=18).value = dedupe_key_formula(r)
        ws.cell(row=r, column=19).value = dedupe_result_formula(r, new_last)
        rewritten += 1

    patched = patch_overview(ov, old_last, new_last)
    dvs = patch_validations(ws, new_last)
    log_row = append_change_log(wb, new_last, [x[0] for x in NEW_ROWS])

    print(f"新增行 {old_last+1}..{new_last}（{len(NEW_ROWS)} 条）")
    print(f"R/S 重写 {rewritten} 行；总览替换 {patched} 格；DV {dvs}")
    print(f"域变更日志新行 {log_row}")
    ws.auto_filter.ref = f"A{HEADER_ROW}:Y{new_last}"
    print(f"autoFilter → {ws.auto_filter.ref}")

    if not args.apply:
        print("[dry-run] 未写盘")
        return 0
    wb.save(LEDGER)
    print(f"[applied] 已写入 {LEDGER}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
