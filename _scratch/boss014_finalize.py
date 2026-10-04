from pathlib import Path
import json,hashlib,sys,shutil
from collections import Counter
from openpyxl import load_workbook
R=Path.cwd();B=R/'assets/art/enemies/bosses/enm_boss_monitor002';sys.path[:0]=[str(R/'scripts'),str(R/'tools/asset_pipeline')]
from ledger_registry import LedgerIndex
from split_asset_ledger import read_source_rows,_row_digest,CONTENT_COLUMNS,col_digest
ID='ENM-BOSS-MONITOR002-3D';ix=LedgerIndex.load(R);p=ix.path_for_category('敌人');rows=read_source_rows(load_workbook(p)['资产主表']);v=next(v for r,v in rows if v[0]==ID)
bp=R/'assets/registry/ledger_split_baseline.json';shutil.copy2(bp,R/'_scratch/boss014_ledger/baseline_before.json');bl=json.loads(bp.read_text(encoding='utf-8'));bl['assets'][ID]={'v':_row_digest(v),'c':'敌人','d':'enemies'}
rows=[]
for d in ix.domains:rows.extend(read_source_rows(load_workbook(R/'assets/registry/ledgers'/d.file)['资产主表']))
bl['column_digests']={str(c):col_digest(rows,c) for c in CONTENT_COLUMNS};bl['category_counts']=dict(Counter(v[2] for r,v in rows));bp.write_text(json.dumps(bl,ensure_ascii=False,indent=1)+'\n',encoding='utf-8')
entry='''## 当前交付 v014（身前平拍与手绘赛博特效）

打开 `source/enm_boss_monitor002_animation_v014.blend`，默认 BOSS002_STUDIO 场景，播放1—55帧。模型母版 `source/enm_boss_monitor002_model_v014.blend`。v012/v013保留回退。

键盘宽面水平向下拍在身前约3.2米处；举高停顿、快速下砸、低位停留和收回保留。修正接触阶段四指握姿；五官以原中心放大22%。idle/move动作曲线与v012完全一致。

image-2制作透明图集 `textures/impact_v014/impact_atlas.png`：手绘白色尖角爆点、贴地冲击环、蓝色挥击笔触、玫红电子碎片。挥击28—34帧点缀，命中33帧触发，44帧前全部消失；独立 `BOSS002_IMPACT_PREVIEW` 集合，角色源场景不含特效。角色仍18,684三角面、64骨、3材质，预览另有3个特效材质。

动作23项与表现5项检查通过，四分之一帧检查无键盘/左手穿地；证据 `previews/keyboard_v014/audit.json`、`presentation_audit.json`。预览 `melee_keyboard.mp4`。仅源级交付，未导出GLB、未接入Godot/伤害事件。

'''
p=B/'README.md';t=p.read_text(encoding='utf-8').replace('## 当前交付 v012','## 历史交付 v012');i=t.index('\n\n')+2;p.write_text(t[:i]+entry+t[i:],encoding='utf-8')
p=R/'docs/v0.1/design/Boss002显示器动画设计.md';t=p.read_text(encoding='utf-8').replace('设计修订r7','设计修订r8').replace('v012动作母版','v014动作母版').replace('宽面平行地面向下拍击','宽面平行地面向下拍击，接触中心位于身前约3.2米').replace('地面冲击环、放射碎屑与短促尘雾','2D手绘尖角爆点、贴地冲击环；挥击和命中短促点缀蓝色/玫红电子干扰');p.write_text(t,encoding='utf-8')
p=R/'docs/v0.1/MODULE_INDEX.md';t=p.read_text(encoding='utf-8').replace('v012已制作idle、move与键盘攻击melee_keyboard、五官浮动及代码UV上升','v014已制作idle、move与身前平拍melee_keyboard，五官放大22%、2D手绘赛博特效预览及代码UV上升');p.write_text(t,encoding='utf-8')
entry='''## 2026-10-05 v014 身前平拍与2D手绘赛博特效

BOSS-STAGES / ASSET-PIPELINE，设计r8。延续用户前一对话v012，v013探索平拍，最终v014按标注将落点移到身前3.2米。宽面水平，攻击握指修正，五官放大22%。image-2生成透明图集，独立预览集合提供手绘爆点/贴地环/蓝色挥击/玫红干扰；首尾无残留。默认打开BOSS002_STUDIO直接播放。

Blender factory-startup、python-exit-code 1：动作23项与表现5项通过，217个四分之一帧采样无键盘/左手穿地；idle/move曲线与v012相同，双母版同骨架。真实Cycles整段预览、关键帧见资产previews/keyboard_v014。已更新同一AssetID及独立制作账本；逐格确认其他资产保持。未执行GLB导出、Godot运行/伤害接入、全工程回归。旧v012/v013保留可回滚。

'''
for rel,e in [('docs/v0.1/development/2026-10-03_boss002_monitor_source.md',entry),('docs/v0.1/development/CHANGELOG.md','## 2026-10-05｜Boss002前方平拍v014\n\n- BOSS-STAGES / ASSET-PIPELINE：键盘身前平拍、五官放大与2D手绘赛博特效，源级28项验收；见[制作记录](2026-10-03_boss002_monitor_source.md)。\n\n')]:
 p=R/rel;t=p.read_text(encoding='utf-8');i=t.index('\n\n')+2;p.write_text(t[:i]+e+t[i:],encoding='utf-8')
print('Ledger baseline and current documentation updated')
