import openpyxl
p = "assets/registry/ledgers/ShellStorm2_场景账本_v001.xlsx"
wb = openpyxl.load_workbook(p, data_only=False)
ws = wb["3D-场景通用"]
hdr = [ws.cell(4,c).value for c in range(1, ws.max_column+1)]
print("HEADER:", hdr)
def show(r):
    print("---------- 3D-场景通用 R%d ----------" % r)
    for c in range(1, ws.max_column+1):
        v = ws.cell(r,c).value
        print("   %-14s = %s" % (hdr[c-1], v))
for r in (102, 104):
    show(r)
# 找 prp_tower_wall_parapet_5m / floor_tile 行做风格参照
for r in range(5, ws.max_row+1):
    v = ws.cell(r,3).value
    if isinstance(v,str) and ("parapet_5m.tscn" in v or "tower_floor_tile_5m.tscn" in v or "tower_wall_solid_5m.tscn" in v):
        show(r)
