from pathlib import Path
import json
import openpyxl

ROOT = Path('I:/工作项目/shellstrom2/ShellStorm2')
BOOKS = ROOT / 'assets/registry/ledgers'
TARGETS = {
    'PRP-BASE-WORKSHOP-STOOL-3D',
    'PRP-BASE-MISSION-COMMAND-CHAIR-3D',
    'PRP-BASE99-TELESCOPIC-LADDER-3D',
}
result = {}
for name in ['ShellStorm2_道具账本_v001.xlsx', 'ShellStorm2_场景账本_v001.xlsx']:
    wb = openpyxl.load_workbook(BOOKS / name, data_only=False)
    book = {'sheets': wb.sheetnames, 'asset_rows': [], 'prefab_sheets': {}}
    if '资产主表' in wb.sheetnames:
        ws = wb['资产主表']
        for row in range(6, ws.max_row + 1):
            asset_id = ws.cell(row, 1).value
            if asset_id in TARGETS:
                book['asset_rows'].append({'sheet': '资产主表', 'row': row, 'values': [ws.cell(row, col).value for col in range(1, 26)]})
    for sheet_name in ['3D-道具', '3D-物品', '3D-场景通用', '3D-设施', '3D-其他']:
        if sheet_name not in wb.sheetnames:
            continue
        ws = wb[sheet_name]
        rows = []
        for row in range(1, ws.max_row + 1):
            vals = [ws.cell(row, col).value for col in range(1, min(ws.max_column, 25) + 1)]
            text = ' | '.join('' if value is None else str(value) for value in vals)
            if any(token in text for token in ['椅', '梯', 'STOOL', 'CHAIR', 'LADDER', '收音机', 'RADIO']):
                rows.append({'row': row, 'values': vals})
        book['prefab_sheets'][sheet_name] = {'max_row': ws.max_row, 'max_column': ws.max_column, 'matches': rows}
    result[name] = book
(ROOT / '_scratch/register_base99_radio_ledger.json').write_text(json.dumps(result, ensure_ascii=False, indent=2, default=str) + '\n', encoding='utf-8')
