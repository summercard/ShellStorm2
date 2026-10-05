from __future__ import annotations

import hashlib
import json
import struct
import subprocess
import sys
from pathlib import Path
from datetime import datetime, timezone

import openpyxl

ROOT = Path('I:/工作项目/shellstrom2/ShellStorm2')
OUT = ROOT / 'outputs/base99_radio_v002'
GLB = ROOT / 'assets/art/props/base_world_3d/components/base99_radio/prp_base99_radio_visual_top3d.glb'
SOURCE = ROOT / 'assets/art/props/base_world_3d/source/base99_radio/prp_base99_radio_source_v002.blend'
PREFAB = ROOT / 'assets/art/props/base_world_3d/runtime/base99_radio/prp_base99_radio_root_top3d.tscn'
ROLLBACK = ROOT / 'outputs/base99_radio_v001/rollback/prp_base99_radio_visual_top3d_v001.glb'
PROP_LEDGER = ROOT / 'assets/registry/ledgers/ShellStorm2_道具账本_v001.xlsx'
MASTER_LEDGER = ROOT / 'assets/registry/ShellStorm2_美术资产台账_v001.xlsx'
BASELINE = ROOT / 'assets/registry/ledger_split_baseline.json'
AID = 'PRP-BASE99-RADIO-3D'


def sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def text(value) -> str:
    return '' if value is None else str(value).strip()


def read_glb_json(path: Path) -> dict:
    data = path.read_bytes()
    magic, version, total = struct.unpack_from('<4sII', data, 0)
    assert magic == b'glTF' and version == 2 and total == len(data)
    offset = 12
    while offset < total:
        length, chunk_type = struct.unpack_from('<II', data, offset)
        chunk = data[offset + 8:offset + 8 + length]
        if chunk_type == 0x4E4F534A:
            return json.loads(chunk.rstrip(b' ').decode('utf-8'))
        offset += 8 + length
    raise AssertionError('GLB JSON chunk missing')


def row_digest(values) -> str:
    return hashlib.sha256('\x1f'.join(text(values[c - 1]) for c in range(1, 26) if c not in (18, 19)).encode('utf-8')).hexdigest()


def sheet_digest(ws) -> str:
    parts = []
    for merged in sorted(str(rng) for rng in ws.merged_cells.ranges):
        parts.append(f'merge:{merged}')
    cells = [
        (cell.row, cell.column, cell.coordinate, cell.value)
        for row in ws.iter_rows()
        for cell in row
        if cell.value not in (None, '')
    ]
    if cells:
        parts.append(
            f'bbox={min(c[0] for c in cells)}:{max(c[0] for c in cells)}'
            f'x{min(c[1] for c in cells)}:{max(c[1] for c in cells)}'
        )
    else:
        parts.append('bbox=empty')
    for _row, _col, coordinate, value in sorted(cells, key=lambda item: (item[0], item[1])):
        parts.append(f'{coordinate}={value!r}')
    return hashlib.sha256('\n'.join(parts).encode('utf-8')).hexdigest()


def run(command: list[str], output_name: str) -> dict:
    result = subprocess.run(command, cwd=ROOT, text=True, encoding='utf-8', errors='replace', capture_output=True)
    (OUT / output_name).write_text(result.stdout + result.stderr, encoding='utf-8')
    return {'command': command, 'returncode': result.returncode, 'log': str(OUT / output_name)}


for path in [GLB, SOURCE, PREFAB, ROLLBACK, PROP_LEDGER, MASTER_LEDGER, BASELINE]:
    assert path.is_file(), path

manifest = json.loads((OUT / 'asset_manifest.json').read_text(encoding='utf-8'))
validator = json.loads((OUT / 'validate_game_prop.json').read_text(encoding='utf-8'))
assert validator['passed'] is True
assert validator['polygon_count'] == 598

prefab_text = PREFAB.read_text(encoding='utf-8')
assert 'metadata/asset_id = "PRP-BASE99-RADIO-3D"' in prefab_text
assert 'metadata/asset_version = "v002"' in prefab_text
assert 'metadata/model_faces = 598' in prefab_text
assert 'metadata/model_triangles = 1116' in prefab_text
assert 'path="res://assets/art/props/base_world_3d/components/base99_radio/prp_base99_radio_visual_top3d.glb"' in prefab_text

wb = openpyxl.load_workbook(PROP_LEDGER, data_only=False)
asset_sheet = wb['资产主表']
prefab_sheet = wb['3D-道具']
change_sheet = wb['域变更日志']
asset_rows = [r for r in range(6, asset_sheet.max_row + 1) if text(asset_sheet.cell(r, 1).value) == AID]
assert asset_rows == [27]
assert text(asset_sheet.cell(27, 13).value) == 'v002'
assert '598面 / 1116三角形' in text(asset_sheet.cell(27, 14).value)
assert text(prefab_sheet.cell(8, 1).value) == AID
assert text(prefab_sheet.cell(8, 16).value) == 'v002'
assert '565a505c9eacc53ef5a8b9b14c5577b64e4226c423ebfffce39d4965a7839f7d' in text(prefab_sheet.cell(8, 17).value)
assert text(change_sheet.cell(12, 1).value) == 'v0.1.6'

baseline = json.loads(BASELINE.read_text(encoding='utf-8'))
old_sheet_digest = baseline['sheet_digests'].get('3D-道具')
current_sheet_digest = sheet_digest(prefab_sheet)
baseline['sheet_digests']['3D-道具'] = current_sheet_digest
BASELINE.write_text(json.dumps(baseline, ensure_ascii=False, indent=1) + '\n', encoding='utf-8')
wb.close()

sha_values = {
    'glb_sha256': sha(GLB),
    'source_blend_sha256': sha(SOURCE),
    'runtime_prefab_sha256': sha(PREFAB),
    'rollback_v001_glb_sha256': sha(ROLLBACK),
}
assert sha_values['glb_sha256'] == manifest['hashes']['glb_sha256']
assert sha_values['source_blend_sha256'] == manifest['hashes']['source_blend_sha256']
assert sha_values['runtime_prefab_sha256'] == manifest['hashes']['runtime_prefab_sha256']
assert sha_values['glb_sha256'] == '565a505c9eacc53ef5a8b9b14c5577b64e4226c423ebfffce39d4965a7839f7d'

report = {
    'asset_id': AID,
    'asset_name': '99F阁楼收音机',
    'version': 'v002',
    'status': '通过',
    'generated_at': datetime.now(timezone.utc).isoformat(),
    'geometry': {
        'visual_faces': 566,
        'visual_triangles': 1068,
        'status_light_faces': 32,
        'status_light_triangles': 48,
        'total_faces': 598,
        'total_triangles': 1116,
        'face_limit': 800,
        'preferred_face_limit': 700,
    },
    'bounds_m': {
        'min': [-0.207, -0.114, 0.0],
        'max': [0.207, 0.114, 0.411],
        'size_blender_xyz': [0.414, 0.228, 0.411],
        'size_godot_xyz': [0.414, 0.411, 0.228],
        'base_plane': 'Z=0 Blender / Y=0 Godot',
    },
    'runtime_contract': {
        'prefab': 'assets/art/props/base_world_3d/runtime/base99_radio/prp_base99_radio_root_top3d.tscn',
        'root': 'ItemRoot',
        'visual': 'ItemRoot/Visual',
        'status_light': 'ItemRoot/StatusLight',
        'script_unchanged': 'src/base3d/Base99Radio3D.gd',
        'music_bus': 'Music',
        'state_cycle': ['off', 'a', 'b', 'off'],
        'formal_placement_unchanged': True,
    },
    'materials': [
        '01_精工金属_紫色骨架',
        '02_细腻哑光_青绿大面',
        '03_清漆反光_紫粉点缀',
        '04_柔和自发光_UI灯光',
    ],
    'palette': {
        'path': 'assets/art/shared/palette/设施低亮多巴胺色盘_10x10_512.png',
        'uv_layer': 'PaletteUV',
        'interpolation': 'Closest',
        'external_only': True,
        'glb_images': len(read_glb_json(GLB).get('images', [])),
        'glb_textures': len(read_glb_json(GLB).get('textures', [])),
    },
    'paths': {
        'source_blend': 'assets/art/props/base_world_3d/source/base99_radio/prp_base99_radio_source_v002.blend',
        'component_glb': 'assets/art/props/base_world_3d/components/base99_radio/prp_base99_radio_visual_top3d.glb',
        'runtime_prefab': 'assets/art/props/base_world_3d/runtime/base99_radio/prp_base99_radio_root_top3d.tscn',
        'rollback_v001_glb': 'outputs/base99_radio_v001/rollback/prp_base99_radio_visual_top3d_v001.glb',
        'preview_closeup': 'outputs/base99_radio_v002/base99_radio_v002_closeup.png',
        'preview_threequarter': 'outputs/base99_radio_v002/base99_radio_v002_threequarter.png',
        'preview_top': 'outputs/base99_radio_v002/base99_radio_v002_top.png',
    },
    'hashes': sha_values,
    'strict_blender_validation': {
        'report': 'outputs/base99_radio_v002/validate_game_prop.json',
        'passed': True,
        'polygon_count': 598,
        'valid_island_polygon_count': 598,
    },
    'ledger': {
        'file': 'assets/registry/ledgers/ShellStorm2_道具账本_v001.xlsx',
        'asset_sheet_row': 27,
        'prefab_sheet': '3D-道具',
        'prefab_sheet_row': 8,
        'change_log_sheet': '域变更日志',
        'change_log_row': 12,
        'baseline_asset_count': baseline['asset_count'],
        'baseline_sheet_digest_updated_for_intentional_prefab_sheet_edit': True,
        'previous_3d_props_sheet_digest': old_sheet_digest,
        'current_3d_props_sheet_digest': current_sheet_digest,
    },
}
(OUT / 'final_acceptance.json').write_text(json.dumps(report, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')

registration = {
    'asset_id': AID,
    'version': 'v002',
    'registration_status': '正式美术已接入',
    'ledger': report['ledger'],
    'changed_scope': [
        '资产主表第27行：版本、规格、Prefab路径、源/组件路径、SHA、变更日志',
        '3D-道具第8行：版本列P8、路径、规格、接口与GLB SHA备注',
        '域变更日志第12行：v0.1.6',
        '无损基线：收音机行指纹与3D-道具专表摘要同步',
    ],
    'unchanged_scope': [
        'v001 Blender源文件',
        'runtime Prefab稳定路径',
        'ItemRoot/Visual/StatusLight运行时节点接口',
        'src/base3d/Base99Radio3D.gd交互与音乐逻辑',
        '正式场景摆位',
    ],
    'backup': 'outputs/base99_radio_v002/backup_before_v002_registration',
    'hashes': sha_values,
}
(OUT / 'registration_evidence.json').write_text(json.dumps(registration, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')

print(json.dumps({'asset_id': AID, 'final_acceptance': str(OUT / 'final_acceptance.json'), 'registration_evidence': str(OUT / 'registration_evidence.json'), 'baseline_sheet_digest_updated': True, 'hashes': sha_values}, ensure_ascii=False, indent=2))
