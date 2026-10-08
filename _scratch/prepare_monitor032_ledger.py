from pathlib import Path
import json,sys,hashlib,shutil
from openpyxl import load_workbook
R=Path.cwd();D=R/'_scratch/monitor032_ledger';D.mkdir(exist_ok=True);B=R/'assets/art/enemies/bosses/enm_boss_monitor002'
sys.path.insert(0,str(R/'scripts'))
from ledger_registry import LedgerIndex
p=LedgerIndex.load(R).path_for_category('敌人');w=load_workbook(p);A='ENM-BOSS-MONITOR002-3D';assert w['资产主表']['A20'].value==A
dual='; '.join((B/f'source/enm_boss_monitor002_{k}_v032.blend').relative_to(R).as_posix() for k in ['model','animation'])
prefab=B/'runtime/enm_boss_monitor002/enm_boss_monitor002_root_top3d.tscn'
note='v032：仅move改为底座平放左右挪动0.46m，机身左右扭转约41°，双臂柔性跟随。1.6s/48帧，其他15剪辑与四技能时序保留。源97采样、Godot509项、48帧真实渲染通过；previews/move_v032/monitor_walk.mp4。远征01首波1只/3波清房/钥匙通行/ STANDARD撤离规则保留。'
lr=w['域变更日志'].max_row+1
jobs=[{'path':str(p),'sheets':{'资产主表':{'M20':'v032','P20':dual+'；SKEL-MONITOR002-005','T20':hashlib.sha256(prefab.read_bytes()).hexdigest(),'V20':'2026-10-08','Y20':note},'域变更日志':{f'{col}{lr}':v for col,v in zip('ABCDEFG',['v032','2026-10-08','Codex',A,note,'docs/v0.1/development/2026-10-08_boss002_grounded_move.md','不改模型几何、玩法或其余动作'])}}}]
states={}
for row in w['敌人动画与状态']:
 if row[0].value==A:states[f'E{row[0].row}']=dual;states[f'I{row[0].row}']=note
jobs.append({'path':str(p),'sheets':{'3D-敌人':{'E10':dual,'O10':'v032','P10':note},'敌人动画与状态':states}})
pp=B/'source/boss002_production_ledger.xlsx';sn=load_workbook(pp).sheetnames[0]
jobs.append({'path':str(pp),'sheets':{sn:{'B4':'v032 贴地横移与机身扭转；正式16剪辑/十二态/四技能/三阶段','B20':'v032','C20':dual,'C21':'source/rig_contract_v032.json','B26':'previews/move_v032/source_audit.json; previews/runtime/flow_report.json','C26':note}}})
(D/'changes.json').write_text(json.dumps(jobs,ensure_ascii=False,indent=2),encoding='utf-8')
shutil.copy2(R/'assets/registry/ledger_split_baseline.json',D/'baseline_before.json')
for name in ['edit.mjs','merge.py']:
 s=(R/'_scratch/monitor_spawn_ledger_20261007'/name).read_text(encoding='utf-8').replace('_scratch/monitor_spawn_ledger_20261007','_scratch/monitor032_ledger')
 (D/name).write_text(s,encoding='utf-8')
print('Prepared v032 existing-row edits, domain log row',lr)
