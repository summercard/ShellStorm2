from pathlib import Path
import json,sys,hashlib,shutil
from openpyxl import load_workbook
R=Path.cwd();D=R/'_scratch/monitor034_ledger';D.mkdir(exist_ok=True);B=R/'assets/art/enemies/bosses/enm_boss_monitor002'
sys.path.insert(0,str(R/'scripts'));from ledger_registry import LedgerIndex
p=LedgerIndex.load(R).path_for_category('敌人');w=load_workbook(p);A='ENM-BOSS-MONITOR002-3D';assert w['资产主表']['A20'].value==A
dual='; '.join((B/f'source/enm_boss_monitor002_{k}_v034.blend').relative_to(R).as_posix() for k in ['model','animation'])
prefab=B/'runtime/enm_boss_monitor002/enm_boss_monitor002_root_top3d.tscn'
note='v034新增activate出场6.4s/192帧：黑屏、代码上刷、断线火花、两次挣扎、弹簧双臂伸出、抓键盘、五官弹出。原16动作曲线和采样数据保留；静态键盘握持挂点校正。首次激活alert承载，休眠续播、保存进度、死亡优先；四技能与三阶段数值不变。'
lr=w['域变更日志'].max_row+1;states={}
for row in w['敌人动画与状态']:
 if row[0].value==A:
  n=row[0].row;states[f'E{n}']=dual;states[f'I{n}']=note
  if row[1].value=='dormant':states[f'D{n}']='首次activate第0帧；已完成则idle'
  if row[1].value=='alert':states[f'D{n}']='首次activate 6.4s；完成后idle'
jobs=[{'path':str(p),'sheets':{'资产主表':{'M20':'v034','P20':dual+'；SKEL-MONITOR002-005','T20':hashlib.sha256(prefab.read_bytes()).hexdigest(),'V20':'2026-10-08','Y20':note},'3D-敌人':{'E10':dual,'F10':'17剪辑；activate6.4s；四技能；三阶段；双母版同64骨架','O10':'v034','P10':note},'敌人动画与状态':states,'域变更日志':{f'{col}{lr}':v for col,v in zip('ABCDEFG',['v034','2026-10-08','Codex',A,note,'docs/v0.1/development/2026-10-08_boss002_activation.md','新增出场；保留v031动作，修正键盘静态挂点'])}}}]
pp=B/'source/boss002_production_ledger.xlsx';sn=load_workbook(pp).sheetnames[0]
jobs.append({'path':str(pp),'sheets':{sn:{'B4':'v034 激活出场；17剪辑/十二态/四技能/三阶段','B20':'v034','C20':dual,'C21':'source/rig_contract_v034.json','B26':'previews/activate_v034/source_audit.json; previews/runtime/flow_report.json','C26':note}}})
cp=R/'docs/v0.1/data/ShellStorm2_游戏内容数据库_v010.xlsx';cw=load_workbook(cp);assert cw['怪物与Boss']['A17'].value=='boss_monitor002'
jobs.append({'path':str(cp),'sheets':{'怪物与Boss':{'L17':'用户激活出场要求；技能设计r3','O17':'17 Blender剪辑；首次activate6.4s（非伤害技能），四技能/十二态；伤害8%韧性击晕4.8s；通电暴击可打断；独立资产ENM-BOSS-MONITOR002-3D'},'Boss002技能设计':{'B21':'未激活：activate第0帧；已出场：idle','B24':'首次activate6.4s；之后idle','B34':'首轮出场进度/完成标记随实例保存；休眠暂停，重进续播；旧档默认已出场。危险技能仍取消。','B36':'docs/v0.1/design/Boss002显示器技能设计.md（r3）；动画设计r28'}}})
(D/'changes.json').write_text(json.dumps(jobs,ensure_ascii=False,indent=2),encoding='utf-8')
shutil.copy2(R/'assets/registry/ledger_split_baseline.json',D/'baseline_before.json')
for name in ['edit.mjs','merge.py']:
 s=(R/'_scratch/monitor033_ledger'/name).read_text(encoding='utf-8').replace('_scratch/monitor033_ledger','_scratch/monitor034_ledger')
 if name=='edit.mjs':s=s.replace("if(i===3)","if(i===2)").replace("range:'A1:H18'","range:'A20:H37'")
 (D/name).write_text(s,encoding='utf-8')
print('Prepared three targeted workbook patches')
