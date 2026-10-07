from pathlib import Path
import sys,json,shutil
from openpyxl import load_workbook
R=Path.cwd();D=R/'_scratch/monitor_spawn_ledger_20261007';D.mkdir(exist_ok=True)
sys.path.insert(0,str(R/'scripts'))
from ledger_registry import LedgerIndex
p=LedgerIndex.load(R).path_for_category('敌人');w=load_workbook(p)
assert w['资产主表']['A20'].value=='ENM-BOSS-MONITOR002-3D'
assert w['3D-敌人']['A10'].value=='ENM-BOSS-MONITOR002-3D'
assert w['域变更日志'].max_row==51
shutil.copy2(p,D/'original.xlsx');shutil.copy2(R/'assets/registry/ledger_split_baseline.json',D/'baseline_before.json')
placement='expedition_01 / floor_00 / boss（源房ID f00_boss）'
summary='远征01正式Boss房投放：box_boss_arena首波生成1只，房间共3波。5200HP、4技能、3阶段。全清后按原钥匙/门策略前往STANDARD撤离；重进不重刷，不发塔楼下行权限。2026-10-07房间流程25项通过；技能专项507项通过。见2026-10-07_expedition01_monitor_boss_spawn.md。'
jobs=[{'path':str(p),'sheets':{'资产主表':{'J20':placement,'V20':'2026-10-07','Y20':w['资产主表']['Y20'].value+'\n投放验收：2026-10-07，box_boss_arena首波1只、3波清房、正式钥匙开门、STANDARD撤离、重进不重刷；25项通过。'},'域变更日志':{'A52':'v031','B52':'2026-10-07','C52':'Codex','D52':'ENM-BOSS-MONITOR002-3D','E52':'补录v031正式接入及远征01房间投放验收。'+summary,'F52':'docs/v0.1/development/2026-10-07_expedition01_monitor_boss_spawn.md','G52':'原资产ID、v031、active保持；未重导模型。'}}},{'path':str(p),'sheets':{'3D-敌人':{'M10':placement+'；box_boss_arena首波1只','P10':summary+'\n原v031：16剪辑/64骨/12态；真实渲染10帧通过。'}}}]
(D/'changes.json').write_text(json.dumps(jobs,ensure_ascii=False,indent=2),encoding='utf-8')
for name in ['edit.mjs','merge.py']:
 s=(R/'_scratch/monitor_ledger'/name).read_text(encoding='utf-8').replace("R/'_scratch/monitor_ledger'","R/'_scratch/monitor_spawn_ledger_20261007'")
 (D/name).write_text(s,encoding='utf-8')
print('Prepared two scoped ledger transactions',p)
