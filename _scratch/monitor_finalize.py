import json,hashlib,sys,subprocess,datetime
from pathlib import Path
from openpyxl import load_workbook
R=Path.cwd();B=R/'assets/art/enemies/bosses/enm_boss_monitor002';D=R/'_scratch/monitor_ledger'
def read(p):return json.loads(p.read_text(encoding='utf-8'))
def save(p,v):p.write_text(json.dumps(v,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def rel(p):return p.relative_to(R).as_posix()
contract=read(B/'source/rig_contract_v031.json')
contract.update(runtime_integration=True,verification='previews/runtime/flow_report.json',preview_vfx_note='Blender preview collections excluded from GLB; runtime filled cyber VFX owned by MonitorBossVfx.',runtime_prefab=rel(B/'runtime/enm_boss_monitor002/enm_boss_monitor002_root_top3d.tscn'))
contract['new_clips']['dead']={'frames':[1,61],'fps':30,'loop':False}
contract['formal_animations_authored'].append('dead') if 'dead' not in contract['formal_animations_authored'] else None
contract['export_note']='Pure visual GLB plus evaluated 30Hz motion JSON; AnimationPlayer indexes 16 clips, preserving constraint stretch and shear. Enemy3D owns collision, state, damage and death.'
save(B/'source/rig_contract_v031.json',contract)
transfer=read(B/'enm_boss_monitor002_transfer_ledger_v031.json')
files=list((B/'components/enm_boss_monitor002').glob('*'));files=[p for p in files if p.suffix not in ['.import','.uid']]
files += list((B/'runtime/enm_boss_monitor002').glob('*.tscn'))
files += [R/'src/enemy3d'/n for n in ['MonitorBossCombat.gd','MonitorBossPresentation.gd','MonitorBossVfx.gd','monitor_expression.gdshader','monitor_code.gdshader']]
files += [B/'source/rig_contract_v031.json',B/'README.md',B/'source/boss002_production_ledger.xlsx']
transfer.update(status='active',files=[{'path':rel(p),'sha256':sha(p),'bytes':p.stat().st_size} for p in sorted(files)],git_baseline=subprocess.check_output(['git','rev-parse','HEAD'],cwd=R,text=True).strip())
transfer['generated_at']=datetime.datetime.now().astimezone().isoformat(timespec='seconds');transfer['generated_by']='Codex'
transfer['sources']=[{'kind':k,'path':p,'sha256':h,'bytes':(R/p).stat().st_size,'skeleton_id':transfer['skeleton_id'],'skeleton_signature':transfer[k+'_signature']} for k,(p,h) in zip(['model','animation'],transfer['source_sha256'].items())]
transfer['parts']=[{'slot_id':'body','variant_id':'default','path':rel(B/'components/enm_boss_monitor002/enm_boss_monitor002_visual_top3d.glb'),'includes':['screen','base','rear_support','spring_arms','gloves','owned_keyboard','owned_cable'],'excludes':['studio','preview_vfx','cameras','lights','collision']}]
transfer['consumers']={'packed_scene':rel(B/'runtime/enm_boss_monitor002/enm_boss_monitor002_root_top3d.tscn'),'controller':'src/enemy3d/Enemy3D.gd','strategy':'src/enemy3d/MonitorBossCombat.gd','presentation':'src/enemy3d/MonitorBossPresentation.gd','state_machine':'Enemy3D.VALID_STATES','verification':['tests/verification/verify_monitor_boss_flow.tscn','tests/verification/verify_monitor_boss_visual.tscn']}
transfer['stage_history']=[{'stage':s,'date':'2026-10-06'} for s in ['authored','exported_pending_godot_validation','validated','active']]
checks=read(B/'previews/runtime/flow_report.json')['checks']
transfer['verification']={'flow':'previews/runtime/flow_report.json','checks':checks,'real_renderer':'previews/runtime/visual_report.json','rendered_frames':sum(1 for p in (B/'previews/runtime/frames').glob('*.png')),'preview':'previews/runtime/monitor_runtime_preview.mp4','runtime_ai_all_four_skills':True,'source_mesh_bounds_tolerance_m':0.008,'source_bone_head_tolerance_m':0.002}
transfer['design']='docs/v0.1/design/Boss002显示器技能设计.md'
transfer['runtime_binding']={'owner':'Enemy3D','content_id':'boss_monitor002','states':12,'clips':16,'skills':4,'phases':3,'location':'expedition_01/floor_00/f00_boss','prefab_root_scale':1.0,'actor_scale':1.05,'bone_alias_rule':'original name if present; otherwise original + _2 (Godot importer disambiguation)','keyboard_attachment':'BoneAttachment3D synchronized after global pose overrides'}
guard=R/'outputs/asset_guard/ENM-BOSS-MONITOR002-3D_enm_boss_monitor002.json'
if guard.exists():transfer['duplicate_gate']=read(guard)
save(B/'enm_boss_monitor002_transfer_ledger_v031.json',transfer);save(B/'enm_boss_monitor002_transfer_ledger.json',transfer)
manifest=read(B/'asset_manifest.json')
for p in [B/'source/enm_boss_monitor002_model_v030.blend',B/'source/enm_boss_monitor002_animation_v030.blend']:assert sha(p)==manifest['files'][p.relative_to(B).as_posix()],'Prior master was changed'
manifest.update(version='v031',status='active',stage='active',model_source=rel(B/'source/enm_boss_monitor002_model_v031.blend'),animation_source=rel(B/'source/enm_boss_monitor002_animation_v031.blend'),runtime_integration=True,runtime_integrated=True,runtime_prefab=contract['runtime_prefab'],content_id='boss_monitor002',current_transfer_ledger=rel(B/'enm_boss_monitor002_transfer_ledger.json'),formal_animations=transfer['clips'],runtime_verification=transfer['verification'],rig_contract='source/rig_contract_v031.json',source='source/enm_boss_monitor002_model_v031.blend')
manifest.setdefault('sha256',{}).update({item['path']:item['sha256'] for item in transfer['files']});manifest['sha256'].update(transfer['source_sha256']);save(B/'asset_manifest.json',manifest)
manifest['formal_animations']=list(transfer['clips']);manifest['clip_specs']=transfer['clips'];manifest['formal_animations_authored']=list(transfer['clips'])
manifest['files'].update({p.relative_to(B).as_posix():sha(p) for p in files if p.is_relative_to(B)})
for name in ['model','animation']:p=B/f'source/enm_boss_monitor002_{name}_v031.blend';manifest['files'][p.relative_to(B).as_posix()]=sha(p)
save(B/'asset_manifest.json',manifest)
# Accept only this asset and its two intentionally edited specialized sheets.
sys.path[:0]=[str(R/'scripts'),str(R/'tools/asset_pipeline')]
from ledger_registry import LedgerIndex
from split_asset_ledger import _row_digest,read_source_rows,sheet_digest,col_digest,CONTENT_COLUMNS
index=LedgerIndex.load(R);p=index.path_for_asset_id(transfer['asset_id'])
wb=load_workbook(p);rows=read_source_rows(wb['资产主表']);row=next(v for _,v in rows if v[0]==transfer['asset_id'])
bp=R/'assets/registry/ledger_split_baseline.json';baseline=read(bp);old=read(D/'baseline_before.json')
baseline['assets'][transfer['asset_id']]['v']=_row_digest(row)
for sheet in ['3D-敌人','敌人动画与状态']:baseline['sheet_digests'][sheet]=sheet_digest(wb[sheet])
assert {k:v for k,v in baseline['assets'].items() if k!=transfer['asset_id']}=={k:v for k,v in old['assets'].items() if k!=transfer['asset_id']},'Do not accept unrelated asset drift'
baseline['targeted_updates']=[u for u in baseline.get('targeted_updates',[]) if not (u.get('asset_id')==transfer['asset_id'] and u.get('version')=='v031')]
collected=[]
for domain in index.domains:
 for _,values in read_source_rows(load_workbook(domain.path)['资产主表']):
  assert _row_digest(values)==baseline['assets'][values[0]]['v'],('Unrelated drift must not enter column baseline',values[0])
  collected.append((values[0],values))
assert {a for a,v in collected}==set(baseline['assets'])
collected.sort(key=lambda x:x[0]);baseline['column_digests']={str(c):col_digest(collected,c) for c in CONTENT_COLUMNS}
baseline['targeted_updates'].append({'date':'2026-10-06','asset_id':transfer['asset_id'],'version':'v031','status':'active','scope':'One asset row plus 3D-敌人 / 敌人动画与状态 sparse edits; union column digests refreshed only after every unrelated row fingerprint matched the prior baseline.'})
save(bp,baseline)
changes=(D/'changes.json').read_text(encoding='utf-8').replace('478项',str(checks)+'项').replace('490项',str(checks)+'项');(D/'changes.json').write_text(changes,encoding='utf-8')
print('Finalized v031 transfer, hashes and targeted baseline')
