import openpyxl
p = "assets/registry/ledgers/ShellStorm2_场景账本_v001.xlsx"
wb = openpyxl.load_workbook(p, data_only=False)
for name in ("总览","账本说明","域变更日志"):
    ws = wb[name]
    print("###### SHEET", name, ws.dimensions, ws.max_row, ws.max_column)
    for r in range(1, min(ws.max_row, 30)+1):
        vals = [ws.cell(r,c).value for c in range(1, ws.max_column+1)]
        vals = [v for v in vals if v is not None]
        if vals: print("  R%03d" % r, vals[:8])
print()
for name in ("3D-场景通用","3D-设施","3D-其他"):
    ws = wb[name]
    print("###### SHEET", name, ws.dimensions, ws.max_row, ws.max_column)
    for r in range(1, min(ws.max_row, 4)+1):
        print("  R%03d" % r, [ws.cell(r,c).value for c in range(1, ws.max_column+1)])
    print("  ... tail:")
    for r in range(max(1, ws.max_row-4), ws.max_row+1):
        print("  R%03d" % r, [ws.cell(r,c).value for c in range(1, ws.max_column+1)])
