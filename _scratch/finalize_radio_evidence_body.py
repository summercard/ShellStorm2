from pathlib import Path
from collections import Counter
import hashlib,json,shutil
import openpyxl
root=Path('I:/工作项目/shellstrom2/ShellStorm2'); books=root/'assets/registry/ledgers'; prop=books/'ShellStorm2_道具账本_v001.xlsx'; master=root/'assets/registry/ShellStorm2_美术资产台账_v001.xlsx'; basep=root/'assets/registry/ledger_split_baseline.json'; out=root/'outputs/base99_radio_v001'; back=out/'backup_20261005'; aid='PRP-BASE99-RADIO-3D'
def t(v): return '' if v is None else str(v).strip()
def rrows(ws): return [(r,[ws.cell(r,c).value for c in range(1,26)]) for r in range(6,ws.max_row+1) if t(ws.cell(r,1).value)]
def rd(v): return hashlib.sha256('\x1f'.join(t(v[c-1]) for c in range(1,26) if c not in (18,19)).encode('utf-8')).hexdigest()
def cd(rows,c):
 o=sorted(rows,key=lambda x:(t(x[1][0]),t(x[1][c-1])))
 return hashlib.sha256('\x1e'.join(t(v[c-1]) for _,v in o).encode('utf-8')).hexdigest()
trace=[]
def ck(x): trace.append(x)
out.mkdir(parents=True,exist_ok=True); back.mkdir(parents=True,exist_ok=True); ck('dirs')
for p in [prop,master,basep]:
 q=back/p.name
 if not q.exists(): shutil.copy2(p,q)
ck('backups')
w=openpyxl.load_workbook(prop,data_only=False); s=w['资产主表']; found=[r for r in range(6,s.max_row+1) if s.cell(r,1).value==aid]; ck('found '+repr(found)); assert found==[27],found
for r in range(6,28):
 s.cell(r,18).value=f'=LOWER(TRIM(C{r})&"|"&TRIM(D{r})&"|"&TRIM(E{r})&"|"&TRIM(F{r})&"|"&TRIM(H{r})&"|"&TRIM(I{r}))'
 s.cell(r,19).value=f'=IF(COUNTIF($R$6:$R$27,R{r})>1,"重复","唯一")'
assert s.max_row==27 and w['3D-道具'].max_row==8 and w['3D-道具'].cell(8,1).value==aid; ck('assert ledger')
assert w['域变更日志'].max_row==11 and w['域变更日志'].cell(11,1).value=='v0.1.5'; ck('assert log')
for d in s.data_validations.dataValidation: assert all(str(x).endswith('27') for x in d.sqref.ranges)
ck('assert dv'); w.save(prop); ck('save prop')
m=openpyxl.load_workbook(master); m['总览']['D18']=21; m['3D Prefab总控']['C7']=4; m['3D Prefab总控']['E7']=2; m['3D Prefab总控']['F7']=2; m['3D Prefab总控']['G7']='桌椅、家具、收音机、容器和摆放装饰物'; m.save(master); ck('save master')
idx=json.loads((root/'assets/registry/ledger_index.json').read_text(encoding='utf-8')); allrows=[]
for d in idx['domains']:
 x=openpyxl.load_workbook(books/d['file'],read_only=False)
 allrows += [(str(v[0]),v) for _,v in rrows(x['资产主表'])]
 x.close()
lookup=dict(allrows); assert aid in lookup; ck('read all rows')
b=json.loads(basep.read_text(encoding='utf-8')); b['assets'][aid]={'v':rd(lookup[aid]),'c':'道具','d':'props'}; b['asset_count']=len(b['assets']); b['captured_at']='2026-10-05'; b['column_digests']={str(c):cd(allrows,c) for c in range(1,26) if c not in (18,19)}; b['category_counts']=dict(sorted(Counter(t(v[2]) for _,v in allrows).items())); basep.write_text(json.dumps(b,ensure_ascii=False,indent=1)+'\n',encoding='utf-8'); ck('save baseline')
e={'asset_id':aid,'ledger':str(prop),'asset_sheet_row':27,'category':'道具','subcategory':'decor_prop','prefab_sheet':'3D-道具','prefab_sheet_row':8,'log_sheet_row':11,'status':'正式美术已接入','model_gate_status':'未全过：StatusLight 命名误报保留；不将模型门禁写成全过。','backup_dir':str(back),'backup_files':[str(back/p.name) for p in [prop,master,basep]],'baseline':{'path':str(basep),'asset_count':b['asset_count'],'captured_at':b['captured_at']},'runtime_contract':{'state_cycle':['off','music_a','music_b','off'],'bus':'Music','input':['鼠标左键','E']},'paths':{'source_blend':'assets/art/props/base_world_3d/source/base99_radio/prp_base99_radio_source_v001.blend','glb':'assets/art/props/base_world_3d/components/base99_radio/prp_base99_radio_visual_top3d.glb','stable_prefab':'assets/art/props/base_world_3d/runtime/base99_radio/prp_base99_radio_root_top3d.tscn','script':'src/base3d/Base99Radio3D.gd','formal_scene':'assets/art/environments/base_facility_3d/runtime/env_base_facility_art_layout_top3d.tscn'}}
(out/'registration_evidence.json').write_text(json.dumps(e,ensure_ascii=False,indent=2)+'\n',encoding='utf-8'); (out/'registration_evidence.txt').write_text('\n'.join([f'AssetID: {aid}',f'账本: {prop}','资产主表行: 27','分类: 道具 / decor_prop / 3D-道具','Prefab分页行: 8','域变更日志行: 11','正式状态: 正式美术已接入','模型门禁: 未全过；StatusLight 命名误报保留，未宣称全过',f'备份目录: {back}','备份文件: ShellStorm2_道具账本_v001.xlsx, ShellStorm2_美术资产台账_v001.xlsx, ledger_split_baseline.json',f'无损基线: {basep}'])+'\n',encoding='utf-8'); ck('evidence')
(root/'_scratch/finalize_radio_checkpoint.txt').write_text('\n'.join(trace)+'\n',encoding='utf-8')
