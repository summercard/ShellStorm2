# -*- coding: utf-8 -*-
"""检索账本里远征01 房型资产的登记现状。只读。"""
import sys
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8")
from openpyxl import load_workbook

LEDGER = Path("assets/registry/ledgers/ShellStorm2_场景账本_v001.xlsx")
wb = load_workbook(LEDGER, data_only=False)

print("=== sheets ===")
for ws in wb.worksheets:
    print("  %-16s rows=%-5d cols=%d" % (ws.title, ws.max_row, ws.max_column))
print()

for name in ("资产主表", "3D-场景通用"):
    ws = wb[name]
    print("#" * 70)
    print("### %s" % name)
    # 表头：找到含"资产ID"的行
    hdr_row = None
    for r in range(1, 8):
        vals = [str(ws.cell(row=r, column=c).value or "") for c in range(1, ws.max_column + 1)]
        if any("资产ID" in v or "AssetID" in v for v in vals):
            hdr_row = r
            break
    print("表头行 =", hdr_row)
    if hdr_row:
        for c in range(1, ws.max_column + 1):
            v = ws.cell(row=hdr_row, column=c).value
            if v:
                print("   col%-3d %s" % (c, v))
    print()
    print("--- 含远征/房型关键词的行 ---")
    keys = ("远征", "expedition", "EXPEDITION", "走廊", "Boss", "BOSS", "竞技场",
            "安全屋", "撤离", "数据库", "通道桥", "办公室", "拐角", "L型", "l_corridor")
    for r in range((hdr_row or 1) + 1, ws.max_row + 1):
        rowvals = [ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)]
        txt = " ".join(str(v) for v in rowvals if v is not None)
        if any(k in txt for k in keys):
            print("row %d:" % r)
            for c in range(1, ws.max_column + 1):
                v = rowvals[c - 1]
                if v is not None and str(v).strip():
                    print("    c%-3d = %s" % (c, str(v)[:110]))
            print()
