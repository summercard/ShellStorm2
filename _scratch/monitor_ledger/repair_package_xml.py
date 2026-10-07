import zipfile,xml.etree.ElementTree as E,shutil
from pathlib import Path
from openpyxl import load_workbook
R=Path.cwd();p=R/'docs/v0.1/data/ShellStorm2_游戏内容数据库_v010.xlsx';before=load_workbook(p)
tmp=R/'_scratch/monitor_ledger/package_validated.xlsx'
with zipfile.ZipFile(p) as original,zipfile.ZipFile(tmp,'w',zipfile.ZIP_DEFLATED) as out:
 for info in original.infolist():
  data=original.read(info.filename)
  if info.filename in ['[Content_Types].xml','xl/_rels/workbook.xml.rels']:
   element=E.fromstring(data);E.register_namespace('',element.tag.split('}')[0][1:]);data=E.tostring(element,encoding='utf-8',xml_declaration=True)
  out.writestr(info,data)
after=load_workbook(tmp)
for name in before.sheetnames:
 for row in before[name]:
  for c in row:assert after[name][c.coordinate].value==c.value
shutil.copy2(tmp,p);print('OPC metadata default namespaces restored; all values and formulas preserved')
