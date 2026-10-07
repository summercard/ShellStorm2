import runpy,zipfile,xml.etree.ElementTree as E,copy,shutil
from pathlib import Path
from openpyxl import load_workbook
R=Path.cwd();D=R/'_scratch/monitor_ledger'
# Reuse helpers without executing the merge loop.
code=(D/'merge.py').read_text(encoding='utf-8').split('for i,job in enumerate(jobs):')[0]
scope={};exec(code,scope);sheets=scope['sheets'];ns=scope['ns'];M=scope['M']
p=R/'docs/v0.1/data/ShellStorm2_游戏内容数据库_v010.xlsx';before=load_workbook(p)
with zipfile.ZipFile(p) as original,zipfile.ZipFile(D/'cells3.xlsx') as generated:
 target=sheets(original)['Boss002技能设计'];src=sheets(generated)['Boss002技能设计']
 style=E.fromstring(original.read('xl/styles.xml'));mapping=scope['styles_map'](style,E.fromstring(generated.read('xl/styles.xml')))
 sheet=E.fromstring(generated.read(src));strings=[]
 if 'xl/sharedStrings.xml' in generated.namelist():strings=[''.join(x.itertext()) for x in E.fromstring(generated.read('xl/sharedStrings.xml'))]
 for c in sheet.findall('.//m:c',ns):
  c.set('s',str(mapping[int(c.get('s','0'))]))
  if c.get('t')=='s':
   value=strings[int(c.find('m:v',ns).text)];c.remove(c.find('m:v',ns));c.set('t','inlineStr');E.SubElement(E.SubElement(c,'{'+M+'}is'),'{'+M+'}t').text=value
 tmp=D/'polished.xlsx'
 with zipfile.ZipFile(tmp,'w',zipfile.ZIP_DEFLATED) as out:
  for info in original.infolist():out.writestr(info,E.tostring(sheet,encoding='utf-8',xml_declaration=True) if info.filename==target else E.tostring(style,encoding='utf-8',xml_declaration=True) if info.filename=='xl/styles.xml' else original.read(info.filename))
after=load_workbook(tmp)
for name in before.sheetnames:
 for row in before[name]:
  for c in row:assert after[name][c.coordinate].value==c.value
shutil.copy2(tmp,p);print('Boss002 new design page formatted; all cell values preserved')
