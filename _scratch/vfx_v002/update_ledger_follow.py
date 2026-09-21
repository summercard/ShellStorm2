## 特效账本更新：FX01-01 枪口花火新增挂点跟随契约 + 尺寸基准 0.8；域变更日志加 v0.1.2
import shutil
import openpyxl

PATH = r"I:\工作项目\shellstrom2\ShellStorm2\assets\registry\ledgers\ShellStorm2_特效账本_v001.xlsx"
BACKUP = r"I:\工作项目\shellstrom2\ShellStorm2\_scratch\vfx_v002\ledger_before_1802.xlsx"
shutil.copy2(PATH, BACKUP)

wb = openpyxl.load_workbook(PATH)
ws = wb["3D-特效"]

g = ws["G8"]  # 功能说明
q = ws["Q8"]  # 备注
assert str(ws["B8"].value).strip() == "VFX-MUZZLE-FLASH-3D", "R8 不是枪口花火行：%s" % ws["B8"].value

g.value = str(g.value).rstrip("。") + (
    "。2026-09-21 补丁：支持 context.follow / follow_local_offset 挂点跟随——"
    "每帧把世界变换贴到挂点（枪口），角色移动、转向与后坐力位移时不残留；未绑定时退化为世界锚定。"
)
q.value = str(q.value or "") + (
    "；2026-09-21 补丁：新增挂点跟随契约（follow），调用方 WeaponModel3D 以 muzzle 挂点为跟随目标；"
    "调用方尺寸基准定档 0.8（视觉体积 80%）。"
)

log = wb["域变更日志"]
log.append([
    "v0.1.2",
    "2026-09-21",
    "条目更新",
    "特效 · 3D-特效",
    "FX01-01 枪口花火：新增 context.follow 挂点跟随契约（角色移动/转向/后坐力时特效始终贴着枪口，"
    "不再残留在开火瞬间的世界坐标）；调用方尺寸基准定为 0.8（视觉体积 80%）；"
    "专项验收补端到端接线断言 + 两轮反向对照",
    "AssetID / Prefab 路径 / 尺寸契约不变；行为为向后兼容增强（未传 follow 时退化为世界锚定）",
    "摩斯拉",
])

wb.save(PATH)

# 回读校验
wb2 = openpyxl.load_workbook(PATH)
ws2 = wb2["3D-特效"]
print("R8 版本列   :", ws2["P8"].value)
print("R8 备注尾部 :", str(ws2["Q8"].value)[-60:])
print("功能说明尾部:", str(ws2["G8"].value)[-60:])
log2 = wb2["域变更日志"]
print("日志末行    :", [str(c.value)[:24] for c in log2[log2.max_row]])
print("DV 保留     :", len(ws2.data_validations.dataValidation), "合并格:", len(ws2.merged_cells.ranges))
