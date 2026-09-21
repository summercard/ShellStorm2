# -*- coding: utf-8 -*-
"""特效账本更新：FX01-01/FX01-02 升 v002 卡通版；新增 FX01-06 飞行子弹纯视觉。

只动 ShellStorm2_特效账本_v001.xlsx：
  1) 3D-特效 分页：row8/row9 的 功能说明/标准尺寸/原点与朝向/版本/备注；
  2) 3D-特效 分页：在 §01 战斗反馈 组内插入一行 FX01-06（原 row13 起下移）；
  3) 总览口径文字（条数 10 -> 11）与进度注记；
  4) 域变更日志 追加 v0.1.1 记录。
资产主表不动（FX01-* 分页条目本就不在门禁计数范围内）。
"""
import copy
import openpyxl

LEDGER = "assets/registry/ledgers/ShellStorm2_特效账本_v001.xlsx"

ROW_MUZZLE = 8
ROW_IMPACT = 9
INSERT_AT = 13

MUZZLE_FUNC = (
    "枪械开火时的瞬时枪口花火（卡通 v002）：铂金亮芯 + 橙金/弹药色立体楔形火焰"
    "（CylinderMesh top_radius=0，Godot 4.x 无 ConeMesh）+ 交叉面尖瓣 + 5 枚飞散火星"
    " + 短冲击环 + 一盏 OmniLight3D（名 MuzzleLight，shadow_enabled=false）。"
    "经 VfxPool3D 按 AssetID 路由。"
)
MUZZLE_SPEC = (
    "亮芯约 0.18m；火焰长 0.7~1.8×size；lifetime 0.16s；"
    "MuzzleLight range 3.5m×size、峰值 light_energy 5.0、寿命末端归零"
)
MUZZLE_ORIGIN = (
    "原点=根节点中心；朝向=Basis.looking_at(context.forward)，local -Z 对齐射击方向"
)
MUZZLE_NOTE = (
    "v002 卡通重做；Prefab=assets/art/vfx/combat_3d/vfx_muzzle_flash_root_top3d.tscn"
    "（去版本化路径，N1/N2 后不再带 _v001）+ 脚本=src/vfx/VfxMuzzleFlash3D.gd；"
    "根节点=VfxMuzzleFlash3D(Node3D)；FX编号 FX01-01；"
    "调用方：WeaponModel3D._spawn_muzzle_effect 经 VfxPool3D.acquire(FX01_MUZZLE_FLASH, …, "
    '{"forward": -global_basis.z})；验收：verify_combat_vfx_toon_v002。'
    "注：旧版备注中的 ..._v001.tscn 路径已废弃。"
)

IMPACT_FUNC = (
    "子弹/投射物命中表面的爆点（卡通 v002）：亮芯 + 冲击环（TorusMesh）"
    " + 7 道错峰星芒（BoxMesh，Quaternion(UP,dir) 放射） + 8 块错峰旋转碎屑；"
    "UNSHADED 半透明材质、透明重叠受控。"
)
IMPACT_SPEC = "亮芯约 0.16m；冲击环 0.2→2.4×size；lifetime 0.32s"
IMPACT_ORIGIN = (
    "原点=根节点中心；朝向=Basis(Quaternion(UP, context.normal))，local +Y 对齐命中法线"
)
IMPACT_NOTE = (
    "v002 卡通重做；Prefab=assets/art/vfx/combat_3d/vfx_impact_root_top3d.tscn"
    "（去版本化路径）+ 脚本=src/vfx/VfxImpact3D.gd；根节点=VfxImpact3D(Node3D)；"
    "FX编号 FX01-02；调用方：Projectile3D._spawn_effect(VfxPool3D.FX01_IMPACT, …) 传 "
    'hit_context={"normal":…, "position":…}；Enemy3D / PlayerMeleeCombat3D 亦经 VfxPool3D 路由；'
    "验收：verify_combat_vfx_toon_v002。注：旧版备注中的 ..._v001.tscn 路径已废弃。"
)

BULLET_ROW = [
    "FX01-06",
    "VFX-BULLET-VISUAL-3D",
    "飞行子弹纯视觉特效",
    "assets/art/vfx/combat_3d/vfx_bullet_visual_root_top3d.tscn",
    None,
    None,
    (
        "飞行中子弹的纯视觉表现（卡通 v002 风格同族）：铂金亮芯胶囊 + 弹色半透明外壳"
        " + 沿弹道拖尾锥（CylinderMesh top_radius=0，宽端贴弹体、尖端向身后）"
        " + 每弹一盏 OmniLight3D（名 ProjectileLight，shadow_enabled=false，"
        "light_energy 1.2、omni_range 1.8）。由 Projectile3D 作为子节点持有，"
        "生命周期随宿主（ProjectilePool3D），不进 VfxPool 自动计时，避免双池 ownership。"
    ),
    "src/vfx/VfxBulletVisual3D.gd",
    "无",
    "无",
    "无",
    "弹芯 CapsuleMesh r=0.045/h=0.20；外壳 r=0.075/h=0.18；拖尾锥底半径 0.06/长 0.5；灯 range 1.8m",
    "原点=宿主原点（local position 保持 0，随 Projectile3D 变换）；朝向=弹体沿 local -Z，拖尾向 local +Z（身后）",
    "玩家/敌人飞行子弹（远程 7 枪 + 各类投射物）",
    "已完成",
    "v001",
    (
        "已实装；Prefab=assets/art/vfx/combat_3d/vfx_bullet_visual_root_top3d.tscn"
        " + 脚本=src/vfx/VfxBulletVisual3D.gd；根节点=VfxBulletVisual3D(Node3D)；"
        "FX编号 FX01-06；调用方：Projectile3D._build_visual() 实例化为子节点并委托 "
        "activate/deactivate/apply_growth/set_trail_visible；"
        "碰撞由 Projectile3D 自持（半径 0.12 球），本视觉 Prefab 内无碰撞体；"
        "验收：verify_combat_vfx_toon_v002。"
    ),
]

CHANGELOG_ROW = [
    "v0.1.1",
    "2026-09-21",
    "条目更新 + 新增",
    "特效 · 3D-特效",
    (
        "FX01-01 枪口花火 / FX01-02 命中爆点 卡通 v002 重做（枪口新增 OmniLight3D "
        "MuzzleLight 与 forward 对齐；命中改为按 context.normal 对齐）；新增 FX01-06 "
        "VFX-BULLET-VISUAL-3D 飞行子弹纯视觉 Prefab（含 ProjectileLight）；"
        "同期修正备注中已废弃的 ..._v001.tscn 路径写法。"
    ),
    "AssetID 不变；仅新增一条 FX01-06；Prefab 路径保持去版本化；资产主表未改动",
    "摩斯拉",
]

NOTE_APPEND = (
    "【2026-09-21 进度】新增 FX01-06 VFX-BULLET-VISUAL-3D（飞行子弹纯视觉，含 ProjectileLight，"
    "由 Projectile3D 持有、不进 VfxPool）；FX01-01 / FX01-02 升 v002 卡通重做"
    "（枪口新增 MuzzleLight、命中改按命中法线对齐）；三类共同验收 verify_combat_vfx_toon_v002。"
)


def _copy_style(src_cell, dst_cell):
    dst_cell._style = copy.copy(src_cell._style)


def main():
    wb = openpyxl.load_workbook(LEDGER)
    ws = wb["3D-特效"]
    assert ws.max_row == 19, "3D-特效 行数与预期不符，拒绝盲写：%d" % ws.max_row

    # --- 1) row8 / row9 升 v002 ---
    for row, func, spec, origin, note in (
        (ROW_MUZZLE, MUZZLE_FUNC, MUZZLE_SPEC, MUZZLE_ORIGIN, MUZZLE_NOTE),
        (ROW_IMPACT, IMPACT_FUNC, IMPACT_SPEC, IMPACT_ORIGIN, IMPACT_NOTE),
    ):
        assert ws.cell(row, 2).value in (
            "VFX-MUZZLE-FLASH-3D",
            "VFX-IMPACT-3D",
        ), "row%d AssetID 与预期不符：%s" % (row, ws.cell(row, 2).value)
        ws.cell(row, 7).value = func      # G 功能说明
        ws.cell(row, 12).value = spec     # L 标准尺寸
        ws.cell(row, 13).value = origin   # M 原点与朝向
        ws.cell(row, 16).value = "v002"   # P 版本
        ws.cell(row, 17).value = note     # Q 备注

    # --- 2) 插入 FX01-06（§01 组内，原 row13 起下移） ---
    style_src_row = 12  # FX01-05 条目行，作为样式模板
    ws.insert_rows(INSERT_AT)
    for col, value in enumerate(BULLET_ROW, start=1):
        cell = ws.cell(INSERT_AT, col)
        cell.value = value
        _copy_style(ws.cell(style_src_row, col), cell)
    # 行高/对齐沿用模板行
    ws.row_dimensions[INSERT_AT].height = ws.row_dimensions[style_src_row].height

    # --- 3) 总览口径与进度注记 ---
    intro = ws.cell(2, 1).value
    assert intro is not None and "本 sheet 共 10 条" in intro, "3D-特效 导语形态与预期不符"
    ws.cell(2, 1).value = intro.replace("本 sheet 共 10 条", "本 sheet 共 11 条") + NOTE_APPEND

    # --- 4) 域变更日志 追加 ---
    log = wb["域变更日志"]
    last = log.max_row
    assert log.cell(last, 1).value == "v0.1.0", "变更日志末行与预期不符：%s" % log.cell(last, 1).value
    new_row = last + 1
    for col, value in enumerate(CHANGELOG_ROW, start=1):
        cell = log.cell(new_row, col)
        cell.value = value
        _copy_style(log.cell(last, col), cell)

    wb.save(LEDGER)
    print("LEDGER_UPDATED: 3D-特效 rows=%d, changelog rows=%d" % (ws.max_row, log.max_row))


if __name__ == "__main__":
    main()
