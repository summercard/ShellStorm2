# -*- coding: utf-8 -*-
"""为本次「房型资产转正」补录无损基线（最小外科编辑）。

口径与先例 tools/ledger_reconcile_baseline_spikeshell.py 一致，理由重申：
  - verify_ledger_split.py 对「基线里没有的 AssetID」直接 fail(asset_not_in_baseline)，
    不区分「静默新增」与「有意新增」⇒ 新增资产必须补录。
  - ⛔ 不能跑 split_asset_ledger.py --baseline-only：它从 ledger_index.json 的 master.path
    （总目录）读《资产主表》，而总目录已不持有资产行 ⇒ 会清空基线。
  - 只加新行指纹、只更新本次改动的两行指纹；其余 413 条逐位不动；再重算列摘要与分类计数。
"""
from __future__ import annotations

import json
import shutil
import sys
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(r"I:\工作项目\shellstrom2\ShellStorm2")
sys.path.insert(0, str(ROOT / "scripts"))
sys.path.insert(0, str(ROOT / "tools" / "asset_pipeline"))

from ledger_registry import LedgerIndex  # noqa: E402
from split_asset_ledger import (  # noqa: E402
    ASSET_SHEET,
    CONTENT_COLUMNS,
    col_digest,
    read_source_rows,
    _row_digest,
)

NEW_ID = "ENV-EXPEDITION-L01-BOSS-ROOM"
UPDATED_IDS = ("ENV-EXPEDITION-L01-OFFICE-ROOM", "ENV-EXPEDITION-L01-BRIDGE-ROOM")
BASELINE = ROOT / "assets/registry/ledger_split_baseline.json"
LEDGER = ROOT / "assets/registry/ledgers/ShellStorm2_场景账本_v001.xlsx"

index = LedgerIndex.load(ROOT)
bl = json.loads(BASELINE.read_text(encoding="utf-8"))
before_count = bl["asset_count"]

if NEW_ID in bl["assets"]:
    raise SystemExit("基线已含 %s —— 本脚本不该重跑" % NEW_ID)

wb = load_workbook(LEDGER)
rows = read_source_rows(wb[ASSET_SHEET])
by_id = {str(v[0]).strip(): (r, v) for r, v in rows}

row_number, values = by_id[NEW_ID]
assert values[2] == "场景", values[2]
digest = _row_digest(values)
bl["assets"][NEW_ID] = {"v": digest, "c": "场景", "d": "scenes"}
print("BASELINE_ADD %s r%d v=%s..." % (NEW_ID, row_number, digest[:16]))

for aid in UPDATED_IDS:
    r, v = by_id[aid]
    old = bl["assets"][aid]["v"]
    new = _row_digest(v)
    assert old != new, "%s 指纹未变，脚本不该改它" % aid
    bl["assets"][aid]["v"] = new
    print("BASELINE_UPDATE %s r%d %s... -> %s..." % (aid, r, old[:12], new[:12]))

bl["asset_count"] = len(bl["assets"])

union = []
totals = {}
for domain in index.domains:
    w = load_workbook(domain.path)
    r = read_source_rows(w[ASSET_SHEET])
    totals[domain.key] = len(r)
    union.extend(r)
bl["column_digests"] = {str(c): col_digest(union, c) for c in CONTENT_COLUMNS}
bl["category_counts"] = dict(sorted(Counter(str(v[2]).strip() for _r, v in union).items()))

shutil.copy2(BASELINE, BASELINE.with_name(BASELINE.name + ".bak_before_roomtype"))
BASELINE.write_text(json.dumps(bl, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")
print("BASELINE_RECONCILED %d -> %d" % (before_count, bl["asset_count"]))
print("ledger_totals", totals)
print("category_counts", bl["category_counts"])
print("BASELINE_ROOMTYPE_DONE")
