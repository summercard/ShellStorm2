# -*- coding: utf-8 -*-
"""修正 P247（源码/策划依据）并同步无损基线。

原因：`check_expedition_room_asset_status.py` 判据 D 的 registered_blob 只拼
「资产主表 P 列（source）」与「3D-场景通用 D 列（glb）」——**不含 O 列文件路径**。
Boss 行原先 P 列只写组件库/工具路径，不含 `boss_room` 字样 ⇒ 判 D 报「boss_room
既没登记也没标未登记」。把房型源 manifest 加进去即解决（这也更符合 P 列语义）。
"""
from __future__ import annotations

import json
import sys
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(r"I:\工作项目\shellstrom2\ShellStorm2")
sys.path.insert(0, str(ROOT / "scripts"))
sys.path.insert(0, str(ROOT / "tools" / "asset_pipeline"))

from ledger_registry import LedgerIndex  # noqa: E402
from split_asset_ledger import (  # noqa: E402
    ASSET_SHEET,
    CONTENT_COLUMNS,
    col_digest,
    read_source_rows,
    _row_digest,
)

LEDGER = ROOT / "assets/registry/ledgers/ShellStorm2_场景账本_v001.xlsx"
BASELINE = ROOT / "assets/registry/ledger_split_baseline.json"
TARGET = "ENV-EXPEDITION-L01-BOSS-ROOM"
NEW_SOURCE = (
    "assets/art/environments/tower_zones/expedition/source/room_types/"
    "boss_room/v002/room_type_manifest.json；"
    "assets/art/environments/tower_zones/expedition/source/common_components/v008/component_plan.json；"
    "component_slug_map.json；归并工具 scripts/blender/regroup_room_type_components.py"
)

wb = load_workbook(LEDGER)
ws = wb[ASSET_SHEET]
by_id = {str(v[0]).strip(): (r, v) for r, v in read_source_rows(ws)}
row_number, _values = by_id[TARGET]
assert row_number == 247, row_number
old_source = str(ws.cell(row_number, 16).value or "")
assert "boss_room" not in old_source, "P 列已含 boss_room，无需再修"
ws.cell(row_number, 16).value = NEW_SOURCE
wb.save(LEDGER)
print("P247_FIXED")

index = LedgerIndex.load(ROOT)
bl = json.loads(BASELINE.read_text(encoding="utf-8"))
wb2 = load_workbook(LEDGER)
by2 = {str(v[0]).strip(): (r, v) for r, v in read_source_rows(wb2[ASSET_SHEET])}
_r, values = by2[TARGET]
old_digest = bl["assets"][TARGET]["v"]
bl["assets"][TARGET]["v"] = _row_digest(values)
print("BASELINE_P_UPDATE %s... -> %s..." % (old_digest[:12], bl["assets"][TARGET]["v"][:12]))

union = []
for domain in index.domains:
    w = load_workbook(domain.path)
    union.extend(read_source_rows(w[ASSET_SHEET]))
bl["column_digests"] = {str(c): col_digest(union, c) for c in CONTENT_COLUMNS}
bl["category_counts"] = dict(sorted(Counter(str(x[2]).strip() for _r, x in union).items()))
BASELINE.write_text(json.dumps(bl, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")
print("BASELINE_REWRITTEN assets=%d" % bl["asset_count"])
print("FIX_P_DONE")
