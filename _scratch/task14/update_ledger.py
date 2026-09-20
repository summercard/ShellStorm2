# -*- coding: utf-8 -*-
"""2026-09-19 天台地板 + 外立面三件：把台账既有三行升为「正式美术已接入」。
只改既有行（不新增 AssetID / 不新增行），并在域变更日志追加一条 v0.1.2。
写完后与原快照逐格比对，确认没有误伤其它单元格。
"""
import shutil
import sys

import openpyxl

LEDGER = r"I:\工作项目\shellstrom2\ShellStorm2\assets\registry\ledgers\ShellStorm2_场景账本_v001.xlsx"
SNAP = r"I:\工作项目\shellstrom2\ShellStorm2\_scratch\task14\ledger_snapshot_before.xlsx"
OUT = r"I:\工作项目\shellstrom2\ShellStorm2\_scratch\task14\ledger_new.xlsx"

FLOOR_NOTE = (
    "2026-09-19：由「仅 Blender 源」升为运行时正式资产，本行不新增 AssetID。"
    "用户要求「天台的地板，帮我导出做一个新的 prefab，替代掉现在的地板，方法是新作一个 prefab，"
    "然后天台的生成代码从指向原来的改成这个」。导出 GLB "
    "assets/art/environments/tower_zones/rooftop/components/env_rooftop_ref_floor_full_top3d.glb"
    "（导出脚本 assets/art/environments/tower_zones/rooftop/source/export_env_rooftop_ref_floor_facade_v001.py），"
    "建立 prefab assets/art/props/dungeon_3d/prp_rooftop_floor_5m.tscn；"
    "TowerFloorStage3D.ROOFTOP_FLOOR_SCENE 接管天台地面，替代原 POLISHED_FLOOR_SCENE。"
    "⚠️ 原点契约与旧地板不同：新件 bottom_center（几何 Y=0..0.30），旧件为几何中心；"
    "_build_floor 改用 _floor_visual_origin_y(mesh) = -AABB.end.y 把顶面对齐到 Y=0，"
    "不再用 -FLOOR_THICKNESS*0.5。实测 234 块 5m 地砖。"
    "Godot 契约断言 tests/verification/verify_rooftop_floor_facade_components.tscn 与 "
    "verify_rooftop_32x32_contract.tscn 通过。"
)

FACADE_SOLID_NOTE = (
    "2026-09-19：由「仅 Blender 源」升为运行时正式资产，本行不新增 AssetID。"
    "用户要求「99 层的外墙 blender 组件中也有，正规导出后变成 prefab，在天台的周边调用然后围起来"
    "（其实是 99 层的外墙）——但是现在没有，所以在 100 层天台的边缘看到底下是空的」。"
    "导出 GLB env_rooftop_ref_facade_solid_top3d.glb，建立 prefab "
    "assets/art/props/dungeon_3d/prp_rooftop_facade_solid_5m.tscn。"
    "摆放口径：天台同一圈、低一层（y=-12，即 99F 楼面标高），外皮与女儿墙共面（内缩 0.15m = 厚度 0.30/2）；"
    "外立面环按「实 1 : 窗 2」每 3 槽一件实墙。"
    "TowerFloorStage3D._build_rooftop_facade_ring 生成实/窗两个 MultiMesh，"
    "并由 FacadeBoundaryCollision_{North/South/West/East} 提供 12m 高、0.30m 厚碰撞。"
    "实测整圈 68 槽（实 24 / 窗 44），四边满铺无缝；幕帘探针 y=-6 水平射线四边均命中 FacadeBoundaryCollision_*。"
)

FACADE_WINDOW_NOTE = (
    "2026-09-19：由「仅 Blender 源」升为运行时正式资产，本行不新增 AssetID。"
    "与 ENV-ROOFTOP-REF-FACADE-SOLID 同批升级：窗墙件占外立面环的另外 2/3（每 3 槽 2 件）。"
    "导出 GLB env_rooftop_ref_facade_window_top3d.glb，建立 prefab "
    "assets/art/props/dungeon_3d/prp_rooftop_facade_window_5m.tscn。"
    "摆放口径同实墙：天台同一圈、低一层（y=-12），外皮与女儿墙共面（内缩 0.15m）。"
    "实测整圈 68 槽（实 24 / 窗 44），四边满铺无缝；"
    "Godot 契约断言 verify_rooftop_32x32_contract.tscn 的 [外立面环] plan==actual 通过。"
)

LOG_DESC = (
    "天台地板与外立面共三件（ENV-ROOFTOP-REF-FLOOR-FULL 完整地砖 / ENV-ROOFTOP-REF-FACADE-SOLID 标准外墙实墙 / "
    "ENV-ROOFTOP-REF-FACADE-WINDOW 标准外墙窗墙）由「仅 Blender 源、未导入 Godot」升为运行时正式资产："
    "从 reference_components/v002 导出 GLB 并建立 prefab assets/art/props/dungeon_3d/prp_rooftop_floor_5m.tscn、"
    "prp_rooftop_facade_solid_5m.tscn、prp_rooftop_facade_window_5m.tscn。TowerFloorStage3D 三处接线："
    "①天台地面由 POLISHED_FLOOR_SCENE 改指 ROOFTOP_FLOOR_SCENE（原点契约由几何中心改为 bottom_center，"
    "改用 AABB 对齐顶面 Y=0）；②天台外墙不再挖楼梯口门洞，西侧围栏缺口闭合（门洞补位墙 3 件 → 0 件，"
    "直段 61 → 64，西侧碰撞由断成两段改为一条连续 Box）；③新增 _build_rooftop_facade_ring："
    "天台同一圈、低一层（y=-12 = 99F 楼面标高）摆 68 槽实/窗外墙（实 1 : 窗 2），"
    "并由 FacadeBoundaryCollision_* 提供 12m 高、0.30m 厚碰撞，补上「从天台边缘往下看到的空档」。"
    "资产主表未增删行；只升 3D-场景通用 既有三行。"
)

LOG_COMPAT = (
    "AssetID 与版本号不变（v002 = 参考组件库版本），仅补 prefab/GLB 路径、碰撞归属与实测记录。"
    "旧地板 POLISHED_FLOOR_SCENE 与旧占位矮墙 prp_tower_wall_parapet_door_5m 保留不动"
    "（后者仍被 DungeonRoom3D 引用且在 PREFAB_CONTRACT 清单内）。本次不新增 AssetID，不新增台账行。"
)

# row -> {col: new_value}
EDITS = {
    97: {
        3: "assets/art/props/dungeon_3d/prp_rooftop_floor_5m.tscn",
        4: "assets/art/environments/tower_zones/rooftop/components/env_rooftop_ref_floor_full_top3d.glb",
        8: "开",
        9: "TowerFloorStage3D",
        10: "BoxShape3D，由 TowerFloorStage3D 的 FloorSupport 按楼层轮廓铺满（Prefab 不再自带碰撞）",
        12: "bottom_center（几何 Y=0..0.30，顶面对齐 Y=0）；+Z 为行方向",
        13: "100F天台地面 / 234块5m地砖满铺（替代 POLISHED_FLOOR）",
        14: "正式美术已接入",
    },
    132: {
        3: "assets/art/props/dungeon_3d/prp_rooftop_facade_solid_5m.tscn",
        4: "assets/art/environments/tower_zones/rooftop/components/env_rooftop_ref_facade_solid_top3d.glb",
        8: "开",
        9: "TowerFloorStage3D",
        10: "BoxShape3D 5×12×0.30，由 FacadeBoundaryCollision_{North/South/West/East} 按槽位生成（Prefab 不再自带碰撞）",
        12: "bottom_center（几何 Y=0..11.90）；外皮与女儿墙共面",
        13: "100F天台外圈（同一圈、低一层 y=-12）/ 实墙 24 槽",
        14: "正式美术已接入",
    },
    133: {
        3: "assets/art/props/dungeon_3d/prp_rooftop_facade_window_5m.tscn",
        4: "assets/art/environments/tower_zones/rooftop/components/env_rooftop_ref_facade_window_top3d.glb",
        8: "开",
        9: "TowerFloorStage3D",
        10: "BoxShape3D 5×12×0.30，由 FacadeBoundaryCollision_{North/South/West/East} 按槽位生成（Prefab 不再自带碰撞）",
        12: "bottom_center（几何 Y=0..11.90）；外皮与女儿墙共面",
        13: "100F天台外圈（同一圈、低一层 y=-12）/ 窗墙 44 槽",
        14: "正式美术已接入",
    },
}
NOTES = {97: FLOOR_NOTE, 132: FACADE_SOLID_NOTE, 133: FACADE_WINDOW_NOTE}


def dump_cells(path):
    wb = openpyxl.load_workbook(path, data_only=False)
    data = {}
    for name in wb.sheetnames:
        ws = wb[name]
        cells = {}
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is not None:
                    cells[cell.coordinate] = cell.value
        data[name] = {
            "max_row": ws.max_row,
            "max_col": ws.max_column,
            "cells": cells,
            "merged": sorted(str(r) for r in ws.merged_cells.ranges),
            "dv": len(ws.data_validations.dataValidation),
        }
    return data


def main():
    wb = openpyxl.load_workbook(LEDGER, data_only=False)
    ws = wb["3D-场景通用"]

    for row, edits in EDITS.items():
        aid = ws.cell(row, 1).value
        assert aid in (
            "ENV-ROOFTOP-REF-FLOOR-FULL",
            "ENV-ROOFTOP-REF-FACADE-SOLID",
            "ENV-ROOFTOP-REF-FACADE-WINDOW",
        ), "row %d asset_id unexpected: %r" % (row, aid)
        for col, value in edits.items():
            ws.cell(row, col).value = value
        old = ws.cell(row, 16).value or ""
        ws.cell(row, 16).value = old + "\n" + NOTES[row]
        print("row %d updated (%s)" % (row, aid))

    log = wb["域变更日志"]
    last = log.max_row
    assert log.cell(last, 1).value == "v0.1.1", "unexpected last log version: %r" % log.cell(last, 1).value
    new_row = last + 1
    log.cell(new_row, 1).value = "v0.1.2"
    log.cell(new_row, 2).value = "2026-09-19"
    log.cell(new_row, 3).value = "资产升版"
    log.cell(new_row, 4).value = "关卡场景"
    log.cell(new_row, 5).value = LOG_DESC
    log.cell(new_row, 6).value = LOG_COMPAT
    log.cell(new_row, 7).value = "摩斯拉"
    print("log row %d appended: v0.1.2" % new_row)

    wb.save(OUT)
    print("saved ->", OUT)

    before = dump_cells(SNAP)
    after = dump_cells(OUT)
    problems = []
    if before.keys() != after.keys():
        problems.append("sheet set changed")
    for name in before:
        b, a = before[name], after[name]
        if b["max_col"] != a["max_col"]:
            problems.append("%s: max_col %s -> %s" % (name, b["max_col"], a["max_col"]))
        if b["merged"] != a["merged"]:
            problems.append("%s: merged changed" % name)
        if b["dv"] != a["dv"]:
            problems.append("%s: data_validation count %s -> %s" % (name, b["dv"], a["dv"]))
        for coord, value in b["cells"].items():
            if a["cells"].get(coord) != value:
                problems.append("%s!%s: %r -> %r" % (name, coord, value, a["cells"].get(coord)))
    # 预期变化：三行的 3/4/8/9/10/12/13/14/16 列 + 域变更日志新行
    expected_prefix = (
        "3D-场景通用!C", "3D-场景通用!D", "3D-场景通用!H", "3D-场景通用!I",
        "3D-场景通用!J", "3D-场景通用!L", "3D-场景通用!M", "3D-场景通用!N",
        "3D-场景通用!P",
    )
    unexpected = []
    for p in problems:
        if "域变更日志!" in p:
            continue
        if p.startswith(expected_prefix):
            continue
        unexpected.append(p)
    print("total diffs:", len(problems))
    for p in unexpected:
        print("  UNEXPECTED:", p)
    expected_diffs = [p for p in problems if p not in unexpected]
    print("expected diffs:", len(expected_diffs))
    if unexpected:
        print("RESULT: FAIL (unexpected changes)")
        return 1
    print("RESULT: OK (only intended cells changed)")
    return 0


if __name__ == "__main__":
    shutil.copyfile(LEDGER, SNAP) if not __import__("os").path.exists(SNAP) else None
    sys.exit(main())
