import openpyxl, collections, sys
p = r"I:\工作项目\shellstrom2\ShellStorm2\assets\registry\ShellStorm2_美术资产台账_v001.xlsx"
wb = openpyxl.load_workbook(p, data_only=True)
print("SHEETS:", wb.sheetnames)
for name in wb.sheetnames:
    ws = wb[name]
    print("=" * 60)
    print(f"[{name}] dims={ws.dimensions} rows={ws.max_row} cols={ws.max_column}")
    # header
    hdr = [c.value for c in ws[1]]
    print("HDR:", hdr)
