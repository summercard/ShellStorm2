"""Apply Artifact-authored wrap alignment to changed cells, preserving original styles."""
from pathlib import Path
import json,zipfile,xml.etree.ElementTree as E,copy,re
from openpyxl import load_workbook
D=Path('_scratch/monitor_spawn_ledger_20261007');jobs=json.loads((D/'changes.json').read_text(encoding='utf-8'));p=Path(jobs[0]['path'])
M='http://schemas.openxmlformats.org/spreadsheetml/2006/main';N={'m':M};E.register_namespace('',M)
def paths(z):
 rel={v.get('Id'):v.get('Target') for v in E.fromstring(z.read('xl/_rels/workbook.xml.rels'))}
 return {v.get('name'):'xl/'+rel[v.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')].lstrip('/').removeprefix('xl/') for v in E.fromstring(z.read('xl/workbook.xml')).find('m:sheets',N)}
with zipfile.ZipFile(p) as z:
 src={i.filename:z.read(i.filename) for i in z.infolist()};mapping=paths(z)
styles=E.fromstring(src['xl/styles.xml']);xfs=styles.find('m:cellXfs',N)
for i,job in enumerate(jobs):
 with zipfile.ZipFile(D/f'cells{i}.xlsx') as gen:
  generated_styles=E.fromstring(gen.read('xl/styles.xml')).find('m:cellXfs',N);gp=paths(gen)
  for sn,changes in job['sheets'].items():
   authored={c.get('r'):c for c in E.fromstring(gen.read(gp[sn])).findall('.//m:c',N)}
   text=src[mapping[sn]].decode()
   for addr in changes:
    match=re.search(r'<c\b[^>]*\br="'+addr+r'"[^>]*>',text);assert match
    old=E.fromstring(match[0][:-1]+'/>');xf=copy.deepcopy(xfs[int(old.get('s','0'))]);align=xf.find('m:alignment',N)
    if align is not None:xf.remove(align)
    generated=generated_styles[int(authored[addr].get('s','0'))].find('m:alignment',N);assert generated is not None
    xf.append(copy.deepcopy(generated));xf.set('applyAlignment','1');new_id=len(xfs);xfs.append(xf)
    tag=re.sub(r' s="\d+"','',match[0]);tag=tag[:-1]+f' s="{new_id}">';text=text[:match.start()]+tag+text[match.end():]
   if sn=='域变更日志':text=re.sub(r'<row\b[^>]*\br="52"[^>]*>', '<row r="52" ht="150" customHeight="1">',text)
   if sn=='3D-敌人':text=re.sub(r'<row\b[^>]*\br="10"[^>]*>', '<row r="10" ht="180" customHeight="1">',text)
   src[mapping[sn]]=text.encode()
xfs.set('count',str(len(xfs)));src['xl/styles.xml']=E.tostring(styles,encoding='utf-8',xml_declaration=True)
temp=D/'polished.xlsx'
with zipfile.ZipFile(temp,'w',zipfile.ZIP_DEFLATED) as z:
 for name,data in src.items():z.writestr(name,data)
before=load_workbook(p);after=load_workbook(temp)
for sn in before.sheetnames:
 for row in before[sn]:
  for c in row:assert c.value==after[sn][c.coordinate].value
p.write_bytes(temp.read_bytes());print('Preserved all values/formulas; changed-cell wrap and two row heights applied')
