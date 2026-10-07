from pathlib import Path
import sys,json
from openpyxl import load_workbook
R=Path.cwd();D=R/'_scratch/monitor_spawn_ledger_20261007';A='ENM-BOSS-MONITOR002-3D'
sys.path[:0]=[str(R/'scripts'),str(R/'tools/asset_pipeline')]
from ledger_registry import LedgerIndex
from split_asset_ledger import _row_digest,read_source_rows,sheet_digest,col_digest,CONTENT_COLUMNS
index=LedgerIndex.load(R);p=R/'assets/registry/ledger_split_baseline.json';b=json.loads(p.read_text(encoding='utf-8'));original=json.loads((D/'baseline_before.json').read_text(encoding='utf-8'))
w=load_workbook(index.path_for_category('敌人'));row=next(v for _,v in read_source_rows(w['资产主表']) if v[0]==A)
b['assets'][A]['v']=_row_digest(row)
if sys.argv[1]=='specialized':b['sheet_digests']['3D-敌人']=sheet_digest(w['3D-敌人'])
assert {k:v for k,v in b['assets'].items() if k!=A}=={k:v for k,v in original['assets'].items() if k!=A}
assert {k:v for k,v in b['sheet_digests'].items() if k!='3D-敌人'}=={k:v for k,v in original['sheet_digests'].items() if k!='3D-敌人'}
rows=[]
for domain in index.domains:
 for _,v in read_source_rows(load_workbook(domain.path)['资产主表']):
  assert _row_digest(v)==b['assets'][v[0]]['v'],('unrelated drift',v[0])
  rows.append((v[0],v))
rows.sort(key=lambda x:x[0]);b['column_digests']={str(c):col_digest(rows,c) for c in CONTENT_COLUMNS}
if sys.argv[1]=='specialized':b.setdefault('targeted_updates',[]).append({'date':'2026-10-07','asset_id':A,'version':'v031','scope':'Expedition01 spawn acceptance: existing master row20, prefab row10 and domain log52; no model or gameplay value changes.'})
p.write_text(json.dumps(b,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
print('Scoped baseline accepted',sys.argv[1],len(rows),'assets; unrelated row digests unchanged')
