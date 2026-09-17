import openpyxl
p = r"I:\工作项目\shellstrom2\ShellStorm2\assets\registry\ShellStorm2_美术资产台账_v001.xlsx"
wb = openpyxl.load_workbook(p, data_only=True)
ws = wb["3D-场景通用"]
for r in range(4, 9):
    vals = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
    vals = [f"{chr(64+c)}:{str(v)[:18]}" if v is not None else "" for c, v in enumerate(vals, 1)]
    print(f"r{r}:", " | ".join([v for v in vals if v]))
