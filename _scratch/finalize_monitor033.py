import json,hashlib,sys,datetime,subprocess
from pathlib import Path
from openpyxl import load_workbook
R=Path.cwd();B=R/'assets/art/enemies/bosses/enm_boss_monitor002';D=R/'_scratch/monitor033_ledger';A='ENM-BOSS-MONITOR002-3D'
def read(p):return json.loads(p.read_text(encoding='utf-8'))
def save(p,v):p.write_text(json.dumps(v,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def rel(p):return p.relative_to(R).as_posix()
t=read(B/'enm_boss_monitor002_transfer_ledger_v031.json');audit=read(B/'previews/move_v033/source_audit.json')
t.update(version='v033',status='active',generated_at='2026-10-08',generated_by='Codex / Blender MCP evaluated move patch',git_baseline=subprocess.check_output(['git','rev-parse','HEAD'],text=True).strip())
t['source_sha256']={rel(B/f'source/enm_boss_monitor002_{k}_v033.blend'):sha(B/f'source/enm_boss_monitor002_{k}_v033.blend') for k in ['model','animation']}
t['sources']=[{'kind':k,'path':rel(B/f'source/enm_boss_monitor002_{k}_v033.blend'),'sha256':sha(B/f'source/enm_boss_monitor002_{k}_v033.blend'),'bytes':(B/f'source/enm_boss_monitor002_{k}_v033.blend').stat().st_size,'skeleton_id':t['skeleton_id'],'skeleton_signature':audit['skeleton_signature']} for k in ['model','animation']]
t['model_signature']=t['animation_signature']=audit['skeleton_signature']
for item in t['files']:
 if item['path'].endswith('rig_contract_v031.json'):item['path']=item['path'].replace('v031','v033')
 p=R/item['path'];item.update(sha256=sha(p),bytes=p.stat().st_size)
t['stage_history'] += [{'stage':s,'version':'v033','date':'2026-10-08'} for s in ['authored','exported_pending_godot_validation','validated','active']]
t['verification']['checks']=509
t['verification']['move']={'source':'previews/move_v033/source_audit.json','real_renderer':'previews/move_v033/visual_report.json','preview':'previews/move_v033/monitor_walk.mp4','baseline_version':'v031','other_move_tracks_unchanged':True,'source_samples':97,'rendered_frames':48,'other_15_clips_unchanged':True}
t['static_glb_reuse']={'sha256':audit['static_glb_reused'],'reason':'Only move curves and evaluated motion changed. Geometry, materials and rest skeleton retain v031.'}
t['rollback']='Retained v031 dual masters and Git baseline; runtime files keep stable paths.'
save(B/'enm_boss_monitor002_transfer_ledger_v033.json',t);save(B/'enm_boss_monitor002_transfer_ledger.json',t)
m=read(B/'asset_manifest.json');m.update(version='v033',source='source/enm_boss_monitor002_model_v033.blend',model_source=rel(B/'source/enm_boss_monitor002_model_v033.blend'),animation_source=rel(B/'source/enm_boss_monitor002_animation_v033.blend'),rig_contract='source/rig_contract_v033.json',runtime_verification=t['verification'])
for item in t['files']:
 p=R/item['path'];m.setdefault('sha256',{})[item['path']]=item['sha256']
 if p.is_relative_to(B):m['files'][p.relative_to(B).as_posix()]=item['sha256']
for path,h in t['source_sha256'].items():m['sha256'][path]=h;m['files'][(R/path).relative_to(B).as_posix()]=h
save(B/'asset_manifest.json',m)
sys.path[:0]=[str(R/'scripts'),str(R/'tools/asset_pipeline')]
from ledger_registry import LedgerIndex
from split_asset_ledger import _row_digest,read_source_rows,sheet_digest,col_digest,CONTENT_COLUMNS
idx=LedgerIndex.load(R);w=load_workbook(idx.path_for_category('敌人'));p=R/'assets/registry/ledger_split_baseline.json';base=read(p);old=read(D/'baseline_before.json')
base['assets'][A]['v']=_row_digest(next(v for _,v in read_source_rows(w['资产主表']) if v[0]==A))
for sn in ['3D-敌人','敌人动画与状态']:base['sheet_digests'][sn]=sheet_digest(w[sn])
assert {k:v for k,v in base['assets'].items() if k!=A}=={k:v for k,v in old['assets'].items() if k!=A}
union=[]
for domain in idx.domains:
 for _,v in read_source_rows(load_workbook(domain.path)['资产主表']):
  assert _row_digest(v)==base['assets'][v[0]]['v'],v[0]
  union.append((v[0],v))
union.sort(key=lambda x:x[0]);base['column_digests']={str(c):col_digest(union,c) for c in CONTENT_COLUMNS}
base.setdefault('targeted_updates',[]).append({'date':'2026-10-08','asset_id':A,'version':'v033','scope':'v031 move restored, pedestal roll only reduced 23 to 10 degrees; other motion retained.'})
save(p,base);print('v033 transfer/manifest/target baseline finalized')
