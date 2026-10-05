from pathlib import Path
from collections import Counter
import hashlib, json, shutil
from copy import copy
from datetime import datetime
import openpyxl

ROOT=Path('I:/工作项目/shellstrom2/ShellStorm2')
BOOKS=ROOT/'assets/registry/ledgers'
PROP=BOOKS/'ShellStorm2_道具账本_v001.xlsx'
MASTER=ROOT/'assets/registry/ShellStorm2_美术资产台账_v001.xlsx'
BASE=ROOT/'assets/registry/ledger_split_baseline.json'
OUT=ROOT/'outputs/base99_radio_v001'
BACK=OUT/'backup_20261005'
ID='PRP-BASE99-RADIO-3D'

def txt(v): return '' if v is None else str(v).strip()
def row_digest(values):
    return hashlib.sha256('\x1f'.join(txt(values[c-1]) for c in range(1,26) if c not in (18,19)).encode()).hexdigest()
def col_digest(rows,col):
    ordered=sorted(rows,key=lambda item:(txt(item[1][0]),txt(item[1][col-1])))
    return hashlib.sha256('\x1e'.join(txt(v[col-1]) for _,v in ordered).encode()).hexdigest()
def rows(ws): return [(r,[ws.cell(r,c).value for c in range(1,26)]) for r in range(6,ws.max_row+1) if txt(ws.cell(r,1).value)]

def main():
    OUT.mkdir(parents=True,exist_ok=True); BACK.mkdir(parents=True,exist_ok=True)
    # 保留首次事务备份；若缺失才复制当前文件。
    for p in [PROP,MASTER,BASE]:
        b=BACK/p.name
        if not b.exists(): shutil.copy2(p,b)
    wb=openpyxl.load_workbook(PROP)
    ws=wb['资产主表']; found=[r for r in range(6,ws.max_row+1) if ws.cell(r,1).value==ID]; assert found==[27],found
    new_row=27
    assert ws.cell(new_row,3).value=='道具' and ws.cell(new_row,4).value=='decor_prop'
    for r in range(6,new_row+1):
        ws.cell(r,18).value=f'=LOWER(TRIM(C{r})&"|"&TRIM(D{r})&"|"&TRIM(E{r})&"|"&TRIM(F{r})&"|"&TRIM(H{r})&"|"&TRIM(I{r}))'
        ws.cell(r,19).value=f'=IF(COUNTIF($R$6:$R${new_row},R{r})>1,"重复","唯一")'
    ov=wb['总览']
    for ref in ('A6','C6','E6','G6','B10','C10','B11','C11'):
        assert f'$27' in txt(ov[ref].value),(ref,ov[ref].value)
    for dv in ws.data_validations.dataValidation:
        assert all(txt(rng).endswith('27') for rng in dv.sqref.ranges),(dv.type,[str(x) for x in dv.sqref.ranges])
    pf=wb['3D-道具']; assert pf.max_row==8 and pf.cell(8,1).value==ID
    log=wb['域变更日志']; assert log.max_row==11 and log.cell(11,1).value=='v0.1.5'
    wb.save(PROP)
    master=openpyxl.load_workbook(MASTER); mws=master['总览']; ctl=master['3D Prefab总控']
    # 账本索引表中，道具是第18行；Prefab总控道具为第7行。
    mws['D18']=21; ctl['C7']=4; ctl['E7']=2; ctl['F7']=2; ctl['G7']='桌椅、家具、收音机、容器和摆放装饰物'; master.save(MASTER)
    idx=json.loads((ROOT/'assets/registry/ledger_index.json').read_text(encoding='utf-8'))
    all_rows=[]
    for d in idx['domains']:
        x=openpyxl.load_workbook(BOOKS/d['file'],read_only=True); all_rows.extend((str(v[0]),v) for _,v in rows(x['资产主表']))
    assert ID in {a for a,_ in all_rows}
    base=json.loads(BASE.read_text(encoding='utf-8')); base['assets'][ID]={'v':row_digest(dict(all_rows)[ID]),'c':'道具','d':'props'}; base['asset_count']=len(base['assets']); base['captured_at']='2026-10-05'; base['column_digests']={str(c):col_digest(all_rows,c) for c in range(1,26) if c not in (18,19)}; base['category_counts']=dict(sorted(Counter(txt(v[2]) for _,v in all_rows).items())); BASE.write_text(json.dumps(base,ensure_ascii=False,indent=1)+'\n',encoding='utf-8')
    evidence={'asset_id':ID,'ledger':str(PROP),'asset_sheet_row':27,'category':'道具','subcategory':'decor_prop','prefab_sheet':'3D-道具','prefab_sheet_row':8,'log_sheet_row':11,'status':'正式美术已接入','model_gate_status':'未全过：StatusLight 命名误报保留；其余模型验收记录通过，不将模型门禁写成全过。','backup_dir':str(BACK),'backup_files':[str(BACK/p.name) for p in [PROP,MASTER,BASE]],'baseline':{'path':str(BASE),'asset_count':base['asset_count'],'captured_at':base['captured_at']},'runtime_contract':{'state_cycle':['off','music_a','music_b','off'],'bus':'Music','input':['鼠标左键','E']},'paths':{'source_blend':'assets/art/props/base_world_3d/source/base99_radio/prp_base99_radio_source_v001.blend','glb':'assets/art/props/base_world_3d/components/base99_radio/prp_base99_radio_visual_top3d.glb','stable_prefab':'assets/art/props/base_world_3d/runtime/base99_radio/prp_base99_radio_root_top3d.tscn','script':'src/base3d/Base99Radio3D.gd','formal_scene':'assets/art/environments/base_facility_3d/runtime/env_base_facility_art_layout_top3d.tscn'}}
    (OUT/'registration_evidence.json').write_text(json.dumps(evidence,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    (OUT/'registration_evidence.txt').write_text('\n'.join([f'AssetID: {ID}',f'账本: {PROP}','资产主表行: 27','分类: 道具 / decor_prop / 3D-道具','Prefab分页行: 8','域变更日志行: 11','正式状态: 正式美术已接入','模型门禁: 未全过；StatusLight 命名误报保留，未宣称全过',f'备份目录: {BACK}','备份文件: ShellStorm2_道具账本_v001.xlsx, ShellStorm2_美术资产台账_v001.xlsx, ledger_split_baseline.json',f'无损基线: {BASE}'])+'\n',encoding='utf-8')
    print('FINISHED',json.dumps(evidence,ensure_ascii=False))
try:
    main()
except Exception:
    import traceback
    Path('I:/工作项目/shellstrom2/ShellStorm2/_scratch/finish_trace.txt').write_text(traceback.format_exc(), encoding='utf-8')
    raise
