from pathlib import Path
import json
import openpyxl
root=Path('I:/工作项目/shellstrom2/ShellStorm2')
p=root/'assets/registry/ledgers/ShellStorm2_道具账本_v001.xlsx'
w=openpyxl.load_workbook(p,data_only=False)
s=w['资产主表']
out={'asset_rows':[(r,[s.cell(r,c).value for c in range(1,26)]) for r in range(6,s.max_row+1) if s.cell(r,1).value=='PRP-BASE99-RADIO-3D'],'asset_max_row':s.max_row,'prefab_max_row':w['3D-道具'].max_row,'log_max_row':w['域变更日志'].max_row}
m=openpyxl.load_workbook(root/'assets/registry/ShellStorm2_美术资产台账_v001.xlsx',data_only=False)
out['master_overview']={k:m['总览'][k].value for k in ['D17','D18']}
q=m['3D Prefab总控']; out['prefab_control']={k:q[k].value for k in ['C7','E7','F7','G7']}
b=json.loads((root/'assets/registry/ledger_split_baseline.json').read_text(encoding='utf-8')); out['baseline']={'asset_count':b['asset_count'],'captured_at':b['captured_at'],'asset':b['assets'].get('PRP-BASE99-RADIO-3D')}
out['evidence_exists']=(root/'outputs/base99_radio_v001/registration_evidence.json').exists()
(root/'_scratch/check_radio_state.json').write_text(json.dumps(out,ensure_ascii=False,indent=2,default=str)+'\n',encoding='utf-8')
