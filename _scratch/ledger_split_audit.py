"""只读审计：重放 verify_ledger_split.py 的行级检查，输出全部红项（不被 [:40] 截断）。

verify_ledger_split.py 的 JSON 只保留 failures[:40]，看不到「门墙行到底有没有被
判 row_content_mutated」。本脚本按同样逻辑重放一遍，按域逐行检查并落全量结果。
"""

import json
import sys
from collections import Counter, defaultdict
from pathlib import Path

PROJECT = Path(r"I:\工作项目\shellstrom2\ShellStorm2")
sys.path.insert(0, str(PROJECT / "tools" / "asset_pipeline"))
sys.path.insert(0, str(PROJECT / "scripts"))

from openpyxl import load_workbook  # noqa: E402

from ledger_registry import LedgerIndex  # noqa: E402
from split_asset_ledger import (  # noqa: E402
    ASSET_SHEET,
    FIRST_DATA_ROW,
    _row_digest,
    _text,
    dedupe_key_formula,
    dedupe_result_formula,
    read_source_rows,
)

TARGET = "ENV-TOWER-WALL-DOOR-5M"

index = LedgerIndex.load(PROJECT)
baseline = json.loads((PROJECT / "assets/registry/ledger_split_baseline.json").read_text(encoding="utf-8"))
expected_assets = baseline["assets"]

failures = []


def fail(kind, **payload):
    failures.append({"kind": kind, **payload})


for domain in index.domains:
    if not domain.path.is_file():
        fail("ledger_missing", domain=domain.key)
        continue
    wb = load_workbook(domain.path)
    if ASSET_SHEET not in wb.sheetnames:
        fail("asset_sheet_missing", domain=domain.key)
        continue
    ws = wb[ASSET_SHEET]
    rows = read_source_rows(ws)
    expected_end = FIRST_DATA_ROW - 1 + len(rows)
    print("domain %-12s path=%-52s rows=%d expected_end=%s" % (
        domain.key, domain.relative_path, len(rows), expected_end))

    for row_number, values in rows:
        asset_id = _text(values[0])
        category = _text(values[2])
        expected = expected_assets.get(asset_id)
        if expected is None:
            fail("asset_not_in_baseline", asset_id=asset_id, domain=domain.key)
        else:
            if expected["c"] != category:
                fail("category_mutated", asset_id=asset_id, domain=domain.key)
            if expected["d"] != domain.key:
                fail("asset_in_wrong_ledger", asset_id=asset_id, domain=domain.key)
            digest_now = _row_digest(values)
            if expected["v"] != digest_now:
                fail("row_content_mutated", asset_id=asset_id, domain=domain.key,
                     baseline_v=expected["v"], actual_v=digest_now,
                     row=list(_text(v) for v in values))

        dedupe_key = _text(ws.cell(row=row_number, column=18).value)
        dedupe_result = _text(ws.cell(row=row_number, column=19).value)
        if dedupe_key != dedupe_key_formula(row_number):
            fail("dedupe_key_formula_wrong", asset_id=asset_id, domain=domain.key,
                 row=row_number, actual=dedupe_key[:160],
                 wanted=dedupe_key_formula(row_number)[:160])
        if dedupe_result != dedupe_result_formula(row_number, expected_end):
            fail("dedupe_result_formula_wrong", asset_id=asset_id, domain=domain.key,
                 row=row_number, actual=type(ws.cell(row=row_number, column=19).value).__name__)

    # 公式实际用到的查重范围（取首末两行的 S 列，看它引用到哪一行）
    if rows:
        first_row, last_row = rows[0][0], rows[-1][0]
        for probe in (first_row, last_row):
            raw = ws.cell(row=probe, column=19).value
            inner = getattr(raw, "text", raw)
            print("   S%-4s -> %s" % (probe, str(inner)[:150]))

counts = Counter(f["kind"] for f in failures)
print("\n=== failure kinds (ALL) ===")
for kind, n in counts.most_common():
    print("%-32s %d" % (kind, n))
print("total = %d" % len(failures))

print("\n=== row_content_mutated (%d) ===" % counts.get("row_content_mutated", 0))
for f in failures:
    if f["kind"] == "row_content_mutated":
        print("  %-34s domain=%-10s base=%s now=%s" % (
            f["asset_id"], f["domain"], f["baseline_v"][:10], f["actual_v"][:10]))

print("\n=== TARGET row ===")
for f in failures:
    if f.get("asset_id") == TARGET:
        print(json.dumps(f, ensure_ascii=False)[:600])

print("\n=== dedupe wrong: domain histogram ===")
hist = defaultdict(int)
for f in failures:
    if f["kind"].startswith("dedupe"):
        hist[(f["kind"], f.get("domain"))] += 1
for key, n in sorted(hist.items()):
    print("  %-32s %-10s %d" % (key[0], key[1], n))

(PROJECT / "_scratch/ledger_split_audit.json").write_text(
    json.dumps({"failures": failures, "counts": dict(counts)}, ensure_ascii=False, indent=2) + "\n",
    encoding="utf-8",
)
print("\nwrote _scratch/ledger_split_audit.json")
