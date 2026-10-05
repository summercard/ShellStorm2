from pathlib import Path
import openpyxl, json
root=Path('I:/工作项目/shellstrom2/ShellStorm2')
all_out=[]
for path in [root/'assets/registry/ledgers/ShellStorm2_道具账本_v001.xlsx', root/'assets/registry/ShellStorm2_美术资产台账_v001.xlsx']:
 wb=openpyxl.load_workbook(path,data_only=False)
 out={'path':str(path)}
 for sn in ['总览','3D Prefab总控']:
  if sn in wb.sheetnames:
   ws=wb[sn]; out[sn]={'cells':{ref:ws[ref].value for ref in ['A6','B6','C6','E6','G6','B10','C10','B11','C11','A7','B7','C7','E7','G7','B12','C12']},'max_row':ws.max_row}
 all_out.append(out)
(root/'_scratch/prop_cells.json').write_text(json.dumps(all_out,ensure_ascii=False,indent=2,default=str)+'\n',encoding='utf-8')
