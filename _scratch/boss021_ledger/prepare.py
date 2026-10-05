import sys,json,hashlib
from pathlib import Path
from openpyxl import load_workbook
R=Path.cwd();sys.path.insert(0,str(R/'scripts'));from ledger_registry import LedgerIndex
ix=LedgerIndex.load(R);p=ix.path_for_category('敌人');w=load_workbook(p);s=w['资产主表'];row=next(r for r in range(1,s.max_row+1) if s.cell(r,1).value=='ENM-BOSS-MONITOR002-3D')
master=load_workbook(R/'assets/registry/ShellStorm2_美术资产台账_v001.xlsx',read_only=True)
for name in ['分类与编码','命名与查重']:
 print(name,[[str(v)[:100] for v in row if v is not None] for row in master[name].iter_rows(values_only=True) if any(v and ('敌人' in str(v) or 'ENM' in str(v)) for v in row)])
B=R/'assets/art/enemies/bosses/enm_boss_monitor002';src=B/'source/enm_boss_monitor002_model_v021.blend';note='v021 源级制作：砸地第65帧增加冰白高亮锯齿爆裂色块及瞬时地面补光，仅持续1帧；保留v020无勾线赛博脉冲；四个角色动作曲线与v017完全相同。保留idle/move。特效为独立预览集合；未导出/未接入Godot。'
changes={f'I{row}':'T Pose / 六表情 / idle / move / melee_keyboard / heavy_spin_slam 3.2秒单次',f'M{row}':'v021',f'O{row}':src.relative_to(R).as_posix(),f'T{row}':hashlib.sha256(src.read_bytes()).hexdigest(),f'P{row}':'source/enm_boss_monitor002_model_v021.blend + enm_boss_monitor002_animation_v021.blend；SKEL-MONITOR002-005',f'V{row}':'2026-10-05',f'Y{row}':note}
logrow=w['域变更日志'].max_row+1
jobs=[{'path':str(p),'sheets':{'资产主表':changes,'域变更日志':{f'{col}{logrow}':value for col,value in zip('ABCDE',['v021','2026-10-05','Codex','ENM-BOSS-MONITOR002-3D',note])}}}]
p=B/'source/boss002_production_ledger.xlsx';w=load_workbook(p);s=w.active;changes={'B4':'v021 重击单帧高亮与锯齿打击形状'}
for r in range(1,s.max_row+1):
 label=s.cell(r,1).value
 if label=='当前源版本':changes.update({f'B{r}':'v021',f'C{r}':'model_v021 + animation_v021；v012/v013保留回退'})
 if label=='绑定验证':changes.update({f'B{r}':'previews/heavy_v021/style_audit.json',f'C{r}':note})
 if label=='面数统计':changes[f'C{r}']='source/rig_contract_v021.json'
jobs.append({'path':str(p),'sheets':{s.title:changes}})
Path('_scratch/boss021_ledger/changes.json').write_text(json.dumps(jobs,ensure_ascii=False),encoding='utf-8')
