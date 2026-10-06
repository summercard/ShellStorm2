import copy
import hashlib
import json
import shutil
import sys
from datetime import datetime
from pathlib import Path
from openpyxl import load_workbook

P = Path('I:/工作项目/shellstrom2/ShellStorm2')
sys.path.insert(0, str(P / 'scripts'))
sys.path.insert(0, str(P / 'tools/asset_pipeline'))
from ledger_registry import LedgerIndex
from split_asset_ledger import read_source_rows, _row_digest, col_digest, CONTENT_COLUMNS, dedupe_key_formula, dedupe_result_formula, sheet_digest

OUT = P / 'outputs/base99_radio_v003'
ID = 'PRP-BASE99-RADIO-3D'
index = LedgerIndex.load(P)
ledger = index.path_for_category('道具')
baseline = P / 'assets/registry/ledger_split_baseline.json'
backup = OUT / 'backup_before_registration'
assert not backup.exists(), '已有事务备份，禁止重跑覆盖'
backup.mkdir()
shutil.copy2(ledger, backup / ledger.name)
shutil.copy2(baseline, backup / baseline.name)

def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()

def cell_snapshot(w):
    return {(ws.title, c.coordinate): c.value for ws in w for row in ws for c in row if c.value is not None}

wb = load_workbook(ledger)
before = cell_snapshot(wb)
ws = wb['资产主表']
ps = wb['3D-道具']
assert ws['A27'].value == ID and ps['A8'].value == ID
assert ps['O4'].value == '版本' and ps['P4'].value == '备注'
assert ws['R27'].value == dedupe_key_formula(27)
assert ws['S27'].value == dedupe_result_formula(27, ws.max_row)
dv_before = [(str(d.sqref), d.formula1) for d in ws.data_validations.dataValidation]
other_rows_before = {v[0]: _row_digest(v) for _, v in read_source_rows(ws) if v[0] != ID}
opt = json.loads((OUT / 'optimization_evidence.json').read_text(encoding='utf-8'))
glb = P / 'assets/art/props/base_world_3d/components/base99_radio/prp_base99_radio_visual_top3d.glb'
prefab = P / 'assets/art/props/base_world_3d/runtime/base99_radio/prp_base99_radio_root_top3d.tscn'
source_rel = 'assets/art/props/base_world_3d/source/base99_radio/prp_base99_radio_source_v003.blend'
assert sha(glb) == opt['glb_sha256']
assert 'BASE99_RADIO_OK' in (OUT / 'radio_window_input_final.log').read_text(encoding='utf-8')
assert 'PLACEMENT_ACCEPTED=true' in (OUT / 'placement_acceptance_final.log').read_text(encoding='utf-8')
assert json.loads((OUT / 'validate_game_prop.json').read_text(encoding='utf-8'))['passed']
spec = '宽0.828m × 深0.456m × 高0.822m；v002各轴2倍，几何烘焙，根scale=1；598面/1116三角形；深墨绿哑光机身/深棕铜饰面；仅StatusLight在A/B发光、off零发光；公共色盘外链、PaletteUV/Closest；WorldCollision/InteractionHitArea分离；左键/E关→A→B→关'
changes = {'M27': 'v003', 'N27': spec, 'P27': source_rel + '; ' + str(glb.relative_to(P)).replace('\\', '/') + '; src/base3d/Base99Radio3D.gd; outputs/base99_radio_v003/optimization_evidence.json', 'T27': sha(prefab), 'V27': datetime(2026, 10, 6), 'Y27': 'v003：只移动收音机至46号BATTERY模块收纳箱顶部；layout=(-1.95,6.97,-13.87143)，Tower world=(-1.95,-5.03,-8.87143)，yaw=10度；柜顶121/121采样承载、无其他可见美术三角形穿插；带窗口鼠标/E与楼层回归通过。源/优化/GLB哈希见optimization_evidence.json；GLB SHA-256=' + sha(glb)}
for cell, value in changes.items():
    ws[cell] = value
ws['R27'] = dedupe_key_formula(27)
ws['S27'] = dedupe_result_formula(27, ws.max_row)
ps_changes = {'E8': source_rel, 'K8': spec, 'L8': '底面Y=0；根ItemRoot居中；local -Z正面；scale=1', 'M8': '99F阁楼北墙/床左侧46号BATTERY柜顶；layout=(-1.95,6.97,-13.87143)，yaw=10度', 'O8': 'v003', 'P8': 'ItemRoot/Visual/StatusLight稳定接口；GLB SHA-256=' + sha(glb) + '；独立optimized派生与验收见outputs/base99_radio_v003/optimization_evidence.json', 'Q8': None}
for cell, value in ps_changes.items():
    ps[cell] = value
log = wb['域变更日志']
row = log.max_row + 1
for c in range(1, 8):
    log.cell(row, c)._style = copy.copy(log.cell(row - 1, c)._style)
for c, value in enumerate(['v0.1.7', '2026-10-06', '99F阁楼收音机v003尺寸与材质摆位', '道具 / 场景可交互道具', '更新radio到v003：各轴2倍，598面/1116三角形；深墨绿/铜棕；46号BATTERY柜顶10度斜摆；保留v001/v002。', '稳定路径、off/A/B交互不变；同步碰撞/音源/Tooltip；真实运行时承载与输入验收；更新仅radio指纹和3D-道具摘要。', 'WorkBuddy'], 1):
    log.cell(row, c, value)
wb.save(ledger)
reopened = load_workbook(ledger)
after = cell_snapshot(reopened)
diffs = [{'sheet': s, 'cell': c, 'before': str(before.get((s, c))), 'after': str(after.get((s, c)))} for s, c in sorted(set(before) | set(after)) if before.get((s, c)) != after.get((s, c))]
allowed = {('资产主表', c) for c in changes} | {('3D-道具', c) for c in ps_changes} | {('域变更日志', reopened['域变更日志'].cell(row, c).coordinate) for c in range(1,8)}
assert all((d['sheet'], d['cell']) in allowed for d in diffs), diffs
assert [(str(d.sqref), d.formula1) for d in reopened['资产主表'].data_validations.dataValidation] == dv_before
assert other_rows_before == {v[0]: _row_digest(v) for _, v in read_source_rows(reopened['资产主表']) if v[0] != ID}
bl = json.loads(baseline.read_text(encoding='utf-8'))
old_bl = copy.deepcopy(bl)
values = next(v for _, v in read_source_rows(reopened['资产主表']) if v[0] == ID)
bl['assets'][ID] = {'v': _row_digest(values), 'c': '道具', 'd': 'props'}
all_rows = []
for domain in index.domains:
    dw = load_workbook(domain.path)
    all_rows.extend(read_source_rows(dw['资产主表']))
bl['column_digests'] = {str(c): col_digest(all_rows, c) for c in CONTENT_COLUMNS}
bl['category_counts'] = {cat: sum(1 for _, v in all_rows if v[2] == cat) for cat in sorted({v[2] for _, v in all_rows})}
bl['sheet_digests']['3D-道具'] = sheet_digest(reopened['3D-道具'])
assert all(bl['assets'][k] == v for k, v in old_bl['assets'].items() if k != ID)
assert bl['asset_count'] == old_bl['asset_count']
baseline.write_text(json.dumps(bl, ensure_ascii=False, indent=1) + '\n', encoding='utf-8')
evidence = {'asset_id': ID, 'ledger': str(ledger), 'changed_cells': diffs, 'other_props_row_digests_unchanged': True, 'data_validations_unchanged': True, 'other_baseline_asset_fingerprints_unchanged': True, 'asset_count': bl['asset_count'], 'old_radio_fingerprint': old_bl['assets'][ID], 'new_radio_fingerprint': bl['assets'][ID], 'old_3d_props_sheet_digest': old_bl['sheet_digests']['3D-道具'], 'new_3d_props_sheet_digest': bl['sheet_digests']['3D-道具'], 'backup': str(backup), 'ledger_sha256': sha(ledger), 'prefab_sha256': sha(prefab)}
(OUT / 'registration_evidence.json').write_text(json.dumps(evidence, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
print('RADIO_V003_LEDGER_TRANSACTION_OK', len(diffs))
