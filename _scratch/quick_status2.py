import openpyxl, collections
p = r"I:\工作项目\shellstrom2\ShellStorm2\assets\registry\ShellStorm2_美术资产台账_v001.xlsx"
wb = openpyxl.load_workbook(p, data_only=True)
for name in ["3D-场景通用", "3D-设施", "3D-道具", "3D-角色", "3D-敌人", "3D-武器", "3D-物品", "3D-特效", "3D-其他"]:
    ws = wb[name]
    print("=" * 70)
    print(f"[{name}] rows={ws.max_row}")
    for r in range(1, 4):
        vals = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
        vals = [str(v)[:14] if v is not None else "" for v in vals]
        print(f"  r{r}:", " | ".join(vals))
