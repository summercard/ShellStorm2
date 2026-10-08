from pathlib import Path
import json,hashlib,sys
from openpyxl import load_workbook
R=Path.cwd();B=R/'assets/art/enemies/bosses/enm_boss_monitor002';D=R/'_scratch/monitor035_ledger';A='ENM-BOSS-MONITOR002-3D'
def read(p):return json.loads(p.read_text(encoding='utf-8'))
def save(p,x):p.write_text(json.dumps(x,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
t=read(B/'enm_boss_monitor002_transfer_ledger.json');t.update(version='v035',generated_by='Codex / Blender MCP screen clearance only')
t['source_sha256']={}
for source in t['sources']:
 source['path']=source['path'].replace('_v034.blend','_v035.blend');p=R/source['path'];source['sha256']=sha(p);source['bytes']=p.stat().st_size;t['source_sha256'][source['path']]=source['sha256']
for item in t['files']:
 item['path']=item['path'].replace('rig_contract_v034.json','rig_contract_v035.json');p=R/item['path'];item.update(sha256=sha(p),bytes=p.stat().st_size)
p=R/'scripts/blender/offset_monitor_screen_v035.py';t['files'].append({'path':p.relative_to(R).as_posix(),'sha256':sha(p),'bytes':p.stat().st_size})
t['verification']['screen_clearance']=read(B/'previews/screen_v035/source_audit.json');t['verification']['screen_clearance']['preview']='previews/screen_v035/monitor_activation.mp4'
t['stage_history'] += [{'stage':s,'version':'v035','date':'2026-10-08'} for s in ['authored','exported_pending_godot_validation','validated','active']]
save(B/'enm_boss_monitor002_transfer_ledger_v035.json',t);save(B/'enm_boss_monitor002_transfer_ledger.json',t)
m=read(B/'asset_manifest.json');m.update(version='v035',source='source/enm_boss_monitor002_model_v035.blend',model_source=t['sources'][0]['path'],animation_source=t['sources'][1]['path'],rig_contract='source/rig_contract_v035.json',runtime_verification=t['verification'])
for item in t['files']:
 p=R/item['path'];m.setdefault('sha256',{})[item['path']]=item['sha256']
 if p.is_relative_to(B):m['files'][p.relative_to(B).as_posix()]=item['sha256']
for path,h in t['source_sha256'].items():m['sha256'][path]=h;m['files'][(R/path).relative_to(B).as_posix()]=h
save(B/'asset_manifest.json',m)
sys.path[:0]=[str(R/'scripts'),str(R/'tools/asset_pipeline')]
from ledger_registry import LedgerIndex
from split_asset_ledger import _row_digest,read_source_rows,sheet_digest,col_digest,CONTENT_COLUMNS
idx=LedgerIndex.load(R);w=load_workbook(idx.path_for_category('敌人'));p=R/'assets/registry/ledger_split_baseline.json';base=read(p);old=read(D/'baseline_before.json')
base['assets'][A]['v']=_row_digest(next(v for _,v in read_source_rows(w['资产主表']) if v[0]==A))
for sn in ['3D-敌人','敌人动画与状态']:base['sheet_digests'][sn]=sheet_digest(w[sn])
assert {k:v for k,v in base['assets'].items() if k!=A}=={k:v for k,v in old['assets'].items() if k!=A}
union=[]
for domain in idx.domains:
 for _,v in read_source_rows(load_workbook(domain.path)['资产主表']):
  assert _row_digest(v)==base['assets'][v[0]]['v'],v[0]
  union.append((v[0],v))
union.sort(key=lambda x:x[0]);base['column_digests']={str(c):col_digest(union,c) for c in CONTENT_COLUMNS};base.setdefault('targeted_updates',[]).append({'date':'2026-10-08','asset_id':A,'version':'v035','scope':'Screen assembly forward clearance only'})
save(p,base);print('V035_FINALIZED')
