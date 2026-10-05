from pathlib import Path
from collections import Counter
import hashlib, json
import openpyxl

ROOT = Path(r'I:/工作项目/shellstrom2/ShellStorm2')
BOOKS = ROOT / 'assets/registry/ledgers'
PROP = BOOKS / 'ShellStorm2_道具账本_v001.xlsx'
BASE = ROOT / 'assets/registry/ledger_split_baseline.json'
INDEX = ROOT / 'assets/registry/ledger_index.json'
GLB = ROOT / 'assets/art/props/base_world_3d/components/base99_radio/prp_base99_radio_visual_top3d.glb'
SOURCE = ROOT / 'assets/art/props/base_world_3d/source/base99_radio/prp_base99_radio_source_v002.blend'
PREFAB = ROOT / 'assets/art/props/base_world_3d/runtime/base99_radio/prp_base99_radio_root_top3d.tscn'
OUT = ROOT / 'outputs/base99_radio_v002'
AID = 'PRP-BASE99-RADIO-3D'

def txt(v): return '' if v is None else str(v).strip()
def sha(p): return hashlib.sha256(p.read_bytes()).hexdigest()
def row_digest(values): return hashlib.sha256('\x1f'.join(txt(values[c-1]) for c in range(1,26) if c not in (18,19)).encode()).hexdigest()
def col_digest(rows, col):
    ordered = sorted(rows, key=lambda item: (txt(item[1][0]), txt(item[1][col-1])) )
    return hashlib.sha256('\x1e'.join(txt(v[col-1]) for _,v in ordered).encode()).hexdigest()

def read_rows(ws):
    return [(row, list(values)) for row, values in enumerate(ws.iter_rows(min_row=6, max_col=25, values_only=True), 6) if txt(values[0])]

wb = openpyxl.load_workbook(PROP, data_only=False)
pf = wb['3D-道具']
print('prefab_headers', [pf.cell(4,c).value for c in range(1,pf.max_column+1)])
print('prefab_before', [pf.cell(8,c).value for c in range(1,pf.max_column+1)])
# P is version in this sheet; Q is notes/hash.
pf.cell(8,16).value = 'v002'
pf.cell(8,17).value = '根节点=ItemRoot；Visual+StatusLight稳定接口；GLB SHA-256=' + sha(GLB)
wb.save(PROP)
wb.close()

# Rebuild only the machine baseline summaries from all domain ledgers without touching specialized sheets.
idx = json.loads(INDEX.read_text(encoding='utf-8'))
all_rows = []
for domain in idx['domains']:
    path = BOOKS / domain['file']
    print('reading', path.name)
    book = openpyxl.load_workbook(path, read_only=True, data_only=False)
    rows = read_rows(book['资产主表'])
    book.close()
    all_rows.extend((txt(values[0]), values) for _, values in rows)
lookup = dict(all_rows)
assert AID in lookup
base = json.loads(BASE.read_text(encoding='utf-8'))
base['assets'][AID] = {'v': row_digest(lookup[AID]), 'c': '道具', 'd': 'props'}
base['asset_count'] = len(base['assets'])
base['captured_at'] = '2026-10-05'
base['column_digests'] = {str(c): col_digest(all_rows, c) for c in range(1,26) if c not in (18,19)}
base['category_counts'] = dict(sorted(Counter(txt(values[2]) for _,values in all_rows).items()))
BASE.write_text(json.dumps(base, ensure_ascii=False, indent=1) + '\n', encoding='utf-8')

manifest = json.loads((OUT/'asset_manifest.json').read_text(encoding='utf-8'))
manifest['hashes'] = {'glb_sha256': sha(GLB), 'source_blend_sha256': sha(SOURCE), 'runtime_prefab_sha256': sha(PREFAB)}
manifest['runtime_prefab_version'] = 'v002'
manifest['ledger'] = {'file':'assets/registry/ledgers/ShellStorm2_道具账本_v001.xlsx', 'asset_sheet_row':27, 'prefab_sheet':'3D-道具', 'prefab_sheet_row':8, 'version_column':'P', 'change_log_row':12}
(OUT/'asset_manifest.json').write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
print(json.dumps({'glb_sha256':sha(GLB),'source_sha256':sha(SOURCE),'prefab_sha256':sha(PREFAB),'baseline_assets':base['asset_count'],'asset_row':27,'prefab_row':8,'change_log_row':12},ensure_ascii=False,indent=2))
