import openpyxl
p = "assets/registry/ledgers/ShellStorm2_场景账本_v001.xlsx"
wb = openpyxl.load_workbook(p, data_only=False)
ws = wb["资产主表"]
for r in range(1, 7):
    vals = [ws.cell(r, c).value for c in range(1, ws.max_column+1)]
    print("ROW%03d" % r, vals)
print()
# find rows mentioning rooftop / parapet / 天台 / wall_solid / door
for r in range(1, ws.max_row+1):
    joined = " | ".join(str(ws.cell(r,c).value) for c in range(1,8))
    if any(k in joined for k in ("天台","女儿墙","parapet","ROOFTOP","wall_solid","wall_door")):
        print("HIT%03d" % r, joined[:260])
