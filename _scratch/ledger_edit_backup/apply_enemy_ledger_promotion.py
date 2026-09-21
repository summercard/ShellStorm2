"""把 ENM-MELEE-FUNGBOAR01（小僵尸）从 2D 口径「程序占位」转正为 3D 口径「已完成」。

四项改动（不做第五项）：
  1) 《资产主表》r6：2D 口径 -> 3D 口径（照精英行 r17 的字段形态）
  2) 《敌人动画与状态》r6:r17：12 态补 动作来源(5)/循环(6)/挂点(7)/首版实现(8)
  3) 《敌人动画与状态》r21：本条敌人 x 状态机的条目行
  4) 《3D-敌人》新增 r7 Prefab 行；《域变更日志》新增 r7

纪律：
  - 只升级既有行，不新增 AssetID（r6-r12 已预登记）。
  - 不碰派生列 R(18) / S(19) 的公式。
  - 写前断言工作表名、行号、表头逐字一致，防止静默写错位置。
  - 写后重跑 scripts/check_asset_registry.py --ledger 敌人 必须绿。

先决条件：runtime tscn 必须已是最终形态（本脚本读它的 sha256 写进主表 O 列）。
"""
from copy import copy
from datetime import datetime
from hashlib import sha256
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
LEDGER = ROOT / "assets/registry/ledgers/ShellStorm2_敌人账本_v001.xlsx"
BASE = "assets/art/enemies/normal_enemy_3d/melee_chaser"
TSCN = f"{BASE}/runtime/enm_melee_fungboar01_root_top3d.tscn"
GLB = f"{BASE}/components/enm_melee_fungboar01_visual_top3d.glb"
MODEL_SRC = f"{BASE}/source/model/enm_melee_fungboar01_model_v002.blend"
ANIM_SRC = f"{BASE}/source/animation/enm_melee_fungboar01_animation_v003.blend"
DISPLAY_NAME = "小僵尸"          # README 记录：用户确认造型后「显示名改为小僵尸」，AssetID/逻辑 ID 不变
LEGACY_NAME = "小菌猪"           # 原登记名，保留为检索别名

assert (ROOT / TSCN).is_file(), f"runtime tscn missing: {TSCN}"
assert (ROOT / GLB).is_file(), f"visual glb missing: {GLB}"
assert (ROOT / MODEL_SRC).is_file(), f"model source missing: {MODEL_SRC}"
assert (ROOT / ANIM_SRC).is_file(), f"animation source missing: {ANIM_SRC}"
SHA = sha256((ROOT / TSCN).read_bytes()).hexdigest()

wb = load_workbook(LEDGER)

# ---- 0) 先决断言：表名 / 表头 / 行号 -------------------------------------------------
for name in ("资产主表", "敌人动画与状态", "3D-敌人", "域变更日志", "总览"):
    assert name in wb.sheetnames, f"sheet missing: {name}"

master_ws = wb["资产主表"]
assert master_ws.cell(5, 1).value == "AssetID", master_ws.cell(5, 1).value
assert master_ws.cell(6, 1).value == "ENM-MELEE-FUNGBOAR01", master_ws.cell(6, 1).value
assert master_ws.cell(6, 2).value == LEGACY_NAME, master_ws.cell(6, 2).value
assert master_ws.cell(6, 11).value == "程序占位", master_ws.cell(6, 11).value
assert str(master_ws.cell(6, 18).value).startswith("=LOWER(TRIM(C6)"), master_ws.cell(6, 18).value

state_ws = wb["敌人动画与状态"]
assert state_ws.cell(5, 2).value == "状态ID", state_ws.cell(5, 2).value
assert state_ws.cell(6, 2).value == "dormant", state_ws.cell(6, 2).value
assert state_ws.cell(17, 2).value == "dead", state_ws.cell(17, 2).value
assert state_ws.cell(21, 2).value == "ENM-MELEE-FUNGBOAR01", state_ws.cell(21, 2).value

prefab_ws = wb["3D-敌人"]
assert prefab_ws.cell(4, 1).value == "AssetID", prefab_ws.cell(4, 1).value
assert prefab_ws.max_row == 6, prefab_ws.max_row
assert prefab_ws.cell(6, 1).value == "ENM-ELITE-RIFT-BOAR-ARMED-3D", prefab_ws.cell(6, 1).value

log_ws = wb["域变更日志"]
assert log_ws.cell(5, 1).value == "台账版本", log_ws.cell(5, 1).value
assert log_ws.cell(6, 1).value == "v0.1.0", log_ws.cell(6, 1).value

# ---- 1) 《资产主表》r6：2D 口径 -> 3D 口径（照 r17 精英行） -------------------------
master = {
    2: DISPLAY_NAME,                                                        # 中文名（显示名）
    6: "root_3d",                                                           # 组件槽
    7: "ENM-ECOSYSTEM-KIT-3D",                                              # 变体父 ID
    8: "Top3D / local -Z 正面",                                             # 视角
    9: "default / chase",                                                   # 状态/动画
    10: "全部战斗关卡普通怪池",                                              # 复用范围
    11: "已完成",                                                           # 制作状态
    13: "v003",                                                             # 版本（取动作母版最新）
    14: "1.2819×0.9535×1.8571m；1 Mesh；1 材质；36 骨；6 剪辑",              # 规格
    15: TSCN,                                                               # 文件路径 = runtime tscn
    16: f"{MODEL_SRC}; {ANIM_SRC}",                                         # 源码依据（模型 + 动作双 blend）
    17: f"近战追击; chaser; {DISPLAY_NAME}; {LEGACY_NAME}（旧登记名）; 普通怪标准样板; ENM-MELEE-FUNGBOAR01",
    20: SHA,                                                                # SHA-256（= runtime tscn）
    21: "Codex",
    22: datetime(2026, 9, 20),
    23: "AI 生成（Tripo）+ Blender 重制",
    25: ("GLB 仅表现；runtime PackedScene 无碰撞，Enemy3D 继续持有 CylinderShape3D、AI、生命、"
         "伤害与掉落。EnemyAvatar3D 以 FORMAL_MELEE_KIND 常量分支自动挂载，普通怪没有 "
         "configure_*_content() 函数。\n"
         "2026-09-20 首个标准普通怪转正：Tripo 源 → Blender 重制（模型 v002 / 动作 v003），"
         "六段剪辑按 12 态采样，verify_melee_zombie_presentation 通过。\n"
         f"显示名迁移：{LEGACY_NAME} → {DISPLAY_NAME}（Enemy3D.configure_from_enemy_data 精确迁移旧存档尾名）；"
         "AssetID 与逻辑 ID 不变。\n"
         "未闭合：碰撞 FOOTPRINT_PROFILES(radius 1.02/height 1.30) 仍为旧程序网格口径，"
         "按新几何应为 0.641/1.857，属玩法数值待拍板；贴图仅 basecolor，无法线贴图。"),
}
for col, value in master.items():
    master_ws.cell(6, col).value = value

# ---- 2) 《敌人动画与状态》r6:r17 —— 12 态补 4 列 ----------------------------------
SOURCE = "程序驱动；melee_chaser 已落地 Blender 六段剪辑"
per_state = {
    "dormant":   ("不播放（采样钉在 0）",        "—",                                "无（钉 0）"),
    "idle":      ("循环",                        "StateVFX 状态环",                  "idle 循环剪辑"),
    "patrol":    ("循环",                        "StateVFX 状态环",                  "walking（唯一用走路的移动状态）"),
    "alert":     ("循环（复用 idle）",           "StateVFX 状态环",                  "idle 循环剪辑"),
    "chase":     ("循环",                        "StateVFX 状态环",                  "running 循环剪辑"),
    "search":    ("循环",                        "StateVFX 状态环",                  "running 循环剪辑"),
    "return":    ("循环",                        "StateVFX 状态环",                  "running 循环剪辑"),
    "telegraph": ("单次（采样至 length×0.48）",  "StateVFX 状态环（前摇必须可读）",  "attack 采样 0.48 前段"),
    "attack":    ("单次（钉 length×0.48）",      "StateVFX 状态环；VFX-IMPACT-3D",   "attack 钉在判定帧"),
    "recovery":  ("单次（0.48 → 1.0 线性）",     "StateVFX 状态环",                  "attack 采样 0.48 后段"),
    "stagger":   ("单次（0.16s 硬直内）",        "受击闪白（材质 emission 覆盖）",   "hurt 按 0.16s 采样"),
    "dead":      ("单次（末帧保持）",            "VFX-IMPACT-3D；表现保留 2.4s",    "dead 采样至剪辑长"),
}
for row in range(6, 18):
    state = str(state_ws.cell(row, 2).value).strip()
    assert state in per_state, f"unexpected state id at r{row}: {state}"
    loop, hook, impl = per_state[state]
    state_ws.cell(row, 5).value = SOURCE
    state_ws.cell(row, 6).value = loop
    state_ws.cell(row, 7).value = hook
    state_ws.cell(row, 8).value = impl

# r21：敌人条目 x 状态机对照（小僵尸行）
state_ws.cell(21, 3).value = DISPLAY_NAME
state_ws.cell(21, 4).value = "Blender 动作母版 v003（36 骨；6 剪辑 idle/walking/running/attack/hurt/dead）"
state_ws.cell(21, 6).value = "骨骼 GLB / runtime Prefab（无碰撞）"
state_ws.cell(21, 7).value = f"有：{ANIM_SRC}"
state_ws.cell(21, 8).value = "已完成（verify_melee_zombie_presentation）"

# ---- 3) 《3D-敌人》新增 r7（样式照 r6） --------------------------------------------
new_row = 7
for col in range(1, 17):
    prefab_ws.cell(new_row, col)._style = copy(prefab_ws.cell(6, col)._style)
prefab = {
    1: "ENM-MELEE-FUNGBOAR01",
    2: DISPLAY_NAME,
    3: TSCN,
    4: GLB,
    5: f"{MODEL_SRC}; {ANIM_SRC}",
    6: "首个标准普通怪表现；六段 Blender 剪辑按 12 态 AI 采样；36 骨",
    7: f"src/enemy3d/EnemyAvatar3D.gd; {BASE}/runtime/melee_chaser_formal_visual.gd",
    8: "关（表现资产）",
    9: "Enemy3D",
    10: "CylinderShape3D（FOOTPRINT_PROFILES）",
    11: "1.2819×0.9535×1.8571m；1 Mesh；1 材质",
    12: "底部可预测；Blender +Y / Godot -Z 正面；根 Scale=1",
    13: "全部战斗关卡普通怪池（melee_chaser）",
    14: "正式美术已接入",
    15: "v003",
    16: ("无碰撞；Enemy3D 持有碰撞/AI/伤害/掉落。碰撞体尺寸沿用旧程序网格口径（1.02/1.30），"
         f"与模型几何（0.641/1.857）不一致，待拍板后再收紧。显示名 {LEGACY_NAME} → {DISPLAY_NAME}。"),
}
for col, value in prefab.items():
    prefab_ws.cell(new_row, col).value = value

# ---- 4) 《域变更日志》新增 r7（样式照 r6） -----------------------------------------
log_row = 7
for col in range(1, 8):
    log_ws.cell(log_row, col)._style = copy(log_ws.cell(6, col)._style)
entry = {
    1: "v0.1.1",
    2: "2026-09-20",
    3: "资产转正",
    4: "敌人",
    5: ("ENM-MELEE-FUNGBOAR01（小僵尸，旧登记名小菌猪）由「程序占位」转正为「已完成」："
        "AssetID 与逻辑 ID(melee_chaser) 不变；《资产主表》r6 迁 3D 口径（组件槽 root_3d、"
        "变体父 ID ENM-ECOSYSTEM-KIT-3D、视角 Top3D/local -Z、规格真实包围盒、文件路径 -> runtime tscn、"
        "SHA-256、模型/动作双 blend）；《敌人动画与状态》12 态补 动作来源/循环/挂点/首版实现 四列，"
        "条目行补骨架与动作母版；《3D-敌人》新增该 Prefab 行。"),
    6: "只升级既有行，未新增 AssetID；其余 6 类普通怪仍为程序占位。碰撞体积未动（玩法数值待拍板）。",
    7: "Codex",
}
for col, value in entry.items():
    log_ws.cell(log_row, col).value = value

wb.save(LEDGER)
print("LEDGER_WRITTEN", LEDGER.name)
print("promoted_asset", "ENM-MELEE-FUNGBOAR01", "| display", DISPLAY_NAME)
print("tscn_sha256", SHA)
print("3d_enemy_row", new_row, "| log_row", log_row, "| state_rows", "6-17 + r21")
