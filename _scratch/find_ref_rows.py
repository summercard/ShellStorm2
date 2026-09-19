import openpyxl
p = "assets/registry/ledgers/ShellStorm2_场景账本_v001.xlsx"
wb = openpyxl.load_workbook(p, data_only=False)
for name in wb.sheetnames:
    ws = wb[name]
    for r in range(1, ws.max_row+1):
        for c in range(1, ws.max_column+1):
            v = ws.cell(r,c).value
            if isinstance(v,str) and "ROOFTOP-REF" in v and "PARAPET" in v:
                print("### %s R%d C%d: %s" % (name, r, c, v))
                break
print("---- 主表中 ENV-ROOFTOP-REF-* 行 ----")
ws = wb["资产主表"]
for r in range(1, ws.max_row+1):
    v = ws.cell(r,1).value
    if isinstance(v,str) and v.startswith("ENV-ROOFTOP-REF"):
        print("  主表 R%d  %s" % (r, v))
print("---- 3D-场景通用 中 ENV-ROOFTOP-REF-* 行 ----")
ws = wb["3D-场景通用"]
cnt=0
for r in range(1, ws.max_row+1):
    v = ws.cell(r,1).value
    if isinstance(v,str) and v.startswith("ENV-ROOFTOP-REF"):
        cnt+=1
        if "PARAPET" in v or "PARAPET" in str(ws.cell(r,2).value):
            print("  R%d  %s | %s | %s | %s | %s" % (r, v, ws.cell(r,2).value, ws.cell(r,14).value, ws.cell(r,15).value, str(ws.cell(r,16).value)[:120]))
print("  总 ENV-ROOFTOP-REF 行数 =", cnt)
