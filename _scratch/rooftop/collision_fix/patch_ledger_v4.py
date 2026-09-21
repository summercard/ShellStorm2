"""四次修正：同步 100F 天台装饰布局在场景账本里的碰撞策略 / .blend SHA / 描述 / 备注 + 域变更日志 v0.1.8。

- 资产主表 row240：col9 状态、col14 规格（0→20 启用碰撞）、col20 SHA-256、col25 备注追加
- 3D-场景通用 row146：col6 功能说明、col8 碰撞开关、col9 碰撞归属、col10 碰撞方式、col16 备注
- 域变更日志 row14：v0.1.8

binary-safe：只改这三个 sheet 的目标单元格，不新增/删除行，改完回读自检。
"""
import hashlib
import shutil
from datetime import date
from pathlib import Path

import openpyxl

XLSX = Path("assets/registry/ledgers/ShellStorm2_场景账本_v001.xlsx")
BLEND = Path(
    "assets/art/environments/tower_zones/rooftop/source/layouts/"
    "100f_decorated_v001/rooftop_100f_decorated_layout_v001.blend"
)
JSON = Path(
    "assets/art/environments/tower_zones/rooftop/source/layouts/"
    "100f_decorated_v001/rooftop_100f_decorated_layout_v001.json"
)
BAK = XLSX.with_name(XLSX.name + ".bak_rooftop_decor_collision_fix")

OLD_SHA = "09e9a917012284710d01fa89216ff49d4137a3b54258096890b99715b40fb198"

NOTE4 = (
    "2026-09-21 四次修正（业主实机报「花盆和花圃没有阻挡」）：绿化 20 件（8 花箱 + 6 大盆栽 + 6 小盆栽）"
    "由 visual_only 改 blocking —— 布局源 add() 新增实例级 collision 策略（词表 visual_only / blocking，"
    "白名单 flowerbox / plant_large / plant_small），运行时 TowerFloorStage3D 新增 "
    "_apply_rooftop_collision_policy()：blocking 件用 TowerGeometry3D.resolve_visual_bounds() 量出**实测可视包络**、"
    "挂一个 StaticBody3D(BlockingCollision, layer=1, mask=0, PROCESS_MODE_ALWAYS) + BoxShape3D(BlockingBox) 代理"
    "挡玩家（尺寸随美术走、不写死常量），其余 462 件仍 visual_only（组件自带碰撞一律先关）。"
    "布局清单 design_intent/validation 记 blocking_collision_count=20 / visual_only_collision_count=462；"
    "校验器新增第 5 层断言（策略词表 + 白名单 slug + 计数 + 清单交叉一致）。"
    "反向对照：把 20 件翻回 visual_only ⇒ 两份运行时探针 exit 1、QA 报 3 条 FAIL；还原后 .blend 与 .json 逐字节一致。"
)

G6_APPEND = (
    "；2026-09-21 四次修正：绿化 20 件（8 花箱+6 大盆栽+6 小盆栽）改 blocking —— 运行时按实测可视包络生成 "
    "StaticBody3D 碰撞代理（layer=1，尺寸随美术自动跟随）挡玩家，其余 462 件仍 visual_only"
)

G16_APPEND = (
    "；2026-09-21 四次修正：绿化 20 件改 blocking（运行时按 TowerGeometry3D.resolve_visual_bounds 实测可视包络"
    "生成 layer=1 BoxShape3D 代理挡玩家），其余 visual_only；collision_policy 词表=visual_only/blocking，"
    "清单设计意图记 blocking_collision_count=20 / visual_only_collision_count=462"
)


def main() -> int:
    if not BAK.exists():
        shutil.copy2(XLSX, BAK)
        print("backup ->", BAK.name)

    new_sha = hashlib.sha256(BLEND.read_bytes()).hexdigest()
    json_sha = hashlib.sha256(JSON.read_bytes()).hexdigest()

    wb = openpyxl.load_workbook(XLSX)
    before = {
        s: (len(wb[s].merged_cells.ranges), len(list(wb[s].tables)))
        for s in wb.sheetnames
    }
    dv_before = {s: len(wb[s].data_validations.dataValidation) for s in wb.sheetnames}

    # ---------- 资产主表 row240 ----------
    ws = wb["资产主表"]
    assert ws.cell(240, 1).value == "ENV-ROOFTOP-DECOR-LAYOUT-100F", ws.cell(240, 1).value

    c9 = ws.cell(240, 9).value
    assert c9 == "static / visual_only", c9
    ws.cell(240, 9).value = "static / per_instance(visual_only+blocking)"

    c14 = ws.cell(240, 14).value
    assert "；0启用碰撞" in c14, c14
    ws.cell(240, 14).value = c14.replace(
        "；0启用碰撞", "；20启用碰撞（绿化：8花箱+6大盆栽+6小盆栽，layer=1 包络代理）", 1
    )

    sha_cell = ws.cell(240, 20).value
    assert sha_cell == OLD_SHA, sha_cell
    ws.cell(240, 20).value = new_sha

    note = ws.cell(240, 25).value
    assert "四次修正" not in note
    ws.cell(240, 25).value = note.rstrip() + NOTE4

    # ---------- 3D-场景通用 row146 ----------
    g = wb["3D-场景通用"]
    assert g.cell(146, 1).value == "ENV-ROOFTOP-DECOR-LAYOUT-100F", g.cell(146, 1).value

    c6 = g.cell(146, 6).value
    assert "四次修正" not in c6, c6
    g.cell(146, 6).value = c6.rstrip() + G6_APPEND

    c8 = g.cell(146, 8).value
    assert c8 == "未创建", c8
    g.cell(146, 8).value = "开"

    c9g = g.cell(146, 9).value
    assert c9g == "引擎（visual_only）", c9g
    g.cell(146, 9).value = "引擎（per_instance：默认 visual_only，绿化 blocking）"

    c10 = g.cell(146, 10).value
    assert c10 == "未创建", c10
    g.cell(146, 10).value = (
        "BoxShape3D 代理（按 TowerGeometry3D.resolve_visual_bounds 实测可视包络；仅绿化 20 件，layer=1）"
    )

    c16 = g.cell(146, 16).value
    assert "四次修正" not in c16, c16
    g.cell(146, 16).value = c16.rstrip() + G16_APPEND

    # ---------- 域变更日志 row14 ----------
    lg = wb["域变更日志"]
    r = 14
    assert lg.cell(r, 1).value in (None, ""), lg.cell(r, 1).value
    lg.cell(r, 1).value = "v0.1.8"
    lg.cell(r, 2).value = date(2026, 9, 21).isoformat()
    lg.cell(r, 3).value = "条目修正"
    lg.cell(r, 4).value = "关卡场景 / 100F天台"
    lg.cell(r, 5).value = (
        "100F 天台装饰布局四次修正（业主实机报「花盆和花圃没有阻挡」）：绿化 20 件（8 花箱 + 6 大盆栽 + "
        "6 小盆栽）由 visual_only 改 blocking —— 布局源 add() 新增实例级 collision 策略（词表 "
        "visual_only/blocking，白名单 flowerbox/plant_large/plant_small），运行时 TowerFloorStage3D 新增 "
        "_apply_rooftop_collision_policy()：blocking 件按 TowerGeometry3D.resolve_visual_bounds() 实测可视包络"
        "挂 StaticBody3D(layer=1)+BoxShape3D 代理挡玩家，尺寸随美术自动跟随；其余 462 件仍 visual_only。"
    )
    lg.cell(r, 6).value = (
        "AssetID / layout_id / layout_version 均不变（原地修正，不新增条目、不新增资产）；账本门禁维持 46；"
        ".blend SHA-256 09e9a917… → %s、清单 .json %s… 同步；结构壳体与玩法不变。" % (new_sha[:8], json_sha[:8])
    )
    lg.cell(r, 7).value = "摩斯拉"

    wb.save(XLSX)

    # ---------- 回读自检 ----------
    wb2 = openpyxl.load_workbook(XLSX)
    after = {
        s: (len(wb2[s].merged_cells.ranges), len(list(wb2[s].tables)))
        for s in wb2.sheetnames
    }
    dv_after = {s: len(wb2[s].data_validations.dataValidation) for s in wb2.sheetnames}
    assert before == after, (before, after)
    assert dv_before == dv_after, (dv_before, dv_after)
    ws2 = wb2["资产主表"]
    assert ws2.cell(240, 9).value == "static / per_instance(visual_only+blocking)"
    assert ws2.cell(240, 20).value == new_sha
    assert wb2["3D-场景通用"].cell(146, 9).value.startswith("引擎（per_instance")
    assert wb2["域变更日志"].cell(14, 1).value == "v0.1.8"

    print("LEDGER_PATCH_V4_OK")
    print("  new_blend_sha256 =", new_sha)
    print("  json_sha256      =", json_sha)
    print("  backup           =", BAK.name)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
