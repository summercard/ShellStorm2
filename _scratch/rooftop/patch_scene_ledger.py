"""三次修正：同步 100F 天台装饰布局在场景账本里的实例数 / .blend SHA / 备注 + 域变更日志。"""
import hashlib
import shutil
import sys
from datetime import date

import openpyxl

P = "assets/registry/ledgers/ShellStorm2_场景账本_v001.xlsx"
BLEND = (
    "assets/art/environments/tower_zones/rooftop/source/layouts/"
    "100f_decorated_v001/rooftop_100f_decorated_layout_v001.blend"
)
BAK = P + ".bak_rooftop_decor_third_fix"

shutil.copy2(P, BAK)
new_sha = hashlib.sha256(open(BLEND, "rb").read()).hexdigest()

wb = openpyxl.load_workbook(P)
ws = wb["资产主表"]

# ① 数量
old = ws.cell(row=240, column=14).value
assert old.startswith("112装饰实例/6组"), old
ws.cell(row=240, column=14).value = old.replace("112装饰实例/6组", "120装饰实例/6组", 1)

# ② .blend SHA256
old_sha = ws.cell(row=240, column=20).value
assert old_sha == "96f914c5f991e1667c2838f2ad079e1b5335246b6a13dda7c3dc1b934d8f7c04", old_sha
ws.cell(row=240, column=20).value = new_sha

# ③ 备注追加
NOTE3 = (
    "2026-09-21 三次修正（业主实机报「水管接下来接到地板上面基本看不到 / 空调机在墙上的状态"
    "需要 90 度旋转让风扇朝外」）：①立管落地 —— pipe_riser 件高 4.945m、原点在管底、顶端是朝 +X "
    "的鹅颈出水口；旧版放 h=5.8 ⇒ 管底悬空 5.8m。现每处摆两段、两段 h 都是 4.945：下段 "
    "rotation_x_deg=180 倒装（包络 0~4.945m，原点即上端）、上段正装（4.945~9.89m），合成一根 "
    "0~9.89m 连续落水管，与 10.65m 环管之间残留 0.76m 由环管本体与支架轨遮住；支架补一段低位"
    "（1.5m）与原有 7.0m 位，覆盖 1.56~10.64m。②墙挂空调倾倒 —— 新增实例朝向分量 rotation_x_deg"
    "（绕自身 X 轴的倾倒）；组件 hvac_small 顶面（Blender 局部 +Z）是出风风扇、前面（局部 −Y）"
    "是进风格栅，故 4 台墙挂机统一 rotation_x_deg=90 让风扇朝外（旧版风扇朝天）。⚠️ Blender 默认 "
    "XYZ 序、Godot 默认 YXZ 序，只有 rz 分量为 0 时 Rz@Rx 与 Ry@Rx 同序、角度可逐值搬运，本布局"
    "遵守该约束（tip 只与 yaw 组合）。③立管/支架按各自墙面给 yaw（旧版立管一律 0 ⇒ 东西墙鹅颈朝"
    "墙里）；6 个墙挂件（4 空调 + 2 通风口）后背统一埋进墙外皮内 0.05m。重放实例 112 → 120"
    "（pipe_risers 8 → 16）。"
)
note = ws.cell(row=240, column=25).value
assert "三次修正" not in note
ws.cell(row=240, column=25).value = note.rstrip() + NOTE3

# ④ 域变更日志
lg = wb["域变更日志"]
r = 13
assert lg.cell(row=r, column=1).value in (None, "")
lg.cell(row=r, column=1).value = "v0.1.7"
lg.cell(row=r, column=2).value = date(2026, 9, 21).isoformat()
lg.cell(row=r, column=3).value = "条目修正"
lg.cell(row=r, column=4).value = "关卡场景 / 100F天台"
lg.cell(row=r, column=5).value = (
    "100F 天台装饰布局三次修正两件事：①立管落地——pipe_riser 件高 4.945m、原点在管底，"
    "旧版放 h=5.8 ⇒ 管底悬空 5.8m；现每处两段（下段 rotation_x_deg=180 倒装），合成 0~9.89m "
    "连续落水管并落地 y=0，与 10.65m 环管残留 0.76m 由环管/支架遮住。②墙挂空调倾倒——新增 "
    "rotation_x_deg 分量（绕自身 X 轴），4 台墙挂机统一 90° 让顶部风扇朝外（旧版风扇朝天）；"
    "Blender XYZ 序与 Godot YXZ 序在 rz=0 时同序、角度逐值搬运。"
)
lg.cell(row=r, column=6).value = (
    "AssetID / layout_id / layout_version 均不变（原地修正，不新增条目）；装饰件仍 visual_only、"
    "0 启用碰撞；结构壳体（234 地砖 / 68 件女儿墙 / 楼梯与碰撞）与玩法不变；重放实例 112 → 120。"
)
lg.cell(row=r, column=7).value = "摩斯拉"

wb.save(P)
print("LEDGER_PATCH_OK new_sha=%s backup=%s" % (new_sha, BAK))
