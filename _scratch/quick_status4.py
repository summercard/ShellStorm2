import openpyxl, collections
p = r"I:\工作项目\shellstrom2\ShellStorm2\assets\registry\ShellStorm2_美术资产台账_v001.xlsx"
wb = openpyxl.load_workbook(p, data_only=True)
sheets = ["3D-场景通用", "3D-设施", "3D-道具", "3D-角色", "3D-敌人", "3D-武器", "3D-物品", "3D-特效", "3D-其他"]
def col_of(ws, name):
    for c in range(1, ws.max_column + 1):
        if ws.cell(4, c).value == name:
            return c
    return None
tot = collections.Counter()
has_glb = collections.Counter()
has_blend = collections.Counter()
per = {}
for s in sheets:
    ws = wb[s]
    cA, cD, cE, cN = col_of(ws, "AssetID"), col_of(ws, "GLB模型路径"), col_of(ws, "Blender源文件"), col_of(ws, "制作状态")
    cnt = collections.Counter(); g = b = 0; rowtot = 0
    for r in range(5, ws.max_row + 1):
        aid = ws.cell(r, cA).value
        if not aid or not str(aid).strip():
            continue
        rowtot += 1
        st = str(ws.cell(r, cN).value or "(空)").strip()
        cnt[st] += 1
        if ws.cell(r, cD).value: g += 1
        if ws.cell(r, cE).value: b += 1
    per[s] = (rowtot, g, b, cnt)
    tot.update(cnt); has_glb[s] = g; has_blend[s] = b
print(f"{'sheet':<12}{'总数':>6}{'有GLB':>7}{'有Blender源':>12}")
for s in sheets:
    t, g, b, _ = per[s]
    print(f"{s:<12}{t:>6}{g:>7}{b:>12}")
print("-" * 40)
print("总计行:", sum(per[s][0] for s in sheets), " 有GLB:", sum(has_glb.values()), " 有Blender源:", sum(has_blend.values()))
print()
print("=== 制作状态汇总 ===")
for k, v in tot.most_common():
    print(f"  {k:<24}{v}")
