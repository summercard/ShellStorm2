from __future__ import annotations

import copy
import hashlib
import json
import shutil
from collections import Counter
from datetime import datetime
from pathlib import Path

import openpyxl

ROOT = Path('I:/工作项目/shellstrom2/ShellStorm2')
OUT = ROOT / 'outputs/base99_radio_v005'
ASSET_ID = 'PRP-BASE99-RADIO-3D'
SOURCE = ROOT / 'assets/art/props/base_world_3d/source/base99_radio/prp_base99_radio_source_v005.blend'
OPTIMIZED = ROOT / 'assets/art/props/base_world_3d/source/base99_radio/export/v005/prp_base99_radio_optimized_v005.blend'
GLB = ROOT / 'assets/art/props/base_world_3d/components/base99_radio/prp_base99_radio_visual_top3d.glb'
PREFAB = ROOT / 'assets/art/props/base_world_3d/runtime/base99_radio/prp_base99_radio_root_top3d.tscn'

import sys
sys.path.insert(0, str(ROOT / 'scripts'))
sys.path.insert(0, str(ROOT / 'tools/asset_pipeline'))
from ledger_registry import LedgerIndex
from split_asset_ledger import CONTENT_COLUMNS, _row_digest, col_digest, read_source_rows, sheet_digest


def sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def text(value) -> str:
    return '' if value is None else str(value).strip()


def rows(ws):
    return read_source_rows(ws)


PROGRESS = OUT / 'ledger_v005_progress.log'

def mark(message: str) -> None:
    PROGRESS.parent.mkdir(parents=True, exist_ok=True)
    with PROGRESS.open('a', encoding='utf-8') as handle:
        handle.write(message + '\n')


def main() -> None:
    mark('start')
    for path in (SOURCE, OPTIMIZED, GLB, PREFAB):
        assert path.is_file(), path
    mark('assets-ok')
    index = LedgerIndex.load(ROOT)
    mark('index-ok')
    ledger = index.path_for_category('道具')
    baseline_path = ROOT / 'assets/registry/ledger_split_baseline.json'
    assert ledger.is_file() and baseline_path.is_file()

    transaction = OUT / ('ledger_v005_transaction_' + datetime.now().strftime('%Y%m%d_%H%M%S_%f'))
    transaction.mkdir(parents=True, exist_ok=True)
    backup_files = [ledger, index.master_path, baseline_path]
    for path in backup_files:
        shutil.copy2(path, transaction / path.name)
    mark('backups-ok')

    source_sha = sha(SOURCE)
    optimized_sha = sha(OPTIMIZED)
    glb_sha = sha(GLB)
    prefab_sha = sha(PREFAB)
    relative = lambda path: path.relative_to(ROOT).as_posix()
    manifest = json.loads((OUT / 'asset_manifest.json').read_text(encoding='utf-8'))
    pixels=json.loads((OUT/'pixel_metrics_runtime.json').read_text(encoding='utf-8'))
    visibility=bool(pixels['visibility_passed'])
    faces, triangles = manifest['faces'], manifest['triangles']
    diameter = manifest['status_light']['diameter_m']
    spec = (
        f'宽0.828m × 深0.456m × 高0.822m；{faces}面/{triangles}三角形<800面；'
        '机身铜橙金属、天线与提手等范围外几何签名保持；单根三节斜金属天线29.35度；'
        '四材质公共色盘外链/PaletteUV/Closest；StatusLight为顶面偏前侧12边低模凸帽，'
        f'直径{diameter}m，替换旧灯且不重复，避开提手/天线；off红常亮、A/B同绿常亮、离开99F红待机；'
        '发光能量1.5；ItemRoot/Visual/ItemRoot/StatusLight实际导入路径稳定；左键/E关→A→B→关及Music逻辑不变；'
        f'root scale=1，bounds不变；可见性门禁={visibility}；原生1280x720像素证据=outputs/base99_radio_v005/pixel_metrics_runtime.json。'
    )

    workbook = openpyxl.load_workbook(ledger)
    asset_sheet = workbook['资产主表']
    prefab_sheet = workbook['3D-道具']
    log_sheet = workbook['域变更日志']
    assert asset_sheet['A27'].value == ASSET_ID
    assert prefab_sheet['A8'].value == ASSET_ID

    changes = {
        'M27': 'v005',
        'N27': spec,
        'P27': '; '.join([
            relative(SOURCE), relative(OPTIMIZED), relative(GLB),
            'src/base3d/Base99Radio3D.gd',
            'outputs/base99_radio_v005/asset_manifest.json',
            'outputs/base99_radio_v005/final_acceptance.json',
        ]),
        'T27': prefab_sha,
        'V27': datetime(2026, 10, 7),
        'Y27': (
            f'v005局部状态灯深化；旧v004源/优化保留；顶面偏前侧12边凸帽直径{diameter}m；'
            f'{faces}面/{triangles}三角；结构/UV/公共色盘与交互结果详见当前final_acceptance；'
            f'正常玩家及全阁楼红绿可辨门禁={visibility}，像素详见pixel_metrics_runtime.json；音乐/摆位/范围外几何不变；'
            f'GLB SHA-256={glb_sha}；source SHA-256={source_sha}；optimized SHA-256={optimized_sha}'
        ),
    }
    for coordinate, value in changes.items():
        asset_sheet[coordinate] = value

    prefab_changes = {
        'E8': relative(SOURCE),
        'K8': spec,
        'O8': 'v005',
        'P8': (
            f'ItemRoot/Visual/StatusLight接口不变；StatusLight顶面偏前侧12边低模凸帽；'
            f'单一灯，无billboard/UI/bloom；源={source_sha}；优化={optimized_sha}；'
            f'GLB={glb_sha}；Prefab={prefab_sha}；证据=outputs/base99_radio_v005/'
        ),
    }
    for coordinate, value in prefab_changes.items():
        prefab_sheet[coordinate] = value
    mark('cells-written')

    workbook.save(ledger)
    mark('ledger-saved')

    after_workbook = openpyxl.load_workbook(ledger, data_only=False)
    mark('after-ledger-loaded')
    after_asset_sheet = after_workbook['资产主表']
    after_prefab_sheet = after_workbook['3D-道具']
    before_workbook=openpyxl.load_workbook(transaction/ledger.name,data_only=False)
    actual_differences=[]
    allowed={('资产主表',c) for c in changes}|{('3D-道具',c) for c in prefab_changes}
    for old_sheet in before_workbook:
        new_sheet=after_workbook[old_sheet.title]
        assert (old_sheet.max_row,old_sheet.max_column)==(new_sheet.max_row,new_sheet.max_column)
        assert str(old_sheet.merged_cells)==str(new_sheet.merged_cells)
        assert [(v.type,v.formula1,str(v.sqref)) for v in old_sheet.data_validations.dataValidation]==[(v.type,v.formula1,str(v.sqref)) for v in new_sheet.data_validations.dataValidation]
        for row in old_sheet:
            for cell in row:
                current=new_sheet[cell.coordinate]
                assert cell._style==current._style
                if cell.value!=current.value:
                    assert (old_sheet.title,cell.coordinate) in allowed,(old_sheet.title,cell.coordinate)
                    actual_differences.append(old_sheet.title+'!'+cell.coordinate)
    assert sha(index.master_path)==sha(transaction/index.master_path.name)
    before_full_raw=(OUT/'full_props_resume_before.log').read_bytes()
    before_full_text=before_full_raw.decode('utf-8') if before_full_raw[:3]==b'\xef\xbb\xbf' else before_full_raw.decode('gbk')
    before_full=json.JSONDecoder().raw_decode(before_full_text[before_full_text.index('{'):])[0]
    debt_unchanged=[]
    for issue in before_full['issues'].get('sha_mismatch',[]):
        assert sha(Path(issue['path']))==issue['actual'],issue['asset_id']
        row=issue['row']
        assert before_workbook['资产主表'].cell(row,20).value==after_asset_sheet.cell(row,20).value==issue['recorded']
        debt_unchanged.append({'asset_id':issue['asset_id'],'row':row,'recorded':issue['recorded'],'actual':issue['actual'],'asset_file_and_ledger_cell_unchanged':True})
    target_values = next(values for _, values in rows(after_asset_sheet) if values[0] == ASSET_ID)
    target_digest = _row_digest(target_values)

    baseline = json.loads(baseline_path.read_text(encoding='utf-8'))
    old_assets = copy.deepcopy(baseline['assets'])
    old_count = baseline['asset_count']
    baseline['assets'][ASSET_ID] = {'v': target_digest, 'c': '道具', 'd': 'props'}
    assert baseline['asset_count'] == old_count
    assert len(baseline['assets']) == len(old_assets)
    for asset_id, record in old_assets.items():
        if asset_id != ASSET_ID:
            assert baseline['assets'][asset_id] == record, asset_id

    all_rows = []
    mark('before-all-rows')
    for domain in index.domains:
        domain_book = openpyxl.load_workbook(domain.path, read_only=False, data_only=False)
        all_rows.extend((asset_id, values) for _, values in rows(domain_book['资产主表']) for asset_id in [text(values[0])])
        domain_book.close()
        mark('domain-' + domain.key)
    baseline['column_digests'] = {str(column): col_digest(all_rows, column) for column in CONTENT_COLUMNS}
    mark('columns-digested')
    baseline['category_counts'] = {
        category: sum(1 for _, values in all_rows if text(values[2]) == category)
        for category in sorted({text(values[2]) for _, values in all_rows})
    }
    baseline['sheet_digests']['3D-道具'] = sheet_digest(after_prefab_sheet)
    baseline['captured_at'] = '2026-10-07'
    baseline_path.write_text(json.dumps(baseline, ensure_ascii=False, indent=1) + '\n', encoding='utf-8')
    mark('baseline-saved')

    evidence = {
        'passed': True,
        'asset_id': ASSET_ID,
        'version': 'v005',
        'ledger': relative(ledger),
        'asset_sheet': '资产主表!27',
        'prefab_sheet': '3D-道具!8',
        'existing_full_props_debt_unchanged':debt_unchanged,
        'actual_cell_differences': actual_differences,
        'other_cells_styles_formulas_dimensions_validations_unchanged': True,
        'master_workbook_unchanged': True,
        'changed_asset_cells': sorted(changes),
        'changed_prefab_cells': sorted(prefab_changes),
        'baseline': relative(baseline_path),
        'baseline_asset_count': baseline['asset_count'],
        'target_row_digest': target_digest,
        'source_sha256': source_sha,
        'optimized_sha256': optimized_sha,
        'glb_sha256': glb_sha,
        'prefab_sha256': prefab_sha,
        'other_asset_fingerprints_unchanged': True,
        'normal_player_visibility_gate': visibility,
        'transaction_backup': transaction.as_posix(),
        'updated_at': '2026-10-07',
    }
    (OUT / 'ledger_v005_evidence.json').write_text(json.dumps(evidence, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    print('RADIO_V005_LEDGER_TRANSACTION_OK')
    print(json.dumps(evidence, ensure_ascii=False))


if __name__ == '__main__':
    main()
