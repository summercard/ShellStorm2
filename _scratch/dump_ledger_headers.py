# -*- coding: utf-8 -*-
"""dump 资产主表表头与若干样例行，确认各列语义。"""
from openpyxl import load_workbook

wb = load_workbook("assets/registry/ledgers/ShellStorm2_场景账本_v001.xlsx")
ws = wb["资产主表"]

print("=== 资产主表 前 5 行 ===")
for r in range(1, 6):
    print("row", r)
    for c in range(1, ws.max_column + 1):
        v = ws.cell(row=r, column=c).value
        if v is not None:
            print("   %-4s = %s" % (ws.cell(row=r, column=c).column_letter, str(v)[:80]))

print()
print("=== 样例行：battle 正式美术房间 / 远征 Boss 房 ===")
for r in (238, 239, 240, 241):
    print("row", r)
    for c in range(1, ws.max_column + 1):
        v = ws.cell(row=r, column=c).value
        if v is not None:
            print("   %-4s = %s" % (ws.cell(row=r, column=c).column_letter, str(v)[:150]))
    print()

print("=== 3D-场景通用 前 5 行表头 ===")
p = wb["3D-场景通用"]
for r in range(1, 6):
    vals = [str(p.cell(row=r, column=c).value)[:26] if p.cell(row=r, column=c).value is not None else "" for c in range(1, p.max_column + 1)]
    print("row", r, "|", " | ".join(vals))
