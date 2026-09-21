import json, hashlib
from pathlib import Path
from openpyxl import load_workbook
root=Path(r"I:\工作项目\shellstrom2\ShellStorm2")
ledger=root/'assets/registry/ledgers/ShellStorm2_敌人账本_v001.xlsx'
wb=load_workbook(ledger, read_only=True, data_only=False)
files={}
for rel in ['assets/art/enemies/normal_enemy_3d/ranged_caster/source/model/enm_ranged_sporeshooter01_model_v002.blend','assets/art/enemies/normal_enemy_3d/ranged_caster/source/animation/enm_ranged_sporeshooter01_animation_v003.blend','assets/art/enemies/normal_enemy_3d/ranged_caster/components/enm_ranged_sporeshooter01_visual_top3d.glb','assets/art/enemies/normal_enemy_3d/ranged_caster/runtime/enm_ranged_sporeshooter01_root_top3d.tscn']:
 p=root/rel; files[rel]={'bytes':p.stat().st_size,'sha256':hashlib.sha256(p.read_bytes()).hexdigest()}
print(json.dumps({'sheets':wb.sheetnames,'rows':{s:wb[s].max_row for s in wb.sheetnames},'files':files},ensure_ascii=False,indent=2))