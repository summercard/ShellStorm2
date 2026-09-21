from openpyxl import load_workbook
p=r'I:\工作项目\shellstrom2\ShellStorm2\\assets\\registry\\ledgers\\ShellStorm2_敌人账本_v001.xlsx'
w=load_workbook(p)
print(w.sheetnames)
[(print(s,w[s].max_row,w[s].max_column)) for s in w.sheetnames]
