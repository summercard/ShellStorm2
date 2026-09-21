from copy import copy
from datetime import datetime
from hashlib import sha256
from pathlib import Path
from openpyxl import load_workbook
ROOT=Path('I:/工作项目/shellstrom2/ShellStorm2'); LEDGER=ROOT/'assets/registry/ledgers/ShellStorm2_敌人账本_v001.xlsx'
BASE='assets/art/enemies/normal_enemy_3d/ranged_caster'; TSCN=f'{BASE}/runtime/enm_ranged_sporeshooter01_root_top3d.tscn'; GLB=f'{BASE}/components/enm_ranged_sporeshooter01_visual_top3d.glb'; MODEL=f'{BASE}/source/model/enm_ranged_sporeshooter01_model_v002.blend'; ANIM=f'{BASE}/source/animation/enm_ranged_sporeshooter01_animation_v003.blend'; SHA=sha256((ROOT/TSCN).read_bytes()).hexdigest()
wb=load_workbook(LEDGER); mw=wb['资产主表']; sw=wb['敌人动画与状态']; pw=wb['3D-敌人']; lw=wb['域变更日志']
assert mw.cell(7,1).value=='ENM-RANGED-SPORESHOOTER01' and mw.cell(7,5).value=='ranged_caster'; assert sw.cell(22,2).value=='ENM-RANGED-SPORESHOOTER01'; assert pw.max_row>=7; assert lw.max_row>=7
vals={2:'保安僵尸',6:'root_3d',7:'ENM-ECOSYSTEM-KIT-3D',8:'Top3D / local -Z 正面',9:'default / chase / shoot',10:'全部战斗关卡远程怪池',11:'已完成',13:'v002',14:'1.2819×0.9535×1.8571m；1 Mesh；1 材质；36 骨；10 剪辑',15:TSCN,16:f'{MODEL}; {ANIM}',17:'远程射击; ranged; 保安僵尸; 孢子射手（旧程序名）; 单手持枪; ENM-RANGED-SPORESHOOTER01',20:SHA,21:'Codex',22:datetime(2026,9,21),23:'桌面 FBX + Blender 共享小僵尸骨架重制',25:'GLB 仅表现；runtime PackedScene 无碰撞；Enemy3D 继续持有 ranged_caster 的 AI、三弹散布、生命、伤害与掉落。枪为双管炮视觉附件，挂 L_Hand，不接玩家武器逻辑。'}
for c,v in vals.items():mw.cell(7,c).value=v
for r in range(6, sw.max_row + 1):
 state=str(sw.cell(r,2).value); sw.cell(r,5).value='程序驱动；ranged_caster 已接入保安僵尸 Blender 视觉与射击剪辑';
 if state in ('telegraph','attack','recovery'): sw.cell(r,8).value='shoot（状态机采样；攻击事件仍由 Enemy3D 发射三弹）'
sw.cell(22,3).value='保安僵尸';sw.cell(22,4).value='Blender 动作母版 v002（复用36骨；6基础剪辑 + armed_idle/walking_armed/running_armed/shoot）';sw.cell(22,6).value='骨骼 GLB / runtime Prefab / 双管炮视觉附件（无碰撞）';sw.cell(22,7).value=f'有：{ANIM}';sw.cell(22,8).value='已接入（verify_security_zombie_presentation）'
for c in range(1, pw.max_column + 1):pw.cell(8,c)._style=copy(pw.cell(7,c)._style)
pref={1:'ENM-RANGED-SPORESHOOTER01',2:'保安僵尸',3:TSCN,4:GLB,5:f'{MODEL}; {ANIM}',6:'远程普通怪正式表现；复用小僵尸36骨与基础动作；单手持枪shoot；双管炮视觉挂 L_Hand',7:f'src/enemy3d/EnemyAvatar3D.gd; {BASE}/runtime/ranged_caster_formal_visual.gd',8:'关（表现资产）',9:'Enemy3D',10:'CylinderShape3D（既有 ranged_caster FOOTPRINT_PROFILES）',11:'1.2819×0.9535×1.8571m；1 Mesh；1 材质；36 骨；10 剪辑',12:'底部可预测；Blender +Y / Godot -Z 正面；根 Scale=1',13:'全部战斗关卡远程怪池（ranged_caster）',14:'正式美术已接入',15:'v002',16:'无碰撞；Enemy3D 保留三弹齐射和所有玩法规则；双管炮仅视觉附件。'}
for c,v in pref.items():pw.cell(8,c).value=v
for c in range(1, lw.max_column + 1):lw.cell(8,c)._style=copy(lw.cell(7,c)._style)
log={1:'v0.1.2',2:'2026-09-21',3:'资产转正',4:'敌人',5:'ENM-RANGED-SPORESHOOTER01（保安僵尸）由程序占位转正为正式3D表现：复用小僵尸36骨与基础动作，新增armed_idle/持枪移动/shoot，双管炮仅作为L_Hand视觉附件；ranged_caster逻辑与三弹发射规则不变。',6:'只升级既有 AssetID 与 ranged_caster 行，未新增逻辑敌人；Prefab无碰撞，Enemy3D继续持有AI、碰撞、生命、伤害、掉落和弹丸。',7:'Codex'}
for c,v in log.items():lw.cell(8,c).value=v
wb.save(LEDGER);print('RANGED_LEDGER_WRITTEN',SHA)
