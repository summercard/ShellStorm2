import json
from pathlib import Path
from openpyxl import load_workbook
p=Path(r"I:\工作项目\shellstrom2\ShellStorm2\assets\registry\ledgers\ShellStorm2_敌人账本_v001.xlsx")
w=load_workbook(p,read_only=True,data_only=False)
for s in w.sheetnames:
    print(s.encode('unicode_escape').decode(), w[s].max_row, w[s].max_column)
for s, coords in {'资产主表':['A7','M7','O7','P7'],'3D-敌人':['A8','C8','D8','E8'],'敌人动画与状态':['B22','C22','D22','G22'],'域变更日志':['A8','E8']}.items():
    if s in w.sheetnames:
        print(s, {c:w[s][c].value for c in coords})