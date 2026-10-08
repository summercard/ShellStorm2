from pathlib import Path
from copy import copy,deepcopy
from datetime import datetime
import json,hashlib,re,shutil,sys,zipfile,xml.etree.ElementTree as ET
import openpyxl
from openpyxl.worksheet.datavalidation import DataValidation
R=Path('I:/工作项目/shellstrom2/ShellStorm2'); O=R/'outputs/base99_radio_music_notes'
sys.path.insert(0,str(R/'scripts')); sys.path.insert(0,str(R/'tools/asset_pipeline'))
from ledger_registry import LedgerIndex
from split_asset_ledger import FIRST_DATA_ROW,HEADER_ROW,CONTENT_COLUMNS,read_source_rows,_row_digest,col_digest,sheet_digest,dedupe_key_formula,dedupe_result_formula
index=LedgerIndex.load(R)
P=index.path_for_category('道具'); V=index.path_for_category('特效'); B=R/'assets/registry/ledger_split_baseline.json'
ID='VFX-RADIO-MUSIC-NOTES-3D'; RID='PRP-BASE99-RADIO-3D'
FX='assets/art/vfx/environment_3d/radio_music_notes/vfx_radio_music_notes_root_top3d.tscn'
RP='assets/art/props/base_world_3d/runtime/base99_radio/prp_base99_radio_root_top3d.tscn'
SPEC='独立radio child纯视觉持续附件；5预制QuadMesh槽位/10三角形；Godot生成四分/八分/双八分音符纹理，无字体依赖；billboard当前相机；淡青绿/暖金unshaded无emission；2.5发/s、寿命1.8s、上飘1.1m、横摆0.09m、尺度0.88~1.06、淡入0.16s、末0.45s淡出；本地挂点Y1.42m；无碰撞/灯光/音频；真实playing驱动，A/B连续，off/离99/exit同步清空；不进任何战斗池。'
def sha(p): return hashlib.sha256(p.read_bytes()).hexdigest()
def dump(p,obj): p.write_bytes((json.dumps(obj,ensure_ascii=False,indent=2,default=str)+'\n').replace('\n','\r\n').encode('utf-8'))
def log(wb,version,kind,scope,description):
 ws=wb['域变更日志']; assert not any(c.value==version for c in ws['A'])
 row=ws.max_row+1
 for c,val in enumerate([version,datetime(2026,10,8),kind,scope,description,'稳定ID/路径不变；主体v005几何与灯帽不变；不改点击或玩法','CodeBuddy'],1):
  ws.cell(row,c).value=val; ws.cell(row,c)._style=copy(ws.cell(row-1,c)._style)
def reconcile(changed,sheets):
 bl=json.loads(B.read_text(encoding='utf-8')); old=deepcopy(bl['assets']); allrows=[]
 for domain in index.domains:
  wb=openpyxl.load_workbook(domain.path,data_only=False)
  for row,vals in read_source_rows(wb['资产主表']):
   allrows.append((row,vals))
   if vals[0] in changed: bl['assets'][vals[0]]={'v':_row_digest(vals),'c':vals[2],'d':domain.key}
  for name in sheets:
   if name in wb: bl['sheet_digests'][name]=sheet_digest(wb[name])
  wb.close()
 for aid,rec in old.items():
  if aid not in changed: assert bl['assets'][aid]==rec
 bl['asset_count']=len(bl['assets']); bl['column_digests']={str(c):col_digest(allrows,c) for c in CONTENT_COLUMNS}
 bl['category_counts']={cat:sum(vals[2]==cat for _,vals in allrows) for cat in sorted({vals[2] for _,vals in allrows})}
 bl['captured_at']='2026-10-08'; dump(B,bl)
 return bl['asset_count']
# 事务1只登记资产主表，域专表另开事务。
pwb=openpyxl.load_workbook(P); vwb=openpyxl.load_workbook(V)
assert pwb['资产主表']['A27'].value==RID
assert not any(vals[0]==ID for domain in index.domains for _,vals in read_source_rows(openpyxl.load_workbook(domain.path)['资产主表']))
for path in [P,V,B]: assert (O/'backup'/path.relative_to(R)).is_file()
ws=pwb['资产主表']; ws['M27']='v005.1'
ws['N27']=str(ws['N27'].value)+' runtime v005.1：模型/Blender/GLB保持v005；新增独立VFX音符附件，预算另计。'+SPEC
ws['T27']=sha(R/RP); ws['V27']=datetime(2026,10,8)
ws['P27']=str(ws['P27'].value)+'; '+FX+'; src/vfx/VfxRadioMusicNotes3D.gd; tests/verification/verify_base99_radio_music_notes.tscn'
ws['Y27']=str(ws['Y27'].value)+' 【runtime v005.1】主体模型v005未重做；收音机持有MusicNotes持续视觉附件；验收=outputs/base99_radio_music_notes/final_acceptance.json'
log(pwb,'v0.1.13','运行时补丁','道具 · 资产主表','PRP-BASE99-RADIO-3D runtime v005.1，模型v005不变；更新Prefab哈希与播放音符附件规格。')
ws=vwb['资产主表']; oldlast=max(r for r,_ in read_source_rows(ws)); new=oldlast+1
values=[ID,'收音机音乐上飘音符','特效','environment','base99_radio_music_notes','root_3d',None,'俯视3D','continuous_music','99F收音机/复用radio附件','正式美术已接入','P1','v001',SPEC,FX,'src/vfx/VfxRadioMusicNotes3D.gd; src/base3d/Base99Radio3D.gd; tests/verification/verify_base99_radio_music_notes.tscn','音符;音乐;收音机;八分;双八分;billboard',dedupe_key_formula(new),dedupe_result_formula(new,new),sha(R/FX),'CodeBuddy',datetime(2026,10,8),'原创Godot ImageTexture/QuadMesh；无第三方字体','FX03-03','宿主唯一生命周期；固定槽位非live节点增殖；off全部槽位隐藏、alpha=0、age=-1；主体599面独立统计。']
for col,val in enumerate(values,1):
 ws.cell(new,col).value=val; ws.cell(new,col)._style=copy(ws.cell(oldlast,col)._style)
ws.row_dimensions[new].height=ws.row_dimensions[oldlast].height
for row in range(FIRST_DATA_ROW,new+1): ws.cell(row,19).value=dedupe_result_formula(row,new)
for ref in ['A6','C6','E6','G6','B10','C10','B11','C11']:
 cell=vwb['总览'][ref]; assert isinstance(cell.value,str) and cell.value.startswith('=')
 cell.value=re.sub(r'(\$[A-Z]+\$)'+str(oldlast)+r'\b',lambda m:m[1]+str(new),cell.value)
validations=[]
for dv in ws.data_validations.dataValidation:
 nd=copy(dv); nd.sqref=' '.join(re.sub(r':([A-Z]+)'+str(oldlast)+r'\b',lambda m:':'+m[1]+str(new),str(rng)) for rng in dv.sqref.ranges); validations.append(nd)
ws.data_validations.dataValidation=validations
ws.auto_filter.ref=f'A{HEADER_ROW}:Y{new}'
log(vwb,'v0.1.8','新增资产主表','特效 · 资产主表','新增VFX-RADIO-MUSIC-NOTES-3D，查重/总览/DV扩至新行；纯视觉radio附件，不进池。')
pwb.save(P); vwb.save(V)
count=reconcile({ID,RID},[])
# 事务2：域专表登记独立PackedScene，并只更新这两张专表的摘要。
tx=O/'prefab_transaction'; tx.mkdir(exist_ok=True)
for path in [P,V,B]: shutil.copy2(path,tx/path.name)
pwb=openpyxl.load_workbook(P); vwb=openpyxl.load_workbook(V)
ws=pwb['3D-道具']; assert ws['A8'].value==RID
ws['O8']='v005.1'; ws['K8']=str(ws['K8'].value)+' runtime补丁：'+SPEC
ws['P8']=str(ws['P8'].value)+'；runtime v005.1，model v005；附件='+ID+'；Prefab SHA='+sha(R/RP)
log(pwb,'v0.1.14','更新Prefab专表','道具 · 3D-道具','收音机runtime v005.1规格与Prefab哈希；源/优化/GLB/灯帽未改。专表独立事务。')
ws=vwb['3D-特效']; pr=max(c.row for c in ws['B'] if c.value)+1
assert ws.cell(pr,1).value is None
pvals=['FX03-03',ID,'收音机音乐上飘音符',FX,'无（预制QuadMesh/ImageTexture）','无（Godot资源生成）',SPEC,'src/vfx/VfxRadioMusicNotes3D.gd','关','无：纯视觉附件','无CollisionObject3D/CollisionShape3D','每音符0.34×0.38m；上飘1.1m','radio本地Y=1.42m；billboard当前相机','收音机MusicNotes子节点','正式美术已接入','v001','独立PackedScene；radio唯一持有，不进池；验收 verify_base99_radio_music_notes。']
for col,val in enumerate(pvals,1): ws.cell(pr,col).value=val; ws.cell(pr,col)._style=copy(ws.cell(pr-1,col)._style)
ws.row_dimensions[pr].height=ws.row_dimensions[pr-1].height
log(vwb,'v0.1.9','新增Prefab专表','特效 · 3D-特效','登记FX03-03收音机音符独立PackedScene，纯视觉无碰撞；专表独立事务更新摘要。')
pwb.save(P); vwb.save(V)
reconcile(set(),['3D-道具','3D-特效'])
# 全工作簿逐格比对，任何未授权差异立即失败。
diffs={}
allowed_props={'M27','N27','P27','T27','V27','Y27'}
allowed_prefab={'O8','K8','P8'}
for path in [P,V]:
 before=openpyxl.load_workbook(O/'backup'/path.relative_to(R)); after=openpyxl.load_workbook(path); changes=[]
 log_last=before['域变更日志'].max_row
 for oldws in before:
  nws=after[oldws.title]; assert str(oldws.merged_cells)==str(nws.merged_cells)
  for row in nws:
   for cell in row:
    old=oldws[cell.coordinate]
    if old.value==cell.value: continue
    allow=False
    if path==P:
     allow=(oldws.title=='资产主表' and cell.coordinate in allowed_props) or (oldws.title=='3D-道具' and cell.coordinate in allowed_prefab) or (oldws.title=='域变更日志' and cell.row>before['域变更日志'].max_row)
    else:
     allow=(oldws.title=='资产主表' and (cell.row==new or cell.column==19)) or (oldws.title=='总览' and cell.coordinate in ['A6','C6','E6','G6','B10','C10','B11','C11']) or (oldws.title=='3D-特效' and cell.row==pr) or (oldws.title=='域变更日志' and cell.row>before['域变更日志'].max_row)
    assert allow,(path.name,oldws.title,cell.coordinate)
    changes.append({'sheet':oldws.title,'cell':cell.coordinate,'before':old.value,'after':cell.value})
 diffs[path.name]=changes
 with zipfile.ZipFile(path) as z:
  dvs=[ET.fromstring(z.read(n)).findall('{*}dataValidations/{*}dataValidation') for n in z.namelist() if re.fullmatch(r'xl/worksheets/sheet\d+.xml',n)]
  if path==V: assert {'C6:C23','K6:K23','L6:L23'}.issubset({d.attrib['sqref'] for group in dvs for d in group})
# 所有既有红项文件和哈希单元格不变。
debt=[]
for domain,path in [('props',P),('vfx',V)]:
 rep=json.JSONDecoder().raw_decode((O/f'testlogs/{domain}_before.log').read_text(encoding='utf-8').lstrip())[0]
 before=openpyxl.load_workbook(O/'backup'/path.relative_to(R)); after=openpyxl.load_workbook(path)
 for issue in rep['issues'].get('sha_mismatch',[]):
  assert sha(Path(issue['path']))==issue['actual']
  assert before['资产主表'].cell(issue['row'],20).value==after['资产主表'].cell(issue['row'],20).value==issue['recorded']
  debt.append(issue)
dump(O/'ledger_evidence.json',{'passed':True,'new_asset_id':ID,'vfx_asset_row':new,'vfx_prefab_row':pr,'radio_asset_row':27,'radio_prefab_row':8,'vfx_count_before':17,'vfx_count_after':18,'baseline_count':count,'differences':diffs,'existing_sha_debt_unchanged':debt,'two_transactions':True,'DV_zip_verified':True,'other_asset_fingerprints_unchanged':True})
print('RADIO_MUSIC_NOTES_LEDGER_OK',count,new,pr)
