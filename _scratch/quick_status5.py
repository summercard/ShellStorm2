import openpyxl
p = r"I:\工作项目\shellstrom2\ShellStorm2\assets\registry\ShellStorm2_美术资产台账_v001.xlsx"
wb = openpyxl.load_workbook(p, data_only=True)
sheets = ["3D-场景通用", "3D-设施", "3D-道具", "3D-角色", "3D-敌人", "3D-武器", "3D-特效"]
KEY = ["未接入", "待玩法映射", "待补", "待拆分"]
for s in sheets:
    ws = wb[s]
    for r in range(5, ws.max_row + 1):
        aid = ws.cell(r, 1).value
        if not aid or not str(aid).strip():
            continue
        st = str(ws.cell(r, 14).value or "")
        if any(k in st for k in KEY):
            print(f"{s:<8} r{r:<4} {str(aid)[:34]:<36} {str(ws.cell(r,2).value or '')[:22]:<24} {st}")
