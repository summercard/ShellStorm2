from pathlib import Path
import json,hashlib,shutil,sys
from openpyxl import load_workbook
R=Path.cwd();B=R/'assets/art/enemies/bosses/enm_boss_monitor002';D=R/'_scratch/monitor035_ledger';D.mkdir(exist_ok=True)
for p in [B/'runtime/enm_boss_monitor002/enm_boss_monitor002_root_top3d.tscn',R/'src/enemy3d/MonitorBossPresentation.gd']:
 s=p.read_text(encoding='utf-8').replace('v034','v035');p.write_text(s,encoding='utf-8')
sys.path.insert(0,str(R/'scripts'));from ledger_registry import LedgerIndex
p=LedgerIndex.load(R).path_for_category('敌人');w=load_workbook(p);A='ENM-BOSS-MONITOR002-3D'
dual='; '.join((B/f'source/enm_boss_monitor002_{k}_v035.blend').relative_to(R).as_posix() for k in ['model','animation'])
note='v035仅屏幕总成沿角色前向平移0.5米，排除下沿支杆穿插。外壳、代码与五官同步；其它骨骼矩阵、旋转/缩放、17剪辑时长、双手、特效及10米激活规则均不变。562项专项通过；真实出场214帧验证。'
states={}
for row in w['敌人动画与状态']:
 if row[0].value==A:states[f'E{row[0].row}']=dual
lr=w['域变更日志'].max_row+1
jobs=[{'path':str(p),'sheets':{'资产主表':{'M20':'v035','P20':dual+'；SKEL-MONITOR002-005','T20':hashlib.sha256((B/'runtime/enm_boss_monitor002/enm_boss_monitor002_root_top3d.tscn').read_bytes()).hexdigest(),'Y20':note},'3D-敌人':{'E10':dual,'O10':'v035','P10':note},'敌人动画与状态':states,'域变更日志':{f'{c}{lr}':v for c,v in zip('ABCDEFG',['v035','2026-10-08','Codex',A,note,'docs/v0.1/development/2026-10-08_boss002_screen_clearance.md','仅屏幕前移'])}}}]
pp=B/'source/boss002_production_ledger.xlsx';sn=load_workbook(pp).sheetnames[0]
jobs.append({'path':str(pp),'sheets':{sn:{'B4':'v035 仅屏幕前移避开支杆；17剪辑/十二态/四技能/三阶段','B20':'v035','C20':dual,'C21':'source/rig_contract_v035.json','B26':'previews/screen_v035/source_audit.json','C26':note}}})
(D/'changes.json').write_text(json.dumps(jobs,ensure_ascii=False,indent=2),encoding='utf-8');shutil.copy2(R/'assets/registry/ledger_split_baseline.json',D/'baseline_before.json')
for name in ['edit.mjs','merge.py']:
 s=(R/'_scratch/monitor034_ledger'/name).read_text(encoding='utf-8').replace('_scratch/monitor034_ledger','_scratch/monitor035_ledger');(D/name).write_text(s,encoding='utf-8')
p=R/'docs/v0.1/design/Boss002显示器动画设计.md';s=p.read_text(encoding='utf-8').replace('状态：v034','状态：v035');s=s.replace('## 1. 角色动作语言','屏幕间距修正：v035屏幕总成沿角色前向平移0.5米，外壳、代码和五官同步避开后支杆；其它骨骼姿态、时序、特效及玩法保持不变。\n\n## 1. 角色动作语言');p.write_text(s,encoding='utf-8')
p=B/'README.md';s=p.read_text(encoding='utf-8');p.write_text('## 当前屏幕间距修正 v035\n\n仅屏幕总成前移0.5米排除支杆穿插，其余设计、动作时序和10米激活逻辑不变。双母版/契约v035，稳定GLB不变，采样仅屏幕平移。预览 `previews/screen_v035/monitor_activation.mp4`。\n\n'+s,encoding='utf-8')
p=R/'docs/v0.1/MODULE_INDEX.md';s=p.read_text(encoding='utf-8').replace('v034正式Prefab','v035正式Prefab');p.write_text(s,encoding='utf-8')
p=R/'docs/v0.1/development/CHANGELOG.md';s=p.read_text(encoding='utf-8');p.write_text('## 2026-10-08 Boss002 屏幕支杆间距修正\n\n仅屏幕前移，其他设计不变。[验证记录](2026-10-08_boss002_screen_clearance.md)。\n\n'+s,encoding='utf-8')
print('SCREEN_CLEARANCE_METADATA_PREPARED')
