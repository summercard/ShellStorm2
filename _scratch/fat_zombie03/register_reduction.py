import sys,json,hashlib,shutil,subprocess
from pathlib import Path
from datetime import datetime
root=Path(__file__).resolve().parents[2];pkg=root/'assets/art/enemies/normal_enemy_3d/fat_zombie03';out=root/'outputs/fat_zombie03_reduction';aid='ENM-NORMAL-FAT-ZOMBIE03'
sys.path[:0]=[str(root/'scripts'),str(root/'tools/asset_pipeline')]
from ledger_registry import LedgerIndex
from split_asset_ledger import read_source_rows,_row_digest,sheet_digest,CONTENT_COLUMNS,col_digest
from openpyxl import load_workbook
index=LedgerIndex.load(root);ledger=index.domain_for_key('enemies').path
sha=lambda p:hashlib.sha256(p.read_bytes()).hexdigest()
rel=lambda p:p.relative_to(root).as_posix()
model=pkg/'source/model/enm_normal_fat_zombie03_model_v004.blend';animation=pkg/'source/animation/enm_normal_fat_zombie03_animation_v009.blend';prefab=pkg/'runtime/enm_normal_fat_zombie03_root_top3d.tscn'
metrics=json.loads((out/'source_comparison.json').read_text());export=json.loads((out/'export.json').read_text())
note='2026-10-09 v012减面：5690→1979三角（减少65.22%）；脸手特征加权保留，512贴图/66骨/13段动画保留；死亡接地形态键重新拟合。模型v004/动作v009；Godot153项及死亡视野真实渲染通过。'
w=load_workbook(ledger);before={v[0]:v for _,v in read_source_rows(w['资产主表'])};ws=w['资产主表'];r=next(r for r,v in read_source_rows(ws) if v[0]==aid)
for c,v in {13:'v012',14:f'源高3.142857m；游戏站立高2.2m；{metrics["vertices"]}顶点/1979三角；66骨；512²贴图；13段动作',16:rel(model)+'; '+rel(animation)+'; docs/v0.1/design/胖子僵尸03动作设计.md',20:sha(prefab),22:datetime(2026,10,9),25:str(ws.cell(r,25).value)+'；'+note}.items():ws.cell(r,c).value=v
log=w['域变更日志'];versions=[str(log.cell(i,1).value).lstrip('v').split('.') for i in range(1,log.max_row+1)];versions=[v for v in versions if len(v)==3 and all(x.isdigit() for x in v)];v=list(max(versions,key=lambda x:tuple(map(int,x))));v[-1]=str(int(v[-1])+1);log.append(['v'+'.'.join(v),'2026-10-09','胖子僵尸减面重新导入','敌人',note,'仅目标资产；动作/状态机/碰撞契约保留','Codex']);w.save(ledger)
w=load_workbook(ledger)
for row in w['3D-敌人'].iter_rows():
 if row[0].value==aid:
  row[4].value=rel(model)+'; '+rel(animation);row[10].value=f'游戏高2.2m；1979三角/{metrics["vertices"]}顶点；66骨；512贴图';row[14].value='v012';row[15].value=str(row[15].value)+'；'+note
for row in w['敌人动画与状态'].iter_rows():
 if row[0].value==aid:row[4].value=rel(animation)
w.save(ledger)
after={v[0]:v for _,v in read_source_rows(w['资产主表'])};assert all(before[k]==after[k] for k in before if k!=aid)
base=root/'assets/registry/ledger_split_baseline.json';bl=json.loads(base.read_text(encoding='utf-8'));bl['assets'][aid]={'v':_row_digest(after[aid]),'c':'敌人','d':'enemies'}
for n in ['3D-敌人','敌人动画与状态']:bl['sheet_digests'][n]=sheet_digest(w[n])
union=[]
for domain in index.domains:union.extend(read_source_rows(load_workbook(domain.path)['资产主表']))
bl['column_digests']={str(c):col_digest(union,c) for c in CONTENT_COLUMNS};base.write_text(json.dumps(bl,ensure_ascii=False,indent=1)+'\n',encoding='utf-8')
tf=pkg/'runtime/character_transfer_ledger.json';t=json.loads(tf.read_text(encoding='utf-8'));t.update(version='v012',model_source=rel(model),animation_source=rel(animation),notes=note,status='active')
for c in t['clips']:c['source']=rel(animation);c['source_validation']=rel(out/'source_comparison.json')
paths=[root/x['path'] for x in t['files'] if '/source/animation/' not in x['path'] and not x['path'].endswith('model_v003.blend')]+[model,animation]
t['files']=[{'path':rel(p),'sha256':sha(p),'bytes':p.stat().st_size} for p in paths]
t['geometry']={'triangles':1979,'vertices':metrics['vertices'],'previous_triangles':5690,'reduction_percent':round((1-1979/5690)*100,2)}
t['validation'].append({'test':'reduction_20261009','passed':True,'checks':153,'source':rel(out/'source_comparison.json'),'export':rel(out/'export.json'),'runtime_log':rel(out/'runtime.log'),'render_log':rel(out/'death_render.log')})
tf.write_text(json.dumps(t,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
for name,target in [('reduce_verify.log','runtime.log'),('reduce_death_render.log','death_render.log'),('reduce_import.log','import.log')]:shutil.copy2(root/'_scratch/fat_zombie03'/name,out/target)
for name in ['falling','impact','bounce','settled']:shutil.copy2(root/f'outputs/fat_zombie03_death_fix/{name}.png',out/f'godot_{name}.png')
record='docs/v0.1/development/2026-10-09_fat_zombie03_reduction.md'
(root/record).write_text('# 胖子僵尸03减面与重新导入\n\nENEMY-AI / ASSET-PIPELINE。用户要求保持外观并降至2000面以下。\n\n'+note+'\n\n拓扑使用脸部与手指优先的加权简化，UV与骨权重插值后归一为最多4影响；腹部压扁形态键插值后以626个死亡姿势采样拟合接地。13段骨动作曲线、GLB动画数据、骨架层级/绑定及材质未改变。仅替换视觉网格及形态键，Enemy3D继续拥有玩法状态、碰撞和回收。\n\n外观验收为正面、侧面、顶视跑步、拍合和倒地六组同机位对照；主要体型与张手剪影保留。减面有局部几何差异，不宣称逐顶点无损。源级表面误差见outputs/fat_zombie03_reduction/source_comparison.json；近景可见细小轮廓简化。\n\nGodot重新导入退出0；153项正式运行检查退出0；死亡视野真实渲染退出0，日志无非预期脚本错误。未运行全项目套件。资产仍使用原稳定GLB/Prefab路径，版本v012，旧母版保留；回滚可从Git恢复运行资产。全局账本/文档/命名门禁结果见同输出目录gates.json，与修改前敌人账本问题作对照。\n',encoding='utf-8')
f=root/'docs/v0.1/feature_registry.json';lines=f.read_text(encoding='utf-8').splitlines(keepends=True)
for i,line in enumerate(lines):
 if '"feature_id":"ENEMY-AI"' in line or '"feature_id":"ASSET-PIPELINE"' in line:
  d=json.loads(line.strip().rstrip(','));d['development_records'].append(record);lines[i]='    '+json.dumps(d,ensure_ascii=False,separators=(',',':'))+',\n'
f.write_text(''.join(lines),encoding='utf-8')
for name,text in [('MODULE_INDEX.md','\nENEMY-AI / ASSET-PIPELINE [胖子僵尸减面](development/2026-10-09_fat_zombie03_reduction.md)：1979三角，保留13段动作，重新导入与死亡渲染通过。\n'),('development/CHANGELOG.md','\n## 2026-10-09｜胖子僵尸减面\n\n5690降至1979三角；模型v004/动作v009、运行v012，重新导入及153项运行检查通过。见[记录](2026-10-09_fat_zombie03_reduction.md)。\n')]:
 f=root/'docs/v0.1'/name;f.write_text(f.read_text(encoding='utf-8')+text,encoding='utf-8')
commands={'structure':['scripts/check_asset_registry.py','--scope','structure'],'full':['scripts/check_asset_registry.py','--scope','full','--ledger','enemies','--json-output',str(out/'ledger_after.json')],'split':['tools/asset_pipeline/verify_ledger_split.py'],'guard':['scripts/asset_guard.py',rel(pkg),'--classify','version_increment'],'docs':['scripts/check_documentation_contracts.py'],'naming':['scripts/check_asset_runtime_naming.py'],'refs':['scripts/check_ledger_refs.py']}
results={}
for name,args in commands.items():
 r=subprocess.run([sys.executable,*args],cwd=root,capture_output=True);results[name]=r.returncode;(out/f'gate_{name}.log').write_bytes(r.stdout+r.stderr);print(name,r.returncode,flush=True)
before=json.loads((root/'_scratch/fat_zombie03/reduce_ledger_before.json').read_text(encoding='utf-8'));after=json.loads((out/'ledger_after.json').read_text(encoding='utf-8'));assert before['issues']==after['issues'];results['existing_ledger_issues_unchanged']=True
(out/'gates.json').write_text(json.dumps(results,indent=2));assert all(results[k]==0 for k in ['structure','split','guard'])
print('REDUCTION_REGISTERED',flush=True)
