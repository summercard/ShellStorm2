import json,sys,shutil,subprocess
from pathlib import Path
from datetime import datetime
from openpyxl import load_workbook
root=Path.cwd();p=root/'_scratch/fat_zombie03';out=root/'outputs/fat_zombie03_death_fix';backup=p/'death_fix_backup';backup.mkdir(exist_ok=True)
sys.path[:0]=[str(root/'scripts'),str(root/'tools/asset_pipeline')]
from ledger_registry import LedgerIndex
from split_asset_ledger import read_source_rows,_row_digest,sheet_digest,CONTENT_COLUMNS,col_digest
index=LedgerIndex.load(root);ledger=index.domain_for_key('enemies').path;baseline=root/'assets/registry/ledger_split_baseline.json';aid='ENM-NORMAL-FAT-ZOMBIE03'
transfer=root/'assets/art/enemies/normal_enemy_3d/fat_zombie03/runtime/character_transfer_ledger.json'
for f in [ledger,baseline,transfer]:
 if not (backup/f.name).exists():shutil.copy2(f,backup/f.name)
assert all(json.loads((out/(name+'.json')).read_text())['passed'] for name in ['logic','render'])
for name in ['death_fix','death_render','death_negative','death_verify_fat_zombie03','death_verify_3d_vision_input_flow','death_verify_melee_zombie_presentation']:shutil.copy2(p/(name+'.log'),out/(name+'.log'))
note='2026-10-04修复游戏内死亡隐藏：尸体退出AI后独立视野组继续判定距离/遮挡，表现帧计时推进2.6秒倒地/腹部接地/回弹后回收；AI休眠不暂停死亡。真实Player3D视野逻辑与Forward+四帧截图通过，负向原失踪条件检出；胖子152项/小僵尸死亡/玩家视野回归通过。'
w=load_workbook(ledger);ws=w['资产主表'];r=next(r for r,v in read_source_rows(ws) if v[0]==aid);ws.cell(r,25).value=str(ws.cell(r,25).value)+'；'+note;ws.cell(r,22).value=datetime(2026,10,4)
log=w['域变更日志'];versions=[str(log.cell(i,1).value).lstrip('v').split('.') for i in range(1,log.max_row+1)];versions=[v for v in versions if len(v)==3 and all(x.isdigit() for x in v)];v=list(max(versions,key=lambda x:tuple(map(int,x))));v[-1]=str(int(v[-1])+1);log.append(['v'+'.'.join(v),'2026-10-04','死亡表现视野与回收修复','敌人',note,'美术源/导出不变；Owner Enemy3D/PlayerVision3D','Codex']);w.save(ledger)
w=load_workbook(ledger)
for sheet in ['3D-敌人','敌人动画与状态']:
 for row in w[sheet].iter_rows():
  if row[0].value==aid and (sheet=='3D-敌人' or row[1].value=='dead'):row[-1].value=str(row[-1].value)+'；'+note
w.save(ledger)
bl=json.loads(baseline.read_text(encoding='utf-8'));values=next(v for _,v in read_source_rows(w['资产主表']) if v[0]==aid);bl['assets'][aid]={'v':_row_digest(values),'c':'敌人','d':'enemies'}
for name in ['3D-敌人','敌人动画与状态']:bl['sheet_digests'][name]=sheet_digest(w[name])
union=[]
for domain in index.domains:union.extend(read_source_rows(load_workbook(domain.path)['资产主表']))
bl['column_digests']={str(c):col_digest(union,c) for c in CONTENT_COLUMNS};baseline.write_text(json.dumps(bl,ensure_ascii=False,indent=1)+'\n',encoding='utf-8')
old=load_workbook(backup/ledger.name);a={v[0]:v for _,v in read_source_rows(old['资产主表'])};b={v[0]:v for _,v in read_source_rows(w['资产主表'])};assert all(a[k]==b[k] for k in a if k!=aid)
t=json.loads(transfer.read_text(encoding='utf-8'));t['validation'].append({'test':'verify_fat_zombie03_death_visibility','date':'2026-10-04','passed':True,'logic':str(out.relative_to(root)/'logic.json'),'render':str(out.relative_to(root)/'render.json'),'negative_expected_exit':1});t['runtime_binding']['death_visibility_group']='enemy_death_visual_3d';t['runtime_binding']['death_clock_owner']='Enemy3D._process';transfer.write_text(json.dumps(t,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
record='docs/v0.1/development/2026-10-04_fat_zombie03_death_visibility.md'
f=root/'docs/v0.1/feature_registry.json';lines=f.read_text(encoding='utf-8').splitlines(keepends=True)
for i,line in enumerate(lines):
 if '"feature_id":"ENEMY-AI"' not in line and '"feature_id":"ASSET-PIPELINE"' not in line:continue
 d=json.loads(line.strip().rstrip(','));d['development_records'].append(record);d['verification'].append({'kind':'scene','id':'verify_fat_zombie03_death_visibility'});lines[i]='    '+json.dumps(d,ensure_ascii=False,separators=(',',':'))+',\n'
f.write_text(''.join(lines),encoding='utf-8')
f=root/'docs/v0.1/MODULE_INDEX.md';f.write_text(f.read_text(encoding='utf-8')+'\nENEMY-AI / ASSET-PIPELINE [死亡视野修复](development/2026-10-04_fat_zombie03_death_visibility.md)：死亡退出AI后仍由独立组参与视野遮挡，表现帧推进倒地至回收；真实Player3D视野与渲染、原失踪条件负向验证。\n',encoding='utf-8')
f=root/'docs/v0.1/development/CHANGELOG.md';s=f.read_text(encoding='utf-8');f.write_text(s.replace('# 游戏设计文档 v0.1 变更记录','# 游戏设计文档 v0.1 变更记录\n\n## 2026-10-04｜死亡动画游戏内被隐藏修复\n\n- 尸体退出AI空间桶后继续由独立视野组管理，避免直接消失；表现时钟独立于物理AI，完整播放2.6秒后回收。\n- 玩家视野、真实渲染、负向及普通怪专项回归通过；见[修复记录](2026-10-04_fat_zombie03_death_visibility.md)。',1),encoding='utf-8')
f=root/record;f.write_text(f.read_text(encoding='utf-8').replace('ENEMY-AI / PLAYER-VISION','ENEMY-AI / ASSET-PIPELINE'),encoding='utf-8')
results={}
for key,args in {'structure':['scripts/check_asset_registry.py','--scope','structure'],'split':['tools/asset_pipeline/verify_ledger_split.py'],'docs':['scripts/check_documentation_contracts.py'],'naming':['scripts/check_asset_runtime_naming.py'],'log':['scripts/check_verification_log.py',str(out/'death_render.log')],'negative_log':['scripts/check_verification_log.py',str(out/'death_negative.log'),'tests/verification/expected_errors/fat_zombie03_death_visibility_negative.txt']}.items():
 r=subprocess.run([sys.executable,*args],cwd=root,capture_output=True);results[key]=r.returncode;(out/f'gate_{key}.log').write_bytes(r.stdout+r.stderr);print(key,r.returncode,flush=True)
results.update(logic=0,real_render=0,negative_expected_exit=1,other_assets_unchanged=True)
(out/'gates.json').write_text(json.dumps(results,indent=2)+'\n',encoding='utf-8');assert all(results[k]==0 for k in ['structure','split','log','negative_log'])
print('DEATH_FIX_REGISTERED')
