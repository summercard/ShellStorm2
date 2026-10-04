import json,hashlib,sys,shutil
from pathlib import Path
from datetime import datetime
from copy import copy
from openpyxl import load_workbook
root=Path.cwd();p=root/'_scratch/fat_zombie03';out=root/'outputs/fat_zombie03_integration';pkg=root/'assets/art/enemies/normal_enemy_3d/fat_zombie03';aid='ENM-NORMAL-FAT-ZOMBIE03'
sys.path[:0]=[str(root/'scripts'),str(root/'tools/asset_pipeline')]
from ledger_registry import LedgerIndex
from split_asset_ledger import read_source_rows,_row_digest,sheet_digest,CONTENT_COLUMNS,col_digest
index=LedgerIndex.load(root);ledger=index.domain_for_key('enemies').path;baseline=root/'assets/registry/ledger_split_baseline.json';content=root/'docs/v0.1/data/ShellStorm2_游戏内容数据库_v010.xlsx'
backup=p/'runtime_binding_backup';backup.mkdir(exist_ok=True)
for f in [ledger,baseline,content,pkg/'runtime/character_transfer_ledger.json']:
 if not (backup/f.name).exists(): shutil.copy2(f,backup/f.name)
assert json.loads((out/'runtime_verification.json').read_text())['passed']
assert json.loads((out/'negative_loop.json').read_text())['failures']==['loop attack']
for log in ['runtime_test','negative_loop','render_forward','preflight_godot','verify_melee_zombie_presentation','verify_security_zombie_presentation','verify_monster_ai_system_complete','verify_enemy_illumination_states']:
 shutil.copy2(p/(log+'.log'),out/(log+'.log'))
sha=lambda f:hashlib.sha256(f.read_bytes()).hexdigest()
rel=lambda f:f.relative_to(root).as_posix()
scene=pkg/'runtime/enm_normal_fat_zombie03_root_top3d.tscn'
s=scene.read_text(encoding='utf-8').replace('"v010"','"v011"').replace('runtime_binding_pending_verification','runtime_verified')
scene.write_text(s,encoding='utf-8')
design=root/'docs/v0.1/design/胖子僵尸03动作设计.md';s=design.read_text(encoding='utf-8')
s=s.replace('玩法状态、转向与局部混合适配待接入。','玩法状态、转向与局部混合已接入，见[运行配置](胖子僵尸03运行配置.md)。')
s=s.replace('以上是设计待接入值，未写运行配置。','已接入运行配置：基准生命696、伤害19、判定1.55m，完整周期约2.5s。')
assert len(s)<5000;design.write_text(s,encoding='utf-8')
note='编号03；v011正式绑定Enemy3D共用12态与13段动画，独立类型fat_zombie03；标准2.2m/512贴图，生命696伤害19，巡逻0.30/追击0.60m/s；F36拍合单次命中、F30锁朝向、视线/范围/扇区判定；普通命中仅上半身，死亡2.6s保留腹部压扁及回弹。专项150项、负向循环故障检出及真实Forward+四图通过。'
tpath=pkg/'runtime/character_transfer_ledger.json';t=json.loads(tpath.read_text(encoding='utf-8'))
t.update(version='v011',status='active',validation_scope='source_and_godot_13_clips; Enemy3D_12_states; spawning_combat_death; real_forward_plus_render',notes=note,validation_status='passed')
t['runtime_binding']={'owner':'Enemy3D','content_id':'fat_zombie03','entity_scene':'scenes/enemies/fat_zombie03.tscn','adapter':rel(pkg/'runtime/fat_zombie03_formal_visual.gd'),'configuration':'docs/v0.1/design/胖子僵尸03运行配置.md','spawn_owner':'MonsterInjector / Dungeon3D','state_map':{'dormant':'idle','idle':'idle / awaken / move_stop / turn_l / turn_r','patrol':'move_start / walking','alert':'alert','chase':'running','search':'running','return':'running','telegraph':'attack[0,1.2]','attack':'attack@1.2','recovery':'attack[1.2,2.5]','stagger':'hurt','dead':'dead'},'ordinary_hit':'hit_light upper bone replacement','death_duration_s':2.6,'collision_owner':'Enemy3D'}
t['consumer']='EnemyAvatar3D formal-normal routing; MonsterInjector explicit type / fixed waves / triggers'
t['validation']=[{'test':'verify_fat_zombie03','passed':True,'checks':150,'report':rel(out/'runtime_verification.json'),'log':rel(out/'runtime_test.log')},{'test':'negative_loop','passed':True,'expected_exit':1,'report':rel(out/'negative_loop.json')},{'test':'real_renderer','passed':True,'renderer':'forward_plus','log':rel(out/'render_forward.log'),'screenshots':[rel(out/(x+'.png')) for x in ['idle','run_top','clap','dead_contact']]}]
t['action_design']['sha256']=sha(design);t['action_design']['character_count']=len(s)
paths=[root/e['path'] for e in t['files']]+[pkg/'runtime/fat_zombie03_formal_visual.gd',root/'scenes/enemies/fat_zombie03.tscn']
t['files']=[{'path':rel(f),'sha256':sha(f)} for f in dict.fromkeys(paths)]
t['remaining_work']=['未加入已有随机关卡池；可在触发盒/固定波次显式调用。','脚步与拍合使用共用命中音效；专属新音效未制作。']
tpath.write_text(json.dumps(t,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
w=load_workbook(ledger);ws=w['资产主表'];row=next(r for r,v in read_source_rows(ws) if v[0]==aid)
for c,v in {9:'12态绑定13段动画；idle/walking/running循环，其他单次；hit_light上半身局部混合',11:'正式美术已接入',13:'v011',16:t['model_source']+'; '+t['animation_source']+'; docs/v0.1/design/胖子僵尸03运行配置.md',20:sha(scene),21:'Codex',22:datetime(2026,10,3),25:note}.items():ws.cell(row,c).value=v
log=w['域变更日志'];versions=[str(log.cell(i,1).value).lstrip('v').split('.') for i in range(1,log.max_row+1)];versions=[x for x in versions if len(x)==3 and all(y.isdigit() for y in x)];v=list(max(versions,key=lambda x:tuple(map(int,x))));v[-1]=str(int(v[-1])+1);log.append(['v'+'.'.join(v),'2026-10-03','胖子僵尸03正式运行配置','敌人',note,'新增kind复用FSM；既有七类玩法默认值保留','Codex']);w.save(ledger)
# 专表独立事务，保留主表和全部其他资产。
w=load_workbook(ledger);ws=w['3D-敌人'];r=next(r for r in range(5,ws.max_row+1) if ws.cell(r,1).value==aid)
for c,v in {6:'正式厚血慢速近战；12态输出单向驱动13段动画',7:'src/enemy3d/Enemy3D.gd; src/enemy3d/EnemyAvatar3D.gd; '+rel(pkg/'runtime/fat_zombie03_formal_visual.gd'),8:'视觉包关闭；实体开启',9:'Enemy3D',10:'玩法CylinderShape3D；视觉包无碰撞',11:'站立2.2m；世界圆柱半径0.77m/高2.2m；66骨；512贴图',13:'scenes/enemies/fat_zombie03.tscn；Dungeon3D固定波次/触发盒',14:'正式美术已接入',15:'v011',16:note}.items():ws.cell(r,c).value=v
ws=w['敌人动画与状态']
for r in range(6,ws.max_row+1):
 if ws.cell(r,1).value!=aid:continue
 name=ws.cell(r,2).value
 ws.cell(r,8).value='已制作/导出/12态绑定/运行验收通过'
 ws.cell(r,9).value=note
 ws.cell(r,7).value={'attack':'Enemy3D：F30锁朝向，F36单次伤害；Root无突进','dead':'死亡2.6秒回收；腹部形变/落地回弹保留','hit_light':'普通命中触发；Waist以上旋转局部替换，保留腿部动作','awaken':'dormant→idle','alert':'alert 0.6秒','turn_l':'idle实际左转触发','turn_r':'idle实际右转触发','move_start':'进入patrol；0.4秒接walking','move_stop':'patrol→idle；0.6秒接idle','walking':'仅patrol；步频按实际速度采样','running':'chase/search/return；按实际速度采样','hurt':'重击stagger 0.8秒','idle':'idle循环；dormant保持首帧'}[name]
w.save(ledger)
bl=json.loads(baseline.read_text(encoding='utf-8'));values=next(v for _,v in read_source_rows(w['资产主表']) if v[0]==aid);bl['assets'][aid]={'v':_row_digest(values),'c':'敌人','d':'enemies'}
for name in ['3D-敌人','敌人动画与状态']:bl['sheet_digests'][name]=sheet_digest(w[name])
union=[]
for d in index.domains:union.extend(read_source_rows(load_workbook(d.path)['资产主表']))
bl['column_digests']={str(c):col_digest(union,c) for c in CONTENT_COLUMNS};baseline.write_text(json.dumps(bl,ensure_ascii=False,indent=1)+'\n',encoding='utf-8')
before=load_workbook(backup/ledger.name);a={v[0]:v for _,v in read_source_rows(before['资产主表'])};b={v[0]:v for _,v in read_source_rows(w['资产主表'])};assert all(a[k]==b[k] for k in a if k!=aid)
wb=load_workbook(content);ws=wb['怪物与Boss'];assert not any(ws.cell(r,1).value=='monster_fat_zombie03' for r in range(5,ws.max_row+1));new=ws.max_row+1
values=['monster_fat_zombie03','胖子僵尸','fat_zombie03',232,0.857142857,19,1.55,1.3,'厚血慢速；张臂抱扑拍合；1.2秒前摇/1.3秒收势；轻击上半身反应','显式固定波次/触发盒/独立场景','通用楼层掉落表；killed一次结算','[已实装]','正式独立怪物，12态/13动画绑定','src/enemy3d/Enemy3D.gd；src/map/MonsterInjector.gd；docs/v0.1/design/胖子僵尸03运行配置.md','D/E为PROFILES源基线；实际第一层生命696、追击0.60m/s、巡逻0.30m/s；内容调用ID=fat_zombie03；资产编号03；未加入既有随机池']
for c,val in enumerate(values,1):ws.cell(new,c).value=val;ws.cell(new,c)._style=copy(ws.cell(8,c)._style)
ws.row_dimensions[new].height=ws.row_dimensions[8].height
if ws.auto_filter.ref:ws.auto_filter.ref=f'A4:O{new}'
wb.save(content)
old=load_workbook(backup/content.name);assert all(list(old[n].values)==list(wb[n].values) for n in old.sheetnames if n!='怪物与Boss');assert list(old['怪物与Boss'].values)==list(wb['怪物与Boss'].values)[:-1]
print('FAT_ZOMBIE03_RUNTIME_LEDGER_OK other_assets_and_content_unchanged=true row=',row,'content_row=',new)
