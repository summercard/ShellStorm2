import sys,json,hashlib,shutil
from pathlib import Path
from datetime import datetime
from openpyxl import load_workbook
p=Path(__file__).parent;root=p.parents[1];pkg=root/'assets/art/enemies/normal_enemy_3d/fat_zombie03';aid='ENM-NORMAL-FAT-ZOMBIE03'
sys.path.insert(0,str(root/'scripts'));sys.path.insert(0,str(root/'tools/asset_pipeline'))
from ledger_registry import LedgerIndex
from split_asset_ledger import read_source_rows,_row_digest,sheet_digest,CONTENT_COLUMNS,col_digest
index=LedgerIndex.load(root);ledger=index.domain_for_key('enemies').path;baseline=root/'assets/registry/ledger_split_baseline.json';backup=p/'idle_walk_v002_backup'
for f,name in [(ledger,'ledger_before_walk.xlsx'),(baseline,'baseline_before_walk.json'),(pkg/'runtime/character_transfer_ledger.json','transfer_before_walk.json')]:
 assert not (backup/name).exists(),'Backup already exists'
 shutil.copy2(f,backup/name)
rel=lambda f:f.relative_to(root).as_posix();sha=lambda f:hashlib.sha256(f.read_bytes()).hexdigest()
report=json.loads((pkg/'previews/idle_walk_v002/validation.json').read_text(encoding='utf-8'))
godot=json.loads((pkg/'previews/idle_walk_v002/godot_validation.json').read_text(encoding='utf-8'));assert godot['passed']
animation=pkg/'source/animation/enm_normal_fat_zombie03_animation_v002.blend';model=pkg/'source/model/enm_normal_fat_zombie03_model_v002.blend';prefab=pkg/'runtime/enm_normal_fat_zombie03_root_top3d.tscn'
note='编号03；v004修正右手指姿态，双手自然张开；idle 3.2s与walking 2.0s均30fps循环并验收。walking短步承重、抬脚6cm、参考速度0.30m/s，Root固定。模型v002/动作v002共享66骨签名、512贴图、站立高2.2m；其余11段待制作，未绑定完整AI或投放。'
transferpath=pkg/'runtime/character_transfer_ledger.json';t=json.loads(transferpath.read_text(encoding='utf-8'));assert t['skeleton_signature']==report['skeleton_signature']
t.update(version='v004',animation_source=rel(animation),validation_scope='idle_walking_source_and_godot_verified; remaining_11_actions_pending',notes=note,clips=[{'name':name,'duration_s':duration,'fps':30,'frame_range':[0,frames],'loop':True,'source':rel(animation),'source_validation':rel(pkg/'previews/idle_walk_v002/validation.json'),'godot_validation':rel(pkg/'previews/idle_walk_v002/godot_validation.json')} for name,duration,frames in [('idle',3.2,96),('walking',2.0,60)]])
t['missing_clips']=[x for x in t['missing_clips'] if x not in ['idle','walking']]
for c in t['planned_clips']:
 if c['name'] in ['idle','walking']:c['status']='authored_exported_verified'
paths=[root/e['path'] for e in t['files'] if 'animation_v001.blend' not in e['path']]+[animation,pkg/'components/enm_normal_fat_zombie03_visual_top3d.glb.import']
t['files']=[{'path':rel(f),'sha256':sha(f)} for f in dict.fromkeys(paths)]
t['action_design']['sha256']=sha(root/t['action_design']['path']);t['action_design']['character_count']=len((root/t['action_design']['path']).read_text(encoding='utf-8'))
t['rollback']='source model v002 retained; _scratch/fat_zombie03/idle_walk_v002_backup contains pre-idle GLB/import/ledger/baseline/transfer'
t.pop('duplicate_gate',None);transferpath.write_text(json.dumps(t,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
w=load_workbook(ledger);ws=w['资产主表'];r=next(r for r,v in read_source_rows(ws) if v[0]==aid)
for c,v in {9:'idle 3.2s / walking 2.0s循环；其余11段待制作',13:'v004',14:'源高3.142857m；游戏站立高2.2m；2867顶点/5690三角；66骨；512²贴图；2已验收动作/13设计动作',16:rel(model)+'; '+rel(animation)+'; docs/v0.1/design/胖子僵尸03动作设计.md',20:sha(prefab),22:datetime(2026,10,3),25:note}.items():ws.cell(r,c).value=v
log=w['域变更日志'];v=str(log.cell(log.max_row,1).value).lstrip('v').split('.');v[-1]=str(int(v[-1])+1);log.append(['v'+'.'.join(v),'2026-10-03','张手修正与行走制作','敌人',note,'骨架/网格/权重/碰撞不变；修正idle并新增walking','Codex']);w.save(ledger)
w=load_workbook(ledger);ws=w['3D-敌人'];r=next(r for r in range(5,ws.max_row+1) if ws.cell(r,1).value==aid)
for c,v in {5:rel(model)+'; '+rel(animation),6:'厚血慢速近战视觉；独立Prefab自动循环idle',14:'待机/行走已验收；其余动作待制作',15:'v004',16:note}.items():ws.cell(r,c).value=v
ws=w['敌人动画与状态']
for r in range(6,ws.max_row+1):
 if ws.cell(r,1).value!=aid:continue
 ws.cell(r,9).value=note
 if ws.cell(r,2).value=='idle':
  for c,v in {4:'自然张手/压重呼吸/左右移重；30fps F0—96；3.2s',5:rel(animation),6:'循环；Godot LOOP_LINEAR',7:'Root固定；双脚锁地',8:'已制作/导出/源与Godot验收；Prefab自动播放'}.items():ws.cell(r,c).value=v
 if ws.cell(r,2).value=='walking':
  for c,v in {4:'短步移重；30fps F0—60；2.0s；参考0.30m/s',5:rel(animation),6:'循环；Godot LOOP_LINEAR',7:'Root固定；F0/F30落脚候选，未绑定声效',8:'已制作/导出/源与Godot验收'}.items():ws.cell(r,c).value=v
w.save(ledger)
bl=json.loads(baseline.read_text(encoding='utf-8'));values=next(v for _,v in read_source_rows(w['资产主表']) if v[0]==aid);bl['assets'][aid]={'v':_row_digest(values),'c':'敌人','d':'enemies'}
for name in ['3D-敌人','敌人动画与状态']:bl['sheet_digests'][name]=sheet_digest(w[name])
union=[]
for d in index.domains:union.extend(read_source_rows(load_workbook(d.path)['资产主表']))
bl['column_digests']={str(c):col_digest(union,c) for c in CONTENT_COLUMNS};baseline.write_text(json.dumps(bl,ensure_ascii=False,indent=1)+'\n',encoding='utf-8')
before=load_workbook(backup/'ledger_before_walk.xlsx');a={v[0]:v for _,v in read_source_rows(before['资产主表'])};b={v[0]:v for _,v in read_source_rows(w['资产主表'])};assert all(a[k]==b[k] for k in a if k!=aid)
print('FAT_ZOMBIE03_IDLE_LEDGER_OK existing_assets_unchanged=true')
