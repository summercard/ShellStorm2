import sys,json,shutil,hashlib
from pathlib import Path
from datetime import datetime
from openpyxl import load_workbook
p=Path(__file__).parent;root=p.parents[1];pkg=root/'assets/art/enemies/normal_enemy_3d/fat_zombie03';aid='ENM-NORMAL-FAT-ZOMBIE03';prefix='enm_normal_fat_zombie03'
sys.path.insert(0,str(root/'scripts'));sys.path.insert(0,str(root/'tools/asset_pipeline'))
from ledger_registry import LedgerIndex
from split_asset_ledger import read_source_rows,_row_digest,sheet_digest,CONTENT_COLUMNS,col_digest
index=LedgerIndex.load(root);ledger=index.domain_for_key('enemies').path;baseline=root/'assets/registry/ledger_split_baseline.json'
backup=p/'v002_backup';backup.mkdir(exist_ok=True)
for f,n in [(ledger,'ledger_before_v002.xlsx'),(baseline,'baseline_before_v002.json')]:
 assert not (backup/n).exists(),'Refuse overwriting transaction backup'
 shutil.copy2(f,backup/n)
rel=lambda f:f.relative_to(root).as_posix();sha=lambda f:hashlib.sha256(f.read_bytes()).hexdigest()
model=pkg/'source/model'/f'{prefix}_model_v002.blend';glb=pkg/'components'/f'{prefix}_visual_top3d.glb';prefab=pkg/'runtime'/f'{prefix}_root_top3d.tscn';tex=pkg/'source/model/textures'/f'{prefix}_basecolor_v002.png'
design=root/'docs/v0.1/design/胖子僵尸03动作设计.md';audit=json.loads((pkg/'source/model/model_audit_v002.json').read_text(encoding='utf-8'))
assert len(design.read_text(encoding='utf-8'))<5000
assert audit['core_missing']==[] and audit['bones']==66 and audit['texture_size']==[512,512]
planned=[('idle',96,True,'压重待机'),('walking',60,True,'沉重巡逻'),('running',36,True,'低速压迫追击'),('attack',75,False,'单臂斜砸'),('hurt',24,False,'重击硬直'),('dead',78,False,'跪塌前扑'),('awaken',36,False,'唤醒'),('alert',18,False,'察觉定身'),('turn_l',24,False,'左转'),('turn_r',24,False,'右转'),('move_start',12,False,'起步'),('move_stop',18,False,'刹停'),('hit_light',9,False,'普通受击')]
note='编号03；v002从原始FBX重建，修正v001连接骨重复变换。36核心骨名齐全+30原有附加骨=66；Root无父级、Hip→Root；保留全部权重。512²贴图；游戏高2.2m（源3.142857m×0.70，kind/variant=1）。动作设计r1为13段，动画尚未制作或绑定，未投放。独立骨架签名，不声称共享静止姿态。'
transferpath=pkg/'runtime/character_transfer_ledger.json';shutil.copy2(transferpath,backup/'character_transfer_ledger_v001.json')
transfer=json.loads(transferpath.read_text(encoding='utf-8'));transfer.update(version='v002',model_source=rel(model),animation_source=None,skeleton_id=audit['skeleton_id'],skeleton_signature=audit['skeleton_signature'],status='exported_pending_godot_validation',validation_scope='model_rig_texture_static_load_verified; animations_design_only',clips=[],missing_clips=[c for c,_,_,_ in planned],planned_clips=[{'name':c,'fps':30,'last_frame':end,'duration_s':end/30,'loop':loop,'status':'designed_not_authored'} for c,end,loop,_ in planned],action_design={'path':rel(design),'revision':'r1','sha256':sha(design),'character_count':len(design.read_text(encoding='utf-8'))},files=[{'path':rel(f),'sha256':sha(f)} for f in [model,tex,glb,prefab]],notes=note,rollback='v001 source retained for history only (known skeleton defect); transaction backups _scratch/fat_zombie03/v002_backup')
transfer.pop('duplicate_gate',None)
transferpath.write_text(json.dumps(transfer,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
for f in ['reopen_audit_v002.json','godot_audit_v002.json']:
 if (p/f).exists():shutil.copy2(p/f,pkg/'previews'/f)
wb=load_workbook(ledger);ws=wb['资产主表'];hits=[r for r,v in read_source_rows(ws) if v[0]==aid];assert len(hits)==1;row=hits[0]
spec='5.2518×2.0550×3.1429m (Blender XYZ,T姿势)；游戏高2.2m；2867顶点/5690三角面；36核心+30附加骨；512²贴图；0已制作动作/13已设计动作'
changes={9:'T_pose / 13动作设计r1；未制作',13:'v002',14:spec,16:rel(model)+'; '+rel(design),20:sha(prefab),22:datetime(2026,10,3),25:note}
for c,v in changes.items():ws.cell(row,c).value=v
log=wb['域变更日志'];last=str(log.cell(log.max_row,1).value);parts=last.lstrip('v').split('.');parts[-1]=str(int(parts[-1])+1)
log.append(['v'+'.'.join(parts),'2026-10-03','骨架修正与动作设计','敌人',aid+'；'+note,'模型与稳定GLB更新；AI/碰撞/伤害未变','Codex'])
wb.save(ledger)
# Separate domain-page transaction with precise baseline refresh.
wb=load_workbook(ledger);ws=wb['3D-敌人'];r=next(r for r in range(5,ws.max_row+1) if ws.cell(r,1).value==aid)
for c,v in {5:rel(model),6:'厚血慢速近战视觉；骨架与512贴图已修正；动作处于设计阶段',11:spec,12:'脚底0；Blender+Y/Godot-Z；Scale=1；游戏高2.2m',14:'骨架与模型已核验；动画待制作',15:'v002',16:note}.items():ws.cell(r,c).value=v
ws=wb['敌人动画与状态']
for clip,end,loop,title in planned:
 found=[r for r in range(6,ws.max_row+1) if ws.cell(r,1).value==aid and ws.cell(r,2).value==clip]
 assert len(found)<=1
 r=found[0] if found else ws.max_row+1
 vals=[aid,clip,'胖子僵尸03 '+title,f'{end/30:.2f}s / 30fps / F0—{end}',rel(design)+' r1；设计完成，Action未制作','循环' if loop else '单次','attack命中F36；Root固定' if clip=='attack' else 'Root固定；事件见设计','设计阶段；未绑定',note]
 for c,v in enumerate(vals,1):ws.cell(r,c).value=v
wb.save(ledger)
bl=json.loads(baseline.read_text(encoding='utf-8'));prior=dict(bl['assets']);values=next(v for _,v in read_source_rows(wb['资产主表']) if v[0]==aid);bl['assets'][aid]={'v':_row_digest(values),'c':'敌人','d':'enemies'}
assert all(bl['assets'][k]==v for k,v in prior.items() if k!=aid)
for s in ['3D-敌人','敌人动画与状态']:bl['sheet_digests'][s]=sheet_digest(wb[s])
union=[]
for d in index.domains:union.extend(read_source_rows(load_workbook(d.path)['资产主表']))
bl['column_digests']={str(c):col_digest(union,c) for c in CONTENT_COLUMNS}
baseline.write_text(json.dumps(bl,ensure_ascii=False,indent=1)+'\n',encoding='utf-8')
before=load_workbook(backup/'ledger_before_v002.xlsx');after=load_workbook(ledger)
assert all(before['资产主表'].cell(r,c).value==after['资产主表'].cell(r,c).value for r in range(6,ws.max_row+1) if r!=row for c in range(1,26))
print('FAT_ZOMBIE03_V002_LEDGER_OK','row',row,'planned',len(planned),'chars',len(design.read_text(encoding='utf-8')))
