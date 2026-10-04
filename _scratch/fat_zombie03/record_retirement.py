from pathlib import Path
from openpyxl import load_workbook
from datetime import datetime
import sys,json,shutil,subprocess
root=Path.cwd();backup=root/'_scratch/fat_zombie03/shielded_retirement_backup';out=root/'outputs/fat_zombie03_retirement';out.mkdir(exist_ok=True)
sys.path[:0]=[str(root/'scripts'),str(root/'tools/asset_pipeline')]
from ledger_registry import LedgerIndex
from split_asset_ledger import read_source_rows,_row_digest,sheet_digest,CONTENT_COLUMNS,col_digest
index=LedgerIndex.load(root);ledger=index.domain_for_key('enemies').path;baseline=root/'assets/registry/ledger_split_baseline.json';content=root/'docs/v0.1/data/ShellStorm2_游戏内容数据库_v010.xlsx'
for p in [ledger,baseline,content]:
 target=backup/p.name
 if not target.exists():shutil.copy2(p,target)
retired='ENM-TANK-SHELLGUARD01';fat='ENM-NORMAL-FAT-ZOMBIE03'
w=load_workbook(ledger);ws=w['资产主表']
for r,values in read_source_rows(ws):
 if values[0] not in [retired,fat]:continue
 if values[0]==retired:
  ws.cell(r,11).value='弃用';ws.cell(r,25).value='2026-10-03用户要求退役；现行盒/随机池/主题池由ENM-NORMAL-FAT-ZOMBIE03接替。旧ID、源与存档兼容保留；不再新投放。'
 else:
  ws.cell(r,25).value=str(ws.cell(r,25).value)+'；已替换壳甲卫兵的远征盒、默认随机池和两个主题怪池；固定种子全关实测10只胖子/0只壳甲。'
 ws.cell(r,22).value=datetime(2026,10,3)
log=w['域变更日志'];versions=[str(log.cell(i,1).value).lstrip('v').split('.') for i in range(1,log.max_row+1)];versions=[x for x in versions if len(x)==3 and all(y.isdigit() for y in x)];v=list(max(versions,key=lambda x:tuple(map(int,x))));v[-1]=str(int(v[-1])+1);log.append(['v'+'.'.join(v),'2026-10-03','壳甲卫兵退役，由03胖子僵尸接替','敌人','远征纵列1只/桥心4只，数量延迟保留；所有现行shielded池改fat_zombie03','旧存档兼容与未实装精英设计保留','Codex']);w.save(ledger)
w=load_workbook(ledger)
for sheet in ['3D-敌人','敌人动画与状态']:
 ws=w[sheet]
 for row in ws.iter_rows():
  if row[0].value==retired:
   # 专表只标记状态与说明，不删历史动作设计。
   if sheet=='敌人动画与状态':row[7].value='退役；新投放由胖子僵尸03接替';row[8].value='旧设计/兼容保留，停止普通怪投放'
  if row[0].value==fat:
   if sheet=='3D-敌人':row[12].value='远征01纵列/桥心盒；默认随机池、锈炉/深渊主题池；独立实体'
   row[-1].value=str(row[-1].value)+'；已接替壳甲卫兵，现行投放启用。'
w.save(ledger)
bl=json.loads(baseline.read_text(encoding='utf-8'))
for r,val in read_source_rows(w['资产主表']):
 if val[0] in [retired,fat]:bl['assets'][val[0]]={'v':_row_digest(val),'c':'敌人','d':'enemies'}
for sheet in ['3D-敌人','敌人动画与状态']:bl['sheet_digests'][sheet]=sheet_digest(w[sheet])
union=[]
for d in index.domains:union.extend(read_source_rows(load_workbook(d.path)['资产主表']))
bl['column_digests']={str(c):col_digest(union,c) for c in CONTENT_COLUMNS};baseline.write_text(json.dumps(bl,ensure_ascii=False,indent=1)+'\n',encoding='utf-8')
wb=load_workbook(content);ws=wb['怪物与Boss']
for r in range(5,ws.max_row+1):
 if ws.cell(r,1).value=='monster_shielded':
  ws.cell(r,10).value='停止现行刷新，原点位/池由fat_zombie03接替';ws.cell(r,12).value='[已退役]';ws.cell(r,13).value='弃用；旧档兼容保留';ws.cell(r,15).value='2026-10-03用户裁定退役；普通怪身份由胖子僵尸03替换。原数值留作历史，未实装壳甲精英设计保留。'
 if ws.cell(r,1).value=='monster_fat_zombie03':
  ws.cell(r,10).value='远征01纵列/桥心盒；原壳甲默认随机池和锈炉/深渊主题池'
  ws.cell(r,15).value=str(ws.cell(r,15).value).replace('未加入既有随机池','已接替壳甲卫兵投放')
wb.save(content)
# 只允许预期两资产和两内容行改变。
old=load_workbook(backup/ledger.name);a={v[0]:v for _,v in read_source_rows(old['资产主表'])};b={v[0]:v for _,v in read_source_rows(w['资产主表'])};assert all(a[k]==b[k] for k in a if k not in [retired,fat])
old=load_workbook(backup/content.name)
for name in old.sheetnames:
 if name!='怪物与Boss':assert list(old[name].values)==list(wb[name].values)
for a,b in zip(old['怪物与Boss'].values,wb['怪物与Boss'].values):
 if a[0] not in ['monster_shielded','monster_fat_zombie03']:assert a==b
for path in ['docs/v0.1/design/远征关卡01设计.md']:
 f=root/path;s=f.read_text(encoding='utf-8');shutil.copy2(f,backup/f.name);s=s.replace('壳甲卫兵','胖子僵尸03').replace('壳甲×','胖子僵尸×').replace('`shielded`','`fat_zombie03`');f.write_text(s,encoding='utf-8')
f=root/'docs/v0.1/06_技术施工_怪物精英与Boss.md';s=f.read_text(encoding='utf-8').replace('六种普通模板：小菌猪、孢子射手、蜂巢怪、壳甲卫兵、炸弹果、地刺虫','六种现行普通模板：小僵尸、保安僵尸、蜂巢怪、胖子僵尸03、炸弹果、地刺虫；壳甲卫兵已退役，旧存档兼容保留');f.write_text(s,encoding='utf-8')
record='docs/v0.1/development/2026-10-03_shielded_retirement.md'
f=root/'docs/v0.1/feature_registry.json';lines=f.read_text(encoding='utf-8').splitlines(keepends=True)
for i,line in enumerate(lines):
 if '"feature_id":"ENEMY-AI"' not in line and '"feature_id":"ASSET-PIPELINE"' not in line:continue
 d=json.loads(line.strip().rstrip(','));d['development_records'].append(record);d['verification'].append({'kind':'scene','id':'verify_shielded_retirement'});lines[i]='    '+json.dumps(d,ensure_ascii=False,separators=(',',':'))+',\n'
f.write_text(''.join(lines),encoding='utf-8')
f=root/'docs/v0.1/MODULE_INDEX.md';s=f.read_text(encoding='utf-8').replace('固定波次和触发盒调用，厚血慢速','固定波次和触发盒调用，已接替壳甲卫兵的现行盒/怪池，厚血慢速');f.write_text(s+'\nENEMY-AI / ASSET-PIPELINE [壳甲卫兵退役](development/2026-10-03_shielded_retirement.md)：现行刷怪由03胖子接替，旧存档及历史设计保留；独立验收verify_shielded_retirement。\n',encoding='utf-8')
f=root/'docs/v0.1/development/CHANGELOG.md';s=f.read_text(encoding='utf-8');f.write_text(s.replace('# 游戏设计文档 v0.1 变更记录','# 游戏设计文档 v0.1 变更记录\n\n## 2026-10-03｜壳甲卫兵退役\n\n- 远征两盒、默认/主题怪池由胖子僵尸03接替；数量、延迟、权重保留。壳甲资产和内容登记退役，旧档兼容保留。见[交付记录](2026-10-03_shielded_retirement.md)。',1),encoding='utf-8')
transfer=root/'assets/art/enemies/normal_enemy_3d/fat_zombie03/runtime/character_transfer_ledger.json';t=json.loads(transfer.read_text(encoding='utf-8'));t['remaining_work']=[x for x in t['remaining_work'] if '随机关卡池' not in x];t['deployment']={'replaces':'shielded','boxes':['box_corridor_column','box_bridge_center'],'default_pools':True,'themes':['rust_foundry','abyss_archive'],'validation':'outputs/fat_zombie03_retirement/runtime.json'};transfer.write_text(json.dumps(t,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
for name in ['retirement','retirement_verify_fat_zombie03','retirement_verify_expedition01_spawn_ramp','retirement_baseline_ramp']:shutil.copy2(root/'_scratch/fat_zombie03'/f'{name}.log',out/f'{name}.log')
results={}
for key,args in {'structure':['scripts/check_asset_registry.py','--scope','structure'],'split':['tools/asset_pipeline/verify_ledger_split.py'],'full':['scripts/check_asset_registry.py','--scope','full','--ledger','enemies','--json-output',str(out/'ledger.json')],'docs':['scripts/check_documentation_contracts.py'],'naming':['scripts/check_asset_runtime_naming.py']}.items():
 r=subprocess.run([sys.executable,*args],capture_output=True,cwd=root);results[key]=r.returncode;(out/f'{key}.log').write_bytes(r.stdout+r.stderr);print(key,r.returncode,flush=True)
results.update(retirement=0,fat_zombie03=0,expedition_ramp=1,expedition_ramp_same_before=True)
(out/'gates.json').write_text(json.dumps(results,indent=2),encoding='utf-8');assert results['structure']==results['split']==0
print('RETIREMENT_RECORDED')
