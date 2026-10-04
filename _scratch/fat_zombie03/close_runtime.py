import sys,json,shutil,re,subprocess
from pathlib import Path
from openpyxl import load_workbook
root=Path.cwd();out=root/'outputs/fat_zombie03_integration';p=root/'_scratch/fat_zombie03';aid='ENM-NORMAL-FAT-ZOMBIE03'
sys.path[:0]=[str(root/'scripts'),str(root/'tools/asset_pipeline')]
from ledger_registry import LedgerIndex
from split_asset_ledger import read_source_rows,_row_digest,sheet_digest,CONTENT_COLUMNS,col_digest
index=LedgerIndex.load(root);ledger=index.domain_for_key('enemies').path
w=load_workbook(ledger)
for sheet in ['资产主表','3D-敌人','敌人动画与状态']:
 ws=w[sheet]
 for row in ws.iter_rows():
  if row[0].value!=aid:continue
  for cell in row:
   if isinstance(cell.value,str) and '专项150项' in cell.value:cell.value=cell.value.replace('专项150项','专项152项')
w.save(ledger)
baseline=root/'assets/registry/ledger_split_baseline.json';bl=json.loads(baseline.read_text(encoding='utf-8'));values=next(v for _,v in read_source_rows(w['资产主表']) if v[0]==aid);bl['assets'][aid]={'v':_row_digest(values),'c':'敌人','d':'enemies'}
for name in ['3D-敌人','敌人动画与状态']:bl['sheet_digests'][name]=sheet_digest(w[name])
union=[]
for d in index.domains:union.extend(read_source_rows(load_workbook(d.path)['资产主表']))
bl['column_digests']={str(c):col_digest(union,c) for c in CONTENT_COLUMNS};baseline.write_text(json.dumps(bl,ensure_ascii=False,indent=1)+'\n',encoding='utf-8')
for path in ['assets/art/enemies/normal_enemy_3d/fat_zombie03/README.md','docs/v0.1/development/2026-10-03_fat_zombie03_runtime_binding.md','docs/v0.1/development/CHANGELOG.md']:
 f=root/path;f.write_text(f.read_text(encoding='utf-8').replace('150项','152项'),encoding='utf-8')
transfer=root/'assets/art/enemies/normal_enemy_3d/fat_zombie03/runtime/character_transfer_ledger.json';t=json.loads(transfer.read_text(encoding='utf-8'));t['notes']=t['notes'].replace('150项','152项');t['validation'][0]['checks']=152;t['runtime_binding']['spawn_catalog']='src/map/SpawnBoxCatalog.gd';transfer.write_text(json.dumps(t,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
shutil.copy2(p/'baseline_full_game.log',out/'baseline_full_game.log');shutil.copy2(p/'final_verify_full_3d_game_flow.log',out/'full_game.log');shutil.copy2(p/'final_verify_3d_enemy_behavior_flow.log',out/'enemy_behavior.log')
before=(out/'baseline_full_game.log').read_text(encoding='utf-8');after=(out/'full_game.log').read_text(encoding='utf-8')
normalize=lambda s:[re.sub(r'budget: \d+','budget: <node_count>',x) for x in s.splitlines() if 'ERROR:' in x]
assert normalize(before)==normalize(after)==['ERROR: Player eight-state whitelist contract regressed','ERROR: 3D level exceeds prototype node budget: <node_count>']
record=root/'docs/v0.1/development/2026-10-03_fat_zombie03_runtime_binding.md'
record.write_text(record.read_text(encoding='utf-8')+'\n扩大回归：普通怪3D行为全场景退出0；全游戏流程退出1，仅玩家八态白名单及关卡节点预算两项失败。隔离副本复制当前源码并只恢复本次改动前的4个既有脚本，复现同两项失败（节点数基线2807/当前2805），证明不是本怪引入；不宣称全工程全绿。触发盒白名单补入fat_zombie03，通用注入入口也能点名生成；专项最终152项。文档门禁在统一Python运行时复测后仅余4个既有未登记测试。\n',encoding='utf-8')
gates=json.loads((out/'gates.json').read_text());gates['full_game']=1;gates['full_game_failures_match_before_binding']=True;gates['enemy_behavior']=0;gates['fat_zombie03']=0;gates['unexpected_errors_in_fat_test']=0;gates['docs_same_runtime']=1
for key,args in {'structure':['scripts/check_asset_registry.py','--scope','structure'],'split':['tools/asset_pipeline/verify_ledger_split.py'],'target_full':['scripts/check_asset_registry.py','--scope','full','--ledger','enemies','--json-output',str(out/'ledger_after.json')]}.items():
 r=subprocess.run([sys.executable,*args],capture_output=True,cwd=root);gates[key]=r.returncode;(out/f'final_{key}.log').write_bytes(r.stdout+r.stderr);print(key,r.returncode)
(out/'gates.json').write_text(json.dumps(gates,indent=2)+'\n',encoding='utf-8')
assert gates['structure']==gates['split']==0
print('FAT_ZOMBIE03_CLOSED')
