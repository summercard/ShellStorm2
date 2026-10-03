import sys,json,shutil,hashlib,re,zipfile
from pathlib import Path
from copy import copy
from collections import Counter
from datetime import datetime
from openpyxl import load_workbook
p=Path(__file__).parent; root=p.parents[1]
sys.path.insert(0,str(root/'scripts')); sys.path.insert(0,str(root/'tools/asset_pipeline'))
from ledger_registry import LedgerIndex
from split_asset_ledger import FIRST_DATA_ROW,HEADER_ROW,dedupe_key_formula,dedupe_result_formula,read_source_rows,_row_digest,CONTENT_COLUMNS,col_digest,sheet_digest
index=LedgerIndex.load(root); domain=index.domain_for_key('enemies'); ledger=domain.path
pkg=root/'assets/art/enemies/normal_enemy_3d/fat_zombie03'; prefix='enm_normal_fat_zombie03'; aid='ENM-NORMAL-FAT-ZOMBIE03'
rel=lambda f:f.relative_to(root).as_posix()
sha=lambda f:hashlib.sha256(f.read_bytes()).hexdigest()
model=pkg/'source/model'/f'{prefix}_model_v001.blend'; visual=pkg/'components'/f'{prefix}_visual_top3d.glb'; prefab=pkg/'runtime'/f'{prefix}_root_top3d.tscn'
audit=json.loads((pkg/'source/model/model_audit.json').read_text(encoding='utf-8'))
shutil.copy2(p/'godot_project/godot_audit.json',pkg/'previews/godot_audit.json')
note='源编号03；模型入库及Godot独立加载通过，非完整怪物验收。65骨Mixamo骨名未对齐项目，缺Root；无动画，待idle/walking/running/attack/hurt/dead六段与动作母版、状态绑定。未配置AI/伤害/掉落/碰撞或投放。T姿势宽度不作为碰撞尺寸。来源许可待确认。'
transfer={'schema_version':1,'asset_id':aid,'content_id':'fat_zombie03','source_number':'03','version':'v001','status':'exported_pending_godot_validation','validation_scope':'static_model_load_passed; full_character_contract_pending','model_source':rel(model),'animation_source':None,'skeleton_signature':audit['skeleton_signature'],'clips':[],'missing_clips':audit['missing_clips'],'files':[{'path':rel(f),'sha256':sha(f)} for f in [model,visual,prefab]],'source_zip':'I:/工作项目/桌面文件备份/僵尸_项目素材_20260922/胖子僵尸.zip','source_zip_sha256':sha(Path('I:/工作项目/桌面文件备份/僵尸_项目素材_20260922/胖子僵尸.zip')),'consumer':'independent PackedScene only; Enemy3D not bound','rollback':'remove newly created fat_zombie03 package; restore ledger_before.xlsx and baseline_before.json','notes':note}
(pkg/'runtime/character_transfer_ledger.json').write_text(json.dumps(transfer,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
shutil.copy2(ledger,p/'ledger_before.xlsx'); baseline=root/'assets/registry/ledger_split_baseline.json'; shutil.copy2(baseline,p/'baseline_before.json')
wb=load_workbook(ledger); roundtrip=p/'roundtrip.xlsx'; wb.save(roundtrip)
def features(f):
 with zipfile.ZipFile(f) as z:
  import xml.etree.ElementTree as ET
  ns={'s':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
  return {n:{tag:[ET.tostring(e,encoding='unicode') for e in ET.fromstring(z.read(n)).findall('s:'+tag,ns)] for tag in ['mergeCells','dataValidations','tableParts']} for n in z.namelist() if n.startswith('xl/worksheets/sheet') and n.endswith('.xml')}
assert features(ledger)==features(roundtrip),'Workbook roundtrip features changed'
main=wb['资产主表']; old=main.max_row; row=old+1
assert aid not in [v[0] for _,v in read_source_rows(main)]
assert main.cell(HEADER_ROW,1).value=='AssetID'
spec='3.1033×1.2143×1.8571m (Blender XYZ,T姿势)；2867顶点/5690三角面；1网格/1材质；65骨；0动画；4096²颜色贴图'
vals=[aid,'胖子僵尸','敌人','normal','fat_zombie03','root_3d',None,'Top3D / local -Z 正面','T_pose / 六段动作待制作','独立新怪资产；未投放','Blender源已完成','P1','v001',spec,rel(prefab),rel(model),'胖子僵尸; fat zombie; 03',dedupe_key_formula(row),dedupe_result_formula(row,row),sha(prefab),'Codex',datetime(2026,10,3),'用户提供Tripo素材；许可待确认','03',note]
for c,v in enumerate(vals,1):main.cell(row,c)._style=copy(main.cell(old,c)._style); main.cell(row,c).value=v
main.row_dimensions[row].height=main.row_dimensions[old].height
for r in range(FIRST_DATA_ROW,row+1):main.cell(r,19).value=dedupe_result_formula(r,row)
for ref in ['A6','C6','E6','G6','B10','C10','B11','C11']:
 cell=wb['总览'][ref]; cell.value=str(cell.value).replace(f'${old}',f'${row}')
for dv in main.data_validations.dataValidation:
 dv.sqref=' '.join(f'{r.min_col and main.cell(r.min_row,r.min_col).column_letter}{r.min_row}:{main.cell(r.max_row,r.max_col).column_letter}{row if r.max_row==old else r.max_row}' for r in dv.sqref.ranges)
log=wb['域变更日志']; log.append(['v0.1.4','2026-10-03','新增素材入库','敌人',aid+' 胖子僵尸源编号03；'+note,'既有资产内容不变；未接入运行池','Codex'])
wb.save(ledger)
# Separate transaction: actual prefab and explicit animation deficits.
wb=load_workbook(ledger); ws=wb['3D-敌人']; ws.append([aid,'胖子僵尸',rel(prefab),rel(visual),rel(model),'独立视觉包装；静态模型可加载',None,'关（表现资产）','Enemy3D','未新增碰撞',spec,'脚底0；Blender +Y / Godot -Z；Scale=1','独立预览；未投放','已入库，待骨架与动画验收','v001',note])
ws=wb['敌人动画与状态']
for clip in audit['missing_clips']:ws.append([aid,clip,'胖子僵尸03','待制作','缺失；原始FBX无Action','循环' if clip in ['idle','walking','running'] else '单次','待定义','未实现',note])
wb.save(ledger)
bl=json.loads(baseline.read_text(encoding='utf-8')); v=next(v for _,v in read_source_rows(wb['资产主表']) if v[0]==aid)
bl['assets'][aid]={'v':_row_digest(v),'c':'敌人','d':'enemies'}; bl['asset_count']=len(bl['assets'])
union=[]
for d in index.domains:union.extend(read_source_rows(load_workbook(d.path)['资产主表']))
bl['column_digests']={str(c):col_digest(union,c) for c in CONTENT_COLUMNS}; bl['category_counts']=dict(sorted(Counter(str(v[2]).strip() for _,v in union).items()))
for s in ['3D-敌人','敌人动画与状态']:bl['sheet_digests'][s]=sheet_digest(wb[s])
baseline.write_text(json.dumps(bl,ensure_ascii=False,indent=1)+'\n',encoding='utf-8')
before=load_workbook(p/'ledger_before.xlsx'); after=load_workbook(ledger)
assert all(before['资产主表'].cell(r,c).value==after['资产主表'].cell(r,c).value for r in range(FIRST_DATA_ROW,old+1) for c in range(1,26) if c!=19)
print('FAT_ZOMBIE03_REGISTERED',row,'roundtrip_preserved=true existing_asset_content_unchanged=true')
