import sys, json, hashlib, shutil
from pathlib import Path
from datetime import datetime
from openpyxl import load_workbook
R = Path(r'I:/工作项目/shellstrom2/ShellStorm2')
sys.path[:0] = [str(R/'scripts'), str(R/'tools/asset_pipeline')]
from ledger_registry import LedgerIndex
from split_asset_ledger import read_source_rows, _row_digest, CONTENT_COLUMNS, col_digest
idx = LedgerIndex.load(R)
ledger = idx.path_for_category('特效')
baseline_path = R/'assets/registry/ledger_split_baseline.json'
out = R/'outputs/bird_flocks_v001'
shutil.copy2(ledger, out/'ledger_before_runtime_finalize.xlsx')
shutil.copy2(baseline_path, out/'baseline_before_runtime_finalize.json')
ids = {'VFX-ENV-BIRDS-FLYBY-3D': ('flyby', 'assets/art/vfx/environment_3d/bird_flocks/runtime/flyby/vfx_env_birds_flyby_root_top3d.tscn'), 'VFX-ENV-BIRDS-GROUND-3D': ('ground', 'assets/art/vfx/environment_3d/bird_flocks/runtime/ground/vfx_env_birds_ground_root_top3d.tscn')}
w = load_workbook(ledger)
s = w['资产主表']
for row in range(6, s.max_row + 1):
    aid = s.cell(row, 1).value
    if aid not in ids:
        continue
    kind, runtime_path = ids[aid]
    s.cell(row, 11).value = '正式美术已接入'
    s.cell(row, 15).value = runtime_path
    s.cell(row, 16).value = 'assets/art/vfx/environment_3d/bird_flocks/source/%s/v001/asset_manifest.json; GLB; Godot PackedScene' % kind
    s.cell(row, 25).value = 'Godot已接入；一次性演出不循环；无碰撞；运行时归属场景特效；' + ('跨塔路线家塔2-3楼群通道。' if kind == 'flyby' else '主塔100F天台开放区。')
    glb = R / ('assets/art/vfx/environment_3d/bird_flocks/components/%s/vfx_env_birds_%s_visual.glb' % (kind, kind))
    s.cell(row, 20).value = hashlib.sha256(glb.read_bytes()).hexdigest()
w.save(ledger)
for kind in ('flyby','ground'):
    manifest_path = R/f'assets/art/vfx/environment_3d/bird_flocks/source/{kind}/v001/asset_manifest.json'
    m = json.loads(manifest_path.read_text(encoding='utf-8'))
    m['status'] = '正式美术已接入'
    m['runtime_integrated'] = True
    m['exported'] = True
    m['optimized_blend'] = f'assets/art/vfx/environment_3d/bird_flocks/source/{kind}/v001/export/v001/vfx_env_birds_{kind}_optimized_v001.blend'
    m['component_glb'] = f'assets/art/vfx/environment_3d/bird_flocks/components/{kind}/vfx_env_birds_{kind}_visual.glb'
    m['runtime_prefab'] = ids['VFX-ENV-BIRDS-FLYBY-3D' if kind == 'flyby' else 'VFX-ENV-BIRDS-GROUND-3D'][1]
    m['placement'] = 'cross_tower_route_between_tower_02_03' if kind == 'flyby' else 'tower_100f_rooftop_open_area'
    manifest_path.write_text(json.dumps(m, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
# Rebuild only the baseline fingerprints from current ledgers, preserving old asset fingerprints.
baseline = json.loads(baseline_path.read_text(encoding='utf-8'))
for row, values in read_source_rows(s):
    if values[0] in ids:
        baseline['assets'][values[0]] = {'v': _row_digest(values), 'c': '特效', 'd': 'vfx'}
union = []
for domain in idx.domains:
    union.extend(read_source_rows(load_workbook(domain.path)['资产主表']))
baseline['asset_count'] = len(baseline['assets'])
baseline['column_digests'] = {str(c): col_digest(union, c) for c in CONTENT_COLUMNS}
baseline['category_counts'] = dict(sorted(__import__('collections').Counter(str(v[2]).strip() for _, v in union).items()))
baseline_path.write_text(json.dumps(baseline, ensure_ascii=False, indent=1) + '\n', encoding='utf-8')
# Update source docs without touching the authored Blender-source history statement.
readme = R/'assets/art/vfx/environment_3d/bird_flocks/README.md'
text = readme.read_text(encoding='utf-8')
text = text.replace('没有引擎导出、Prefab 或运行接入。', '已导出 GLB、创建稳定 PackedScene 并接入游戏运行场景。')
text = text.replace('账本通过 `ledger_index.json` 的 vfx 域定位，状态为 `Blender源已完成`。', '账本通过 `ledger_index.json` 的 vfx 域定位，状态为 `正式美术已接入`；运行时归类为场景特效。')
readme.write_text(text, encoding='utf-8')
print('BIRDS_RUNTIME_FINALIZED', sorted(ids))
