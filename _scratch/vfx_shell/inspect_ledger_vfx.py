from pathlib import Path
import openpyxl

root = Path(r"I:\工作项目\shellstrom2\ShellStorm2")
ledger = root / "assets/registry/ledgers/ShellStorm2_特效账本_v001.xlsx"
wb = openpyxl.load_workbook(ledger)
print("SHEETS", wb.sheetnames)

fx = wb["3D-特效"]
print("FX dims", fx.dimensions, "max_row", fx.max_row, "max_col", fx.max_column)
for r in range(1, min(fx.max_row, 20) + 1):
    vals = []
    for c in range(1, fx.max_column + 1):
        v = fx.cell(r, c).value
        if v is not None:
            s = str(v).replace("\n", "\\n")
            if len(s) > 60:
                s = s[:60] + "..."
            vals.append(f"{c}:{s}")
    if vals:
        print(f"FX-R{r}", " | ".join(vals))

master = wb["资产主表"]
print("MASTER dims", master.dimensions, "max_row", master.max_row)
for r in range(1, 8):
    print(f"MASTER-HDR R{r}", [str(master.cell(r, c).value)[:18] if master.cell(r, c).value is not None else None for c in range(1, master.max_column + 1)])
for r in range(6, master.max_row + 1):
    if master.cell(r, 1).value == "VFX-SHELL-CASING-3D":
        print("MASTER SHELL ROW", r)
        for c in range(1, master.max_column + 1):
            print(f"   M{r}C{c}", repr(master.cell(r, c).value)[:200])

log = wb["域变更日志"]
print("LOG dims", log.dimensions, "max_row", log.max_row)
for r in range(1, log.max_row + 1):
    print(f"LOG R{r}", [str(log.cell(r, c).value)[:40] if log.cell(r, c).value is not None else None for c in range(1, 9)])
