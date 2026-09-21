from pathlib import Path
import openpyxl
p = Path(r"I:\工作项目\shellstrom2\ShellStorm2\assets\registry\ledgers\ShellStorm2_特效账本_v001.xlsx")
wb = openpyxl.load_workbook(p)
ws = wb["资产主表"]
last = 22
for row in range(6, last + 1):
    ws.cell(row, 18).value = f'=LOWER(TRIM(C{row})&"|"&TRIM(D{row})&"|"&TRIM(E{row})&"|"&TRIM(F{row})&"|"&TRIM(H{row})&"|"&TRIM(I{row}))'
    ws.cell(row, 19).value = f'=IF(COUNTIF($R$6:$R${last},R{row})>1,"重复","唯一")'
ov = wb["总览"]
for cell in ("A6", "C6", "E6", "G6"):
    ov[cell].value = ov[cell].value.replace("$21", "$22")
for row in range(10, 11):
    for col in ("B", "C"):
        ov[f"{col}{row}"].value = ov[f"{col}{row}"].value.replace("$21", "$22")
wb.save(p)
check = openpyxl.load_workbook(p, data_only=False)
print('formulas', check['资产主表']['S6'].value, check['资产主表']['S22'].value)
print('overview', check['总览']['A6'].value, check['总览']['C6'].value, check['总览']['B10'].value)
