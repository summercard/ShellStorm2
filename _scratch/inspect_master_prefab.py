from pathlib import Path
import json
import openpyxl
root=Path('I:/工作项目/shellstrom2/ShellStorm2')
wb=openpyxl.load_workbook(root/'assets/registry/ShellStorm2_美术资产台账_v001.xlsx',data_only=False)
out={}
for sheet_name in ['总览','3D Prefab总控']:
 ws=wb[sheet_name]
 out[sheet_name]={'max_row':ws.max_row,'max_column':ws.max_column,'rows':[[ws.cell(r,c).value for c in range(1,ws.max_column+1)] for r in range(1,ws.max_row+1)]}
(root/'_scratch/master_prefab_context.json').write_text(json.dumps(out,ensure_ascii=False,indent=2,default=str)+'\n',encoding='utf-8')
