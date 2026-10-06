import json,zipfile,xml.etree.ElementTree as E,re,shutil,copy,sys
from pathlib import Path
from openpyxl import load_workbook
R=Path.cwd();D=R/'_scratch/boss028_ledger';N='http://schemas.openxmlformats.org/spreadsheetml/2006/main';ns={'m':N}
jobs=json.loads((D/'changes.json').read_text(encoding='utf-8'))
def sheets(z):
 rels={e.attrib['Id']:e.attrib['Target'] for e in E.fromstring(z.read('xl/_rels/workbook.xml.rels'))}
 return {s.attrib['name']:('xl/'+rels[s.attrib['{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id']].lstrip('/').removeprefix('xl/')) for s in E.fromstring(z.read('xl/workbook.xml')).find('m:sheets',ns)}
for i,job in enumerate(jobs):
 p=Path(job['path']);backup=D/('before_'+p.name);shutil.copy2(p,backup);before=load_workbook(p)
 replacements={}
 with zipfile.ZipFile(p) as original,zipfile.ZipFile(D/f'cells{i}.xlsx') as generated:
  paths=sheets(original);newpaths=sheets(generated)
  strings=[]
  if 'xl/sharedStrings.xml' in generated.namelist():strings=[''.join(e.itertext()) for e in E.fromstring(generated.read('xl/sharedStrings.xml'))]
  for name,changes in job['sheets'].items():
   text=original.read(paths[name]).decode();root=E.fromstring(generated.read(newpaths[name]));cells={c.attrib['r']:c for c in root.findall('.//m:c',ns)}
   for addr,value in changes.items():
    cell=copy.deepcopy(cells[addr]);old=re.search(r'<c\b[^>]*\br="'+addr+r'"[^>]*(?:/>|>.*?</c>)',text)
    if old:
     style=re.search(r'\bs="([^"]+)"',old.group())
     if style:cell.set('s',style[1])
     else:cell.attrib.pop('s',None)
    else:cell.attrib.pop('s',None)
    if cell.get('t')=='s':
     val=strings[int(cell.find('m:v',ns).text)];cell.remove(cell.find('m:v',ns));cell.set('t','inlineStr');E.SubElement(E.SubElement(cell,'{'+N+'}is'),'{'+N+'}t').text=val
    xml=E.tostring(cell,encoding='unicode').replace('ns0:','').replace(':ns0','')
    if old:text=text[:old.start()]+xml+text[old.end():]
    else:
     row=re.search(r'\d+',addr)[0];match=re.search(r'<row\b[^>]*\br="'+row+r'"[^>]*>.*?</row>',text)
     if match:text=text[:match.end()-6]+xml+text[match.end()-6:]
     else:text=text.replace('</sheetData>',f'<row r="{row}">{xml}</row></sheetData>')
   replacements[paths[name]]=text.encode()
  temp=D/f'merged{i}.xlsx'
  with zipfile.ZipFile(temp,'w',zipfile.ZIP_DEFLATED) as out:
   for info in original.infolist():out.writestr(info,replacements.get(info.filename,original.read(info.filename)))
 after=load_workbook(temp)
 for name in before.sheetnames:
  expected=job['sheets'].get(name,{})
  for row in before[name]:
   for cell in row:
    assert after[name][cell.coordinate].value==expected.get(cell.coordinate,cell.value),(name,cell.coordinate)
 for name,changes in job['sheets'].items():
  for addr,value in changes.items():assert after[name][addr].value==value
 assert p.read_bytes()==backup.read_bytes(),'Concurrent ledger edit; regenerate patch'
 shutil.copy2(temp,p)
 print('Patched only requested cells:',p.name)
