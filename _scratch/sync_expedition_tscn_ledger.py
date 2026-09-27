from copy import copy
from datetime import date
import json
import sys
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(r"I:\工作项目\shellstrom2\ShellStorm2")
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "scripts"))
from ledger_registry import LedgerIndex
from tools.asset_pipeline.split_asset_ledger import (
    CONTENT_COLUMNS,
    _row_digest,
    col_digest,
    dedupe_key_formula,
    dedupe_result_formula,
    read_source_rows,
    sheet_digest,
)

LEDGER = ROOT / "assets/registry/ledgers/ShellStorm2_场景账本_v001.xlsx"
BASELINE = ROOT / "assets/registry/ledger_split_baseline.json"
ASSET_IDS = {
    "ENV-EXPEDITION-L01-L-CORRIDOR",
    "ENV-EXPEDITION-L01-OFFICE-ROOM",
    "ENV-EXPEDITION-L01-BRIDGE-ROOM",
    "ENV-EXPEDITION-L01-ROOM-STD-25X25",
    "ENV-EXPEDITION-L01-ROOM-DB-70X50",
    "ENV-EXPEDITION-L01-BOSS-ROOM",
}
NOTE = (
    "2026-09-28：远征关卡01完成每房独立静态 TSCN 改造。"
    "start、room_01…room_10、boss、extraction 各由显式 static_layout_scene_path 加载 "
    "runtime/room_instances/f00_<room_id>/room_static_layout.tscn；场景只拥有顶层 PackedScene "
    "prefab 实例根，不递归改 prefab 内部 owner，避免内部网格/碰撞展开。"
    "静态 TSCN 接管视觉与固定碰撞，代码继续负责动态门、触发器、刷怪、导航、流送、存档、交互与撤离。"
    "专项验收 EXPEDITION_ROOM_STATIC_SCENES_OK checks=209 rooms=13、"
    "ROOM_TYPE_COMPONENT_REPLAY_OK checks=4540，退出码 0 且无 RID/ObjectDB/resource 泄漏。"
    "完整 verify_expedition_level01_flow 仍因数据库房/撤离房周界碰撞及入口安全房旧快照兼容字段失败；"
    "本次不把程序壳体转正为正式美术资产。"
)

wb = load_workbook(LEDGER, data_only=False)
ws = wb["资产主表"]
found = set()
for row in range(6, ws.max_row + 1):
    asset_id = str(ws.cell(row, 1).value or "").strip()
    if asset_id not in ASSET_IDS:
        continue
    found.add(asset_id)
    old_note = str(ws.cell(row, 25).value or "").strip()
    if NOTE not in old_note:
        ws.cell(row, 25).value = old_note + ("\n" if old_note else "") + NOTE
    ws.cell(row, 22).value = date.today().isoformat()
assert found == ASSET_IDS, sorted(ASSET_IDS - found)
for row in range(6, ws.max_row + 1):
    ws.cell(row, 18).value = dedupe_key_formula(row)
    ws.cell(row, 19).value = dedupe_result_formula(row, ws.max_row)

log = wb["域变更日志"]
existing_rows = [row for row in range(1, log.max_row + 1) if log.cell(row, 1).value == "v0.1.22"]
new_row = existing_rows[0] if existing_rows else log.max_row + 1
if not existing_rows:
    for col in range(1, log.max_column + 1):
        log.cell(new_row, col)._style = copy(log.cell(new_row - 1, col)._style)
        log.cell(new_row, col).number_format = log.cell(new_row - 1, col).number_format
    log.row_dimensions[new_row].height = log.row_dimensions[new_row - 1].height
values = [
    "v0.1.22",
    date.today().isoformat(),
    "运行时交付形态同步",
    "关卡场景 / 远征关卡01",
    "远征关卡01的13个具体房间完成独立静态TSCN化：每房一份room_static_layout.tscn，静态视觉与固定碰撞由顶层PackedScene prefab实例持有；运行时显式下发static_layout_scene_path，代码保留动态门、触发器、刷怪、导航、流送、存档、交互与撤离。专项静态场景/组件重放验收分别209/4540项通过，无RID、ObjectDB或资源泄漏。",
    "只更新6条既有远征房型资产的备注与更新时间，不新增AssetID、不改变制作状态/版本/路径/SHA。撤离房只有程序壳体TSCN，正式房型源仍未开始，不登记为正式资产。完整整关流程仍有数据库房/撤离房周界碰撞及入口安全房旧快照兼容红项。无损基线同步更新6行指纹、列摘要与分类计数；摘要锁死表不动。",
    "摩斯拉",
]
for col, value in enumerate(values, 1):
    log.cell(new_row, col).value = value
wb.save(LEDGER)

index = LedgerIndex.load(ROOT)
old = json.loads(BASELINE.read_text(encoding="utf-8"))
assets = {}
collected = []
sheet_digests = {}
category_counts = {}
for domain in index.domains:
    domain_wb = load_workbook(domain.path, data_only=False)
    for _, row_values in read_source_rows(domain_wb["资产主表"]):
        asset_id = str(row_values[0]).strip()
        category = str(row_values[2]).strip()
        assets[asset_id] = {"v": _row_digest(row_values), "c": category, "d": domain.key}
        collected.append((asset_id, row_values))
        category_counts[category] = category_counts.get(category, 0) + 1
    for sheet in domain.sheet_scope:
        if sheet in domain_wb.sheetnames:
            sheet_digests[sheet] = sheet_digest(domain_wb[sheet])
collected.sort(key=lambda item: item[0])
old.update({
    "source": "assets/registry/ledgers (current runtime-truth reconciliation)",
    "captured_at": date.today().isoformat(),
    "asset_count": len(assets),
    "assets": assets,
    "column_digests": {str(col): col_digest(collected, col) for col in CONTENT_COLUMNS},
    "sheet_digests": sheet_digests,
    "category_counts": category_counts,
})
BASELINE.write_text(json.dumps(old, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")
print("EXPEDITION_TSCN_LEDGER_SYNC_OK", len(found), "log_row", new_row)
