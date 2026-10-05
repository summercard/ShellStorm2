from pathlib import Path
import json
import openpyxl

ROOT = Path('I:/工作项目/shellstrom2/ShellStorm2')
BOOK = ROOT / 'assets/registry/ledgers/ShellStorm2_道具账本_v001.xlsx'
wb = openpyxl.load_workbook(BOOK, data_only=False)
out = {}
for sheet_name in ['总览', '资产主表', '3D-道具', '3D-物品', '域变更日志']:
    ws = wb[sheet_name]
    item = {'max_row': ws.max_row, 'max_column': ws.max_column}
    if sheet_name in ['资产主表', '3D-道具', '3D-物品', '域变更日志']:
        item['rows'] = []
        for r in range(1, ws.max_row + 1):
            vals = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
            if r <= 7 or r >= max(1, ws.max_row - 4) or any(k in ' | '.join('' if v is None else str(v) for v in vals) for k in ['TV', '电视', '椅', '梯', '收音机', 'RADIO', 'v0.1']):
                item['rows'].append({'row': r, 'values': vals})
    else:
        item['cells'] = {ref: ws[ref].value for ref in ['A6','B6','C6','E6','G6','B10','C10','B11','C11'] if ref in ws}
    if sheet_name == '资产主表':
        item['validations'] = [{'type': dv.type, 'formula1': dv.formula1, 'sqref': [str(r) for r in dv.sqref.ranges]} for dv in wb[sheet_name].data_validations.dataValidation]
    out[sheet_name] = item
(ROOT / '_scratch/radio_registration_context.json').write_text(json.dumps(out, ensure_ascii=False, indent=2, default=str) + '\n', encoding='utf-8')
