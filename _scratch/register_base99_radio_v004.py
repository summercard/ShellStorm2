import copy,hashlib,json,shutil,sys
from pathlib import Path
from datetime import datetime
from openpyxl import load_workbook
P=Path('I:/工作项目/shellstrom2/ShellStorm2');O=P/'outputs/base99_radio_v004'
sys.path.insert(0,str(P/'scripts'));sys.path.insert(0,str(P/'tools/asset_pipeline'))
from ledger_registry import LedgerIndex
from split_asset_ledger import read_source_rows,_row_digest,col_digest,CONTENT_COLUMNS,dedupe_key_formula,dedupe_result_formula,sheet_digest
index=LedgerIndex.load(P);ledger=index.path_for_category('道具');baseline=P/'assets/registry/ledger_split_baseline.json'
sha=lambda p:hashlib.sha256(p.read_bytes()).hexdigest()
m=json.loads((O/'asset_manifest.json').read_text(encoding='utf-8'))
assert json.loads((O/'validate_source_final.json').read_text(encoding='utf-8'))['passed']
assert json.loads((O/'validate_optimized_final.json').read_text(encoding='utf-8'))['passed']
assert 'BASE99_RADIO_OK' in (O/'verify_window_final.log').read_text(encoding='utf-8')
assert 'PLACEMENT_ACCEPTED=true' in (O/'placement.log').read_text(encoding='utf-8')
frozen=json.loads((O/'before_hashes.json').read_text(encoding='utf-8'))
for path in (ledger,baseline):
 assert sha(path)==frozen[path.relative_to(P).as_posix()], '登记文件已发生后续修改，停止事务：'+str(path)
for key,path in [('source',Path(m['source_blend'])),('optimized',Path(m['optimized_blend'])),('glb',Path(m['component_glb']))]:
 assert sha(path)==m['hashes_sha256'][key], '资产与最终清单哈希不一致：'+key
assert json.loads((O/'optimization_evidence.json').read_text(encoding='utf-8'))['passed']
backup=O/'backup_before_registration';assert not backup.exists();backup.mkdir()
shutil.copy2(ledger,backup/ledger.name);shutil.copy2(baseline,backup/baseline.name)
ID=m['asset_id'];wb=load_workbook(ledger)
def snap(w):return {(ws.title,c.coordinate):c.value for ws in w for row in ws for c in row if c.value is not None}
before=snap(wb);ws=wb['资产主表'];ps=wb['3D-道具']
assert ws['A27'].value==ID and ps['A8'].value==ID and ws['M27'].value=='v003'
assert ws['R27'].value==dedupe_key_formula(27) and ws['S27'].value==dedupe_result_formula(27,ws.max_row)
other={v[0]:_row_digest(v) for _,v in read_source_rows(ws) if v[0]!=ID}
dvs=[(str(d.sqref),d.formula1) for d in ws.data_validations.dataValidation]
prefab=P/m['runtime_prefab_target'];source=Path(m['source_blend']).relative_to(P).as_posix();opt=Path(m['optimized_blend']).relative_to(P).as_posix()
spec='宽0.828m × 深0.456m × 高0.822m；646面/1264三角形<800面；复古铜橙喷漆主壳、深色栅格、金属护框旋钮；单根三节斜金属天线29.35度；四材质金属M0.88/R0.32，哑光M0.04/R0.66，清漆M0.16/R0.16，灯M0/R0.38/能量1.5；仅独立小圆灯发光，off红常亮、A/B同绿常亮、离楼红待机；公共色盘外链/PaletteUV/Closest，UV换格白乘色；ItemRoot/Visual/StatusLight稳定；左键/E关→A→B→关音乐逻辑与摆位不变'
changes={'M27':'v004','N27':spec,'P27':source+'; '+opt+'; src/base3d/Base99Radio3D.gd; outputs/base99_radio_v004/asset_manifest.json; outputs/base99_radio_v004/optimization_evidence.json','T27':sha(prefab),'V27':datetime(2026,10,7),'Y27':'v004材质及灯表现深化；46号BATTERY柜顶与两个layout完全不改，local=(-1.95,6.97,-13.87143)/yaw10度；旧source/optimized/公共色盘哈希不变；红格(4,8)/绿格(5,5)左下0-based，GLTF翻V后UV1_offset=(0.1,0.3)；GLB SHA-256='+m['hashes_sha256']['glb']+'；source SHA-256='+m['hashes_sha256']['source']+'；optimized SHA-256='+m['hashes_sha256']['optimized']}
pschanges={'E8':source,'K8':spec,'O8':'v004','P8':'ItemRoot/Visual/StatusLight接口不变；独立Antenna；仅小灯UV红绿格切换，因公共色盘无白格不使用染色叠乘；源='+m['hashes_sha256']['source']+'；优化='+m['hashes_sha256']['optimized']+'；GLB='+m['hashes_sha256']['glb']+'；证据outputs/base99_radio_v004/；近远景off/a/b PNG与同相机v003对照'}
for c,v in changes.items():ws[c]=v
for c,v in pschanges.items():ps[c]=v
log=wb['域变更日志'];r=log.max_row+1
last=str(log.cell(r-1,1).value);nums=last.lstrip('v').split('.');nums[-1]=str(int(nums[-1])+1)
for c in range(1,8):log.cell(r,c)._style=copy.copy(log.cell(r-1,c)._style)
for c,v in enumerate(['v'+'.'.join(nums),'2026-10-07','99F收音机v004铜橙金属与常驻红绿灯','道具 / decor_prop','646面/1264三角；替换单根斜伸缩天线；off红、A/B同绿，离楼停止音乐但红灯待机；旧源与摆位不改','仅更新radio登记指纹、3D-道具摘要；严检UV和真实窗口交互截图通过','WorkBuddy'],1):log.cell(r,c,v)
wb.save(ledger);w=load_workbook(ledger);after=snap(w)
diffs=[{'sheet':s,'cell':c,'before':str(before.get((s,c))),'after':str(after.get((s,c)))} for s,c in sorted(set(before)|set(after)) if before.get((s,c))!=after.get((s,c))]
allowed={('资产主表',c) for c in changes}|{('3D-道具',c) for c in pschanges}|{('域变更日志',log.cell(r,c).coordinate) for c in range(1,8)}
assert all((d['sheet'],d['cell']) in allowed for d in diffs)
assert other=={v[0]:_row_digest(v) for _,v in read_source_rows(w['资产主表']) if v[0]!=ID}
assert dvs==[(str(d.sqref),d.formula1) for d in w['资产主表'].data_validations.dataValidation]
bl=json.loads(baseline.read_text(encoding='utf-8'));old=copy.deepcopy(bl)
values=next(v for _,v in read_source_rows(w['资产主表']) if v[0]==ID)
bl['assets'][ID]={'v':_row_digest(values),'c':'道具','d':'props'}
rows=[]
for domain in index.domains:rows.extend(read_source_rows(load_workbook(domain.path)['资产主表']))
bl['column_digests']={str(c):col_digest(rows,c) for c in CONTENT_COLUMNS}
bl['category_counts']={cat:sum(1 for _,v in rows if v[2]==cat) for cat in sorted({v[2] for _,v in rows})}
bl['sheet_digests']['3D-道具']=sheet_digest(w['3D-道具'])
assert all(bl['assets'][k]==v for k,v in old['assets'].items() if k!=ID)
assert bl['asset_count']==old['asset_count']
baseline.write_text(json.dumps(bl,ensure_ascii=False,indent=1)+'\n',encoding='utf-8')
m['hashes_sha256']['prefab']=sha(prefab)
(O/'asset_manifest.json').write_text(json.dumps(m,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
report={'passed':True,'asset_id':ID,'ledger':ledger.as_posix(),'changed_cells':diffs,'other_asset_rows_unchanged':True,'other_asset_baseline_fingerprints_unchanged':True,'formulas_data_validations_unchanged':True,'asset_count':bl['asset_count'],'hashes_sha256':m['hashes_sha256'],'backup':backup.as_posix()}
(O/'registration_evidence.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
print('RADIO_V004_REGISTRATION_OK',len(diffs))
