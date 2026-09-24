# -*- coding: utf-8 -*-
"""检索场景账本中与 l_corridor / 走廊 / 远征 相关的行，并 dump 资产主表表头。"""
import sys
import openpyxl
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

LEDGER = "assets/registry/ledgers/ShellStorm2_场景账本_v001.xlsx"
wb = load_workbook(LEDGER)

ws = wb["资产主表"]
print("### 《资产主表》表头 ###")
for c in range(1, ws.max_column + 1):
    v = ws.cell(row=4, column=c).value
    v2 = ws.cell(row=5, column=c).value
    print("  col%-3d %-18s | %s" % (c, c and openpyxl.utils.get_column_letter(c), v))
print()

KEYS = ("CORRIDOR", "corridor", "走廊", "L型", "L型走廊")
print("### 命中 l_corridor / 走廊 的行 ###")
hits = 0
for r in range(5, ws.max_row + 1):
    text = " ".join(str(ws.cell(row=r, column=c).value or "") for c in range(1, 26))
    if any(k in text for k in KEYS):
        hits += 1
        print("row %-4d A=%-42s B=%s" % (r, ws.cell(row=r, column=1).value, ws.cell(row=r, column=2).value))
        print("        E=%s  L=%s  O=%s" % (ws.cell(row=r, column=5).value,
                                            ws.cell(row=r, column=12).value,
                                            ws.cell(row=r, column=15).value))
if not hits:
    print("  （无）")
print()

print("### 命中 expedition / EXPEDITION / 远征 的行 ###")
for r in range(5, ws.max_row + 1):
    text = " ".join(str(ws.cell(row=r, column=c).value or "") for c in range(1, 26))
    if "EXPEDITION" in text.upper() or "远征" in text:
        print("row %-4d A=%-42s B=%s" % (r, ws.cell(row=r, column=1).value, ws.cell(row=r, column=2).value))
print()

print("### 资产主表末尾 6 行（看最新登记） ###")
for r in range(max(6, ws.max_row - 5), ws.max_row + 1):
    print("row %-4d A=%-42s B=%-30s L=%s O=%s" % (
        r, ws.cell(row=r, column=1).value, ws.cell(row=r, column=2).value,
        ws.cell(row=r, column=12).value, ws.cell(row=r, column=15).value))
print()

for name in ("3D-场景通用", "3D-设施"):
    s = wb[name]
    print("### 《%s》表头 + 命中行 ###" % name)
    for c in range(1, s.max_column + 1):
        print("  col%-3d %s" % (c, s.cell(row=4, column=c).value))
    for r in range(5, s.max_row + 1):
        text = " ".join(str(s.cell(row=r, column=c).value or "") for c in range(1, 17))
        if any(k in text for k in KEYS):
            print("  HIT row %-4d A=%s B=%s" % (r, s.cell(row=r, column=1).value, s.cell(row=r, column=2).value))
    print()
