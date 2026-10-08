from pathlib import Path
import json,shutil,sys
from openpyxl import load_workbook
R=Path.cwd();D=R/'_scratch/monitor_distance_ledger';D.mkdir(exist_ok=True)
sys.path.insert(0,str(R/'scripts'));from ledger_registry import LedgerIndex
p=LedgerIndex.load(R).path_for_category('敌人');w=load_workbook(p)
note='首次激活需存活玩家距Boss根节点≤10世界米；范围外黑屏静止。触发后完整播放activate6.4秒再战斗，离开范围不取消；休眠及读档续播。562项专项、29项正式远征流程通过。'
row=w['域变更日志'].max_row+1
jobs=[{'path':str(p),'sheets':{'资产主表':{'Y20':w['资产主表']['Y20'].value+'\n'+note},'3D-敌人':{'P10':w['3D-敌人']['P10'].value+'\n'+note},'域变更日志':{f'{c}{row}':v for c,v in zip('ABCDEFG',['v034规则修订','2026-10-08','Codex','ENM-BOSS-MONITOR002-3D',note,'docs/v0.1/development/2026-10-08_boss002_activation_distance.md','动画数据保持不变'])}}}]
cp=R/'docs/v0.1/data/ShellStorm2_游戏内容数据库_v010.xlsx';cw=load_workbook(cp)
jobs.append({'path':str(cp),'sheets':{'怪物与Boss':{'O17':cw['怪物与Boss']['O17'].value+'；首次10米触发，出场播完后战斗','L17':'用户10米激活要求；技能设计r4'},'Boss002技能设计':{'B24':'存活玩家≤10米首次触发activate6.4s，结束后战斗','B34':'10米外黑屏静止；触发后离开范围继续播完；started/进度/完成标记随实例保存；休眠暂停，旧档兼容。','B36':'docs/v0.1/design/Boss002显示器技能设计.md（r4）；动画设计r28'}}})
(D/'changes.json').write_text(json.dumps(jobs,ensure_ascii=False,indent=2),encoding='utf-8')
shutil.copy2(R/'assets/registry/ledger_split_baseline.json',D/'baseline_before.json')
for name in ['edit.mjs','merge.py']:
 s=(R/'_scratch/monitor034_ledger'/name).read_text(encoding='utf-8').replace('_scratch/monitor034_ledger','_scratch/monitor_distance_ledger').replace('if(i===2)','if(i===1)')
 (D/name).write_text(s,encoding='utf-8')
p=R/'docs/v0.1/design/Boss002显示器技能设计.md';s=p.read_text(encoding='utf-8').replace('设计r3','设计r4');p.write_text(s,encoding='utf-8')
for file in ['Boss002显示器动画设计.md','远征关卡01设计.md']:
 p=R/'docs/v0.1/design'/file;s=p.read_text(encoding='utf-8')
 s=s.replace('首次正式激活由Enemy3D','存活玩家接近Boss根节点10世界米内才触发首次出场；触发后离开范围也播完，之后才战斗。首次正式激活由Enemy3D')
 s=s.replace('首波显示器Boss沿既有进房激活流程播放activate 6.4秒','首波显示器Boss沿既有进房流程生成，存活玩家接近其根节点10世界米内才播放activate 6.4秒；10米外黑屏静止，触发后离开范围仍播完再战斗')
 p.write_text(s,encoding='utf-8')
p=R/'docs/v0.1/development/CHANGELOG.md';s=p.read_text(encoding='utf-8');p.write_text('## 2026-10-08 Boss002 十米激活\n\n玩家进入10米后播放完整出场，再进入战斗。[验证记录](2026-10-08_boss002_activation_distance.md)。\n\n'+s,encoding='utf-8')
p=R/'docs/v0.1/MODULE_INDEX.md';s=p.read_text(encoding='utf-8').replace('activate出场、原版侧翘横移move','10米触发activate后战斗、原版侧翘横移move');p.write_text(s,encoding='utf-8')
print('Prepared distance rule ledger and design updates')
