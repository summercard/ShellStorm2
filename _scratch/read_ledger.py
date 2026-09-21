import os
out = []
def log(*a):
    s = " ".join(str(x) for x in a)
    out.append(s.encode("ascii", "replace").decode("ascii"))

LEDGER = "I:/工作项目/shellstrom2/ShellStorm2/assets/registry/ledgers/ShellStorm2_敌人账本_v001.xlsx"
log("LEDGER exists:", os.path.isfile(LEDGER))
log("LEDGER path:", LEDGER)

try:
    import openpyxl
    log("openpyxl OK", openpyxl.__version__)
except Exception as e:
    log("NO openpyxl:", e)
    openpyxl = None

if openpyxl:
    wb = openpyxl.load_workbook(LEDGER, data_only=True)
    log("SHEETS:", wb.sheetnames)
    for ws in wb.worksheets:
        log("---- SHEET: %s  dims=%s" % (ws.title, ws.dimensions))
        # print header row
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        header = rows[0]
        log("  HEADER: " + " | ".join(str(h) for h in header))
        # print all rows that mention ranged/cast/保安/security
        kw = ["ranged", "cast", "保安", "security", "guard", "ENM-RANGED", "RANGED"]
        for ri, row in enumerate(rows[1:], start=2):
            line = " | ".join("" if c is None else str(c) for c in row)
            if any(k.lower() in line.lower() for k in kw):
                log("  ROW%d: %s" % (ri, line))

with open("I:/工作项目/shellstrom2/ShellStorm2/_scratch/read_ledger_result.log", "w", encoding="ascii") as f:
    f.write("\n".join(out) + "\n")
