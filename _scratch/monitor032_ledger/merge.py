"""Merge only Artifact Tool-authored cells/new sheet; preserve native workbook parts."""
import json,zipfile,xml.etree.ElementTree as E,re,copy,shutil,hashlib,sys
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string
R=Path.cwd();D=R/'_scratch/monitor032_ledger';M='http://schemas.openxmlformats.org/spreadsheetml/2006/main';REL='http://schemas.openxmlformats.org/officeDocument/2006/relationships';P='http://schemas.openxmlformats.org/package/2006/relationships';CT='http://schemas.openxmlformats.org/package/2006/content-types';ns={'m':M}
E.register_namespace('',M);E.register_namespace('r',REL)
jobs=json.loads((D/'changes.json').read_text(encoding='utf-8'))
def package_xml(element,namespace):
 E.register_namespace('',namespace)
 result=E.tostring(element,encoding='utf-8',xml_declaration=True)
 E.register_namespace('',M)
 return result
def sheets(z):
 rels={e.attrib['Id']:e.attrib['Target'] for e in E.fromstring(z.read('xl/_rels/workbook.xml.rels'))}
 return {s.attrib['name']:'xl/'+rels[s.attrib['{'+REL+'}id']].lstrip('/').removeprefix('xl/') for s in E.fromstring(z.read('xl/workbook.xml')).find('m:sheets',ns)}
def styles_map(old,new):
 maps={}
 for tag in ['fonts','fills','borders']:
  a=old.find('m:'+tag,ns);b=new.find('m:'+tag,ns);offset=len(a);maps[tag]={i:i+offset for i in range(len(b))}
  for item in b:a.append(copy.deepcopy(item))
  a.set('count',str(len(a)))
 a=old.find('m:numFmts',ns)
 if a is None:a=E.Element('{'+M+'}numFmts',{'count':'0'});old.insert(0,a)
 used=[int(item.get('numFmtId')) for item in a];nextid=max([163]+used)+1;formats={}
 b=new.find('m:numFmts',ns)
 if b is not None:
  for item in b:
   formats[int(item.get('numFmtId'))]=nextid;v=copy.deepcopy(item);v.set('numFmtId',str(nextid));a.append(v);nextid+=1
 a.set('count',str(len(a)))
 a=old.find('m:cellXfs',ns);b=new.find('m:cellXfs',ns);offset=len(a);result={i:i+offset for i in range(len(b))}
 for item in b:
  v=copy.deepcopy(item)
  for key,tag in [('fontId','fonts'),('fillId','fills'),('borderId','borders')]:v.set(key,str(maps[tag][int(v.get(key,'0'))]))
  if int(v.get('numFmtId','0')) in formats:v.set('numFmtId',str(formats[int(v.get('numFmtId'))]))
  v.set('xfId','0');a.append(v)
 a.set('count',str(len(a)))
 return result
for i,job in enumerate(jobs):
 if len(sys.argv)>1 and str(i) not in sys.argv[1:]:continue
 path=Path(job['path']);backup=D/f'before_job{i}.xlsx';shutil.copy2(path,backup);before=load_workbook(path)
 replacements={};new_sheets=job.get('new_sheets',[])
 with zipfile.ZipFile(path) as original,zipfile.ZipFile(D/f'cells{i}.xlsx') as generated:
  paths=sheets(original);newpaths=sheets(generated)
  strings=[]
  if 'xl/sharedStrings.xml' in generated.namelist():strings=[''.join(x.itertext()) for x in E.fromstring(generated.read('xl/sharedStrings.xml'))]
  stylemap={}
  if new_sheets:
   styles=E.fromstring(original.read('xl/styles.xml'));stylemap=styles_map(styles,E.fromstring(generated.read('xl/styles.xml')));replacements['xl/styles.xml']=E.tostring(styles,encoding='utf-8',xml_declaration=True)
  for name,changes in job['sheets'].items():
   root=E.fromstring(generated.read(newpaths[name]));cells={c.attrib['r']:c for c in root.findall('.//m:c',ns)}
   for c in cells.values():
    if c.get('t')=='s':
     value=strings[int(c.find('m:v',ns).text)];c.remove(c.find('m:v',ns));c.set('t','inlineStr');E.SubElement(E.SubElement(c,'{'+M+'}is'),'{'+M+'}t').text=value
   if name in new_sheets:
    for c in cells.values():c.set('s',str(stylemap[int(c.get('s','0'))]))
    number=max([int(re.search(r'sheet(\d+)\.xml',p)[1]) for p in paths.values() if re.search(r'sheet(\d+)\.xml',p)])+1
    target=f'xl/worksheets/sheet{number}.xml';paths[name]=target;replacements[target]=E.tostring(root,encoding='utf-8',xml_declaration=True)
    wb=E.fromstring(original.read('xl/workbook.xml'));relationships=E.fromstring(original.read('xl/_rels/workbook.xml.rels'));types=E.fromstring(original.read('[Content_Types].xml'))
    rid='rId'+str(max(int(re.search(r'\d+',r.get('Id'))[0]) for r in relationships)+1)
    E.SubElement(wb.find('m:sheets',ns),'{'+M+'}sheet',{'name':name,'sheetId':str(max(int(s.get('sheetId')) for s in wb.find('m:sheets',ns))+1),'{'+REL+'}id':rid})
    E.SubElement(relationships,'{'+P+'}Relationship',{'Id':rid,'Type':REL+'/worksheet','Target':f'worksheets/sheet{number}.xml'})
    E.SubElement(types,'{'+CT+'}Override',{'PartName':'/'+target,'ContentType':'application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml'})
    replacements.update({'xl/workbook.xml':E.tostring(wb,encoding='utf-8',xml_declaration=True),'xl/_rels/workbook.xml.rels':package_xml(relationships,P),'[Content_Types].xml':package_xml(types,CT)})
    continue
   text=original.read(paths[name]).decode();previous=before[name]
   for address,value in changes.items():
    cell=copy.deepcopy(cells[address]);old=re.search(r'<c\b[^>]*\br="'+address+r'"[^>]*(?:/>|>.*?</c>)',text,re.S)
    if old:
     style=re.search(r'\bs="([^"]+)"',old.group());cell.set('s',style[1] if style else '0')
    else:
     rownum=int(re.search(r'\d+',address)[0]);colnum=column_index_from_string(re.match(r'[A-Z]+',address)[0])
     style=previous.cell(min(rownum-1,previous.max_row),colnum).style_id;cell.set('s',str(style))
    xml=E.tostring(cell,encoding='unicode')
    if old:text=text[:old.start()]+xml+text[old.end():]
    else:
     row=re.search(r'\d+',address)[0];match=re.search(r'<row\b[^>]*\br="'+row+r'"[^>]*>.*?</row>',text)
     if match:text=text[:match.end()-6]+xml+text[match.end()-6:]
     else:
      height=' ht="120" customHeight="1"' if name in ['3D-敌人','怪物与Boss'] else ' ht="90" customHeight="1"' if name=='敌人动画与状态' else ''
      text=text.replace('</sheetData>',f'<row r="{row}"{height}>{xml}</row></sheetData>')
   maxrow=max([previous.max_row]+[int(re.search(r'\d+',c)[0]) for c in changes])
   text=re.sub(r'(<dimension\b[^>]*\bref="[A-Z]+\d+:[A-Z]+)\d+(\")',lambda m:m[1]+str(maxrow)+m[2],text)
   text=re.sub(r'(<autoFilter\b[^>]*\bref="[A-Z]+\d+:[A-Z]+)\d+(\")',lambda m:m[1]+str(maxrow)+m[2],text)
   # Add canonical pipeline status choices to this domain's existing validation only.
   if name=='资产主表' and i==2:
    stages='design_only,authored,exported_pending_godot_validation,validated,active'
    text=re.sub(r'(<dataValidation\b[^>]*sqref="K[^>]*>.*?<formula1>)(.*?)(</formula1>)',lambda m:m[1]+m[2].rstrip('&quot;').rstrip('"')+','+stages+'&quot;'+m[3] if 'active' not in m[2] else m[0],text)
   replacements[paths[name]]=text.encode()
  tmp=D/f'merged{i}.xlsx'
  with zipfile.ZipFile(tmp,'w',zipfile.ZIP_DEFLATED) as out:
   for info in original.infolist():out.writestr(info,replacements.get(info.filename,original.read(info.filename)))
   for name,data in replacements.items():
    if name not in original.namelist():out.writestr(name,data)
 after=load_workbook(tmp)
 # openpyxl normalizes empty borders on merged cells after adding a new style.
 # Verify the original native style definitions and cell style references instead.
 with zipfile.ZipFile(backup) as original,zipfile.ZipFile(tmp) as final:
  old_styles=E.fromstring(original.read('xl/styles.xml'));final_styles=E.fromstring(final.read('xl/styles.xml'))
  for tag in ['fonts','fills','borders','cellXfs']:
   old_entries=old_styles.find('m:'+tag,ns);new_entries=final_styles.find('m:'+tag,ns)
   assert all(E.tostring(a)==E.tostring(b) for a,b in zip(old_entries,new_entries)),('native styles changed',tag)
  for name,target in sheets(original).items():
   old_cells={c.get('r'):c.get('s','0') for c in E.fromstring(original.read(target)).findall('.//m:c',ns)}
   new_cells={c.get('r'):c.get('s','0') for c in E.fromstring(final.read(target)).findall('.//m:c',ns)}
   assert all(new_cells.get(address)==style for address,style in old_cells.items()),('existing style reference changed',name)
 for name in before.sheetnames:
  expected=job['sheets'].get(name,{})
  for row in before[name]:
   for c in row:
    actual=after[name][c.coordinate];assert actual.value==expected.get(c.coordinate,c.value),(name,c.coordinate,'value changed outside scope')
 for name,changes in job['sheets'].items():
  for address,value in changes.items():assert after[name][address].value==value,(name,address)
 assert path.read_bytes()==backup.read_bytes(),'Concurrent workbook edit; regenerate'
 shutil.copy2(tmp,path);print('Patched stage/workbook',i,path.name)
(D/'preservation.json').write_text(json.dumps({'jobs':len(jobs),'unrelated_cell_values_and_formulas_unchanged':True,'existing_styles_unchanged':True,'new_sheet':'Boss002技能设计'},ensure_ascii=False,indent=2),encoding='utf-8')
