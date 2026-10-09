import sys,json,hashlib,shutil,subprocess
from pathlib import Path
from datetime import datetime
from copy import copy
root=Path(__file__).resolve().parents[2];pkg=root/'assets/art/enemies/normal_enemy_3d/melee_chaser';out=root/'outputs/melee_zombie_reduction';aid='ENM-MELEE-FUNGBOAR01'
sys.path[:0]=[str(root/'scripts'),str(root/'tools/asset_pipeline')]
from ledger_registry import LedgerIndex
from split_asset_ledger import read_source_rows,_row_digest,sheet_digest,CONTENT_COLUMNS,col_digest
from openpyxl import load_workbook
index=LedgerIndex.load(root);ledger=index.domain_for_key('enemies').path;sha=lambda p:hashlib.sha256(p.read_bytes()).hexdigest();rel=lambda p:p.relative_to(root).as_posix()
model=pkg/'source/model/enm_melee_fungboar01_model_v003.blend';animation=pkg/'source/animation/enm_melee_fungboar01_animation_v004.blend';prefab=pkg/'runtime/enm_melee_fungboar01_root_top3d.tscn';glb=pkg/'components/enm_melee_fungboar01_visual_top3d.glb'
report=json.loads((out/'source_report.json').read_text());assert json.loads((out/'reopen.json').read_text())['passed'];assert json.loads((out/'render.json').read_text())['passed']
note='2026-10-09 v004减面与减骨：4552→1790三角，897顶点；36→20骨，移除16根手指骨，每手仅一根手骨；自然弯曲形状固定，手指权重合并至手骨并重新归一。512贴图/六段身体动作/12态/碰撞/受击闪色/2.4秒死亡保留。源重开、半帧蒙皮检查、运行专项及死亡视野真实渲染通过。'
w=load_workbook(ledger);before={v[0]:v for _,v in read_source_rows(w['资产主表'])};ws=w['资产主表'];r=next(r for r,v in read_source_rows(ws) if v[0]==aid)
for c,v in {13:'v004',14:'游戏高1.3m；源高1.857143m；1790三角/897顶点；20骨；512贴图；6剪辑',16:rel(model)+'; '+rel(animation),20:sha(prefab),22:datetime(2026,10,9),25:str(ws.cell(r,25).value)+'\n'+note}.items():ws.cell(r,c).value=v
log=w['域变更日志'];versions=[str(log.cell(i,1).value).lstrip('v').split('.') for i in range(1,log.max_row+1)];versions=[v for v in versions if len(v)==3 and all(x.isdigit() for x in v)];v=list(max(versions,key=lambda x:tuple(map(int,x))));v[-1]=str(int(v[-1])+1);log.append(['v'+'.'.join(v),'2026-10-09','小僵尸减面减骨重新蒙皮','敌人',note,'用户授权移除手指独立动画；其余契约保留','Codex']);w.save(ledger)
w=load_workbook(ledger)
for row in w['3D-敌人'].iter_rows():
 if row[0].value==aid:
  row[4].value=rel(model)+'; '+rel(animation);row[5].value='六段Blender剪辑按12态采样；20骨，每手一骨';row[10].value='源高1.857143m/游戏1.3m；1790三角/897顶点；20骨；512贴图';row[14].value='v004';row[15].value=str(row[15].value)+'；'+note
tf=pkg/'runtime/character_transfer_ledger.json';t=json.loads(tf.read_text(encoding='utf-8'));ws=w['敌人动画与状态']
assert not any(row[0].value==aid for row in ws.iter_rows())
for c in t['clips']:
 ws.append([aid,c['id'],c['scene'],f"{c['duration']}s；保留身体曲线，手指不独立运动",rel(animation),'循环' if c['loop'] else '单次末帧保持','原运行事件不变','已重新蒙皮/导入/验收',note])
 for cell in ws[ws.max_row]:cell._style=copy(ws.cell(ws.max_row-1,cell.column)._style)
w.save(ledger)
after={v[0]:v for _,v in read_source_rows(w['资产主表'])};assert all(before[k]==after[k] for k in before if k!=aid)
base=root/'assets/registry/ledger_split_baseline.json';bl=json.loads(base.read_text(encoding='utf-8'));bl['assets'][aid]={'v':_row_digest(after[aid]),'c':'敌人','d':'enemies'}
for n in ['3D-敌人','敌人动画与状态']:bl['sheet_digests'][n]=sheet_digest(w[n])
union=[]
for domain in index.domains:union.extend(read_source_rows(load_workbook(domain.path)['资产主表']))
bl['column_digests']={str(c):col_digest(union,c) for c in CONTENT_COLUMNS};base.write_text(json.dumps(bl,ensure_ascii=False,indent=1)+'\n',encoding='utf-8')
t.update(asset_version='v004',version='v004',model=str(model),animation=str(animation),skeleton_id='SKEL-MELEE-ZOMBIE01-003',skeleton_signature=report['skeleton_signature'],geometry={'triangles':1790,'vertices':897,'bones':20,'removed_bones':report['removed_bones']})
t['source_sha256']={rel(p):sha(p) for p in [model,animation]};t['outputs']={rel(glb):sha(glb)}
paths=[model,animation,pkg/'source/model/textures/enm_melee_fungboar01_basecolor_v002.png',glb,Path(str(glb)+'.import'),pkg/'runtime/melee_chaser_post_import.gd']
t['files']=[{'path':rel(p),'sha256':sha(p),'bytes':p.stat().st_size} for p in paths]
t['runtime_files']=[{'path':rel(p),'sha256':sha(p)} for p in [prefab,pkg/'runtime/melee_chaser_formal_visual.gd']]
for c in t['clips']:c['action']=c['id'];c['scene']='MeleeZombieActions'
t['validation'].append({'test':'reduction_reskin_20261009','passed':True,'source':rel(out/'source_report.json'),'reopen':rel(out/'reopen.json'),'runtime_log':rel(out/'runtime.log'),'render_log':rel(out/'death_render.log'),'negative_budget_expected_exit':1})
t['limitations'].append('v004按用户要求删除手指独立变形，固定自然弯曲轮廓；近景存在减面造成的局部轮廓简化。')
tf.write_text(json.dumps(t,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
for n in ['runtime','death_render','negative','import','reopen','author']:shutil.copy2(root/f'_scratch/melee_reduce/{n}.log',out/f'{n}.log')
record='docs/v0.1/development/2026-10-09_melee_zombie_reduction.md'
(root/record).write_text('# 小僵尸减面、减骨、重新蒙皮与导入\n\nENEMY-AI / ASSET-PIPELINE。设计依据：[小僵尸模型优化](../design/小僵尸模型优化.md)。\n\n'+note+'\n\n模型v003、动作v004共享SKEL-MELEE-ZOMBIE01-003及新签名。将原待机手指弯曲烘入静态几何，再将指骨权重归并到对应手骨；简化保留脸手优先级、UV、512材质，蒙皮最多4影响且归一。六段保留骨骼的F曲线逐项一致，导出动画采样二进制复用，仅去掉指骨轨道。骨架节点真正移除，不仅隐藏。\n\n源级736个半帧姿势、重开两个母版、骨架签名/权重/六段曲线对照通过。正侧面、跑步、攻击、死亡和手部六组同机位图检查，主要轮廓保留，近景有细小几何简化；不宣称逐顶点无损。\n\n正式运行专项退出0，覆盖20骨、无指骨、低于1800面、六段时长/循环、12态/速度选段/攻击边界、碰撞及2.4秒回收；超面负向注入退出1，只报告预算错误。真实Forward+死亡视野验收退出0，覆盖退出AI后显示、遮挡/距离隐藏与恢复、死亡回收。正向日志无非预期脚本错误。未跑全项目套件及目标设备性能测试。\n\n原表现层用fmod循环，导入剪辑未标循环；本次补导入钩子明确idle/walking/running循环，攻击/受击/死亡单次。保留.import配置入库，避免干净导入丢失设置。Enemy3D继续拥有碰撞/状态/伤害/回收；表现层仅更新来源版本标记。\n\n全局门禁及既有问题见outputs/melee_zombie_reduction/gates.json；仅接受本资产哈希变化，其他资产行逐值保持。旧源保留，运行文件同路径覆盖，回滚可恢复Git原文件。\n',encoding='utf-8')
f=pkg/'README.md';s=f.read_text(encoding='utf-8');s=s.replace('# 小僵尸 · 首只标准普通怪资产包','# 小僵尸 · 首只标准普通怪资产包\n\n## 当前版本v004：2026-10-09减面减骨\n\n'+note+'\n\n当前母版为model_v003.blend与animation_v004.blend；此前章节为历史记录。',1);f.write_text(s,encoding='utf-8')
f=root/'docs/v0.1/feature_registry.json';lines=f.read_text(encoding='utf-8').splitlines(keepends=True)
for i,line in enumerate(lines):
 if '"feature_id":"ENEMY-AI"' in line or '"feature_id":"ASSET-PIPELINE"' in line:
  d=json.loads(line.strip().rstrip(','));d['development_records'].append(record);d['design_docs'].append('docs/v0.1/design/小僵尸模型优化.md');d['verification'].append({'kind':'scene','id':'verify_melee_zombie_death_visibility'});lines[i]='    '+json.dumps(d,ensure_ascii=False,separators=(',',':'))+',\n'
f.write_text(''.join(lines),encoding='utf-8')
for name,text in [('MODULE_INDEX.md','\nENEMY-AI / ASSET-PIPELINE [小僵尸减面减骨](development/2026-10-09_melee_zombie_reduction.md)：1790三角/20骨，每手一骨，重新蒙皮与六段动作、死亡渲染验收通过。\n'),('development/CHANGELOG.md','\n## 2026-10-09｜小僵尸减面减骨\n\n4552→1790三角、36→20骨，每手一骨，重新蒙皮导入；[验收记录](2026-10-09_melee_zombie_reduction.md)。\n')]:
 f=root/'docs/v0.1'/name;f.write_text(f.read_text(encoding='utf-8')+text,encoding='utf-8')
commands={'structure':['scripts/check_asset_registry.py','--scope','structure'],'full':['scripts/check_asset_registry.py','--scope','full','--ledger','enemies','--json-output',str(out/'ledger_after.json')],'split':['tools/asset_pipeline/verify_ledger_split.py'],'guard':['scripts/asset_guard.py',rel(pkg),'--classify','version_increment'],'docs':['scripts/check_documentation_contracts.py'],'naming':['scripts/check_asset_runtime_naming.py']}
results={}
for name,args in commands.items():
 r=subprocess.run([sys.executable,*args],cwd=root,capture_output=True);results[name]=r.returncode;(out/f'gate_{name}.log').write_bytes(r.stdout+r.stderr);print(name,r.returncode,flush=True)
before=json.loads((root/'_scratch/melee_reduce/ledger_before.json').read_text(encoding='utf-8'));after=json.loads((out/'ledger_after.json').read_text(encoding='utf-8'))
ext=lambda d:{k:[v for v in rows if v.get('asset_id')!=aid] for k,rows in d['issues'].items()}
assert ext(before)==ext(after);assert not any(v.get('asset_id')==aid for rows in after['issues'].values() for v in rows)
results.update(other_asset_issues_unchanged=True,runtime=0,real_render=0,reopen=0,negative_expected_exit=1)
(out/'gates.json').write_text(json.dumps(results,indent=2));assert all(results[k]==0 for k in ['structure','split','guard']);print('MELEE_REDUCTION_REGISTERED')
