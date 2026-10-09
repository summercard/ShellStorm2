from __future__ import annotations

import copy
import hashlib
import json
import shutil
from collections import Counter
from datetime import datetime
from pathlib import Path

import openpyxl

ROOT = Path('I:/工作项目/shellstrom2/ShellStorm2')
LEDGER = ROOT / 'assets/registry/ledgers/ShellStorm2_道具账本_v001.xlsx'
MASTER = ROOT / 'assets/registry/ShellStorm2_美术资产台账_v001.xlsx'
BASELINE = ROOT / 'assets/registry/ledger_split_baseline.json'
OUT = ROOT / 'outputs/base99_radio_sound_v001'
BACKUP = OUT / 'backup_20261009'
ASSET_ID = 'PRP-BASE99-RADIO-3D'
UPDATED_AT = datetime(2026, 10, 9)


def text(value) -> str:
    return '' if value is None else str(value).strip()


def row_values(ws, row: int, width: int = 25):
    return [ws.cell(row, col).value for col in range(1, width + 1)]


def row_digest(values) -> str:
    return hashlib.sha256('\x1f'.join(text(values[col - 1]) for col in range(1, 26) if col not in (18, 19)).encode('utf-8')).hexdigest()


def col_digest(rows, col: int) -> str:
    ordered = sorted(rows, key=lambda item: (text(item[0]), text(item[1][col - 1])))
    return hashlib.sha256('\x1e'.join(text(values[col - 1]) for _, values in ordered).encode('utf-8')).hexdigest()


def read_asset_rows(workbook):
    ws = workbook['资产主表']
    return [(row, row_values(ws, row)) for row in range(6, ws.max_row + 1) if text(ws.cell(row, 1).value)]


def sheet_digest(ws) -> str:
    parts = [f'merge:{rng}' for rng in sorted(str(r) for r in ws.merged_cells.ranges)]
    cells = [
        (cell.row, cell.column, cell.coordinate, cell.value)
        for row in ws.iter_rows()
        for cell in row
        if cell.value not in (None, '')
    ]
    if cells:
        parts.append(f'bbox={min(c[0] for c in cells)}:{max(c[0] for c in cells)}x{min(c[1] for c in cells)}:{max(c[1] for c in cells)}')
    else:
        parts.append('bbox=empty')
    parts.extend(f'{coordinate}={value!r}' for _, _, coordinate, value in sorted(cells, key=lambda item: (item[0], item[1])))
    return hashlib.sha256('\n'.join(parts).encode('utf-8')).hexdigest()


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    BACKUP.mkdir(parents=True, exist_ok=True)
    for path in (LEDGER, MASTER, BASELINE):
        backup = BACKUP / path.name
        if not backup.exists():
            shutil.copy2(path, backup)

    workbook = openpyxl.load_workbook(LEDGER)
    asset_sheet = workbook['资产主表']
    prefab_sheet = workbook['3D-道具']
    log_sheet = workbook['域变更日志']
    assert asset_sheet['A27'].value == ASSET_ID
    assert prefab_sheet['A8'].value == ASSET_ID

    audio_spec = (
        '空间声音契约v001：唯一AudioStreamPlayer3D；Music总线；'
        'unit_size=18.0m；max_distance=78.0m；attenuation_model=inverse_square；'
        '99F正式边界Rect2(-50,-35,100,80)内source volume=0dB；'
        '边界外起始-18dB并按4dB/m继续衰减，最低-42dB；不复制播放；'
        '收音机离开99F停播/返回后手动开启，模型摆位、音符表现与音乐控制不变。'
    )
    evidence_ref = 'outputs/base99_radio_sound_v001/sound_acceptance.json'
    ledger_evidence_ref = 'outputs/base99_radio_sound_v001/ledger_evidence.json'

    old_n27 = text(asset_sheet['N27'].value)
    old_y27 = text(asset_sheet['Y27'].value)
    old_k8 = text(prefab_sheet['K8'].value)
    old_p8 = text(prefab_sheet['P8'].value)
    asset_sheet['N27'] = old_n27 + ' ' + audio_spec
    asset_sheet['Y27'] = old_y27 + '；' + audio_spec + f' 验收={evidence_ref}。'
    prefab_sheet['K8'] = old_k8 + '；' + audio_spec
    prefab_sheet['P8'] = old_p8 + f'；空间声音契约={audio_spec}；验收={evidence_ref}；账本证据={ledger_evidence_ref}'

    log_row = log_sheet.max_row + 1
    for col in range(1, 8):
        log_sheet.cell(log_row, col)._style = copy.copy(log_sheet.cell(log_row - 1, col)._style)
    log_values = [
        'v0.1.17', UPDATED_AT, '99F收音机空间声音v001', '道具 / 场景可交互道具',
        f'PRP-BASE99-RADIO-3D runtime v005.1 增加可量化3D空间声音契约：{audio_spec}',
        '仅更新资产主表第27行、3D-道具第8行、域变更日志和baseline；不改模型、摆位、音符表现、音乐控制或唯一AudioStreamPlayer3D结构。新sound test通过；原radio headless鼠标路径与music notes超时按原始日志单独记录。',
        'CodeBuddy',
    ]
    for col, value in enumerate(log_values, 1):
        log_sheet.cell(log_row, col).value = value
    workbook.save(LEDGER)

    # 重新读取确认：除本事务白名单外，三张受影响分页的非目标单元格内容保持不变。
    after = openpyxl.load_workbook(LEDGER, data_only=False)
    before = openpyxl.load_workbook(BACKUP / LEDGER.name, data_only=False)
    allowed = {
        ('资产主表', 'N27'), ('资产主表', 'Y27'),
        ('3D-道具', 'K8'), ('3D-道具', 'P8'),
    }
    changed = []
    for sheet_name in ('资产主表', '3D-道具', '域变更日志'):
        old_ws = before[sheet_name]
        new_ws = after[sheet_name]
        for row in range(1, max(old_ws.max_row, new_ws.max_row) + 1):
            for col in range(1, max(old_ws.max_column, new_ws.max_column) + 1):
                old_cell = old_ws.cell(row, col)
                new_cell = new_ws.cell(row, col)
                if old_cell.value != new_cell.value:
                    coord = new_cell.coordinate
                    if sheet_name == '域变更日志' and row == log_row:
                        changed.append({'sheet': sheet_name, 'cell': coord, 'before': old_cell.value, 'after': new_cell.value})
                    else:
                        assert (sheet_name, coord) in allowed, (sheet_name, coord)
                        changed.append({'sheet': sheet_name, 'cell': coord, 'before': old_cell.value, 'after': new_cell.value})

    baseline = json.loads(BASELINE.read_text(encoding='utf-8'))
    index = json.loads((ROOT / 'assets/registry/ledger_index.json').read_text(encoding='utf-8'))
    all_rows = []
    for domain in index['domains']:
        path = ROOT / domain['file'] if 'file' in domain else ROOT / 'assets/registry/ledgers' / domain['file']
        if not path.exists():
            path = ROOT / 'assets/registry/ledgers' / domain['file']
        wb = openpyxl.load_workbook(path, read_only=True, data_only=False)
        all_rows.extend((asset_id, values) for _, values in read_asset_rows(wb) for asset_id in [text(values[0])])
        wb.close()
    target = next(values for asset_id, values in all_rows if asset_id == ASSET_ID)
    old_count = baseline['asset_count']
    baseline['assets'][ASSET_ID] = {'v': row_digest(target), 'c': '道具', 'd': 'props'}
    assert len(baseline['assets']) == old_count
    baseline['asset_count'] = len(baseline['assets'])
    baseline['column_digests'] = {str(col): col_digest(all_rows, col) for col in range(1, 26) if col not in (18, 19)}
    baseline['category_counts'] = dict(sorted(Counter(text(values[2]) for _, values in all_rows).items()))
    baseline['sheet_digests']['3D-道具'] = sheet_digest(after['3D-道具'])
    baseline['captured_at'] = '2026-10-09'
    BASELINE.write_text(json.dumps(baseline, ensure_ascii=False, indent=1) + '\n', encoding='utf-8')

    result = {
        'passed': True,
        'asset_id': ASSET_ID,
        'updated_at': UPDATED_AT.strftime('%Y-%m-%d'),
        'ledger': 'assets/registry/ledgers/ShellStorm2_道具账本_v001.xlsx',
        'asset_sheet': '资产主表!27',
        'prefab_sheet': '3D-道具!8',
        'change_log_sheet': f'域变更日志!{log_row}',
        'changed_cells': changed,
        'allowed_content_scope': ['runtime audio spatial specification', 'ledger evidence', 'baseline evidence'],
        'forbidden_content_unchanged': ['model placement', 'music notes presentation', 'music control logic', 'duplicate playback'],
        'audio_spatial_contract': {
            'source_count': 1,
            'node_type': 'AudioStreamPlayer3D',
            'bus': 'Music',
            'unit_size_m': 18.0,
            'max_distance_m': 78.0,
            'attenuation_model': 'inverse_square',
            'base99_rect2': 'Rect2(-50,-35,100,80)',
            'outside_edge_drop_db': -18.0,
            'outside_drop_db_per_m': 4.0,
            'outside_floor_db': -42.0,
        },
        'baseline': {
            'path': 'assets/registry/ledger_split_baseline.json',
            'backup': f'outputs/base99_radio_sound_v001/backup_20261009/{BASELINE.name}',
            'asset_count': baseline['asset_count'],
            'captured_at': baseline['captured_at'],
        },
        'backup_dir': 'outputs/base99_radio_sound_v001/backup_20261009',
        'backup_files': [
            f'outputs/base99_radio_sound_v001/backup_20261009/{LEDGER.name}',
            f'outputs/base99_radio_sound_v001/backup_20261009/{MASTER.name}',
            f'outputs/base99_radio_sound_v001/backup_20261009/{BASELINE.name}',
        ],
    }
    (OUT / 'ledger_evidence.json').write_text(json.dumps(result, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    (OUT / 'logs/ledger_update.log').write_text(json.dumps(result, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    print(json.dumps(result, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    main()
