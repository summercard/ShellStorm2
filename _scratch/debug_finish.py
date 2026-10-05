from pathlib import Path
import json, traceback
root=Path('I:/工作项目/shellstrom2/ShellStorm2'); p=root/'_scratch/debug_finish_result.txt'
try:
 import openpyxl
 vals=[]
 w=openpyxl.load_workbook(root/'assets/registry/ledgers/ShellStorm2_道具账本_v001.xlsx',data_only=False); s=w['资产主表']; vals += ['found='+repr([r for r in range(6,s.max_row+1) if s.cell(r,1).value=='PRP-BASE99-RADIO-3D']),'max='+str(s.max_row),'dv='+repr([(d.type,[str(x) for x in d.sqref.ranges]) for d in s.data_validations.dataValidation]),'formulas='+repr([(r,s.cell(r,18).value,s.cell(r,19).value) for r in [6,26,27]),'prefab_max='+str(w['3D-道具'].max_row)+' prefab_id='+repr(w['3D-道具'].cell(8,1).value),'log_max='+str(w['域变更日志'].max_row)+' log_id='+repr(w['域变更日志'].cell(11,1).value)]
 m=openpyxl.load_workbook(root/'assets/registry/ShellStorm2_美术资产台账_v001.xlsx',data_only=False); vals += ['master_D18='+repr(m['总览']['D18'].value),'control='+repr([m['3D Prefab总控'][x].value for x in ['C7','E7','F7','G7']])]
 b=json.loads((root/'assets/registry/ledger_split_baseline.json').read_text(encoding='utf-8')); vals += ['base_asset='+repr(b['assets'].get('PRP-BASE99-RADIO-3D')),'base_count='+str(b['asset_count']),'evidence='+str((root/'outputs/base99_radio_v001/registration_evidence.json').exists())]
 p.write_text('\n'.join(vals),encoding='utf-8')
except Exception: p.write_text(traceback.format_exc(),encoding='utf-8'); raise
