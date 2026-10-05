from pathlib import Path
import json
import openpyxl
root=Path('I:/工作项目/shellstrom2/ShellStorm2')
path=root/'assets/registry/ShellStorm2_美术资产台账_v001.xlsx'
wb=openpyxl.load_workbook(path,data_only=False)
out={}
for name in ['分类与编码','命名与查重','3D Prefab总控']:
    ws=wb[name]
    rows=[]
    for r in range(1,ws.max_row+1):
        vals=[ws.cell(r,c).value for c in range(1,ws.max_column+1)]
        text=' | '.join('' if v is None else str(v) for v in vals)
        if name=='3D Prefab总控' or any(k in text for k in ['PRP','场景道具','道具','decor_prop','3D-道具','Prefab']):
            rows.append({'row':r,'values':vals})
    out[name]={'max_row':ws.max_row,'max_column':ws.max_column,'rows':rows}
idx=json.loads((root/'assets/registry/ledger_index.json').read_text(encoding='utf-8'))
out['ledger_index']={'scenes':idx['domains'][2],'props':idx['domains'][3]}
(root/'_scratch/radio_rules.json').write_text(json.dumps(out,ensure_ascii=False,indent=2,default=str)+'\n',encoding='utf-8')
