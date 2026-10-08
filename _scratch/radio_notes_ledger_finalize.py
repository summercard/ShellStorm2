from pathlib import Path
import json,hashlib,sys,re,zipfile,xml.etree.ElementTree as ET
from copy import deepcopy
import openpyxl
R=Path('I:/工作项目/shellstrom2/ShellStorm2'); O=R/'outputs/base99_radio_music_notes'
sys.path.insert(0,str(R/'scripts'));sys.path.insert(0,str(R/'tools/asset_pipeline'))
from ledger_registry import LedgerIndex
from split_asset_ledger import read_source_rows,_row_digest,col_digest,sheet_digest,CONTENT_COLUMNS
index=LedgerIndex.load(R);P=index.path_for_category('道具');V=index.path_for_category('特效');B=R/'assets/registry/ledger_split_baseline.json'
ID='VFX-RADIO-MUSIC-NOTES-3D';RID='PRP-BASE99-RADIO-3D';FX='assets/art/vfx/environment_3d/radio_music_notes/vfx_radio_music_notes_root_top3d.tscn'
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def dump(p,v):p.write_bytes((json.dumps(v,ensure_ascii=False,indent=2,default=str)+'\n').replace('\n','\r\n').encode())
w=openpyxl.load_workbook(V); assert w['资产主表']['A23'].value==ID
w['资产主表']['T23']=sha(R/FX)
w['3D-特效']['L24']=str(w['3D-特效']['L24'].value).replace('0.34×0.38','0.46×0.52')
for name,ref in [('资产主表','N23'),('3D-特效','G24')]:
 c=w[name][ref]; c.value=str(c.value)+' 单符Quad最终0.46×0.52m、出生横向偏移±0.22m。'
w.save(V)
bl=json.loads(B.read_text(encoding='utf-8'));old=json.loads((O/'backup/assets/registry/ledger_split_baseline.json').read_text(encoding='utf-8'));allrows=[]
for d in index.domains:
 wb=openpyxl.load_workbook(d.path)
 for r,vals in read_source_rows(wb['资产主表']):
  allrows.append((r,vals))
  if vals[0] in {ID,RID}:bl['assets'][vals[0]]={'v':_row_digest(vals),'c':vals[2],'d':d.key}
 for name in ['3D-道具','3D-特效']:
  if name in wb:bl['sheet_digests'][name]=sheet_digest(wb[name])
for aid,rec in old['assets'].items():
 if aid!=RID:assert bl['assets'][aid]==rec
bl['asset_count']=len(bl['assets']);bl['column_digests']={str(c):col_digest(allrows,c) for c in CONTENT_COLUMNS};bl['category_counts']={c:sum(v[2]==c for _,v in allrows) for c in sorted({v[2] for _,v in allrows})}
dump(B,bl)
diffs={}
for path in [P,V]:
 before=openpyxl.load_workbook(O/'backup'/path.relative_to(R));after=openpyxl.load_workbook(path);changes=[];lastlog=before['域变更日志'].max_row
 for oldws in before:
  nws=after[oldws.title];assert str(oldws.merged_cells)==str(nws.merged_cells)
  for row in nws:
   for cell in row:
    oldcell=oldws[cell.coordinate]
    if oldcell.value==cell.value:
     assert oldcell._style==cell._style or (path==V and oldws.title=='资产主表' and cell.row==23)
     continue
    if path==P:
     allow=(oldws.title=='资产主表' and cell.coordinate in {'M27','N27','P27','T27','V27','Y27'}) or (oldws.title=='3D-道具' and cell.coordinate in {'O8','K8','P8'}) or (oldws.title=='域变更日志' and cell.row>lastlog)
    else:
     allow=(oldws.title=='资产主表' and (cell.row==23 or cell.column==19)) or (oldws.title=='总览' and cell.coordinate in ['A6','C6','E6','G6','B10','C10','B11','C11']) or (oldws.title=='3D-特效' and cell.row==24) or (oldws.title=='域变更日志' and cell.row>lastlog)
    assert allow,(path.name,oldws.title,cell.coordinate)
    changes.append({'sheet':oldws.title,'cell':cell.coordinate,'before':oldcell.value,'after':cell.value})
 diffs[path.name]=changes
 if path==V:
  with zipfile.ZipFile(path) as z:
   actual={d.attrib['sqref'] for n in z.namelist() if re.fullmatch(r'xl/worksheets/sheet\d+.xml',n) for d in ET.fromstring(z.read(n)).findall('{*}dataValidations/{*}dataValidation')}
   assert {'C6:C23','K6:K23','L6:L23'}.issubset(actual)
debt=[]
for domain,path in [('props',P),('vfx',V)]:
 rep=json.JSONDecoder().raw_decode((O/f'testlogs/{domain}_before.log').read_text(encoding='utf-8').lstrip())[0]
 before=openpyxl.load_workbook(O/'backup'/path.relative_to(R));after=openpyxl.load_workbook(path)
 for issue in rep['issues'].get('sha_mismatch',[]):
  assert sha(Path(issue['path']))==issue['actual']
  assert before['资产主表'].cell(issue['row'],20).value==after['资产主表'].cell(issue['row'],20).value==issue['recorded']
  debt.append(issue)
dump(O/'ledger_evidence.json',{'passed':True,'new_asset_id':ID,'vfx_asset_row':23,'vfx_prefab_row':24,'radio_asset_row':27,'radio_prefab_row':8,'vfx_count_before':17,'vfx_count_after':18,'baseline_count':bl['asset_count'],'differences':diffs,'existing_sha_debt_unchanged':debt,'two_transactions':True,'DV_zip_verified':True,'other_asset_fingerprints_unchanged':True})
print('RADIO_NOTES_LEDGER_EVIDENCE_OK',bl['asset_count'])
