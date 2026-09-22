# -*- coding: utf-8 -*-
"""把 ENV-TOWER-FLOOR-TILE-5M（资产主表第 54 行）从 v002 升到 v003。

只改单元格文字，不增删行列（该页带 AssetMasterTable 表格对象）。
改前自动备份 .bak_floor_tile_v003。
"""
import shutil
import sys
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(r"I:\工作项目\shellstrom2\ShellStorm2")
LEDGER = ROOT / "assets/registry/ledgers/ShellStorm2_场景账本_v001.xlsx"
BACKUP = LEDGER.with_name(LEDGER.name + ".bak_floor_tile_v003")

ASSET_ID = "ENV-TOWER-FLOOR-TILE-5M"
SHEET = "资产主表"

NEW_O = ("assets/art/environments/tower_descent_3d/runtime/floor_tile_5m/"
         "env_tower_floor_tile_5m_root_top3d.tscn")
NEW_P = "; ".join([
    "assets/art/environments/tower_descent_3d/source/floor_tile_5m/env_tower_floor_tile_5m_source_v003.blend",
    "assets/art/environments/tower_descent_3d/components/floor_tile_5m/env_tower_floor_tile_5m_top3d.glb",
    "scripts/blender/build_tower_floor_tile_5m_v003.py",
    "assets/art/props/dungeon_3d/prp_tower_floor_tile_5m.tscn",
    "tests/verification/verify_tower_journey_polish.tscn",
])
NEW_N = ("5×5×0.304m；740三角面；2材质；共享MultiMesh；行走面锚点+0.1500不变；"
         "装饰件反共面阶梯+0.1520/+0.1540（XY重叠可见面高度差≥2mm）")
NEW_Y_APPEND = (
    "2026-09-22：修复细节级闪面（z-fighting）升 v003。v002 顶面四家族"
    "（哑光面板/角部紧固座/紧固槽/检修盖板）同落 +0.1500 且 XY 互叠（跨零件共面 13 对、"
    "0.3745㎡，最严重一处为检修盖板整片 0.324㎡），相机贴地观察时逐像素闪烁；"
    "另一并存缺陷是检修盖框顶面 +0.1480，比面板低 2mm 被永久埋没不可见。"
    "v003 采用反共面阶梯：行走面锚点 +0.1500 不动（TowerFloorStage3D 固定 "
    "-FLOOR_THICKNESS*0.5=-0.15 偏移契约不受影响），盖框/紧固座/警示短标抬到 +0.1520，"
    "紧固槽/盖板格栅条抬到 +0.1540，并将检修盖缝改为 2mm 真实凹槽（通透负空间）。"
    "本件与 prp_tower_floor_tile_5m.tscn 共用同一个已去版本化的 GLB，"
    "故 98F 标准砖与远征/98F 抛光砖同时受益；99F 自持地砖不受影响。"
)


def main() -> int:
    if not LEDGER.is_file():
        print("LEDGER_MISSING %s" % LEDGER)
        return 2
    shutil.copy2(LEDGER, BACKUP)
    print("BACKUP_OK %s (%d bytes)" % (BACKUP.name, BACKUP.stat().st_size))

    wb = load_workbook(LEDGER, data_only=False)
    ws = wb[SHEET]
    target_row = None
    for row in range(6, ws.max_row + 1):
        if str(ws.cell(row, 1).value or "").strip() == ASSET_ID:
            target_row = row
            break
    if target_row is None:
        print("ROW_NOT_FOUND %s" % ASSET_ID)
        return 3

    before = {
        "M": ws.cell(target_row, 13).value,
        "N": ws.cell(target_row, 14).value,
        "O": ws.cell(target_row, 15).value,
        "P": ws.cell(target_row, 16).value,
        "V": ws.cell(target_row, 22).value,
        "Y": ws.cell(target_row, 25).value,
    }
    for key, value in before.items():
        print("BEFORE %s = %s" % (key, value))

    old_y = str(before["Y"] or "").rstrip()
    new_y = (old_y + "\n" + NEW_Y_APPEND) if old_y else NEW_Y_APPEND

    ws.cell(target_row, 13).value = "v003"
    ws.cell(target_row, 14).value = NEW_N
    ws.cell(target_row, 15).value = NEW_O
    ws.cell(target_row, 16).value = NEW_P
    ws.cell(target_row, 22).value = "2026-09-22 00:00:00"
    ws.cell(target_row, 25).value = new_y

    wb.save(LEDGER)
    print("SAVED row=%d" % target_row)

    # 回读校验
    wb2 = load_workbook(LEDGER, data_only=False)
    ws2 = wb2[SHEET]
    assert str(ws2.cell(target_row, 1).value).strip() == ASSET_ID, "row shifted!"
    for col, name in ((13, "M"), (14, "N"), (15, "O"), (16, "P"), (22, "V")):
        print("AFTER %s = %s" % (name, ws2.cell(target_row, col).value))
    print("AFTER Y = %s" % ws2.cell(target_row, 25).value)
    print("TABLE_OBJECTS %s" % list(ws2.tables.keys()))
    print("ROW54_ASSET_ID %s" % ws2.cell(54, 1).value)
    print("LEDGER_V003_ROW_OK")
    return 0


if __name__ == "__main__":
    sys.exit(main())
