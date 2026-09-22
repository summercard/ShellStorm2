"""Read-only ledger health check before/after the expedition01 registration."""
from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(r"I:\工作项目\shellstrom2\ShellStorm2")
LEDGER = ROOT / "assets/registry/ledgers/ShellStorm2_场景账本_v001.xlsx"

wb = load_workbook(LEDGER)
print("sheets:", wb.sheetnames)

main = wb["资产主表"]
print("\n[资产主表] max_row =", main.max_row, "max_col =", main.max_column)
print("header row 5:", [main.cell(row=5, column=c).value for c in range(1, main.max_column + 1)])
for r in range(max(6, main.max_row - 4), main.max_row + 1):
    print("  row", r, "|", main.cell(row=r, column=1).value, "|", main.cell(row=r, column=11).value)

prefab = wb["3D-场景通用"]
print("\n[3D-场景通用] max_row =", prefab.max_row, "max_col =", prefab.max_column)
print("header row 5:", [prefab.cell(row=5, column=c).value for c in range(1, prefab.max_column + 1)])
for r in range(max(6, prefab.max_row - 3), prefab.max_row + 1):
    print("  row", r, "|", prefab.cell(row=r, column=1).value, "|", prefab.cell(row=r, column=14).value)

cl = wb["域变更日志"]
print("\n[域变更日志] max_row =", cl.max_row, "max_col =", cl.max_column)
for r in range(max(2, cl.max_row - 3), cl.max_row + 1):
    print("  row", r, "|", cl.cell(row=r, column=1).value, "|", cl.cell(row=r, column=3).value)

ov = wb["总览"]
print("\n[总览] max_row =", ov.max_row, "max_col =", ov.max_column)
hits = []
for row in ov.iter_rows():
    for cell in row:
        v = cell.value
        if isinstance(v, str) and ("资产主表" in v or "3D-场景通用" in v):
            hits.append((cell.coordinate, v[:110]))
for c, v in hits:
    print("  ", c, "|", v)

# duplicate scan
ids = {}
for r in range(6, main.max_row + 1):
    v = main.cell(row=r, column=1).value
    if v:
        ids.setdefault(v, []).append(r)
dups = {k: v for k, v in ids.items() if len(v) > 1}
print("\n[资产主表] duplicate AssetIDs:", dups if dups else "none")
print("[资产主表] total id rows:", len(ids))
