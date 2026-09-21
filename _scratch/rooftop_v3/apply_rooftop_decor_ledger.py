#!/usr/bin/env python3
"""把 100F 天台装饰接入的事实写回场景账本。

只编辑 assets/registry/ledgers/ShellStorm2_场景账本_v001.xlsx：
  1. 资产主表 库行（ENV-ROOFTOP-REFERENCE-COMPONENT-LIBRARY）：44包->47包、方案A、SHA、时间。
  2. 3D-场景通用 专页：11 类装饰组件由「Blender源已完成」升级为已导出+已接入。
  3. 3D-场景通用 专页：新增 100F 天台装饰布局 登记行。
  4. 域变更日志：追加 v0.1.4。

AssetID 全部沿用既有行；只有「布局」是全新条目，不与任何既有 AssetID 冲突。
"""
from __future__ import annotations

import copy
import hashlib
import json
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
LEDGER = ROOT / "assets" / "registry" / "ledgers" / "ShellStorm2_场景账本_v001.xlsx"
CATALOG = (
    ROOT
    / "assets/art/environments/tower_zones/rooftop/source/reference_components/v002"
    / "component_packages_v002/catalog.json"
)
LIB_BLEND = "assets/art/environments/tower_zones/rooftop/source/reference_components/v002/天台区块_参考组件库_v002.blend"
LAYOUT_BLEND = "assets/art/environments/tower_zones/rooftop/source/layouts/100f_decorated_v001/rooftop_100f_decorated_layout_v001.blend"
LAYOUT_JSON = "assets/art/environments/tower_zones/rooftop/source/layouts/100f_decorated_v001/rooftop_100f_decorated_layout_v001.json"
RUNTIME_MANIFEST = "assets/art/environments/tower_zones/rooftop/runtime/rooftop_100f_decorated_runtime_manifest.json"
GLB_DIR = "assets/art/environments/tower_zones/rooftop/components"
TSCN_DIR = "assets/art/environments/tower_zones/rooftop/runtime"
STAGE_SCRIPT = "src/world3d/TowerFloorStage3D.gd"
EXPECTED_END = 239  # 资产主表 查重范围上界（本账本行数不变）
TODAY_SERIAL = 46286  # 2026-09-21（Excel 日期序列）

# 11 类装饰组件：slug -> (账本行号, 中文名, 布局分组说明, 实例数)
DECOR = {
    "hvac_small": (113, "小型空调机组", "房屋四向外墙墙面（避开门洞）", 4),
    "hvac_vent": (115, "小通风口", "房屋南北墙高处", 2),
    "pipe_straight": (118, "直管段", "房屋四周闭环水管", 32),
    "pipe_elbow": (119, "转角弯管", "闭环水管四角", 4),
    "pipe_tee": (120, "三通管", "闭环水管灌溉分支", 2),
    "pipe_riser": (121, "立管下水管", "闭环水管至地面立管", 4),
    "pipe_bracket": (122, "管道支架", "立管墙面固定", 4),
    "ivy": (125, "墙面攀爬藤蔓", "房屋墙面分段挂藤", 11),
    "flowerbox": (126, "长条花箱", "房屋四周贴墙长条花箱", 6),
    "plant_large": (127, "大盆栽", "房屋周边高点盆栽", 6),
    "plant_small": (128, "小盆栽", "房屋周边低点盆栽", 6),
}


def sha12(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()[:12]


def main() -> int:
    catalog = {row["slug"]: row for row in json.loads(CATALOG.read_text(encoding="utf-8"))}
    wb = load_workbook(LEDGER)
    main_sheet = wb["资产主表"]
    page = wb["3D-场景通用"]
    log = wb["域变更日志"]

    # ---------- 1. 资产主表：库行升级 ----------
    exported = [slug for slug, row in catalog.items() if row.get("exported")]
    integrated = [slug for slug, row in catalog.items() if row.get("runtime_integrated")]
    lib_sha = hashlib.sha256((ROOT / LIB_BLEND).read_bytes()).hexdigest()
    row = main_sheet[238]
    assert row[0].value == "ENV-ROOFTOP-REFERENCE-COMPONENT-LIBRARY", row[0].value
    row[13].value = "47包 / 4共享材质；女儿墙方案A 0.80m；%d件已导出（含11件装饰接入）" % len(exported)
    row[9].value = "100F独立组件制作；装饰件已按布局接入运行场景"
    row[19].value = lib_sha
    row[21].value = TODAY_SERIAL
    row[24].value = (
        "2026-09-20 女儿墙 1.80m→0.80m（方案A平整墙板）写回主库，三件破损变种入 02_女儿墙；"
        "2026-09-21 组件库升至 47 包；11 类装饰（空调/通风口/水管/藤蔓/花圃盆栽）导出 GLB 并接入 100F 装饰布局。"
    )
    print("MAIN_ROW_238_OK packages=47 exported=%d integrated=%d sha=%s..." % (len(exported), len(integrated), lib_sha[:12]))

    # ---------- 2. 专页：11 类装饰组件升级 ----------
    updated = []
    for slug, (line, name, placement, count) in DECOR.items():
        pkg = catalog[slug]
        cell_row = page[line]
        assert cell_row[0].value == pkg["package_id"], (line, cell_row[0].value, pkg["package_id"])
        glb_rel = "%s/env_rooftop_ref_%s_top3d.glb" % (GLB_DIR, slug)
        tscn_rel = "%s/%s.tscn" % (TSCN_DIR, slug)
        assert (ROOT / glb_rel).is_file(), glb_rel
        assert (ROOT / tscn_rel).is_file(), tscn_rel
        cell_row[2].value = tscn_rel
        cell_row[3].value = glb_rel
        cell_row[4].value = LIB_BLEND
        cell_row[5].value = (
            "Godot视觉装饰PackedScene；100F天台%s（布局内 %d 个实例），由 TowerFloorStage3D 按装饰布局清单实例化"
            % (placement, count)
        )
        cell_row[6].value = STAGE_SCRIPT
        cell_row[7].value = "无"
        cell_row[8].value = "引擎（visual_only）"
        cell_row[9].value = "未创建"
        cell_row[12].value = "100F天台装饰布局 / %s" % slug
        cell_row[13].value = "正式美术已接入"
        cell_row[15].value = (
            "父库=ENV-ROOFTOP-REFERENCE-COMPONENT-LIBRARY；2026-09-21 从 v002 主库导出 GLB(%s) 并生成 runtime PackedScene；"
            "接入 100F 装饰布局＝113实例/7组/0启用碰撞；布局源=%s" % (sha12(ROOT / glb_rel), LAYOUT_JSON)
        )
        updated.append(cell_row[0].value)
    print("PAGE_DECOR_ROWS_OK count=%d ids=%s" % (len(updated), ",".join(updated)))

    # ---------- 3. 专页：库行说明补方案A ----------
    lib_row = page[96]
    assert lib_row[0].value == "ENV-ROOFTOP-REFERENCE-COMPONENT-LIBRARY", lib_row[0].value
    lib_row[5].value = (
        "47独立包（10类）；女儿墙已改方案A 0.80m 并含3件破损变种；"
        "11 类装饰（空调/通风口/水管/藤蔓/花圃盆栽）已导出并接入 100F 装饰布局。"
    )
    lib_row[15].value = (
        "2026-09-20 方案A 0.80m 写回主库；2026-09-21 库内 47 包，%d 件已导出，其中 11 件装饰已接入运行时；"
        "其余包仍为 Blender 源状态（未导出 GLB）。" % len(exported)
    )
    print("PAGE_LIBRARY_ROW_OK")

    # ---------- 4. 专页：新增装饰布局登记行 ----------
    new_row = page.max_row + 1
    assert page.cell(row=new_row, column=1).value is None, "目标行非空: %d" % new_row
    style_source = page[page.max_row]
    layout_values = [
        "ENV-ROOFTOP-DECOR-LAYOUT-100F",
        "100F天台装饰布局",
        "",
        "",
        LAYOUT_BLEND,
        "Blender组件实例布局源；只摆放共享组件实例、不拥有几何；113个装饰实例＝房屋墙体16+屋顶16+空调6+绿化18+藤蔓11+水管38+立管支架8",
        STAGE_SCRIPT,
        "未创建",
        "引擎（visual_only）",
        "未创建",
        "90×80m天台；外围68件（64直段+4外角）；234块5m地砖（开口=中庭6×6+西侧楼梯口3×6）",
        "世界原点对齐 ROOFTOP_WORLD_RECT；Blender Z-up / -Y；Godot=(Blender X, Blender Z, -Blender Y)",
        "100F天台 / TowerFloorStage3D 运行时按布局清单重放",
        "正式美术已接入",
        "v001",
        "清单=%s；运行时登记=%s；验收=verify_rooftop_32x32_contract + probe_rooftop_decorated_layout + probe_rooftop_decorated_stage_only；"
        "全部实例缩放=1、启用碰撞=0；结构壳体（地砖/女儿墙）仍由 TowerFloorStage3D 负责，布局不重复生成。"
        % (LAYOUT_JSON, RUNTIME_MANIFEST),
    ]
    for col, value in enumerate(layout_values, start=1):
        cell = page.cell(row=new_row, column=col, value=value)
        source = style_source[col - 1]
        cell._style = copy.copy(source._style)
    print("PAGE_LAYOUT_ROW_OK row=%d id=%s" % (new_row, layout_values[0]))

    # ---------- 5. 域变更日志 v0.1.4 ----------
    last = log.max_row
    assert log.cell(row=last, column=1).value == "v0.1.3", log.cell(row=last, column=1).value
    log_row = last + 1
    log_values = [
        "v0.1.4",
        "2026-09-21",
        "资产升版+新条目",
        "关卡场景 / 100F天台",
        "天台参考组件库升为 47 包（女儿墙方案A 0.80m + 3件破损变种）；"
        "11 类装饰组件（HVAC-SMALL/VENT、PIPE-STRAIGHT/ELBOW/TEE/RISER/BRACKET、IVY、FLOWERBOX、PLANT-LARGE/SMALL）"
        "由「Blender源已完成」升级为「正式美术已接入」：填 runtime PackedScene 与 GLB 路径、碰撞归零；"
        "新增 ENV-ROOFTOP-DECOR-LAYOUT-100F 登记 100F 天台装饰布局（113 实例 / 7 组）。",
        "AssetID 全部沿用既有行，无新增组件 ID；仅「装饰布局」为新条目。装饰件无碰撞、无新玩法语义，"
        "原有地砖/女儿墙/楼梯结构与碰撞不变；旧账本留档 " + LEDGER.name + ".bak_rooftop_decor_ledger。",
        "摩斯拉",
    ]
    for col, value in enumerate(log_values, start=1):
        cell = log.cell(row=log_row, column=col, value=value)
        source = log.cell(row=last, column=col)
        cell._style = copy.copy(source._style)
    print("LOG_ROW_OK row=%d version=%s" % (log_row, log_values[0]))

    wb.save(LEDGER)
    print("LEDGER_SAVED %s bytes=%d" % (LEDGER.name, LEDGER.stat().st_size))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
