#!/usr/bin/env python3
"""第二阶段：把「100F 天台装饰布局」补成《资产主表》正式条目。

README 规定「资产行只落《资产主表》，3D-* 分页只是视图」，而 100F 天台装饰布局
目前只在专页登记。本脚本按上一批新增门扇行的同一先例，把它补进主表，并把随之
必须扩容的校验/条件格式/筛选/总览公式从 $239 扩到 $240。
"""
from __future__ import annotations

import copy
import hashlib
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
LEDGER = ROOT / "assets" / "registry" / "ledgers" / "ShellStorm2_场景账本_v001.xlsx"
LAYOUT_BLEND = "assets/art/environments/tower_zones/rooftop/source/layouts/100f_decorated_v001/rooftop_100f_decorated_layout_v001.blend"
LAYOUT_JSON = "assets/art/environments/tower_zones/rooftop/source/layouts/100f_decorated_v001/rooftop_100f_decorated_layout_v001.json"
OLD_END = 239
NEW_END = 240
TODAY_SERIAL = 46286


def main() -> int:
    wb = load_workbook(LEDGER)
    sheet = wb["资产主表"]
    assert sheet.max_row == OLD_END, sheet.max_row
    assert sheet.cell(row=NEW_END, column=1).value is None, "R240 非空，先人工确认"

    layout_sha = hashlib.sha256((ROOT / LAYOUT_BLEND).read_bytes()).hexdigest()

    row = NEW_END
    values = {
        1: "ENV-ROOFTOP-DECOR-LAYOUT-100F",
        2: "100F天台装饰布局",
        3: "场景",
        4: "environment_layout_3d",
        5: "rooftop_100f_decorated_layout",
        6: "root_3d",
        7: None,
        8: "Blender Z-up / -Y",
        9: "static / visual_only",
        10: "100F天台专用；只摆共享组件实例，不拥有几何",
        11: "正式美术已接入",
        12: "P1",
        13: "v001",
        14: "113装饰实例/7组；外围68件（64直段+4外角）；234块5m地砖；0启用碰撞",
        15: LAYOUT_BLEND,
        16: LAYOUT_JSON + "; src/world3d/TowerFloorStage3D.gd",
        17: "天台;装饰布局;空调;通风口;水管;藤蔓;花圃盆栽;100F",
        18: '=LOWER(TRIM(C%d)&"|"&TRIM(D%d)&"|"&TRIM(E%d)&"|"&TRIM(F%d)&"|"&TRIM(H%d)&"|"&TRIM(I%d))' % ((row,) * 6),
        19: '=IF(COUNTIF($R$6:$R$%d,R%d)>1,"重复","唯一")' % (NEW_END, row),
        20: layout_sha,
        21: "摩斯拉",
        22: TODAY_SERIAL,
        23: "用户要求「用组件库现有件装饰天台」；Blender 组件实例布局",
        24: "100F_decorated_v001",
        25: (
            "运行时按布局清单重放：装饰组=房屋墙体16+屋顶16+空调与通风口6+绿化18+藤蔓11+水管38+立管支架8＝113实例；"
            "Godot 坐标=(Blender X, Blender Z, -Blender Y)；全部实例缩放=1且启用碰撞=0；"
            "结构壳体（234地砖/68件女儿墙/楼梯与碰撞）仍由 TowerFloorStage3D 负责，本布局不重复生成。"
            "运行时登记=assets/art/environments/tower_zones/rooftop/runtime/rooftop_100f_decorated_runtime_manifest.json；"
            "验收=verify_rooftop_32x32_contract + probe_rooftop_decorated_layout + probe_rooftop_decorated_stage_only。"
        ),
    }
    style_source = sheet[OLD_END]
    for col in range(1, 26):
        cell = sheet.cell(row=row, column=col, value=values[col])
        cell._style = copy.copy(style_source[col - 1]._style)

    # ---- 派生公式范围 239 -> 240 ----
    rewritten = 0
    for r in range(6, NEW_END + 1):
        cell = sheet.cell(row=r, column=19)
        text = cell.value
        if isinstance(text, str) and "$%d" % OLD_END in text:
            cell.value = text.replace("$%d" % OLD_END, "$%d" % NEW_END)
            rewritten += 1
    print("DEDUPE_RANGE_REWRITTEN rows=%d" % rewritten)
    # 旧区间 6..239 的 234 行都该被改写；新行自身的公式本来就是 $240，不计入
    assert rewritten == OLD_END - 6 + 1, rewritten

    # ---- 数据校验 / 条件格式 / 自动筛选 ----
    for dv in sheet.data_validations.dataValidation:
        ref = str(dv.sqref)
        if str(OLD_END) in ref:
            dv.sqref = ref.replace(str(OLD_END), str(NEW_END))
    cf_added = 0
    for rng in list(sheet.conditional_formatting):
        ref = str(rng.sqref)
        if str(OLD_END) in ref:
            for rule in rng.rules:
                sheet.conditional_formatting.add(ref.replace(str(OLD_END), str(NEW_END)), copy.copy(rule))
                cf_added += 1
    print("CF_RULES_EXTENDED count=%d" % cf_added)
    if sheet.auto_filter.ref:
        sheet.auto_filter.ref = sheet.auto_filter.ref.replace(str(OLD_END), str(NEW_END))

    # ---- 其它页公式里的 $239 ----
    overview_fixed = 0
    for ws in wb.worksheets:
        for row_cells in ws.iter_rows():
            for cell in row_cells:
                if isinstance(cell.value, str) and "$%d" % OLD_END in cell.value:
                    cell.value = cell.value.replace("$%d" % OLD_END, "$%d" % NEW_END)
                    overview_fixed += 1
    print("SHEET_FORMULAS_FIXED cells=%d" % overview_fixed)

    wb.save(LEDGER)
    print("MAIN_ASSET_ROW_ADDED row=%d id=%s sha=%s..." % (row, values[1], layout_sha[:12]))
    print("LEDGER_SAVED bytes=%d" % LEDGER.stat().st_size)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
