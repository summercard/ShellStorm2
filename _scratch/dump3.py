import openpyxl
p = "assets/registry/ledgers/ShellStorm2_场景账本_v001.xlsx"
wb = openpyxl.load_workbook(p, data_only=False)
ws = wb["资产主表"]
hdr = [ws.cell(5,c).value for c in range(1, ws.max_column+1)]
def show(r):
    print("---------- ROW %d ----------" % r)
    for c in range(1, ws.max_column+1):
        v = ws.cell(r,c).value
        if v is not None:
            print("   %-28s = %s" % (hdr[c-1], v))
for r in (56, 57, 70, 101, 238):
    show(r)
