import sys,json,hashlib
from pathlib import Path
from openpyxl import load_workbook
R=Path.cwd();sys.path.insert(0,str(R/'scripts'));from ledger_registry import LedgerIndex
ix=LedgerIndex.load(R);p=ix.path_for_category('敌人');w=load_workbook(p);s=w['资产主表'];row=next(r for r in range(1,s.max_row+1) if s.cell(r,1).value=='ENM-BOSS-MONITOR002-3D')
master=load_workbook(R/'assets/registry/ShellStorm2_美术资产台账_v001.xlsx',read_only=True)
for name in ['分类与编码','命名与查重']:
 print(name,[[str(v)[:100] for v in row if v is not None] for row in master[name].iter_rows(values_only=True) if any(v and ('敌人' in str(v) or 'ENM' in str(v)) for v in row)])
B=R/'assets/art/enemies/bosses/enm_boss_monitor002';src=B/'source/enm_boss_monitor002_model_v026.blend';note='v026 补齐剩余五类十剪辑：special四段、hurt、stun三段、turn_left/right。肩侧带动双手，前后错拍与腕部回弹；插地锁点及坐地接触专项检查。15个正式动作源级制作，未导出或接入Godot。'
changes={f'I{row}':'T Pose / 六表情 / 15正式剪辑：idle move melee_keyboard melee_cable heavy_spin_slam special四段 hurt stun三段 turn_left/right',f'M{row}':'v026',f'O{row}':src.relative_to(R).as_posix(),f'T{row}':hashlib.sha256(src.read_bytes()).hexdigest(),f'P{row}':'source/enm_boss_monitor002_model_v026.blend + enm_boss_monitor002_animation_v026.blend；SKEL-MONITOR002-005',f'V{row}':'2026-10-05',f'Y{row}':note}
logrow=w['域变更日志'].max_row+1
jobs=[{'path':str(p),'sheets':{'资产主表':changes,'域变更日志':{f'{col}{logrow}':value for col,value in zip('ABCDE',['v026','2026-10-05','Codex','ENM-BOSS-MONITOR002-3D',note])}}}]
p=B/'source/boss002_production_ledger.xlsx';w=load_workbook(p);s=w.active;changes={'B4':'v026 数据线近战弧形甩击'}
for r in range(1,s.max_row+1):
 label=s.cell(r,1).value
 if label=='当前源版本':changes.update({f'B{r}':'v026',f'C{r}':'model_v026 + animation_v026；v012/v013保留回退'})
 if label=='绑定验证':changes.update({f'B{r}':'previews/remaining_v026/audit.json',f'C{r}':note})
 if label=='面数统计':changes[f'C{r}']='source/rig_contract_v026.json'
jobs.append({'path':str(p),'sheets':{s.title:changes}})
Path('_scratch/boss026_ledger/changes.json').write_text(json.dumps(jobs,ensure_ascii=False),encoding='utf-8')
