# -*- coding: utf-8 -*-
"""把 L 型走廊登记进场景账本，并同步无损基线。

登记内容
  《资产主表》      第 242 行  ENV-EXPEDITION-L01-L-CORRIDOR
  《3D-场景通用》   第 148 行  Prefab 分页对齐行
  《总览》          A6/C6/E6/G6 与 B10..C13 的 资产主表 区间 $241 -> $242
  DataValidation    C/K/L 三列 sqref C6:C240 -> C6:C242（顺带修正上一行未覆盖的遗留）
  《域变更日志》    第 19 行    v0.1.13
  基线               assets/registry/ledger_split_baseline.json

派生列 R/S 由 split_asset_ledger 的唯一真源公式生成，不写字面量。
"""

from __future__ import annotations

import copy
import json
import shutil
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
sys.path.insert(0, str(ROOT / "tools" / "asset_pipeline"))

from openpyxl import load_workbook  # noqa: E402
from ledger_registry import LedgerIndex  # noqa: E402
from split_asset_ledger import (  # noqa: E402
    ASSET_SHEET,
    CONTENT_COLUMNS,
    FIRST_DATA_ROW,
    _row_digest,
    _text,
    col_digest,
    dedupe_key_formula,
    dedupe_result_formula,
    read_source_rows,
    sheet_digest,
)

LEDGER = ROOT / "assets/registry/ledgers/ShellStorm2_场景账本_v001.xlsx"
BASELINE = ROOT / "assets/registry/ledger_split_baseline.json"

ASSET_ID = "ENV-EXPEDITION-L01-L-CORRIDOR"
NAME_ZH = "远征关卡01 L型走廊 数据连廊 45×40m 房间种类源"
BLEND_REL = (
    "assets/art/environments/tower_zones/expedition/source/room_types/"
    "l_corridor/v003/L型走廊种类_数据连廊_45x40m_v003.blend"
)
SHA = "9c323876f7165f656e016049135652a96c83ed0adb45f4f28c49fc50cfb5adae"
BUILD_PY = (
    "assets/art/environments/tower_zones/expedition/source/room_types/"
    "l_corridor/v003/widen_l_corridor_v003.py"
)

NOTE = (
    "由 battle 区块房间种类源迁入（原 ENV-BATTLE-L-CORRIDOR-TYPE，原路径 "
    "tower_zones/battle/source/room_types/l_corridor/；业主 2026-09-24 裁决归属远征关卡01，"
    "battle 侧不留副本）。远征01 版图与白盒目录均无走廊白模，尺寸与门位按参考图推定。"
    "v001/v002 为 10m 宽走廊（外包络 45×35m / 28 块地砖 / 88、111 个组件包），"
    "v003 拓宽为 15m（45×40m / 42 块地砖 / 121 个组件包 / 4 个门位：西1 北2 南1）。"
    "随迁统一内部结构：component_packages_vNNN→component_packages、facility→facilities、"
    "散放渲染图入 renders/。共享件引用由 ENV-BATTLE-COMMON-* 更正为 shared 库现有 "
    "ENV-SHARED-GENERIC-*。未导 GLB、未接 Godot。"
)

MAIN_ROW_VALUES = {
    1: ASSET_ID,
    2: NAME_ZH,
    3: "场景",
    4: "l_corridor",
    5: "expedition_01_l_corridor",
    6: "room_type",
    7: None,
    8: "Top3D / Blender Z-up",
    9: "scene_art / reference_derived",
    10: "远征关卡01 专用；连接房间之间的 L 型通廊房间种类源",
    11: "Blender源已完成",
    12: "P1",
    13: "v003",
    14: "45×40×11.9m；横臂 9×3 + 竖臂 3×5 个 5m 槽；走廊净宽 15m（三砖）；42 块地砖 / 121 个独立组件包 / 4 个门位",
    15: BLEND_REL,
    16: BUILD_PY,
    17: "远征关卡01;L型走廊;数据连廊;房间种类;45×40m;走廊;服务器机柜;运维通道",
    20: SHA,
    21: "摩斯拉",
    22: "2026-09-24",
    23: "用户设计（参考图推定尺寸；业主 2026-09-24 裁决归属远征关卡01）",
    24: "—",
    25: NOTE,
}

PREFAB_ROW_VALUES = {
    1: ASSET_ID,
    2: NAME_ZH,
    3: "未制作",
    4: "未制作",
    5: BLEND_REL,
    6: (
        "远征关卡01 L型走廊房间种类源；45×40m（横臂 9×3 + 竖臂 3×5 个 5m 槽）；"
        "走廊净宽 15m 即三块 5m 地砖并列；42 块地砖 / 121 个独立组件包；"
        "门位 4 个（西1、北2、南1），门墙槽留空由通用门墙组件补位。"
    ),
    7: "无",
    8: "无",
    9: "无",
    10: "无",
    11: "45×40×11.9m",
    12: "房间局部坐标：X/Y 居中；底面 Z=0；Godot -Z",
    13: "远征关卡01 / Expedition",
    14: "房间种类源已完成；QA PASS",
    15: "v003",
    16: (
        "由 battle 区块迁入（原 l_corridor_room_type_v003.blend）；无白模，尺寸按参考图推定；"
        "v001/v002（10m 宽）保留在同目录历史版本；未导 GLB、未接 Godot。"
    ),
}

CHANGELOG_ROW_VALUES = {
    1: "v0.1.13",
    2: "2026-09-24",
    3: "归属迁移",
    4: "关卡场景 / 远征关卡01",
    5: (
        "新增 ENV-EXPEDITION-L01-L-CORRIDOR（远征关卡01 L型走廊 数据连廊 45×40m 房间种类源）："
        "原为 battle 区块房间种类源（ENV-BATTLE-L-CORRIDOR-TYPE），业主 2026-09-24 裁决归属远征关卡01，"
        "整目录迁至 tower_zones/expedition/source/room_types/l_corridor/，battle 侧不留副本。"
        "随迁统一：component_packages_vNNN→component_packages、facility→facilities、"
        "渲染图入 renders/、源改名 L型走廊种类_数据连廊_<尺寸>m_vNNN.blend；"
        "共享件引用由 ENV-BATTLE-COMMON-* 更正为 shared 库现有 ENV-SHARED-GENERIC-*。"
        "同批登记《资产主表》第242行与《3D-场景通用》第148行。"
        "v001/v002 为 10m 宽（45×35m / 88、111 包），v003 为 15m 宽（45×40m / 121 包）。未导 GLB、未接 Godot。"
    ),
    6: (
        "新增条目，不改动既有 AssetID / 路径 / 哈希；battle 区块的 room_types 目录随迁后为空。"
        "远征01 版图与白盒目录仍无走廊白模，接入前需远征白盒补走廊席位。"
        "《资产主表》DV 此前落后一行（C6:C240 未覆盖第241行），本次一并扩到 C6:C242。"
    ),
    7: "摩斯拉",
}


def copy_style(src, dst):
    dst._style = copy.copy(src._style)


def main() -> int:
    # ---- 0. 备份 ----
    bak = LEDGER.with_name(LEDGER.name + ".bak_l_corridor_expedition")
    if not bak.exists():
        shutil.copy2(LEDGER, bak)
    bak_base = BASELINE.with_name(BASELINE.name + ".bak_before_l_corridor_expedition")
    if not bak_base.exists():
        shutil.copy2(BASELINE, bak_base)
    print("backup:", bak.name, "|", bak_base.name)

    wb = load_workbook(LEDGER)
    report = {}

    # ---- 1. 《资产主表》: 新行 + 重算派生列 ----
    ws = wb[ASSET_SHEET]
    template = ws.max_row              # 241
    new_row = template + 1             # 242
    for col, value in MAIN_ROW_VALUES.items():
        cell = ws.cell(row=new_row, column=col)
        cell.value = value
        copy_style(ws.cell(row=template, column=col), cell)
    ws.cell(row=new_row, column=18).value = dedupe_key_formula(new_row)
    ws.cell(row=new_row, column=19).value = dedupe_result_formula(new_row, new_row)
    copy_style(ws.cell(row=template, column=18), ws.cell(row=new_row, column=18))
    copy_style(ws.cell(row=template, column=19), ws.cell(row=new_row, column=19))
    if ws.row_dimensions[template].height:
        ws.row_dimensions[new_row].height = ws.row_dimensions[template].height

    rewritten = 0
    for row in range(FIRST_DATA_ROW, new_row + 1):
        want = dedupe_result_formula(row, new_row)
        if ws.cell(row=row, column=19).value != want:
            ws.cell(row=row, column=19).value = want
            rewritten += 1
    report["main_row"] = new_row
    report["dedupe_rows_rewritten"] = rewritten

    # ---- 2. 《总览》: 资产主表区间 241 -> 242 ----
    ov = wb["总览"]
    touched = []
    for row in ov.iter_rows():
        for cell in row:
            v = cell.value
            if isinstance(v, str) and "资产主表" in v and "$241" in v:
                cell.value = v.replace("$241", "$242")
                touched.append(cell.coordinate)
    report["overview_cells"] = touched

    # ---- 3. DataValidation 扩到新末行 ----
    dv_changed = []
    for dv in ws.data_validations.dataValidation:
        s = str(dv.sqref)
        if s.endswith(":C240") or s.endswith(":K240") or s.endswith(":L240"):
            dv.sqref = s.replace("240", str(new_row))
            dv_changed.append(f"{s} -> {dv.sqref}")
    report["data_validations"] = dv_changed

    # ---- 4. 《3D-场景通用》: Prefab 分页对齐行 ----
    p = wb["3D-场景通用"]
    p_template = p.max_row
    p_row = p_template + 1
    for col, value in PREFAB_ROW_VALUES.items():
        cell = p.cell(row=p_row, column=col)
        cell.value = value
        copy_style(p.cell(row=p_template, column=col), cell)
    report["prefab_row"] = p_row

    # ---- 5. 《域变更日志》 ----
    cl = wb["域变更日志"]
    cl_template = cl.max_row
    cl_row = cl_template + 1
    for col, value in CHANGELOG_ROW_VALUES.items():
        cell = cl.cell(row=cl_row, column=col)
        cell.value = value
        copy_style(cl.cell(row=cl_template, column=col), cell)
    report["changelog_row"] = cl_row

    values = tuple(ws.cell(row=new_row, column=c).value for c in range(1, 26))

    wb.save(LEDGER)
    print("ledger saved:", json.dumps(report, ensure_ascii=False))

    # ---- 6. 同步无损基线 ----
    index = LedgerIndex.load(ROOT)
    baseline = json.loads(BASELINE.read_text(encoding="utf-8"))

    baseline["assets"][ASSET_ID] = {
        "v": _row_digest(values),
        "c": "场景",
        "d": "scenes",
    }
    baseline["asset_count"] = len(baseline["assets"])
    counts = baseline.get("category_counts", {})
    counts["场景"] = counts.get("场景", 0) + 1
    baseline["category_counts"] = counts

    # 重新采集 9 域并集，重算列级/专表摘要
    collected = []
    for domain in index.domains:
        wb2 = load_workbook(domain.path)
        for row_number, vals in read_source_rows(wb2[ASSET_SHEET]):
            collected.append((_text(vals[0]), vals))
    baseline["column_digests"] = {
        str(col): col_digest(collected, col) for col in CONTENT_COLUMNS
    }
    scene_wb = load_workbook(LEDGER)
    for sheet in ("3D-场景通用", "3D-设施", "3D-其他"):
        baseline["sheet_digests"][sheet] = sheet_digest(scene_wb[sheet])

    BASELINE.write_text(
        json.dumps(baseline, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    print(
        "baseline updated: asset_count=%d scene=%d 3D-场景通用=%s"
        % (baseline["asset_count"], counts["场景"], baseline["sheet_digests"]["3D-场景通用"][:16])
    )
    print("TOKEN_LCORRIDOR_LEDGER_REGISTER_DONE")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
