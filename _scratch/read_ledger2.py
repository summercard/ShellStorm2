import os
import openpyxl

LEDGER = "I:/工作项目/shellstrom2/ShellStorm2/assets/registry/ledgers/ShellStorm2_敌人账本_v001.xlsx"
wb = openpyxl.load_workbook(LEDGER, data_only=True)

with open("I:/工作项目/shellstrom2/ShellStorm2/_scratch/ledger_ranged.txt", "w", encoding="utf-8") as f:
    for ws in wb.worksheets:
        f.write("\n==== SHEET: %s (%s) ====\n" % (ws.title, ws.dimensions))
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        for ri, row in enumerate(rows, start=1):
            line = " | ".join("" if c is None else str(c) for c in row)
            if any(k in line for k in ["RANGED", "SPORESHOOTER", "ranged_caster", "MELEE", "FUNGBOAR", "melee_chaser"]):
                f.write("R%d: %s\n" % (ri, line))
    # also dump 3D-敌人 prefab sheet full
    f.write("\n==== 3D-敌人 full ====\n")
    try:
        ws3 = wb["3D-敌人"]
        for ri, row in enumerate(ws3.iter_rows(values_only=True), start=1):
            line = " | ".join("" if c is None else str(c) for c in row)
            if line.strip():
                f.write("R%d: %s\n" % (ri, line))
    except Exception as e:
        f.write("no 3D-敌人: %s\n" % e)
print("OK")
