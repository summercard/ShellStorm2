import openpyxl, json, sys
p = "assets/registry/ledgers/ShellStorm2_场景账本_v001.xlsx"
wb = openpyxl.load_workbook(p, data_only=False)
print("SHEETS:", wb.sheetnames)
ws = wb["资产主表"]
print("dims:", ws.dimensions, "max_row", ws.max_row, "max_col", ws.max_column)
hdr = [c.value for c in ws[1]]
print("HEADER:")
for i,h in enumerate(hdr,1):
    print("  col%02d %s" % (i,h))
