from pathlib import Path
import hashlib
import json
import re
import shutil
from datetime import datetime
from copy import copy
import openpyxl
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(r'I:/工作项目/shellstrom2/ShellStorm2')
GLB = ROOT / 'assets/art/props/base_world_3d/components/base99_radio/prp_base99_radio_visual_top3d.glb'
SOURCE = ROOT / 'assets/art/props/base_world_3d/source/base99_radio/prp_base99_radio_source_v002.blend'
PREFAB = ROOT / 'assets/art/props/base_world_3d/runtime/base99_radio/prp_base99_radio_root_top3d.tscn'
PROP = ROOT / 'assets/registry/ledgers/ShellStorm2_道具账本_v001.xlsx'
MASTER = ROOT / 'assets/registry/ShellStorm2_美术资产台账_v001.xlsx'
BASELINE = ROOT / 'assets/registry/ledger_split_baseline.json'
OUT = ROOT / 'outputs/base99_radio_v002'
BACKUP = OUT / 'backup_before_v002_registration'
AID = 'PRP-BASE99-RADIO-3D'

def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()

def text(v):
    return '' if v is None else str(v).strip()

def row_digest(values):
    return hashlib.sha256('\x1f'.join(text(values[c-1]) for c in range(1,26) if c not in (18,19)).encode()).hexdigest()

def col_digest(rows, col):
    ordered = sorted(rows, key=lambda item: (text(item[1][0]), text(item[1][col-1])))
    return hashlib.sha256('\x1e'.join(text(v[col-1]) for _, v in ordered).encode()).hexdigest()

def rows(ws):
    return [(r,[ws.cell(r,c).value for c in range(1,26)]) for r in range(6,ws.max_row+1) if text(ws.cell(r,1).value)]

def dedupe_key(r):
    return f'=LOWER(TRIM(C{r})&"|"&TRIM(D{r})&"|"&TRIM(E{r})&"|"&TRIM(F{r})&"|"&TRIM(H{r})&"|"&TRIM(I{r}))'

def dedupe_result(r,last):
    return f'=IF(COUNTIF($R$6:$R${last},R{r})>1,"重复","唯一")'

def copy_dv(ws, last):
    vals = list(ws.data_validations.dataValidation)
    ws.data_validations.dataValidation = []
    for old in vals:
        new = DataValidation(type=old.type, formula1=old.formula1, formula2=old.formula2, allow_blank=old.allow_blank, showErrorMessage=old.showErrorMessage, showInputMessage=old.showInputMessage, error=old.error, errorTitle=old.errorTitle, prompt=old.prompt, promptTitle=old.promptTitle)
        for rng in old.sqref.ranges:
            new.add(re.sub(r'([A-Z]+)\d+$', rf'\g<1>{last}', str(rng)))
        ws.add_data_validation(new)

for p in [PROP, MASTER, BASELINE, PREFAB]:
    BACKUP.mkdir(parents=True, exist_ok=True)
    shutil.copy2(p, BACKUP / p.name)
sha_glb = sha(GLB)
sha_prefab = sha(PREFAB)
sha_source = sha(SOURCE)

prefab = PREFAB.read_text(encoding='utf-8')
prefab = prefab.replace('metadata/asset_version = "v001"', 'metadata/asset_version = "v002"')
prefab = re.sub(r'metadata/model_sha256 = "[^"]*"\n', '', prefab)
prefab = prefab.replace('metadata/bounds_size_m = Vector3(0.414, 0.228, 0.411)', 'metadata/bounds_size_m = Vector3(0.414, 0.228, 0.411)\nmetadata/model_sha256 = "' + sha_glb + '"\nmetadata/model_faces = 598\nmetadata/model_triangles = 1116')
PREFAB.write_text(prefab, encoding='utf-8')
sha_prefab = sha(PREFAB)

wb = openpyxl.load_workbook(PROP, data_only=False)
ws = wb['资产主表']
found = [r for r in range(6,ws.max_row+1) if ws.cell(r,1).value == AID]
assert found == [27], found
r = 27
ws.cell(r,13).value = 'v002'
ws.cell(r,14).value = '0.414m × 0.228m × 0.411m；Visual+StatusLight 598面 / 1116三角形（Visual 566面 / 1068三角形，StatusLight 32面 / 48三角形）；独立视觉GLB/PackedScene；四材质角色、公共色盘外链、PaletteUV/Closest、GLB不嵌纹理；WorldCollision与点击区域分离；AudioStreamPlayer3D走Music总线；左键/E循环关→基地音乐A→基地音乐B→关'
ws.cell(r,15).value = 'assets/art/props/base_world_3d/runtime/base99_radio/prp_base99_radio_root_top3d.tscn'
ws.cell(r,16).value = '; '.join(['assets/art/props/base_world_3d/source/base99_radio/prp_base99_radio_source_v002.blend','assets/art/props/base_world_3d/components/base99_radio/prp_base99_radio_visual_top3d.glb','src/base3d/Base99Radio3D.gd','assets/art/environments/base_facility_3d/runtime/env_base_facility_art_layout_top3d.tscn'])
ws.cell(r,20).value = sha_prefab
ws.cell(r,25).value = '2026-10-05：v002复古低模改造；保留v001源可回滚；598面/1116三角形；Prefab路径、ItemRoot/Visual/StatusLight接口、脚本、正式摆位与音乐逻辑不变；GLB SHA-256=' + sha_glb + '；源SHA-256=' + sha_source
last = 27
for rr in range(6,last+1):
    ws.cell(rr,18).value = dedupe_key(rr)
    ws.cell(rr,19).value = dedupe_result(rr,last)
copy_dv(ws,last)
ws.auto_filter.ref = f'A5:X{last}'
pf = wb['3D-道具']
assert pf.cell(8,1).value == AID
pf.cell(8,3).value = 'assets/art/props/base_world_3d/runtime/base99_radio/prp_base99_radio_root_top3d.tscn'
pf.cell(8,4).value = 'assets/art/props/base_world_3d/components/base99_radio/prp_base99_radio_visual_top3d.glb'
pf.cell(8,5).value = 'assets/art/props/base_world_3d/source/base99_radio/prp_base99_radio_source_v002.blend'
pf.cell(8,11).value = '0.414m × 0.228m × 0.411m；598面 / 1116三角形；-Z'
pf.cell(8,16).value = 'v002'
pf.cell(8,17).value = '根节点=ItemRoot；Visual+StatusLight稳定接口；GLB SHA-256=' + sha_glb
log = wb['域变更日志']
log_row = log.max_row + 1
for c in range(1,8): log.cell(log_row,c)._style = copy(log.cell(log_row-1,c)._style)
for c,v in enumerate(['v0.1.6','2026-10-05','99F阁楼收音机v002低模改造','道具 / 场景可交互道具','更新 PRP-BASE99-RADIO-3D 至 v002；598面/1116三角形；保留v001回滚源；稳定Prefab路径与ItemRoot/Visual/StatusLight接口不变。','同步Prefab版本与GLB SHA；更新规格、源路径、哈希和无损基线；不改脚本、正式场景摆位和音乐逻辑。','Codex'],1): log.cell(log_row,c).value = v
wb.save(PROP)

master = openpyxl.load_workbook(MASTER)
master['3D Prefab总控']['C7'] = 4
master['3D Prefab总控']['E7'] = 2
master['3D Prefab总控']['F7'] = 2
master['3D Prefab总控']['G7'] = '桌椅、家具、收音机、容器和摆放装饰物'
master.save(MASTER)

idx = json.loads((ROOT/'assets/registry/ledger_index.json').read_text(encoding='utf-8'))
all_rows=[]
for d in idx['domains']:
    book=openpyxl.load_workbook(ROOT/'assets/registry/ledgers'/d['file'], read_only=True)
    all_rows.extend((str(v[0]),v) for _,v in rows(book['资产主表']))
    book.close()
lookup=dict(all_rows)
base=json.loads(BASELINE.read_text(encoding='utf-8'))
base['assets'][AID]={'v':row_digest(lookup[AID]),'c':'道具','d':'props'}
base['asset_count']=len(base['assets'])
base['captured_at']='2026-10-05'
base['column_digests']={str(c):col_digest(all_rows,c) for c in range(1,26) if c not in (18,19)}
from collections import Counter
base['category_counts']=dict(sorted(Counter(text(v[2]) for _,v in all_rows).items()))
BASELINE.write_text(json.dumps(base,ensure_ascii=False,indent=1)+'\n',encoding='utf-8')

manifest=json.loads((OUT/'asset_manifest.json').read_text(encoding='utf-8'))
manifest['hashes']={'glb_sha256':sha_glb,'source_blend_sha256':sha_source,'runtime_prefab_sha256':sha_prefab}
manifest['runtime_prefab_version']='v002'
manifest['ledger']={'file':'assets/registry/ledgers/ShellStorm2_道具账本_v001.xlsx','asset_sheet_row':27,'prefab_sheet':'3D-道具','prefab_sheet_row':8,'change_log_row':log_row}
(OUT/'asset_manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
print(json.dumps({'asset_id':AID,'glb_sha256':sha_glb,'source_sha256':sha_source,'prefab_sha256':sha_prefab,'asset_sheet_row':27,'prefab_sheet_row':8,'log_row':log_row,'baseline_assets':base['asset_count']},ensure_ascii=False,indent=2))
