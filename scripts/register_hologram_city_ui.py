"""Promote the existing level-select UI row; preserve every unrelated asset fingerprint."""
from pathlib import Path
from copy import copy
from collections import Counter
import datetime
import hashlib
import json
import shutil
import sys
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'tools' / 'asset_pipeline'))
import split_asset_ledger as schema

index = json.loads((ROOT / 'assets/registry/ledger_index.json').read_text(encoding='utf-8'))
domain = next(d for d in index['domains'] if d['key'] == 'ui')
path = ROOT / index['ledger_dir'] / domain['file']
baseline_path = ROOT / 'assets/registry/ledger_split_baseline.json'
backup = ROOT / '_scratch/hologram_city_ledger_r3_before'
backup.mkdir(exist_ok=True)
for file in [path, baseline_path]:
    target = backup / file.name
    if not target.exists():
        shutil.copy2(file, target)
wb = openpyxl.load_workbook(path)
ws = wb[index['asset_sheet']]
rows = schema.read_source_rows(ws)
row = next(r for r, values in rows if values[0] == 'UI-SCREEN-LEVEL-SELECT')
before = {(cell.row, cell.column): cell.value for values in ws for cell in values}
asset_path = 'assets/art/ui/expedition_city/hologram_city.tscn'
updates = {
    2: '远征全息城市选关', 8: '世界空间3D近景', 9: '主楼先起/双入口错峰/波浪扫描/方块粒子/倒序收起',
    10: '99F中央全息平台', 11: '原型已接入', 12: 'P0', 13: 'v002',
    14: '主体约3.1m宽；42座外缘错落楼群；径向密度衰减；GPU逐块升降；2座楼顶入口',
    15: asset_path, 16: 'scenes/RogueMapSelectMenu.gd; src/ui/HologramCity3D.gd; docs/v0.1/design/远征全息城市交互设计.md',
    17: '远征情报室;全息城市;镜头拉近;远征01;测试99;3D选关',
    20: hashlib.sha256((ROOT / asset_path).read_bytes()).hexdigest(),
    21: 'Codex', 22: datetime.datetime(2026, 9, 29),
    23: '项目原创程序几何；用户提供风格参考',
    25: '2026-09-29：按用户参考制作Box城市；提前1秒展开，继承原环境至近景后渐暗；主楼优先/入口错峰/方块过渡；同场景世界空间入口。真实渲染/点击/键盘/手柄/返程通过，参考图像素级一致性未签署。',
}
for column, value in updates.items():
    ws.cell(row, column).value = value
assert all(cell.value == before.get((cell.row,cell.column)) for values in ws for cell in values if cell.row != row)
log = wb['域变更日志']
if not any(log.cell(r,1).value == 'v0.2.1' for r in range(6,log.max_row+1)):
    r = log.max_row + 1
    values = ['v0.2.1','2026-09-29','选关表现替换','UI-SCREEN-LEVEL-SELECT','原弹窗改为平台锚定全息城市及专属镜头','ID与出发携带物契约保持','Codex']
    for c,value in enumerate(values,1):
        log.cell(r,c).value = value
        log.cell(r,c)._style = copy(log.cell(6,c)._style)
wb.save(path)
check = openpyxl.load_workbook(path)
new_values = next(v for _,v in schema.read_source_rows(check[index['asset_sheet']]) if v[0]=='UI-SCREEN-LEVEL-SELECT')
baseline = json.loads(baseline_path.read_text(encoding='utf-8'))
baseline['assets']['UI-SCREEN-LEVEL-SELECT']['v'] = schema._row_digest(new_values)
all_rows = []
for d in index['domains']:
    book = openpyxl.load_workbook(ROOT / index['ledger_dir'] / d['file'])
    all_rows.extend(schema.read_source_rows(book[index['asset_sheet']]))
baseline['column_digests'] = {str(c):schema.col_digest(all_rows,c) for c in schema.CONTENT_COLUMNS}
baseline['category_counts'] = dict(Counter(str(v[2]) for _,v in all_rows))
baseline_path.write_bytes((json.dumps(baseline,ensure_ascii=False,indent=1)+'\n').replace('\n','\r\n').encode('utf-8'))
print('HOLOGRAM_CITY_LEDGER_UPDATED row=%s assets=%s unrelated_rows=unchanged' % (row,len(rows)))
