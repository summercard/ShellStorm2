"""Promote only the axe grip correction, keeping other ledger identities frozen."""
from pathlib import Path
from copy import copy
from datetime import datetime
from collections import Counter
import sys,json,hashlib,re,ast
import openpyxl
ROOT=Path(__file__).resolve().parents[1]
sys.path[:0]=[str(ROOT/'scripts'),str(ROOT/'tools/asset_pipeline')]
from workspace_paths import work_path
from ledger_registry import LedgerIndex
from split_asset_ledger import read_source_rows,_row_digest,sheet_digest,col_digest,CONTENT_COLUMNS
INDEX=LedgerIndex.load(ROOT)
BASELINE=ROOT/'assets/registry/ledger_split_baseline.json'
OUT=work_path('sessions/2026-10-17_axe_edge_v042/placeholder').parent
B=Path('assets/art/characters/player/chr_player_capsule01_3d/variants/bunny01')
SOURCE=Path('art_source/characters/player/chr_player_capsule01_3d/variants/bunny01/source/animation/chr_bunny01_animation_v042.blend')
ASSET='CHR-PLY-CAPSULE01-3D-BUNNY01'
PREFAB=B/'production/v021/runtime/chr_bunny01_root_v021.tscn'
TRANSFER=B/'character_transfer_ledger.json'
DOC='docs/v0.1/development/2026-10-17_axe_edge_v042.md'
tree=ast.parse((ROOT/'scripts/register_bunny_weapon_idles_v025.py').read_text(encoding='utf-8'))
exec(compile(ast.Module(body=[n for n in tree.body if isinstance(n,ast.FunctionDef)],type_ignores=[]),'<ledger-helpers>','exec'))
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def addlog(wb,title,detail):
 s=wb['域变更日志'];nums=[int(m[1]) for row in s.values if (m:=re.fullmatch(r'v0\.1\.(\d+)',str(row[0])))]
 append_styled(s,[f'v0.1.{max(nums,default=0)+1}','2026-10-17',title,'角色握持角度',detail,DOC,'Codex'],6)
results={n:json.loads((OUT/'after'/(n+'.json')).read_text(encoding='utf-8')) for n in ['import','verify_player3d_melee_motion_visual','verify_player3d_melee_special_visual_visual','verify_player3d_melee_styles']}
assert all(v['exit_code']==0 and not v['unexpected_errors'] for v in results.values())
comparison=json.loads((OUT/'export_comparison.json').read_text());assert len(comparison['changed'])==14 and comparison['retained']==114
source_report=json.loads((OUT/'source_validation.json').read_text());assert len(source_report['clips'])==37
data=json.loads((ROOT/TRANSFER).read_text(encoding='utf-8'));assert data['version']=='v042'
assert all(sha(ROOT/e['path'])==e['sha256'] for e in data['files'])
data.update(status='active',validation_status='targeted_pass_with_declared_limits',validated_at='2026-10-17',verification=results,source_validation=str(OUT/'source_validation.json'),export_comparison=str(OUT/'export_comparison.json'),new_clip_count=0,corrected_clip_count=14,unchanged_clip_count=114,retained_legacy_clip_count=91,stage_history=['authored','exported_pending_godot_validation','validated','active'],limits=['Full project suite and mobile not executed','Balance remains provisional','Non-melee family charge not included'],expected_errors=[],unexpected_errors=[])
data['files'].append(dict(path=PREFAB.as_posix(),sha256=sha(ROOT/PREFAB),bytes=(ROOT/PREFAB).stat().st_size))
(ROOT/TRANSFER).write_text(json.dumps(data,ensure_ascii=False,indent=2),encoding='utf-8')
p=ROOT/B/'components/chr_bunny01_motion/asset_manifest.json';m=json.loads(p.read_text(encoding='utf-8'));m.update(asset_version='v042',source=SOURCE.as_posix(),clips=data['clips'],files=data['files'],limits=data['limits']);p.write_text(json.dumps(m,ensure_ascii=False,indent=2),encoding='utf-8')

PATH=INDEX.path_for_category('角色');wb=openpyxl.load_workbook(PATH);s=wb['资产主表'];row=next(r for r,v in read_source_rows(s) if v[0]==ASSET)
assert s.cell(row,13).value=='v041'
for c,v in {13:'v042',16:str(s.cell(row,16).value)+'; '+SOURCE.as_posix(),20:sha(ROOT/PREFAB),22:datetime(2026,10,17),25:str(s.cell(row,25).value)+'；v042修正斧子握持：绕柄轴-90°，刃面进入下劈平面、银刃领先；14斧动作仅WeaponSocket旋转改变，其余114动作逐值不变；'+DOC}.items():s.cell(row,c).value=v
s=wb['动画与状态']
for rr in range(1,s.max_row+1):
 if str(s.cell(rr,2).value).startswith('melee_axe_'):s.cell(rr,12).value='v042 active；源银刃-X/柄+Y，绕柄-90°使刃朝劈砍轨迹；接触拍刃朝下，双掌/阶段连续/真实渲染通过；'+SOURCE.as_posix()
for label,p in [('v042斧握持修正源',SOURCE),('v042正式中转',TRANSFER),('v042运行采样',B/'components/chr_bunny01_motion/anim_bunny01_library.json'),('v042动画GLB',B/'components/chr_bunny01_motion/anim_bunny01_library.glb')]:append_styled(wb['角色中转记录'],[label,p.as_posix(),sha(ROOT/p),(ROOT/p).stat().st_size],6)
s=wb['3D-角色']
for rr in range(1,s.max_row+1):
 if s.cell(rr,3).value==PREFAB.as_posix():
  s.cell(rr,5).value=SOURCE.as_posix()+'；模型v021保持';s.cell(rr,15).value='v042';s.cell(rr,16).value='128动作保持；仅斧14动作握持绕柄-90°校正；'+DOC
addlog(wb,'斧子刃口朝向修正v042','银刃沿竖向下劈方向领先；源、GLB、运行采样与真实渲染一致。');wb.save(PATH)
baseline_write(changed_asset=ASSET,sheets=['动画与状态','角色中转记录','3D-角色'])
PATH=INDEX.path_for_category('武器');wb=openpyxl.load_workbook(PATH);s=wb['资产主表'];aid='WPN-MELEE-BP-WARAXE';row=next(r for r,v in read_source_rows(s) if v[0]==aid)
s.cell(row,25).value=str(s.cell(row,25).value)+'；2026-10-17角色动作v042修正握持绕柄-90°，银刃沿下劈领先；武器模型未改；'+DOC
s.cell(row,22).value=datetime(2026,10,17);addlog(wb,'战斧持握刃向校正','角色源14动作修正，武器源模型保持。');wb.save(PATH);baseline_write(changed_asset=aid)
p=ROOT/'assets/registry/current_source_index.json';idx=json.loads(p.read_text(encoding='utf-8'));e=next(e for e in idx['sources'] if e['path'].endswith('/chr_bunny01_animation_v041.blend'));e.update(path=SOURCE.as_posix(),version=[42],sha256=sha(ROOT/SOURCE),bytes=(ROOT/SOURCE).stat().st_size);p.write_text(json.dumps(idx,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
print('AXE_V042_REGISTERED')
