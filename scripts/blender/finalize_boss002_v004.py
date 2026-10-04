from pathlib import Path
import sys,json,hashlib
from collections import Counter
from openpyxl import load_workbook
R=Path(__file__).resolve().parents[2];sys.path[:0]=[str(R/'scripts'),str(R/'tools/asset_pipeline')]
from ledger_registry import LedgerIndex
from split_asset_ledger import read_source_rows,_row_digest,CONTENT_COLUMNS,col_digest
B=R/'assets/art/enemies/bosses/enm_boss_monitor002';ID='ENM-BOSS-MONITOR002-3D';src=B/'source/enm_boss_monitor002_source_v004.blend';audit=json.loads((B/'previews/expressions_v004/audit.json').read_text());assert audit['passed']
ix=LedgerIndex.load(R);p=ix.path_for_category('敌人');w=load_workbook(p);s=w['资产主表'];row=next(r for r,v in read_source_rows(s) if v[0]==ID)
s.cell(row,13,'v004');s.cell(row,15,src.relative_to(R).as_posix());s.cell(row,20,hashlib.sha256(src.read_bytes()).hexdigest());s.cell(row,14,'18144三角面；六表情图集；无骨骼');s.cell(row,9,'T Pose / default,suspicious,angry,sleepy,taunting,glitched');s.cell(row,25,'v004：按参考图修正默认并补齐6种表情；ExpressionController.expression_index切换，24帧恒定插值预览；图片已内嵌。JSON记录各部件裁切与UV规则。身体不变，未接入Godot。')
w['域变更日志'].append(['v004','2026-10-03','Codex',ID,'六表情图集与可切换预览，保持18144面']);w.save(p)
blp=R/'assets/registry/ledger_split_baseline.json';bl=json.loads(blp.read_text(encoding='utf-8'));v=next(v for r,v in read_source_rows(s) if v[0]==ID);bl['assets'][ID]={'v':_row_digest(v),'c':'敌人','d':'enemies'}
rows=[]
for d in ix.domains:rows.extend(read_source_rows(load_workbook(R/'assets/registry/ledgers'/d.file)['资产主表']))
bl['column_digests']={str(c):col_digest(rows,c) for c in CONTENT_COLUMNS};bl['category_counts']=dict(Counter(v[2] for r,v in rows));blp.write_text(json.dumps(bl,ensure_ascii=False,indent=1)+'\n',encoding='utf-8')
p=B/'source/boss002_production_ledger.xlsx';w=load_workbook(p);s=w.active;s['C12']='六表情透明图集；三个平面6三角面；属性切换';s['D12']='v004已制作'
for rr in range(1,s.max_row+1):
    if s.cell(rr,1).value=='当前源版本':s.cell(rr,2,'v004');s.cell(rr,3,'source/enm_boss_monitor002_source_v004.blend')
    if s.cell(rr,1).value=='面数统计':s.cell(rr,2,'previews/expressions_v004/audit.json')
s.append(['表情库','source/expression_library_v004.json','0默认 1怀疑 2愤怒 3困倦 4挑衅 5故障']);s.append(['表情预览','帧1/25/49/73/97/121','只驱动贴图UV与剪片尺寸，无骨骼']);w.save(p)
m=json.loads((B/'asset_manifest.json').read_text());m['version']='v004';m['source']=src.relative_to(B).as_posix();m['expression_library']='source/expression_library_v004.json';m['files']={p.relative_to(B).as_posix():hashlib.sha256(p.read_bytes()).hexdigest() for p in B.rglob('*') if p.is_file() and p.name!='asset_manifest.json'};(B/'asset_manifest.json').write_text(json.dumps(m,ensure_ascii=False,indent=2),encoding='utf-8');print('BOSS002_V004_LEDGER_UPDATED')
