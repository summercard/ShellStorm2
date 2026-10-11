# -*- coding: utf-8 -*-
"""把已登记新怪的中文名 棘背兽 -> 暗影兽（项目既定的 openpyxl 外科式编辑）。

⚠️ 为什么不是本地 editor_sdk：实测它 save_file 后把《资产主表》
   R 列公式写成 "==LOWER(...)"（多一个等号）、S 列写成数组公式（openpyxl 读为
   ArrayFormula），check_asset_registry + verify_ledger_split 共报 28 条失败
   ⇒ 破坏台账公式契约。已 close_file 释放脏实例并回滚磁盘文件。

改点：
  账本《资产主表》r18   B(中文名) / Q(关键词)
  账本《域变更日志》r9  E(说明文字里的怪物名)
  内容库《怪物与Boss》r15 B(名称) / N(事实源里写错的账本行号 r15 -> r18)
基线的 r18 行指纹与列摘要同步补录（不改 AssetID）。
"""
from __future__ import annotations

import json
import shutil
import sys
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(r"I:\工作项目\shellstrom2\ShellStorm2")
sys.path.insert(0, str(ROOT / "scripts"))
sys.path.insert(0, str(ROOT / "tools" / "asset_pipeline"))
from ledger_registry import LedgerIndex  # noqa: E402
from split_asset_ledger import (  # noqa: E402
    ASSET_SHEET,
    CONTENT_COLUMNS,
    col_digest,
    read_source_rows,
    _row_digest,
    _text,
)

LEDGER = ROOT / "assets/registry/ledgers/ShellStorm2_敌人账本_v001.xlsx"
CONTENT = ROOT / "docs/v0.1/data/ShellStorm2_游戏内容数据库_v010.xlsx"
BASELINE = ROOT / "assets/registry/ledger_split_baseline.json"

NEW_ID = "ENM-BEAST-SPIKESHELL01"
OLD_NAME = "棘背兽"
NEW_NAME = "暗影兽"

# ======================= 1) 回滚（撤销本地编辑器的破坏）====================
for src in (LEDGER, CONTENT):
    bak = src.with_name(src.name + ".bak_rename_shadow")
    assert bak.exists(), f"缺备份 {bak}"
    shutil.copy2(bak, src)
    print(f"ROLLBACK {src.name} <- {bak.name}")

wb = load_workbook(LEDGER)
main = wb["资产主表"]
r6 = str(main.cell(6, 18).value)
s6 = str(main.cell(6, 19).value)
assert r6.startswith("=LOWER") and not r6.startswith("=="), f"R6 公式仍异常: {r6[:80]}"
assert "COUNTIF" in s6 and s6.startswith("=") and not s6.startswith("=="), f"S6 公式仍异常: {s6[:80]}"
print("回滚自检 OK：R6=%s... | S6=%s..." % (r6[:40], s6[:40]))

# ======================= 2) 前置断言 =====================================
log = wb["域变更日志"]
assert main.cell(18, 1).value == NEW_ID, main.cell(18, 1).value
assert main.cell(18, 2).value == OLD_NAME, main.cell(18, 2).value
assert str(main.cell(18, 17).value).startswith(OLD_NAME), main.cell(18, 17).value
assert log.cell(9, 1).value == "v0.1.3", log.cell(9, 1).value
assert OLD_NAME in str(log.cell(9, 5).value)

# ======================= 3) 改名（账本）==================================
main.cell(18, 2).value = NEW_NAME
kw = str(main.cell(18, 17).value)
main.cell(18, 17).value = kw.replace(OLD_NAME, NEW_NAME).replace("刺背", "暗影; 刺背", 1)
print("Q18:", main.cell(18, 17).value)

log.cell(9, 5).value = str(log.cell(9, 5).value).replace(OLD_NAME, NEW_NAME)
wb.save(LEDGER)
print("LEDGER_WRITTEN B18=%r" % main.cell(18, 2).value)

# ======================= 4) 改名（内容库）+ 修行号笔误 =====================
cwb = load_workbook(CONTENT)
mb = cwb["怪物与Boss"]
assert mb.cell(15, 1).value == "monster_spike_shell", mb.cell(15, 1).value
assert mb.cell(15, 2).value == OLD_NAME, mb.cell(15, 2).value
mb.cell(15, 2).value = NEW_NAME
src = str(mb.cell(15, 14).value)
assert "r15 =" in src, src[:120]
mb.cell(15, 14).value = src.replace("r15 =", "r18 =")
cwb.save(CONTENT)
print("CONTENT_WRITTEN B15=%r | N15=%r" % (mb.cell(15, 2).value, mb.cell(15, 14).value))

# ======================= 5) 基线补录（r18 指纹 + 列摘要）===================
index = LedgerIndex.load(ROOT)
bl = json.loads(BASELINE.read_text(encoding="utf-8"))
assert NEW_ID in bl["assets"], "基线里应有 r18 条目"

wb2 = load_workbook(LEDGER)
rows = read_source_rows(wb2[ASSET_SHEET])
hits = [(r, v) for r, v in rows if _text(v[0]) == NEW_ID]
assert len(hits) == 1, f"命中 {len(hits)} 行"
row_no, values = hits[0]
old_v, new_v = bl["assets"][NEW_ID]["v"], _row_digest(values)
bl["assets"][NEW_ID]["v"] = new_v

union, totals = [], {}
for domain in index.domains:
    w = load_workbook(domain.path)
    rr = read_source_rows(w[ASSET_SHEET])
    totals[domain.key] = len(rr)
    union.extend(rr)
bl["column_digests"] = {str(c): col_digest(union, c) for c in CONTENT_COLUMNS}
bl["category_counts"] = dict(sorted(Counter(_text(v[2]) for _r, v in union).items()))
bl["asset_count"] = len(bl["assets"])

shutil.copy2(BASELINE, BASELINE.with_name(BASELINE.name + ".bak_before_rename_shadow"))
BASELINE.write_text(json.dumps(bl, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")
print(f"BASELINE r18 digest {old_v[:12]} -> {new_v[:12]} | asset_count={bl['asset_count']}")
print("ledger_totals", totals)
print("\nDONE", NEW_ID, "->", NEW_NAME, f"| 账本 r{row_no} | 内容库 r15")
