from pathlib import Path
import sys,json,hashlib
from collections import Counter
from openpyxl import load_workbook
R=Path(__file__).resolve().parents[2];sys.path[:0]=[str(R/'scripts'),str(R/'tools/asset_pipeline')]
from ledger_registry import LedgerIndex
from split_asset_ledger import read_source_rows,_row_digest,CONTENT_COLUMNS,col_digest
B=R/'assets/art/enemies/bosses/enm_boss_monitor002';ID='ENM-BOSS-MONITOR002-3D';src=B/'source/enm_boss_monitor002_source_v002.blend'
ix=LedgerIndex.load(R);p=ix.path_for_category('敌人');w=load_workbook(p);s=w['资产主表'];row=next(r for r,v in read_source_rows(s) if v[0]==ID)
s.cell(row,13,'v002');s.cell(row,15,src.relative_to(R).as_posix());s.cell(row,20,hashlib.sha256(src.read_bytes()).hexdigest());s.cell(row,9,'T Pose / 四指平展手掌 / 无骨骼');s.cell(row,25,'v002：每手1拇指+3手指，手掌朝下平展。保留原手持道具为独立展示件，当前张掌不握持。无骨骼权重动作；未接入Godot。统计见previews/hands_v002/audit.json。')
w['域变更日志'].append(['v002','2026-10-03','Codex',ID,'四指手套改为平展姿势；保留v001源']);w.save(p)
blp=R/'assets/registry/ledger_split_baseline.json';bl=json.loads(blp.read_text(encoding='utf-8'));v=next(v for r,v in read_source_rows(s) if v[0]==ID);bl['assets'][ID]={'v':_row_digest(v),'c':'敌人','d':'enemies'}
rows=[]
for d in ix.domains:rows.extend(read_source_rows(load_workbook(R/'assets/registry/ledgers'/d.file)['资产主表']))
bl['column_digests']={str(c):col_digest(rows,c) for c in CONTENT_COLUMNS};bl['category_counts']=dict(Counter(v[2] for r,v in rows));blp.write_text(json.dumps(bl,ensure_ascii=False,indent=1)+'\n',encoding='utf-8')
p=B/'source/boss002_production_ledger.xlsx';w=load_workbook(p);s=w.active;s['D14']='v002已制作';s['C14']='每手1拇指+3手指，平展掌心朝下';s.append(['当前源版本','v002','source/enm_boss_monitor002_source_v002.blend']);s.append(['面数统计','previews/hands_v002/audit.json','含修改器求值三角面；不含摄影环境']);w.save(p)
m=json.loads((B/'asset_manifest.json').read_text());m['version']='v002';m['source']=src.relative_to(B).as_posix();m['files']={p.relative_to(B).as_posix():hashlib.sha256(p.read_bytes()).hexdigest() for p in B.rglob('*') if p.is_file() and p.name!='asset_manifest.json'};(B/'asset_manifest.json').write_text(json.dumps(m,ensure_ascii=False,indent=2),encoding='utf-8')
print('BOSS002_V002_LEDGER_UPDATED')
