from pathlib import Path
import sys,json,hashlib,shutil
from collections import Counter
from openpyxl import load_workbook
R=Path(__file__).resolve().parents[2];sys.path[:0]=[str(R/'scripts'),str(R/'tools/asset_pipeline')]
from ledger_registry import LedgerIndex
from split_asset_ledger import read_source_rows,_row_digest,CONTENT_COLUMNS,col_digest
B=R/'assets/art/enemies/bosses/enm_boss_monitor002';ID='ENM-BOSS-MONITOR002-3D';src=B/'source/enm_boss_monitor002_model_v007.blend';audit=json.loads((B/'previews/rig_v007/audit.json').read_text(encoding='utf-8'));assert audit['passed'];contract=json.loads((B/'source/rig_contract_v007.json').read_text(encoding='utf-8'))
ix=LedgerIndex.load(R);p=ix.path_for_category('敌人');w=load_workbook(p);s=w['资产主表'];before={v[0]:v for r,v in read_source_rows(s)};row=next(r for r,v in read_source_rows(s) if v[0]==ID)
backup=R/'_scratch/boss002_v007_ledger_backup';backup.mkdir(exist_ok=True);shutil.copy2(p,backup/p.name);shutil.copy2(R/'assets/registry/ledger_split_baseline.json',backup/'ledger_split_baseline.json')
values={13:'v007',15:src.relative_to(R).as_posix(),20:hashlib.sha256(src.read_bytes()).hexdigest(),14:'18320三角面；52骨；每臂6段；3材质；中心后轴独立屏幕旋转',9:'T Pose / 六表情状态 / 10类动作设计，正式动作待制作',16:'source/enm_boss_monitor002_model_v007.blend + enm_boss_monitor002_animation_v007.blend；SKEL-MONITOR002-003；用户2026-10-04中心轴/多段臂/3材质要求',22:'2026-10-04',25:'v007 authored_rigged。五官下移；动画设计见docs/v0.1/design/Boss002显示器动画设计.md，10类15剪辑尚未制作。中心后轴连接双臂，屏幕独立旋转；六段弹簧臂FK弯曲+伸缩；身体色盘UV、六表情单图集、屏幕向上UV滚动，共3材质。33项源级检查通过。未导出/未接入Godot。'}
for c,v in values.items():s.cell(row,c,v)
w['域变更日志'].append(['v007','2026-10-04','Codex',ID,'五官按红框下移，10类/15剪辑动画设计，正式动作未制作；33项检查通过']);w.save(p)
after={v[0]:v for r,v in read_source_rows(load_workbook(p)['资产主表'])};assert all(before[k]==after[k] for k in before if k!=ID)
blp=R/'assets/registry/ledger_split_baseline.json';bl=json.loads(blp.read_text(encoding='utf-8'));bl['assets'][ID]={'v':_row_digest(after[ID]),'c':'敌人','d':'enemies'}
rows=[]
for d in ix.domains:rows.extend(read_source_rows(load_workbook(R/'assets/registry/ledgers'/d.file)['资产主表']))
bl['column_digests']={str(c):col_digest(rows,c) for c in CONTENT_COLUMNS};bl['category_counts']=dict(Counter(v[2] for r,v in rows));blp.write_text(json.dumps(bl,ensure_ascii=False,indent=1)+'\n',encoding='utf-8')
p=B/'source/boss002_production_ledger.xlsx';shutil.copy2(p,backup/p.name);w=load_workbook(p);s=w.active;s['B3']='authored_rigged';s['B4']='v007中心轴/6段双臂/3材质双母版';s['C13']='每臂6段FK；stretch_L/R控制弹簧长度；根随rear_axle';s['C17']='3段柔性支撑延伸到屏幕中心后轴'
for rr in range(1,s.max_row+1):
 label=s.cell(rr,1).value
 if label=='当前源版本':s.cell(rr,2,'v007');s.cell(rr,3,'模型model_v007 + 动作animation_v007')
 if label=='面数统计':s.cell(rr,2,'18,320三角面');s.cell(rr,3,'source/rig_contract_v007.json')
 if label=='骨架':s.cell(rr,2,'SKEL-MONITOR002-003');s.cell(rr,3,'52骨，双母版同签名；左右臂各6段')
 if label=='绑定验证':s.cell(rr,2,'previews/rig_v007/audit.json');s.cell(rr,3,'33项通过；5个QA动作不是正式游戏动作')
s.append(['v007材质','共3个','身体色盘UV / 六表情共用图集 / 屏幕代码']);s.append(['v007控制','rear_axle / monitor_spin / arm_01..06.L/R','stretch_L/R伸缩；code_scroll正数向上滚动']);w.save(p)
m=json.loads((B/'asset_manifest.json').read_text(encoding='utf-8'));m.update(version='v007',stage='authored_rigged',source=src.relative_to(B).as_posix(),rigged=True,triangles=contract['triangles'],skeleton_id=contract['skeleton_id'],rig_contract='source/rig_contract_v007.json',animation_source='source/enm_boss_monitor002_animation_v007.blend',material_count=3,materials=contract['materials']);m['files']={p.relative_to(B).as_posix():hashlib.sha256(p.read_bytes()).hexdigest() for p in B.rglob('*') if p.is_file() and p.name!='asset_manifest.json' and p.suffix!='.blend1'};(B/'asset_manifest.json').write_text(json.dumps(m,ensure_ascii=False,indent=2),encoding='utf-8');print('BOSS002_V007_LEDGER_UPDATED_OTHER_ROWS_UNCHANGED')

