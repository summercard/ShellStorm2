# -*- coding: utf-8 -*-
"""为新增资产补录「无损基线」，并修正域变更日志里对基线的描述。

背景（别重蹈覆辙）：
  - verify_ledger_split.py L119-121 对「基线里没有的 AssetID」直接 fail(asset_not_in_baseline)，
    它**不区分**「被静默新增」和「有意新增」⇒ 新增资产必须同步补录基线，否则门禁长期红。
  - ⛔ 不能跑 split_asset_ledger.py --baseline-only 来重建：
    它从 ledger_index.json 的 master.path(= 总目录) 读《资产主表》，而总目录**已不持有资产行**
    ⇒ 会把 409 行基线清空，反而制造 409 条 asset_lost。
  - 基线自身记录的 source 就是 "assets/registry/ledgers (current runtime-truth reconciliation)"，
    即「从当前分账本重建」是既有口径。本次做**最小编辑**：只加新增那一条，
    其余 409 条指纹逐位不动（历史投放/篡改仍能被抓），再重算列摘要与分类计数。
"""
from __future__ import annotations

import json
import shutil
import sys
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(r"I:\工作项目\shellstrom2\ShellStorm2")
sys.path.insert(0, str(ROOT / "scripts"))
sys.path.insert(0, str(ROOT / "tools" / "asset_pipeline"))

from ledger_registry import LedgerIndex  # noqa: E402
from split_asset_ledger import (  # noqa: E402
    ASSET_SHEET,
    CONTENT_COLUMNS,
    col_digest,
    read_source_rows,
    _row_digest,
)

NEW_ID = "ENM-BEAST-SPIKESHELL01"
BASELINE = ROOT / "assets/registry/ledger_split_baseline.json"
LEDGER = ROOT / "assets/registry/ledgers/ShellStorm2_敌人账本_v001.xlsx"

index = LedgerIndex.load(ROOT)
bl = json.loads(BASELINE.read_text(encoding="utf-8"))

if NEW_ID in bl["assets"]:
    print("基线已含该 AssetID，无需补录")
else:
    # ---- 取新行的内容指纹 ----
    wb = load_workbook(LEDGER)
    rows = read_source_rows(wb[ASSET_SHEET])
    hits = [(r, v) for r, v in rows if str(v[0]).strip() == NEW_ID]
    assert len(hits) == 1, f"账本里 {NEW_ID} 命中 {len(hits)} 次"
    row_number, values = hits[0]
    digest = _row_digest(values)
    assert values[2] == "敌人", values[2]

    before_count = bl["asset_count"]
    bl["assets"][NEW_ID] = {"v": digest, "c": "敌人", "d": "enemies"}
    bl["asset_count"] = len(bl["assets"])

    # ---- 重算列摘要 + 分类计数（其余资产指纹保持原值）----
    union = []
    totals = {}
    for domain in index.domains:
        w = load_workbook(domain.path)
        r = read_source_rows(w[ASSET_SHEET])
        totals[domain.key] = len(r)
        union.extend(r)
    bl["column_digests"] = {str(c): col_digest(union, c) for c in CONTENT_COLUMNS}
    bl["category_counts"] = dict(sorted(Counter(str(v[2]).strip() for _r, v in union).items()))

    shutil.copy2(BASELINE, BASELINE.with_name(BASELINE.name + ".bak_before_spikeshell"))
    BASELINE.write_text(json.dumps(bl, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")
    print(f"BASELINE_RECONCILED {before_count} -> {bl['asset_count']}  row={row_number}  v={digest[:16]}...")
    print("ledger_totals", totals)

# ---- 修正域变更日志 r9 第 6 列的错误描述（原先写成「仍 exit 0」，实际会 fail）----
wb = load_workbook(LEDGER)
log = wb["域变更日志"]
assert log.cell(9, 1).value == "v0.1.3", log.cell(9, 1).value
log.cell(9, 6).value = (
    "新增资产行 ⇒ 其余 11 行内容未变，仅《总览》区间随之扩到 $18（属预期）。"
    "无损基线 verify_ledger_split 原会 fail(asset_not_in_baseline) —— 已按既有"
    "「从当前分账本重建」口径做最小补录（只加本行指纹，其余 409 条逐位不变，"
    "列摘要与分类计数重算），补录前基线备份见 ledger_split_baseline.json.bak_before_spikeshell。"
)
wb.save(LEDGER)
print("LOG_UPDATED r9c6")
