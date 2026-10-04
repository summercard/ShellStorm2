from pathlib import Path
import sys,json,hashlib
from collections import Counter
from openpyxl import load_workbook
R=Path(__file__).resolve().parents[2];sys.path[:0]=[str(R/'scripts'),str(R/'tools/asset_pipeline')]
from ledger_registry import LedgerIndex
from split_asset_ledger import read_source_rows,_row_digest,CONTENT_COLUMNS,col_digest
B=R/'assets/art/enemies/bosses/enm_boss_monitor002';ID='ENM-BOSS-MONITOR002-3D';src=B/'source/enm_boss_monitor002_model_v005.blend';audit=json.loads((B/'previews/rig_v005/audit.json').read_text());assert audit['passed'];contract=json.loads((B/'source/rig_contract_v005.json').read_text())
ix=LedgerIndex.load(R);p=ix.path_for_category('敌人');w=load_workbook(p);s=w['资产主表'];row=next(r for r,v in read_source_rows(s) if v[0]==ID)
s.cell(row,13,'v005');s.cell(row,15,src.relative_to(R).as_posix());s.cell(row,20,hashlib.sha256(src.read_bytes()).hexdigest());s.cell(row,14,'18308三角面；41骨；弹簧伸缩/三段支撑/六表情状态');s.cell(row,9,'T Pose / 六表情状态 / 5个QA绑定测试动作');s.cell(row,16,'source/enm_boss_monitor002_model_v005.blend + enm_boss_monitor002_animation_v005.blend；SKEL-MONITOR002-001；用户2026-10-03绑定要求');s.cell(row,25,'v005 authored_rigged。权重归一，弹簧伸缩、支撑弯曲、屏幕转动、五官跟位置不跟旋转、手指弯曲、键盘手骨挂载均实测。双母版骨架签名一致。表情取消轮播，状态持续保持。未导出/未接入Godot。')
w['域变更日志'].append(['v005','2026-10-03','Codex',ID,'41骨蒙皮与双母版；六表情状态；21项源级验收通过']);w.save(p)
blp=R/'assets/registry/ledger_split_baseline.json';bl=json.loads(blp.read_text(encoding='utf-8'));v=next(v for r,v in read_source_rows(s) if v[0]==ID);bl['assets'][ID]={'v':_row_digest(v),'c':'敌人','d':'enemies'}
rows=[]
for d in ix.domains:rows.extend(read_source_rows(load_workbook(R/'assets/registry/ledgers'/d.file)['资产主表']))
bl['column_digests']={str(c):col_digest(rows,c) for c in CONTENT_COLUMNS};bl['category_counts']=dict(Counter(v[2] for r,v in rows));blp.write_text(json.dumps(bl,ensure_ascii=False,indent=1)+'\n',encoding='utf-8')
p=B/'source/boss002_production_ledger.xlsx';w=load_workbook(p);s=w.active;s['B3']='authored_rigged';s['B4']='T Pose骨骼蒙皮；模型/动作双母版；六表情状态';s['C12']='六表情状态，关闭轮播；旋转时五官跟位置不跟旋转';s['D12']='v005已绑定';s['C13']='手端控制骨驱动伸缩；线圈半径保持';s['C14']='4指两节骨蒙皮，每手1拇指+3手指';s['C15']='键盘独立道具，挂prop_socket.L';s['C16']='5节数据线骨链';s['C17']='3段柔性支撑蒙皮；底座固定'
for rr in range(1,s.max_row+1):
    label=s.cell(rr,1).value
    if label=='当前源版本':s.cell(rr,2,'v005');s.cell(rr,3,'模型model_v005 + 动作animation_v005')
    if label=='面数统计':s.cell(rr,2,'18,308三角面');s.cell(rr,3,'source/rig_contract_v005.json')
    if label=='延期项目':s.cell(rr,2,'GLB约束烘焙、Godot导入与正式动作未执行')
    if label=='表情预览':s.cell(rr,2,'改为持久状态');s.cell(rr,3,'ExpressionController.expression_index；不再随时间轮播')
s.append(['骨架','SKEL-MONITOR002-001','41骨，双母版同签名']);s.append(['绑定验证','previews/rig_v005/audit.json','21项通过；5个QA动作不是正式游戏动作']);w.save(p)
m=json.loads((B/'asset_manifest.json').read_text());m.update(version='v005',stage='authored_rigged',source=src.relative_to(B).as_posix(),rigged=True,triangles=contract['triangles'],skeleton_id=contract['skeleton_id'],rig_contract='source/rig_contract_v005.json',animation_source='source/enm_boss_monitor002_animation_v005.blend');m['files']={p.relative_to(B).as_posix():hashlib.sha256(p.read_bytes()).hexdigest() for p in B.rglob('*') if p.is_file() and p.name!='asset_manifest.json'};(B/'asset_manifest.json').write_text(json.dumps(m,ensure_ascii=False,indent=2),encoding='utf-8');print('BOSS002_V005_LEDGER_UPDATED')
