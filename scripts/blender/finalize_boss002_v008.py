from pathlib import Path
import sys,json,hashlib,shutil
from collections import Counter
from openpyxl import load_workbook
R=Path(__file__).resolve().parents[2];sys.path[:0]=[str(R/'scripts'),str(R/'tools/asset_pipeline')]
from ledger_registry import LedgerIndex
from split_asset_ledger import read_source_rows,_row_digest,CONTENT_COLUMNS,col_digest
B=R/'assets/art/enemies/bosses/enm_boss_monitor002';ID='ENM-BOSS-MONITOR002-3D';src=B/'source/enm_boss_monitor002_model_v008.blend';audit=json.loads((B/'previews/idle_v008/audit.json').read_text(encoding='utf-8'));assert audit['passed'];contract=json.loads((B/'source/rig_contract_v008.json').read_text(encoding='utf-8'))
ix=LedgerIndex.load(R);p=ix.path_for_category('敌人');w=load_workbook(p);s=w['资产主表'];before={v[0]:v for r,v in read_source_rows(s)};row=next(r for r,v in read_source_rows(s) if v[0]==ID)
backup=R/'_scratch/boss002_v008_ledger_backup';backup.mkdir(exist_ok=True);shutil.copy2(p,backup/p.name);shutil.copy2(R/'assets/registry/ledger_split_baseline.json',backup/'ledger_split_baseline.json')
values={13:'v008',15:src.relative_to(R).as_posix(),20:hashlib.sha256(src.read_bytes()).hexdigest(),14:'18320三角面；52骨；每臂6段；3材质；中心后轴独立屏幕旋转',9:'T Pose模型 / 六表情状态 / 正式idle 3.2秒循环',16:'source/enm_boss_monitor002_model_v008.blend + enm_boss_monitor002_animation_v008.blend；SKEL-MONITOR002-002；用户2026-10-04中心轴/多段臂/3材质要求',22:'2026-10-04',25:'v008 authored_rigged。五官恢复v006位置；正式idle已制作，3.2秒30fps循环，其余正式动作未制作。中心后轴连接双臂，屏幕独立旋转；六段弹簧臂FK弯曲+伸缩；身体色盘UV、六表情单图集、屏幕向上UV滚动，共3材质。9项源级检查通过。未导出/未接入Godot。'}
for c,v in values.items():s.cell(row,c,v)
w['域变更日志'].append(['v008','2026-10-04','Codex',ID,'五官恢复；正式idle循环，弧线下垂双臂和滞后回摆；9项检查通过']);w.save(p)
after={v[0]:v for r,v in read_source_rows(load_workbook(p)['资产主表'])};assert all(before[k]==after[k] for k in before if k!=ID)
blp=R/'assets/registry/ledger_split_baseline.json';bl=json.loads(blp.read_text(encoding='utf-8'));bl['assets'][ID]={'v':_row_digest(after[ID]),'c':'敌人','d':'enemies'}
rows=[]
for d in ix.domains:rows.extend(read_source_rows(load_workbook(R/'assets/registry/ledgers'/d.file)['资产主表']))
bl['column_digests']={str(c):col_digest(rows,c) for c in CONTENT_COLUMNS};bl['category_counts']=dict(Counter(v[2] for r,v in rows));blp.write_text(json.dumps(bl,ensure_ascii=False,indent=1)+'\n',encoding='utf-8')
p=B/'source/boss002_production_ledger.xlsx';shutil.copy2(p,backup/p.name);w=load_workbook(p);s=w.active;s['B3']='authored_rigged';s['B4']='v008中心轴/6段双臂/3材质双母版';s['C13']='每臂6段FK；stretch_L/R控制弹簧长度；根随rear_axle';s['C17']='3段柔性支撑延伸到屏幕中心后轴'
for rr in range(1,s.max_row+1):
 label=s.cell(rr,1).value
 if label=='当前源版本':s.cell(rr,2,'v008');s.cell(rr,3,'模型model_v008 + 动作animation_v008')
 if label=='面数统计':s.cell(rr,2,'18,320三角面');s.cell(rr,3,'source/rig_contract_v008.json')
 if label=='骨架':s.cell(rr,2,'SKEL-MONITOR002-002');s.cell(rr,3,'52骨，双母版同签名；左右臂各6段')
 if label=='绑定验证':s.cell(rr,2,'previews/idle_v008/audit.json');s.cell(rr,3,'9项通过；正式idle；其他正式动画待制作')
s.append(['v008材质','共3个','身体色盘UV / 六表情共用图集 / 屏幕代码']);s.append(['v008控制','rear_axle / monitor_spin / arm_01..06.L/R','stretch_L/R伸缩；code_scroll正数向上滚动']);w.save(p)
m=json.loads((B/'asset_manifest.json').read_text(encoding='utf-8'));m.update(version='v008',stage='authored_rigged',source=src.relative_to(B).as_posix(),rigged=True,triangles=contract['triangles'],skeleton_id=contract['skeleton_id'],rig_contract='source/rig_contract_v008.json',animation_source='source/enm_boss_monitor002_animation_v008.blend',material_count=3,materials=contract['materials']);m['files']={p.relative_to(B).as_posix():hashlib.sha256(p.read_bytes()).hexdigest() for p in B.rglob('*') if p.is_file() and p.name!='asset_manifest.json' and p.suffix!='.blend1'};(B/'asset_manifest.json').write_text(json.dumps(m,ensure_ascii=False,indent=2),encoding='utf-8');print('BOSS002_V008_LEDGER_UPDATED_OTHER_ROWS_UNCHANGED')

