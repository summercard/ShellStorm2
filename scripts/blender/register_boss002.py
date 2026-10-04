from pathlib import Path
import sys, json, copy, re, hashlib, shutil
from datetime import datetime
from collections import Counter
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill, Alignment
ROOT=Path(__file__).resolve().parents[2]
sys.path[:0]=[str(ROOT/'scripts'),str(ROOT/'tools/asset_pipeline')]
from ledger_registry import LedgerIndex
from split_asset_ledger import FIRST_DATA_ROW, read_source_rows, dedupe_key_formula, dedupe_result_formula, _row_digest, CONTENT_COLUMNS, col_digest
ID='ENM-BOSS-MONITOR002-3D'
SLUG='enm_boss_monitor002'
BASE=ROOT/'assets/art/enemies/bosses'/SLUG
index=LedgerIndex.load(ROOT)
p=index.path_for_category('敌人')
wb=load_workbook(p); ws=wb['资产主表']
found=[r for r,v in read_source_rows(ws) if v[0]==ID]
finish='--finish' in sys.argv
if not found:
    backup=ROOT/'_scratch/boss002_before'; backup.mkdir(parents=True,exist_ok=True)
    shutil.copy2(p,backup/p.name)
    shutil.copy2(ROOT/'assets/registry/ledger_split_baseline.json',backup/'ledger_split_baseline.json')
    for d in index.domains:
        w=load_workbook(index.path_for_domain(d),read_only=False) if hasattr(index,'path_for_domain') else load_workbook(ROOT/'assets/registry/ledgers'/d.file,read_only=False)
        assert not any(v[0]==ID for r,v in read_source_rows(w['资产主表']))
    old=max(r for r,v in read_source_rows(ws)); row=old+1
    vals=[ID,'Boss002 显示器 MONITOR.EXE','敌人','boss','boss_monitor002','root_3d',None,'Top3D / Blender +Y','T Pose / 无骨骼','Boss002 独立美术源','待制作','P1','v001','创作高度约3.4m；运行尺寸待定',None,'用户2026-10-03参考图；src/enemy3d/BossContentCatalog.gd','Boss002;monitor;显示器;平面五官',None,None,None,'Codex',datetime(2026,10,3),'用户设计图；原创程序建模','BOSS002_MONITOR','用户指定不绑骨；无动作母版；未接入运行池']
    for c,v in enumerate(vals,1):
        ws.cell(row,c,v); ws.cell(row,c)._style=copy.copy(ws.cell(old,c)._style)
    ws.row_dimensions[row].height=ws.row_dimensions[old].height
    for r in range(FIRST_DATA_ROW,row+1):
        ws.cell(r,18,dedupe_key_formula(r)); ws.cell(r,19,dedupe_result_formula(r,row))
    for s in [wb['总览']]:
        for cells in s:
            for cell in cells:
                if cell.data_type=='f': cell.value=re.sub(r'(\$[A-Z]+\$)'+str(old)+r'\b',lambda m:m[1]+str(row),cell.value)
    for dv in ws.data_validations.dataValidation:
        ranges=[]
        for rr in dv.sqref.ranges:
            if rr.max_row==old: rr.max_row=row
            ranges.append(str(rr))
        dv.sqref=' '.join(ranges)
    ws.auto_filter.ref=f'A5:Y{row}'
    wb['域变更日志'].append(['v0.1.新增Boss002',datetime(2026,10,3),'Codex',ID,'新增002显示器Boss，仅T Pose无骨骼模型，独立制作明细账；尚无运行时映射'])
else: row=found[0]
source=BASE/'source'/f'{SLUG}_source_v001.blend'
if finish:
    assert source.exists()
    ws.cell(row,11,'Blender源已完成'); ws.cell(row,15,source.relative_to(ROOT).as_posix())
    ws.cell(row,20,hashlib.sha256(source.read_bytes()).hexdigest())
    ws.cell(row,25,'authored_unrigged：T Pose模型完成；用户要求暂不绑骨。无动画、无GLB、未接入Godot。制作明细账见source/boss002_production_ledger.xlsx；验证见previews/source_audit.json。')
wb.save(p)
blp=ROOT/'assets/registry/ledger_split_baseline.json'; bl=json.loads(blp.read_text(encoding='utf-8'))
values=next(v for r,v in read_source_rows(ws) if v[0]==ID)
bl['assets'][ID]={'v':_row_digest(values),'c':'敌人','d':'enemies'}; bl['asset_count']=len(bl['assets'])
allrows=[]
for d in index.domains:
    path=ROOT/'assets/registry/ledgers'/d.file
    allrows.extend(read_source_rows(load_workbook(path,read_only=False)['资产主表']))
bl['column_digests']={str(c):col_digest(allrows,c) for c in CONTENT_COLUMNS}
bl['category_counts']=dict(Counter(v[2] for r,v in allrows))
blp.write_text(json.dumps(bl,ensure_ascii=False,indent=1)+'\n',encoding='utf-8')
for d in ['source','previews','reference']: (BASE/d).mkdir(parents=True,exist_ok=True)
if not finish:
    shutil.copy2(Path('C:/Users/ZHUANG~1/AppData/Local/Temp/codex-clipboard-e519b7fb-586c-4cf5-ab72-06c991e33dd5.png'),BASE/'reference/design.png')
book=Workbook(); s=book.active; s.title='Boss002制作明细'
rows=[['Boss 002 · MONITOR.EXE 制作账本'],['资产编号',ID],['阶段','authored_unrigged' if finish else 'design_only'],['制作范围','T Pose，无骨骼，无动画，Blender美术源'],['主登记','敌人分账本 / 资产主表 / '+str(row)],['创作约定','米制；脚底原点；Blender +Y 正面；根缩放1'],['消费者','未来 BossContentCatalog → Enemy3D；目前未建立此Boss内容映射'],['运行尺寸','未指定；本轮不改碰撞或玩法'],[],['部件','集合','制作要求','状态'],['显示器','01_MONITOR','竖屏、银灰边框、黑色后壳、绿色代码','已制作' if finish else '待制作'],['漂浮五官','02_FACE','独立平面剪片：异形大眼、圆眼、红嘴','已制作' if finish else '待制作'],['双臂','03_ARMS','后部连接；水平线圈；T Pose','已制作' if finish else '待制作'],['手套','04_GLOVES','白色卡通手套与袖口','已制作' if finish else '待制作'],['键盘','05_KEYBOARD','独立按键，握持关系可编辑','已制作' if finish else '待制作'],['数据线','06_CABLE','长线垂落回环与绿色插头','已制作' if finish else '待制作'],['底座','07_STAND','旋转铰链、支柱、弧形底座','已制作' if finish else '待制作'],[],['延期项目','骨架/权重/动作双母版/GLB/Godot接入均未执行，按用户范围延期。']]
for r in rows:s.append(r)
for cells in s:
    for c in cells:c.font=Font(name='Microsoft YaHei',size=11,color='223344'); c.alignment=Alignment(vertical='center',wrap_text=True)
s['A1'].font=Font(name='Microsoft YaHei',size=16,bold=True); s.merge_cells('A1:D1')
for c in s[10]:c.fill=PatternFill('solid',fgColor='DCE8DE');c.font=Font(name='Microsoft YaHei',size=11,bold=True)
for col,width in [('A',23),('B',44),('C',53),('D',15)]:s.column_dimensions[col].width=width
for r in range(1,s.max_row+1):s.row_dimensions[r].height=32
s.sheet_view.showGridLines=False;s.freeze_panes='B11';s.auto_filter.ref='A10:D17'
book.save(BASE/'source/boss002_production_ledger.xlsx')
print('REGISTERED',ID,'row',row,'finish',finish)

