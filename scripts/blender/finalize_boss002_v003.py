from pathlib import Path
import sys,json,hashlib
from collections import Counter
from openpyxl import load_workbook
R=Path(__file__).resolve().parents[2];sys.path[:0]=[str(R/'scripts'),str(R/'tools/asset_pipeline')]
from ledger_registry import LedgerIndex
from split_asset_ledger import read_source_rows,_row_digest,CONTENT_COLUMNS,col_digest
B=R/'assets/art/enemies/bosses/enm_boss_monitor002';ID='ENM-BOSS-MONITOR002-3D';src=B/'source/enm_boss_monitor002_source_v003.blend'
audit=json.loads((B/'previews/optimized_v003/audit.json').read_text());assert audit['triangles']<20000
ix=LedgerIndex.load(R);p=ix.path_for_category('敌人');w=load_workbook(p);s=w['资产主表'];row=next(r for r,v in read_source_rows(s) if v[0]==ID)
s.cell(row,13,'v003');s.cell(row,15,src.relative_to(R).as_posix());s.cell(row,20,hashlib.sha256(src.read_bytes()).hexdigest());s.cell(row,14,f"{audit['triangles']}三角面；圆头四指；2张image-2贴图；无骨骼");s.cell(row,9,'T Pose / 圆头四指平展 / 无骨骼');s.cell(row,25,'v003低模：全模型18144三角面，每手2222；屏幕文字和五官用gpt-image-2贴图，键盘无字符；纹理已内嵌blend。保留v001/v002。未接入Godot。')
w['域变更日志'].append(['v003','2026-10-03','Codex',ID,'圆头手指；整套低于20000三角面；image-2代码与表情贴图；键盘减面无字']);w.save(p)
blp=R/'assets/registry/ledger_split_baseline.json';bl=json.loads(blp.read_text(encoding='utf-8'));v=next(v for r,v in read_source_rows(s) if v[0]==ID);bl['assets'][ID]={'v':_row_digest(v),'c':'敌人','d':'enemies'}
rows=[]
for d in ix.domains:rows.extend(read_source_rows(load_workbook(R/'assets/registry/ledgers'/d.file)['资产主表']))
bl['column_digests']={str(c):col_digest(rows,c) for c in CONTENT_COLUMNS};bl['category_counts']=dict(Counter(v[2] for r,v in rows));blp.write_text(json.dumps(bl,ensure_ascii=False,indent=1)+'\n',encoding='utf-8')
p=B/'source/boss002_production_ledger.xlsx';w=load_workbook(p);s=w.active
s['C11']='竖屏银灰边框；image-2绿色代码贴图';s['C12']='image-2透明五官贴图，三个平面共6三角面';s['C14']='1拇指+3手指，圆头平展，每手2222三角面';s['D14']='v003已优化';s['C15']='低模按键，无字符；整键盘2184三角面'
for rr in range(1,s.max_row+1):
    if s.cell(rr,1).value=='当前源版本':s.cell(rr,2,'v003');s.cell(rr,3,'source/enm_boss_monitor002_source_v003.blend')
    if s.cell(rr,1).value=='面数统计':s.cell(rr,2,'previews/optimized_v003/audit.json');s.cell(rr,3,'整套18144三角面；无未应用修改器')
s.append(['贴图来源','gpt-image-2','source/textures_v003/；已打包至blend']);w.save(p)
m=json.loads((B/'asset_manifest.json').read_text());m['version']='v003';m['source']=src.relative_to(B).as_posix();m['triangles']=audit['triangles'];m['texture_model']='gpt-image-2';m['files']={p.relative_to(B).as_posix():hashlib.sha256(p.read_bytes()).hexdigest() for p in B.rglob('*') if p.is_file() and p.name!='asset_manifest.json'};(B/'asset_manifest.json').write_text(json.dumps(m,ensure_ascii=False,indent=2),encoding='utf-8');print('BOSS002_V003_LEDGER_UPDATED')

