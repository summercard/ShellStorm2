"""Finalize verified motion integration without rebasing unrelated ledger assets."""
from pathlib import Path
from copy import copy
from collections import Counter
from datetime import datetime
import ast,hashlib,json,shutil,subprocess,sys,os
import openpyxl
ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'scripts'));sys.path.insert(0,str(ROOT/'tools/asset_pipeline'))
from ledger_registry import LedgerIndex
from split_asset_ledger import read_source_rows,_row_digest,sheet_digest,col_digest,CONTENT_COLUMNS
INDEX=LedgerIndex.load(ROOT);PATH=INDEX.path_for_category('角色');BASELINE=ROOT/'assets/registry/ledger_split_baseline.json'
OUT=ROOT/'outputs/character_pipeline/carry_v031';BACKUP=ROOT/'_scratch/bunny_carry_v031_ledger_before'
ASSET='CHR-PLY-CAPSULE01-3D-BUNNY01';B=Path('assets/art/characters/player/chr_player_capsule01_3d/variants/bunny01')
SOURCE=B/'source/animation/chr_bunny01_animation_v031.blend';TRANSFER=B/'character_transfer_ledger.json'
PREFAB=B/'production/v021/runtime/chr_bunny01_root_v021.tscn';MANIFEST=B/'components/chr_bunny01_motion/asset_manifest.json'
tree=ast.parse((ROOT/'scripts/register_bunny_weapon_idles_v025.py').read_text(encoding='utf-8'))
exec(compile(ast.Module(body=[n for n in tree.body if isinstance(n,ast.FunctionDef)],type_ignores=[]),'<ledger-helpers>','exec'))
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
tests=['import','verify_player3d_directional_motion','verify_player3d_firing_motion','verify_player3d_directional_motion_visual','verify_player3d_animation_flow','verify_player3d_idle_animation_flow','verify_player3d_weapon_pose_collision_flow','verify_player3d_diy_flow','verify_player3d_avatar_bounds','verify_player3d_lower_body_socket_flow','verify_player3d_state_gallery_flow']
results={n:json.loads((OUT/(n+'.json')).read_text(encoding='utf-8')) for n in tests}
assert all(v['exit_code']==0 and not v['unexpected_errors'] for v in results.values())
data=json.loads((ROOT/TRANSFER).read_text(encoding='utf-8'))
assert all(sha(ROOT/e['path'])==e['sha256'] for e in data['files'])
assert len(data['clips'])==76
assert len(json.loads((OUT/'source_validation.json').read_text()))==18
assert not BACKUP.exists();gates('before');BACKUP.mkdir(parents=True)
shutil.copy2(PATH,BACKUP/PATH.name);shutil.copy2(BASELINE,BACKUP/BASELINE.name)
data.update(status='active',validation_status='targeted_pass_with_declared_limits',validated_at='2026-10-07',classification='version_increment',stage_history=['authored','exported_pending_godot_validation','validated','active'],verification={n:dict(exit_code=v['exit_code'],unexpected_errors=v['unexpected_errors'],user_data_isolated_before_autoload=True) for n,v in results.items()},limits=['No family reload/charge or heavy-melee action authored','Other longguns support offsets differ: 0.029-0.108m at default scale','No exact world foot lock or 2D directional blend','No full project or mobile performance suite'],runtime_default_scale=.8,runtime_default_height_m=1.2,weapon_grip_contract='Exported palm offsets relative to existing cuff/wrist HandJoint; carry orientation from source; no hand IK')
for p in [PREFAB,Path('src/player3d/CharacterMotionLibrary3D.gd'),Path('src/player3d/PlayerAvatar3D.gd'),Path('src/player3d/Player3D.gd')]:
    data['files'].append(dict(path=p.as_posix(),sha256=sha(ROOT/p),bytes=(ROOT/p).stat().st_size))
(ROOT/TRANSFER).write_text(json.dumps(data,ensure_ascii=False,indent=2),encoding='utf-8')
(ROOT/MANIFEST).write_text(json.dumps(dict(asset_id=ASSET,asset_version='v031',assembly_version='v021',status=data['status'],source=SOURCE.as_posix(),transfer=TRANSFER.as_posix(),prefab=PREFAB.as_posix(),skeleton_id=data['skeleton_id'],skeleton_sha256=data['skeleton_sha256'],clips=data['clips'],files=data['files'],limits=data['limits']),ensure_ascii=False,indent=2),encoding='utf-8')
before=openpyxl.load_workbook(PATH);wb=openpyxl.load_workbook(PATH);ws=wb['资产主表']
r=next(r for r,v in read_source_rows(ws) if v[0]==ASSET)
changes={9:'10 existing states / angled longgun-machinegun carry / 27 raised firing variants',11:'active',13:'v031',14:'创作高1.5m；默认运行倍率0.8/静止高1.2m；碰撞仍由Player3D拥有。',16:ws.cell(r,16).value+'; '+SOURCE.as_posix()+'; '+TRANSFER.as_posix()+'; '+MANIFEST.as_posix(),20:sha(ROOT/PREFAB),21:'Codex',22:datetime(2026,10,7),25:ws.cell(r,25).value+'；v031正式接入18条长枪机枪斜持修订并保留58条其他剪辑，长枪机枪小幅斜持、机枪收近3.5cm创作尺寸；身体脚步耳朵、短枪和射击不变。专项通过，命名/文档全局既有红项保留；其他长枪支撑差异和世界滑步、换弹蓄力近战待完善。见2026-10-07_player_motion_carry_v031.md。'}
for c,v in changes.items():ws.cell(r,c).value=v
lr=append_styled(wb['域变更日志'],['v0.1.16','2026-10-07','长枪机枪斜持与机枪收近','角色 / Bunny01','v031动作库76条；保持v021模型装配和碰撞；掌心挂点与源枪朝向。','专项通过；换弹蓄力近战/逐枪支撑/锁地待完善。','Codex'],17)
wb.save(PATH);baseline_write(changed_asset=ASSET)
# Specialized sheets are an independent transaction.
wb=openpyxl.load_workbook(PATH);ws=wb['动画与状态'];ar=[];newrows=[]
library=json.loads((ROOT/B/'components/chr_bunny01_motion/anim_bunny01_library.json').read_text())['clips']
for state,clip in library.items():
    if not clip.get('source_action'):continue
    matches=[i for i in range(1,ws.max_row+1) if ws.cell(i,2).value==state]
    note=f"v031 active；{clip['duration']:.1f}秒循环；CharacterMotionLibrary3D采样；{SOURCE.as_posix()}；换弹蓄力近战/锁地待完善。"
    if matches:
        assert len(matches)==1
        row=matches[0];ws.cell(row,12).value=note
    else:
        row=append_styled(ws,['表现动作变体（非顶层状态）',state,state,'moving + family/speed/local direction','Player3D状态所有者','独立原地脚步；无根运动','反移动方向拖曳/错相回弹','源动作采样，不用IK','主掌心跟随；双手类别公开支撑误差','无碰撞修改','无新增VFX',note],64);newrows.append(row)
    ar.append(row)
tr=[]
for label,p in [('v031动作源',SOURCE),('v031正式中转',TRANSFER),('v031运行采样库',B/'components/chr_bunny01_motion/anim_bunny01_library.json'),('v031标准GLB',B/'components/chr_bunny01_motion/anim_bunny01_library.glb')]:
    tr.append(append_styled(wb['角色中转记录'],[label,p.as_posix(),sha(ROOT/p),(ROOT/p).stat().st_size],112))
ws=wb['3D-角色'];pr=[]
for row in range(1,ws.max_row+1):
    if ws.cell(row,3).value==PREFAB.as_posix():
        pr.append(row);ws.cell(row,5).value=SOURCE.as_posix()+'；模型保持v021'
        ws.cell(row,7).value='src/player3d/PlayerAvatar3D.gd; src/player3d/CharacterMotionLibrary3D.gd'
        ws.cell(row,11).value='源高1.5m，默认倍率0.8/高1.2m；Player3D碰撞不变'
        ws.cell(row,14).value='active';ws.cell(row,15).value='v031'
        ws.cell(row,16).value='模型装配v021；动作v031；专项通过；其他长枪支撑差异、脚底锁地、换弹蓄力近战待完善。'+TRANSFER.as_posix()
wb.save(PATH);baseline_write(sheets=['动画与状态','角色中转记录','3D-角色'])
after=openpyxl.load_workbook(PATH);diff=[]
for name in before.sheetnames:
    a,b=before[name],after[name]
    for row in b:
        for c in row:
            if a.cell(c.row,c.column).value!=c.value:
                assert (name=='资产主表' and c.row==r and c.column in changes) or (name=='动画与状态' and c.row in ar and (c.row in newrows or c.column==12)) or (name=='角色中转记录' and c.row in tr) or (name=='域变更日志' and c.row==lr) or (name=='3D-角色' and c.row in pr and c.column in [5,7,11,14,15,16]),(name,c.coordinate)
                diff.append([name,c.coordinate])
    assert str(a.data_validations)==str(b.data_validations)
(OUT/'ledger_edit_report.json').write_text(json.dumps(dict(changed_cells=diff,animation_rows=ar,transfer_rows=tr,prefab_rows=pr),ensure_ascii=False,indent=2),encoding='utf-8')
gates('after');print('V031_REGISTERED',len(ar),tr,pr)
