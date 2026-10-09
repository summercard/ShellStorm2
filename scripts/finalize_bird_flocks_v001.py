import sys,json,hashlib,shutil
from pathlib import Path
from collections import Counter
from openpyxl import load_workbook
R=Path(__file__).resolve().parents[1]
sys.path[:0]=[str(R/'scripts'),str(R/'tools/asset_pipeline')]
from ledger_registry import LedgerIndex
from split_asset_ledger import read_source_rows,_row_digest,CONTENT_COLUMNS,col_digest,sheet_digest
idx=LedgerIndex.load(R); p=idx.path_for_category('特效'); b=R/'assets/registry/ledger_split_baseline.json'; out=R/'outputs/bird_flocks_v001'
w=load_workbook(p); s=w['资产主表']; baseline=json.loads(b.read_text(encoding='utf-8')); ids=['VFX-ENV-BIRDS-FLYBY-3D','VFX-ENV-BIRDS-GROUND-3D']
for r,values in read_source_rows(s):
    if values[0] in ids:s.cell(r,20,hashlib.sha256((R/values[14]).read_bytes()).hexdigest())
w.save(p)
for _,values in read_source_rows(s):
    if values[0] in ids:baseline['assets'][values[0]]={'v':_row_digest(values),'c':'特效','d':'vfx'}
union=[]
for d in idx.domains:union.extend(read_source_rows(load_workbook(d.path)['资产主表']))
baseline['column_digests']={str(c):col_digest(union,c) for c in CONTENT_COLUMNS}; baseline['category_counts']=dict(sorted(Counter(str(v[2]).strip() for _,v in union).items()))
b.write_text(json.dumps(baseline,ensure_ascii=False,indent=1)+'\n',encoding='utf-8')
before=load_workbook(out/'ledger_before.xlsx')
for row,values in read_source_rows(before['资产主表']):
    for col in CONTENT_COLUMNS:assert s.cell(row,col).value==values[col-1]
assert sheet_digest(before['3D-特效'])==sheet_digest(w['3D-特效'])
old=json.loads((out/'baseline_before.json').read_text(encoding='utf-8'))
assert all(baseline['assets'][k]==v for k,v in old['assets'].items())
for kind in ('flyby','ground'):
    folder=R/f'assets/art/vfx/environment_3d/bird_flocks/source/{kind}/v001'; qa=folder/'qa'; qa.mkdir(exist_ok=True)
    for name in (f'{kind}_motion_qa.json',f'{kind}_palette_qa.json'):shutil.copy2(out/name,qa/name)
    path=folder/'asset_manifest.json'; m=json.loads(path.read_text(encoding='utf-8')); m['source_sha256']=hashlib.sha256((R/m['source_blend']).read_bytes()).hexdigest(); path.write_text(json.dumps(m,ensure_ascii=False,indent=2),encoding='utf-8')
log=R/'docs/v0.1/development/CHANGELOG.md'; content=log.read_text(encoding='utf-8')
entry='## 2026-10-09 鸟群场景特效 Blender 源\n\n两套各7只的低模鸟群：12秒空中掠过、24秒落地停留再起飞；场景共享哑光材质与冷白色格。可编辑骨骼关键帧、源文件、完整预览和专项检查已交付，未接入Godot。[制作与验收记录](2026-10-09_bird_flock_blender_sources.md)。\n\n'
if '## 2026-10-09 鸟群场景特效 Blender 源' not in content:log.write_text(entry+content,encoding='utf-8')
index=R/'docs/v0.1/MODULE_INDEX.md'; text=index.read_text(encoding='utf-8'); lines=text.splitlines()
for i,line in enumerate(lines):
    if line.startswith('| VFX-POOL |') and 'bird_flock_blender_sources' not in line:lines[i]=line.rstrip().removesuffix('|').rstrip()+'；[鸟群双套Blender源](development/2026-10-09_bird_flock_blender_sources.md)已完成，未接入运行时 |'
index.write_text('\n'.join(lines)+'\n',encoding='utf-8')
print('FINAL_HASH_SYNC_OK; previous assets and prefab table preserved')
