"""Narrow ledger transaction: source-only scene kit; no fictitious Prefab rows."""
from pathlib import Path
import sys,json,hashlib,shutil,re
from copy import copy
from collections import Counter
from openpyxl import load_workbook
R=Path(__file__).resolve().parents[1];sys.path.insert(0,str(R/'tools/asset_pipeline'))
from split_asset_ledger import FIRST_DATA_ROW,read_source_rows,_row_digest,col_digest,CONTENT_COLUMNS,dedupe_key_formula,dedupe_result_formula
from ledger_registry import LedgerIndex
I=LedgerIndex.load(R);D=next(d for d in I.domains if d.key=='scenes')
O=R/'assets/art/environments/master_office_3d/source/env_father_office/v002';A='ENV-BATTLE-FATHER-OFFICE-SOURCE'
story='--story' in sys.argv
if story:O=R/'assets/art/environments/master_office_3d/source/env_block00_story_rooms/v001';A='ENV-BATTLE-BLOCK00-STORY-ROOMS-SOURCE'
assert json.loads((O/'task_acceptance.json').read_text(encoding='utf8'))['passed']
assert json.loads((O/'material_acceptance.json').read_text(encoding='utf8'))['passed']
F=O/('env_block00_story_rooms_source_v001.blend' if story else 'env_father_office_source_v002.blend');sha=hashlib.sha256(F.read_bytes()).hexdigest()
B=R/('outputs/block00_story_rooms_20261009/ledger_backup' if story else 'outputs/father_office_20261009/ledger_backup');B.mkdir(parents=True,exist_ok=True)
base=R/'assets/registry/ledger_split_baseline.json'
for f in [D.path,base]:
 if not (B/f.name).exists():shutil.copy2(f,B/f.name)
w=load_workbook(D.path);s=w['资产主表'];rows=read_source_rows(s)
assert not any(v[0]==A for _,v in rows),'Asset already registered; inspect before updating'
last=max(r for r,v in rows);nr=last+1
values=[A,'98F 父亲办公室·破损休息区原始源','场景','room','father_office','root',None,'Top3D','ruined','98F 区块00最内侧 master_office','Blender源已完成','P1','v002','15×20m；可见墙高11.9m；24组件/270实例；4共享材质；PaletteUV逐面验收',str(F.relative_to(R)).replace('\\','/'),'source/art/blender/master_office_layout/source/block_00_master_office_layout_v002.layout.json; src/world3d/Block00MasterOfficeLayout3D.gd','98F;父亲办公室;沙发;破墙;father_office',None,None,sha,'Codex','2026-10-09','按用户提供参考图原创建模；参考图权利由提供方确认','master_office','仅Blender源。24组件作为源内组成部分，可独立打开；无GLB/无PackedScene/未接入Godot。技术验收通过；用户视觉待复核。']
if story:
 values[1]='98F 会议室与走廊门厅·工业雕塑废墟源';values[3]='environment_kit_3d';values[4]='block00_story_rooms';values[9]='98F会议室40×15m、走廊5×20m、门厅15×15m；外链父亲办公室';values[12]='v001';values[13]='9新增组件+24外链组件；395实例；12雕塑共6976三角；中央6m净空；拆门后5m开口；四材质';values[16]='98F;会议室;苏联工业雕塑;走廊;门厅;壁画;废墟;断电线';values[22]='按用户文字要求原创建模；办公室风格沿用用户参考';values[23]='block00_story_rooms';values[24]='Blender源已完成；9个独立组件源与6张预览，42,587面PaletteUV验收通过。第三/第四房门墙和门扇在新美术源中拆除，旧运行时布局未改。办公室Library Link引用v002且SHA保持。未导出GLB/未创建Prefab/未接入Godot；用户视觉待复核。'
for c,v in enumerate(values,1):s.cell(nr,c,v);s.cell(nr,c)._style=copy(s.cell(last,c)._style)
s.row_dimensions[nr].height=s.row_dimensions[last].height
for r in range(FIRST_DATA_ROW,nr+1):s.cell(r,18,dedupe_key_formula(r));s.cell(r,19,dedupe_result_formula(r,nr))
for row in w['总览']:
 for cell in row:
  if isinstance(cell.value,str) and cell.value.startswith('='):cell.value=re.sub(r'(\$[A-Z]+\$)'+str(last)+r'\b',lambda m:m[1]+str(nr),cell.value)
for dv in s.data_validations.dataValidation:
 for rng in dv.sqref.ranges:
  if rng.max_row==last:rng.max_row=nr
s.auto_filter.ref=f'A5:Y{nr}'
log=w['域变更日志'];prev=str(log.cell(log.max_row,1).value);match=re.match(r'v(\d+)\.(\d+)\.(\d+)',prev);version=f'v{match[1]}.{match[2]}.{int(match[3])+1}' if match else 'v0.1.50'
log.append([version,'2026-10-09','新增办公室Blender源',A,values[-1],'新增资产与基线定点补录；保留既有资产与专表','Codex'])
w.save(D.path)
check=load_workbook(D.path);assert check['资产主表'].cell(nr,1).value==A
# Keep every preexisting content cell unchanged except the required S formula span.
for r,old in rows:
 for col in CONTENT_COLUMNS:assert check['资产主表'].cell(r,col).value==old[col-1],(r,col)
bl=json.loads(base.read_text(encoding='utf8'));new=dict(read_source_rows(check['资产主表']))[nr]
bl['assets'][A]={'v':_row_digest(new),'c':'场景','d':'scenes'};bl['asset_count']=len(bl['assets'])
allrows=[]
for domain in I.domains:allrows.extend(read_source_rows(load_workbook(domain.path,read_only=False,data_only=False)['资产主表']))
bl['column_digests']={str(c):col_digest(allrows,c) for c in CONTENT_COLUMNS};bl['category_counts']=dict(Counter(v[2] for _,v in allrows))
base.write_text(json.dumps(bl,ensure_ascii=False,indent=1)+'\n',encoding='utf8')
print('REGISTERED',A,'row',nr,'sha',sha)
