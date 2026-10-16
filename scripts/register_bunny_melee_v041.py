"""Register verified melee actions and five VFX; preserve unrelated ledger rows."""
from pathlib import Path
from copy import copy
from collections import Counter
from datetime import datetime
import sys,json,hashlib,shutil,re
import openpyxl

R=Path(__file__).resolve().parents[1]
sys.path[:0]=[str(R/'scripts'),str(R/'tools/asset_pipeline')]
from workspace_paths import work_path
from ledger_registry import LedgerIndex
from split_asset_ledger import read_source_rows,_row_digest,sheet_digest,col_digest,CONTENT_COLUMNS,FIRST_DATA_ROW,dedupe_key_formula,dedupe_result_formula
INDEX=LedgerIndex.load(R)
OUT=work_path('sessions/2026-10-16_melee_v041/placeholder').parent
BACKUP=R.parent/'ShellStorm2-backup/2026-10-16_melee_v041/registration_before'
B=Path('assets/art/characters/player/chr_player_capsule01_3d/variants/bunny01')
SOURCE=Path('art_source/characters/player/chr_player_capsule01_3d/variants/bunny01/source/animation/chr_bunny01_animation_v041.blend')
TRANSFER=B/'character_transfer_ledger.json'
PREFAB=B/'production/v021/runtime/chr_bunny01_root_v021.tscn'
ASSET='CHR-PLY-CAPSULE01-3D-BUNNY01'
DOC='docs/v0.1/development/2026-10-16_melee_styles_v041.md'
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def backup(p):
 target=BACKUP/p.relative_to(R);target.parent.mkdir(parents=True,exist_ok=True)
 if not target.exists():shutil.copy2(p,target)
def append(ws,values,template=6):
 row=max(c.row for cells in ws for c in cells if c.value is not None)+1
 for i,value in enumerate(values,1):
  c=ws.cell(row,i,value);c._style=copy(ws.cell(template,i)._style);c.alignment=copy(ws.cell(template,i).alignment)
 ws.row_dimensions[row].height=70
 return row
def log(wb,title,detail):
 ws=wb['域变更日志'];versions=[]
 for row in ws.values:
  m=re.fullmatch(r'v0\.1\.(\d+)',str(row[0]))
  if m:versions.append(int(m[1]))
 append(ws,[f'v0.1.{max(versions,default=0)+1}','2026-10-16',title,'近战差异 v041',detail,DOC,'Codex'])

names=['import','verify_player3d_melee_motion_visual','verify_player3d_melee_styles','verify_player3d_melee_special_visual_visual','verify_3d_melee_combat_flow','verify_3d_melee_feedback_flow','verify_player3d_directional_motion','verify_player3d_weapon_switch','verify_player3d_animation_flow','verify_player3d_weapon_pose_collision_flow']
results={n:json.loads((OUT/'after'/(n+'.json')).read_text(encoding='utf-8')) for n in names}
assert all(r['exit_code']==0 and not r['unexpected_errors'] for r in results.values())
source_report=json.loads((OUT/'source_validation.json').read_text());assert len(source_report['clips'])==37
data=json.loads((R/TRANSFER).read_text(encoding='utf-8'));assert data['version']=='v041'
assert all(sha(R/e['path'])==e['sha256'] for e in data['files'])
clips=json.loads((R/B/'components/chr_bunny01_motion/anim_bunny01_library.json').read_text())['clips']
assert len(clips)==128
for key in ['melee_bat_idle','melee_blade_idle','melee_axe_idle','melee_blade_spin']:
 first,last=clips[key]['frames'][0],clips[key]['frames'][-1]
 for bone,a in first.items():
  b=last[bone];assert max(abs(x-y) for x,y in zip(a['p'],b['p']))<1e-5
  assert abs(abs(sum(x*y for x,y in zip(a['q'],b['q'])))-1)<1e-5
data.update(status='active',validation_status='targeted_pass_with_declared_limits',validated_at='2026-10-16',verification=results,source_validation=str(OUT/'source_validation.json'),new_clip_count=37,retained_legacy_clip_count=91,retained_clips=91,glb_clip_count=37,stage_history=['authored','exported_pending_godot_validation','validated','active'],limits=['Full project suite and mobile device performance not executed','Balance tuning remains provisional','Existing non-melee family charge action is not included'],unexpected_errors=[],expected_errors=[],not_executed=['Full project suite','Mobile device performance','Full expedition save/resume'],preview=str(OUT/'preview/contact_sheet.jpg'))
data['files'].append(dict(path=PREFAB.as_posix(),sha256=sha(R/PREFAB),bytes=(R/PREFAB).stat().st_size))
data['retained_clip_provenance']['rollback_library']=str(R.parent/'ShellStorm2-backup/2026-10-16_melee_v040/before_edits'/B/'components/chr_bunny01_motion/anim_bunny01_library.json')
backup(R/TRANSFER);(R/TRANSFER).write_text(json.dumps(data,ensure_ascii=False,indent=2),encoding='utf-8')
manifest=R/B/'components/chr_bunny01_motion/asset_manifest.json';backup(manifest)
m=json.loads(manifest.read_text(encoding='utf-8'));m.update(asset_version='v041',status='active',source=SOURCE.as_posix(),clips=data['clips'],files=data['files'],limits=data['limits']);manifest.write_text(json.dumps(m,ensure_ascii=False,indent=2),encoding='utf-8')

changes={};sheet_changes=[]
for cat in ['角色','武器','敌人','特效']:
 path=INDEX.path_for_category(cat);backup(path);before=sha(path);wb=openpyxl.load_workbook(path);ws=wb['资产主表']
 if cat=='角色':
  row=next(r for r,v in read_source_rows(ws) if v[0]==ASSET)
  assert ws.cell(row,13).value=='v039','Already registered; inspect before retry'
  for col,val in {9:'128 clips / 37 distinct melee actions / six melee phases',13:'v041',16:str(ws.cell(row,16).value)+'; '+SOURCE.as_posix(),20:sha(R/PREFAB),22:datetime(2026,10,16),25:str(ws.cell(row,25).value)+'；v041：棍普通挥击、斧慢速下劈与蓄力击退、剑交替横扫与持续旋风。37段Blender近战动作、91条旧动作保真；专项与真实渲染通过，详见'+DOC}.items():ws.cell(row,col).value=val
  changes[ASSET]='characters'
  for state,c in clips.items():
   if not state.startswith('melee_'):continue
   append(wb['动画与状态'],['近战表现（非移动顶层状态）',state,state,'PlayerMeleeCombat3D phase/attack_mode/spin_progress','玩法状态机',str(c['duration'])+'s；按真实进度采样','躯干发力、头部跟随；旋风整身转圈','左掌距握点0.18m','右掌对齐WeaponSocket','根静止/单位缩放/原碰撞','五种独立近战VFX；有效扣血局部停顿','v041 active；源与运行握点、阶段连续、真实渲染通过；'+SOURCE.as_posix()])
  for label,p in [('v041动作源',SOURCE),('v041中转',TRANSFER),('v041运行采样库',B/'components/chr_bunny01_motion/anim_bunny01_library.json'),('v041近战增量GLB',B/'components/chr_bunny01_motion/anim_bunny01_library.glb')]:append(wb['角色中转记录'],[label,p.as_posix(),sha(R/p),(R/p).stat().st_size])
  for cells in wb['3D-角色']:
   if cells[2].value==PREFAB.as_posix():
    rr=cells[0].row
    for c,v in {5:SOURCE.as_posix()+'；模型v021保持',14:'active',15:'v041',16:'37近战+91保留动作；五特效；'+DOC}.items():wb['3D-角色'].cell(rr,c).value=v
  sheet_changes+=['动画与状态','角色中转记录','3D-角色']
  log(wb,'三类近战37动作与命中反馈','三类攻击差异、敌方硬直、玩家局部1/2/4帧停顿；模型与移动动作保持。')
 elif cat=='武器':
  descriptions={'bp_baseball_bat':'普通三连挥棒；1帧停顿；现有风痕','bp_waraxe':'慢速头顶下劈；按住1秒蓄满，松开强击退；普通2帧/蓄力4帧停顿；蓄力环/竖劈/冲击','bp_greatblade':'点按左右交替三连；按住0.35秒旋风，0.55秒一圈/0.25秒伤害；松开结束；横扫/旋风特效'}
  for row,v in read_source_rows(ws):
   if v[4] in descriptions:
    ws.cell(row,25).value=str(v[24])+'；2026-10-16角色动作v041：'+descriptions[v[4]]+'；模型版本不变；'+DOC
    ws.cell(row,22).value=datetime(2026,10,16);changes[v[0]]='weapons'
  log(wb,'三类近战玩法与角色动作差异','模型复用；角色源、命中反馈及VFX按各武器独立接线。')
 elif cat=='敌人':
  s=wb['敌人动画与状态']
  row=next(r for r in range(6,s.max_row+1) if s.cell(r,2).value=='stagger')
  s.cell(row,9).value=str(s.cell(row,9).value or '')+'；近战有效扣血后普通怪0.20s，第三击/蓄力斧0.30s硬直，max不累加；允许击退/重力，死亡优先；Boss/监控器/精英逃跑保持原抗打断。Enemy3D.apply_melee_stagger。'+DOC
  sheet_changes+=['敌人动画与状态'];log(wb,'近战受击硬直','增加公开近战硬直接口；不改敌人模型动作源与Boss策略。')
 else:
  specs=[('AXE-CHARGE','axe_charge','VfxAxeCharge3D','战斧蓄力环','收缩橙金双环，随玩家，持有至释放'),('AXE-CLEAVE','axe_cleave','VfxAxeCleave3D','战斧下劈弧光','竖向橙金弧，随玩家朝向，有效阶段消退'),('AXE-IMPACT','axe_impact','VfxAxeImpact3D','战斧触地冲击','落点扩散环与8枚碎屑，世界锚定'),('BLADE-SWEEP','blade_sweep','VfxBladeSweep3D','大剑横扫弧光','蓝白弧，第二击反向，随玩家'),('BLADE-WHIRLWIND','blade_whirlwind','VfxBladeWhirlwind3D','大剑旋风环','蓝白双层断续环，每0.55秒旋转，释放/中断归还')]
  old_last=max(r for r,v in read_source_rows(ws))
  for i,(suffix,slug,cls,title,description) in enumerate(specs,24):
   aid='VFX-MELEE-'+suffix+'-3D';prefab=f'assets/art/vfx/combat_3d/vfx_{slug}_root_top3d.tscn';script=f'src/vfx/{cls}.gd'
   assert not any(v[0]==aid for _,v in read_source_rows(ws))
   append(ws,[aid,title,'特效','combat','melee_'+slug,'root_3d',None,'俯视3D',slug,'三类近战','已完成','P0','v001','静态Prefab几何+Shader；无碰撞',prefab,script+'；'+DOC,title,None,None,sha(R/prefab),'Codex',datetime(2026,10,16),'原创Godot几何与Shader',f'FX01-{i}',description+'；池化/跟随/回收/真实渲染通过；无伤害逻辑'],old_last)
   append(wb['3D-特效'],[f'FX01-{i}',aid,title,prefab,None,None,description,script,'无','无','无','基础半径1m×size','local -Z前方；中心为特效原点','PlayerMeleeCombat3D → VfxPool3D','完成','v001',DOC],8)
   changes[aid]='vfx'
  last=max(r for r,v in read_source_rows(ws))
  for row,v in read_source_rows(ws):ws.cell(row,18).value=dedupe_key_formula(row);ws.cell(row,19).value=dedupe_result_formula(row,last)
  for cells in wb['总览']:
   for c in cells:
    if c.data_type=='f':c.value=re.sub(r'(\$[A-Z]+\$)'+str(old_last)+r'\b',lambda m:m[1]+str(last),c.value)
  for dv in ws.data_validations.dataValidation:
   ranges=list(dv.sqref.ranges);dv.sqref=' '.join(f'{r.min_col and openpyxl.utils.get_column_letter(r.min_col)}{r.min_row}:{openpyxl.utils.get_column_letter(r.max_col)}{last if r.max_row==old_last else r.max_row}' for r in ranges)
  sheet_changes+=['3D-特效'];log(wb,'五种近战专属VFX','FX01-24至28；独立Prefab，池化跟随与显式结束，动作机唯一调用。')
 assert sha(path)==before,'Concurrent workbook edit; retry safely'
 wb.save(path)

# Baseline: update only explicitly changed identities/sheets; other row hashes remain frozen.
base=R/'assets/registry/ledger_split_baseline.json';backup(base);raw=base.read_bytes();bl=json.loads(raw);all_rows=[]
for domain in INDEX.domains:
 wb=openpyxl.load_workbook(domain.path);rows=read_source_rows(wb['资产主表']);all_rows.extend(rows)
 for _,v in rows:
  if v[0] in changes:bl['assets'][v[0]]={'v':_row_digest(v),'c':v[2],'d':changes[v[0]]}
 for s in sheet_changes:
  if s in wb.sheetnames:bl['sheet_digests'][s]=sheet_digest(wb[s])
 wb.close()
bl['asset_count']=len(bl['assets']);bl['category_counts']=dict(sorted(Counter(v[2] for _,v in all_rows).items()));bl['column_digests']={str(c):col_digest(all_rows,c) for c in CONTENT_COLUMNS}
assert base.read_bytes()==raw;base.write_text(json.dumps(bl,ensure_ascii=False,indent=1)+'\n',encoding='utf-8')
index_path=R/'assets/registry/current_source_index.json';backup(index_path);idx=json.loads(index_path.read_text(encoding='utf-8'))
entry=next(e for e in idx['sources'] if e['path'].endswith('/chr_bunny01_animation_v039.blend'));entry.update(path=SOURCE.as_posix(),bytes=(R/SOURCE).stat().st_size,version=[41],sha256=sha(R/SOURCE));index_path.write_text(json.dumps(idx,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')

# Content workbook describes actual inputs; no economy/ownership/formula changes.
path=R/'docs/v0.1/data/ShellStorm2_游戏内容数据库_v010.xlsx';backup(path);wb=openpyxl.load_workbook(path);ws=wb['武器']
descriptions={'bp_baseball_bat':'普通三连挥棒；按下立即攻击；1帧停顿','bp_waraxe':'慢速下劈；按住蓄力1秒上限，松开强击退；普通2帧/蓄力4帧停顿','bp_greatblade':'点按松开左右横扫三连；按住0.35秒旋风；每0.25秒伤害×0.55，松开0.3秒收势'}
for row in range(5,ws.max_row+1):
 key=ws.cell(row,1).value
 if key in descriptions:
  for c,v in {8:descriptions[key],17:'[已实装并专项验收] 六态近战动作机；含充能/持续旋转与打断',18:'点按/按住由攻击键决定；快照提供charge_ratio/spin_progress',19:descriptions[key]+'；Blender v041角色动作；五种专属VFX',20:'[专项通过] 源37动作、运行/真实渲染；完整游戏与移动端未验；'+DOC}.items():ws.cell(row,c).value=v
  if key=='bp_waraxe':ws.cell(row,7).value+='；普通预备0.48/0.52/0.65s，有效0.14/0.16/0.20s，收势0.40/0.46/0.65s；蓄力伤害1.2→2/击退1.4→2.6'
for row in range(5,wb['掉落物品'].max_row+1):
 key=str(wb['掉落物品'].cell(row,1).value).replace('weapon_','bp_')
 if key in descriptions:wb['掉落物品'].cell(row,9).value=descriptions[key]
wb.save(path)
(OUT/'registration_report.json').write_text(json.dumps({'assets':changes,'sheets':sheet_changes,'tests':names,'source_actions':37,'retained_clips':91},ensure_ascii=False,indent=2),encoding='utf-8')
print('MELEE_V041_REGISTERED',len(changes))
