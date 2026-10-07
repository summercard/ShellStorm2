"""Register source-only idles on the existing character identity in two transactions.

Uses the project's ledger-specific openpyxl contract. Does not promote runtime
status, change the active Prefab, add asset identities, or rebase unrelated rows.
"""
from pathlib import Path
from copy import copy
from collections import Counter
from datetime import datetime
import hashlib
import json
import os
import shutil
import subprocess
import sys

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'scripts'))
sys.path.insert(0, str(ROOT / 'tools/asset_pipeline'))
from ledger_registry import LedgerIndex
from split_asset_ledger import read_source_rows, _row_digest, sheet_digest, col_digest, CONTENT_COLUMNS

OUT = ROOT / 'outputs/character_pipeline/weapon_idle_v025'
BACKUP = ROOT / '_scratch/bunny_weapon_idles_v025_ledger_before'
INDEX = LedgerIndex.load(ROOT)
PATH = INDEX.path_for_category('角色')
BASELINE = ROOT / 'assets/registry/ledger_split_baseline.json'
ASSET = 'CHR-PLY-CAPSULE01-3D-BUNNY01'
SOURCE = 'assets/art/characters/player/chr_player_capsule01_3d/variants/bunny01/source/animation/chr_bunny01_animation_v025.blend'
TRANSFER = SOURCE.replace('chr_bunny01_animation_v025.blend', 'chr_bunny01_weapon_idles_v025.json')


def gates(tag):
    commands = {
        'structure': ['scripts/check_asset_registry.py', '--scope', 'structure'],
        'characters': ['scripts/check_asset_registry.py', '--scope', 'full', '--ledger', 'characters'],
        'split': ['tools/asset_pipeline/verify_ledger_split.py'],
        'docs': ['scripts/check_documentation_contracts.py'],
    }
    result = {}
    for name, args in commands.items():
        env = dict(os.environ, PYTHONUTF8='1', PYTHONIOENCODING='utf-8')
        if name == 'docs':
            # This historical checker invokes literal python3. Keep its checks
            # unchanged but route child processes through the same environment.
            code = "import subprocess,sys,runpy; original=subprocess.run; subprocess.run=lambda a,*p,**k: original(([sys.executable]+a[1:]) if isinstance(a,list) and a and a[0]=='python3' else a,*p,**k); runpy.run_path('scripts/check_documentation_contracts.py',run_name='__main__')"
            command = [sys.executable, '-X', 'utf8', '-c', code]
        else:
            command = [sys.executable, '-X', 'utf8', *args]
        p = subprocess.run(command, cwd=ROOT, capture_output=True, text=True, encoding='utf-8', env=env, timeout=120)
        (OUT / (tag + '_' + name + '.log')).write_text(p.stdout + p.stderr, encoding='utf-8')
        result[name] = p.returncode
    (OUT / (tag + '_gates.json')).write_text(json.dumps(result, indent=2), encoding='utf-8')
    print(tag, result, flush=True)
    return result


def append_styled(ws, values, template):
    row = max(c.row for cells in ws for c in cells if c.value is not None) + 1
    for col, value in enumerate(values, 1):
        c = ws.cell(row, col, value)
        c._style = copy(ws.cell(template, col)._style)
        c.alignment = copy(ws.cell(template, col).alignment)
    ws.row_dimensions[row].height = max(ws.row_dimensions[template].height or 30, 60)
    return row


def baseline_write(changed_asset=None, sheets=()):
    # Re-read at each transaction so unrelated already-saved work is preserved.
    before = BASELINE.read_bytes()
    baseline = json.loads(before)
    wb = openpyxl.load_workbook(PATH)
    if changed_asset:
        values = next(v for _, v in read_source_rows(wb['资产主表']) if v[0] == changed_asset)
        baseline['assets'][changed_asset]['v'] = _row_digest(values)
    for sheet in sheets:
        baseline['sheet_digests'][sheet] = sheet_digest(wb[sheet])
    rows = []
    for domain in INDEX.domains:
        # read_source_rows uses random cell access; read_only would reparse
        # the entire worksheet for every cell and make large ledgers quadratic.
        dw = openpyxl.load_workbook(domain.path)
        rows.extend(read_source_rows(dw['资产主表']))
        dw.close()
    baseline['column_digests'] = {str(c): col_digest(rows, c) for c in CONTENT_COLUMNS}
    baseline['category_counts'] = dict(sorted(Counter(v[2] for _, v in rows).items()))
    assert BASELINE.read_bytes() == before, 'Concurrent baseline edit; retry transaction safely'
    BASELINE.write_text(json.dumps(baseline, ensure_ascii=False, indent=1) + '\n', encoding='utf-8')


if '--gates-only' in sys.argv:
    gates('after')
    raise SystemExit(0)

resume = '--resume-main-transaction' in sys.argv
assert not BACKUP.exists() or resume, 'Already registered or backup exists; inspect instead of overwriting'
assert (ROOT / SOURCE).exists() and (ROOT / TRANSFER).exists()
validation = json.loads((OUT / 'validation.json').read_text(encoding='utf-8'))
assert set(validation['clips']) == {'sidearm_idle', 'longgun_idle', 'machinegun_idle'}
if not resume:
    gates('before')
    BACKUP.mkdir(parents=True)
    shutil.copy2(PATH, BACKUP / PATH.name)
    shutil.copy2(BASELINE, BACKUP / BASELINE.name)
before = openpyxl.load_workbook(BACKUP / PATH.name)
wb = openpyxl.load_workbook(PATH)
ws = wb['资产主表']
row = next(r for r, v in read_source_rows(ws) if v[0] == ASSET)
assert row == 13
if resume:
    assert SOURCE in ws.cell(row, 16).value
    logrow = next(c.row for cells in wb['域变更日志'] for c in cells if c.column == 1 and c.value == 'v0.1.10')
else:
    assert 'v025' not in (ws.cell(row, 16).value or '')
    ws.cell(row, 16).value += '; ' + SOURCE + '; ' + TRANSFER
    ws.cell(row, 22).value = datetime(2026, 10, 7)
    ws.cell(row, 25).value += '；2026-10-07新增v025站立持枪待机sidearm_idle/longgun_idle/machinegun_idle，源级authored，待导出接线；保留原14条动作、模型与骨架，四方向移动另待制作。运行路径/版本/状态保持，不代表新动作已启用。'
    logrow = append_styled(wb['域变更日志'], ['v0.1.10', '2026-10-07', '三类站立持枪待机源动画', '角色 / Bunny01 / 动画',
        '短枪单手朝上3.2秒；长枪胸前斜持3.6秒；机枪低位承重4秒。现有AssetID增加v025源，原14动作保留。',
        '仅源级authored，未导出/导入/改接线；活动版本与Prefab不变，主表与动作专表分事务登记。', 'Codex'], wb['域变更日志'].max_row)
    wb.save(PATH)
baseline_write(changed_asset=ASSET)

# Separate transaction for the guarded animation and transfer sheets.
wb = openpyxl.load_workbook(PATH)
assert not any(c.value == 'sidearm_idle' for cells in wb['动画与状态'] for c in cells), 'Specialized transaction already completed'
animrows = []
for family, title, duration, left, right in [
    ('sidearm', '短枪站立待机', '3.2', '左手自然放松', '右手单手握枪，枪口朝上，位于脸侧外方'),
    ('longgun', '长枪站立待机', '3.6', '左手托SupportHandSocket', '右手主握，步枪斜持胸前，枪托不贴腮'),
    ('machinegun', '机枪站立待机', '4.0', '左手托SupportHandSocket承重', '右手主握，双手低位承重，枪口略朝下')]:
    animrows.append(append_styled(wb['动画与状态'], [
        '表现动作变体（非顶层状态）', family + '_idle', title, 'idle + weapon_family=' + family, '由原状态机决定',
        duration + '秒循环；双脚站稳，根不移，独立呼吸曲线', '轻微呼吸跟随及耳部滞后', left, right,
        '保持现有组件契约', '无新增VFX/UI',
        'v025源已完成/authored；保存重开、四分之一帧骨链/双握点/循环检查通过。未导入Godot；' + SOURCE], 6))
transferrows = []
for label, path in [('v025动作母版（authored，未接线）', SOURCE), ('v025源级中转记录（待导出）', TRANSFER)]:
    data = (ROOT / path).read_bytes()
    transferrows.append(append_styled(wb['角色中转记录'], [label, path, hashlib.sha256(data).hexdigest(), len(data)], 5))
wb.save(PATH)
baseline_write(sheets=['动画与状态', '角色中转记录'])

after = openpyxl.load_workbook(PATH)
diff = []
for name in before.sheetnames:
    a, b = before[name], after[name]
    for cells in b:
        for cell in cells:
            old = a.cell(cell.row, cell.column).value
            if old != cell.value:
                diff.append((name, cell.coordinate))
                if name == '资产主表':
                    assert cell.coordinate in {'P13', 'V13', 'Y13'}
                elif name == '动画与状态':
                    assert cell.row in animrows
                elif name == '角色中转记录':
                    assert cell.row in transferrows
                elif name == '域变更日志':
                    assert cell.row == logrow
                else:
                    raise AssertionError(('Unrelated cell changed', name, cell.coordinate))
    assert set(map(str, a.merged_cells.ranges)) == set(map(str, b.merged_cells.ranges))
    assert str(a.data_validations) == str(b.data_validations)
report = dict(asset_id=ASSET, main_row=row, animation_rows=animrows, transfer_rows=transferrows, log_row=logrow,
              changed_cells=diff, runtime_fields_unchanged=True, backup=str(BACKUP.relative_to(ROOT)))
(OUT / 'ledger_edit_report.json').write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding='utf-8')
gates('after')
print('BUNNY_IDLE_LEDGER_REGISTERED', report, flush=True)
