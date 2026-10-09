from pathlib import Path
import json,sys,hashlib,shutil,re
from collections import Counter
from openpyxl import load_workbook
R=Path(__file__).resolve().parents[1];sys.path.insert(0,str(R/'tools/asset_pipeline'))
from split_asset_ledger import read_source_rows,_row_digest,col_digest,CONTENT_COLUMNS,sheet_digest
from ledger_registry import LedgerIndex
I=LedgerIndex.load(R);D=next(d for d in I.domains if d.key=='scenes');O=R/'outputs/block00_story_rooms_20261009/layout_adjustment'
base=R/'assets/registry/ledger_split_baseline.json';sha=lambda p:hashlib.sha256(p.read_bytes()).hexdigest()
for p in [D.path,base]:
 if not (O/p.name).exists():shutil.copy2(p,O/p.name)
P=R/'assets/art/environments/master_office_3d/source/env_block00_story_rooms/export/v001/import_manifest.json';m=json.loads(P.read_text('utf8'));qa=json.loads((R/'outputs/block00_story_rooms_20261009/facility_verification.json').read_text('utf8'));assert qa['passed']
changed={'ENV-BLOCK00-MASTER-OFFICE-FACILITY-LAYOUT':'沙发及抱枕/毯子按沙发基点缩放70%；地毯顶面0.091m，距旧地砖可视顶面8mm，无碰撞。','ENV-BLOCK00-LOBBY-FACILITY-LAYOUT':'壁画移至地面，向后倾斜22度并侧倾6度；实测网格最低点0.091m，不穿地。','ENV-BLOCK00-CORRIDOR-FACILITY-LAYOUT':'第三间取消墙面灯开关，保留常亮照明；通道和设施摆位不变。'}
w=load_workbook(D.path);s=w['资产主表'];old=read_source_rows(s)
for r,v in old:
 if v[0] in changed:
  if 'CORRIDOR' not in v[0]:s.cell(r,13,'v002');s.cell(r,20,sha(R/v[14]))
  s.cell(r,25,changed[v[0]]+' 原始Blend与组件GLB保持不变；Godot正式布局为摆位事实源。')
log=w['域变更日志'];match=re.match(r'v(\d+)\.(\d+)\.(\d+)',str(log.cell(log.max_row,1).value));ver=f'v{match[1]}.{match[2]}.{int(match[3])+1}';log.append([ver,'2026-10-09','98F按用户反馈调整尺寸/落地/开关',';'.join(changed),'沙发70%、地毯降低、三房去开关、四房壁画落地','局部改动，保留组件与源','Codex'])
w.save(D.path)
bl=json.loads(base.read_text('utf8'))
for r,v in read_source_rows(s):
 if v[0] in changed:bl['assets'][v[0]]={'v':_row_digest(v),'c':v[2],'d':'scenes'}
rows=[]
for domain in I.domains:rows.extend(read_source_rows(load_workbook(domain.path)['资产主表']))
bl['column_digests']={str(c):col_digest(rows,c) for c in CONTENT_COLUMNS};bl['category_counts']=dict(Counter(v[2] for _,v in rows));base.write_text(json.dumps(bl,ensure_ascii=False,indent=1)+'\n',encoding='utf8')
# Separate intentional Prefab-sheet transaction.
w=load_workbook(D.path);sheet=w['3D-场景通用']
for row in sheet:
 aid=row[0].value
 if aid in changed:
  n=row[0].row
  if 'CORRIDOR' not in aid:sheet.cell(n,15,'v002')
  sheet.cell(n,16,changed[aid])
w.save(D.path);bl['sheet_digests']['3D-场景通用']=sheet_digest(sheet);base.write_text(json.dumps(bl,ensure_ascii=False,indent=1)+'\n',encoding='utf8')
for d in m['rooms']:d['scene_sha256']=sha(R/d['scene_path'])
m['acceptance']=qa;m['authored_layout_adjustments']=changed
P.write_text(json.dumps(m,ensure_ascii=False,indent=2),encoding='utf8')
for r,v in old:
 for c in CONTENT_COLUMNS:
  if v[0] in changed and c in [13,20,25]:continue
  assert s.cell(r,c).value==v[c-1]
doc=R/'docs/v0.1/development/2026-10-09_block00_facilities_import.md';s=doc.read_text('utf8');s+='\n## 同日用户反馈调整\n\n沙发实例及关联抱枕、盖毯统一以沙发基点缩放至70%，碰撞随实例同步缩小；公共组件与原始Blend不改。地毯顶面从0.228m降低至0.091m，距现有地砖可视顶部0.083m保留8mm，不新增碰撞。第三间不再生成灯开关，保留常亮照明。第四间画从悬挂改成向后倾22°、侧倾6°、最低网格顶点0.091m的斜靠地面状态。办公室/门厅布局版本v002，稳定路径不变。新增反馈验收后完整专项259项通过，真实渲染241项通过；场景账本与Prefab页定点同步。\n';doc.write_text(s,encoding='utf8')
for rel in ['assets/art/environments/master_office_3d/source/env_father_office/README.md','assets/art/environments/master_office_3d/source/env_block00_story_rooms/README.md']:
 p=R/rel;s=p.read_text('utf8');s+='\n后续状态：已按用户要求完成29组件独立优化导入与4房设施TSCN接入。源Blend不变；正式摆位及同日70%沙发/地毯/拆开关/落地壁画调整以Godot runtime/room_instances为准，详见项目docs/v0.1/development/2026-10-09_block00_facilities_import.md。\n';p.write_text(s,encoding='utf8')
print('ADJUSTMENT_LEDGER_SYNCED',qa['checks'])
