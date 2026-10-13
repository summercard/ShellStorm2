"""Promote the verified v038 motion replacement on the existing character ID."""
from pathlib import Path
from copy import copy
from collections import Counter
from datetime import datetime
import sys, ast, json, hashlib, os, subprocess
import openpyxl

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'scripts'))
sys.path.insert(0,str(ROOT/'tools/asset_pipeline'))
from workspace_paths import work_path
from ledger_registry import LedgerIndex
from split_asset_ledger import read_source_rows,_row_digest,sheet_digest,col_digest,CONTENT_COLUMNS
INDEX=LedgerIndex.load(ROOT)
PATH=INDEX.path_for_category('角色')
BASELINE=ROOT/'assets/registry/ledger_split_baseline.json'
OUT=work_path('sessions/2026-10-14_player_jog_v038/placeholder').parent
BACKUP=ROOT.parent/'ShellStorm2-backup/2026-10-14_player_jog_v038/before_edits'
B=Path('assets/art/characters/player/chr_player_capsule01_3d/variants/bunny01')
ASSET='CHR-PLY-CAPSULE01-3D-BUNNY01'
TRANSFER=B/'character_transfer_ledger.json'
PREFAB=B/'production/v021/runtime/chr_bunny01_root_v021.tscn'
MANIFEST=B/'components/chr_bunny01_motion/asset_manifest.json'
SOURCE=Path('art_source/characters/player/chr_player_capsule01_3d/variants/bunny01/source/animation/chr_bunny01_animation_v038.blend')
tree=ast.parse((ROOT/'scripts/register_bunny_weapon_idles_v025.py').read_text(encoding='utf-8'))
exec(compile(ast.Module(body=[n for n in tree.body if isinstance(n,ast.FunctionDef)],type_ignores=[]),'<ledger-helpers>','exec'))
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()

names=['import','verify_player3d_directional_motion_visual','verify_player3d_firing_motion','verify_player3d_firing_motion_visual','verify_player3d_reload_motion','verify_player3d_reload_motion_visual','verify_player3d_weapon_switch','verify_player3d_animation_flow','verify_player3d_weapon_pose_collision_flow','verify_player3d_diy_flow','verify_player3d_avatar_bounds','verify_player3d_head_accessory_flow','verify_player3d_lower_body_socket_flow','verify_player3d_state_gallery_flow']
results={n:json.loads((OUT/'after'/(n+'.json')).read_text(encoding='utf-8')) for n in names}
assert all(v['exit_code']==0 and not v['unexpected_errors'] for v in results.values())
source_validation=json.loads((OUT/'source_validation.json').read_text())
assert len(source_validation['clips'])==56
comparison=json.loads((OUT/'export_comparison.json').read_text())
assert len(comparison['changed_clips'])==56 and comparison['retained_clips']==35
data=json.loads((ROOT/TRANSFER).read_text(encoding='utf-8'))
assert data['version']=='v038'
assert all(sha(ROOT/e['path'])==e['sha256'] for e in data['files'])
old=json.loads((BACKUP/TRANSFER).read_text(encoding='utf-8'))
data.update(status='active',validation_status='targeted_pass_with_declared_limits',validated_at='2026-10-14',stage_history=['authored','exported_pending_godot_validation','validated','active'],verification=results,source_validation=str(OUT/'source_validation.json'),export_comparison=str(OUT/'export_comparison.json'),preview=str(OUT/'preview/jog.gif'),limits=old.get('limits',[]),expected_errors=[],unexpected_errors=[],not_executed=['Full project suite','Mobile device performance','Full expedition save/resume'],runtime_default_scale=.8,runtime_default_height_m=1.2,weapon_grip_contract=old.get('weapon_grip_contract','authored palm offsets, no runtime IK'))
data['files'].append(dict(path=PREFAB.as_posix(),sha256=sha(ROOT/PREFAB),bytes=(ROOT/PREFAB).stat().st_size))
data['retained_clip_provenance']['rollback_library']=str(BACKUP/B/'components/chr_bunny01_motion/anim_bunny01_library.json')
(ROOT/TRANSFER).write_text(json.dumps(data,ensure_ascii=False,indent=2),encoding='utf-8')
manifest=json.loads((ROOT/MANIFEST).read_text(encoding='utf-8'))
manifest.update(asset_version='v038',assembly_version='v021',status='active',source=SOURCE.as_posix(),transfer=TRANSFER.as_posix(),prefab=PREFAB.as_posix(),clips=data['clips'],files=data['files'],limits=data['limits'])
(ROOT/MANIFEST).write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')

# Main row transaction: no new identity, no formula or validation changes.
before=openpyxl.load_workbook(PATH)
wb=openpyxl.load_workbook(PATH);ws=wb['资产主表']
row=next(r for r,v in read_source_rows(ws) if v[0]==ASSET)
assert ws.cell(row,13).value=='v037','Already registered; inspect before rerunning'
note='；2026-10-14 v038：56条四方向慢走/小跑/移动射击增强头部起伏、前进前探和躯干轴向扭转；双手随动、射击稳枪。35条非目标剪辑保持，下肢/根差异<0.000001；源56条及14项导入/运行/真实渲染检查通过。完整项目与移动端未验，历史头像哈希/命名债务另列。'
changes={9:'10 existing states / 56 refined directional locomotion and firing loops / reload and switch overlays',11:'active',13:'v038',16:str(ws.cell(row,16).value)+'; '+SOURCE.as_posix()+'; '+TRANSFER.as_posix(),20:sha(ROOT/PREFAB),21:'Codex',22:datetime(2026,10,14),25:str(ws.cell(row,25).value)+note}
for c,v in changes.items():ws.cell(row,c).value=v
logrow=append_styled(wb['域变更日志'],['v0.1.23','2026-10-14','四方向移动自然小跑 v038','角色 / Bunny01','56条头部起伏及躯干扭转；三类移动射击稳枪；v021模型、91条总库保持。','源及角色专项通过；全工程/移动端未验。','Codex'],28)
wb.save(PATH);baseline_write(changed_asset=ASSET)

# Specialized sheets are a separate transaction.
wb=openpyxl.load_workbook(PATH);ws=wb['动画与状态'];animation_rows=[]
for state in comparison['changed_clips']:
    matches=[r for r in range(1,ws.max_row+1) if ws.cell(r,2).value==state]
    assert len(matches)==1,(state,matches)
    rr=matches[0];animation_rows.append(rr)
    ws.cell(rr,12).value='v038 active；头部起伏/前进前探/躯干竖轴扭转；移动射击保持枪口朝向。'+SOURCE.as_posix()+'；56循环源检查、运行四方向1881检查及关联专项通过。'
transfer_rows=[]
for label,p in [('v038动作源',SOURCE),('v038正式中转',TRANSFER),('v038运行采样库',B/'components/chr_bunny01_motion/anim_bunny01_library.json'),('v038标准GLB',B/'components/chr_bunny01_motion/anim_bunny01_library.glb')]:
    transfer_rows.append(append_styled(wb['角色中转记录'],[label,p.as_posix(),sha(ROOT/p),(ROOT/p).stat().st_size],149))
ws=wb['3D-角色'];prefab_rows=[]
for rr in range(1,ws.max_row+1):
    if ws.cell(rr,3).value==PREFAB.as_posix():
        prefab_rows.append(rr)
        ws.cell(rr,5).value=SOURCE.as_posix()+'；模型保持v021'
        ws.cell(rr,14).value='active';ws.cell(rr,15).value='v038'
        ws.cell(rr,16).value='模型v021；动作v038，56条四方向自然小跑与移动射击。14项导入/角色专项/真实渲染通过；历史未完项不变。'+TRANSFER.as_posix()
wb.save(PATH);baseline_write(sheets=['动画与状态','角色中转记录','3D-角色'])
after=openpyxl.load_workbook(PATH);diff=[]
for name in before.sheetnames:
    a,b=before[name],after[name]
    assert str(a.data_validations)==str(b.data_validations)
    for cells in b:
        for c in cells:
            if a.cell(c.row,c.column).value!=c.value:
                assert (name=='资产主表' and c.row==row and c.column in changes) or (name=='域变更日志' and c.row==logrow) or (name=='动画与状态' and c.row in animation_rows and c.column==12) or (name=='角色中转记录' and c.row in transfer_rows) or (name=='3D-角色' and c.row in prefab_rows and c.column in [5,14,15,16]),(name,c.coordinate)
                diff.append([name,c.coordinate])
(OUT/'ledger_edit_report.json').write_text(json.dumps(dict(changed_cells=diff,asset_id=ASSET,formula_and_validation_preserved=True),ensure_ascii=False,indent=2),encoding='utf-8')
index_path=ROOT/'assets/registry/current_source_index.json'
index=json.loads(index_path.read_text(encoding='utf-8'))
entry=next(e for e in index['sources'] if e['path'].endswith('/chr_bunny01_animation_v037.blend'))
entry.update(path=SOURCE.as_posix(),bytes=(ROOT/SOURCE).stat().st_size,version=[38],sha256=sha(ROOT/SOURCE))
index['date']='2026-10-14'
index_path.write_text(json.dumps(index,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
print('V038_REGISTERED',len(animation_rows),len(diff))
