"""Promote the verified v039 motion replacement on the existing character ID."""
from pathlib import Path
from copy import copy
from collections import Counter
from datetime import datetime
import sys, ast, json, hashlib, os, subprocess
import openpyxl

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'scripts'))
sys.path.insert(0,str(ROOT/'tools/asset_pipeline'))
from workspace_paths import work_path
from ledger_registry import LedgerIndex
from split_asset_ledger import read_source_rows,_row_digest,sheet_digest,col_digest,CONTENT_COLUMNS
INDEX=LedgerIndex.load(ROOT)
PATH=INDEX.path_for_category('角色')
BASELINE=ROOT/'assets/registry/ledger_split_baseline.json'
OUT=work_path('sessions/2026-10-15_player_motion_v039/placeholder').parent
BACKUP=ROOT.parent/'ShellStorm2-backup/2026-10-15_player_motion_v039/before_edits'
B=Path('assets/art/characters/player/chr_player_capsule01_3d/variants/bunny01')
ASSET='CHR-PLY-CAPSULE01-3D-BUNNY01'
TRANSFER=B/'character_transfer_ledger.json'
PREFAB=B/'production/v021/runtime/chr_bunny01_root_v021.tscn'
MANIFEST=B/'components/chr_bunny01_motion/asset_manifest.json'
SOURCE=Path('art_source/characters/player/chr_player_capsule01_3d/variants/bunny01/source/animation/chr_bunny01_animation_v039.blend')
tree=ast.parse((ROOT/'scripts/register_bunny_weapon_idles_v025.py').read_text(encoding='utf-8'))
exec(compile(ast.Module(body=[n for n in tree.body if isinstance(n,ast.FunctionDef)],type_ignores=[]),'<ledger-helpers>','exec'))
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()

names=['import','verify_player3d_directional_motion_visual','verify_player3d_firing_motion','verify_player3d_firing_motion_visual','verify_player3d_reload_motion','verify_player3d_reload_motion_visual','verify_player3d_weapon_switch','verify_player3d_animation_flow','verify_player3d_weapon_pose_collision_flow','verify_player3d_diy_flow','verify_player3d_avatar_bounds','verify_player3d_head_accessory_flow','verify_player3d_lower_body_socket_flow','verify_player3d_state_gallery_flow','verify_player3d_idle_animation_flow']
results={n:json.loads((OUT/'after'/(n+'.json')).read_text(encoding='utf-8')) for n in names}
assert all(v['exit_code']==0 and not v['unexpected_errors'] for v in results.values())
source_validation=json.loads((OUT/'source_validation.json').read_text())
assert len(source_validation['clips'])==63
comparison=json.loads((OUT/'export_comparison.json').read_text())
assert len(comparison['changed_clips'])==63 and comparison['retained_clips']==28
data=json.loads((ROOT/TRANSFER).read_text(encoding='utf-8'))
assert data['version']=='v039'
assert all(sha(ROOT/e['path'])==e['sha256'] for e in data['files'])
old=json.loads((BACKUP/TRANSFER).read_text(encoding='utf-8'))
data.update(status='active',validation_status='targeted_pass_with_declared_limits',validated_at='2026-10-15',stage_history=['authored','exported_pending_godot_validation','validated','active'],verification=results,source_validation=str(OUT/'source_validation.json'),export_comparison=str(OUT/'export_comparison.json'),preview=str(OUT/'preview/jog.gif'),limits=old.get('limits',[])+['No exact world foot lock or 2D direction blend','Other longguns retain their existing support socket differences','Family charge and heavy melee remain separate unfinished actions'],expected_errors=[],unexpected_errors=[],not_executed=['Full project suite','Mobile device performance','Full expedition save/resume'],runtime_default_scale=.8,runtime_default_height_m=1.2,weapon_grip_contract=old.get('weapon_grip_contract','authored palm offsets, no runtime IK'))
data['files']=[e for e in data['files'] if e['path']!=PREFAB.as_posix()]
data['files'].append(dict(path=PREFAB.as_posix(),sha256=sha(ROOT/PREFAB),bytes=(ROOT/PREFAB).stat().st_size))
data['retained_clip_provenance']['rollback_library']=str(BACKUP/B/'components/chr_bunny01_motion/anim_bunny01_library.json')
(ROOT/TRANSFER).write_text(json.dumps(data,ensure_ascii=False,indent=2),encoding='utf-8')
manifest=json.loads((ROOT/MANIFEST).read_text(encoding='utf-8'))
manifest.update(asset_version='v039',assembly_version='v021',status='active',source=SOURCE.as_posix(),transfer=TRANSFER.as_posix(),prefab=PREFAB.as_posix(),clips=data['clips'],files=data['files'],limits=data['limits'])
(ROOT/MANIFEST).write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')

# Main row transaction: no new identity, no formula or validation changes.
resume='--resume-specialized' in sys.argv
before=openpyxl.load_workbook(BACKUP/PATH.relative_to(ROOT) if resume else PATH)
wb=openpyxl.load_workbook(PATH);ws=wb['资产主表']
row=next(r for r,v in read_source_rows(ws) if v[0]==ASSET)
assert ws.cell(row,13).value in (('v039',) if resume else ('v037','v038')),'Already registered; inspect before rerunning'
note='；2026-10-15 v039：63条四方向慢走/跑步/移动射击及7待机增强头部起伏与躯干扭转；双手随动、射击稳枪。28条非目标剪辑保持，下肢/根差异<0.000001；源63条及15项导入/运行/真实渲染检查通过。完整项目与移动端未验，历史头像哈希/命名债务另列。'
changes={9:'10 existing states / 56 directional loops / 7 expressive idles / reload and switch overlays',11:'active',13:'v039',16:str(ws.cell(row,16).value)+'; '+SOURCE.as_posix()+'; '+TRANSFER.as_posix(),20:sha(ROOT/PREFAB),21:'Codex',22:datetime(2026,10,15),25:str(ws.cell(row,25).value)+note}
if not resume:
    for c,v in changes.items():ws.cell(row,c).value=v
    logrow=append_styled(wb['域变更日志'],['v0.1.24','2026-10-15','跑步与待机加强 v039','角色 / Bunny01','63条动作起伏与扭转加强；三类移动射击稳枪；v021模型、91条总库保持。','源及角色专项通过；全工程/移动端未验。','Codex'],28)
    wb.save(PATH);baseline_write(changed_asset=ASSET)
else:
    logrow=next(c.row for cells in wb['域变更日志'] for c in cells if c.column==3 and c.value=='跑步与待机加强 v039')

# Specialized sheets are a separate transaction.
wb=openpyxl.load_workbook(PATH);ws=wb['动画与状态'];animation_rows=[]
for state in comparison['changed_clips']:
    matches=[r for r in range(1,ws.max_row+1) if ws.cell(r,2).value==state]
    if state=='idle':
        # The same sheet also contains an enemy idle row; only the player owner applies.
        matches=[r for r in matches if '八态StateMachine' in str(ws.cell(r,12).value)]
    assert len(matches)==1,(state,matches)
    rr=matches[0];animation_rows.append(rr)
    ws.cell(rr,12).value='v039 active；头部起伏/前进前探/躯干竖轴扭转；移动射击保持枪口朝向。'+SOURCE.as_posix()+'；63循环源检查及运行方向/待机/关联专项通过。'
    if state=='idle':
        for c,v in {6:'v039：无枪idle源2.4秒；三枪型使用独立sidearm_idle/longgun_idle/machinegun_idle；全局1.25倍率。',7:'头部起伏与延迟跟随，耳部错相摆动；不缩放头部。',8:'双手随源躯干扭转；单手短枪及双手长枪/机枪分别制作。',9:'右掌对齐源WeaponSocket，支撑手同枪体随动；射击使用独立稳枪动作，不启用程序手部覆盖。',10:'脚部支撑保持；身体重心转移与竖轴扭转。'}.items():ws.cell(rr,c).value=v
transfer_rows=[]
for label,p in [('v039动作源',SOURCE),('v039正式中转',TRANSFER),('v039运行采样库',B/'components/chr_bunny01_motion/anim_bunny01_library.json'),('v039标准GLB',B/'components/chr_bunny01_motion/anim_bunny01_library.glb')]:
    transfer_rows.append(append_styled(wb['角色中转记录'],[label,p.as_posix(),sha(ROOT/p),(ROOT/p).stat().st_size],149))
ws=wb['3D-角色'];prefab_rows=[]
for rr in range(1,ws.max_row+1):
    if ws.cell(rr,3).value==PREFAB.as_posix():
        prefab_rows.append(rr)
        ws.cell(rr,5).value=SOURCE.as_posix()+'；模型保持v021'
        ws.cell(rr,14).value='active';ws.cell(rr,15).value='v039'
        ws.cell(rr,16).value='模型v021；动作v039，四方向跑步与7待机加强，保留全局1.25倍率。15项导入/角色专项/真实渲染通过；历史未完项不变。'+TRANSFER.as_posix()
wb.save(PATH);baseline_write(sheets=['动画与状态','角色中转记录','3D-角色'])
after=openpyxl.load_workbook(PATH);diff=[]
for name in before.sheetnames:
    a,b=before[name],after[name]
    assert str(a.data_validations)==str(b.data_validations)
    for cells in b:
        for c in cells:
            if a.cell(c.row,c.column).value!=c.value:
                assert (name=='资产主表' and c.row==row and c.column in changes) or (name=='域变更日志' and c.row==logrow) or (name=='动画与状态' and c.row in animation_rows and (c.column==12 or c.row==6 and c.column in [6,7,8,9,10])) or (name=='角色中转记录' and c.row in transfer_rows) or (name=='3D-角色' and c.row in prefab_rows and c.column in [5,14,15,16]),(name,c.coordinate)
                diff.append([name,c.coordinate])
(OUT/'ledger_edit_report.json').write_text(json.dumps(dict(changed_cells=diff,asset_id=ASSET,formula_and_validation_preserved=True),ensure_ascii=False,indent=2),encoding='utf-8')
index_path=ROOT/'assets/registry/current_source_index.json'
index=json.loads(index_path.read_text(encoding='utf-8'))
entries=[e for e in index['sources'] if e['path'].endswith(('/chr_bunny01_animation_v037.blend','/chr_bunny01_animation_v038.blend'))]
assert entries
entry=entries[-1]
index['sources']=[e for e in index['sources'] if e not in entries[:-1]]
entry.update(path=SOURCE.as_posix(),bytes=(ROOT/SOURCE).stat().st_size,version=[39],sha256=sha(ROOT/SOURCE))
index['date']='2026-10-15'
index_path.write_text(json.dumps(index,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
print('V039_REGISTERED',len(animation_rows),len(diff))
