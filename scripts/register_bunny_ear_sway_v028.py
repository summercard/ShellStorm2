"""Source-only ear refinement registration, scoped two-phase ledger edits."""
from pathlib import Path
from copy import copy
from collections import Counter
import ast,hashlib,json,shutil,subprocess,sys,os
import openpyxl
ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'scripts'));sys.path.insert(0,str(ROOT/'tools/asset_pipeline'))
from ledger_registry import LedgerIndex
from split_asset_ledger import read_source_rows,_row_digest,sheet_digest,col_digest,CONTENT_COLUMNS
INDEX=LedgerIndex.load(ROOT);PATH=INDEX.path_for_category('角色');BASELINE=ROOT/'assets/registry/ledger_split_baseline.json'
OUT=ROOT/'outputs/character_pipeline/directional_v028';BACKUP=ROOT/'_scratch/bunny_ear_sway_v028_ledger_before'
ASSET='CHR-PLY-CAPSULE01-3D-BUNNY01'
SOURCE='assets/art/characters/player/chr_player_capsule01_3d/variants/bunny01/source/animation/chr_bunny01_animation_v028.blend'
TRANSFER=SOURCE.replace('chr_bunny01_animation_v028.blend','chr_bunny01_directional_v028.json')
tree=ast.parse((ROOT/'scripts/register_bunny_weapon_idles_v025.py').read_text(encoding='utf-8'))
exec(compile(ast.Module(body=[n for n in tree.body if isinstance(n,ast.FunctionDef)],type_ignores=[]),'<ledger-helpers>','exec'))
data=json.loads((ROOT/TRANSFER).read_text(encoding='utf-8'));assert len(data['ear_validation'])==18 and data['version']=='v028'
assert hashlib.sha256((ROOT/SOURCE).read_bytes()).hexdigest()==data['sha256']
assert not BACKUP.exists();gates('before');BACKUP.mkdir(parents=True)
shutil.copy2(PATH,BACKUP/PATH.name);shutil.copy2(BASELINE,BACKUP/BASELINE.name)
before=openpyxl.load_workbook(PATH);wb=openpyxl.load_workbook(PATH);ws=wb['资产主表']
r=next(r for r,v in read_source_rows(ws) if v[0]==ASSET)
ws.cell(r,16).value+='; '+SOURCE+'; '+TRANSFER
ws.cell(r,25).value+='；v028为v027的18移动循环补耳朵反方向拖曳/回弹，其余曲线保持，源级authored待导入。'
lr=append_styled(wb['域变更日志'],['v0.1.13','2026-10-07','18条移动循环耳朵跟随','角色 / Bunny01','v028仅补耳朵旋转；保留v027旧动作、持枪和脚步。','源级authored；运行版本和Prefab不变。','Codex'],17)
wb.save(PATH);baseline_write(changed_asset=ASSET)
wb=openpyxl.load_workbook(PATH);ar=[];tr=[];ws=wb['动画与状态']
for state,sp in data['specs'].items():
    rows=[row for row in range(1,ws.max_row+1) if ws.cell(row,2).value==state]
    assert len(rows)==1,(state,rows)
    row=rows[0];ar.append(row)
    ws.cell(row,7).value='v028耳朵反移动方向飘动，错相回弹，耳根固定'
    ws.cell(row,12).value=f"v028 authored；{sp['period']/60:.1f}秒循环；仅耳旋转更新，持枪脚步沿用v027；未导入Godot。{SOURCE}"
for label,p in [('v028耳朵跟随动作母版（authored）',SOURCE),('v028源级中转（待导出）',TRANSFER)]:
    raw=(ROOT/p).read_bytes();tr.append(append_styled(wb['角色中转记录'],[label,p,hashlib.sha256(raw).hexdigest(),len(raw)],110))
wb.save(PATH);baseline_write(sheets=['动画与状态','角色中转记录'])
after=openpyxl.load_workbook(PATH);diff=[]
for name in before.sheetnames:
    a,b=before[name],after[name]
    for row in b:
        for c in row:
            if a.cell(c.row,c.column).value!=c.value:
                assert (name=='资产主表' and c.row==r and c.column in (16,25)) or (name=='动画与状态' and c.row in ar and c.column in (7,12)) or (name=='角色中转记录' and c.row in tr) or (name=='域变更日志' and c.row==lr)
                diff.append([name,c.coordinate])
    assert str(a.data_validations)==str(b.data_validations)
(OUT/'ledger_edit_report.json').write_text(json.dumps(dict(changed_cells=diff,animation_rows=ar,transfer_rows=tr),ensure_ascii=False,indent=2),encoding='utf-8')
gates('after');print('V028_LEDGER_DONE',ar,tr)
