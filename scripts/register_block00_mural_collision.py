"""壁画补碰撞的账本登记（独立事务）：资产主表行 + 3D-设施分页 + 无损基线 + 域变更日志。

只动 `PRP-BLOCK00-CONSTRUCTIVIST-MURAL` 一行；其他资产原内容逐格保持。
运行前置：`scripts/add_block00_mural_collision.py` 已执行，且导入清单哈希与磁盘一致。
"""
from pathlib import Path
import json
import hashlib
import shutil
import sys
from copy import copy
from collections import Counter

R = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(R / 'tools/asset_pipeline'))
from split_asset_ledger import (  # noqa: E402
    CONTENT_COLUMNS,
    FIRST_DATA_ROW,
    _row_digest,
    col_digest,
    read_source_rows,
    sheet_digest,
)
from ledger_registry import LedgerIndex  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

ASSET_ID = 'PRP-BLOCK00-CONSTRUCTIVIST-MURAL'
DATE = '2026-10-12'
O = R / 'outputs/block00_mural_collision_20261012'
O.mkdir(parents=True, exist_ok=True)

M = R / 'art_source/environments/master_office_3d/source/env_block00_story_rooms/export/v001/import_manifest.json'
manifest = json.loads(M.read_text('utf8'))
d = next(x for x in manifest['components'] if x['slug'] == 'constructivist_mural')
sha = lambda p: hashlib.sha256(Path(p).read_bytes()).hexdigest()
prefab_sha = sha(R / d['prefab_path'])
assert prefab_sha == d['prefab_sha256'], '导入清单哈希与磁盘不一致，先跑 add_block00_mural_collision.py'
assert d['collision_policy'] == 'fitted_box'

index = LedgerIndex.load(R)
domain = next(x for x in index.domains if x.key == 'scenes')
backup = O / ('ledger_%s_before.xlsx' % domain.path.stem)
if not backup.exists():
    shutil.copy2(domain.path, backup)

w = load_workbook(domain.path)
sheet = w['资产主表']
rows = read_source_rows(sheet)
hit = [(r, v) for r, v in rows if v[0] == ASSET_ID]
assert len(hit) == 1, hit
row, values = hit[0]
assert 'visual_only' in str(values[13]), ('资产主表已更新，勿重跑', values[13])

spec = str(values[13]).replace('visual_only', 'fitted_box')
sheet.cell(row, 14, spec)
sheet.cell(row, 20, prefab_sha)
sheet.cell(row, 22, DATE)
sheet.cell(row, 25, str(values[24]) + ' 2026-10-12 用户反馈补碰撞：原 visual_only，玩家可穿过斜靠地面的画板；'
                                  '改为 fitted_box，盒体随实例倾斜，底面即实测最低网格点 0.091m。')

facility = w['3D-设施']
fhits = [(r, [facility.cell(r, c).value for c in range(1, facility.max_column + 1)])
         for r in range(5, facility.max_row + 1) if facility.cell(r, 1).value == ASSET_ID]
assert len(fhits) == 1, fhits
frow, fvals = fhits[0]
assert fvals[9] == 'visual_only', ('设施分页已更新，勿重跑', fvals[9])
facility.cell(frow, 8, '开')
facility.cell(frow, 9, 'Prefab自身')
facility.cell(frow, 10, 'fitted_box')
facility.cell(frow, 16, str(fvals[15]) + ' 2026-10-12 补 fitted_box 碰撞（10.3×6.28×0.37m 画板包围盒）。')

log = w['域变更日志']
prev = str(log.cell(log.max_row, 1).value)
import re  # noqa: E402
match = re.match(r'v(\d+)\.(\d+)\.(\d+)', prev)
assert match, prev
version = 'v%s.%s.%d' % (match.group(1), match.group(2), int(match.group(3)) + 1)
log.append([version, DATE, '第四间门厅壁画补碰撞', 'PRP-BLOCK00-CONSTRUCTIVIST-MURAL',
            '用户反馈"98F 第四间掉下来的壁画没有阻挡"；原 visual_only 玩家可穿过，改为 fitted_box',
            '只改 Godot 稳定 Prefab 与清单碰撞字段；Blender 源、GLB、布局与实例姿态不变', '摩斯拉'])
w.save(domain.path)
print('LEDGER_ROW_UPDATED', version, 'row', row, 'facility_row', frow)

# ---- 无损基线 -------------------------------------------------------------
baseline_path = R / 'assets/registry/ledger_split_baseline.json'
baseline = json.loads(baseline_path.read_text('utf8'))
new_rows = dict(read_source_rows(sheet))
baseline['assets'][ASSET_ID] = {'v': _row_digest(new_rows[row]), 'c': new_rows[row][2], 'd': 'scenes'}
baseline['asset_count'] = len(baseline['assets'])
all_rows = []
for dom in index.domains:
    all_rows.extend(read_source_rows(load_workbook(dom.path)['资产主表']))
baseline['column_digests'] = {str(c): col_digest(all_rows, c) for c in CONTENT_COLUMNS}
baseline['category_counts'] = dict(Counter(v[2] for _, v in all_rows))
w2 = load_workbook(domain.path)
baseline['sheet_digests']['3D-设施'] = sheet_digest(w2['3D-设施'])
baseline_path.write_text(json.dumps(baseline, ensure_ascii=False, indent=1) + '\n', encoding='utf8')
print('BASELINE_UPDATED', ASSET_ID, baseline['assets'][ASSET_ID]['v'][:16])
