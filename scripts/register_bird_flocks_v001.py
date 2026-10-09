"""Register only the two authored bird sources, preserving existing ledger content."""
import sys,json,shutil,hashlib,re
from pathlib import Path
from copy import copy
from datetime import datetime
from collections import Counter
from openpyxl import load_workbook
R=Path(__file__).resolve().parents[1]
sys.path[:0]=[str(R/'scripts'),str(R/'tools/asset_pipeline')]
from ledger_registry import LedgerIndex
from split_asset_ledger import FIRST_DATA_ROW,HEADER_ROW,read_source_rows,dedupe_key_formula,dedupe_result_formula,_row_digest,CONTENT_COLUMNS,col_digest,sheet_digest
idx=LedgerIndex.load(R); path=idx.path_for_category('特效'); baseline=R/'assets/registry/ledger_split_baseline.json'
out=R/'outputs/bird_flocks_v001'; out.mkdir(exist_ok=True)
assert not (out/'ledger_before.xlsx').exists(), 'Registration already started; inspect existing backup.'
shutil.copy2(path,out/'ledger_before.xlsx'); shutil.copy2(baseline,out/'baseline_before.json')
w=load_workbook(path); s=w['资产主表']; old=s.max_row; old_rows=read_source_rows(s); special=sheet_digest(w['3D-特效'])
bl=json.loads(baseline.read_text(encoding='utf-8')); old_assets=dict(bl['assets'])
newids=[]
for kind,name,action,seconds in [('flyby','鸟群空中掠过','flyby',12),('ground','鸟群落地停留再起飞','land_walk_peck_takeoff',24)]:
    folder=R/f'assets/art/vfx/environment_3d/bird_flocks/source/{kind}/v001'; m=json.loads((folder/'asset_manifest.json').read_text(encoding='utf-8')); aid=m['asset_id']
    assert aid not in bl['assets']; newids.append(aid); p=R/m['source_blend']; row=s.max_row+1
    values=[aid,name,'特效','environment','env_birds_'+kind,'flock_root',None,'3D',action,'室外场景环境氛围','Blender源已完成','P1','v001',f'7只；{seconds}秒；30fps；每鸟296三角面/11骨骼；独立错峰关键帧；共享场景哑光材质/PaletteUV R10C10冷白；无碰撞',p.relative_to(R).as_posix(),'docs/v0.1/development/2026-10-09_bird_flock_blender_sources.md','鸟群;场景特效;低模;白色;拍翼;落地;啄地',dedupe_key_formula(row),None,hashlib.sha256(p.read_bytes()).hexdigest(),'Codex',datetime(2026,10,9),'原创程序建模与人工动作编排；项目共享材质','02_游戏输出_鸟群资产包','仅Blender源；单次演出不循环；未生成Prefab/未接入Godot；90展示集合不导出']
    for col,val in enumerate(values,1):
        cell=s.cell(row,col,val); cell._style=copy(s.cell(old,col)._style)
    s.row_dimensions[row].height=s.row_dimensions[old].height
last=s.max_row
for row in range(FIRST_DATA_ROW,last+1):s.cell(row,19,dedupe_result_formula(row,last))
for row in w['总览']:
    for cell in row:
        if cell.data_type=='f':cell.value=re.sub(r'(\$[A-Z]+\$)'+str(old)+r'\b',lambda m:m.group(1)+str(last),cell.value)
for dv in s.data_validations.dataValidation:
    ranges=[]
    for item in dv.sqref.ranges:
        item=copy(item)
        if item.max_row==old:item.max_row=last
        ranges.append(str(item))
    dv.sqref=' '.join(ranges)
if s.auto_filter.ref:s.auto_filter.ref=f'A{HEADER_ROW}:Y{last}'
log=w['域变更日志']; lastver=str(log.cell(log.max_row,1).value); parts=lastver.lstrip('v').split('.'); parts[-1]=str(int(parts[-1])+1)
log.append(['v'+'.'.join(parts),datetime(2026,10,9),'新增Blender源资产','特效 / 资产主表','新增两套7鸟环境动画；同步查重、总览、下拉范围和无损基线。','仅源阶段；3D-特效Prefab专表保持不变。','Codex'])
w.save(path)
check=load_workbook(path)
assert sheet_digest(check['3D-特效'])==special
for row,values in old_rows:
    for col in CONTENT_COLUMNS:assert check['资产主表'].cell(row,col).value==values[col-1],(row,col)
for row,values in read_source_rows(check['资产主表']):
    if values[0] in newids:bl['assets'][values[0]]={'v':_row_digest(values),'c':'特效','d':'vfx'}
assert all(bl['assets'][key]==value for key,value in old_assets.items())
union=[]
for domain in idx.domains:union.extend(read_source_rows(load_workbook(domain.path)['资产主表']))
bl['asset_count']=len(bl['assets']); bl['column_digests']={str(c):col_digest(union,c) for c in CONTENT_COLUMNS}; bl['category_counts']=dict(sorted(Counter(str(v[2]).strip() for _,v in union).items()))
baseline.write_text(json.dumps(bl,ensure_ascii=False,indent=1)+'\n',encoding='utf-8')
print('BIRDS_REGISTERED',newids,'existing content preserved; prefab sheet unchanged')
