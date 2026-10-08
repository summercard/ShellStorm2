"""Register the 18 source-only directional loops on the existing character."""
from pathlib import Path
from copy import copy
from collections import Counter
import ast,hashlib,json,shutil,subprocess,sys,os
import openpyxl
ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'scripts'));sys.path.insert(0,str(ROOT/'tools/asset_pipeline'))
from ledger_registry import LedgerIndex
from split_asset_ledger import read_source_rows,_row_digest,sheet_digest,col_digest,CONTENT_COLUMNS
INDEX=LedgerIndex.load(ROOT);PATH=INDEX.path_for_category('角色');BASELINE=ROOT/'assets/registry/ledger_split_baseline.json'
OUT=ROOT/'outputs/character_pipeline/directional_v027';BACKUP=ROOT/'_scratch/bunny_directional_v027_ledger_before'
ASSET='CHR-PLY-CAPSULE01-3D-BUNNY01'
SOURCE='assets/art/characters/player/chr_player_capsule01_3d/variants/bunny01/source/animation/chr_bunny01_animation_v027.blend'
TRANSFER=SOURCE.replace('chr_bunny01_animation_v027.blend','chr_bunny01_directional_v027.json')
tree=ast.parse((ROOT/'scripts/register_bunny_weapon_idles_v025.py').read_text(encoding='utf-8'))
exec(compile(ast.Module(body=[n for n in tree.body if isinstance(n,ast.FunctionDef)],type_ignores=[]),'<ledger-helpers>','exec'))
data=json.loads((ROOT/TRANSFER).read_text(encoding='utf-8'));assert len(data['clips'])==18
assert not BACKUP.exists();gates('before');BACKUP.mkdir(parents=True)
shutil.copy2(PATH,BACKUP/PATH.name);shutil.copy2(BASELINE,BACKUP/BASELINE.name)
before=openpyxl.load_workbook(PATH);wb=openpyxl.load_workbook(PATH);ws=wb['资产主表']
r=next(r for r,v in read_source_rows(ws) if v[0]==ASSET)
ws.cell(r,16).value+='; '+SOURCE+'; '+TRANSFER
ws.cell(r,25).value+='；v027按v026持枪姿态新增三枪型×慢走/正常移动×左移/右移/后退18循环，源级authored，待导入；前进另待制作。'
lr=append_styled(wb['域变更日志'],['v0.1.12','2026-10-07','18条持枪三方向移动循环','角色 / Bunny01','v027，保留v026姿态与旧动作；各方向独立步态，根不移动。','源级authored；运行版本与Prefab不变。','Codex'],17)
wb.save(PATH);baseline_write(changed_asset=ASSET)
wb=openpyxl.load_workbook(PATH);ar=[];tr=[]
for state,s in data['specs'].items():
    ar.append(append_styled(wb['动画与状态'],['表现移动变体（非顶层状态）',state,state,'moving + 类型/速度/局部方向','由玩法状态机决定','重心轻摆，面朝向保持','沿用v026头耳','跟随枪型握持关系','v026主握与枪体朝向保持','不新增','无新增',f"v027 authored；{s['period']/60:.1f}秒循环；未导入Godot。{SOURCE}"],61))
for label,p in [('v027三方向动作母版（authored）',SOURCE),('v027源级中转（待导出）',TRANSFER)]:
    raw=(ROOT/p).read_bytes();tr.append(append_styled(wb['角色中转记录'],[label,p,hashlib.sha256(raw).hexdigest(),len(raw)],108))
wb.save(PATH);baseline_write(sheets=['动画与状态','角色中转记录'])
after=openpyxl.load_workbook(PATH);diff=[]
for name in before.sheetnames:
    a,b=before[name],after[name]
    for row in b:
        for c in row:
            if a.cell(c.row,c.column).value!=c.value:
                assert (name=='资产主表' and c.coordinate in ('P13','Y13')) or (name=='动画与状态' and c.row in ar) or (name=='角色中转记录' and c.row in tr) or (name=='域变更日志' and c.row==lr)
                diff.append([name,c.coordinate])
    assert str(a.data_validations)==str(b.data_validations)
(OUT/'ledger_edit_report.json').write_text(json.dumps(dict(changed_cells=diff,animation_rows=ar,transfer_rows=tr),ensure_ascii=False,indent=2),encoding='utf-8')
gates('after');print('V027_LEDGER_DONE',ar,tr)
