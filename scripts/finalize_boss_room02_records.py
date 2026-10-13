"""Finalize level99 binding ledger, source manifests and current design facts."""
import json,sys,hashlib,re
from pathlib import Path
from copy import copy
from collections import Counter
from openpyxl import load_workbook
R=Path(__file__).resolve().parents[1];sys.path[:0]=[str(R/'scripts'),str(R/'tools/asset_pipeline')]
from ledger_registry import LedgerIndex
from split_asset_ledger import FIRST_DATA_ROW,CONTENT_COLUMNS,read_source_rows,_row_digest,col_digest,sheet_digest,dedupe_key_formula,dedupe_result_formula
from workspace_paths import work_path
W=work_path('outputs/boss_room02_import/final_records.json').parent;D=R/'art_source/environments/tower_zones/expedition/source/room_types/boss_room_B/v027';S=json.loads((W/'spec.json').read_text('utf8'));ix=LedgerIndex.load(R);ledger=ix.path_for_category('场景');baseline=R/'assets/registry/ledger_split_baseline.json';sha=lambda p:hashlib.sha256(p.read_bytes()).hexdigest()
def write(p,d):p.parent.mkdir(parents=True,exist_ok=True);p.write_text(json.dumps(d,ensure_ascii=False,indent=2)+'\n',encoding='utf8')
aid='ENV-EXPEDITION-L99-BOSS-ROOM-02';scene='assets/art/environments/tower_zones/expedition/runtime/room_instances/99/f00_boss_static_layout.tscn';master=next(D.glob('*.blend')).relative_to(R).as_posix();wb=load_workbook(ledger);ws=wb[ix.asset_sheet];ids={e['asset_id'] for e in json.loads((W/'registration.json').read_text('utf8'))['registered_assets']}
oldlast=ws.max_row;old_cells={(c.row,c.column):c.value for row in ws.iter_rows() for c in row}
assert aid not in {str(v[0]) for _,v in read_source_rows(ws)}
usage='远征关卡01 Boss房02独立场景；测试关卡99 boss正式引用同一组件。01原Boss房保留；01新房后续接关卡。'
for row,values in read_source_rows(ws):
 if str(values[0]) in ids:ws.cell(row,10,usage)
values=[aid,'测试关卡99 Boss房02 深井房装配','场景','room','99_boss','room_instance','ENV-EXPEDITION-L01-BOSS-ROOM-02','Top3D / Godot Y-up','scene_art','测试关卡99 boss；共享远征01 Boss房02稳定场景','正式美术已接入','P1','v027','80×100m；中央25×25m，双桥20×5m；46组件652实例',scene,master,'测试99;Boss房02;深井;远征01共享',None,None,sha(R/scene),'Codex','2026-10-13','用户授权替换测试关卡99 Boss房','boss','四房主链不变；组件承重让位、门扇由原RoomDoor3D独占、Boss增援点落平台。']
new=oldlast+1
for col,value in enumerate(values,1):c=ws.cell(new,col,value);c._style=copy(ws.cell(oldlast,col)._style)
ws.cell(new,18,dedupe_key_formula(new))
for row in range(FIRST_DATA_ROW,new+1):ws.cell(row,19,dedupe_result_formula(row,new))
for row in wb['总览'].iter_rows():
 for c in row:
  if isinstance(c.value,str) and '资产主表' in c.value:c.value=re.sub(r'\$([A-Z]+)\$'+str(oldlast)+r'\b',lambda m:'$'+m[1]+'$'+str(new),c.value)
for dv in ws.data_validations.dataValidation:dv.sqref=' '.join(re.sub(r'([A-Z]+)'+str(oldlast)+r'\b',lambda m:m[1]+str(new),str(rng)) for rng in dv.sqref.ranges)
ws.auto_filter.ref=f'A5:Y{new}';log=wb['域变更日志'];ver=str(log.cell(log.max_row,1).value);m=re.search(r'(\d+)$',ver);log.append([ver[:m.start()]+str(int(m[1])+1),'2026-10-13','Codex','测试99 Boss房替换','新增共享场景实例登记；仅前轮47行用途扩展至99','空间适配与原Boss身份/增援/奖励/门策略专项通过'])
assert all(ws.cell(r,c).value==v for (r,c),v in old_cells.items() if c not in (19,) and not (c==10 and str(ws.cell(r,1).value) in ids))
wb.save(ledger);bl=json.loads(baseline.read_text('utf8'))
for row,v in read_source_rows(ws):
 if str(v[0]) in ids|{aid}:bl['assets'][str(v[0])]={'v':_row_digest(v),'c':'场景','d':'scenes'}
union=[]
for domain in ix.domains:union.extend(read_source_rows(load_workbook(domain.path)[ix.asset_sheet]))
bl['asset_count']=len(bl['assets']);bl['column_digests']={str(c):col_digest(union,c) for c in CONTENT_COLUMNS};bl['category_counts']=dict(sorted(Counter(str(v[2]).strip() for _,v in union).items()));write(baseline,bl)
# Separate prefab-page append. Existing approved entries change only their usage column.
wb=load_workbook(ledger);ps=wb['3D-场景通用'];oldps={(c.row,c.column):c.value for row in ps.iter_rows() for c in row};last=ps.max_row
for r in range(5,last+1):
 if str(ps.cell(r,1).value) in ids:ps.cell(r,13,usage)
pv=[aid,'测试关卡99 Boss房02 深井房装配',scene,'无整屋GLB；引用共享组件',master,'测试99的boss静态场景引用01 Boss房02','', '有','component_prefabs','component_owned','80×100m；平台25×25m','房间中心；端口local(-40,-22.5)/(40,-22.5)','测试关卡99 boss','正式美术已接入','v027','房间视觉稳定引用；原首领/增援参数保留，空间挂点适配中央平台。']
for col,v in enumerate(pv,1):c=ps.cell(last+1,col,v);c._style=copy(ps.cell(last,col)._style)
assert all(ps.cell(r,c).value==v for (r,c),v in oldps.items() if not (c==13 and str(ps.cell(r,1).value) in ids))
wb.save(ledger);bl=json.loads(baseline.read_text('utf8'));bl['sheet_digests']['3D-场景通用']=sheet_digest(ps);write(baseline,bl)
runtime=R/'assets/art/environments/tower_zones/expedition/runtime/room_instances'
write(runtime/'99/boss_room_manifest.json',{'schema':'shellstorm2.expedition.room_runtime_manifest.v001','room_id':'boss','level_id':'99','asset_id':aid,'version':'v027','formal_scene':'res://'+scene,'scene_sha256':sha(R/scene),'shared_room_scene':S['runtime_scene'],'shared_asset_id':S['asset_id'],'components':46,'instances':652,'runtime_connected':True,'dimensions_m':[80,100],'support_owner':'component_prefabs','floor_stage_hole':True,'entry_exit_ports':[[-40,0,-22.5],[40,0,-22.5]],'boss_content_id':'boss_monitor002','encounter_owner':'TowerDescent3D existing encounter','test_scene':'res://scripts/probe_boss_room02_level99.tscn','acceptance':['verify_test_level_99_flow','probe_boss_room02_level99'],'source':master})
cat=json.loads((D/'component_catalog.json').read_text('utf8'))
for c in cat['components']:
 d=next(x for x in S['components'] if x['slug']==c['slug']);c.update(runtime_connected=True,registration_status='正式美术已接入',collision_owner=d['collision_owner'],runtime_scene='res://'+d['prefab_path'],runtime_glb='res://'+d['glb_path'],optimized_source=d['optimized_path'],optimized_sha256=d['optimized_sha256'])
 p=D/'component_packages'/c['slug']/'asset_manifest.json';package=json.loads(p.read_text('utf8'));package.update({k:c[k] for k in ['runtime_connected','registration_status','collision_owner','runtime_scene','runtime_glb','optimized_source','optimized_sha256']});write(p,package)
write(D/'component_catalog.json',cat);p=D/'room_type_manifest.json';d=json.loads(p.read_text('utf8'));d.update(source_status='Blender art and 02b imports verified',runtime_connected=True,formal_godot_scene_modified=True,remaining_draft=[],registered_asset_id=S['asset_id'],formal_scene=S['runtime_scene'],level99_scene='res://'+scene,component_count=46,instance_count=652,scope='Complete sealed deep-shaft room, optimized component imports and independent Godot assembly');write(p,d)
# Current facts, with all component AssetIDs covered by the existing design status gate.
p=R/'docs/v0.1/design/远征关卡01设计.md';lines=p.read_text('utf8').splitlines()
for i,line in enumerate(lines):
 if line.startswith('| 6b |'):lines[i]='| 6b | Boss房02 深井计算机房（`boss_room_B`，新增独立变体） | `art_source/environments/tower_zones/expedition/source/room_types/boss_room_B/v027/`；46优化组件与独立TSCN | `ENV-EXPEDITION-L01-BOSS-ROOM-02` v027；共享实例 `ENV-EXPEDITION-L99-BOSS-ROOM-02` | ✅ 正式美术已接入独立场景及测试99；01旧Boss房保留，01新房后续接关卡 | ✅ 独立场景与99已接入；01替换待接关卡 |'
 if line.startswith('深井变体 B（第 `6b` 行）'):lines[i]='Boss房02（第 `6b` 行）使用 v027 深井变体。平台25×25m、双桥各20×5m，围护80×100m。组件GLB和Prefab已导入，正式布局由 `f00_boss_room_02_static_layout.tscn` 持有；独立运行入口为 `scenes/Expedition01BossRoom02_3D.tscn`。01原Boss房01保留，后续接关卡选用02；测试关卡99当前boss已引用02。详见[Boss房02装配契约](远征关卡01_Boss房02装配.md)。'
 if line.startswith('- **深井变体 B 美术现状**'):lines[i]='- **Boss房02美术现状**：46组件、652实例，v027原始源与独立优化源均保留；01独立场景及测试99正式房间引用同一稳定组件。三屏使用用户授权外部贴图，屏幕扫描材质可动画；100踏步灯＋80竖灯独立控制，角色逐灯触发仍待玩法接线。见[装配契约](远征关卡01_Boss房02装配.md)。'
text='\n'.join(lines)+'\n';anchor='> 上表按 **8个房型**';at=text.index(anchor)
table='<details>\n<summary>Boss房02的46个组件登记（共享于01独立房和测试99）</summary>\n\n| 编号 | 组件 | 制作源 | AssetID | 账本制作状态 | 级别 |\n|---|---|---|---|---|---|\n'
for n,d in enumerate(S['components'],1):table+=f'| 6b-{n:02d} | {d["name_zh"]} | v027独立源＋optimized | `{d["asset_id"]}` | ✅ 正式美术已接入 | ✅ 已接入 |\n'
text=text[:at]+table+'\n</details>\n\n'+text[at:];p.write_text(text,encoding='utf8')
contract='''# 远征关卡01 Boss房02装配契约

FeatureID：ASSET-PIPELINE / WORLD-BLOCKS。来源：用户参考图及2026-10-13导入、独立TSCN、测试99替换要求。版本：源v027，游戏0.1.0。

正式布局由Godot TSCN拥有。原Boss房01保留；Boss房02为新增独立变体，01后续接关卡选用02。测试关卡99的运行房间ID仍为boss，当前已引用02。

- 房间资产：`ENV-EXPEDITION-L01-BOSS-ROOM-02`；99实例：`ENV-EXPEDITION-L99-BOSS-ROOM-02`。
- 正式布局：`assets/art/environments/tower_zones/expedition/runtime/room_instances/expedition_01/f00_boss_room_02_static_layout.tscn`。
- 可直接F6运行：`scenes/Expedition01BossRoom02_3D.tscn`，引用正式布局并附验收相机与环境。
- 99绑定：`assets/art/environments/tower_zones/expedition/runtime/room_instances/99/f00_boss_static_layout.tscn`；由既有`TowerDescent3D.expedition_static_layout_path("99","boss")`解析。
- 组件：46个独立GLB和PackedScene，652实例。稳定运行路径无版本号，组件来源与优化清单为`art_source/environments/tower_zones/expedition/source/room_types/boss_room_B/v027/import_manifest.json`。
- Blender房间坐标变换为Godot `(x,z,-(y+22.5))`，足迹中心(0,-22.5,0)、平台走面Z=0。围护80×100m、Z=-36..36；平台25×25m、侧桥各20×5m。Godot侧端口local(-40,0,-22.5)/(40,0,-22.5)，高位入口local(0,12,50)。
- 平台、桥、栏杆、门框、壳体及楼梯由组件Prefab拥有碰撞；楼梯使用连续凸斜坡，门框分段碰撞，纯景观无交互碰撞。99楼层台座在整个房间足迹让出视觉与承重，避免填平深井、共面地板及两套承重。
- 99的Boss/普通敌人刷怪候选取25块平台地板标记；导航由房间TSCN声明实际平台与桥面多边形，运行导航仍由DungeonRoom3D创建。深井空白区域不作导航面。
- 已连接侧门的外门扇视觉挂到同一边唯一RoomDoor3D，内门预览上移打开并关闭独立碰撞。未连接的高位入口保留密封门扇。门权限、清房、钥匙、命运及动作状态仍由旧运行系统负责；模型开口5.02×6.225m仅调整表现包络和门碰撞。
- 三屏授权贴图是公共色盘之外的明确例外；GLB仍不携带images/textures，外部运行屏图最长边1024。扫描shader使用TIME动画，电子画面内容仍为贴图；未来可替换动画贴图/视频输入。
- 100踏步灯和80栏杆竖灯独立Prefab，可调用`set_lit(bool)`，实例带`chase_index`。这里交付控制接口，不包含角色下楼触发实现。
- 失败时停止导入或装配，不回退旧美术伪称完成；原始Blend及外链不变证据、源/优化/GLB哈希由导入清单追溯。

账本由ledger_index.json解析场景域，资产主表与3D-场景通用登记房间和每组件。验收入口：`scripts/verify_boss_room02_runtime.gd`、`scripts/probe_boss_room02_level99.tscn`、`tests/verification/verify_test_level_99_flow.tscn`。交付证据见[开发记录](../development/2026-10-13_boss_room02_godot_import.md)。
'''
(R/'docs/v0.1/design/远征关卡01_Boss房02装配.md').write_text(contract,encoding='utf8')
p=R/'docs/v0.1/design/新关卡设计表_测试关卡99.md';t=p.read_text('utf8');t=t.replace('| 3 | `boss_45x45` | `BOSS_ROOM` | 45 × 45 | north / south / east / west | 0 |','| 3 | `boss_deep_shaft_80x100` | `BOSS_ROOM` | 80 × 100 | west / east，lane −22.5 | 1（Boss房02） |')
t=t.replace('| 1 | `entry` | 入口安全房 | `SAFE_ROOM` | `stair_entry` | | `safe_15x15` | 15 × 15 | 2.5, 2.5 | |','| 1 | `entry` | 入口安全房 | `SAFE_ROOM` | `stair_entry` | | `safe_15x15` | 15 × 15 | −72.5, −62.5 | |').replace('25 × 25 | 2.5, −32.5','25 × 25 | −52.5, −62.5').replace('`boss_45x45` | 45 × 45 | 2.5, −87.5','`boss_deep_shaft_80x100` | 80 × 100 | 0, −40').replace('25 × 25 | 2.5, −132.5','25 × 25 | 52.5, −62.5')
t=t.replace('无。本关全部房间之间都修走廊。','三条主链边使用显式端口直接对接，端点世界坐标重合；Boss房内部的两座20m桥属于Boss房02组件。连接顺序仍为entry→room_01→boss→extraction。')
t=t.replace('Boss 房 45.0 m','Boss 房侧墙100.0 m');start=t.index('**门槽实测结果**');end=t.index('\n---',start)
t=t[:start]+'''**当前门口几何**

| 房间 | 门 | 方位 | 局部门坐标x,z(m) | 所连房间 |
|---|---|---|---|---|
| entry | P1 | east | (7.5,0) | room_01 |
| room_01 | P1 | west | (−12.5,0) | entry |
| room_01 | P2 | east | (12.5,0) | boss |
| boss | P1 | west | (−40,−22.5) | room_01 |
| boss | P2 | east | (40,−22.5) | extraction |
| extraction | P1 | west | (−12.5,0) | boss |

每条边端口相合、朝外相反，不生成额外跨房走廊。门实体仍由父端唯一持有，子端引用同一实体。Boss房的中央平台局部中心(0,−22.5)、尺寸25×25m；Boss出生点与平台中心一致。增援挂点x=±10、z=−32.5/−27.5/−22.5/−17.5/−12.5，高度3m；数量、间隔、激活、上限与原计划一致。旧45×45模板仅保留历史源，不再被本关L1引用。

房间资源与组件所有权见[Boss房02装配契约](远征关卡01_Boss房02装配.md)。深井支撑由Prefab独占，楼层板在80×100足迹让位，导航只落在实际平台与桥面。99环境雾密度0.003，避免原0.04遮住深井远景。
'''+t[end:]
t=t.replace('通过（15 / 25 / 45）','通过（15 / 25 / 80 / 100）').replace('通过（x ∈ −20 ~ 25，y ∈ −145 ~ 10）','通过（x ∈ −80 ~ 65，z ∈ −90 ~ 10）').replace('通过（横向偏移 0 m）','通过（显式端口世界坐标相合）').replace('通过（15 / 20 / 10 m）','不适用（显式端口直接对接，净距0）').replace('由两端门中心距推导','显式端口相合，本关不生成跨房走廊').replace('boss_45x45.json','boss_deep_shaft_80x100.json')
t=t.replace('**校验命令与结果**','**原白盒历史校验命令与结果（Boss房02现行证据见开发记录）**');p.write_text(t,encoding='utf8')
p=R/'docs/v0.1/MODULE_INDEX.md';t=p.read_text('utf8');p.write_text('ASSET-PIPELINE / WORLD-BLOCKS：[远征Boss房02装配](design/远征关卡01_Boss房02装配.md)。v027，46组件/652实例；01原Boss房01保留，02独立TSCN已导入；测试关卡99 boss正式引用02。组件承重、门表现与走行空间按场景契约接线。见[交付记录](development/2026-10-13_boss_room02_godot_import.md)。\n\n'+t,encoding='utf8')
print('FINAL_RECORDS_UPDATED assets48, main_row',new,'prefab_row',last+1)
