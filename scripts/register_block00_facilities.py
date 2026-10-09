"""Two narrow ledger transactions: asset rows, then actual Prefab sheets."""
from pathlib import Path
import sys,json,hashlib,shutil,re
from copy import copy
from collections import Counter
from openpyxl import load_workbook
R=Path(__file__).resolve().parents[1];sys.path.insert(0,str(R/'tools/asset_pipeline'))
from split_asset_ledger import FIRST_DATA_ROW,read_source_rows,_row_digest,col_digest,CONTENT_COLUMNS,dedupe_key_formula,dedupe_result_formula,sheet_digest
from ledger_registry import LedgerIndex
I=LedgerIndex.load(R);D=next(d for d in I.domains if d.key=='scenes')
P=R/'assets/art/environments/master_office_3d/source/env_block00_story_rooms/export/v001/import_manifest.json';m=json.loads(P.read_text('utf8'))
qa=json.loads((R/'outputs/block00_story_rooms_20261009/facility_verification.json').read_text('utf8'));assert qa['passed'] and qa['checks']>=253
sha=lambda p:hashlib.sha256(p.read_bytes()).hexdigest()
assert all(sha(Path(p))==h for p,h in m['protected_sources'].items())
B=R/'outputs/block00_story_rooms_20261009/runtime_ledger_backup';B.mkdir(exist_ok=True)
baseline=R/'assets/registry/ledger_split_baseline.json'
for p in [D.path,baseline]:assert not (B/p.name).exists();shutil.copy2(p,B/p.name)
w=load_workbook(D.path);s=w['资产主表'];old=read_source_rows(s);last=max(r for r,v in old);existing={v[0] for r,v in old}
entries=list(m['components'])
names={'master_office':'父亲办公室','meeting_room':'会议室','corridor':'第三间走廊','lobby':'第四间门厅'}
for d in m['rooms']:
 rid='ENV-BLOCK00-'+d['room_id'].replace('_','-').upper()+'-FACILITY-LAYOUT'
 p=R/d['scene_path'];text=p.read_text('utf8');text=text.replace('metadata/room_id =',f'metadata/asset_id = "{rid}"\nmetadata/asset_version = "v001"\nmetadata/room_id =',1);p.write_text(text,encoding='utf8')
 entries.append({'asset_id':rid,'slug':d['room_id']+'_facility_layout','name_zh':names[d['room_id']]+'设施布局','version':'v001','prefab_path':d['scene_path'],'source_path':'assets/art/environments/master_office_3d/source/env_block00_story_rooms/v001/env_block00_story_rooms_source_v001.blend','bounds_size_m':None,'collision_policy':'referenced_prefabs','glb_path':'','layout':True,'instances':d['instances']})
assert not existing.intersection(d['asset_id'] for d in entries),'Already registered; do not rerun initialization.'
updates={}
for d in entries:
 last+=1;cat='场景' if d['asset_id'].startswith('ENV-') else '场景道具';sub='room' if d.get('layout') else 'wall' if cat=='场景' else 'facility'
 notes='98F已接入；逐件优化/源SHA保护/导入/保真/实际房间验收见 '+str(P.relative_to(R)).replace('\\','/')
 values=[d['asset_id'],'98F '+d['name_zh'],cat,sub,'block00_'+d['slug'],'root',None,'Top3D','ruined','98F 区块00四房；可复用组件','正式美术已接入','P1',d['version'],str(d.get('bounds_size_m') or str(d.get('instances'))+'设施实例')+'；'+d['collision_policy'],d['prefab_path'],'src/world3d/Block00MasterOfficeLayout3D.gd','98F;设施;废墟;'+d['slug'],None,None,sha(R/d['prefab_path']),'Codex','2026-10-09','用户参考风格与文字要求；项目原创建模','master_office',notes]
 for c,value in enumerate(values,1):s.cell(last,c,value);s.cell(last,c)._style=copy(s.cell(FIRST_DATA_ROW,c)._style)
 updates[d['asset_id']]=last
for r,v in old:
 if v[0] in ['ENV-BATTLE-FATHER-OFFICE-SOURCE','ENV-BATTLE-BLOCK00-STORY-ROOMS-SOURCE']:
  s.cell(r,25,'原始Blender源保持不变。29组件经独立optimized源导出并包装，328设施实例接入98F四房；普通墙地砖沿用原有组件，破墙单独替换北墙视觉；三四房拆门贯通。优化与验收清单：'+str(P.relative_to(R)).replace('\\','/'));updates[v[0]]=r
for r in range(FIRST_DATA_ROW,last+1):s.cell(r,18,dedupe_key_formula(r));s.cell(r,19,dedupe_result_formula(r,last))
oldlast=max(r for r,v in old)
for row in w['总览']:
 for cell in row:
  if isinstance(cell.value,str) and cell.value.startswith('='):cell.value=re.sub(r'(\$[A-Z]+\$)'+str(oldlast)+r'\b',lambda match:match[1]+str(last),cell.value)
for dv in s.data_validations.dataValidation:
 ranges=[str(x) for x in dv.sqref.ranges];dv.sqref=' '.join(re.sub(r'(?<=[A-Z])'+str(oldlast)+r'$',str(last),x) for x in ranges)
s.auto_filter.ref=f'A5:Y{last}'
def log(message):
 log=w['域变更日志'];prev=str(log.cell(log.max_row,1).value);match=re.match(r'v(\d+)\.(\d+)\.(\d+)',prev);ver=f'v{match[1]}.{match[2]}.{int(match[3])+1}' if match else 'v0.1.1';log.append([ver,'2026-10-09',message,'98F 区块00设施',','.join(updates),'只修改本批资产与专表；源文件不变','Codex'])
log('29组件及4正式房间设施布局接入登记');w.save(D.path)
bl=json.loads(baseline.read_text('utf8'));newrows=dict(read_source_rows(s))
for aid,r in updates.items():bl['assets'][aid]={'v':_row_digest(newrows[r]),'c':newrows[r][2],'d':'scenes'}
bl['asset_count']=len(bl['assets']);allrows=[]
for domain in I.domains:allrows.extend(read_source_rows(load_workbook(domain.path)['资产主表']))
bl['column_digests']={str(c):col_digest(allrows,c) for c in CONTENT_COLUMNS};bl['category_counts']=dict(Counter(v[2] for _,v in allrows));baseline.write_text(json.dumps(bl,ensure_ascii=False,indent=1)+'\n',encoding='utf8')
# Separate transaction: every Prefab exists; update only the two intentional sheet digests.
w=load_workbook(D.path)
for d in entries:
 sheet=w['3D-场景通用' if d['asset_id'].startswith('ENV-') else '3D-设施'];n=sheet.max_row+1
 vals=[d['asset_id'],'98F '+d['name_zh'],d['prefab_path'],d['glb_path'] or None,d['source_path'],'纯美术组件/布局；不附加虚假交互',None,'开' if d['collision_policy'] in ['fitted_box','pedestal_box'] else '无','Prefab自身' if d['collision_policy'] in ['fitted_box','pedestal_box'] else '原房间边界' if d['slug']=='wall_fractured' else '无',d['collision_policy'],str(d.get('bounds_size_m') or '房间局部米制布局'),'底面中心；Blender XY/Z → Godot XZ/Y；根scale=1','98F '+str([r['room_id'] for r in m['rooms']]),'正式美术已接入',d['version'],'优化/保真/哈希与QA：'+str(P.relative_to(R)).replace('\\','/')]
 for c,val in enumerate(vals,1):sheet.cell(n,c,val);sheet.cell(n,c)._style=copy(sheet.cell(n-1,c)._style)
 sheet.auto_filter.ref=f'A4:P{n}'
log('实际Prefab专表登记（独立事务）');w.save(D.path)
for name in ['3D-场景通用','3D-设施']:bl['sheet_digests'][name]=sheet_digest(w[name])
baseline.write_text(json.dumps(bl,ensure_ascii=False,indent=1)+'\n',encoding='utf8')
check=load_workbook(D.path)['资产主表']
for r,values in old:
 for c in CONTENT_COLUMNS:
  if values[0] in updates and c==25:continue
  assert check.cell(r,c).value==values[c-1],(r,c)
m['runtime_integrated']=True;m['acceptance']=qa;m['registered_assets']=[d['asset_id'] for d in entries]
for d in m['components']:d['prefab_sha256']=sha(R/d['prefab_path'])
for d in m['rooms']:d['scene_sha256']=sha(R/d['scene_path'])
P.write_text(json.dumps(m,ensure_ascii=False,indent=2),encoding='utf8')
print('BLOCK00_LEDGER_REGISTERED',len(entries),'preserved_existing_rows',len(old))
