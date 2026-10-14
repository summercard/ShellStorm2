#!/usr/bin/env python3
"""Materialize content-ID navigation ledgers; never write asset rows or gameplay data."""
from __future__ import annotations

import argparse
import csv
import json
import math
import os
import re
import sys
import unicodedata
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / 'scripts'))
from ledger_registry import LedgerIndex

INDEX = ROOT / 'assets/registry/ledger_index.json'
REGISTRY = 'src/base/ItemRegistry.gd'
CONTENT = 'docs/v0.1/data/ShellStorm2_游戏内容数据库_v010.xlsx'
DROP_CSV = 'docs/v0.1/data/怪物掉落表_地图x怪物.csv'
PLAN = 'data/level_plans/expedition_01/level_plan.json'
WEAPON_SHEET = '背包物品调用索引'
HEADERS = ['内容ID', '名称', '种类', '子类', '唯一登记账本', '内容定义入口', '内容数据库定位',
           '怪物掉落来源', '池声明入口', '背包接收入口', '使用/装配入口', '模型调用入口',
           '当前模型表现', '关联AssetID（引用）', '资产登记定位（引用）', '功能模块', '审查备注']


def definitions():
    """Read each literal ItemRegistry dictionary without executing GDScript."""
    source = (ROOT / REGISTRY).read_text()
    result = {}
    for match in re.finditer(r'"id"\s*:\s*"([^"]+)"', source):
        item_id = match[1]
        start = source.rfind('{', 0, match.start())
        depth, quoted, escaped = 0, False, False
        end = None
        for pos in range(start, len(source)):
            ch = source[pos]
            if quoted:
                if escaped:
                    escaped = False
                elif ch == '\\':
                    escaped = True
                elif ch == '"':
                    quoted = False
            elif ch == '"':
                quoted = True
            elif ch == '{':
                depth += 1
            elif ch == '}':
                depth -= 1
                if depth == 0:
                    end = pos + 1
                    break
        if end is None or item_id in result:
            raise ValueError(f'Invalid/duplicate literal: {item_id}')
        block = source[start:end]
        fields = dict(re.findall(r'"(name|type|subtype|use_action|assembly_id)"\s*:\s*"([^"]*)"', block))
        if not fields.get('type') or not fields.get('name'):
            raise ValueError(f'Incomplete item: {item_id}')
        fields['line'] = source[:match.start()].count('\n') + 1
        pools = re.search(r'"floor_loot_weights"\s*:\s*\{([^}]*)\}', block)
        fields['pools'] = re.findall(r'"([^"]+)"\s*:', pools[1]) if pools else []
        result[item_id] = fields
    return result


def weapon_owned(item_id, item):
    return item['type'] in ('weapon', 'attachment') or item.get('subtype') == 'bullet'


def authored_model_paths():
    """Read the actual centralized preloads so regenerated navigation stays current."""
    result = {}
    for rel, names in [('src/world3d/ItemModelFactory3D.gd', ['GROUND_SCENES']),
                       ('src/combat3d/WeaponModel3D.gd', ['GUN_VISUAL_SCENES', 'MELEE_VISUAL_SCENES'])]:
        source = (ROOT / rel).read_text(encoding='utf8')
        for name in names:
            block = re.search(r'const\s+' + name + r'\s*:=\s*\{(.*?)\n\}', source, re.S)
            if block:
                result.update(re.findall(r'"([^"]+)"\s*:\s*preload\("res://([^"]+)"\)', block[1]))
    return result


def link(path, location=''):
    return path + ('#' + location if location else '')


def expected():
    index = LedgerIndex.load(ROOT)
    inventory = index.raw['content_ledgers']['inventory_items']
    items = definitions()
    book = openpyxl.load_workbook(ROOT / CONTENT, read_only=True)
    content_rows = {str(row[0]): n for n, row in enumerate(book['掉落物品'].iter_rows(min_row=5, values_only=True), 5) if row[0]}
    book.close()
    csv_rows = list(csv.DictReader((ROOT / DROP_CSV).open(encoding='utf-8-sig')))
    plan = json.loads((ROOT / PLAN).read_text())['monster_drop_table']['rows']
    if [(r['怪物ID'], r['物品ID']) for r in csv_rows] != [(r['monster_id'], r['item_id']) for r in plan]:
        raise ValueError('CSV/runtime plan item sequence differs')
    asset_refs = {}
    asset_id_refs = {}
    authored = authored_model_paths()
    for domain in index.domains:
        wb = openpyxl.load_workbook(domain.path, read_only=True)
        for n, row in enumerate(wb[index.asset_sheet].iter_rows(min_row=6, values_only=True), 6):
            if row[0] and row[4]:
                asset_refs.setdefault(str(row[4]), []).append((str(row[0]), link(domain.relative_path, f'资产主表!A{n}')))
                asset_id_refs[str(row[0])] = (str(row[0]), link(domain.relative_path, f'资产主表!A{n}'))
        wb.close()
    rows = {'物品道具': [], WEAPON_SHEET: [], '掉落归属导航': [], '非背包与待设计': []}
    for item_id, item in items.items():
        weapon = weapon_owned(item_id, item)
        owner = index.domain_for_key('weapons').relative_path if weapon else inventory['path']
        sheet = WEAPON_SHEET if weapon else '物品道具'
        action = item.get('use_action', '')
        if item['type'] == 'weapon':
            use = link('src/world3d/Dungeon3D.gd', '_equip_weapon_from_inventory')
            model = 'weapon；真实武器快照重建'
        elif item['type'] == 'attachment' or item.get('subtype') == 'bullet':
            use = link('src/world3d/Dungeon3D.gd', '_install_attachment_from_inventory' if item['type'] == 'attachment' else '_install_weapon_module_from_item')
            model = 'attachment' if item['type'] == 'attachment' else 'bullet'
        elif item.get('subtype') == 'backpack':
            use = link('src/world3d/Dungeon3D.gd', '_equip_backpack_from_inventory')
            model = 'backpack；专用PackedScene，账本仍待制作'
        elif item_id == 'item_beacon':
            use = link('src/world3d/Dungeon3D.gd', '_on_inventory_item_clicked')
            model = 'beacon；程序几何'
        elif item_id == 'item_room_key':
            use = link('src/world3d/Dungeon3D.gd', '_consume_room_key')
            model = 'key；程序几何'
        elif item_id == 'item_ammo_pack':
            use = link('src/world3d/Dungeon3D.gd', '_on_inventory_item_clicked')
            model = 'ammo；程序罐体'
        else:
            use = link('src/game/ItemUseHandler.gd', f'apply / {action}')
            model = {'heal': 'heal；程序罐体', 'refill_ammo': 'ammo；程序罐体',
                     'unlock_blueprint': 'blueprint；程序几何'}.get(action, 'crate；通用箱体，尚无专用外观')
        refs = [] if item['type'] == 'blueprint' else asset_refs.get(item_id, asset_refs.get(item.get('assembly_id', ''), []))
        model_key = '__blueprint__' if item['type'] == 'blueprint' else item_id
        authored_path = authored.get(model_key, authored.get(item.get('assembly_id', '')))
        if authored_path:
            prefab = (ROOT / authored_path).read_text(encoding='utf8')
            identity = re.search(r'metadata/asset_id\s*=\s*"([^"]+)"', prefab)
            if identity and identity[1] in asset_id_refs:
                refs = [asset_id_refs[identity[1]]]
            model = '已接入3D资产；地面/图标共用：' + authored_path
        occurrences = [f'{DROP_CSV}#CSV行{n} / {r["怪物ID"]}' for n, r in enumerate(csv_rows, 2) if r['物品ID'] == item_id]
        note = '已注册；本表只导航，堆叠/数值/掉落概率由原定义拥有'
        if not occurrences:
            note += '；远征01作者掉落表未列此项（不可视为无其他产出）'
        if item_id not in content_rows:
            note += '；内容数据库缺行，运行注册表有此ID'
        if not refs:
            note += '；无对应资产主表登记，不新增AssetID'
        if item_id in ('item_battery_s', 'item_battery_l', 'item_cell_pack') and not authored_path:
            note += '；电池使用动作存在，模型仍落入crate兜底'
        if item.get('subtype') == 'flashlight_module':
            note += '；使用需设施上下文，入包不等于可在战局切换'
        pools = item['pools']
        row = [item_id, item['name'], item['type'], item.get('subtype', ''), owner,
               link(REGISTRY, f'L{item["line"]}'), link(CONTENT, f'掉落物品!A{content_rows[item_id]}') if item_id in content_rows else '缺行',
               '\n'.join(occurrences) or '未在当前CSV列出；池来源见下一列',
               link(REGISTRY, f'L{item["line"]}') + ' / ' + (', '.join(pools) or '无池声明'),
               link('src/game/InventoryModule.gd', 'add_item'), use,
               link('src/world3d/ItemModelFactory3D.gd', 'create_model / get_model_kind'), model,
               '\n'.join(r[0] for r in refs) or '未登记', '\n'.join(r[1] for r in refs) or '未登记',
               'WORLD-LOOT / REWARD-SERVICE / INVENTORY-SLOTS' + (' / WEAPON-OWNERSHIP' if weapon else ''), note]
        rows[sheet].append(row)
    locations = {row[0]: (inventory['path'], '物品道具', n) for n, row in enumerate(rows['物品道具'], 5)}
    locations.update({row[0]: (index.domain_for_key('weapons').relative_path, WEAPON_SHEET, n) for n, row in enumerate(rows[WEAPON_SHEET], 5)})
    for n, row in enumerate(csv_rows, 2):
        item_id = row['物品ID']
        if item_id.startswith('@currency:'):
            target = link('src/world3d/Dungeon3D.gd', '_grant_run_currency')
            owner = '货币；不进入背包'
        else:
            if item_id not in locations:
                raise ValueError(f'Unrouted drop: {item_id}')
            path, sheet, dest = locations[item_id]
            target, owner = link(path, f'{sheet}!A{dest}'), '武器账本' if sheet == WEAPON_SHEET else '物品道具账本'
        rows['掉落归属导航'].append([link(DROP_CSV, f'CSV行{n}'), row['关卡ID'], row['怪物ID'], item_id, row['物品'], owner, target,
                                     link(PLAN, f'monster_drop_table.rows[{n-2}]'),
                                     link('src/rewards/MonsterDropTable.gd', 'compile'),
                                     link('src/rewards/RuntimeRewardCoordinator.gd', 'configure_level_drop_table / resolve_kill')])
    for item_id, n in content_rows.items():
        if item_id not in items:
            rows['非背包与待设计'].append([item_id, link(CONTENT, f'掉落物品!A{n}'), '内容表设计条目；ItemRegistry未注册；不可列为当前可入包物品'])
    rows['非背包与待设计'].append(['@currency:extraction_points', link('src/world3d/Dungeon3D.gd', '_grant_run_currency'), '魂直接入货币余额，不占背包格'])
    rows['功能调用链'] = [
        ['内容定义', 'ItemRegistry', link(REGISTRY, 'get_item'), '物品字段/堆叠/使用动作与池成员的唯一运行定义'],
        ['掉落编译', 'MonsterDropTable', link('src/rewards/MonsterDropTable.gd', 'compile'), '关卡掉落配置校验；整表失败拒绝部分生效'],
        ['击杀调度', 'RuntimeRewardCoordinator', link('src/rewards/RuntimeRewardCoordinator.gd', 'resolve_kill'), '关卡作者表优先；未配置怪/精英/Boss按RewardService既有路由'],
        ['通用池', 'MonsterInjector', link('src/map/MonsterInjector.gd', 'drop_spec_for'), '备弹使用item_ammo_pack；不得为本账本建立新掉落ID'],
        ['池成员', 'RewardPoolRegistry', link('src/rewards/RewardPoolRegistry.gd', 'declared_chance'), '成员来自ItemRegistry；废弃combat/spawn池不可当作可达性证明'],
        ['地面发放', 'RewardSink', link('src/rewards/RewardSink.gd', 'apply_ground'), '发放口确认生成结果；不持有背包规则'],
        ['地面实例', 'Dungeon3D', link('src/world3d/Dungeon3D.gd', '_spawn_loot_items'), '按既有房间边界、支撑面和散布规则落地'],
        ['拾取申请', 'GroundLootPickup3D', link('src/world3d/GroundLootPickup3D.gd', 'accept_pickup'), '事务确认后播放收取动画；模型读取ItemModelFactory3D'],
        ['拾取事务', 'Dungeon3D', link('src/world3d/Dungeon3D.gd', '_on_ground_loot_requested'), '魂走货币；物品走add_item，满包保留、部分接收保留余量'],
        ['背包事实', 'InventoryModule', link('src/game/InventoryModule.gd', 'add_item'), '唯一格位所有者；武器实例去重，其他物品按stack_max堆叠'],
        ['使用入口', 'Dungeon3D', link('src/world3d/Dungeon3D.gd', '_on_inventory_item_clicked'), '先分流换弹/武器/背包/手电/配件/子弹；普通use_action进入ItemUseHandler'],
        ['效果处理', 'ItemUseHandler', link('src/game/ItemUseHandler.gd', 'apply'), '药水heal；电池restore_flashlight_charge；拒绝效果时不扣物品'],
        ['共用模型', 'ItemModelFactory3D', link('src/world3d/ItemModelFactory3D.gd', 'create_model'), '地面与背包图标共用工厂；非专用物品落入crate'],
        ['背包图标', 'ItemModelIcon3D', link('src/ui/ItemModelIcon3D.gd', 'configure'), '图标读取同一物品字典与模型工厂'],
    ]
    return index, inventory, rows


def format_sheet(ws, title, description, headers, rows, table_name):
    ws.append([title]); ws.append([description]); ws.append([]); ws.append(headers)
    for row in rows:
        ws.append(row)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=min(4, len(headers)))
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=min(4, len(headers)))
    ws.row_dimensions[1].height = 30; ws.row_dimensions[2].height = 44
    ws.freeze_panes = 'C5'
    ws.sheet_view.showGridLines = False
    for row in ws:
        for cell in row:
            cell.alignment = Alignment(vertical='top', wrap_text=True)
            cell.font = Font(name='Microsoft YaHei', size=11)
    ws['A1'].font = Font(name='Microsoft YaHei', size=14, bold=True, color='243B53')
    ws.row_dimensions[4].height = 32
    for rownum in (4,):
        for cell in ws[rownum]:
            cell.fill = PatternFill('solid', fgColor='243B53')
            cell.font = Font(name='Microsoft YaHei', size=12, bold=True, color='FFFFFF')
    for col in range(1, len(headers)+1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 28 if col <= 4 else 46
    for n in range(5, ws.max_row+1):
        lines = 1
        for cell in ws[n]:
            width = ws.column_dimensions[cell.column_letter].width - 2
            height = sum(max(1, math.ceil(sum(2 if unicodedata.east_asian_width(ch) in 'WF' else 1 for ch in line) / width))
                         for line in str(cell.value or '').split('\n'))
            lines = max(lines, height)
        ws.row_dimensions[n].height = max(42, lines * 17 + 8)
        for cell in ws[n]:
            if isinstance(cell.value, str) and '\n' not in cell.value and ('.gd#' in cell.value or '.xlsx#' in cell.value):
                path, location = cell.value.split('#', 1)
                cell.hyperlink = os.path.relpath(ROOT / path, ROOT / 'assets/registry/ledgers') + ('#'+location if '.xlsx' in path else '')
                cell.font = Font(name='Microsoft YaHei', size=11, color='1565C0', underline='single')
    table = Table(displayName=table_name, ref=f'A4:{openpyxl.utils.get_column_letter(len(headers))}{ws.max_row}')
    table.tableStyleInfo = TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)
    ws.add_table(table)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--write', action='store_true', help='Create navigation workbook and weapon-owned index; use only after backing up existing files')
    args = parser.parse_args()
    index, inventory, rows = expected()
    target = ROOT / inventory['path']
    if args.write:
        wb = openpyxl.Workbook(); wb.remove(wb.active)
        format_sheet(wb.create_sheet('物品道具'), '物品道具账本', '可掉落到地面、拾取进入背包的物品；内容ID唯一。子弹模块、配件及成品武器归武器账本；通用弹药归道具。', HEADERS, rows['物品道具'], 'InventoryItems')
        format_sheet(wb.create_sheet('掉落归属导航'), '怪物掉落归属导航', '每行对应原CSV一行，仅指向登记所有者与运行投影。概率、数量与权重仍以原作者表及运行配置为准。',
                     ['作者表行', '关卡ID', '怪物ID', '内容ID', '名称', '归属', '登记位置', '运行投影', '编译入口', '调度入口'], rows['掉落归属导航'], 'DropNavigation')
        format_sheet(wb.create_sheet('非背包与待设计'), '非背包产出与待设计', '货币与未注册设计条目显式排除，不冒充当前可入包物品。',
                     ['内容ID', '定义/消费入口', '审查结论'], rows['非背包与待设计'], 'ExcludedItems')
        format_sheet(wb.create_sheet('功能调用链'), '物品功能调用链', '账本仅导航；运行时不读取此XLSX，不另设物品、掉落或背包状态所有者。',
                     ['阶段', '模块/所有者', '代码入口', '契约与失败处理'], rows['功能调用链'], 'ItemCallChain')
        target.parent.mkdir(parents=True, exist_ok=True); wb.save(target)
        weapon_path = index.domain_for_key('weapons').path
        wb = openpyxl.load_workbook(weapon_path)
        if WEAPON_SHEET in wb:
            del wb[WEAPON_SHEET]
        format_sheet(wb.create_sheet(WEAPON_SHEET), '武器账本 · 背包物品调用索引', '成品武器、子弹模块、枪械配件唯一归属本账本；内容ID导航页，不新增资产主表行。', HEADERS, rows[WEAPON_SHEET], 'WeaponInventoryItems')
        wb.save(weapon_path)
    # Compare saved source-derived values, uniqueness, ownership and unchanged source contracts.
    for path, sheets in [(target, ['物品道具', '掉落归属导航', '非背包与待设计', '功能调用链']), (index.domain_for_key('weapons').path, [WEAPON_SHEET])]:
        wb = openpyxl.load_workbook(path, read_only=True)
        if path == target and '资产主表' in wb.sheetnames:
            raise ValueError('Inventory workbook must not contain 资产主表')
        for sheet in sheets:
            saved = [list(row) for row in wb[sheet].iter_rows(min_row=5, values_only=True)]
            normalized = lambda rs: [[v if v is not None else '' for v in r] for r in rs]
            if normalized(saved) != normalized(rows[sheet]):
                raise ValueError(f'Stale navigation: {path.name}#{sheet}; review source changes before --write')
        wb.close()
    item_ids = {r[0] for r in rows['物品道具']}
    weapon_ids = {r[0] for r in rows[WEAPON_SHEET]}
    assert not item_ids & weapon_ids
    assert item_ids | weapon_ids == set(definitions())
    for row in rows['功能调用链']:
        path, symbol = row[2].split('#', 1)
        assert re.search(r'func\s+' + re.escape(symbol) + r'\s*\(', (ROOT / path).read_text()), row[2]
    for row in rows['物品道具'] + rows[WEAPON_SHEET]:
        assert (ROOT / row[5].split('#')[0]).is_file()
        assert (ROOT / row[9].split('#')[0]).is_file()
        use_path, symbol = row[10].split('#', 1)
        use_source = (ROOT / use_path).read_text()
        assert re.search(r'func\s+' + re.escape(symbol.split(' / ')[0]) + r'\s*\(', use_source), row[10]
        assert (ROOT / row[11].split('#')[0]).is_file()
    print(f'INVENTORY_ITEM_LEDGER_OK items={len(item_ids)} weapons={len(weapon_ids)} drop_rows={len(rows["掉落归属导航"])} excluded={len(rows["非背包与待设计"])}')


if __name__ == '__main__':
    main()
