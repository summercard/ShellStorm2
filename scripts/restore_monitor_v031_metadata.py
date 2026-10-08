"""Restore the active Boss002 metadata and ledger rows to the verified v031 release."""
import hashlib
import json
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path("I:/工作项目/shellstrom2/ShellStorm2")
BOSS = ROOT / "assets/art/enemies/bosses/enm_boss_monitor002"
ASSET_ID = "ENM-BOSS-MONITOR002-3D"

v031_ledger_path = BOSS / "enm_boss_monitor002_transfer_ledger_v031.json"
v031_ledger = json.loads(v031_ledger_path.read_text(encoding="utf-8"))

# The active transfer record must describe the actual active runtime asset.
(BOSS / "enm_boss_monitor002_transfer_ledger.json").write_text(
    json.dumps(v031_ledger, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
)

# Keep the manifest's historical hashes, but repoint all active fields and active output hashes.
manifest_path = BOSS / "asset_manifest.json"
manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
manifest["version"] = "v031"
manifest["stage"] = "active"
manifest["source"] = "source/enm_boss_monitor002_model_v031.blend"
manifest["model_source"] = "assets/art/enemies/bosses/enm_boss_monitor002/source/enm_boss_monitor002_model_v031.blend"
manifest["animation_source"] = "assets/art/enemies/bosses/enm_boss_monitor002/source/enm_boss_monitor002_animation_v031.blend"
manifest["rig_contract"] = "source/rig_contract_v031.json"
manifest["runtime_verification"] = v031_ledger["verification"]
for item in v031_ledger["files"]:
    path = item["path"]
    manifest.setdefault("sha256", {})[path] = item["sha256"]
    if path.startswith(str(BOSS.relative_to(ROOT)).replace("\\", "/") + "/"):
        manifest.setdefault("files", {})[path.split("enm_boss_monitor002/", 1)[1]] = item["sha256"]
for path, digest in v031_ledger["source_sha256"].items():
    manifest.setdefault("sha256", {})[path] = digest
    manifest.setdefault("files", {})[path.split("enm_boss_monitor002/", 1)[1]] = digest
manifest_path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

# Existing Boss rows stay in place.  Only their current-version/source/evidence cells return to v031.
enemy_book = ROOT / "assets/registry/ledgers/ShellStorm2_敌人账本_v001.xlsx"
workbook = load_workbook(enemy_book)
sources = "; ".join([
    "assets/art/enemies/bosses/enm_boss_monitor002/source/enm_boss_monitor002_model_v031.blend",
    "assets/art/enemies/bosses/enm_boss_monitor002/source/enm_boss_monitor002_animation_v031.blend",
])
note = "恢复v031老版走路；底座23°侧翘、原横移与手臂节奏。其余15剪辑、四技能、三阶段与远征01投放不变。"

master = workbook["资产主表"]
assert master["A20"].value == ASSET_ID
master["M20"] = "v031"
master["P20"] = sources + "；SKEL-MONITOR002-005"
master["V20"] = "2026-10-08"
master["Y20"] = note

enemy3d = workbook["3D-敌人"]
assert enemy3d["A10"].value == ASSET_ID
enemy3d["E10"] = sources
enemy3d["O10"] = "v031"
enemy3d["P10"] = note

states = workbook["敌人动画与状态"]
for row in range(1, states.max_row + 1):
    if states.cell(row, 1).value == ASSET_ID:
        states.cell(row, 5).value = sources
        states.cell(row, 9).value = note

log = workbook["域变更日志"]
log.append(["v031", "2026-10-08", "Codex", ASSET_ID, note,
            "docs/v0.1/development/2026-10-08_boss002_restore_v031.md", "用户要求回退，保留v032/v033历史源文件"])
workbook.save(enemy_book)

production_book = BOSS / "source/boss002_production_ledger.xlsx"
production = load_workbook(production_book)
sheet = production[production.sheetnames[0]]
sheet["B4"] = "v031 正式16剪辑/十二态/四技能/三阶段"
sheet["B20"] = "v031"
sheet["C20"] = sources
sheet["C21"] = "source/rig_contract_v031.json"
sheet["B26"] = "previews/runtime/flow_report.json; previews/runtime/visual_report.json"
sheet["C26"] = note
production.save(production_book)

print("RESTORED_V031_METADATA", hashlib.sha256((BOSS / "components/enm_boss_monitor002/monitor_motion.json").read_bytes()).hexdigest())
