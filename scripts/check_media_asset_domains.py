#!/usr/bin/env python3
"""Validate the P2 UI/SFX/music asset-domain split without mutating files."""

from __future__ import annotations

import json
import re
import sys
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

EXPECTED = {
    "ui": {
        "name": "UI",
        "file": "ShellStorm2_UI账本_v001.xlsx",
        "category": "UI",
        "skill": "ui-asset-pipeline",
        "count": 16,
        "tests": ["verify_hud_presenter_3d", "verify_pause_game_save_reset_flow", "verify_gamepad_input_flow"],
    },
    "audio": {
        "name": "音效",
        "file": "ShellStorm2_音效账本_v001.xlsx",
        "category": "音效",
        "skill": "audio-sfx-asset-pipeline",
        "count": 48,
        "tests": ["verify_requested_experience_upgrade_flow", "verify_3d_melee_feedback_flow"],
    },
    "music": {
        "name": "音乐",
        "file": "ShellStorm2_音乐账本_v001.xlsx",
        "category": "音乐",
        "skill": "music-asset-pipeline",
        "count": 9,
        "tests": ["verify_music_system"],
    },
}


def main() -> int:
    root = Path(__file__).resolve().parents[1]
    issues: list[str] = []
    index = json.loads((root / "assets/registry/ledger_index.json").read_text(encoding="utf-8"))
    manifest_path = root / "assets/registry/media_domain_split_manifest.json"
    if not manifest_path.is_file():
        issues.append("Missing media migration manifest")
        manifest = {}
    else:
        manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    baseline = json.loads((root / "assets/registry/ledger_split_baseline.json").read_text(encoding="utf-8"))
    baseline_assets = baseline.get("assets", {})

    domains = {item["key"]: item for item in index.get("domains", [])}
    runner = (root / "scripts/run_verification_suite.sh").read_text(encoding="utf-8")
    registered = set(re.findall(r"^\s+(verify_[a-zA-Z0-9_]+)\s*$", runner, re.MULTILINE))
    totals: Counter[str] = Counter()
    asset_ids: set[str] = set()

    for key, expected in EXPECTED.items():
        domain = domains.get(key)
        if domain is None:
            issues.append(f"Missing ledger domain: {key}")
            continue
        for field, wanted in (
            ("name", expected["name"]),
            ("file", expected["file"]),
            ("primary_skill", expected["skill"]),
        ):
            if domain.get(field) != wanted:
                issues.append(f"{key}.{field}: expected {wanted!r}, got {domain.get(field)!r}")
        if domain.get("categories") != [expected["category"]]:
            issues.append(f"{key}.categories must be [{expected['category']!r}]")

        skill_path = root / ".codex/skills" / expected["skill"] / "SKILL.md"
        if not skill_path.is_file():
            issues.append(f"Missing skill: {skill_path.relative_to(root)}")
        else:
            skill_text = skill_path.read_text(encoding="utf-8")
            if "## 验收" not in skill_text and "## Acceptance" not in skill_text:
                issues.append(f"Skill lacks an acceptance section: {expected['skill']}")
            if re.search(r"\bTODO\b|待补|待定", skill_text, re.IGNORECASE):
                issues.append(f"Skill contains unresolved placeholder: {expected['skill']}")

        ledger_path = root / "assets/registry/ledgers" / expected["file"]
        if not ledger_path.is_file():
            issues.append(f"Missing ledger: {ledger_path.relative_to(root)}")
            continue
        workbook = load_workbook(ledger_path, read_only=True, data_only=False)
        if "资产主表" not in workbook.sheetnames:
            issues.append(f"Missing 资产主表: {expected['file']}")
            continue
        rows = []
        for values in workbook["资产主表"].iter_rows(min_row=6, max_col=25, values_only=True):
            if values[0] not in (None, ""):
                rows.append(values)
        totals[key] = len(rows)
        if len(rows) != expected["count"]:
            issues.append(f"{key} row count: expected {expected['count']}, got {len(rows)}")
        for values in rows:
            asset_id = str(values[0]).strip()
            category = str(values[2]).strip()
            subtype = str(values[3] or "").strip().lower()
            runtime_path = str(values[14] or "").strip()
            if asset_id in asset_ids:
                issues.append(f"Duplicate media AssetID: {asset_id}")
            asset_ids.add(asset_id)
            if category != expected["category"]:
                issues.append(f"{asset_id}: category {category!r} is outside {key}")
            if key == "music" and subtype != "bgm":
                issues.append(f"{asset_id}: music row subtype must be bgm")
            if key == "audio" and subtype == "bgm":
                issues.append(f"{asset_id}: BGM must not remain in the SFX ledger")
            if key == "audio" and runtime_path and not runtime_path.startswith("src/assets/audio/sfx/"):
                issues.append(f"{asset_id}: SFX runtime path is outside src/assets/audio/sfx/")
            if key == "music" and runtime_path and not runtime_path.startswith("assets/audio/music/"):
                issues.append(f"{asset_id}: music runtime path is outside assets/audio/music/")
            baseline_row = baseline_assets.get(asset_id)
            if baseline_row is None:
                issues.append(f"{asset_id}: absent from the frozen pre-split baseline")
            else:
                if baseline_row.get("d") != key:
                    issues.append(f"{asset_id}: current baseline domain must be {key!r}")

        for test in expected["tests"]:
            if test not in registered:
                issues.append(f"Acceptance test is not registered: {test}")
            if not (root / f"tests/verification/{test}.gd").is_file():
                issues.append(f"Acceptance test script is missing: {test}")

    legacy = root / "assets/registry/ledgers/ShellStorm2_表现资源账本_v001.xlsx"
    if legacy.exists():
        issues.append(f"Legacy combined ledger still exists: {legacy.relative_to(root)}")
    routes = {item.get("target_domain"): item.get("expected_count") for item in manifest.get("routes", [])}
    for key, expected in EXPECTED.items():
        if routes.get(key) != expected["count"]:
            issues.append(f"Migration manifest count mismatch for {key}: {routes.get(key)!r}")

    result = {
        "domains": list(EXPECTED),
        "ledger_totals": dict(totals),
        "asset_count": sum(totals.values()),
        "registered_acceptance_tests": sorted(set().union(*(set(item["tests"]) for item in EXPECTED.values()))),
        "issues": issues,
    }
    print(json.dumps(result, ensure_ascii=False, indent=2))
    if issues:
        print(f"MEDIA_ASSET_DOMAINS_FAILED issues={len(issues)}")
        return 1
    print("MEDIA_ASSET_DOMAINS_OK domains=3 assets=73")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
