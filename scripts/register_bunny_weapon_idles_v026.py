"""Narrow source revision registration; preserve existing runtime and formulas."""
from pathlib import Path
from copy import copy
from collections import Counter
from datetime import datetime
import ast, hashlib, json, shutil, subprocess, sys, os
import openpyxl

ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT/'scripts'))
sys.path.insert(0,str(ROOT/'tools/asset_pipeline'))
from ledger_registry import LedgerIndex
from split_asset_ledger import read_source_rows,_row_digest,sheet_digest,col_digest,CONTENT_COLUMNS
INDEX=LedgerIndex.load(ROOT)
PATH=INDEX.path_for_category('角色')
BASELINE=ROOT/'assets/registry/ledger_split_baseline.json'
OUT=ROOT/'outputs/character_pipeline/weapon_idle_v026'
BACKUP=ROOT/'_scratch/bunny_weapon_idles_v026_ledger_before'
ASSET='CHR-PLY-CAPSULE01-3D-BUNNY01'
SOURCE='assets/art/characters/player/chr_player_capsule01_3d/variants/bunny01/source/animation/chr_bunny01_animation_v026.blend'
TRANSFER=SOURCE.replace('chr_bunny01_animation_v026.blend','chr_bunny01_weapon_idles_v026.json')
tree=ast.parse((ROOT/'scripts/register_bunny_weapon_idles_v025.py').read_text(encoding='utf-8'))
nodes=[n for n in tree.body if isinstance(n,ast.FunctionDef)]
exec(compile(ast.Module(body=nodes,type_ignores=[]),'<ledger-helpers>','exec'))
assert json.loads((ROOT/TRANSFER).read_text(encoding='utf-8'))['version']=='v026'
assert not BACKUP.exists()
gates('before')
BACKUP.mkdir(parents=True)
shutil.copy2(PATH,BACKUP/PATH.name)
shutil.copy2(BASELINE,BACKUP/BASELINE.name)
before=openpyxl.load_workbook(PATH)
wb=openpyxl.load_workbook(PATH)
ws=wb['资产主表']
row=next(r for r,v in read_source_rows(ws) if v[0]==ASSET)
ws.cell(row,16).value+='; '+SOURCE+'; '+TRANSFER
ws.cell(row,25).value+='；v026按用户正侧参考调整短枪前上倾、三枪前移，握点同步；源级待导入，v025保留。'
logrow=append_styled(wb['域变更日志'],['v0.1.11','2026-10-07','持枪站立待机位置修订','角色 / Bunny01','v026短枪参考位置、长枪机枪前移，三枪头部/躯干相交采样检查；旧源保留。','authored；运行路径、状态和版本保持。','Codex'],16)
wb.save(PATH)
baseline_write(changed_asset=ASSET)
wb=openpyxl.load_workbook(PATH)
for row in (61,62,63):
    assert wb['动画与状态'].cell(row,2).value in ('sidearm_idle','longgun_idle','machinegun_idle')
    wb['动画与状态'].cell(row,12).value='v026持枪位置按正侧参考修订，三类枪整体前移，双握点跟随；源级authored，未导入Godot。'+SOURCE
rows=[]
for label,p in [('v026动作母版（authored，未接线）',SOURCE),('v026源级中转（待导出）',TRANSFER)]:
    data=(ROOT/p).read_bytes()
    rows.append(append_styled(wb['角色中转记录'],[label,p,hashlib.sha256(data).hexdigest(),len(data)],106))
wb.save(PATH)
baseline_write(sheets=['动画与状态','角色中转记录'])
after=openpyxl.load_workbook(PATH)
diff=[]
for name in before.sheetnames:
    a,b=before[name],after[name]
    for cells in b:
        for c in cells:
            if a.cell(c.row,c.column).value!=c.value:
                assert (name=='资产主表' and c.coordinate in ('P13','Y13')) or (name=='动画与状态' and c.coordinate in ('L61','L62','L63')) or (name=='角色中转记录' and c.row in rows) or (name=='域变更日志' and c.row==logrow),(name,c.coordinate)
                diff.append([name,c.coordinate])
    assert str(a.data_validations)==str(b.data_validations)
    assert set(map(str,a.merged_cells.ranges))==set(map(str,b.merged_cells.ranges))
(OUT/'ledger_edit_report.json').write_text(json.dumps(diff,ensure_ascii=False,indent=2),encoding='utf-8')
gates('after')
print('V026_LEDGER_DONE',diff)
