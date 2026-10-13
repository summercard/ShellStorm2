"""Register approved component imports and standalone expedition01 boss room02."""
import json,hashlib,shutil,sys,re
from pathlib import Path
from copy import copy
from collections import Counter
from openpyxl import load_workbook
from openpyxl.worksheet.datavalidation import DataValidation
R=Path(__file__).resolve().parents[1];sys.path[:0]=[str(R/'scripts'),str(R/'tools/asset_pipeline')]
from ledger_registry import LedgerIndex
from split_asset_ledger import FIRST_DATA_ROW,CONTENT_COLUMNS,read_source_rows,_row_digest,col_digest,sheet_digest,dedupe_key_formula,dedupe_result_formula
from workspace_paths import work_path
W=work_path('outputs/boss_room02_import/registration.json').parent;D=R/'art_source/environments/tower_zones/expedition/source/room_types/boss_room_B/v027';S=json.loads((W/'spec.json').read_text('utf8'));B=R.parent/'ShellStorm2-backup/2026-10-13_boss_room02_import/before_edits';ix=LedgerIndex.load(R)
sha=lambda p:hashlib.sha256(p.read_bytes()).hexdigest()
def backup(p):
 q=B/p.relative_to(R);q.parent.mkdir(parents=True,exist_ok=True)
 if not q.exists():shutil.copy2(p,q)
def write(p,d):p.parent.mkdir(parents=True,exist_ok=True);p.write_text(json.dumps(d,ensure_ascii=False,indent=2)+'\n',encoding='utf8')
assert json.loads((W/'godot_acceptance.json').read_text('utf8'))['passed']
protected=json.loads((W/'protected_baseline.json').read_text('utf8'));assert all(sha(Path(p))==h for p,h in protected.items())
entries=[]
for d in S['components']:
 d['runtime_integrated']=True;d['fidelity_acceptance']='Fixed front/back/side/game contact sheets reviewed; exact bbox; MAE <0.011/255';d['prefab_sha256']=sha(R/d['prefab_path']);d['source_unchanged']=True
 entries.append({'asset_id':d['asset_id'],'name_zh':d['name_zh'],'slug':d['slug'],'source':d['source_path'],'prefab':d['prefab_path'],'glb':d['glb_path'],'collision':d['collision_policy'],'owner':d['collision_owner'],'size':d['bounds_size_m_godot'],'notes':'46 component budget. Original/optimized/export evidence: '+(D/'import_manifest.json').relative_to(R).as_posix(),'function_script':'src/world3d/BossRoomDisplay3D.gd' if d['slug'].startswith('screen_') or d['slug'] in ('stair_chase_light','stair_pillar_light') else ''})
master=next(D.glob('*.blend'));scene=S['runtime_scene'].removeprefix('res://')
entries.append({'asset_id':S['asset_id'],'name_zh':'远征关卡01 Boss房02 深井计算机房','slug':'boss_room_02','source':master.relative_to(R).as_posix(),'prefab':scene,'glb':'','collision':'component_owned','owner':'component_prefabs','size':[80,72,100],'notes':'New independent room, 46 definitions / 652 instances. Original boss01 retained; future level connection selects room02. No encounter, FSM, navigation or save changes.','function_script':''})
ledger=ix.path_for_category('场景');baseline=R/'assets/registry/ledger_split_baseline.json';backup(ledger);backup(baseline)
wb=load_workbook(ledger);ws=wb[ix.asset_sheet];old_end=ws.max_row;old_cells={(c.row,c.column):c.value for row in ws.iter_rows() for c in row};existing={str(v[0]).strip() for _,v in read_source_rows(ws)}
assert not any(e['asset_id'] in existing for e in entries),'Registration already exists: inspect and update instead of appending.'
for e in entries:
 row=ws.max_row+1
 values=[e['asset_id'],e['name_zh'],'场景','room' if e['slug']=='boss_room_02' else 'environment_module_3d','expedition_01_boss_room02_'+e['slug'],'room_variant' if e['slug']=='boss_room_02' else e['slug'],None,'Top3D / Godot Y-up','scene_art','远征关卡01 Boss房02；后续接关卡使用02，01保留','正式美术已接入','P1','v027','×'.join('%g'%v for v in e['size'])+'m',e['prefab'],e['source'],'远征关卡01;Boss房02;深井;'+e['slug'],None,None,sha(R/e['prefab']),'Codex','2026-10-13','用户参考图制作；Image-2屏幕为已授权贴图例外',e['slug'],e['notes']]
 for col,value in enumerate(values,1):c=ws.cell(row,col,value);c._style=copy(ws.cell(old_end,col)._style)
 ws.cell(row,18,dedupe_key_formula(row))
last=ws.max_row
for row in range(FIRST_DATA_ROW,last+1):ws.cell(row,19,dedupe_result_formula(row,last))
for row in wb['总览'].iter_rows():
 for c in row:
  if isinstance(c.value,str) and '资产主表' in c.value:c.value=re.sub(r'\$([A-Z]+)\$'+str(old_end)+r'\b',lambda m:'$'+m[1]+'$'+str(last),c.value)
for dv in ws.data_validations.dataValidation:
 ranges=[]
 for rng in dv.sqref.ranges:
  text=str(rng);text=re.sub(r'([A-Z]+)'+str(old_end)+r'\b',lambda m:m[1]+str(last),text);ranges.append(text)
 dv.sqref=' '.join(ranges)
ws.auto_filter.ref=f'A5:Y{last}'
log=wb['域变更日志'];lv=str(log.cell(log.max_row,1).value);match=re.search(r'(\d+)$',lv);version=lv[:match.start()]+str(int(match[1])+1) if match else 'v0.1.1';log.append([version,'2026-10-13','Codex','Boss房02运行资产登记','46组件＋独立房间TSCN，47行；原01保留','源/优化/运行哈希齐全；Godot专项3081检查通过，已有账本门禁债务单列'])
# Main-table transaction: existing content cells remain exact; only formula spans change.
assert all(ws.cell(r,c).value==v for (r,c),v in old_cells.items() if c not in (18,19))
wb.save(ledger);bl=json.loads(baseline.read_text('utf8'))
for _,values in read_source_rows(ws):
 if str(values[0]) in {e['asset_id'] for e in entries}:bl['assets'][str(values[0])]={'v':_row_digest(values),'c':'场景','d':'scenes'}
bl['asset_count']=len(bl['assets']);union=[]
for domain in ix.domains:union.extend(read_source_rows(load_workbook(domain.path)[ix.asset_sheet]))
bl['column_digests']={str(c):col_digest(union,c) for c in CONTENT_COLUMNS};bl['category_counts']=dict(sorted(Counter(str(v[2]).strip() for _,v in union).items()));write(baseline,bl)
# Separate prefab-page transaction. Prove every old cell survives before updating this page digest.
wb=load_workbook(ledger);ps=wb['3D-场景通用'];end=ps.max_row;old_page={(c.row,c.column):c.value for row in ps.iter_rows() for c in row};old_digest=sheet_digest(ps)
for e in entries:
 row=ps.max_row+1;values=[e['asset_id'],e['name_zh'],e['prefab'],e['glb'] or '无整屋GLB；引用独立组件',e['source'],'Boss房02独立运行组件与布局',e['function_script'],'有' if e['owner']!='none' else '无',e['owner'],e['collision'],'×'.join('%g'%v for v in e['size'])+'m','底部中心；Blender(x,y,z)→Godot(x,z,-y)，房间中心重定位一次','远征关卡01 Boss房02','正式美术已接入','v027',e['notes']]
 for col,value in enumerate(values,1):c=ps.cell(row,col,value);c._style=copy(ps.cell(end,col)._style)
assert all(ps.cell(r,c).value==v for (r,c),v in old_page.items())
wb.save(ledger);bl=json.loads(baseline.read_text('utf8'));bl['sheet_digests']['3D-场景通用']=sheet_digest(ps);write(baseline,bl)
write(W/'registration.json',{'registered_assets':entries,'main_rows':[old_end+1,last],'prefab_rows':[end+1,ps.max_row],'old_content_unchanged':True,'prefab_page_digest_before':old_digest,'prefab_page_baseline_before':json.loads((B/baseline.relative_to(R)).read_text('utf8'))['sheet_digests']['3D-场景通用'],'note':'Prefab page already had baseline drift before this task; existing page cells remain unchanged. New authorized page digest captured in separate transaction.'})
import_manifest=json.loads((D/'import_manifest.json').read_text('utf8'));import_manifest['components']=S['components'];import_manifest['runtime_integrated']=True;import_manifest['fidelity_views']=['front','back','side','game'];write(D/'import_manifest.json',import_manifest)
# Runtime AssetID resolver, additive only. Keep existing component IDs and room01 references.
p=R/'assets/art/environments/tower_zones/shared/runtime/shell_component_catalog.json';backup(p);c=json.loads(p.read_text('utf8'))
for d in S['components']:
 assert not any(x['component_id']==d['asset_id'] for x in c['components'])
 c['components'].append({'component_id':d['asset_id'],'aliases':[],'slug':d['slug'],'kind':'exclusive','room_type':'BOSS_ROOM','room_variant':'deep_shaft_02','category':'Boss房02组件','slot_role':'room_type_component','source_package_id':d['asset_id'],'prefab_asset_id':d['asset_id'],'prefab_path':'res://'+d['prefab_path'],'bounds_size_m_godot':d['bounds_size_m_godot'],'collision_owner':d['collision_owner'],'collision_policy':d['collision_policy'],'source_version':'v027','note':'Dedicated reusable Boss room02 component; placement owned by formal room TSCN.'})
c['component_count']=len(c['components']);c.setdefault('room_type_sources',[]).append((D/'component_catalog.json').relative_to(R).as_posix());write(p,c)
roomdir=R/'assets/art/environments/tower_zones/expedition/runtime/room_instances/expedition_01'
manifest={'schema':'shellstorm2.expedition.room_runtime_manifest.v001','room_id':'boss_room_02','room_number':'02','level_id':'expedition_01','block_id':'expedition','room_type':'BOSS_ROOM','variant_id':'deep_shaft_02','asset_id':S['asset_id'],'version':'v027','formal_scene':S['runtime_scene'],'scene_sha256':sha(R/scene),'source_blend':master.relative_to(R).as_posix(),'source_sha256':sha(master),'assembly_route':'room_type_layout_replay','layout_owner':'Godot TSCN after initialization','coordinate_contract':S['coordinate_contract'],'component_definitions':46,'instance_count':652,'components':[{k:d[k] for k in ('asset_id','slug','prefab_path','glb_path','glb_sha256','optimized_path','optimized_sha256','source_sha256_before','source_sha256_after','collision_policy')} for d in S['components']],'ports':[{'port_id':'west','position_m':[-40,0,-22.5],'outward':[-1,0,0]},{'port_id':'east','position_m':[40,0,-22.5],'outward':[1,0,0]},{'port_id':'entrance','position_m':[0,12,50],'outward':[0,0,1]}],'collision_owner':'component_prefabs','level_connection':'not replacing existing boss01 now; user selected room02 for future level connection','runtime_integrated':True,'acceptance':'scripts/verify_boss_room02_runtime.gd','ledger_ref':ledger.relative_to(R).as_posix()+'#3D-场景通用'}
write(roomdir/'boss_room_02_manifest.json',manifest)
write(roomdir/'boss_room_variants.json',{'level_id':'expedition_01','selection_owner':'Level assembly; explicit user choice 2026-10-13','variants':[{'room_number':'01','scene':'res://assets/art/environments/tower_zones/expedition/runtime/room_instances/expedition_01/f00_boss_static_layout.tscn','current_level_binding':True,'preserve':True},{'room_number':'02','scene':S['runtime_scene'],'asset_id':S['asset_id'],'selected_for_next_level_connection':True,'current_level_binding':False}],'note':'Catalog records the selection; current TowerDescent path resolver is unchanged. No claim of automatic variant switching.'})
p=R/'assets/registry/current_source_index.json';backup(p);index=json.loads(p.read_text('utf8'));paths={x['path'] for x in index['sources']}
for d in S['components']:
 rel=d['optimized_path'];f=R/rel
 if rel not in paths:index['sources'].append({'path':rel,'bytes':f.stat().st_size,'family':rel.replace('/v027/','/current/').replace('_v027.blend','.blend'),'version':[27],'role':'optimized','sha256':sha(f)})
write(p,index)
print('BOSS_ROOM02_REGISTERED',len(entries),'component resolver',c['component_count'])
