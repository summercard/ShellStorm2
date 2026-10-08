"""三款背包的显式登记事务；不吸收其他资产的历史漂移。"""
from pathlib import Path
from copy import copy, deepcopy
from collections import Counter
from datetime import datetime
import sys, json, hashlib, shutil, re, zipfile
import openpyxl
from openpyxl.worksheet.cell_range import MultiCellRange

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / 'scripts'))
sys.path.insert(0, str(ROOT / 'tools/asset_pipeline'))
from ledger_registry import LedgerIndex
from split_asset_ledger import FIRST_DATA_ROW, HEADER_ROW, read_source_rows, dedupe_key_formula, dedupe_result_formula, _row_digest, CONTENT_COLUMNS, col_digest, sheet_digest

OUT = ROOT / 'outputs/backpacks_20261008/ledger'
BASE = ROOT / 'assets/registry/ledger_split_baseline.json'
PARENT = 'ITM-EQUIPMENT-BACKPACK-3D'
ROLES = ['01_精工金属_紫色骨架', '02_细腻哑光_青绿大面', '03_清漆反光_紫粉点缀', '04_柔和自发光_UI灯光']
SPECS = [('small', 2, '小号粉色圆顶背包', [0.56, 0.28, 0.60]), ('medium', 4, '中号黄色单扣背包', [0.66, 0.33, 0.74]), ('large', 8, '大号青色卷铺盖背包', [0.76, 0.38, 0.88])]

def sha(p):
    return hashlib.sha256(p.read_bytes()).hexdigest()

def rel(p):
    return p.relative_to(ROOT).as_posix()

def cells(wb):
    return {(w.title, c.coordinate): c.value for w in wb for row in w for c in row if c.value is not None}

def append(w, values, template):
    row = max(c.row for rr in w for c in rr if c.column == 1 and c.value is not None) + 1
    for col, value in enumerate(values, 1):
        w.cell(row, col)._style = copy(w.cell(template, col)._style)
        w.cell(row, col).value = value
    w.row_dimensions[row].height = 64
    return row

def contract(spec, index):
    size, slots, name, dims = spec
    slug = 'backpack_' + size
    logic = 'prp_' + slug
    folder = ROOT / 'assets/art/items_weapons/props' / slug
    return {'schema_version': 1, 'asset_id': 'ITM-EQUIPMENT-BACKPACK-' + size.upper() + '-3D',
            'parent_asset_id': PARENT, 'display_name_zh': name, 'category': '道具', 'subclass': 'equipment',
            'logical_id': 'equipment_backpack_' + str(slots), 'extra_slots': slots,
            'domain_ledger': rel(index.path_for_category('道具')), 'ledger_sheet': '3D-物品',
            'dimensions_m': dict(zip(['width', 'depth', 'height'], dims)),
            'dimension_source': '用户批准的整包宽深高；含提手、前袋与大包卷铺盖',
            'origin_and_facing': '整包包围盒中心原点；Godot +Z前袋朝外，+Y向上；Blender -Y前袋朝外，+Z向上',
            'collision_intent': {'enabled': False, 'owner': '无', 'method': '无'},
            'function_note': '仅表现；地面拾取/UI图标/角色BackpackSocket共用；容量和存档规则不变',
            'animation_tier': 'tier_0_static', 'animation_reason': '封闭静态装备；本任务无开盖动作，无骨架、形变及自身动画',
            'source_collection': 'PRP_' + name + '_中文资产管理', 'anchor_contract': ['ItemRoot'],
            'material_roles': ROLES, 'material_override_reason': '用户明确要求场景四材质标准名；不使用道具专属别名',
            'palette': 'assets/art/shared/palette/设施低亮多巴胺色盘_10x10_512.png',
            'source_dir': rel(folder / 'source'), 'source_blend': rel(folder / 'source' / (logic + '_source_v001.blend')),
            'optimized_blend': rel(folder / 'source/export/v001' / (logic + '_optimized_v001.blend')),
            'stable_glb': rel(folder / 'components' / (logic + '_visual_top3d.glb')),
            'stable_prefab': rel(folder / 'runtime' / (logic + '_root_top3d.tscn')),
            'max_triangles': 500, 'shoulder_straps': False, 'front_clasp_straps': size != 'small', 'top_handle': True,
            'source_version': 'v001', 'reference_image': 'C:/Users/zhuangmenghong/.workbuddy/clipboard-images/clipboard-2026-10-08T09-16-34-177Z-40fdd968.jpg',
            'template_resolution': {'requested_missing': 'assets/art/items_weapons/_templates/prp/prp_template_source_v001.blend',
                'search_result': '仅找到角色模板，与静态道具契约不符，不挪用角色骨架',
                'replacement': 'assets/art/items_weapons/_templates/prp/prp_template_source_v001.blend',
                'action': '按06五集合、公共色盘、PaletteUV及ItemRoot补建静态分支模板；不覆盖历史文件；tier_0删除示例骨架'}}

def main():
    phase = sys.argv[1]
    assert phase in ['register', 'prefab', 'promote']
    index = LedgerIndex.load(ROOT)
    path = index.path_for_category('道具')
    OUT.mkdir(parents=True, exist_ok=True)
    assert not (OUT / ('before_' + phase + '.xlsx')).exists(), '事务已执行，禁止覆盖备份'
    shutil.copy2(path, OUT / ('before_' + phase + '.xlsx'))
    shutil.copy2(BASE, OUT / ('before_' + phase + '_baseline.json'))
    wb = openpyxl.load_workbook(path)
    baseline = json.loads(BASE.read_text('utf-8'))
    old = deepcopy(baseline)
    before = cells(wb)
    w = wb['资产主表']
    records = [contract(s, index) for s in SPECS]
    changed_ids = {PARENT} | {c['asset_id'] for c in records}
    if phase == 'register':
        assert not any(v[0] in changed_ids - {PARENT} for _, v in read_source_rows(w))
        parent_row = next(r for r, v in read_source_rows(w) if v[0] == PARENT)
        assert w.cell(parent_row, 11).value == '程序占位'
        w.cell(parent_row, 2).value = '可装备背包3D资产族'
        w.cell(parent_row, 9).value = 'asset_family / tier_0_static'
        w.cell(parent_row, 11).value = 'design_only'
        w.cell(parent_row, 14).value = '资产族索引；小/中/大三款各自独立源、GLB和PackedScene；非独立几何'
        w.cell(parent_row, 22).value = datetime(2026, 10, 8)
        w.cell(parent_row, 25).value = '保留原AssetID与equipment_backpack逻辑族；三款子资产通过G列归属；父项无生产路径和SHA，不再冒充程序几何。'
        last = max(r for r, _ in read_source_rows(w))
        for c in records:
            folder = ROOT / c['source_dir']
            folder.mkdir(parents=True, exist_ok=True)
            (folder / '.gdignore').write_text('', encoding='utf-8')
            (folder.parent / 'asset_contract.json').write_text(json.dumps(c, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
            d = c['dimensions_m']
            dim = f"{d['width']:.2f}×{d['depth']:.2f}×{d['height']:.2f}m"
            append(w, [c['asset_id'], c['display_name_zh'], '道具', 'equipment', c['logical_id'], 'root_3d', PARENT, '俯视3D / local +Z正面', 'tier_0_static', c['function_note'], '待制作', 'P0', 'v001', dim + '；整包中心原点；<=500三角；无背部肩带；场景四材质/公共色盘', None, 'src/world3d/ItemModelFactory3D.gd; tests/verification/verify_backpack_equipment_flow.tscn', 'backpack;' + c['display_name_zh'], None, None, None, 'WorkBuddy', datetime(2026,10,8), '用户参考图重建；项目公共色盘', c['source_collection'], '新增身份，非历史已有ID；尺寸经用户批准；静态闭合表现；以asset_contract.json冻结。'], parent_row)
        end = max(r for r, _ in read_source_rows(w))
        for r, _ in read_source_rows(w):
            w.cell(r,18).value = dedupe_key_formula(r)
            w.cell(r,19).value = dedupe_result_formula(r,end)
        for ref in ['A6','C6','E6','G6','B10','C10','B11','C11']:
            cell = wb['总览'][ref]
            assert f'${last}' in cell.value
            cell.value = cell.value.replace(f'${last}', f'${end}')
        for dv in w.data_validations.dataValidation:
            ranges = []
            for rr in dv.sqref.ranges:
                rr = copy(rr)
                if rr.min_row == FIRST_DATA_ROW: rr.max_row = end
                ranges.append(str(rr))
            dv.sqref = MultiCellRange(' '.join(ranges))
        w.auto_filter.ref = f'A{HEADER_ROW}:Y{end}'
    elif phase == 'prefab':
        p = wb['3D-物品']
        assert p['A4'].value == 'AssetID' and not p['A5'].value
        for c in records:
            assert (ROOT / c['stable_prefab']).exists()
            d = c['dimensions_m']
            append(p, [c['asset_id'], c['display_name_zh'], c['stable_prefab'], c['stable_glb'], c['source_blend'], c['function_note'], 'src/world3d/ItemModelFactory3D.gd', '无','无','无',f"{d['width']:.2f}×{d['depth']:.2f}×{d['height']:.2f}m",c['origin_and_facing'],'世界掉落/UI图标/玩家BackpackSocket','待制作','v001','真实PackedScene骨架已创建；视觉制作及验收待完成；禁止视为正式资产。'],4)
        p['A2'] = '独立物品PackedScene；3款背包已登记，等待制作与验收。'
        baseline['sheet_digests']['3D-物品'] = sheet_digest(p)
    else:
        acceptance = json.loads((OUT.parent / 'final_acceptance.json').read_text('utf-8'))
        assert acceptance['ready_for_promotion'] is True
        for c in records:
            r = next(r for r,v in read_source_rows(w) if v[0] == c['asset_id'])
            manifest = json.loads((ROOT / c['source_dir']).parent.joinpath('source_manifest.json').read_text('utf-8'))
            w.cell(r,11).value = '正式美术已接入'
            w.cell(r,15).value = c['stable_prefab']
            w.cell(r,20).value = sha(ROOT / c['stable_prefab'])
            w.cell(r,25).value = f"{manifest['triangles']}三角；无肩带；源/优化/导出及逐面PaletteUV通过；工厂/背负/世界掉落/UI共用同一资产；证据outputs/backpacks_20261008/final_acceptance.json。"
        p = wb['3D-物品']
        for r in range(5,8):
            p.cell(r,14).value = '正式美术已接入'
            p.cell(r,16).value = '静态纯视觉，无骨架/动画/玩法碰撞；source_manifest.json记录源与GLB哈希及面数；真实渲染和专项验收通过。'
        p['A2'] = '独立物品PackedScene；3款背包已正式接入；共用工厂服务背负、掉落和UI。'
        baseline['sheet_digests']['3D-物品'] = sheet_digest(p)
    log = wb['域变更日志']
    versions = [str(c.value) for row in log for c in row if c.column == 1 and re.fullmatch(r'v0\.1\.\d+', str(c.value))]
    patch = max([int(v.rsplit('.',1)[1]) for v in versions] or [0]) + 1
    append(log,[f'v0.1.{patch}','2026-10-08','三款无肩带背包：' + phase,'道具 / equipment',','.join(c['asset_id'] for c in records),'仅授权背包；其他资产内容和指纹不变。','WorkBuddy'],log.max_row)
    for r, values in read_source_rows(w):
        if values[0] in changed_ids:
            baseline['assets'][values[0]] = {'v': _row_digest(values),'c':'道具','d':'props'}
    baseline['asset_count'] = len(baseline['assets'])
    union = []
    for domain in index.domains:
        union.extend(read_source_rows((wb if domain.key == 'props' else openpyxl.load_workbook(domain.path))['资产主表']))
    baseline['column_digests'] = {str(c):col_digest(union,c) for c in CONTENT_COLUMNS}
    baseline['category_counts'] = dict(sorted(Counter(v[2] for _,v in union).items()))
    assert all(baseline['assets'][k] == v for k,v in old['assets'].items() if k not in changed_ids)
    after = cells(wb)
    changes = sorted(k for k in set(before)|set(after) if before.get(k) != after.get(k))
    allowed_sheets = {'总览','资产主表','域变更日志'} if phase == 'register' else {'3D-物品','域变更日志'} if phase == 'prefab' else {'资产主表','3D-物品','域变更日志'}
    assert all(k[0] in allowed_sheets for k in changes)
    oldrows = {v[0]:v for _,v in read_source_rows(openpyxl.load_workbook(OUT / ('before_' + phase + '.xlsx'))['资产主表'])}
    for _,v in read_source_rows(w):
        if v[0] not in changed_ids:
            assert all(v[i] == oldrows[v[0]][i] for i in range(25) if i not in (17,18))
    wb.save(path)
    assert cells(openpyxl.load_workbook(path)) == after
    BASE.write_text(json.dumps(baseline,ensure_ascii=False,indent=1)+'\n',encoding='utf-8')
    (OUT / (phase + '_transaction.json')).write_text(json.dumps({'phase':phase,'changed_cells':changes,'historical_fingerprints_preserved':True,'ledger_sha256':sha(path)},ensure_ascii=False,indent=2),encoding='utf-8')
    if phase == 'register':
        with zipfile.ZipFile(path) as z:
            xml = '\n'.join(z.read(n).decode('utf-8') for n in z.namelist() if n.startswith('xl/worksheets/') and n.endswith('.xml'))
            assert all(f'{col}6:{col}{end}' in xml for col in ['C','K','L'])
    print('BACKPACK_LEDGER_TRANSACTION_OK',phase,len(changes))

if __name__ == '__main__':
    main()
