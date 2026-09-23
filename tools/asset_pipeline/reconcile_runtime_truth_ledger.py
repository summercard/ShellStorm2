#!/usr/bin/env python3
"""Reconcile split asset ledgers with current runtime truth.

Policy (2026-09-23):
1. A runtime asset currently loaded by the game is authoritative over stale ledger paths/hashes.
2. One physical shared asset has one primary registry row; semantic slots stay in notes/contracts.
3. The formal vending-machine prefab is the runtime/vending_machine scene.
"""
from __future__ import annotations

import hashlib
import json
import re
import sys
from datetime import date
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "scripts"))
from ledger_registry import LedgerIndex  # noqa: E402

sys.path.insert(0, str(Path(__file__).resolve().parent))
from split_asset_ledger import (  # noqa: E402
    CONTENT_COLUMNS,
    _row_digest,
    col_digest,
    dedupe_key_formula,
    dedupe_result_formula,
    read_source_rows,
    sheet_digest,
)

FIRST_DATA_ROW = 6
DONE_STATUSES = (
    "已完成", "原型已接入", "正式美术已接入", "已优化并正式接入",
    "Blender源已完成", "已导入；优化完成", "白模源已完成；QA PASS",
)
SKIP_HASH_STATUSES = {"待制作", "程序占位", "弃用"}

REMOVE_SHARED_ROWS = {
    # These are slots/configuration concepts backed by another registered physical file.
    "CHR-PLY-CAPSULE01-HAND-3D",
    "CHR-PLY-CAPSULE01-SCARF-3D",
    "CHR-PLY-CAPSULE01-WEAPON-SOCKET-3D",
    "CHR-PLY-CAPSULE01-FOOT-3D",
    "CHR-PLY-CAPSULE01-EYES-3D",
    "CHR-PLY-CAPSULE01-EARS-3D",
    "CHR-PLY-CAPSULE01-TAIL-STUB-3D",
    "CHR-PLY-CAPSULE01-BODY-DIY-3D",
    "CHR-PLY-CAPSULE01-HEAD-DIY-3D",
    "CHR-PLY-CAPSULE01-HAND-DIY-3D",
    "CHR-PLY-CAPSULE01-HAT-DIY-3D",
    "CHR-PLY-CAPSULE01-GLASSES-DIY-3D",
    "CHR-PLY-CAPSULE01-FEET-DIY-3D",
    "CHR-PLY-CAPSULE01-BACKPACK-SOCKET-3D",
    "CHR-PLY-CAPSULE01-3D-BUNNY01-LOWER-BODY-SOCKET",
    "VFX-HEALTH-VIGNETTE",
    "UI-SCREEN-INVENTORY",
}

TOWER_PLACEHOLDERS = {
    "ENV-TOWER-STAIR-FLOOR-LOWER",
    "ENV-TOWER-STAIR-FLOOR-MID",
    "ENV-TOWER-STAIR-FLOOR-UPPER",
    "ENV-TOWER-STAIR-FLIGHT-ADJUSTABLE",
    "ENV-TOWER-CORNER-COLUMN-05M",
}

VENDING_ID = "PRP-BASE-VENDING-MACHINE-3D"
VENDING_PATH = "assets/art/props/base_world_3d/runtime/vending_machine/prp_base_vending_machine_root_top3d.tscn"
PLAYER_BASE_ID = "CHR-PLY-CAPSULE01"
PLAYER_BASE_PATH = "assets/art/characters/player/chr_player_capsule01_3d/chr_player_capsule01_root_top3d_v001.tscn"


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def stable_runtime_path(raw: str) -> str | None:
    candidate = re.sub(r"_v\d{3}(?=\.(?:tscn|tres|glb|gltf)$)", "", raw)
    if candidate != raw and (ROOT / candidate).is_file():
        return candidate
    return None


def remove_ids_from_sheet(ws) -> int:
    removed = 0
    for row in range(ws.max_row, 1, -1):
        if str(ws.cell(row, 1).value or "").strip() in REMOVE_SHARED_ROWS:
            ws.delete_rows(row, 1)
            removed += 1
    return removed


def refresh_overview(workbook, end: int) -> None:
    overview = workbook["总览"]
    overview["A6"] = f"=COUNTA('资产主表'!$A$6:$A${end})"
    overview["C6"] = "=" + "+".join(
        f'COUNTIF(\'资产主表\'!$K$6:$K${end},"{status}")' for status in DONE_STATUSES
    )
    overview["E6"] = (
        f'=COUNTIF(\'资产主表\'!$K$6:$K${end},"待制作")+'
        f'COUNTIF(\'资产主表\'!$K$6:$K${end},"程序占位")'
    )
    overview["G6"] = f'=COUNTIF(\'资产主表\'!$S$6:$S${end},"重复")'
    row = 10
    statuses = "+".join(f'(\'资产主表\'!$K$6:$K${end}="{status}")' for status in DONE_STATUSES)
    while overview.cell(row, 1).value not in (None, ""):
        overview.cell(row, 2).value = f"=COUNTIF('资产主表'!$C$6:$C${end},A{row})"
        overview.cell(row, 3).value = f"=SUMPRODUCT(('资产主表'!$C$6:$C${end}=A{row})*({statuses}))"
        row += 1


def reconcile_domain(path: Path) -> dict[str, int]:
    workbook = load_workbook(path, data_only=False)
    main = workbook["资产主表"]
    removed = remove_ids_from_sheet(main)
    # Keep secondary contract pages free of rows whose physical identity was consolidated.
    for ws in workbook.worksheets:
        if ws.title not in {"资产主表", "总览", "账本说明", "域变更日志"}:
            remove_ids_from_sheet(ws)

    paths = hashes = statuses = 0
    for row in range(FIRST_DATA_ROW, main.max_row + 1):
        asset_id = str(main.cell(row, 1).value or "").strip()
        if not asset_id:
            continue
        if asset_id == VENDING_ID:
            main.cell(row, 15).value = VENDING_PATH
            paths += 1
        elif asset_id == PLAYER_BASE_ID:
            main.cell(row, 15).value = PLAYER_BASE_PATH
            paths += 1
        else:
            raw = str(main.cell(row, 15).value or "").strip()
            replacement = stable_runtime_path(raw)
            if replacement:
                main.cell(row, 15).value = replacement
                paths += 1

        if asset_id in TOWER_PLACEHOLDERS:
            main.cell(row, 11).value = "程序占位"
            main.cell(row, 12).value = "P2"
            note = str(main.cell(row, 25).value or "").strip()
            marker = "2026-09-23：连续爬塔当前非主要玩法，保留白盒设计与源文件，降为程序占位/P2；未来重新开放时再恢复专项验收。"
            if marker not in note:
                main.cell(row, 25).value = (note + "\n" + marker).strip()
            statuses += 1

        status = str(main.cell(row, 11).value or "").strip()
        raw = str(main.cell(row, 15).value or "").strip()
        target = ROOT / raw if raw else None
        if status not in SKIP_HASH_STATUSES and target and target.is_file():
            actual = sha256(target)
            if str(main.cell(row, 20).value or "").strip().lower() != actual:
                main.cell(row, 20).value = actual
                hashes += 1

    asset_rows = [
        row for row in range(FIRST_DATA_ROW, main.max_row + 1)
        if str(main.cell(row, 1).value or "").strip()
    ]
    end = max(asset_rows, default=FIRST_DATA_ROW - 1)
    for row in range(FIRST_DATA_ROW, end + 1):
        if main.cell(row, 1).value:
            main.cell(row, 18).value = dedupe_key_formula(row)
            main.cell(row, 19).value = dedupe_result_formula(row, end)
    refresh_overview(workbook, end)
    workbook.save(path)
    return {
        "removed": removed,
        "paths": paths,
        "hashes": hashes,
        "statuses": statuses,
        "rows": len(asset_rows),
    }


def update_master_counts(index: LedgerIndex, counts: dict[str, int]) -> None:
    workbook = load_workbook(index.master_path, data_only=False)
    ws = workbook["分账本索引"]
    for offset, domain in enumerate(index.domains):
        ws.cell(FIRST_DATA_ROW + offset, 7).value = counts[domain.key]
    ws["A17"] = f"快照日期 {date.today().isoformat()}；条数为当前运行时真源对账后的快照，权威数量在各分账本《总览》。新增资产先确认归属域，再进对应分账本登记。"
    workbook.save(index.master_path)


def refresh_current_baseline(index: LedgerIndex) -> None:
    old_path = ROOT / "assets/registry/ledger_split_baseline.json"
    old = json.loads(old_path.read_text(encoding="utf-8"))
    assets = {}
    collected = []
    sheet_digests = {}
    category_counts = {}
    for domain in index.domains:
        workbook = load_workbook(domain.path, data_only=False)
        for _row_number, values in read_source_rows(workbook["资产主表"]):
            asset_id = str(values[0]).strip()
            category = str(values[2]).strip()
            assets[asset_id] = {"v": _row_digest(values), "c": category, "d": domain.key}
            collected.append((asset_id, values))
            category_counts[category] = category_counts.get(category, 0) + 1
        for sheet in domain.sheet_scope:
            if sheet in workbook.sheetnames:
                sheet_digests[sheet] = sheet_digest(workbook[sheet])
    collected.sort(key=lambda item: item[0])
    old.update({
        "source": "assets/registry/ledgers (current runtime-truth reconciliation)",
        "captured_at": date.today().isoformat(),
        "asset_count": len(assets),
        "assets": assets,
        "column_digests": {str(col): col_digest(collected, col) for col in CONTENT_COLUMNS},
        "sheet_digests": sheet_digests,
        "category_counts": category_counts,
    })
    old_path.write_text(json.dumps(old, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")


def main() -> int:
    index = LedgerIndex.load(ROOT)
    counts: dict[str, int] = {}
    for domain in index.domains:
        result = reconcile_domain(domain.path)
        counts[domain.key] = result["rows"]
        print(domain.key, result)
    update_master_counts(index, counts)
    refresh_current_baseline(index)
    print("RUNTIME_TRUTH_LEDGER_RECONCILE_OK", sum(counts.values()))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
