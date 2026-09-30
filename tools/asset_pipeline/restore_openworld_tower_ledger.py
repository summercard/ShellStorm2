#!/usr/bin/env python3
"""Restore Tower02 v003 and Tower03 v001 Blender-source rows after a ledger reset."""
from __future__ import annotations

import hashlib
import json
import shutil
import sys
from copy import copy
from datetime import date
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "scripts"))
from ledger_registry import LedgerIndex
from split_asset_ledger import CONTENT_COLUMNS, _row_digest, col_digest, dedupe_key_formula, dedupe_result_formula, read_source_rows

LEDGER = ROOT / "assets/registry/ledgers/ShellStorm2_场景账本_v001.xlsx"
BASELINE = ROOT / "assets/registry/ledger_split_baseline.json"
MASTER = ROOT / "assets/registry/ShellStorm2_美术资产台账_v001.xlsx"
TOWER02 = ROOT / "assets/art/environments/open_world/source/tower_02/v003/塔2_施工高楼_70x50m_v003.blend"
TOWER03 = ROOT / "assets/art/environments/open_world/source/tower_03/v001/塔楼03_设备天台办公楼_v001.blend"


def sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def main() -> None:
    assert TOWER02.is_file() and TOWER03.is_file()
    wb = load_workbook(LEDGER)
    ws = wb["资产主表"]
    assert ws.cell(248, 1).value == "ENV-OPENWORLD-TOWER02"
    assert ws.max_row in (248, 249)
    if ws.max_row == 249:
        assert ws.cell(249, 1).value == "ENV-OPENWORLD-TOWER03"
    backup = ROOT / "_scratch/restore_openworld_tower_ledger_backup"
    backup.mkdir(parents=True, exist_ok=True)
    for path in (LEDGER, BASELINE, MASTER):
        target = backup / path.name
        if not target.exists():
            shutil.copy2(path, target)

    src02 = TOWER02.relative_to(ROOT).as_posix()
    src03 = TOWER03.relative_to(ROOT).as_posix()
    ws.cell(248, 13, "v003")
    ws.cell(248, 14, "主体70×50m；92独立资产包；灰色屋顶、错层楼板与高低立柱；3台塔吊各自独立合集，每台7组成件；4角色公共色盘")
    ws.cell(248, 15, src02)
    ws.cell(248, 17, "塔2;施工高楼;塔吊独立合集;错层屋顶;灰色顶面;v003")
    ws.cell(248, 20, sha(TOWER02))
    ws.cell(248, 22, date.today())
    ws.cell(248, 25, "Blender原始文件已完成；3台塔吊为各自独立合集。未导出GLB、未接入Godot、未制作碰撞/LOD。")

    for col in range(1, 26):
        ws.cell(249, col)._style = copy(ws.cell(248, col)._style)
    ws.row_dimensions[249].height = ws.row_dimensions[248].height
    values = {
        1: "ENV-OPENWORLD-TOWER03",
        2: "开放世界 塔楼03 设备天台办公楼 70×44m",
        3: "场景",
        4: "environment_kit_3d",
        5: "open_world_tower_03",
        6: "building_source",
        8: "Top3D / Blender Z-up",
        9: "office_rooftop",
        10: "开放世界独立塔楼，与主塔、塔2平级",
        11: "Blender源已完成",
        12: "P1",
        13: "v001",
        14: "主体70×44m；屋顶高40m、机房高8m；217独立资产包；天台双层机房、风机、风管、通信桅杆和卫星天线",
        15: src03,
        16: "用户参考图1（天台布局）及图2（整体结构）；assets/art/environments/open_world/source/tower_03/v001/catalog.json",
        17: "塔楼03;设备天台;办公楼;风机;风管;通信桅杆;卫星天线;70×44m",
        20: sha(TOWER03),
        21: "Codex",
        22: date.today(),
        23: "用户提供参考图，仅据图建模",
        24: "塔楼03",
        25: "仅Blender原始文件；按用户确认尺寸制作。未导出GLB、未接入Godot、未制作碰撞/LOD。",
    }
    for col, value in values.items():
        ws.cell(249, col, value)
    for row in range(6, 250):
        ws.cell(row, 18, dedupe_key_formula(row))
        ws.cell(row, 19, dedupe_result_formula(row, 249))

    overview = wb["总览"]
    for cell in ("A6", "C6", "E6", "G6", "B10", "C10", "B11", "C11", "B12", "C12"):
        value = overview[cell].value
        assert isinstance(value, str) and "$248" in value, (cell, value)
        overview[cell] = value.replace("$248", "$249").replace("==", "=")
    dvs = list(ws.data_validations.dataValidation)
    ws.data_validations.dataValidation.clear()
    for dv in dvs:
        new = DataValidation(type=dv.type, formula1=dv.formula1, formula2=dv.formula2, allow_blank=dv.allow_blank)
        new.error = dv.error
        new.errorTitle = dv.errorTitle
        new.prompt = dv.prompt
        new.promptTitle = dv.promptTitle
        new.showErrorMessage = dv.showErrorMessage
        new.showInputMessage = dv.showInputMessage
        for rng in dv.sqref.ranges:
            new.add(str(rng).replace("248", "249"))
        ws.add_data_validation(new)
    log = wb["域变更日志"]
    last = log.max_row
    for col in range(1, log.max_column + 1):
        log.cell(last + 1, col)._style = copy(log.cell(last, col)._style)
    log.cell(last + 1, 1, "v0.1.25")
    log.cell(last + 1, 2, date.today())
    log.cell(last + 1, 3, "复原开放世界塔楼登记")
    log.cell(last + 1, 4, "关卡场景 / Blender源")
    log.cell(last + 1, 5, "塔2更新至v003并恢复塔楼03 v001源文件登记；仅Blender源，未登记运行时Prefab。")
    wb.save(LEDGER)

    index = LedgerIndex.load(ROOT)
    baseline = json.loads(BASELINE.read_text(encoding="utf-8"))
    all_rows = []
    category_counts = {}
    for domain in index.domains:
        book = load_workbook(domain.path, read_only=True, data_only=False)
        for _row, record in read_source_rows(book["资产主表"]):
            asset_id = str(record[0]).strip()
            category = str(record[2]).strip()
            all_rows.append((asset_id, record))
            category_counts[category] = category_counts.get(category, 0) + 1
            if asset_id in {"ENV-OPENWORLD-TOWER02", "ENV-OPENWORLD-TOWER03"}:
                baseline["assets"][asset_id] = {"v": _row_digest(record), "c": category, "d": domain.key}
    all_rows.sort(key=lambda item: item[0])
    baseline["asset_count"] = len(baseline["assets"])
    baseline["column_digests"] = {str(col): col_digest(all_rows, col) for col in CONTENT_COLUMNS}
    baseline["category_counts"] = category_counts
    baseline["captured_at"] = date.today().isoformat()
    BASELINE.write_text(json.dumps(baseline, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")

    master = load_workbook(MASTER)
    index_sheet = master["分账本索引"]
    for offset, domain in enumerate(index.domains):
        if domain.key == "scenes":
            index_sheet.cell(6 + offset, 7, ws.max_row - 5)
    master.save(MASTER)
    print("OPENWORLD_TOWER_LEDGER_RESTORED", ws.max_row, baseline["asset_count"])


if __name__ == "__main__":
    main()
