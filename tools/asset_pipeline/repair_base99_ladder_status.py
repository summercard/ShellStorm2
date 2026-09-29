"""Synchronize the v002 ladder ledger row and its lossless baseline."""
from pathlib import Path
import hashlib
import json
import sys

import openpyxl

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
sys.path.insert(0, str(HERE))
from split_asset_ledger import _row_digest, col_digest, read_source_rows  # noqa: E402

books = ROOT / "assets/registry/ledgers"
path = books / "ShellStorm2_道具账本_v001.xlsx"
workbook = openpyxl.load_workbook(path)
sheet = workbook["资产主表"]
assert sheet.cell(26, 1).value == "PRP-BASE99-TELESCOPIC-LADDER-3D"
assert sheet.cell(26, 11).value in ("正式资源已接入", "正式美术已接入")
sheet.cell(26, 11).value = "正式美术已接入"
sheet.cell(26, 13).value = "v002"
sheet.cell(26, 14).value = "双段6m直梯；上端E放下；展开后两端E自动攀爬、反向输入折返；顶部穿东门至100F平台、底部落99F阁楼；场景实体碰撞；展开状态存档"
sheet.cell(26, 16).value = str(sheet.cell(26, 16).value).replace("prp_base99_telescopic_ladder_source_v001.blend", "prp_base99_telescopic_ladder_source_v002.blend")
sheet.cell(26, 25).value = "2026-09-29：梯子贴合99F东墙并朝向阁楼；上端跨东门落100F平台。Blender v002 直接复用基地母版的四种既有材质，不新增或调整材质参数。展开一次后永久保持并写入基地存档。"
runtime = ROOT / str(sheet.cell(26, 15).value)
sheet.cell(26, 20).value = hashlib.sha256(runtime.read_bytes()).hexdigest()
workbook.save(path)
baseline_path = ROOT / "assets/registry/ledger_split_baseline.json"
baseline = json.loads(baseline_path.read_text(encoding="utf-8"))
record = tuple(sheet.cell(26, col).value for col in range(1, 26))
baseline["assets"][record[0]]["v"] = _row_digest(record)
index = json.loads((ROOT / "assets/registry/ledger_index.json").read_text(encoding="utf-8"))
all_rows = []
for domain in index["domains"]:
    other = openpyxl.load_workbook(books / domain["file"], read_only=True)
    all_rows.extend(read_source_rows(other["资产主表"]))
for col in (11, 13, 14, 16, 20, 25):
    baseline["column_digests"][str(col)] = col_digest(all_rows, col)
baseline_path.write_text(json.dumps(baseline, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")
print("BASE99_LADDER_STATUS_REPAIRED")
