"""Record the seat collision-follow and tip-recovery runtime fix in the props ledger."""

from pathlib import Path
import hashlib
import json

import openpyxl

from split_asset_ledger import _row_digest, col_digest, read_source_rows


ROOT = Path(__file__).resolve().parents[2]
index = json.loads((ROOT / "assets/registry/ledger_index.json").read_text(encoding="utf-8"))
books = ROOT / "assets/registry/ledgers"
props_domain = next(domain for domain in index["domains"] if domain["key"] == "props")
book_path = books / props_domain["file"]
workbook = openpyxl.load_workbook(book_path)
sheet = workbook["资产主表"]
targets = {
    18: "PRP-BASE-WORKSHOP-STOOL-3D",
    19: "PRP-BASE-MISSION-COMMAND-CHAIR-3D",
}
for row, asset_id in targets.items():
    assert sheet.cell(row, 1).value == asset_id
    spec = str(sheet.cell(row, 14).value)
    if "45°倾倒自动离座" not in spec:
        sheet.cell(row, 14).value = spec + "；45°倾倒自动离座，E扶正后可再次乘坐；角色专用碰撞体随刚体同步移动"
    runtime = ROOT / str(sheet.cell(row, 15).value)
    sheet.cell(row, 20).value = hashlib.sha256(runtime.read_bytes()).hexdigest()
    note = str(sheet.cell(row, 25).value)
    if "修复PlayerBlocker旧位置残留" not in note:
        sheet.cell(row, 25).value = note + "\n2026-09-29：修复PlayerBlocker旧位置残留；双椅倾角超过45°自动让角色离座，再次E交互扶正；Godot物理回归验证。"
workbook.save(book_path)

baseline_path = ROOT / "assets/registry/ledger_split_baseline.json"
baseline = json.loads(baseline_path.read_text(encoding="utf-8"))
for row, asset_id in targets.items():
    record = tuple(sheet.cell(row, col).value for col in range(1, 26))
    baseline["assets"][asset_id]["v"] = _row_digest(record)
all_rows = []
for domain in index["domains"]:
    other = openpyxl.load_workbook(books / domain["file"], read_only=True)
    all_rows.extend(read_source_rows(other["资产主表"]))
for col in (14, 20, 25):
    baseline["column_digests"][str(col)] = col_digest(all_rows, col)
baseline_path.write_text(json.dumps(baseline, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")
print("PUSHABLE_SEAT_COLLISION_LEDGER_REPAIRED")
