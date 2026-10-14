"""Register the completed 20-asset batch and preserve unrelated asset fingerprints."""
import sys,json,shutil,re,hashlib
from pathlib import Path
from copy import copy,deepcopy
from collections import Counter
from datetime import datetime
import openpyxl
from openpyxl.worksheet.cell_range import MultiCellRange
ROOT=Path(__file__).resolve().parents[2]
sys.path[:0]=[str(ROOT/'scripts'),str(ROOT/'tools/asset_pipeline')]
from workspace_paths import work_path
from ledger_registry import LedgerIndex
from split_asset_ledger import read_source_rows,_row_digest,col_digest,CONTENT_COLUMNS,sheet_digest,dedupe_key_formula,dedupe_result_formula,FIRST_DATA_ROW,HEADER_ROW
OUT=work_path('sessions/2026-10-14_remaining_items/ledger_transaction.json').parent
REPORT=json.loads((OUT/'build_report.json').read_text('utf8'))
assert json.loads((OUT/'runtime_result.json').read_text('utf8'))['passed']
BACKUP=ROOT.parent/'ShellStorm2-backup/2026-10-14_remaining_items/before_edits'
def backup(p):
    dest=BACKUP/p.relative_to(ROOT);dest.parent.mkdir(parents=True,exist_ok=True)
    if not dest.exists():shutil.copy2(p,dest)
def write(p,d):p.write_text(json.dumps(d,ensure_ascii=False,indent=2)+'\n',encoding='utf8')
index=LedgerIndex.load(ROOT);bp=ROOT/'assets/registry/ledger_split_baseline.json';backup(bp)
baseline=json.loads(bp.read_text('utf8'));old=deepcopy(baseline)
books={d.key:openpyxl.load_workbook(d.path) for d in index.domains}
target={x['asset_id'] for x in REPORT}
source_index_path=ROOT/'assets/registry/current_source_index.json';backup(source_index_path);sources=json.loads(source_index_path.read_text('utf8'))
for domain in ['props','weapons']:
    path=index.domain_for_key(domain).path;backup(path);wb=books[domain];ws=wb['资产主表'];old_last=max(r for r,v in read_source_rows(ws))
    for x in REPORT:
        if x['domain']!=domain:continue
        aid=x['asset_id'];row=next((r for r,v in read_source_rows(ws) if v[0]==aid),max(r for r,v in read_source_rows(ws))+1)
        ground=x['model_variants']['ground'];preview=x['model_variants']['preview'];dims=x['godot_dimensions_m'];logical=x['logical_id']
        if x['kind'] in ['bat','blade','axe']:logical='bp_'+logical.removeprefix('weapon_')
        if x['kind']=='scroll':logical='blueprint_shared'
        fields={1:aid,2:x['name'],3:'道具' if domain=='props' else '武器',4:'melee' if x['kind'] in ['bat','blade','axe'] else ('bullet_module' if x['logical_id'].startswith('mod_') else 'item'),
            5:logical,6:'root',8:'top3d',9:'tier_0_static',10:'地面/图标；preview独立源；'+','.join(x['content_ids']),11:'已优化并正式接入',12:'P1',13:'v001',
            14:f'Godot XYZ={dims}m；ground {ground["triangles"]}三角；公共色盘；原点{x["origin"]}',15:ground['prefab'],
            16:'src/world3d/ItemModelFactory3D.gd; src/combat3d/WeaponModel3D.gd; docs/v0.1/design/物品道具双模型制作规范.md',17:x['name']+';'+','.join(x['content_ids']),
            20:ground['prefab_sha256'],21:'Codex',22:datetime(2026,10,14),23:'用户风格延续；程序制作原创组件；项目公共色盘',24:x['name']+'_中文资产管理',
            25:'高精度源='+preview['source_blend']+'；地面源='+ground['source_blend']+'；清单包含两角色哈希与派生关系；Godot实际工厂/图标/28内容映射通过；'+('近战保持握点、伤害归原玩法' if x['kind'] in ['bat','blade','axe'] else '拾取碰撞归GroundLootPickup3D')}
        if ws.cell(row,1).value is None:
            for c in range(1,26):ws.cell(row,c)._style=copy(ws.cell(old_last,c)._style)
        for c,v in fields.items():ws.cell(row,c,v)
        x['status']='已优化并正式接入';x['validation']={'runtime_gallery':True,'content_mapping':True,'ground_budget':True,'notes':'see development/2026-10-14_remaining_item_models.md'}
        folder=(ROOT/ground['source_blend']).parent.parent
        write(folder/'asset_contract.json',x);write(folder/'source_manifest.json',x)
        write((ROOT/ground['prefab']).parent.parent/'asset_manifest.json',x)
        for role,variant in x['model_variants'].items():
            path_source=ROOT/variant['source_blend'];entry={'path':variant['source_blend'],'bytes':path_source.stat().st_size,'family':variant['source_blend'].replace('_v001',''),'version':[1],'role':role,'sha256':variant['sha256']}
            at=next((i for i,s in enumerate(sources['sources']) if s['path']==entry['path']),None)
            if at is None:sources['sources'].append(entry)
            else:sources['sources'][at]=entry
    last=max(r for r,v in read_source_rows(ws))
    for row,v in read_source_rows(ws):ws.cell(row,18,dedupe_key_formula(row));ws.cell(row,19,dedupe_result_formula(row,last))
    for a in ['A6','C6','E6','G6','B10','C10','B11','C11']:
        cell=wb['总览'][a]
        if isinstance(cell.value,str):cell.value=cell.value.replace('$'+str(old_last),'$'+str(last))
    for dv in ws.data_validations.dataValidation:
        ranges=[]
        for rr in dv.sqref.ranges:
            rr=copy(rr)
            if rr.min_row==FIRST_DATA_ROW:rr.max_row=last
            ranges.append(str(rr))
        dv.sqref=MultiCellRange(' '.join(ranges))
    ws.auto_filter.ref=f'A{HEADER_ROW}:Y{last}'
    log=wb['域变更日志'];versions=[int(str(a[0].value).split('.')[-1]) for a in log if re.fullmatch(r'v0\.1\.\d+',str(a[0].value))]
    log.append([f'v0.1.{max(versions or [0])+1}','2026-10-14','剩余程序化物品20模型/28内容ID','双模型/GLB/PackedScene/调用链','100–200三角；蓝图共用；真实渲染及映射通过','已优化并正式接入','Codex'])
    for row,v in read_source_rows(ws):
        if v[0] in target:baseline['assets'][v[0]]={'v':_row_digest(v),'c':v[2],'d':domain}
    wb.save(index.domain_for_key(domain).path)
baseline['asset_count']=len(baseline['assets'])
union=[rv for wb in books.values() for rv in read_source_rows(wb['资产主表'])]
baseline['column_digests']={str(c):col_digest(union,c) for c in CONTENT_COLUMNS};baseline['category_counts']=dict(sorted(Counter(v[2] for _,v in union).items()))
assert all(baseline['assets'][k]==v for k,v in old['assets'].items() if k not in target)
write(bp,baseline);write(source_index_path,sources)
# Separate transaction: actual runtime prefabs, one row per stable path.
for domain,sheet in [('props','3D-物品'),('weapons','3D-武器')]:
    path=index.domain_for_key(domain).path;wb=openpyxl.load_workbook(path);ws=wb[sheet]
    for x in REPORT:
        if x['domain']!=domain:continue
        # Existing melee prototype rows become the new ground entries; no second live claim.
        for role,v in x['model_variants'].items():
            row=next((r for r in range(5,ws.max_row+1) if ws.cell(r,3).value==v['prefab']),None)
            if row is None and role=='ground':row=next((r for r in range(5,ws.max_row+1) if ws.cell(r,1).value==x['asset_id'] and 'melee_3d/' in str(ws.cell(r,3).value)),None)
            if row is None:row=ws.max_row+1
            values=[x['asset_id'],x['name']+' / '+role,v['prefab'],v['glb'],v['source_blend'],'纯视觉；游戏规则由原系统持有','src/world3d/ItemModelFactory3D.gd; src/combat3d/WeaponModel3D.gd','无','外部玩法','外部拾取球/原近战扫掠',str(x['godot_dimensions_m'])+'；'+str(v['triangles'])+'三角',x['origin']+'；Godot Y上，近战-Z前',','.join(x['content_ids']),'已优化并正式接入','v001','公共材质/外链色盘；source_manifest记录哈希和双角色']
            for c,val in enumerate(values,1):ws.cell(row,c)._style=copy(ws.cell(5,c)._style);ws.cell(row,c,val)
    ws.auto_filter.ref=f'A4:P{ws.max_row}';baseline['sheet_digests'][sheet]=sheet_digest(ws);wb.save(path)
write(bp,baseline)
backup(ROOT/'assets/registry/ledgers/ShellStorm2_物品道具账本_v001.xlsx')
write(OUT/'ledger_transaction.json',{'asset_ids':sorted(target),'new_asset_count':len(baseline['assets'])-len(old['assets']),'unrelated_fingerprints_preserved':True,'prefab_pages':['3D-物品','3D-武器']})
print('REMAINING_LEDGER_REGISTERED',len(target))
