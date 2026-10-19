"""Add sourced bird, fan/blower and wood footsteps for actual world playback."""
from pathlib import Path
from copy import copy
import csv,datetime,hashlib,json,re,shutil,subprocess,sys,zipfile
import numpy as np
import openpyxl

ROOT=Path(__file__).resolve().parents[2]
sys.path.insert(0,str(ROOT/'scripts'))
from workspace_paths import work_path
sys.path.insert(0,str(Path(__file__).parent))
from build_natural_world_audio import repeat_crossfade
SESSION=work_path('sessions/2026-10-19_world_audio_runtime/assets.json').parent
BANK=ROOT/'art_source/audio/sfx/recorded_world_library/source'
OUT=ROOT/'src/assets/audio/sfx'
sha=lambda p:hashlib.sha256(p.read_bytes()).hexdigest()
def save(p,v,indent=2):p.write_text(json.dumps(v,ensure_ascii=False,indent=indent)+'\n',encoding='utf-8')

def main():
 cp=OUT/'sfx_catalog.json';c=json.loads(cp.read_text(encoding='utf-8'));assert 'outdoor_bird' not in c['sound_sets']
 source_index=json.loads((BANK/'source_index.json').read_text(encoding='utf-8'))
 bird=BANK/'syncopika_birdchirping071414.wav';shutil.copy2(SESSION/'intake/birdchirping071414.wav',bird)
 source_index[bird.name]={'path':bird.relative_to(ROOT).as_posix(),'sha256':sha(bird),'credit':'syncopika','source_url':'https://opengameart.org/content/bird-chirping-sounds','license':'CC0-1.0'}
 impact=zipfile.ZipFile(work_path('sessions/2026-10-19_natural_world_audio/intake/impact.zip'))
 for var in range(4):
  p=BANK/f'impact_footstep_wood_{var:03d}.ogg';p.write_bytes(impact.read(f'Audio/footstep_wood_{var:03d}.ogg'))
  source_index[p.name]={'path':p.relative_to(ROOT).as_posix(),'sha256':sha(p),'credit':'Kenney','source_url':'https://kenney.nl/assets/impact-sounds','license':'CC0-1.0'}
 plans=[]
 for var in range(4):plans.append(('footstep_wood','木地板脚步',var,BANK/f'impact_footstep_wood_{var:03d}.ogg',0,.38,False,14000,-20,'player.footstep.wood'))
 for var in range(3):plans.append(('outdoor_bird','户外鸟叫',var,bird,var*.92,.95,False,15000,-20,'amb.outdoor_bird'))
 plans.extend([('amb_fan','风扇运转',0,BANK/'fan.wav',0,8,True,4500,-20,'amb.fan'),('amb_blower','鼓风机气流',0,BANK/'vent.mp3',0,8,True,9000,-19,'amb.blower')])
 assets=[]
 for key,label,var,raw,start,duration,loop,cutoff,peak,event in plans:
  data=subprocess.check_output(['ffmpeg','-v','error','-ss',str(start),'-i',str(raw),'-ac','1','-ar','44100','-af',f'highpass=f=70,lowpass=f={cutoff}','-f','f32le','-'])
  x=np.frombuffer(data,dtype='<f4').astype(np.float64);x-=x.mean();n=round(duration*44100)
  if loop:
   f=22050;x=repeat_crossfade(x,n+f);w=np.linspace(0,1,f);x=np.concatenate([x[f:-f],x[-f:]*(1-w)+x[:f]*w])
   gain=min(10**(-29/20)/max(float(np.sqrt(np.mean(x*x))),1e-9),10**(peak/20)/max(abs(x)))
  else:
   x=x[:n];active=np.flatnonzero(abs(x)>max(float(np.max(abs(x)))*.02,.0001));trim=max(0,int(active[0])-44) if len(active) else 0
   x=x[trim:];x=np.pad(x,(0,max(0,n-len(x))))[:n];x[:44]*=np.linspace(0,1,44);x[-1103:]*=np.linspace(1,0,1103);gain=10**(peak/20)/max(float(np.max(abs(x))),1e-9)
  x*=gain;temp=SESSION/'raw.f32';temp.write_bytes(x.astype('<f4').tobytes())
  asset_id='AUD-SFX-'+key.upper().replace('_','-')+('' if var==0 else f'-VAR{var+1:02d}')
  source=ROOT/'art_source/audio/sfx'/asset_id/'source';source.mkdir(parents=True,exist_ok=True)
  filename=key+('' if var==0 else f'_alt{var+1:02d}')+'_v001';master=source/(filename+'.wav');runtime=OUT/(filename+'.ogg')
  subprocess.run(['ffmpeg','-v','error','-y','-f','f32le','-ar','44100','-ac','1','-i',str(temp),'-c:a','pcm_s24le',str(master)],check=True)
  subprocess.run(['ffmpeg','-v','error','-y','-i',str(master),'-c:a','libvorbis','-q:a','6',str(runtime)],check=True)
  record=source_index[raw.name]
  a={'asset_id':asset_id,'key':key,'variant':var,'source_path':master.relative_to(ROOT).as_posix(),'runtime_path':'res://'+runtime.relative_to(ROOT).as_posix(),'version':1,'duration_sec':duration,'source_format':'PCM24 WAV','runtime_format':'OGG Vorbis q6','sample_rate_hz':44100,'source_channels':1,'runtime_channels':1,'sha256':sha(runtime),'source_sha256':sha(master),'production_revision':'world-environment-r6','source_method':'Sourced recordings / real-world environment playback','license_source':record['source_url']+' | '+record['credit'],'license_scope':record['license'],'status':'imported_pending_listening','spatial_mode':'world_3d','loop':loop,'peak_target_db':peak,'event_ids':[event],'sound_set_id':'SFXSET-'+key.upper().replace('_','-')}
  assets.append(a);save(source/'recipe.json',{'renderer':'tools/asset_pipeline/build_world_environment_assets.py','asset':a,'input':record,'start_sec':start,'duration_sec':duration,'lowpass_hz':cutoff,'gain':gain,'loop_crossfade_sec':.5 if loop else 0})
  if key not in c['sound_sets']:
   c['sound_sets'][key]={'id':a['sound_set_id'],'paths':[],'loop':loop,'spatial_mode':'world_3d','bus':'Ambience' if key.startswith('amb_') or key=='outdoor_bird' else 'SFX','max_voices':6 if loop else 1 if key=='outdoor_bird' else 4,'priority':'P2'}
  c['sound_sets'][key]['paths'].append(a['runtime_path']);c['assets'].append(a)
  c['events'][event]={'sound_set':key,'owner':'WorldAudioEnvironment3D','trigger_fact':'audible support step / nearby enabled device / outdoor bird interval','ai_signal':'unchanged','binding_status':'domain_bound','acceptance_case':'verify_world_environment_audio_flow'}
 assert len(assets)==9
 c['events']['amb.rooftop_wind']={'sound_set':'amb_expedition_open','owner':'WorldAudioEnvironment3D','binding_status':'domain_bound','acceptance_case':'verify_world_environment_audio_flow'}
 bound_keys={'footstep_concrete','footstep_metal','footstep_debris','footstep_wood','light_on','light_off','amb_expedition_open','amb_fan','amb_blower','outdoor_bird'}
 for event in c['events'].values():
  if event['sound_set'] in bound_keys:event.update(binding_status='domain_bound',acceptance_case='verify_world_environment_audio_flow')
 policy=c['world_style_policy'];policy['natural_world_sets'].extend(['footstep_wood','outdoor_bird','amb_fan','amb_blower']);policy.update(changed_assets=93,environment_loops=9,revision='natural-world-r5 + world-environment-r6')
 c['revision']='world-environment-r6 + natural-world-r5 + gun-families-r4 + system cues'
 save(cp,c);save(ROOT/'art_source/audio/sfx/world_audio_policy.json',policy);save(BANK/'source_index.json',source_index)
 with (BANK/'CREDITS.md').open('a',encoding='utf-8') as f:f.write('\n- Bird chirps recorded by syncopika, CC0-1.0: https://opengameart.org/content/bird-chirping-sounds ; edits: excerpts, mono, filtering, fades, normalization. Wood footsteps: Kenney Impact Sounds, CC0-1.0.\n')
 mf=OUT/'audio_manifest.csv'
 with mf.open(encoding='utf-8',newline='') as f:r=csv.DictReader(f);fields=r.fieldnames;rows=list(r)
 for a in assets:rows.append({'key':a['key'],'asset_id':a['asset_id'],'source_master':a['source_path'],'runtime_file':Path(a['runtime_path']).name,'status':'正式环境触发已接入；听感待验','usage':a['key'],'duration_sec':a['duration_sec'],'sample_rate_hz':44100,'master_channels':1,'runtime_channels':1,'runtime_resource_path':a['runtime_path'],'sha256':a['sha256'],'sound_set_id':a['sound_set_id'],'event_ids':';'.join(a['event_ids'])})
 with mf.open('w',encoding='utf-8',newline='') as f:w=csv.DictWriter(f,fieldnames=fields);w.writeheader();w.writerows(rows)
 update_ledger(assets,c,bound_keys)
 p=ROOT/'art_source/audio/sfx/library_contract.json';d=json.loads(p.read_text(encoding='utf-8'));d.update(design_revision='AUDIO-SFX r5',production_revision='world-environment-r6',world_style_policy=policy)
 for key in ['footstep_wood','outdoor_bird','amb_fan','amb_blower']:
  own=[a for a in assets if a['key']==key];d['sounds'].append({'key':key,'name':key,'variants':len(own),'events':own[0]['event_ids'],'production_override':{'renderer':'tools/asset_pipeline/build_world_environment_assets.py','revision':'world-environment-r6','source_of_current_duration_and_version':'src/assets/audio/sfx/sfx_catalog.json'}})
 save(p,d);save(SESSION/'assets.json',{'new_assets':assets,'bound_keys':sorted(bound_keys),'asset_count':166,'sound_set_count':108,'audio_ledger_count':178})
 print('Installed 9 assets / 4 sets; 166 OGGs, 108 sets; audio ledger 178')

def update_ledger(assets,c,bound_keys):
 from split_asset_ledger import read_source_rows,_row_digest,CONTENT_COLUMNS,col_digest,dedupe_key_formula,dedupe_result_formula
 from openpyxl.worksheet.cell_range import MultiCellRange
 idx=json.loads((ROOT/'assets/registry/ledger_index.json').read_text(encoding='utf-8'));domain=next(d for d in idx['domains'] if d['key']=='audio');p=ROOT/idx['ledger_dir']/domain['file'];w=openpyxl.load_workbook(p);s=w[idx['asset_sheet']];old_end=s.max_row;ids={s.cell(r,1).value:r for r in range(6,old_end+1)};template=ids['AUD-SFX-FOOTSTEP-CONCRETE'];vals=[v.value for v in s[template]]
 labels={'footstep_wood':'木地板脚步','outdoor_bird':'户外鸟叫','amb_fan':'风扇运转','amb_blower':'鼓风机气流'}
 for a in assets:
  assert a['asset_id'] not in ids
  r=s.max_row+1;ids[a['asset_id']]=r
  for col,value in enumerate(vals,1):s.cell(r,col,value);s.cell(r,col)._style=copy(s.cell(template,col)._style)
  parent=None if a['variant']==0 else next(v['asset_id'] for v in assets if v['key']==a['key'] and v['variant']==0)
  row={1:a['asset_id'],2:labels[a['key']]+('' if a['variant']==0 else f'（变体{a["variant"]+1:02d}）'),3:'音效',4:'ambience' if a['loop'] else 'sfx',5:a['key'],6:'root' if a['variant']==0 else f'var{a["variant"]+1:02d}',7:parent,10:';'.join(a['event_ids']),11:'原型已接入',12:'P2',13:'v001',14:f'44100Hz / mono / PCM24 master / Vorbis q6 / {a["duration_sec"]}s',15:a['runtime_path'].removeprefix('res://'),16:a['source_path'],17:a['key'],20:a['sha256'],21:'Codex audio integration',22:datetime.datetime(2026,10,19),23:a['license_source']+' | '+a['license_scope'],24:a['source_path']}
  for col,value in row.items():s.cell(r,col,value)
 changed=set()
 for r in range(6,s.max_row+1):
  if s.cell(r,5).value in bound_keys and s.cell(r,11).value!='弃用':
   s.cell(r,11,'原型已接入');s.cell(r,22,datetime.datetime(2026,10,19));s.cell(r,25,'WorldAudioEnvironment3D / RoomLightSwitch3D 正式场景触发已接入；验收 verify_world_environment_audio_flow；人工听感/Android待验');changed.add(s.cell(r,1).value)
 last=s.max_row
 for r in range(6,last+1):s.cell(r,18,dedupe_key_formula(r));s.cell(r,19,dedupe_result_formula(r,last))
 for sheet in w:
  if sheet==s:continue
  for row in sheet:
   for cell in row:
    if cell.data_type=='f' and '资产主表' in cell.value:cell.value=re.sub(r'(\$?[A-Y]\$?)'+str(old_end)+r'\b',lambda m:m[1]+str(last),cell.value)
 for dv in s.data_validations.dataValidation:dv.sqref=MultiCellRange(' '.join(re.sub(r'(:[A-Z]+)'+str(old_end)+r'\b',lambda m:m[1]+str(last),str(v)) for v in dv.sqref.ranges))
 s.auto_filter.ref=f'A5:Y{last}'
 log=w['域变更日志'];r=log.max_row+1
 for col,value in enumerate(['v0.2.3',datetime.datetime(2026,10,19),'环境音正式触发','新增9条、4组；四类脚步/鸟/风扇/鼓风机','天台/户外风声与鸟叫、真实设备位置循环、地板脚步、电灯点击、楼梯间回声正式接入；复用既有风声与开关素材。','系统提示干声；七类枪种声音及AI刺激不变；人工听感/Android待验。','Codex'],1):log.cell(r,col,value);log.cell(r,col)._style=copy(log.cell(6,col)._style)
 log.cell(r,2).number_format='yyyy-mm-dd';log.row_dimensions[r].height=80;w.save(p)
 bp=ROOT/'assets/registry/ledger_split_baseline.json';b=json.loads(bp.read_text(encoding='utf-8'))
 for _,v in read_source_rows(s):
  if v[0] in changed:b['assets'][v[0]]={'v':_row_digest(v),'c':v[2],'d':'audio'}
 allrows=[]
 for domain in idx['domains']:
  other=openpyxl.load_workbook(ROOT/idx['ledger_dir']/domain['file']);allrows.extend(read_source_rows(other[idx['asset_sheet']]))
 b['asset_count']=len(b['assets']);b['column_digests']={str(c):col_digest(allrows,c) for c in CONTENT_COLUMNS};b['category_counts']={'音效':178,**{k:v for k,v in b['category_counts'].items() if k!='音效'}};save(bp,b,1)

if __name__=='__main__':main()
