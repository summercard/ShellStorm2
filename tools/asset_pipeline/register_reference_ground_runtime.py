"""Register actual ground exports; keep runtime acceptance pending until requested."""
import hashlib
import json
import re
import shutil
import sys
from collections import Counter
from copy import copy, deepcopy
from datetime import datetime
from pathlib import Path
import openpyxl

ROOT = Path(__file__).resolve().parents[2]
sys.path[:0] = [str(ROOT / 'scripts'), str(ROOT / 'tools/asset_pipeline')]
from workspace_paths import work_path
from ledger_registry import LedgerIndex
from split_asset_ledger import read_source_rows, _row_digest, col_digest, CONTENT_COLUMNS, sheet_digest

def write(path, data):
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2) + '\n', encoding='utf8')

BACKUP = ROOT.parent / 'ShellStorm2-backup/2026-10-14_item_ground_import/before_edits'
def backup(path):
    dest = BACKUP / path.relative_to(ROOT)
    dest.parent.mkdir(parents=True, exist_ok=True)
    if not dest.exists():
        shutil.copy2(path, dest)

report = json.loads(work_path('sessions/2026-10-14_item_ground_import/export_report.json').read_text('utf8'))
index = LedgerIndex.load(ROOT)
bp = ROOT / 'assets/registry/ledger_split_baseline.json'
backup(bp)
baseline = json.loads(bp.read_text('utf8'))
old = deepcopy(baseline)
books = {d.key: openpyxl.load_workbook(d.path) for d in index.domains}
status = 'exported_pending_godot_validation'
target_ids = {x['asset_id'] for x in report}

# Transaction 1: update existing identities and content navigation, no new asset IDs.
for x in report:
    domain = 'props' if x['asset_id'].startswith('ITM-') else 'weapons'
    ws = books[domain]['资产主表']
    row = next(r for r, v in read_source_rows(ws) if v[0] == x['asset_id'])
    fields = {11: status, 15: x['prefab'], 20: x['prefab_sha256'], 22: datetime(2026, 10, 14),
        10: 'ground地面简模已接线；preview高模待导入；枪上装配保持原表现',
        16: 'src/world3d/ItemModelFactory3D.gd; src/world3d/GroundLootPickup3D.gd; docs/v0.1/design/物品道具双模型制作规范.md',
        25: f"ground={x['ground_triangles']}三角；GLB={x['glb']}；ground源={x['ground_source']} SHA256={x['ground_sha256']}；preview源={x['preview_source']} SHA256={x['preview_sha256']}；源不变；资源已导入且地面入口已接线，运行/视觉验收未执行，未签署正式可用。"}
    for col, value in fields.items():
        ws.cell(row, col, value)
    source_folder = (ROOT / x['ground_source']).parent.parent
    for name in ('asset_contract.json', 'source_manifest.json'):
        path = source_folder / name
        backup(path)
        data = json.loads(path.read_text('utf8'))
        data['model_variants']['ground'].update(stable_glb=x['glb'], stable_prefab=x['prefab'],
            glb_sha256=x['glb_sha256'], prefab_sha256=x['prefab_sha256'], status=status)
        data['runtime_integration'] = 'ground已导出导入并接线；preview及枪上装配未切换；运行/视觉验收未执行'
        data['status'] = status
        data['function_note'] = 'ground仅视觉；拾取/弹跳/碰撞由GroundLootPickup3D持有，倍率0.504不烘入源'
        write(path, data)
for domain in ('props', 'weapons'):
    wb = books[domain]
    path = index.domain_for_key(domain).path
    backup(path)
    log = wb['域变更日志']
    versions = [int(str(row[0].value).split('.')[-1]) for row in log if re.fullmatch(r'v0\.1\.\d+', str(row[0].value))]
    log.append([f'v0.1.{max(versions or [0])+1}', '2026-10-14', '地面简模正式路径与代码接线',
        '8件物品ground角色', 'GLB/PackedScene/公共色盘导入；同一AssetID；预览待导入', status, 'Codex'])
    for row, values in read_source_rows(wb['资产主表']):
        if values[0] in target_ids:
            baseline['assets'][values[0]] = {'v': _row_digest(values), 'c': values[2], 'd': domain}
    wb.save(path)
navpath = ROOT / 'assets/registry/ledgers/ShellStorm2_物品道具账本_v001.xlsx'
backup(navpath)
nav = openpyxl.load_workbook(navpath)
by_id = {x['asset_id']: x for x in report}
for wb, sheet in ((nav, '物品道具'), (books['weapons'], '背包物品调用索引')):
    ws = wb[sheet]
    for row in ws:
        aid = row[13].value if len(row) > 13 else None
        if aid in by_id:
            x = by_id[aid]
            ws.cell(row[0].row, 17, '2026-10-14：地面简模已导入且接线；预览高模待导入；运行/视觉验收未执行。ground=' + x['prefab'])
nav.save(navpath)
books['weapons'].save(index.domain_for_key('weapons').path)
union = [rv for wb in books.values() for rv in read_source_rows(wb['资产主表'])]
baseline['column_digests'] = {str(c): col_digest(union, c) for c in CONTENT_COLUMNS}
baseline['category_counts'] = dict(sorted(Counter(v[2] for _, v in union).items()))
if any(baseline['assets'][k] != v for k, v in old['assets'].items() if k not in target_ids):
    raise RuntimeError('Unrelated baseline changed')
write(bp, baseline)

# Transaction 2: register only real PackedScenes, update only their sheet digests.
for domain, sheet in (('props', '3D-物品'), ('weapons', '3D-武器')):
    path = index.domain_for_key(domain).path
    wb = openpyxl.load_workbook(path)
    ws = wb[sheet]
    for x in report:
        if ('props' if x['asset_id'].startswith('ITM-') else 'weapons') != domain:
            continue
        if not (ROOT / x['prefab']).is_file():
            raise RuntimeError('Missing prefab')
        row = next((r for r in range(5, ws.max_row + 1) if ws.cell(r, 3).value == x['prefab']), ws.max_row + 1)
        contract = json.loads(((ROOT / x['ground_source']).parent.parent / 'asset_contract.json').read_text('utf8'))
        dims = contract['dimensions_m']
        vals = [x['asset_id'], x['name'] + '（地面）', x['prefab'], x['glb'], x['ground_source'],
            '仅ground视觉；同一物品身份，外部负责拾取/弹跳/收集', 'src/world3d/ItemModelFactory3D.gd; src/world3d/GroundLootPickup3D.gd',
            'Prefab无；外部拾取触发', 'GroundLootPickup3D', '外部SphereShape3D；无实体阻挡',
            f"{dims['width']}×{dims['depth']}×{dims['height']}m；{x['ground_triangles']}三角；外部倍率0.504",
            contract['origin_and_facing'], '世界地面掉落/拾取；非背包预览与枪上装配', status, 'v001',
            '源保持；公共场景色盘；资源导入完成，未执行运行/视觉验收']
        for col, value in enumerate(vals, 1):
            ws.cell(row, col)._style = copy(ws.cell(5, col)._style)
            ws.cell(row, col, value)
    ws.auto_filter.ref = f'A4:P{ws.max_row}'
    ws.cell(2, 1, '一行一个独立PackedScene；地面模型使用ground角色，状态以各行记录为准。')
    wb.save(path)
    baseline['sheet_digests'][sheet] = sheet_digest(ws)
write(bp, baseline)
write(work_path('sessions/2026-10-14_item_ground_import/ledger_transaction.json'),
    {'assets': sorted(target_ids), 'status': status, 'prefab_pages': ['3D-物品', '3D-武器'], 'unrelated_asset_fingerprints_preserved': True})
print('GROUND_RUNTIME_REGISTERED', len(report))
