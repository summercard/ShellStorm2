"""Second ledger transaction: update the character animation/state specialty sheet."""

from copy import copy
from pathlib import Path
import json
import shutil
import sys

import openpyxl

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
sys.path.insert(0, str(HERE))
from split_asset_ledger import sheet_digest  # noqa: E402

BOOK = ROOT / "assets/registry/ledgers/ShellStorm2_角色账本_v001.xlsx"
BASELINE = ROOT / "assets/registry/ledger_split_baseline.json"
BACKUP = ROOT / "_scratch/seat_state_sheet_before_20260929"


def main() -> None:
    BACKUP.mkdir(parents=True, exist_ok=True)
    shutil.copy2(BOOK, BACKUP / BOOK.name)
    shutil.copy2(BASELINE, BACKUP / BASELINE.name)
    workbook = openpyxl.load_workbook(BOOK)
    sheet = workbook["动画与状态"]
    assert sheet["B18"].value is None and sheet["A19"].value == "层级"
    assert "八态" in sheet["A3"].value
    sheet["A3"] = sheet["A3"].value.replace("八态", "九态")
    sheet["E6"] = sheet["E6"].value + ",seated"
    sheet["E7"] = sheet["E7"].value + ",seated"
    sheet["A36"] = sheet["A36"].value.replace("八态", "九态")
    sheet["K38"] = sheet["K38"].value.replace("八态", "九态")
    sheet.unmerge_cells("A18:L18")
    seated = [
        "顶层", "seated", "座椅乘坐 / 滑行", "idle/moving，经 E 交互且椅子空闲",
        "idle,moving,locked,hurt,falling,dead；E 安全离座",
        "v022：原始 Blender 动画源新增 1.2 秒坐姿轻摆循环；不驱动椅子位移。",
        "头部随躯干微调，耳朵保持轻微呼吸。",
        "双手放低稳定姿态；乘坐时不射击或挥砍。",
        "右手与 WeaponSocket 随坐姿，不触发持枪动作覆盖。",
        "随坐姿缓动", "E 离座提示；场景墙/地面实体碰撞",
        "Player3DSeatedState 驱动；PushableSeat3D 刚体受方向输入施力，乘客跟随座点；退出前做地面与胶囊空间查询。",
    ]
    for col, value in enumerate(seated, 1):
        target = sheet.cell(18, col)
        target._style = copy(sheet.cell(13, col)._style)
        target.value = value
    sheet.row_dimensions[18].height = sheet.row_dimensions[13].height
    log = workbook["域变更日志"]
    row = log.max_row + 1
    values = [
        "v0.1.2", "2026-09-29", "九态坐姿映射同步", "角色 / 动画与状态",
        "Player3D 八态扩为九态，补 seated 乘坐滑行状态的来源、去向、Blender v022 动作及安全离座契约。",
        "仅改《动画与状态》专表和变更日志；资产主表内容不变，专表无损摘要同步更新。", "Codex",
    ]
    for col, value in enumerate(values, 1):
        log.cell(row, col)._style = copy(log.cell(row - 1, col)._style)
        log.cell(row, col).value = value
    workbook.save(BOOK)
    baseline = json.loads(BASELINE.read_text(encoding="utf-8"))
    baseline["sheet_digests"]["动画与状态"] = sheet_digest(workbook["动画与状态"])
    BASELINE.write_text(json.dumps(baseline, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")
    print("SEATED_STATE_SHEET_UPDATED")


if __name__ == "__main__":
    main()
