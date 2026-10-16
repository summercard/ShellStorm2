"""Register item palette and source rebinding, preserving unrelated ledger fingerprints."""
from pathlib import Path
from copy import copy, deepcopy
from collections import Counter
import sys,json,re,os
import openpyxl
from openpyxl.worksheet.cell_range import MultiCellRange
ROOT=Path(__file__).resolve().parents[2]
sys.path[:0]=[str(ROOT/'scripts'),str(ROOT/'tools/asset_pipeline')]
from workspace_paths import work_path
from item_palette import SOURCE,TARGET,RECIPE,IMPORT_SCRIPT,ASSET_ID,sha,MATERIAL_ROLES,MATERIAL_PATHS,SATURATION_MULTIPLIER,VALUE_MULTIPLIER,ICON_FILL_ENERGY
from ledger_registry import LedgerIndex
from split_asset_ledger import read_source_rows,_row_digest,col_digest,CONTENT_COLUMNS,sheet_digest,dedupe_key_formula,dedupe_result_formula,FIRST_DATA_ROW,HEADER_ROW
SESSION=os.environ.get('ITEM_PALETTE_SESSION','2026-10-15_item_palette')
OUT=work_path('sessions/'+SESSION+'/ledger_result.json').parent
validation=json.loads((OUT/'runtime_result.json').read_text(encoding='utf8'));assert validation['passed']
sources=json.loads((OUT/'blender_result.json').read_text(encoding='utf8'))['sources']
hashes={s['old_sha256']:s['sha256'] for s in sources}
runtime=json.loads((OUT/'runtime_input.json').read_text(encoding='utf8'))
target={r['asset_id'] for r in runtime['rows']}
palette_contract=dict(path=TARGET,sha256=sha(ROOT/TARGET),recipe=RECIPE,import_script=IMPORT_SCRIPT,saturation_multiplier=SATURATION_MULTIPLIER,value_multiplier=VALUE_MULTIPLIER,material_roles=MATERIAL_ROLES,material_paths=MATERIAL_PATHS,icon_shadow_fill_energy=ICON_FILL_ENERGY,geometry_uv_unchanged=True,validated_on='2026-10-15')
def write(p,data):p.write_text(json.dumps(data,ensure_ascii=False,indent=2)+'\n',encoding='utf8')
def convert(value,key=''):
    if isinstance(value,dict):return {k:convert(v,k) for k,v in value.items()}
    if isinstance(value,list):return [convert(v,key) for v in value]
    if isinstance(value,str):
        if 'before' not in key:
            for a,b in hashes.items():value=value.replace(a,b)
        for a,b in zip(['01_精工金属_紫色骨架','02_细腻哑光_青绿大面','03_清漆反光_紫粉点缀','04_柔和自发光_UI灯光'],MATERIAL_ROLES):value=value.replace(a,b)
        return value.replace(SOURCE,TARGET)
    return value
manifest_count=0
for base in [ROOT/'art_source/items_weapons',ROOT/'assets/art/items_weapons']:
    for p in base.rglob('*.json'):
        if p.name not in ['asset_contract.json','source_manifest.json','asset_manifest.json']:continue
        d=json.loads(p.read_text(encoding='utf8'))
        if d.get('asset_id') not in target:continue
        d=convert(d);d['palette_contract']=palette_contract
        if 'palette_sha256' in d:d['palette_sha256']=sha(ROOT/TARGET)
        if 'source_hashes_unchanged' in d:d['source_hashes_unchanged']=False
        write(p,d);manifest_count+=1
# Current editable sources, including optimized mothers and the reusable item template.
p=ROOT/'assets/registry/current_source_index.json';d=json.loads(p.read_text(encoding='utf8'))
for source in sources:
    at=next((x for x in d['sources'] if x['path']==source['path']),None)
    if at is None:
        at=dict(path=source['path'],family=source['path'],version=[1]);d['sources'].append(at)
    at.update(sha256=source['sha256'],bytes=source['bytes'])
recipe_entry=next((x for x in d['sources'] if x['path']==RECIPE),None)
if recipe_entry is None:
    recipe_entry=dict(path=RECIPE,family='item_palette_recipe',version=[1],role='palette_recipe');d['sources'].append(recipe_entry)
recipe_entry.update(sha256=sha(ROOT/RECIPE),bytes=(ROOT/RECIPE).stat().st_size)
write(p,d)
index=LedgerIndex.load(ROOT)
bp=ROOT/'assets/registry/ledger_split_baseline.json';bl=json.loads(bp.read_text(encoding='utf8'));old=deepcopy(bl)
books={domain.key:openpyxl.load_workbook(domain.path) for domain in index.domains}
modified=[]
for domain in ['props','weapons']:
    wb=books[domain];ws=wb['资产主表'];old_last=max(r for r,v in read_source_rows(ws))
    for r,values in read_source_rows(ws):
        if values[0] not in target:continue
        for c,cell in enumerate(values,1):
            if isinstance(cell,str):ws.cell(r,c,convert(cell))
        ws.cell(r,22,'2026-10-15')
        note=str(ws.cell(r,25).value or '')
        ws.cell(r,25,note+'\n【2026-10-15独立四材质】'+TARGET+'；HSV S/V=原场景×1.40×1.15=1.61封顶1；四个道具专用材质；图标本色补光0.45；高低/优化源同步；场景原材质不变；运行及图标验收通过。')
        modified.append(values[0])
    if domain=='props':
        row=next((r for r,v in read_source_rows(ws) if v[0]==ASSET_ID),old_last+1)
        fields=[ASSET_ID,'物品专用高亮高饱和色板','道具','palette','item_high_visibility_palette','texture','','shared','static','31物品资产/地面/图标/预览/背包','已完成','P1','v001','512×512 RGBA；10×10；sRGB HSV S/V各×1.61（1.40后再增15%）并封顶；H/Alpha不变',TARGET,RECIPE+'; '+IMPORT_SCRIPT,'物品;色板;高亮;高饱和;palette','','',sha(ROOT/TARGET),'Codex','2026-10-15','从项目场景色板按用户要求确定性派生','','只登记贴图资产，不登记虚假Prefab。配方与生成器可重建，Nearest/无mipmap/无损压缩。']
        for c,value in enumerate(fields,1):ws.cell(row,c)._style=copy(ws.cell(old_last,c)._style);ws.cell(row,c,value)
        for i,path in enumerate(MATERIAL_PATHS):
            aid='ITM-MATERIAL-'+['METAL','MATTE','CLEARCOAT','EMISSIVE'][i]
            target.add(aid)
            row=next((r for r,v in read_source_rows(ws) if v[0]==aid),max(r for r,v in read_source_rows(ws))+1)
            values=[aid,MATERIAL_ROLES[i],'道具','material','item_material_'+str(i+1),'material','','shared','static','31物品资产；世界与预览；图标派生实例','已完成','P1','v001','四角色独立资源；物品专用色板；图标补光能量下限0.45',path,IMPORT_SCRIPT+'; src/ui/ItemIconMaterialPreset.gd','物品;独立材质;图标;补光','','',sha(ROOT/path),'Codex','2026-10-15','项目原创材质参数','','不引用场景材质；图标使用surface override副本，不修改世界材质或新增角色。']
            for c,value in enumerate(values,1):ws.cell(row,c)._style=copy(ws.cell(old_last,c)._style);ws.cell(row,c,value)
    last=max(r for r,v in read_source_rows(ws))
    for r,v in read_source_rows(ws):ws.cell(r,18,dedupe_key_formula(r));ws.cell(r,19,dedupe_result_formula(r,last))
    for addr in ['A6','C6','E6','G6','B10','C10','B11','C11']:
        cell=wb['总览'][addr]
        if isinstance(cell.value,str):cell.value=cell.value.replace('$'+str(old_last),'$'+str(last))
    for dv in ws.data_validations.dataValidation:
        ranges=[]
        for rr in dv.sqref.ranges:
            rr=copy(rr)
            if rr.min_row==FIRST_DATA_ROW:rr.max_row=last
            ranges.append(str(rr))
        dv.sqref=MultiCellRange(' '.join(ranges))
    ws.auto_filter.ref=f'A{HEADER_ROW}:Y{last}'
    log=wb['域变更日志'];versions=[int(str(row[0].value).split('.')[-1]) for row in log if re.fullmatch(r'v0\.1\.\d+',str(row[0].value))]
    log.append([f'v0.1.{max(versions or [0])+1}','2026-10-15','物品独立四材质与图标补光','HSV S/V×1.61；31资产；51GLB；64Blender；四个独立材质','同灯光开关补光对比通过，世界材质无污染','已完成','Codex'])
    for r,v in read_source_rows(ws):
        if v[0] in target|{ASSET_ID}:bl['assets'][v[0]]={'v':_row_digest(v),'c':v[2],'d':domain}
    wb.save(index.domain_for_key(domain).path)
# Separate existing-Prefab metadata transaction; no new scene rows for a texture.
for domain,sheet in [('props','3D-物品'),('weapons','3D-武器')]:
    wb=books[domain];ws=wb[sheet]
    for r in range(5,ws.max_row+1):
        if ws.cell(r,1).value in target:
            for c in range(1,ws.max_column+1):
                if isinstance(ws.cell(r,c).value,str):ws.cell(r,c,convert(ws.cell(r,c).value))
            ws.cell(r,16,str(ws.cell(r,16).value or '')+'；当前物品独立四材质、S/V×1.61，'+TARGET)
    bl['sheet_digests'][sheet]=sheet_digest(ws)
    wb.save(index.domain_for_key(domain).path)
bl['asset_count']=len(bl['assets'])
union=[rv for wb in books.values() for rv in read_source_rows(wb['资产主表'])]
bl['column_digests']={str(c):col_digest(union,c) for c in CONTENT_COLUMNS}
bl['category_counts']=dict(sorted(Counter(v[2] for _,v in union).items()))
assert all(bl['assets'][k]==v for k,v in old['assets'].items() if k not in target|{ASSET_ID})
write(bp,bl)
# Add the shared palette subclass to the existing ITM coding rule.
p=index.master_path
wb=openpyxl.load_workbook(p);ws=wb['分类与编码']
for row in ws:
    if row[0].value=='ITM':
        if 'palette' not in str(row[2].value):row[2].value=str(row[2].value)+' / palette（物品共享色板，assets/art/shared/palette/）'
        if 'material' not in str(row[2].value):row[2].value=str(row[2].value)+' / material（物品独立四材质）'
wb.save(p)
write(OUT/'ledger_result.json',dict(updated_assets=sorted(modified),new_asset=ASSET_ID,manifest_count=manifest_count,source_count=len(sources),unrelated_fingerprints_preserved=True))
print('ITEM_PALETTE_LEDGER_OK',len(modified),manifest_count,len(sources))
