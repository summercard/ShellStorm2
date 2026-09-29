"""Record v004 swivel-top / rolling-base 99F chair revisions."""

from copy import copy
from datetime import datetime
from pathlib import Path
import hashlib
import json
import shutil
import sys

import openpyxl

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
sys.path.insert(0, str(HERE))
from split_asset_ledger import _row_digest, col_digest, read_source_rows  # noqa: E402

BOOKS = ROOT / "assets/registry/ledgers"
BOOK = BOOKS / "ShellStorm2_道具账本_v001.xlsx"
BASELINE = ROOT / "assets/registry/ledger_split_baseline.json"
BACKUP = ROOT / "_scratch/swivel_chairs_ledger_before_20260929"


def main() -> None:
    workbook = openpyxl.load_workbook(BOOK)
    sheet = workbook["资产主表"]
    targets = {
        "PRP-BASE-WORKSHOP-STOOL-3D": ("workshop_stool", "维修圆凳"),
        "PRP-BASE-MISSION-COMMAND-CHAIR-3D": ("mission_command_chair", "战术指挥椅"),
    }
    rows = {str(sheet.cell(r, 1).value): r for r in range(6, sheet.max_row + 1)}
    for asset_id in targets:
        assert asset_id in rows
        assert sheet.cell(rows[asset_id], 13).value != "v004", asset_id
    BACKUP.mkdir(parents=True, exist_ok=True)
    shutil.copy2(BOOK, BACKUP / BOOK.name)
    shutil.copy2(BASELINE, BACKUP / BASELINE.name)
    for asset_id, (slug, label) in targets.items():
        row = rows[asset_id]
        runtime = ROOT / str(sheet.cell(row, 15).value)
        source = f"assets/art/props/base_world_3d/source/{slug}/prp_base_{slug}_source_v004.blend"
        base = f"assets/art/props/base_world_3d/components/{slug}/prp_base_{slug}_base_visual_v004.glb"
        swivel = f"assets/art/props/base_world_3d/components/{slug}/prp_base_{slug}_swivel_visual_v004.glb"
        assert runtime.is_file() and all((ROOT / path).is_file() for path in [source, base, swivel])
        sheet.cell(row, 9).value = "default / pushable / seated / swivel / driveable"
        sheet.cell(row, 13).value = "v004"
        sheet.cell(row, 14).value = "分体座面旋转与轮脚底座滑行；乘坐限速=角色步速1.3倍；松手惯性衰减；离座防弹射；实体场景碰撞"
        sheet.cell(row, 16).value = "; ".join([source, base, swivel])
        sheet.cell(row, 20).value = hashlib.sha256(runtime.read_bytes()).hexdigest()
        sheet.cell(row, 22).value = datetime(2026, 9, 29)
        sheet.cell(row, 25).value = str(sheet.cell(row, 25).value) + (
            f"\n2026-09-29：{label} v004 从 v003 Blender 源分离旋转上部与滑行底座，GLB 双件；"
            "角色碰撞改走独立 PlayerBlocker 代理以消除离座刚体弹射。"
            "乘坐速度=玩家步速×1.3，松手后短距滑行并衰减；Godot 物理回归通过。"
        )
    log = workbook["域变更日志"]
    next_row = log.max_row + 1
    assert log.cell(next_row - 1, 1).value == "v0.1.1"
    values = [
        "v0.1.2", "2026-09-29", "99F旋转椅分体与离座防弹射", "道具 / 基地座椅",
        "两把椅子从 v003 源派生 v004；轮脚底座物理滑行、座面/靠背独立转向；乘坐1.3倍步速与松手惯性。",
        "更新两条既有资产的版本、双GLB来源、物理规格、SHA与备注；AssetID和运行时PackedScene路径不变。", "Codex",
    ]
    for col, value in enumerate(values, 1):
        log.cell(next_row, col)._style = copy(log.cell(next_row - 1, col)._style)
        log.cell(next_row, col).value = value
    workbook.save(BOOK)

    baseline = json.loads(BASELINE.read_text(encoding="utf-8"))
    for asset_id in targets:
        row = rows[asset_id]
        values = tuple(sheet.cell(row, col).value for col in range(1, 26))
        baseline["assets"][asset_id]["v"] = _row_digest(values)
    index = json.loads((ROOT / "assets/registry/ledger_index.json").read_text(encoding="utf-8"))
    all_rows = []
    for domain in index["domains"]:
        other = openpyxl.load_workbook(BOOKS / domain["file"])
        all_rows.extend(read_source_rows(other["资产主表"]))
    for col in (9, 13, 14, 16, 20, 22, 25):
        baseline["column_digests"][str(col)] = col_digest(all_rows, col)
    BASELINE.write_text(json.dumps(baseline, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")
    print("SWIVEL_CHAIRS_LEDGER_UPDATED", len(targets))


if __name__ == "__main__":
    main()
