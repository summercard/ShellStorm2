import json,sys,hashlib,shutil
from pathlib import Path
from copy import deepcopy
from collections import Counter
import openpyxl
ROOT=Path(__file__).resolve().parents[2];sys.path[:0]=[str(ROOT/'scripts'),str(ROOT/'tools/asset_pipeline')]
from workspace_paths import work_path
from ledger_registry import LedgerIndex
from split_asset_ledger import read_source_rows,_row_digest,col_digest,CONTENT_COLUMNS
out=work_path('items_reference_20261014/ground/ledger/transaction.json').parent
report=json.loads((out.parent/'report.json').read_text('utf8'))
assert json.loads((out.parent/'validation.json').read_text('utf8'))['passed']
idx=LedgerIndex.load(ROOT);bp=ROOT/'assets/registry/ledger_split_baseline.json';baseline=json.loads(bp.read_text('utf8'));old=deepcopy(baseline);shutil.copy2(bp,out/'baseline_before.json')
books={d.key:openpyxl.load_workbook(d.path) for d in idx.domains}
changed=set();sourceindex=ROOT/'assets/registry/current_source_index.json';si=json.loads(sourceindex.read_text('utf8'))
for x in report:
 folder=(ROOT/x['preview_source']).parent.parent
 for filename in ['asset_contract.json','source_manifest.json']:
  p=folder/filename;shutil.copy2(p,out/(x['slug']+'_'+filename));v=json.loads(p.read_text('utf8'))
  v['model_variants']={'preview':{'source_blend':x['preview_source'],'sha256':x['preview_sha256'],'triangles':x['preview_triangles'],'usage':'近距离预览'},'ground':{'source_blend':x['ground_source'],'sha256':x['ground_sha256'],'polygons':x['ground_polygons'],'triangles':x['ground_triangles'],'usage':'地面掉落/拾取显示','derived_from_sha256':x['preview_sha256']}}
  v['runtime_integration']='两种源均完成，尚未导出或接入用途选择';p.write_text(json.dumps(v,ensure_ascii=False,indent=2)+'\n',encoding='utf8')
 domain='props' if x['asset_id'].startswith('ITM') else 'weapons';w=books[domain]['资产主表'];r=next(r for r,v in read_source_rows(w) if v[0]==x['asset_id'])
 w.cell(r,15).value=x['preview_source']+'; '+x['ground_source']
 # SHA column remains the first path's digest, both exact digests retained in notes and manifest.
 w.cell(r,20).value=x['preview_sha256']
 w.cell(r,14).value=str(w.cell(r,14).value)+'；preview高模；ground '+str(x['ground_triangles'])+'三角面（100–200预算）'
 w.cell(r,25).value=str(w.cell(r,25).value)+' 双模型源：preview='+str(x['preview_triangles'])+'三角；ground='+str(x['ground_triangles'])+'三角；地面源SHA256='+x['ground_sha256']+'；高模哈希保持；未接入运行时。'
 changed.add(x['asset_id'])
 p=ROOT/x['ground_source'];si['sources'].append({'path':x['ground_source'],'bytes':p.stat().st_size,'family':x['ground_source'].replace('_v001',''),'version':[1],'role':'ground','sha256':x['ground_sha256']})
for domain in ['props','weapons']:
 wb=books[domain];path=idx.domain_for_key(domain).path;shutil.copy2(path,out/path.name)
 log=wb['域变更日志'];log.append(['v0.1.100','2026-10-14','物品道具双模型源','preview / ground','现有高模保留；地面100–200三角面','源已完成；运行时用途选择未接入','Codex'])
 for r,v in read_source_rows(wb['资产主表']):
  if v[0] in changed:baseline['assets'][v[0]]={'v':_row_digest(v),'c':v[2],'d':domain}
 wb.save(path)
union=[rv for wb in books.values() for rv in read_source_rows(wb['资产主表'])]
baseline['column_digests']={str(c):col_digest(union,c) for c in CONTENT_COLUMNS};baseline['category_counts']=dict(sorted(Counter(v[2] for _,v in union).items()))
assert all(baseline['assets'][k]==v for k,v in old['assets'].items() if k not in changed)
bp.write_text(json.dumps(baseline,ensure_ascii=False,indent=1)+'\n',encoding='utf8');sourceindex.write_text(json.dumps(si,ensure_ascii=False,indent=2)+'\n',encoding='utf8')
for path,sheet in [(ROOT/'assets/registry/ledgers/ShellStorm2_物品道具账本_v001.xlsx','物品道具'),(idx.domain_for_key('weapons').path,'背包物品调用索引')]:
 wb=openpyxl.load_workbook(path);w=wb[sheet]
 for row in w:
  if len(row)>13 and row[13].value in changed:
   cell=w.cell(row[0].row,17);cell.value=str(cell.value or '')+'；preview/ground双模型源已完成，ground100–200三角，运行时用途分流待接入。'
 wb.save(path)
(out/'transaction.json').write_text(json.dumps({'changed_assets':sorted(changed),'unrelated_fingerprints_preserved':True,'high_sources_unchanged':True},ensure_ascii=False,indent=2),encoding='utf8')
print('DUAL_MODEL_LEDGER_OK',len(changed))
