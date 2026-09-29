"""Register the 99F telescopic ladder and v024 climb clip without rebuilding ledgers."""

from collections import Counter
from copy import copy
from datetime import datetime
from pathlib import Path
import hashlib
import json
import re
import shutil
import sys

import openpyxl
from openpyxl.worksheet.datavalidation import DataValidation

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
sys.path.insert(0, str(HERE))
from split_asset_ledger import (  # noqa: E402
    CONTENT_COLUMNS, FIRST_DATA_ROW, _row_digest, col_digest,
    dedupe_key_formula, dedupe_result_formula, read_source_rows,
)

BOOKS = ROOT / "assets/registry/ledgers"
BASELINE = ROOT / "assets/registry/ledger_split_baseline.json"
BACKUP = ROOT / "_scratch/base99_telescopic_ladder_ledger_before_20260929"
ASSET_ID = "PRP-BASE99-TELESCOPIC-LADDER-3D"
TODAY = datetime(2026, 9, 29)


def sha(path: str) -> str:
    return hashlib.sha256((ROOT / path).read_bytes()).hexdigest()


def get_row(sheet, asset_id):
    matches = [r for r in range(FIRST_DATA_ROW, sheet.max_row + 1) if sheet.cell(r, 1).value == asset_id]
    assert len(matches) == 1, (asset_id, matches)
    return matches[0]


def append_log(workbook, title, scope, details):
    sheet = workbook["域变更日志"]
    old = sheet.max_row
    old_version = str(sheet.cell(old, 1).value)
    match = re.fullmatch(r"v(\d+)\.(\d+)\.(\d+)", old_version)
    assert match, old_version
    version = f"v{match[1]}.{match[2]}.{int(match[3]) + 1}"
    values = [version, "2026-09-29", title, scope, details, "更新资产主表与运行时事实源；既有历史资源保留。", "Codex"]
    for col, value in enumerate(values, 1):
        sheet.cell(old + 1, col)._style = copy(sheet.cell(old, col)._style)
        sheet.cell(old + 1, col).value = value


def main():
    index = json.loads((ROOT / "assets/registry/ledger_index.json").read_text(encoding="utf-8"))
    prop_path = BOOKS / "ShellStorm2_道具账本_v001.xlsx"
    scene_path = BOOKS / "ShellStorm2_场景账本_v001.xlsx"
    char_path = BOOKS / "ShellStorm2_角色账本_v001.xlsx"
    prop = openpyxl.load_workbook(prop_path)
    scene = openpyxl.load_workbook(scene_path)
    char = openpyxl.load_workbook(char_path)
    assert all(name in prop.sheetnames for name in ("资产主表", "总览", "域变更日志"))
    assert all(prop["资产主表"].cell(r, 1).value != ASSET_ID for r in range(FIRST_DATA_ROW, prop["资产主表"].max_row + 1))
    old_stair_row = get_row(scene["资产主表"], "BPK-BASE99-EAST-UPPER-TRANSITION-STAIR")
    bunny_row = get_row(char["资产主表"], "CHR-PLY-CAPSULE01-3D-BUNNY01")
    BACKUP.mkdir(parents=True, exist_ok=True)
    for path in (prop_path, scene_path, char_path, BASELINE):
        shutil.copy2(path, BACKUP / path.name)

    sheet = prop["资产主表"]
    row = sheet.max_row + 1
    template = get_row(sheet, "PRP-BASE-WORKSHOP-STOOL-3D")
    for col in range(1, 26):
        sheet.cell(row, col)._style = copy(sheet.cell(template, col)._style)
    sheet.row_dimensions[row].height = sheet.row_dimensions[template].height
    source = "assets/art/props/base_world_3d/source/base99_telescopic_ladder/prp_base99_telescopic_ladder_source_v001.blend"
    upper = "assets/art/props/base_world_3d/components/base99_telescopic_ladder/prp_base99_telescopic_ladder_upper_visual_top3d.glb"
    lower = "assets/art/props/base_world_3d/components/base99_telescopic_ladder/prp_base99_telescopic_ladder_lower_visual_top3d.glb"
    runtime = "assets/art/props/base_world_3d/runtime/base99_telescopic_ladder/prp_base99_telescopic_ladder_root_top3d.tscn"
    for path in (source, upper, lower, runtime):
        assert (ROOT / path).is_file(), path
    values = {
        1: ASSET_ID, 2: "99F东阁楼双段伸缩直梯", 3: "道具", 4: "decor_prop",
        5: "base99_rooftop_telescopic_ladder", 6: "root", 7: None,
        8: "Top3D / 固定位置", 9: "retracted / deploying / deployed / climbing",
        10: "TowerDescent99F / base100_rooftop", 11: "正式美术已接入", 12: "P1", 13: "v001",
        14: "双段6m直梯；上端E放下；展开后两端E自动攀爬、反向输入折返；顶部穿东门至100F平台、底部落99F阁楼；场景实体碰撞；展开状态存档",
        15: runtime,
        16: "; ".join((source, upper, lower, "src/base3d/Base99TelescopicLadder3D.gd", "src/player3d/states/Player3DClimbingState.gd")),
        17: "99F;100F;伸缩直梯;可使用道具;攀爬;双段梯", 20: sha(runtime),
        21: "Codex", 22: TODAY, 23: "用户指定Blender资产", 24: "99F-100F-EAST-LADDER",
        25: "2026-09-29：替换东侧12米斜梯，西侧原楼梯保留；Blender源含Deploy_LowerSection规范动画；运行时固定交互、自动攀爬及存档验证。",
    }
    for col, value in values.items():
        sheet.cell(row, col).value = value
    for current in range(FIRST_DATA_ROW, row + 1):
        sheet.cell(current, 18).value = dedupe_key_formula(current)
        sheet.cell(current, 19).value = dedupe_result_formula(current, row)
    for address in ("A6", "C6", "E6", "G6", "B10", "C10", "B11", "C11"):
        cell = prop["总览"][address]
        assert f"${row - 1}" in cell.value, (address, cell.value)
        cell.value = cell.value.replace(f"${row - 1}", f"${row}")
    for validation in list(sheet.data_validations.dataValidation):
        ranges = [str(value) for value in validation.sqref.ranges]
        new_validation = DataValidation(
            type=validation.type, formula1=validation.formula1, formula2=validation.formula2,
            allow_blank=validation.allow_blank, showErrorMessage=validation.showErrorMessage,
            showInputMessage=validation.showInputMessage, error=validation.error,
            errorTitle=validation.errorTitle, prompt=validation.prompt,
            promptTitle=validation.promptTitle,
        )
        sheet.data_validations.dataValidation.remove(validation)
        sheet.add_data_validation(new_validation)
        for value in ranges:
            new_validation.add(re.sub(rf"(?<=:)([A-Z]+){row - 1}$", rf"\g<1>{row}", value))
    sheet.auto_filter.ref = f"A5:X{row}"
    append_log(prop, "99F双段伸缩直梯正式接入", "道具 / 基地交互道具", "新增伸缩直梯 AssetID；双段GLB、固定碰撞、展开动画、攀爬状态与存档接入。")

    old = scene["资产主表"]
    old.cell(old_stair_row, 11).value = "弃用"
    old.cell(old_stair_row, 22).value = TODAY
    old.cell(old_stair_row, 25).value = str(old.cell(old_stair_row, 25).value) + "\n2026-09-29：99F东侧旧斜梯从运行时布局移除；历史源与Prefab保留；由PRP-BASE99-TELESCOPIC-LADDER-3D替换。"
    append_log(scene, "东侧旧斜梯退役", "场景 / 基地结构", "BPK-BASE99-EAST-UPPER-TRANSITION-STAIR从运行时移除并标记弃用，保留历史源。")

    cs = char["资产主表"]
    animation = "assets/art/characters/player/chr_player_capsule01_3d/variants/bunny01/production/v024/source/animation/chr_bunny01_animation_v024.blend"
    transfer = "assets/art/characters/player/chr_player_capsule01_3d/variants/bunny01/production/v024/character_transfer_ledger_v024.json"
    assert (ROOT / animation).is_file() and (ROOT / transfer).is_file()
    cs.cell(bunny_row, 9).value = "default / seated / climbing"
    cs.cell(bunny_row, 16).value = str(cs.cell(bunny_row, 16).value) + "; " + animation + "; " + transfer
    cs.cell(bunny_row, 22).value = TODAY
    cs.cell(bunny_row, 25).value = str(cs.cell(bunny_row, 25).value) + "\n2026-09-29：v024原角色动画母版新增climbing交替手脚攀爬循环，映射到Player3D climbing状态。"
    append_log(char, "兔子角色攀爬状态与动作", "角色 / Bunny01", "v024 Blender动作母版与14剪辑库新增climbing；旧模型与其余动作保持。")

    for workbook, path in ((prop, prop_path), (scene, scene_path), (char, char_path)):
        workbook.save(path)
    baseline = json.loads(BASELINE.read_text(encoding="utf-8"))
    touched = {ASSET_ID, "BPK-BASE99-EAST-UPPER-TRANSITION-STAIR", "CHR-PLY-CAPSULE01-3D-BUNNY01"}
    all_rows = []
    for domain in index["domains"]:
        book = openpyxl.load_workbook(BOOKS / domain["file"], read_only=True)
        for number, record in read_source_rows(book["资产主表"]):
            all_rows.append((number, record))
            asset_id = str(record[0])
            if asset_id in touched:
                baseline["assets"][asset_id] = {
                    "v": _row_digest(record), "c": str(record[2]), "d": domain["key"],
                }
    baseline["captured_at"] = "2026-09-29"
    baseline["asset_count"] = len(baseline["assets"])
    baseline["column_digests"] = {str(col): col_digest(all_rows, col) for col in CONTENT_COLUMNS}
    baseline["category_counts"] = dict(sorted(Counter(str(record[2]) for _, record in all_rows).items()))
    BASELINE.write_text(json.dumps(baseline, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")
    check = openpyxl.load_workbook(prop_path)
    assert check["资产主表"].cell(row, 1).value == ASSET_ID
    assert all(str(dv.sqref).endswith(f"{row}") for dv in check["资产主表"].data_validations.dataValidation)
    print("BASE99_TELESCOPIC_LADDER_LEDGER_UPDATED", row)


if __name__ == "__main__":
    main()
