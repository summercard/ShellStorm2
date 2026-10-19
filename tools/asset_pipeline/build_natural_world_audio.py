"""Replace all non-system/non-weapon sound sets with sourced recording edits.

Input archives and downloaded recordings live in the external 2026-10-19 session.
Every preserved system/weapon asset is hash-locked; no runtime synthesis is used.
"""
from pathlib import Path
from copy import copy
import csv, datetime, hashlib, json, shutil, subprocess, sys, zipfile
import numpy as np
import openpyxl

ROOT=Path(__file__).resolve().parents[2]
sys.path.insert(0,str(ROOT/'scripts'))
from workspace_paths import work_path
SESSION=work_path('sessions/2026-10-19_natural_world_audio/report.json').parent
INTAKE=SESSION/'intake'
OUT=ROOT/'src/assets/audio/sfx'
BANK=ROOT/'art_source/audio/sfx/recorded_world_library/source'
BACKUP=ROOT.parent/'ShellStorm2-backup/2026-10-19_natural_world_audio/original_tree'
RATE=44100
REV='natural-world-r5'
sha=lambda p:hashlib.sha256(p.read_bytes()).hexdigest()

def save(p,v,indent=2):
 p.parent.mkdir(parents=True,exist_ok=True)
 p.write_text(json.dumps(v,ensure_ascii=False,indent=indent)+'\n',encoding='utf-8')

def backup(p):
 if p.exists():
  target=BACKUP/p.relative_to(ROOT);target.parent.mkdir(parents=True,exist_ok=True)
  if not target.exists():shutil.copy2(p,target)

class Sources:
 def __init__(self):self.records={};self.pcm={};BANK.mkdir(parents=True,exist_ok=True)
 def archive(self,pack,name):
  key=pack+'_'+Path(name).name
  if key in self.records:return key
  z=zipfile.ZipFile(INTAKE/(pack+'.zip'));p=BANK/key;p.write_bytes(z.read(name))
  credit,url,license={
   'impact':('Kenney','https://kenney.nl/assets/impact-sounds','CC0-1.0'),
   'rpg':('Kenney','https://kenney.nl/assets/rpg-audio','CC0-1.0'),
   'monster':('StarNinjas','https://opengameart.org/content/16-monster-growls','CC0-1.0'),
   'human':('HaelDB','https://opengameart.org/content/male-gruntyelling-sounds','CC0-1.0'),
  }[pack]
  self.records[key]={'path':p.relative_to(ROOT).as_posix(),'sha256':sha(p),'credit':credit,'source_url':url,'license':license,'archive_member':name}
  return key
 def direct(self,name):
  if name in self.records:return name
  p=BANK/name;shutil.copy2(INTAKE/name,p)
  credit,url,license={
   'vent.mp3':('MidFag','https://opengameart.org/content/ventilation-version2','CC-BY-3.0'),
   'fan.wav':('dklon','https://opengameart.org/content/broken-fan-loop','CC-BY-3.0'),
   'wind.wav':('remaxim','https://opengameart.org/content/short-wind-sound','CC0-1.0'),
  }[name]
  self.records[name]={'path':p.relative_to(ROOT).as_posix(),'sha256':sha(p),'credit':credit,'source_url':url,'license':license}
  return name
 def openspades(self,name):
  key='openspades_'+name
  if key in self.records:return key
  p=BANK/key;shutil.copy2(INTAKE/'openspades'/name,p)
  lines=(ROOT/'art_source/audio/sfx/recording_credits_20261019.md').read_text(encoding='utf-8').splitlines()
  original=name.replace('_','/')
  match=next(l for l in lines if 'Sounds/'+original+' |' in l)
  fields=[v.strip() for v in match.split('|')]
  self.records[key]={'path':p.relative_to(ROOT).as_posix(),'sha256':sha(p),'credit':fields[4],'license':fields[3],
    'source_url':fields[2].split('>')[0].lstrip('<'),'distribution_url':'https://github.com/DeathByDenim/openspades-free-pak'}
  return key
 def decode(self,key):
  if key not in self.pcm:
   p=ROOT/self.records[key]['path']
   data=subprocess.check_output(['ffmpeg','-v','error','-i',str(p),'-ac','1','-ar',str(RATE),'-af','highpass=f=35','-f','f32le','-'])
   x=np.frombuffer(data,dtype='<f4').astype(np.float64);self.pcm[key]=x-x.mean()
  return self.pcm[key]

def layer(key,start=0,duration=None,offset=0,gain=1):
 return dict(source=key,start_sec=start,duration_sec=duration,offset_sec=offset,gain=gain)

def plan(s,key,var):
 r=lambda n:s.archive('rpg','Audio/'+n+'.ogg')
 i=lambda n:s.archive('impact','Audio/'+n+'.ogg')
 m=lambda n:s.archive('monster',f'monster.{n}.ogg')
 h=lambda n:s.archive('human',f'yelling sounds/3grunt{n}.wav')
 o=lambda n:s.openspades(n+'.opus')
 v=var%3;idx=f'{var%5:03d}'
 if key.startswith('amb_'):
  vent=s.direct('vent.mp3');fan=s.direct('fan.wav');wind=s.direct('wind.wav')
  mix,cutoff={
   'amb_base_facility':([(vent,.8),(fan,.15)],3000),
   'amb_workshop':([(fan,.85),(vent,.35)],6500),
   'amb_recovery':([(vent,.55),(wind,.15)],1200),
   'amb_expedition_room':([(wind,.45),(vent,.35)],2500),
   'amb_expedition_open':([(wind,1)],8500),
   'amb_extraction_beacon':([(fan,.55),(wind,.3)],4200),
   'amb_boss_arena':([(wind,.55),(fan,.65)],1600),
  }[key]
  return [layer(k,gain=g) for k,g in mix],10,cutoff,-21
 if key=='enemy_hit':return [layer(o(f'Weapons_Impacts_Flesh{v+1}')),layer(i('impactSoft_medium_'+idx),gain=.35)],.36,13000,-14
 if key=='enemy_die':return [layer(m([6,11,16][v])),layer(i('impactSoft_heavy_'+idx),offset=.15,gain=.5)],1.0,11000,-14
 if key=='player_hit':return [layer(h(v+1),duration=.45,gain=.6),layer(i('impactPunch_medium_'+idx),gain=.5)],.55,13000,-14
 if key=='player_dash':return [layer(r('cloth'+str(v+1))),layer(o('Weapons_Spade_Miss'),gain=.45)],.42,13000,-17
 if key=='crit_hit':return [layer(o(f'Weapons_Impacts_Flesh{v+1}')),layer(i('impactPunch_heavy_'+idx),gain=.65)],.4,14000,-13
 if key=='melee_swing':return [layer(o('Weapons_Spade_Miss')),layer(r('cloth'+str(v+1)),gain=.6)],.42,14000,-16
 if key in ['melee_impact','melee_heavy_impact']:
  return [layer(i('impactPunch_heavy_'+idx)),layer(i('impactMetal_medium_'+idx),gain=.25),layer(r('cloth'+str(v+1)),gain=.2)],.65 if key=='melee_heavy_impact' else .45,13000,-13
 if key=='melee_target_hit':return [layer(o(f'Weapons_Impacts_Flesh{v+1}')),layer(i('impactSoft_medium_'+idx),gain=.35)],.36,13000,-16
 if key=='projectile_impact':return [layer(i('impactMining_'+idx)),layer(i('impactWood_light_'+idx),gain=.45)],.4,13000,-16
 if key=='footstep_concrete':return [layer(i('footstep_concrete_'+idx))],.36,14000,-20
 if key=='footstep_metal':return [layer(i('footstep_concrete_'+idx)),layer(i('impactPlate_light_'+idx),gain=.3)],.38,14000,-20
 if key=='footstep_debris':return [layer(i('footstep_grass_'+idx)),layer(i('impactMining_'+idx),gain=.25)],.4,14000,-20
 if key=='player_land':return [layer(i('impactSoft_heavy_'+idx)),layer(i('footstep_concrete_'+idx),gain=.45)],.5,12000,-17
 if key=='player_down':return [layer(h(4),duration=.6,gain=.6),layer(i('impactSoft_heavy_003'),offset=.2,gain=.8)],.85,12000,-15
 if key=='player_death':return [layer(h(6),duration=.8,gain=.5),layer(i('impactSoft_heavy_004'),offset=.2,gain=.8),layer(r('cloth4'),offset=.25,gain=.3)],1.15,11000,-15
 if key=='door_open':return [layer(r('metalLatch'),gain=.7),layer(r('doorOpen_2'),offset=.05)],1.4,12000,-16
 if key=='door_close':return [layer(r('doorClose_2')),layer(i('impactMetal_medium_003'),offset=.4,gain=.2)],1.1,12000,-16
 if key=='door_locked':return [layer(r('metalLatch'))],.35,13000,-18
 if key=='container_open':return [layer(r('beltHandle1')),layer(r('metalLatch'),offset=.08,gain=.6),layer(r('creak1'),offset=.12,gain=.5)],.85,11000,-18
 if key=='container_search':return [layer(r('handleSmallLeather')),layer(r('clothBelt2'),offset=.1,gain=.5)],.65,13000,-21
 if key=='container_cancel':return [layer(r('beltHandle2'))],.32,13000,-20
 if key=='melee_cancel':return [layer(r('drawKnife2'),gain=.6),layer(r('clothBelt'),gain=.4)],.35,12000,-18
 if key in ['flashlight_on','light_on','radio_on']:return [layer(r('metalClick'))],.22,14000,-20
 if key in ['flashlight_off','light_off','radio_off']:return [layer(r('metalLatch'),duration=.17)],.24,14000,-20
 if key=='facility_activate':return [layer(r('metalLatch'),gain=.3),layer(s.direct('fan.wav'),duration=1.5,offset=.06,gain=.45)],1.7,6500,-19
 if key=='enemy_alert':return [layer(m(9))],.85,11000,-16
 if key=='enemy_search':return [layer(m(2),gain=.65)],.85,10000,-19
 if key=='enemy_telegraph':return [layer(m([3,5,12][v]))],.85,12000,-15
 if key=='enemy_attack':return [layer(m([4,7,10][v]),gain=.7),layer(r('cloth'+str(v+1)),gain=.4)],.65,13000,-16
 if key=='enemy_summon':return [layer(o('Weapons_Grenade_Debris'),gain=.5),layer(m(11),gain=.45)],1.25,9000,-16
 if key=='enemy_heal':return [layer(r('clothBelt'),gain=.5),layer(m(1),gain=.25)],.95,6500,-19
 if key=='enemy_avoid_light':return [layer(m(15))],.6,13000,-17
 if key=='enemy_charge':return [layer(m(14),gain=.7),layer(r('cloth4'),gain=.3)],1.0,9000,-15
 if key=='enemy_shield':return [layer(i('impactPlate_heavy_002')),layer(r('clothBelt2'),gain=.4)],.75,13000,-16
 if key=='enemy_explosion':return [layer(o('Weapons_Grenade_Explode'+str(v%2+1))),layer(o('Weapons_Grenade_Debris'),offset=.12,gain=.3+v*.1)],1.5,14000,-12
 raise ValueError('Unplanned physical sound '+key)

def repeat_crossfade(x,n):
 overlap=min(round(.2*RATE),len(x)//4)
 assert overlap>0
 out=x.copy();w=np.linspace(0,1,overlap)
 while len(out)<n:
  out=np.concatenate([out[:-overlap],out[-overlap:]*(1-w)+x[:overlap]*w,x[overlap:]])
 return out[:n]

def render(s,asset,layers,duration,cutoff,peak):
 looping=asset['loop'];n=round(duration*RATE);f=round(.5*RATE) if looping else 0
 mix=np.zeros(n+f)
 for item in layers:
  x=s.decode(item['source']);start=round(item['start_sec']*RATE);x=x[start:]
  if item['duration_sec'] is not None:x=x[:round(item['duration_sec']*RATE)]
  if looping:x=repeat_crossfade(x,len(mix))
  else:
   x=x.copy();threshold=max(np.max(abs(x))*.02,.00001);active=np.flatnonzero(abs(x)>threshold)
   if len(active):x=x[max(0,int(active[0])-44):]
   fade=min(441,len(x)//4)
   if fade:x[-fade:]*=np.linspace(1,0,fade)
  offset=round(item['offset_sec']*RATE);length=min(len(x),len(mix)-offset)
  mix[offset:offset+length]+=x[:length]*item['gain']
 raw=SESSION/'mix.f32';filtered=SESSION/'filtered.f32';raw.write_bytes(mix.astype('<f4').tobytes())
 subprocess.run(['ffmpeg','-v','error','-y','-f','f32le','-ar',str(RATE),'-ac','1','-i',str(raw),'-af',f'lowpass=f={cutoff}','-f','f32le',str(filtered)],check=True)
 x=np.frombuffer(filtered.read_bytes(),dtype='<f4').astype(np.float64)
 if looping:
  # Rotate an overlap across the file boundary: equal-power is avoided here
  # because correlated fan/wind takes should not gain 3 dB at every splice.
  w=np.linspace(0,1,f);x=np.concatenate([x[f:-f],x[-f:]*(1-w)+x[:f]*w])
  gain=min(10**(-31/20)/max(np.sqrt(np.mean(x*x)),1e-9),10**(peak/20)/max(abs(x)))
 else:
  active=np.flatnonzero(abs(x)>max(np.max(abs(x))*.02,.00001));trim=max(0,int(active[0])-44) if len(active) else 0
  if trim:x=np.concatenate([x[trim:],np.zeros(trim)])
  x[:44]*=np.linspace(0,1,44);fade=min(round(.025*RATE),len(x)//4);x[-fade:]*=np.linspace(1,0,fade)
  gain=10**(peak/20)/max(float(np.max(abs(x))),1e-9)
 return (x*gain).astype('<f4'),gain

def main():
 scope=json.loads((SESSION/'scope.json').read_text(encoding='utf-8'))
 cp=OUT/'sfx_catalog.json';catalog=json.loads(cp.read_text(encoding='utf-8'))
 assert not catalog.get('world_style_policy'),'Natural-world audio already installed'
 protected={a['asset_id']:a.copy() for a in catalog['assets'] if a['key'] not in scope['replace']}
 target=[a for a in catalog['assets'] if a['key'] in scope['replace']];assert len(target)==84
 backup(cp);s=Sources();changes={};path_map={}
 # Resolve the complete selection before editing any current asset.
 plans={a['asset_id']:plan(s,a['key'],a['variant']) for a in target}
 save(SESSION/'frozen_plan.json',plans)
 for asset in target:
  layers,duration,cutoff,peak=plans[asset['asset_id']]
  x,gain=render(s,asset,layers,duration,cutoff,peak)
  old_runtime=ROOT/asset['runtime_path'].removeprefix('res://');old_source=ROOT/asset['source_path'];source=old_source.parent
  for p in source.rglob('*'):
   if p.is_file():backup(p)
  backup(old_runtime)
  version=asset['version']+1;name=old_runtime.stem.rsplit('_v',1)[0]+f'_v{version:03d}'
  master=source/(name+'.wav');runtime=OUT/(name+'.ogg');raw=SESSION/'normalized.f32';raw.write_bytes(x.tobytes())
  subprocess.run(['ffmpeg','-v','error','-y','-f','f32le','-ar',str(RATE),'-ac','1','-i',str(raw),'-c:a','pcm_s24le',str(master)],check=True)
  subprocess.run(['ffmpeg','-v','error','-y','-i',str(master),'-c:a','libvorbis','-q:a','6',str(runtime)],check=True)
  original=asset['runtime_path'];new='res://'+runtime.relative_to(ROOT).as_posix();path_map[original]=new
  license_items=[s.records[l['source']] for l in layers]
  asset.update(source_path=master.relative_to(ROOT).as_posix(),runtime_path=new,version=version,duration_sec=duration,
   sha256=sha(runtime),source_sha256=sha(master),license_source='; '.join(sorted({v['source_url']+' | '+v['credit'] for v in license_items})),
   license_scope='; '.join(sorted({v['license'] for v in license_items})),production_revision=REV,
   source_method='Sourced recordings / Foley edit and linear layering; no tone synthesis',peak_target_db=peak,status='imported_pending_listening')
  asset.pop('seed',None)
  save(source/'recipe.json',{'renderer':'tools/asset_pipeline/build_natural_world_audio.py','asset':asset,'layers':layers,
   'sources':{l['source']:s.records[l['source']] for l in layers},'lowpass_hz':cutoff,'gain':gain,
   'loop_crossfade_sec':.5 if asset['loop'] else 0,'mix_policy':'44.1kHz mono PCM24; one-shots peak normalized; beds RMS -31dBFS with -21dBFS peak ceiling'})
  assert source.resolve().is_relative_to((ROOT/'art_source/audio/sfx').resolve())
  for p in list(source.iterdir()):
   if p.is_file() and p.name not in [master.name,'recipe.json']:p.unlink()
  stems=source/'stems'
  if stems.exists():assert stems.resolve().parent==source.resolve();shutil.rmtree(stems)
  old_runtime.unlink();changes[asset['asset_id']]=asset
  print(asset['key'],asset['variant'],'ready',flush=True)
 save(BANK/'source_index.json',s.records)
 for pack in ['impact','rpg']:
  z=zipfile.ZipFile(INTAKE/(pack+'.zip'));(BANK/(pack+'_LICENSE.txt')).write_bytes(z.read('License.txt'))
 lines=['# Recorded world audio sources','', 'Edits: extraction, layering, filters, gain, fades, looping. No claim that creature voices are real animals.','']
 for v in s.records.values():lines.append(f'- {v["path"]} — {v["credit"]}; {v["license"]}; {v["source_url"]}')
 (BANK/'CREDITS.md').write_text('\n'.join(lines)+'\n',encoding='utf-8')
 for sound_set in catalog['sound_sets'].values():sound_set['paths']=[path_map.get(p,p) for p in sound_set['paths']]
 policy={'revision':REV,'system_sets':scope['keep_system'],'weapon_sets':scope['keep_weapons'],'natural_world_sets':scope['replace'],
  'preserved_asset_hashes':{k:a['sha256'] for k,a in protected.items()},'changed_assets':len(changes),'environment_loops':7}
 catalog['world_style_policy']=policy;catalog['revision']='natural-world-r5 + gun-families-r4 + system cues'
 save(cp,catalog);save(ROOT/'art_source/audio/sfx/world_audio_policy.json',policy)
 for a in catalog['assets']:
  if a['asset_id'] in protected:assert a==protected[a['asset_id']] and sha(ROOT/a['runtime_path'].removeprefix('res://'))==a['sha256']
 manager=ROOT/'src/core/AudioManager.gd';backup(manager);text=manager.read_text(encoding='utf-8')
 for a,b in path_map.items():text=text.replace(a,b)
 manager.write_text(text,encoding='utf-8')
 mf=OUT/'audio_manifest.csv';backup(mf)
 with mf.open(encoding='utf-8',newline='') as f:r=csv.DictReader(f);fields=r.fieldnames;rows=list(r)
 for row in rows:
  a=changes.get(row['asset_id'])
  if a:row.update(source_master=a['source_path'],runtime_file=Path(a['runtime_path']).name,runtime_resource_path=a['runtime_path'],sha256=a['sha256'],duration_sec=a['duration_sec'],status='自然材质/录音替换；人工听感待验')
 with mf.open('w',encoding='utf-8',newline='') as f:w=csv.DictWriter(f,fieldnames=fields);w.writeheader();w.writerows(rows)
 update_ledger(changes)
 contract=ROOT/'art_source/audio/sfx/library_contract.json';backup(contract);d=json.loads(contract.read_text(encoding='utf-8'))
 d.update(design_revision='AUDIO-SFX r4',production_revision=REV,world_style_policy=policy,
  source_method='Natural recording/Foley world and ambience; recorded gun families; original DSP retained only for explicit system cue allowlist')
 for spec in d['sounds']:
  if spec['key'] in scope['replace']:spec['production_override']={'revision':REV,'renderer':'tools/asset_pipeline/build_natural_world_audio.py','source_of_current_duration_and_version':'src/assets/audio/sfx/sfx_catalog.json'}
 save(contract,d);save(SESSION/'report.json',{'replaced_sets':47,'replaced_assets':84,'protected_assets':len(protected),'sources':len(s.records),'assets':list(changes.values())})
 print('NATURAL_WORLD_AUDIO_INSTALLED: 47 sets / 84 assets; 73 system and weapon assets unchanged')

def update_ledger(changes):
 sys.path.insert(0,str(ROOT/'tools/asset_pipeline'))
 from split_asset_ledger import read_source_rows,_row_digest,CONTENT_COLUMNS,col_digest
 idx=json.loads((ROOT/'assets/registry/ledger_index.json').read_text(encoding='utf-8'));domain=next(d for d in idx['domains'] if d['key']=='audio')
 p=ROOT/idx['ledger_dir']/domain['file'];backup(p);w=openpyxl.load_workbook(p);ws=w[idx['asset_sheet']];count=0
 for row in range(6,ws.max_row+1):
  a=changes.get(ws.cell(row,1).value)
  if not a:continue
  status=ws.cell(row,11).value
  vals={13:f'v{a["version"]:03d}',14:f'44100Hz / mono / PCM24 master / Vorbis q6 / {a["duration_sec"]}s',
   15:a['runtime_path'].removeprefix('res://'),16:a['source_path'],20:a['sha256'],22:datetime.datetime(2026,10,19),
   23:a['license_source']+' | '+a['license_scope'],24:a['source_path'],25:'自然材质/录音替换；'+('已接入；' if status=='原型已接入' else '素材就绪、正式领域触发仍待接入；')+'人工听感及Android待验；verify_audio_library_flow；docs/v0.1/development/2026-10-19_natural_world_audio.md'}
  for col,value in vals.items():ws.cell(row,col,value)
  count+=1
 assert count==84
 log=w['域变更日志'];r=log.max_row+1
 for col,value in enumerate(['v0.2.2',datetime.datetime(2026,10,19),'世界音效自然化','47组84条；全部7组环境循环','替换环境、脚步、受击、近战、怪物、门锁、容器和设备的旧DSP；29组系统提示与28组枪械声音保持。','AssetID、事件键、104组157条总数、玩法与触发保持；人工听感及Android待验。','Codex'],1):
  log.cell(r,col,value);log.cell(r,col)._style=copy(log.cell(6,col)._style);log.cell(r,col).alignment=openpyxl.styles.Alignment(wrap_text=True,vertical='center')
 log.cell(r,2).number_format='yyyy-mm-dd';log.row_dimensions[r].height=85
 desc=w['账本说明'];r=desc.max_row+1;desc.cell(r,1,'环境/实体声音风格');desc.cell(r,2,'47组84条自然Foley/录音，含7组无缝环境循环；29组系统提示风格保留；既有七类枪声不变。人工听感与Android待验。');desc.row_dimensions[r].height=52;desc.cell(r,2).alignment=openpyxl.styles.Alignment(wrap_text=True)
 w.save(p)
 bp=ROOT/'assets/registry/ledger_split_baseline.json';backup(bp);bl=json.loads(bp.read_text(encoding='utf-8'))
 for _,values in read_source_rows(ws):
  if values[0] in changes:bl['assets'][values[0]]['v']=_row_digest(values)
 allrows=[]
 for domain in idx['domains']:
  other=openpyxl.load_workbook(ROOT/idx['ledger_dir']/domain['file']);allrows.extend(read_source_rows(other[idx['asset_sheet']]))
 bl['column_digests']={str(c):col_digest(allrows,c) for c in CONTENT_COLUMNS};save(bp,bl,1)

if __name__=='__main__':main()
