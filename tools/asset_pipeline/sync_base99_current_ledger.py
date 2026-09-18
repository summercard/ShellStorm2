"""Synchronize the authoritative Base99 rows with the current asset graph.

只作用于《基地资产包》所在的分账本（由 assets/registry/ledger_index.json 解析）。
"""
from __future__ import annotations

import argparse
import hashlib
import json
import re
import sys
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[2] / "scripts"))
from ledger_registry import LedgerIndex  # noqa: E402
from split_asset_ledger import dedupe_key_formula, dedupe_result_formula  # noqa: E402


DONE_STATUSES = (
    "已完成",
    "原型已接入",
    "正式美术已接入",
    "已优化并正式接入",
    "Blender源已完成",
    "已导入；优化完成",
)


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def rel(root: Path, path: Path) -> str:
    return str(path.relative_to(root)).replace("\\", "/")


def manifests_by_slug(root: Path) -> dict[str, tuple[int, Path, dict]]:
    result: dict[str, tuple[int, Path, dict]] = {}
    for path in (root / "source/art/blender/base_facility_layout/component_packages").rglob("asset_manifest.json"):
        data = json.loads(path.read_text(encoding="utf-8"))
        slug = str(data.get("asset_slug", "")).strip()
        if not slug:
            continue
        match = re.search(r"v(\d{3})", str(data.get("version", "")))
        version = int(match.group(1)) if match else -1
        current = result.get(slug)
        if current is None or (version, str(path)) > (current[0], str(current[1])):
            result[slug] = (version, path, data)
    return result


def set_formula(ws, row: int, end: int) -> None:
    # 公式形状只由 split_asset_ledger 定义一次，避免这里另抄一份后悄悄失配。
    ws.cell(row, 18).value = dedupe_key_formula(row)
    ws.cell(row, 19).value = dedupe_result_formula(row, end)


def update_workbook(root: Path, dry_run: bool) -> dict:
    # 账本已按域拆开：本工具只动 base99 家族，全部落在「基地资产包」大类所属的分账本里。
    # 路径一律经 ledger_index.json 解析，不再硬编码单体账本文件名。
    index = LedgerIndex.load(root)
    target = index.domain_for_category("基地资产包")
    workbook_path = target.path
    workbook = load_workbook(workbook_path, data_only=False)
    main = workbook["资产主表"]
    overview = workbook["总览"]
    manifest_map = manifests_by_slug(root)
    master = "source/art/blender/base_facility_layout/source/base_facility_runtime_layout_hq_v026.blend"
    updated = defaultdict(int)
    updated["ledger"] = target.relative_path

    explicit = {
        "ENV-BASE99-WALL-DOOR-5X9": (
            "assets/art/environments/base_facility_3d/runtime/env_base99_wall_door_5x9/env_base99_wall_door_5x9_root_top3d_v003.tscn",
            "v003",
            "已完成",
            master + "; source/art/blender/base_facility_layout/export/v025/base_facility_runtime_layout_hq-v025-door_wall_palette.blend; assets/art/environments/base_facility_3d/components/env_base99_wall_door_5x9/env_base99_wall_door_5x9_visual_top3d_v003.glb",
        ),
        "ENV-BASE99-ART-LAYOUT-3D": (
            "assets/art/environments/tower_zones/base/runtime/zone_base_v002.tscn",
            "v002",
            "原型已接入",
            master + "; assets/art/environments/tower_zones/base/runtime/zone_base_v002.tscn",
        ),
        "ENV-TOWER-CORNER-L-5M": (
            "assets/art/environments/base_facility_3d/runtime/env_base99_corner_l_5m/env_base99_corner_l_5m_root_top3d_v005.tscn",
            "v005",
            "正式美术已接入",
            master + "; source/art/blender/base_facility_layout/component_packages/architecture/base_corner_l_5m/base_corner_l_5m_source_v024.blend; assets/art/environments/base_facility_3d/components/env_base99_corner_l_5m/env_base99_corner_l_5m_visual_top3d_v002.glb",
        ),
        "ENV-BASE99-WALL-CONTENTS-V021": (
            "assets/art/environments/base_facility_3d/runtime/env_base99_wall_contents_v021/env_base99_wall_contents_root_top3d_v003.tscn",
            "v003",
            "已导入；优化完成",
            master + "; assets/art/environments/base_facility_3d/source/env_base99_wall_contents_v021_manifest.json",
        ),
        "ENV-BASE99-REMAINING-FACILITIES-V021": (
            "assets/art/environments/base_facility_3d/runtime/env_base99_remaining_facilities_v021/env_base99_remaining_facilities_root_top3d_v004.tscn",
            "v004",
            "已导入；优化完成",
            master + "; assets/art/environments/base_facility_3d/source/env_base99_remaining_facilities_v021_manifest.json",
        ),
    }

    for row in range(6, main.max_row + 1):
        asset_id = str(main.cell(row, 1).value or "").strip()
        if "BASE99" not in asset_id.upper() and asset_id != "ENV-TOWER-CORNER-L-5M":
            continue

        # 原来是硬编码 `row == 235`（单体账本行号）。拆账本后行号必然改变，
        # 所以改按 AssetID 认行 —— 行号会漂，资产身份不会。
        if asset_id == "ENV-BASE99-MODULAR-KIT-3D":
            main.cell(row, 11).value = "弃用"
            main.cell(row, 15).value = None
            main.cell(row, 16).value = "聚合目录条目，不属于独立 PackedScene；由布局与 BPK 子项登记"
            main.cell(row, 20).value = None
            updated["deprecated"] += 1
            continue

        if asset_id.startswith("BPK-BASE99-"):
            slug = asset_id.removeprefix("BPK-BASE99-").lower().replace("-", "_")
            manifest = manifest_map.get(slug)
            if manifest is None:
                raise RuntimeError(f"Missing manifest for {asset_id} ({slug})")
            _version, manifest_path, manifest_data = manifest
            main.cell(row, 15).value = rel(root, manifest_path)
            main.cell(row, 16).value = master + "; " + rel(root, manifest_path)
            main.cell(row, 13).value = str(manifest_data.get("version", main.cell(row, 13).value or "v001"))
            updated["manifest_rows"] += 1

        if asset_id in explicit:
            path, version, status, source = explicit[asset_id]
            main.cell(row, 15).value = path
            main.cell(row, 13).value = version
            main.cell(row, 11).value = status
            main.cell(row, 16).value = source
            updated["explicit_rows"] += 1

        path_value = str(main.cell(row, 15).value or "").strip()
        if path_value and not path_value.startswith("="):
            path = root / path_value
            if path.is_file():
                main.cell(row, 20).value = sha256(path)
                updated["sha_rows"] += 1

    end = main.max_row
    for row in range(6, end + 1):
        set_formula(main, row, end)
    updated["formula_rows"] = end - 5

    overview["A6"] = f"=COUNTA('资产主表'!$A$6:$A${end})"
    active_formula = "+".join(f'COUNTIF(\'资产主表\'!$K$6:$K${end},"{status}")' for status in DONE_STATUSES)
    overview["C6"] = "=" + active_formula
    overview["E6"] = f'=COUNTIF(\'资产主表\'!$K$6:$K${end},"待制作")+COUNTIF(\'资产主表\'!$K$6:$K${end},"程序占位")'
    overview["G6"] = f'=COUNTIF(\'资产主表\'!$S$6:$S${end},"重复")'
    # 分账本《总览》的大类行只到本域的大类数为止（场景账本 = 3 行），不能照抄单体账本的 10..18。
    row = 10
    status_sum = "+".join(f'(\'资产主表\'!$K$6:$K${end}="{status}")' for status in DONE_STATUSES)
    while overview.cell(row, 1).value not in (None, ""):
        overview.cell(row, 2).value = f"=COUNTIF('资产主表'!$C$6:$C${end},A{row})"
        overview.cell(row, 3).value = f"=SUMPRODUCT(('资产主表'!$C$6:$C${end}=A{row})*({status_sum}))"
        row += 1

    if not dry_run:
        workbook.save(workbook_path)
    return dict(updated)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--project", type=Path, default=Path(__file__).resolve().parents[2])
    parser.add_argument("--dry-run", action="store_true")
    args = parser.parse_args()
    result = update_workbook(args.project.resolve(), args.dry_run)
    print("BASE99_LEDGER_SYNC", result, "dry_run=" + str(args.dry_run))


if __name__ == "__main__":
    main()
