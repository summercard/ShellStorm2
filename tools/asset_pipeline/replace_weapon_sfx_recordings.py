"""Install the reviewed CC0 recording selection, preserving unrelated sound sets.

Inputs are downloaded to workspace_paths' 2026-10-19 session beforehand.
No procedural oscillators, pitch sweeps or synthetic chimes are used.
"""
import csv
import hashlib
import json
import shutil
import subprocess
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / 'scripts'))
from workspace_paths import work_path

SESSION = work_path('sessions/2026-10-19_realistic_sfx/inputs.json').parent
RUNTIME = ROOT / 'src/assets/audio/sfx'
BACKUP = ROOT.parent / 'ShellStorm2-backup/2026-10-19_realistic_sfx/original_tree'
SHA = lambda p: hashlib.sha256(p.read_bytes()).hexdigest()


def ffmpeg(*args):
    subprocess.run(['ffmpeg', '-hide_banner', '-loglevel', 'error', '-y', *map(str, args)], check=True)


def backup(p):
    if p.exists():
        dest = BACKUP / p.relative_to(ROOT)
        dest.parent.mkdir(parents=True, exist_ok=True)
        if not dest.exists():
            shutil.copy2(p, dest)


def main():
    catalog_path = RUNTIME / 'sfx_catalog.json'
    catalog = json.loads(catalog_path.read_text(encoding='utf-8'))
    if catalog.get('weapon_profiles'):
        raise SystemExit('Gun-family profiles are installed; do not overwrite them with the earlier generic selection.')
    # File, excerpt start, maximum duration. Every shot is one trigger, never a burst.
    selections = {
        'pistol_fire': [('Rifle_FireLocal.opus', 0, .36), ('Rifle_Fire.opus', 0, .34), ('Rifle_Fire1.opus', 0, .37)],
        'rifle_fire': [('SMG_V2Local%d.opus' % i, 0, .43) for i in (1, 2, 3)],
        'smg_fire': [('SMG_V2Local%d.opus' % i, 0, .23) for i in (2, 3, 4)],
        'shotgun_fire': [('Shotgun_V2AmbienceSmall.opus', 0, .70), ('Rifle_Fire3.opus', 0, .66), ('Rifle_Fire2.opus', 0, .52)],
        'sniper_fire': [('Rifle_Fire3.opus', 0, .77), ('Rifle_Fire2.opus', 0, .55), ('Rifle_Fire1.opus', 0, .45)],
        'laser_fire': [('SMG_V2Local%d.opus' % i, 0, .32) for i in (1, 3, 4)],
        'reload': [('Rifle_Reload.opus', 0, 1.58)],
        'reload_action': [('SMG_Reload.opus', .43, .62)],
        'reload_end': [('Rifle_RaiseLocal.opus', .05, .32)],
        'ammo_empty': [('SMG_RaiseLocal.opus', .06, .10)],
        'item_pickup': [('Rifle_RaiseLocal.opus', 0, .39)],
        'weapon_swap': [('SMG_RaiseLocal.opus', 0, .39)],
        'weapon_attach': [('Rifle_RaiseLocal.opus', .05, .30)],
    }
    changed = []
    replacements = {}
    for asset in catalog['assets']:
        key = asset['key']
        if key not in selections:
            continue
        if asset.get('production_revision') == 'recorded-weapons-r3':
            raise RuntimeError('Already installed; restore backup before reauthoring')
        filename, start, duration = selections[key][asset['variant']]
        raw = SESSION / 'recordings' / filename
        assert raw.is_file(), raw
        old_source = ROOT / asset['source_path']
        source_dir = old_source.parent
        backup(old_source)
        backup(source_dir / 'recipe.json')
        old_runtime = ROOT / asset['runtime_path'].removeprefix('res://')
        backup(old_runtime)
        version = asset['version'] + 1
        suffix = '_v%03d' % version
        new_source = old_source.with_name(old_source.stem.rsplit('_v', 1)[0] + suffix + '.wav')
        new_runtime = RUNTIME / (old_runtime.stem.rsplit('_v', 1)[0] + suffix + '.ogg')
        recording = source_dir / ('recording_' + filename)
        shutil.copy2(raw, recording)
        stage = SESSION / ('edit_' + new_source.name)
        filters = 'highpass=f=75,lowpass=f=14500,afade=t=in:st=0:d=0.001,afade=t=out:st=%.5f:d=0.025' % (duration-.025)
        ffmpeg('-ss', start, '-i', raw, '-t', duration, '-ac', 1, '-ar', 44100, '-af', filters, '-c:a', 'pcm_s24le', stage)
        # Peak normalization retains waveform dynamics, with headroom for automatic fire.
        import numpy as np
        pcm = subprocess.check_output(['ffmpeg','-v','error','-i',str(stage),'-f','f32le','-ac', '1','-'])
        peak = float(np.max(np.abs(np.frombuffer(pcm, dtype='<f4'))))
        target = -10 if key.endswith('_fire') else -14
        gain = 10**(target/20) / max(peak, .00001)
        ffmpeg('-i', stage, '-af', 'volume=%.9f' % gain, '-c:a', 'pcm_s24le', new_source)
        ffmpeg('-i', new_source, '-c:a', 'libvorbis', '-q:a', 6, new_runtime)
        actual_duration = float(subprocess.check_output(['ffprobe','-v','error','-show_entries','format=duration','-of','default=nw=1:nk=1',str(new_runtime)]))
        old_path = asset['runtime_path']
        new_path = 'res://' + new_runtime.relative_to(ROOT).as_posix()
        replacements[old_path] = new_path
        credit = 'Ben Jaszczak, Brian Nelson, Kevin Heras, Matthew Nanney'
        source_url = 'https://opengameart.org/content/the-free-firearm-sound-library'
        if 'RaiseLocal' in filename:
            credit = 'LFA' if filename.startswith('SMG') else 'SpringySpringo'
            source_url = 'https://opengameart.org/content/' + ('equipment-clicks-iii' if filename.startswith('SMG') else 'gun-reload-sounds')
        elif filename == 'Rifle_Reload.opus':
            credit = 'zer0_sol'
            source_url = 'https://opengameart.org/content/handgun-reload-sound-effect'
        elif filename == 'SMG_Reload.opus':
            credit = 'LFA'
            source_url = 'https://opengameart.org/content/equipment-clicks-iii'
        asset.update(source_path=new_source.relative_to(ROOT).as_posix(), runtime_path=new_path,
                     version=version, duration_sec=actual_duration, sha256=SHA(new_runtime),
                     source_sha256=SHA(new_source), license_source=source_url + ' | ' + credit,
                     license_scope='CC0-1.0; edited recording; OpenSpades free pak distribution',
                     peak_target_db=target, status='imported_pending_listening', production_revision='recorded-weapons-r3')
        asset.pop('seed', None)
        recipe = {'renderer': 'tools/asset_pipeline/replace_weapon_sfx_recordings.py', 'asset': asset,
                  'input': recording.name, 'input_sha256': SHA(recording), 'start_sec': start,
                  'duration_limit_sec': duration, 'filters': filters, 'gain': gain,
                  'download_url': 'https://raw.githubusercontent.com/DeathByDenim/openspades-free-pak/main/Sounds/Weapons/' + filename.replace('_','/',1),
                  'license': 'CC0-1.0', 'attribution': credit, 'source_page': source_url}
        (source_dir / 'recipe.json').write_text(json.dumps(recipe, ensure_ascii=False, indent=2)+'\n', encoding='utf-8')
        for old in [old_source, old_runtime]:
            old.unlink()
        stems = source_dir / 'stems'
        if stems.exists():
            for p in stems.rglob('*'):
                if p.is_file():
                    backup(p)
            shutil.rmtree(stems)
        changed.append(asset)
    for sound_set in catalog['sound_sets'].values():
        sound_set['paths'] = [replacements.get(p,p) for p in sound_set['paths']]
    backup(catalog_path)
    catalog['revision'] = 'industrial-r2 + recorded-weapons-r3'
    catalog_path.write_text(json.dumps(catalog, ensure_ascii=False, indent=2)+'\n', encoding='utf-8')
    manager = ROOT/'src/core/AudioManager.gd'
    backup(manager)
    text = manager.read_text(encoding='utf-8')
    for old,new in replacements.items():
        text = text.replace(old,new)
    manager.write_text(text, encoding='utf-8')
    manifest = RUNTIME/'audio_manifest.csv'
    backup(manifest)
    with manifest.open(encoding='utf-8',newline='') as f:
        reader=csv.DictReader(f); fields=reader.fieldnames; rows=list(reader)
    by_id={a['asset_id']:a for a in changed}
    for row in rows:
        a=by_id.get(row['asset_id'])
        if a:
            row.update(source_master=a['source_path'],runtime_file=Path(a['runtime_path']).name,
                       runtime_resource_path=a['runtime_path'],sha256=a['sha256'],duration_sec=a['duration_sec'],status='录音替换；听感待验收')
    with manifest.open('w',encoding='utf-8',newline='') as f:
        writer=csv.DictWriter(f,fieldnames=fields);writer.writeheader();writer.writerows(rows)
    # Existing rows only; preserve formulas, validations and all unrelated cells.
    import openpyxl
    index=json.loads((ROOT/'assets/registry/ledger_index.json').read_text(encoding='utf-8'))
    domain=next(d for d in index['domains'] if d['key']=='audio')
    ledger=ROOT/index['ledger_dir']/domain['file'];backup(ledger)
    wb=openpyxl.load_workbook(ledger); ws=wb[index['asset_sheet']]
    cell_changes=[]
    for row in range(6,ws.max_row+1):
        a=by_id.get(ws.cell(row,1).value)
        if not a: continue
        for col,value in {13:'v%03d'%a['version'],15:a['runtime_path'].removeprefix('res://'),16:a['source_path'],20:a['sha256'],23:a['license_source']+' | CC0-1.0',25:'真实录音剪辑替换；人工听感及Android待验；verify_audio_library_flow'}.items():
            cell_changes.append([row,col,str(ws.cell(row,col).value),value]);ws.cell(row,col,value)
    assert len(cell_changes)==len(changed)*6
    wb.save(ledger)
    sys.path.insert(0,str(ROOT/'tools/asset_pipeline'))
    from split_asset_ledger import read_source_rows,_row_digest
    baseline=ROOT/'assets/registry/ledger_split_baseline.json';backup(baseline)
    bl=json.loads(baseline.read_text(encoding='utf-8'))
    for _,values in read_source_rows(ws):
        if values[0] in by_id: bl['assets'][values[0]]['v']=_row_digest(values)
    baseline.write_text(json.dumps(bl,ensure_ascii=False,indent=1)+'\n',encoding='utf-8')
    shutil.copy2(SESSION/'openspades_credits.md',ROOT/'art_source/audio/sfx/recording_credits_20261019.md')
    work_path('sessions/2026-10-19_realistic_sfx/changed_assets.json').write_text(json.dumps({'assets':changed,'cell_changes':cell_changes},ensure_ascii=False,indent=2),encoding='utf-8')
    print('Replaced',len(changed),'assets in',len(selections),'sets')


if __name__ == '__main__':
    main()
