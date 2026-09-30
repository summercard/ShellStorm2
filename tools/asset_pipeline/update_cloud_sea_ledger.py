"""Upgrade the existing cloud AssetID using separate main-row and Prefab-sheet transactions.
Preserve every other asset fingerprint and every unrelated workbook cell.
"""
from pathlib import Path
from copy import deepcopy
from collections import Counter
import argparse,json,hashlib,sys,shutil,subprocess
import openpyxl
ROOT=Path(__file__).resolve().parents[2]
sys.path.insert(0,str(ROOT/'scripts'))
sys.path.insert(0,str(ROOT/'tools/asset_pipeline'))
from ledger_registry import LedgerIndex
from split_asset_ledger import read_source_rows,_row_digest,col_digest,CONTENT_COLUMNS,sheet_digest
ID='VFX-ENV-CLOUD-SEA-3D'
BASE=ROOT/'assets/registry/ledger_split_baseline.json'
PREFAB=ROOT/'assets/art/vfx/environment_3d/cloud_sea/vfx_env_cloud_sea_root_top3d.tscn'
OUT=ROOT/'outputs/outdoor_clouds/ledger_v002'
SPEC='连续世界密度云海476×90×536m，Y=-105～-15，云顶圆润起伏、云内遮光、下层薄雾、楼边透出下方楼层。24块漂浮小云，低档12块；高/中/低96/64/40步体积采样、4/3/2步定向透射。image-2原创1024²灰度密度，128³ RGB周期云结构，256×80×288 R8建筑避让；无碰撞。'
ORIGIN='项目原创Shader/三维程序密度；gpt-image-2原创灰度云贴图，用户参考图仅作风格参考'
NOTE='v002品质优先版，尚未专项优化。主场景OutdoorClouds持有；主塔全高度禁云，700建筑边界遮罩随最新布局重烘焙。实际天台520m及99F145m视距均验证云可见、亮部有余量；GPU实际纹理采样、暂停、60秒流动、失效遮罩拒绝通过。原AssetID升级，不新增行。'
def snapshots(wb):
 return {(w.title,c.coordinate):(c.value,c.style_id,c.number_format) for w in wb for row in w for c in row if c.value is not None}
def gates(label):
 for script,args in [('scripts/check_asset_registry.py',['--project-root',str(ROOT),'--scope','structure']),('tools/asset_pipeline/verify_ledger_split.py',['--project-root',str(ROOT)])]:
  p=subprocess.run([sys.executable,'-X','utf8',str(ROOT/script),*args],cwd=OUT,text=True,encoding='utf-8',capture_output=True)
  (OUT/(label+'_'+Path(script).stem+'.log')).write_text(p.stdout+p.stderr,encoding='utf-8')
  assert p.returncode==0,p.stdout+p.stderr
def main():
 parser=argparse.ArgumentParser();parser.add_argument('phase',choices=['main','prefab']);phase=parser.parse_args().phase
 OUT.mkdir(parents=True,exist_ok=True);index=LedgerIndex.load(ROOT);path=index.domain_for_key('vfx').path
 shutil.copy2(path,OUT/('before_'+phase+'.xlsx'));shutil.copy2(BASE,OUT/('before_'+phase+'_baseline.json'))
 wb=openpyxl.load_workbook(path);before=snapshots(wb);bl=json.loads(BASE.read_text('utf-8'));old=deepcopy(bl);allowed=set()
 if phase=='main':
  w=wb['资产主表'];r=next(c.row for row in w for c in row if c.column==1 and c.value==ID)
  updates={2:'主塔柔软体积云海与漂浮小云',11:'正式美术已接入',13:'v002',14:SPEC,16:'src/vfx/VfxCloudSea3D.gd; docs/v0.1/design/outdoor_cloud_sea.md; source/v002',17:'云海;体积云;光遇风格;柔软;漂浮;薄雾;室外',20:hashlib.sha256(PREFAB.read_bytes()).hexdigest(),23:ORIGIN,25:NOTE}
  for c,v in updates.items():w.cell(r,c).value=v;allowed.add((w.title,w.cell(r,c).coordinate))
  row_values=next(v for rid,v in read_source_rows(w) if v[0]==ID)
  bl['assets'][ID]['v']=_row_digest(row_values)
  rows=[]
  for domain in index.domains:rows.extend(read_source_rows((wb if domain.key=='vfx' else openpyxl.load_workbook(domain.path))['资产主表']))
  bl['column_digests']={str(c):col_digest(rows,c) for c in CONTENT_COLUMNS}
  bl['category_counts']=dict(sorted(Counter(v[2] for _,v in rows).items()))
 else:
  w=wb['3D-特效'];r=next(c.row for row in w for c in row if c.column==2 and c.value==ID)
  updates={3:'主塔柔软体积云海与漂浮小云',7:'室外常驻连续云海、下沉雾层、柔软漂浮小云',12:SPEC,13:ORIGIN,15:'正式美术已接入',16:'v002',17:'源：assets/art/vfx/environment_3d/cloud_sea/source/v002；Manifest：assets/art/vfx/environment_3d/cloud_sea/asset_manifest.json；'+NOTE}
  for c,v in updates.items():w.cell(r,c).value=v;allowed.add((w.title,w.cell(r,c).coordinate))
  bl['sheet_digests']['3D-特效']=sheet_digest(w)
 for key,v in old['assets'].items():
  if phase=='main' and key==ID:continue
  assert bl['assets'][key]==v,key
 assert bl['asset_count']==old['asset_count']
 changes=[key for key,v in before.items() if snapshots(wb).get(key)!=v]
 assert set(changes)<=allowed,changes
 tmp=path.with_suffix('.cloud_v002_tmp.xlsx');wb.save(tmp);tmp.replace(path)
 saved=openpyxl.load_workbook(path);assert snapshots(saved)==snapshots(wb)
 BASE.write_text(json.dumps(bl,ensure_ascii=False,indent=1)+'\n',encoding='utf-8')
 (OUT/(phase+'_transaction.json')).write_text(json.dumps({'asset_id':ID,'row':r,'changes':changes,'other_asset_fingerprints_preserved':len(old['assets'])-1},ensure_ascii=False,indent=2),encoding='utf-8')
 gates(phase);print('CLOUD_V002_LEDGER_OK phase='+phase+' row='+str(r))
if __name__=='__main__':main()
