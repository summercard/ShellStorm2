"""Register only the accepted cross-tower assembly and its edited main scene."""
import json
import re
import shutil
from collections import Counter
from copy import copy
from datetime import date
from pathlib import Path
from openpyxl import load_workbook
from register_openworld_tower_runtime import (
    ROOT, LedgerIndex, FIRST_DATA_ROW, CONTENT_COLUMNS, _row_digest,
    col_digest, dedupe_key_formula, dedupe_result_formula, read_source_rows,
    sheet_digest, sha, write_json,
)


def main():
    index = LedgerIndex.load(ROOT)
    domain = next(d for d in index.domains if d.key == 'scenes')
    baseline_path = ROOT / 'assets/registry/ledger_split_baseline.json'
    route = 'assets/art/environments/open_world/runtime/cross_tower_route/env_cross_tower_route_root_top3d.tscn'
    report = json.loads((ROOT / Path(route).parent / 'acceptance.json').read_text(encoding='utf8'))
    assert report['checks'] >= 308 and not report['failures']
    assert report['renderer'] != 'headless' and report['scene_sha256'] == sha(ROOT / route)
    backup = ROOT.parent / '_scratch/cross_tower_ledger_backup'
    backup.mkdir(parents=True, exist_ok=True)
    for path in [domain.path, index.master_path, baseline_path]:
        if not (backup / path.name).exists():
            shutil.copy2(path, backup / path.name)
    wb = load_workbook(domain.path)
    ws = wb['资产主表']
    old = ws.max_row
    aid = 'ENV-OPENWORLD-CROSS-TOWER-ROUTE'
    known = {ws.cell(r, 1).value: r for r in range(FIRST_DATA_ROW, old + 1)}
    row = known.get(aid, old + 1)
    if aid not in known:
        for c in range(1, 26):
            ws.cell(row, c)._style = copy(ws.cell(old, c)._style)
        ws.row_dimensions[row].height = ws.row_dimensions[old].height
    values = {
        1: aid, 2: '主塔100层跨塔天桥', 3: '场景', 4: 'environment_module_3d',
        5: 'openworld_rooftop_cross_tower_route', 6: 'cross_tower_route',
        8: 'Top3D / Godot Y-up', 9: 'default', 10: '主塔100F外部路线；独立PackedScene',
        11: '已导入；优化完成', 12: 'P1', 13: 'v002',
        14: '原塔吊32m起重臂；两段各64m；检修走道净宽1.65m；楼间净距30m；塔2不可进入；塔3行走面0m',
        15: route, 16: 'docs/v0.1/design/rooftop_cross_tower_route.md',
        17: '天台；桥；塔2；塔3', 20: sha(ROOT / route), 21: 'Codex',
        22: date(2026, 9, 30), 23: '用户示意图；项目既有组件复用',
        24: 'cross_tower_route',
        25: '独立路线包装拥有承重及阻挡碰撞；未修改原塔楼/塔吊Prefab、源Blend和材质；不用minimap或新增MultiMesh。',
    }
    for c, value in values.items():
        ws.cell(row, c, value)
    changed = {aid}
    # The only existing resource intentionally edited in this transaction.
    main_id = 'ENV-TOWER-DESCENT-KIT-3D'
    main_row = known[main_id]
    assert ws.cell(main_row, 15).value == 'scenes/TowerDescent3D.tscn'
    ws.cell(main_row, 20, sha(ROOT / 'scenes/TowerDescent3D.tscn'))
    ws.cell(main_row, 22, date(2026, 9, 30))
    changed.add(main_id)
    last = ws.max_row
    for r in range(FIRST_DATA_ROW, last + 1):
        ws.cell(r, 18, dedupe_key_formula(r))
        ws.cell(r, 19, dedupe_result_formula(r, last))
    for overview_row in wb['总览']:
        for cell in overview_row:
            if cell.data_type == 'f':
                cell.value = re.sub(r'(\$[A-Z]+\$)' + str(old) + r'\b', lambda m: m[1] + str(last), cell.value)
    for dv in ws.data_validations.dataValidation:
        ranges = []
        for item in dv.sqref.ranges:
            item = copy(item)
            if item.max_row == old:
                item.max_row = last
            ranges.append(str(item))
        dv.sqref = ' '.join(ranges)
    for cf in list(ws.conditional_formatting):
        rules = ws.conditional_formatting[cf]
        ranges = []
        for item in cf.sqref.ranges:
            item = copy(item)
            if item.max_row == old:
                item.max_row = last
            ranges.append(str(item))
        del ws.conditional_formatting[' '.join(str(item) for item in cf.sqref.ranges)]
        for rule in rules:
            ws.conditional_formatting.add(' '.join(ranges), rule)
    ws.auto_filter.ref = f'A{FIRST_DATA_ROW-1}:Y{last}'
    log = wb['域变更日志']
    message = '主塔100F北侧开5m桥口；两段桥改用塔2原塔吊起重臂（两节32m/段）；跨低塔2到等高塔3；无新增材质。'
    if log.cell(log.max_row, 5).value != message:
        bits = log.cell(log.max_row, 1).value.split('.')
        bits[-1] = str(int(bits[-1]) + 1)
        log.append(['.'.join(bits), date(2026, 9, 30), '吊臂跨塔天桥接入', '关卡场景 / PackedScene', message, f"{report['checks']}项真实渲染与物理验收；仅更新两项身份。", 'Codex'])
    else:
        log.cell(log.max_row, 6, f"{report['checks']}项真实渲染与物理验收；仅更新两项身份。")
    wb.save(domain.path)
    baseline = json.loads(baseline_path.read_text(encoding='utf8'))
    saved = load_workbook(domain.path)
    for _, vals in read_source_rows(saved['资产主表']):
        if vals[0] in changed:
            baseline['assets'][vals[0]] = {'v': _row_digest(vals), 'c': vals[2], 'd': 'scenes'}
    all_rows = []
    for d in index.domains:
        all_rows.extend((vals[0], vals) for _, vals in read_source_rows(load_workbook(d.path)['资产主表']))
    all_rows.sort(key=lambda item: item[0])
    baseline['asset_count'] = len(baseline['assets'])
    baseline['category_counts'] = dict(Counter(vals[2] for _, vals in all_rows))
    baseline['column_digests'] = {str(c): col_digest(all_rows, c) for c in CONTENT_COLUMNS}
    write_json(baseline_path, baseline, 1)
    # Separate, explicit Prefab-page transaction. No other locked sheet is edited.
    page = saved['3D-场景通用']
    page_rows = {page.cell(r, 1).value: r for r in range(1, page.max_row + 1)}
    r = page_rows.get(aid, page.max_row + 1)
    template = page.max_row
    for c in range(1, 17):
        page.cell(r, c)._style = copy(page.cell(template, c)._style)
    page.row_dimensions[r].height = page.row_dimensions[template].height
    prefab = [aid, '主塔100层跨塔天桥', route, None, 'Godot原生装配；源塔楼和桥梁组件稳定Prefab',
              'CrossTowerRoute独立路线', None, '开', '静态碰撞', '分段BoxShape3D',
              values[14], 'Godot Y-up；根缩放1', 'Blocks/Rooftop/CrossTowerRoute',
              '已导入；优化完成', 'v002', '原塔吊吊臂及自带检修走道；普通独立实例；无新增材质与minimap；塔2不可进入。']
    for c, value in enumerate(prefab, 1):
        page.cell(r, c).value = value
    page.auto_filter.ref = f'A{FIRST_DATA_ROW-1}:P{page.max_row}'
    saved.save(domain.path)
    baseline['sheet_digests']['3D-场景通用'] = sheet_digest(page)
    write_json(baseline_path, baseline, 1)
    master = load_workbook(index.master_path)
    for offset, d in enumerate(index.domains):
        if d.key == 'scenes':
            master['分账本索引'].cell(FIRST_DATA_ROW + offset, 7, last - FIRST_DATA_ROW + 1)
    for cells in master['3D Prefab总控']:
        if any(c.value == '3D-场景通用' for c in cells):
            actual = [vals for vals in page.values if isinstance(vals[0], str) and re.fullmatch(r'[A-Z0-9]+(?:-[A-Z0-9]+)+', vals[0])]
            cells[2].value = len(actual)
            cells[3].value = sum(bool(vals[3]) for vals in actual)
            cells[4].value = sum(bool(vals[6]) for vals in actual)
            cells[5].value = sum(vals[7] == '开' for vals in actual)
    master.save(index.master_path)
    print('CROSS_TOWER_LEDGER_OK', row, 'updated', sorted(changed))


if __name__ == '__main__':
    main()
