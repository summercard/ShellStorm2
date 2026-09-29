"""Record the two 99F driveable seats and Bunny seated animation without adding AssetIDs."""

from collections import Counter
from datetime import datetime
from pathlib import Path
import hashlib
import json
import shutil
import sys

import openpyxl

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
sys.path.insert(0, str(HERE))
from split_asset_ledger import CONTENT_COLUMNS, _row_digest, col_digest, read_source_rows  # noqa: E402

BOOKS = ROOT / "assets/registry/ledgers"
BACKUP = ROOT / "_scratch/seat_ledger_before_20260929"
BASELINE = ROOT / "assets/registry/ledger_split_baseline.json"
TODAY = datetime(2026, 9, 29)


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def update_book(filename: str, changes: dict[str, dict[int, object]], log: tuple[str, ...]) -> None:
    path = BOOKS / filename
    workbook = openpyxl.load_workbook(path)
    sheet = workbook["资产主表"]
    for asset_id, fields in changes.items():
        found = [r for r in range(6, sheet.max_row + 1) if sheet.cell(r, 1).value == asset_id]
        assert len(found) == 1, (filename, asset_id, found)
        row = found[0]
        for col, value in fields.items():
            sheet.cell(row, col).value = value
        sheet.cell(row, 22).value = TODAY
    log_sheet = workbook["域变更日志"]
    next_row = log_sheet.max_row + 1
    for col in range(1, 8):
        log_sheet.cell(next_row, col)._style = log_sheet.cell(next_row - 1, col)._style
    for col, value in enumerate(log, 1):
        log_sheet.cell(next_row, col).value = value
    workbook.save(path)


def main() -> None:
    for filename, ids in [
        ("ShellStorm2_角色账本_v001.xlsx", ["CHR-PLY-CAPSULE01", "CHR-PLY-CAPSULE01-3D-BUNNY01"]),
        ("ShellStorm2_道具账本_v001.xlsx", ["PRP-BASE-WORKSHOP-STOOL-3D", "PRP-BASE-MISSION-COMMAND-CHAIR-3D"]),
    ]:
        sheet = openpyxl.load_workbook(BOOKS / filename, read_only=True)["资产主表"]
        for row in sheet.iter_rows(min_row=6, max_col=9):
            if row[0].value in ids:
                assert "seated" not in str(row[8].value), f"already updated: {row[0].value}"
    BACKUP.mkdir(parents=True, exist_ok=True)
    for name in ["ShellStorm2_角色账本_v001.xlsx", "ShellStorm2_道具账本_v001.xlsx"]:
        shutil.copy2(BOOKS / name, BACKUP / name)
    shutil.copy2(BASELINE, BACKUP / BASELINE.name)

    animation = "assets/art/characters/player/chr_player_capsule01_3d/variants/bunny01/production/v022/source/animation/chr_bunny01_animation_v022.blend"
    transfer = "assets/art/characters/player/chr_player_capsule01_3d/variants/bunny01/production/v022/character_transfer_ledger_v022.json"
    character_book = "ShellStorm2_角色账本_v001.xlsx"
    character = openpyxl.load_workbook(BOOKS / character_book)
    character_sheet = character["资产主表"]
    char_changes = {}
    for row, asset_id in [(6, "CHR-PLY-CAPSULE01"), (13, "CHR-PLY-CAPSULE01-3D-BUNNY01")]:
        assert character_sheet.cell(row, 1).value == asset_id
        old_source = str(character_sheet.cell(row, 16).value)
        old_note = str(character_sheet.cell(row, 25).value)
        char_changes[asset_id] = {
            9: "default / seated",
            16: old_source + "; " + animation + "; " + transfer,
            20: sha256(ROOT / str(character_sheet.cell(row, 15).value)),
            25: old_note + "\n2026-09-29：v021 模型及原12动作保持不变；新增 v022 Blender 坐姿滑行循环 seated，Player3D 九态状态机与 99F 两把可驾驶椅接入；verify_pushable_base_chairs 通过。",
        }
    update_book(character_book, char_changes, (
        "v0.1.1", "2026-09-29", "玩家坐姿动作接入", "角色 / Bunny01",
        "v022 动画源新增 seated 循环，原 12 动作逐剪辑一致；玩法状态机新增 seated。",
        "更新既有玩家与 Bunny 根资产的状态、源文件、备注和日期；模型/Prefab 路径不变，SHA 同步当前文件。", "Codex",
    ))

    prop_book = "ShellStorm2_道具账本_v001.xlsx"
    prop = openpyxl.load_workbook(BOOKS / prop_book)
    prop_sheet = prop["资产主表"]
    prop_changes = {}
    for row, asset_id, source in [
        (18, "PRP-BASE-WORKSHOP-STOOL-3D", "workshop_stool"),
        (19, "PRP-BASE-MISSION-COMMAND-CHAIR-3D", "mission_command_chair"),
    ]:
        assert prop_sheet.cell(row, 1).value == asset_id
        runtime = str(prop_sheet.cell(row, 15).value)
        assert (ROOT / runtime).is_file()
        old_note = str(prop_sheet.cell(row, 25).value)
        prop_changes[asset_id] = {
            9: "default / pushable / seated / driveable",
            14: "独立GLB/PackedScene；实体刚体与场景共用碰撞层；E乘坐/离开；轻微滚轮滑行",
            16: f"assets/art/props/base_world_3d/source/{source}/prp_base_{source}_source_v003.blend",
            20: sha256(ROOT / runtime),
            25: old_note + "\n2026-09-29：99F 基地座椅接入 PushableSeat3D，角色碰撞可推动；E 上下座，方向输入通过刚体受力滑行；场景墙/地面碰撞保留，安全下座查询；verify_pushable_base_chairs 通过。",
        }
    update_book(prop_book, prop_changes, (
        "v0.1.1", "2026-09-29", "99F座椅乘坐与驾驶", "道具 / 基地座椅",
        "维修圆凳与战术指挥椅改为可推刚体、可乘坐驾驶；与场景碰撞交互。",
        "更新两条既有资产的状态、规格、v003 源、SHA、备注和日期；AssetID 与运行时路径不变。", "Codex",
    ))

    baseline = json.loads(BASELINE.read_text(encoding="utf-8"))
    index = json.loads((ROOT / "assets/registry/ledger_index.json").read_text(encoding="utf-8"))
    all_rows = []
    touched = set(char_changes) | set(prop_changes)
    for domain in index["domains"]:
        book = BOOKS / domain["file"]
        workbook = openpyxl.load_workbook(book, read_only=True)
        for _row_number, values in read_source_rows(workbook["资产主表"]):
            all_rows.append((_row_number, values))
            asset_id = str(values[0])
            if asset_id in touched:
                baseline["assets"][asset_id]["v"] = _row_digest(values)
    baseline["captured_at"] = "2026-09-29"
    baseline["asset_count"] = len(baseline["assets"])
    baseline["column_digests"] = {str(col): col_digest(all_rows, col) for col in CONTENT_COLUMNS}
    baseline["category_counts"] = dict(sorted(Counter(str(values[2]) for _, values in all_rows).items()))
    BASELINE.write_text(json.dumps(baseline, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")
    print("SEATED_CHAIRS_LEDGER_UPDATED", len(touched))


if __name__ == "__main__":
    main()
