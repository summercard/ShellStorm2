"""Record the v023 seated pose and the 99F chair runtime regression fixes."""
from copy import copy
from datetime import datetime
from pathlib import Path
import hashlib
import json
import shutil
import sys

import openpyxl

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
sys.path.insert(0, str(HERE))
from split_asset_ledger import _row_digest, col_digest, read_source_rows, sheet_digest  # noqa: E402

BOOKS = ROOT / "assets/registry/ledgers"
BASELINE = ROOT / "assets/registry/ledger_split_baseline.json"
BACKUP = ROOT / "_scratch/seat_regression_ledger_before_20260929"
TODAY = datetime(2026, 9, 29)


def _log(workbook, version, title, scope, summary):
    sheet = workbook["域变更日志"]
    row = sheet.max_row + 1
    assert sheet.cell(row - 1, 1).value == "v0.1.2"
    values = [version, "2026-09-29", title, scope, summary,
              "更新既有资产记录与验收契约；保留 AssetID 和稳定 PackedScene 路径。", "Codex"]
    for col, value in enumerate(values, 1):
        sheet.cell(row, col)._style = copy(sheet.cell(row - 1, col)._style)
        sheet.cell(row, col).value = value


def main():
    character_path = BOOKS / "ShellStorm2_角色账本_v001.xlsx"
    prop_path = BOOKS / "ShellStorm2_道具账本_v001.xlsx"
    BACKUP.mkdir(parents=True, exist_ok=True)
    for path in (character_path, prop_path, BASELINE):
        shutil.copy2(path, BACKUP / path.name)

    character = openpyxl.load_workbook(character_path)
    sheet = character["资产主表"]
    old_base = "assets/art/characters/player/chr_player_capsule01_3d/variants/bunny01/production/v022"
    new_base = old_base.replace("v022", "v023")
    for row, asset_id in ((6, "CHR-PLY-CAPSULE01"), (13, "CHR-PLY-CAPSULE01-3D-BUNNY01")):
        assert sheet.cell(row, 1).value == asset_id
        source = str(sheet.cell(row, 16).value)
        assert source.count(old_base) == 2 and new_base not in source
        sheet.cell(row, 16).value = source.replace(old_base, new_base)
        sheet.cell(row, 20).value = hashlib.sha256((ROOT / str(sheet.cell(row, 15).value)).read_bytes()).hexdigest()
        sheet.cell(row, 22).value = TODAY
        sheet.cell(row, 25).value = str(sheet.cell(row, 25).value) + (
            "\n2026-09-29：v023 Blender 坐姿源将双脚向座面前方伸出；角色乘坐时可正常射击，"
            "相机补偿座面抬升，离座恢复世界碰撞。"
        )
    state = character["动画与状态"]
    assert state["B18"].value == "seated" and "不射击" in str(state["H18"].value)
    state["F18"] = "v023：原始 Blender 动画源迭代 1.2 秒坐姿轻摆循环；双脚前伸，不驱动椅子位移。"
    state["H18"] = "双脚前伸；双手随坐姿与武器挂点，乘坐时仍可正常射击。"
    state["I18"] = "右手与 WeaponSocket 随坐姿；开火沿用既有射击/弹药判定。"
    state["K18"] = "E 离座提示；世界碰撞离座恢复；镜头补偿座面升高"
    _log(character, "v0.1.3", "坐姿动作与乘坐射击修正", "角色 / Bunny01",
         "v023 双脚前伸 Blender 动作；乘坐仍可开枪，相机与离座碰撞恢复。")
    character.save(character_path)

    props = openpyxl.load_workbook(prop_path)
    sheet = props["资产主表"]
    for row, asset_id in ((18, "PRP-BASE-WORKSHOP-STOOL-3D"),
                          (19, "PRP-BASE-MISSION-COMMAND-CHAIR-3D")):
        assert sheet.cell(row, 1).value == asset_id
        sheet.cell(row, 20).value = hashlib.sha256((ROOT / str(sheet.cell(row, 15).value)).read_bytes()).hexdigest()
        sheet.cell(row, 22).value = TODAY
        sheet.cell(row, 25).value = str(sheet.cell(row, 25).value) + (
            "\n2026-09-29：v004 分体 GLB 导入设置补回设施共享色盘后处理；"
            "99F 实例验证底座和旋转上部均绑定色盘；离座碰撞与镜头回归通过。"
        )
    _log(props, "v0.1.3", "99F座椅材质与上下座回归", "道具 / 基地座椅",
         "两把椅子的分体视觉恢复共享材质；验证上下座碰撞、镜头高度及座中射击。")
    props.save(prop_path)

    baseline = json.loads(BASELINE.read_text(encoding="utf-8"))
    for row in (6, 13):
        values = tuple(character["资产主表"].cell(row, col).value for col in range(1, 26))
        baseline["assets"][str(values[0])]["v"] = _row_digest(values)
    for row in (18, 19):
        values = tuple(props["资产主表"].cell(row, col).value for col in range(1, 26))
        baseline["assets"][str(values[0])]["v"] = _row_digest(values)
    baseline["sheet_digests"]["动画与状态"] = sheet_digest(state)
    index = json.loads((ROOT / "assets/registry/ledger_index.json").read_text(encoding="utf-8"))
    all_rows = []
    for domain in index["domains"]:
        book = openpyxl.load_workbook(BOOKS / domain["file"], read_only=True)
        all_rows.extend(read_source_rows(book["资产主表"]))
    for col in (16, 20, 22, 25):
        baseline["column_digests"][str(col)] = col_digest(all_rows, col)
    BASELINE.write_text(json.dumps(baseline, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")
    print("SEAT_REGRESSION_LEDGER_UPDATED")


if __name__ == "__main__":
    main()
