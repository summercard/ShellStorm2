"""Author recording-based gun families and reconcile only their ledger rows.

Seven families, nine actual gun IDs; no inference from combat statistics.
Input takes: workspace session intake (Tabasco) and previous CC0 Foley intake.
"""
from pathlib import Path
from copy import copy
from collections import Counter
import csv
import datetime
import hashlib
import json
import re
import shutil
import subprocess
import sys
import numpy as np
import openpyxl

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT/'scripts'))
from workspace_paths import work_path
SESSION = work_path('sessions/2026-10-19_weapon_profiles/result.json').parent
OLD = work_path('sessions/2026-10-19_realistic_sfx/inputs.json').parent
OUT = ROOT/'src/assets/audio/sfx'
BACKUP = ROOT.parent/'ShellStorm2-backup/2026-10-19_weapon_profiles/original_tree'
RATE = 44100
sha = lambda p: hashlib.sha256(p.read_bytes()).hexdigest()
FAMILIES = ['pistol','rifle','machinegun','shotgun','sniper','launcher','charge']
NAMES = dict(zip(FAMILIES, ['手枪','步枪','机枪','霰弹枪','狙击枪','榴弹筒','蓄力炮']))


def save_json(p, value, indent=2):
    p.parent.mkdir(parents=True, exist_ok=True)
    p.write_text(json.dumps(value, ensure_ascii=False, indent=indent)+'\n', encoding='utf-8')


def backup(p):
    if p.exists():
        target=BACKUP/p.relative_to(ROOT);target.parent.mkdir(parents=True,exist_ok=True)
        assert target.resolve().is_relative_to(BACKUP.resolve())
        if not target.exists(): shutil.copy2(p,target)


def run(*args, **kwargs):
    return subprocess.run(['ffmpeg','-v','error','-y',*map(str,args)],check=True,**kwargs)


def input_path(name):
    p=SESSION/name
    if not p.exists(): p=OLD/name
    if not p.exists(): p=OLD/'recordings'/name
    assert p.is_file(),p
    return p


# Each layer: source take, start, length, offset, relative gain.
def layer(name,start,length,offset=0,gain=1): return (name,start,length,offset,gain)


def author(asset, layers, duration, cutoff, peak_db):
    source=ROOT/'art_source/audio/sfx'/asset['asset_id']/'source'
    source.mkdir(parents=True,exist_ok=True)
    if 'runtime_path' in asset:
        old_runtime=ROOT/asset['runtime_path'].removeprefix('res://')
        old_master=ROOT/asset['source_path']
        for p in [old_runtime,old_master,source/'recipe.json']: backup(p)
        version=asset['version']+1
    else: version=1
    variant=asset['variant']
    filename=asset['key']+('' if variant==0 else '_alt%02d'%(variant+1))+'_v%03d'%version
    mix=np.zeros(round(duration*RATE),dtype=np.float64)
    inputs=[]
    for name,start,length,offset,gain in layers:
        raw=input_path(name)
        recorded=source/('recording_'+raw.name)
        shutil.copy2(raw,recorded)
        args=['ffmpeg','-v','error','-ss',str(start),'-i',str(raw),'-t',str(length),'-ac','1','-ar',str(RATE),'-af','highpass=f=75,afftdn=nr=12:nf=-35','-f','f32le','-']
        x=np.frombuffer(subprocess.check_output(args),dtype='<f4').astype(np.float64)
        n=min(len(x),len(mix)-round(offset*RATE));begin=round(offset*RATE)
        if n>0:
            x=x[:n];fade=min(1103,n//4);x[:44]*=np.linspace(0,1,min(44,n));x[-fade:]*=np.linspace(1,0,fade)
            mix[begin:begin+n]+=x*gain
        if raw.name in ['cz.wav','sks.wav','shotty.wav','mosin.wav']:
            credit='Vincent Sevedge / Tabasco'
            url='https://opengameart.org/content/gunshot-sounds'
            license='CC-BY-3.0 (archive notice; page also offers CC0; retain stricter attribution)'
        elif 'gunreload' in name or 'assaultrifle' in name or name=='shotguncock.wav' or name=='Rifle_RaiseLocal.opus':
            credit='SpringySpringo';url='https://opengameart.org/content/gun-reload-sounds';license='CC0-1.0'
        elif name=='SMG_Reload.opus' or name=='SMG_RaiseLocal.opus':
            credit='LFA';url='https://opengameart.org/content/equipment-clicks-iii';license='CC0-1.0'
        else:
            credit='Ben Jaszczak, Brian Nelson, Kevin Heras, Matthew Nanney'
            url='https://opengameart.org/content/the-free-firearm-sound-library';license='CC0-1.0'
        inputs.append(dict(file=recorded.name,sha256=sha(recorded),start_sec=start,length_sec=length,offset_sec=offset,gain=gain,credit=credit,source_page=url,license=license))
    pcm=SESSION/(filename+'_mix.f32');pcm.write_bytes(mix.astype('<f4').tobytes())
    filtered=SESSION/(filename+'_filtered.f32')
    run('-f','f32le','-ar',RATE,'-ac',1,'-i',pcm,'-af',f'lowpass=f={cutoff},afade=t=out:st={duration-.03}:d=0.03','-f','f32le',filtered)
    x=np.frombuffer(filtered.read_bytes(),dtype='<f4').copy()
    # Remove recording lead-in and denoiser latency so a trigger sounds immediately.
    active=np.flatnonzero(np.abs(x)>max(float(np.max(np.abs(x)))*.03,.0001))
    trim=max(0,int(active[0])-44) if len(active) else 0
    if trim:x=np.concatenate([x[trim:],np.zeros(trim,dtype='<f4')])
    filtered.write_bytes(x.tobytes())
    gain=10**(peak_db/20)/max(float(np.max(np.abs(x))),1e-9)
    master=source/(filename+'.wav');runtime=OUT/(filename+'.ogg')
    run('-f','f32le','-ar',RATE,'-ac',1,'-i',filtered,'-af',f'volume={gain:.9f}','-c:a','pcm_s24le',master)
    run('-i',master,'-c:a','libvorbis','-q:a',6,runtime)
    asset.update(source_path=master.relative_to(ROOT).as_posix(),runtime_path='res://'+runtime.relative_to(ROOT).as_posix(),
        version=version,duration_sec=duration,source_format='PCM24 WAV',runtime_format='OGG Vorbis q6',sample_rate_hz=RATE,
        source_channels=1,runtime_channels=1,sha256=sha(runtime),source_sha256=sha(master),peak_target_db=peak_db,
        license_source='; '.join(sorted({v['source_page']+' | '+v['credit'] for v in inputs})),
        license_scope='; '.join(sorted({v['license'] for v in inputs})),status='imported_pending_listening',
        production_revision='gun-families-r4',loop=False)
    asset.pop('seed',None)
    save_json(source/'recipe.json',dict(renderer='tools/asset_pipeline/build_weapon_audio_profiles.py',asset=asset,inputs=inputs,
        duration_sec=duration,lowpass_hz=cutoff,peak_db=peak_db,normalization_gain=gain,leading_trim_samples=trim,
        input_filter='highpass=75Hz; afftdn nr=12 nf=-35; 1ms/25ms layer fades; 30ms output fade; no pitch shift'))
    # Archive superseded sources, including previous recording takes. Never delete unbacked files.
    keep={master.name,'recipe.json'}|{v['file'] for v in inputs}
    for p in list(source.iterdir()):
        if p.is_file() and p.name not in keep:
            backup(p);p.unlink()
    if 'old_runtime' in locals() and old_runtime!=runtime: old_runtime.unlink(missing_ok=True)
    return asset


def main():
    cp=OUT/'sfx_catalog.json';catalog=json.loads(cp.read_text(encoding='utf-8'))
    assert not catalog.get('weapon_profiles'),'Profiles already installed'
    backup(cp)
    changes=[];new=[];replace={}
    gun_keys={'pistol':'pistol_fire','rifle':'rifle_fire','machinegun':'smg_fire','shotgun':'shotgun_fire','sniper':'sniper_fire','launcher':'launcher_fire','charge':'laser_fire'}
    # Separate gun recordings: CZ-52, SKS, pump shotgun, Mosin; AK-47 machinegun takes retain r3.
    takes={
        'pistol':[(.218,.38),(2.81,.38),(4.04,.38)],
        'rifle':[(.33,.40),(2.29,.40),(5.99,.40)],
        'sniper':[(.418,.85),(1.658,.85),(3.558,.85)]}
    for family in ['pistol','rifle','shotgun','sniper','charge','launcher']:
        key=gun_keys[family]
        for var in range(3):
            asset=next((a for a in catalog['assets'] if a['key']==key and a['variant']==var),None)
            if asset is None:
                asset={'asset_id':'AUD-SFX-LAUNCHER-FIRE'+('' if var==0 else '-VAR%02d'%(var+1)), 'key':key,'variant':var,
                       'sound_set_id':'SFXSET-LAUNCHER-FIRE','event_ids':['combat.weapon.fire.launcher'],'spatial_mode':'world_3d'}
                new.append(asset)
            old=asset.get('runtime_path')
            if family in takes:
                name={'pistol':'cz','rifle':'sks','sniper':'mosin'}[family]+'.wav';start,duration=takes[family][var]
                layers=[layer(name,start,duration)];cutoff=14000
            elif family=='shotgun':
                duration=.55;layers=[layer('shotty.wav',.13,.55)];cutoff=[14000,12500,15500][var]
            elif family=='launcher':
                duration=.66;cutoff=[2600,2900,3100][var]
                layers=[layer('shotty.wav',.13,.55),layer('shotguncock.wav',.18,.23,.025,.22)]
            else:
                duration=.9;cutoff=[5200,5600,6000][var]
                layers=[layer('mosin.wav',takes['sniper'][var][0],.85),layer('assaultriflereload1.wav',.90,.24,.025,.13)]
            author(asset,layers,duration,cutoff,-10)
            if old: replace[old]=asset['runtime_path']
            changes.append(asset)
        catalog['sound_sets'][key]={'id':'SFXSET-'+key.upper().replace('_','-'),'paths':[a['runtime_path'] for a in changes if a['key']==key],
            'loop':False,'spatial_mode':'world_3d','bus':'SFX','max_voices':8,'priority':'P1'}
    reloads={
        'pistol':([layer('gunreload1.wav',.03,1.20)],1.24),
        'rifle':([layer('assaultriflereload1.wav',.04,1.45)],1.49),
        'machinegun':([layer('SMG_Reload.opus',.03,2.02)],2.06),
        'shotgun':([layer('gunreload1.wav',.42,.20),layer('shotguncock.wav',.01,.45,.32)],.81),
        'sniper':([layer('shotguncock.wav',.02,.44),layer('Rifle_RaiseLocal.opus',.05,.30,.55)],.89),
        'launcher':([layer('assaultriflereload1.wav',.17,.30),layer('shotguncock.wav',.16,.27,.60)],.92),
        'charge':([layer('SMG_Reload.opus',.44,.65),layer('shotguncock.wav',.03,.42,.80)],1.26)}
    pickups={
        'pistol':([layer('gunreload1.wav',.15,.22)],.26),
        'rifle':([layer('assaultriflereload1.wav',.20,.34)],.38),
        'machinegun':([layer('SMG_Reload.opus',.55,.35)],.39),
        'shotgun':([layer('shotguncock.wav',.01,.44)],.48),
        'sniper':([layer('shotguncock.wav',.23,.23),layer('Rifle_RaiseLocal.opus',.05,.13,.25,.35)],.42),
        'launcher':([layer('assaultriflereload1.wav',.34,.22),layer('shotguncock.wav',.21,.18,.20,.6)],.42),
        'charge':([layer('SMG_Reload.opus',1.3,.3),layer('Rifle_RaiseLocal.opus',.07,.19,.29,.55)],.52)}
    for action,selection in [('reload',reloads),('pickup',pickups)]:
        for family,(layers,duration) in selection.items():
            key=action+'_'+family
            asset={'asset_id':'AUD-SFX-'+key.upper().replace('_','-'),'key':key,'variant':0,
                   'sound_set_id':'SFXSET-'+key.upper().replace('_','-'),'spatial_mode':'listener_follow',
                   'event_ids':['combat.weapon.reload_start' if action=='reload' else 'world.pickup.item']}
            new.append(asset);author(asset,layers,duration,14000,-14);changes.append(asset)
            catalog['sound_sets'][key]={'id':asset['sound_set_id'],'paths':[asset['runtime_path']],'loop':False,
                'spatial_mode':'listener_follow','bus':'SFX','max_voices':2,'priority':'P1' if action=='reload' else 'P2'}
    assert len(new)==17 and len(changes)==32
    for a in new: catalog['assets'].append(a)
    for s in catalog['sound_sets'].values():s['paths']=[replace.get(p,p) for p in s['paths']]
    for family in FAMILIES:
        profile={'family':family,'fire':gun_keys[family],'reload':'reload_'+family,'pickup':'pickup_'+family}
        catalog.setdefault('weapon_profiles',{})['bp_'+family]=profile
    catalog['weapon_profiles']['bp_starter_rifle']=catalog['weapon_profiles']['bp_rifle'].copy()
    catalog['weapon_profiles']['bp_sprinkler']=catalog['weapon_profiles']['bp_machinegun'].copy()
    catalog['revision']='gun-families-r4 + retained industrial-r2 sets'
    catalog['events']['combat.weapon.fire.launcher']=dict(catalog['events']['combat.weapon.fire.sniper'],sound_set='launcher_fire',owner='WeaponModel3D')
    save_json(cp,catalog)
    manager=ROOT/'src/core/AudioManager.gd';backup(manager);text=manager.read_text(encoding='utf-8')
    for a,b in replace.items():text=text.replace(a,b)
    manager.write_text(text,encoding='utf-8')
    manifest=OUT/'audio_manifest.csv';backup(manifest)
    with manifest.open(encoding='utf-8',newline='') as f:
        reader=csv.DictReader(f);fields=reader.fieldnames;rows=list(reader)
    by_id={a['asset_id']:a for a in changes}
    by_row={r['asset_id']:r for r in rows}
    for a in changes:
        row=by_row.get(a['asset_id'])
        if row is None:row={k:'' for k in fields};rows.append(row)
        row.update(key=a['key'],asset_id=a['asset_id'],source_master=a['source_path'],runtime_file=Path(a['runtime_path']).name,
            status='枪种录音已接入；听感待验',usage=a['key'],duration_sec=a['duration_sec'],sample_rate_hz=RATE,master_channels=1,runtime_channels=1,
            runtime_resource_path=a['runtime_path'],sha256=a['sha256'],sound_set_id=a['sound_set_id'],event_ids=';'.join(a['event_ids']))
    with manifest.open('w',encoding='utf-8',newline='') as f:
        writer=csv.DictWriter(f,fieldnames=fields);writer.writeheader();writer.writerows(rows)
    index=json.loads((ROOT/'assets/registry/ledger_index.json').read_text(encoding='utf-8'))
    domain=next(d for d in index['domains'] if d['key']=='audio')
    ledger=ROOT/index['ledger_dir']/domain['file'];backup(ledger)
    wb=openpyxl.load_workbook(ledger);ws=wb[index['asset_sheet']];old_end=ws.max_row
    ids={ws.cell(r,1).value:r for r in range(6,old_end+1)}
    template=ids['AUD-SFX-PISTOL-FIRE'];template_values=[c.value for c in ws[template]]
    # Frozen new IDs are appended without disturbing existing row identities.
    for asset in new:
        assert asset['asset_id'] not in ids
        row=ws.max_row+1;ids[asset['asset_id']]=row
        for col,value in enumerate(template_values,1):
            ws.cell(row,col,value);ws.cell(row,col)._style=copy(ws.cell(template,col)._style)
    for asset in changes:
        row=ids[asset['asset_id']];key=asset['key']
        if key.endswith('_fire'):
            family=next(f for f,k in gun_keys.items() if k==key);label=NAMES[family]+'开火'
        else:
            action,family=key.split('_',1);label=NAMES[family]+('换弹' if action=='reload' else '拾取')
        if asset['variant']:label+='（变体%02d）'%(asset['variant']+1)
        vals={1:asset['asset_id'],2:label,3:'音效',4:'sfx',5:key,6:'root' if asset['variant']==0 else 'var%02d'%(asset['variant']+1),
              10:';'.join(asset['event_ids']),11:'原型已接入',12:'P1',13:'v%03d'%asset['version'],
              14:f'44100Hz / mono / PCM24 master / Vorbis q6 / {asset["duration_sec"]}s',15:asset['runtime_path'].removeprefix('res://'),
              16:asset['source_path'],17:key,20:asset['sha256'],21:'Codex recording edit',22:datetime.datetime(2026,10,19),
              23:asset['license_source']+' | '+asset['license_scope'],24:asset['source_path'],25:'枪种独立SoundSet；人工听感/Android待验；verify_audio_library_flow'}
        for col,value in vals.items():ws.cell(row,col,value)
    sys.path.insert(0,str(ROOT/'tools/asset_pipeline'))
    from split_asset_ledger import dedupe_key_formula,dedupe_result_formula,read_source_rows,_row_digest,CONTENT_COLUMNS,col_digest
    last=ws.max_row
    for row in range(6,last+1):
        ws.cell(row,18,dedupe_key_formula(row));ws.cell(row,19,dedupe_result_formula(row,last))
    # Preserve existing overview formulas and expand only their terminal row references.
    for s in wb:
        if s.title==ws.title:continue
        for row in s:
            for cell in row:
                if cell.data_type=='f' and '资产主表' in cell.value:
                    cell.value=re.sub(r'(\$?[A-Y]\$?)'+str(old_end)+r'\b',lambda m:m[1]+str(last),cell.value)
    from openpyxl.worksheet.cell_range import MultiCellRange
    for dv in ws.data_validations.dataValidation:
        dv.sqref=MultiCellRange(' '.join(re.sub(r'(:[A-Z]+)'+str(old_end)+r'\b',lambda m:m[1]+str(last),str(r)) for r in dv.sqref.ranges))
    ws.auto_filter.ref=f'A5:Y{last}';wb.save(ledger)
    baseline=ROOT/'assets/registry/ledger_split_baseline.json';backup(baseline)
    bl=json.loads(baseline.read_text(encoding='utf-8'))
    for _,values in read_source_rows(ws):
        if values[0] in by_id:bl['assets'][values[0]]={'v':_row_digest(values),'c':values[2],'d':'audio'}
    collected=[]
    for d in index['domains']:
        other=openpyxl.load_workbook(ROOT/index['ledger_dir']/d['file'])
        collected.extend(read_source_rows(other[index['asset_sheet']]))
    collected.sort(key=lambda x:x[0]);bl['asset_count']=len(bl['assets'])
    bl['column_digests']={str(c):col_digest(collected,c) for c in CONTENT_COLUMNS}
    bl['category_counts']=dict(Counter(v[2] for _,v in collected));save_json(baseline,bl,1)
    contract=ROOT/'art_source/audio/sfx/library_contract.json';backup(contract)
    d=json.loads(contract.read_text(encoding='utf-8'));d.update(design_revision='AUDIO-SFX r3',production_revision='gun-families-r4',weapon_profiles=catalog['weapon_profiles'],source_method='Seven gun families use recording edits (Tabasco archive attribution CC-BY-3.0 plus credited CC0 firearm/Foley); other sets retain original industrial-r2 DSP')
    for a in changes:
        spec=next((s for s in d['sounds'] if s['key']==a['key']),None)
        if spec is None:spec={'key':a['key'],'name':a['key'],'events':a['event_ids'],'variants':3 if a['key']=='launcher_fire' else 1};d['sounds'].append(spec)
        spec['production_override']={'revision':'gun-families-r4','renderer':'tools/asset_pipeline/build_weapon_audio_profiles.py','source_of_current_duration_and_version':'src/assets/audio/sfx/sfx_catalog.json'}
    save_json(contract,d)
    save_json(SESSION/'result.json',{'changed_assets':changes,'new_asset_count':len(new),'sound_sets':len(catalog['sound_sets']),'asset_count':len(catalog['assets'])})
    (ROOT/'art_source/audio/sfx/tabasco_recording_license.txt').write_text('Gunshot Sounds — Vincent Sevedge / Tabasco\nhttps://opengameart.org/content/gunshot-sounds\nArchive copyright notice: Copyright (c) 2009 Vincent Sevedge.\nArchive license: Creative Commons Attribution 3.0 Unported, https://creativecommons.org/licenses/by/3.0/\nThe web page offers CC0; ShellStorm2 retains the archive attribution and CC-BY-3.0 terms.\nEdits: transient extraction, noise reduction, filtering, fades, peak normalization; launcher/charge layered with separately credited CC0 Foley.\n',encoding='utf-8')
    print('Gun profiles installed:',len(catalog['weapon_profiles']),'IDs;',len(changes),'authored assets;',len(new),'new assets')


if __name__=='__main__':main()
