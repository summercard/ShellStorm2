"""Delete the retired equipment_steam_fx package from current assets and ledgers."""
from __future__ import annotations

import argparse
import json
import re
import shutil
from pathlib import Path

from openpyxl import load_workbook


TARGET_SLUG = "equipment_steam_fx"
TARGET_CN = "63_设备蒸汽动效组_资产包"
TARGET_ASSET_ID = "BPK-BASE99-EQUIPMENT-STEAM-FX"
DONE_STATUSES = ("已完成", "原型已接入", "正式美术已接入", "已优化并正式接入", "Blender源已完成", "已导入；优化完成")


def assert_inside(root: Path, path: Path) -> None:
    resolved = path.resolve()
    if root.resolve() not in resolved.parents:
        raise RuntimeError(f"Refusing path outside project: {resolved}")


def remove_json_targets(value):
    if isinstance(value, dict):
        if any(str(value.get(key, "")) in {TARGET_SLUG, TARGET_CN} for key in ("slug", "asset_slug", "package", "display_name")):
            return None
        cleaned = {}
        for key, child in value.items():
            if key == TARGET_SLUG:
                continue
            result = remove_json_targets(child)
            if result is not None:
                cleaned[key] = result
        if "package_count" in cleaned:
            for key in ("packages", "glbs"):
                if isinstance(cleaned.get(key), (list, dict)):
                    cleaned["package_count"] = len(cleaned[key])
                    break
        return cleaned
    if isinstance(value, list):
        return [item for child in value if (item := remove_json_targets(child)) is not None]
    return value


def remove_godot_references(root: Path) -> list[str]:
    changed = []
    aggregate_root = root / "assets/art/environments/base_facility_3d/runtime/env_base99_remaining_facilities_v021"
    for path in sorted(aggregate_root.glob("env_base99_remaining_facilities_root_top3d_v*.tscn")):
        text = path.read_text(encoding="utf-8")
        if TARGET_SLUG not in text and TARGET_CN not in text:
            continue
        lines = text.splitlines()
        removed_ext = sum(1 for line in lines if line.startswith("[ext_resource") and TARGET_SLUG in line)
        lines = [line for line in lines if TARGET_SLUG not in line and TARGET_CN not in line]
        if removed_ext:
            def decrement(match):
                return f"[gd_scene load_steps={max(1, int(match.group(1)) - removed_ext)} format=3]"
            lines = [re.sub(r"\[gd_scene load_steps=(\d+) format=3\]", decrement, line) for line in lines]
        text = "\n".join(lines).rstrip() + "\n"
        text = re.sub(
            r"(metadata/optimized_package_count = )(\d+)",
            lambda m: m.group(1) + str(max(0, int(m.group(2)) - 1)),
            text,
        )
        path.write_text(text, encoding="utf-8")
        changed.append(str(path.relative_to(root)))

    script = root / "scripts/blender/organize_base_facility_runtime_layout_hq_v007.py"
    if script.is_file():
        text = script.read_text(encoding="utf-8")
        updated = "\n".join(line for line in text.splitlines() if TARGET_CN not in line).rstrip() + "\n"
        if updated != text:
            script.write_text(updated, encoding="utf-8")
            changed.append(str(script.relative_to(root)))
    return changed


def remove_ledger_row(root: Path) -> int | None:
    path = root / "assets/registry/ShellStorm2_美术资产台账_v001.xlsx"
    workbook = load_workbook(path, data_only=False)
    main = workbook["资产主表"]
    target_row = next((row for row in range(6, main.max_row + 1) if str(main.cell(row, 1).value or "") == TARGET_ASSET_ID), None)
    if target_row is None:
        return None
    main.delete_rows(target_row, 1)
    end = main.max_row
    for row in range(6, end + 1):
        main.cell(row, 18).value = f'=LOWER(TRIM(C{row})&"|"&TRIM(D{row})&"|"&TRIM(E{row})&"|"&TRIM(F{row})&"|"&TRIM(H{row})&"|"&TRIM(I{row}))'
        main.cell(row, 19).value = f'=IF(COUNTIF($R$6:$R${end},R{row})>1,"重复","唯一")'
    overview = workbook["总览"]
    overview["A6"] = f"=COUNTA('资产主表'!$A$6:$A${end})"
    overview["C6"] = "=" + "+".join(f'COUNTIF(\'资产主表\'!$K$6:$K${end},"{status}")' for status in DONE_STATUSES)
    overview["E6"] = f'=COUNTIF(\'资产主表\'!$K$6:$K${end},"待制作")+COUNTIF(\'资产主表\'!$K$6:$K${end},"程序占位")'
    overview["G6"] = f'=COUNTIF(\'资产主表\'!$S$6:$S${end},"重复")'
    for row in range(10, 19):
        overview.cell(row, 2).value = f"=COUNTIF('资产主表'!$C$6:$C${end},A{row})"
        statuses = "+".join(f'(\'资产主表\'!$K$6:$K${end}="{status}")' for status in DONE_STATUSES)
        overview.cell(row, 3).value = f"=SUMPRODUCT(('资产主表'!$C$6:$C${end}=A{row})*({statuses}))"
    workbook.save(path)
    return target_row


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--project", type=Path, default=Path(__file__).resolve().parents[2])
    parser.add_argument("--dry-run", action="store_true")
    args = parser.parse_args()
    root = args.project.resolve()
    deleted_dirs = []
    target_dirs = [
        root / "source/art/blender/base_facility_layout/component_packages/support" / TARGET_SLUG,
        root / "assets/art/environments/base_facility_3d/components/env_base99_remaining_facilities_v021" / TARGET_SLUG,
        root / "assets/art/environments/base_facility_3d/runtime/env_base99_remaining_facilities_v021" / TARGET_SLUG,
    ]
    for directory in target_dirs:
        assert_inside(root, directory)
        if directory.exists():
            deleted_dirs.append(str(directory.relative_to(root)))
            if not args.dry_run:
                shutil.rmtree(directory)

    json_changed = []
    json_files = [
        root / "assets/art/asset_import_manifest_v001.json",
        *sorted((root / "assets/art/environments/base_facility_3d").rglob("*.json")),
        *sorted((root / "source/art/blender/base_facility_layout/export").rglob("*.json")),
    ]
    for path in json_files:
        before = path.read_text(encoding="utf-8")
        data = json.loads(before)
        cleaned = remove_json_targets(data)
        after = json.dumps(cleaned, ensure_ascii=False, indent=2) + "\n"
        if after != before:
            json_changed.append(str(path.relative_to(root)))
            if not args.dry_run:
                path.write_text(after, encoding="utf-8")

    if not args.dry_run:
        global_path = root / "assets/art/asset_import_manifest_v001.json"
        global_data = json.loads(global_path.read_text(encoding="utf-8"))
        palette = global_data.get("shared_scene_facility_palette", {})
        palette["managed_blend_sources"] = len(list((root / "source/art/blender/base_facility_layout/source").glob("base_facility_runtime_layout_hq_v*.blend")))
        palette["managed_glbs"] = len(list((root / "assets/art/environments/base_facility_3d/components").rglob("*.glb")))
        palette["source_controlled_glb_import_contracts"] = palette["managed_glbs"]
        global_data["base99_source_policy"]["current_master_source"] = "source/art/blender/base_facility_layout/source/base_facility_runtime_layout_hq_v026.blend"
        global_path.write_text(json.dumps(global_data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

    godot_changed = [] if args.dry_run else remove_godot_references(root)
    target_row = None if args.dry_run else remove_ledger_row(root)
    print("REMOVE_EQUIPMENT_STEAM_FX", {"directories": deleted_dirs, "json_files": json_changed, "godot_files": godot_changed, "excel_row": target_row, "dry_run": args.dry_run})


if __name__ == "__main__":
    main()
