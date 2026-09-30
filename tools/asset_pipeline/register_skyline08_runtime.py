"""仅登记本批SKYLINE及授权路线修改；保全无关资产内容与基线指纹。"""
import hashlib
import json
import re
import sys
from collections import Counter
from copy import copy
from datetime import date
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / 'scripts'))
sys.path.insert(0, str(Path(__file__).resolve().parent))
from ledger_registry import LedgerIndex
from split_asset_ledger import FIRST_DATA_ROW, CONTENT_COLUMNS, _row_digest, col_digest, dedupe_key_formula, dedupe_result_formula, read_source_rows, sheet_digest

BASE = ROOT / 'assets/art/environments/open_world'
OUT = ROOT / 'outputs/skyline08_import_20260930'

def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()

def write_json(path, value):
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2) + '\n', encoding='utf8', newline='\r\n')

def main():
    acceptance = json.loads((OUT / 'acceptance_render.json').read_text(encoding='utf8'))
    assert not acceptance['failures'] and acceptance['display'] != 'headless'
    assert acceptance['existing_material_resources'] == json.loads((OUT / 'before_runtime.json').read_text(encoding='utf8'))['old_material_resources']
    route = 'assets/art/environments/open_world/runtime/cross_tower_route/env_cross_tower_route_root_top3d.tscn'
    assert acceptance['route_sha256'] == sha(ROOT / route)
    manifest = json.loads((BASE / 'source/skyline_08/export/v004/export_manifest.json').read_text(encoding='utf8'))
    assert sha(ROOT / manifest['source']) == manifest['source_sha256']
    index = LedgerIndex.load(ROOT)
    domain = next(d for d in index.domains if d.key == 'scenes')
    baseline_path = ROOT / 'assets/registry/ledger_split_baseline.json'
    runtime_manifest_path = BASE / 'runtime/skyline_08/asset_manifest.json'
    runtime_manifest = json.loads(runtime_manifest_path.read_text(encoding='utf8'))
    runtime_manifest.update(export_manifest='assets/art/environments/open_world/source/skyline_08/export/v004/export_manifest.json', acceptance='outputs/skyline08_import_20260930/acceptance.json', runtime_integrated=True, runtime_status='已导入；优化完成', material_policy='无新材质缓存；实例入树直接绑定塔2既有四角色材质资源', ledger=domain.path.relative_to(ROOT).as_posix(), ledger_sheet='3D-场景通用')
    write_json(runtime_manifest_path, runtime_manifest)
    assert sha(domain.path) == sha(OUT / 'before' / domain.path.name), '账本出现并发修改，停止写入'
    assert sha(baseline_path) == sha(OUT / 'before' / baseline_path.name), '基线出现并发修改'
    wb = load_workbook(domain.path)
    ws = wb['资产主表']
    old_last = ws.max_row
    original_rows = {v[0]:v for _,v in read_source_rows(ws)}
    known = {ws.cell(r,1).value:r for r in range(FIRST_DATA_ROW,old_last+1)}
    root_prefab = 'assets/art/environments/open_world/runtime/skyline_08/env_skyline_08_root_top3d.tscn'
    records = [dict(asset_id='ENV-OPENWORLD-SKYLINE08',display_name='SKYLINE大楼 8层 精细天台',prefab=root_prefab,glb='',slug='building_root')]+manifest['records']
    changed = {'ENV-OPENWORLD-CROSS-TOWER-ROUTE'}
    for record in records:
        aid = record['asset_id']
        assert (ROOT / record['prefab']).is_file()
        row = known.get(aid)
        if row is None:
            row = ws.max_row + 1
            for c in range(1,26): ws.cell(row,c)._style = copy(ws.cell(old_last,c)._style)
            ws.row_dimensions[row].height = ws.row_dimensions[old_last].height
            values = {1:aid,2:'SKYLINE '+record['display_name'],3:'场景',4:'environment_module_3d',5:'open_world_skyline_08',6:record['slug'],7:'ENV-OPENWORLD-SKYLINE08',8:'Top3D / Godot Y-up',9:'default',10:'大地图场景景观；纯表现独立组件',12:'P1',16:'docs/v0.1/development/2026-09-30_skyline08_building_source.md',17:record['display_name'],23:'用户参考图；项目自制模型',24:record['slug']}
            for c,v in values.items(): ws.cell(row,c,v)
        else:
            ws.cell(row,6,'building_root')
            ws.cell(row,8,'Top3D / Godot Y-up')
            ws.cell(row,10,'大地图场景景观；独立稳定PackedScene')
        spec = '主体32×24m；8层×4m；完整Godot XYZ包络33.020×45.680×25.935m；252组件；四角色全部复用塔2旧材质' if aid=='ENV-OPENWORLD-SKYLINE08' else '×'.join('%.3f'%v for v in record['dimensions'])+'m (Blender XYZ)'
        values = {11:'已导入；优化完成',13:'v004',14:spec,15:record['prefab'],20:sha(ROOT/record['prefab']),21:'WorkBuddy',22:date(2026,9,30),25:'仅景观表现；无碰撞/导航/交互；源：'+manifest['source']+'；旧材质子资源直接引用，无新增材质。最终根坐标(10,-100.5,-90)；按用户指定摆位保留包络交叠，未私自改尺寸。'}
        for c,v in values.items(): ws.cell(row,c,v)
        changed.add(aid)
    row = known['ENV-OPENWORLD-CROSS-TOWER-ROUTE']
    ws.cell(row,20,sha(ROOT/route))
    ws.cell(row,22,date(2026,9,30))
    ws.cell(row,25,str(ws.cell(row,25).value)+' 2026-09-30：仅Tower2根+X20m并新增Skyline08于塔2原根-X10m；Tower3、桥及碰撞不动。')
    last = ws.max_row
    for r in range(FIRST_DATA_ROW,last+1):
        ws.cell(r,18,dedupe_key_formula(r))
        ws.cell(r,19,dedupe_result_formula(r,last))
    for line in wb['总览']:
        for cell in line:
            if cell.data_type=='f': cell.value=re.sub(r'(\$[A-Z]+\$)'+str(old_last)+r'\b',lambda m:m[1]+str(last),cell.value)
    for dv in ws.data_validations.dataValidation:
        ranges=[]
        for region in dv.sqref.ranges:
            region=copy(region)
            if region.max_row==old_last: region.max_row=last
            ranges.append(str(region))
        dv.sqref=' '.join(ranges)
    for cf in list(ws.conditional_formatting):
        rules=ws.conditional_formatting[cf]
        ranges=[]
        for region in cf.sqref.ranges:
            region=copy(region)
            if region.max_row==old_last: region.max_row=last
            ranges.append(str(region))
        del ws.conditional_formatting[' '.join(str(x) for x in cf.sqref.ranges)]
        for rule in rules: ws.conditional_formatting.add(' '.join(ranges),rule)
    ws.auto_filter.ref=f'A{FIRST_DATA_ROW-1}:Y{last}'
    log=wb['域变更日志']
    bits=str(log.cell(log.max_row,1).value).split('.');bits[-1]=str(int(bits[-1])+1)
    log.append(['.'.join(bits),date(2026,9,30),'SKYLINE v004景观运行接入','关卡场景 / 大地图景观','252组件与整楼入口；直接复用塔2四材质；塔2东20m、新景观塔2原点西10m；塔3保持。','根坐标遵照用户要求；视觉包络交叠未擅自修正；无碰撞。','WorkBuddy'])
    wb.save(domain.path)
    baseline=json.loads(baseline_path.read_text(encoding='utf8'))
    fingerprints=dict(baseline['assets'])
    actual_rows=read_source_rows(load_workbook(domain.path)['资产主表'])
    for _,values in actual_rows:
        if values[0] in changed:
            baseline['assets'][values[0]]={'v':_row_digest(values),'c':'场景','d':'scenes'}
        elif values[0] in original_rows:
            assert all(values[c-1]==original_rows[values[0]][c-1] for c in CONTENT_COLUMNS), '无关资产内容被改动'
    all_rows=[]
    for d in index.domains:
        all_rows.extend(read_source_rows(load_workbook(d.path)['资产主表']))
    baseline['asset_count']=len(baseline['assets'])
    baseline['category_counts']=dict(Counter(v[2] for _,v in all_rows))
    baseline['column_digests']={str(c):col_digest(all_rows,c) for c in CONTENT_COLUMNS}
    assert all(baseline['assets'][k]==v for k,v in fingerprints.items() if k not in changed)
    write_json(baseline_path,baseline)
    # 第二个事务只登记真实存在的PackedScene，并独立更新专表摘要。
    wb=load_workbook(domain.path)
    page=wb['3D-场景通用']
    old_page=page.max_row
    page_before={r:tuple(page.cell(r,c).value for c in range(1,17)) for r in range(1,old_page+1)}
    for record in records:
        r=page.max_row+1
        for c in range(1,17): page.cell(r,c)._style=copy(page.cell(old_page,c)._style)
        page.row_dimensions[r].height=page.row_dimensions[old_page].height
        values=[record['asset_id'],record['display_name'],record['prefab'],record['glb'] or None,manifest['source'],'仅表现',None,'无','无','无','完整模型原尺寸；米制','底部中心；Blender -Y→Godot +Z','大地图场景景观','已导入；优化完成','v004','所有表面复用塔2旧材质；不含玩法碰撞；正式布局归Godot TSCN。']
        for c,v in enumerate(values,1): page.cell(r,c).value=v
    assert all(tuple(page.cell(r,c).value for c in range(1,17))==v for r,v in page_before.items())
    page.auto_filter.ref=f'A{FIRST_DATA_ROW-1}:P{page.max_row}'
    wb.save(domain.path)
    baseline['sheet_digests']['3D-场景通用']=sheet_digest(page)
    write_json(baseline_path,baseline)
    assert sha(index.master_path)==sha(OUT/'before'/index.master_path.name), '总目录并发修改'
    master=load_workbook(index.master_path)
    for offset,d in enumerate(index.domains):
        if d.key=='scenes': master['分账本索引'].cell(FIRST_DATA_ROW+offset,7,last-FIRST_DATA_ROW+1)
    for line in master['3D Prefab总控']:
        if any(c.value=='3D-场景通用' for c in line):
            rows=[v for v in page.values if isinstance(v[0],str) and re.fullmatch(r'[A-Z0-9]+(?:-[A-Z0-9]+)+',v[0])]
            line[2].value=len(rows);line[3].value=sum(bool(v[3]) for v in rows);line[4].value=sum(bool(v[6]) for v in rows);line[5].value=sum(v[7]=='开' for v in rows)
    master.save(index.master_path)
    catalog_path=BASE/'source/skyline_08/v004/catalog.json'
    catalog=json.loads(catalog_path.read_text(encoding='utf8'))
    catalog.update(runtime_status='已导入；优化完成',runtime_integrated=True,runtime_prefab=root_prefab,category='大地图场景景观')
    lookup={r['slug']:r for r in manifest['records']}
    for package in catalog['packages']:
        record=lookup[package['slug']]
        package.update(exported=True,runtime_integrated=True,runtime_prefab=record['prefab'],runtime_glb=record['glb'],runtime_asset_id=record['asset_id'],collision_status='none_visual_only')
        p=catalog_path.parent/'component_packages'/package['category']/package['slug']/'asset_manifest.json'
        previous=json.loads(p.read_text(encoding='utf8'))
        previous.update(exported=True,runtime_integrated=True,runtime_prefab=record['prefab'],runtime_glb=record['glb'],runtime_asset_id=record['asset_id'],collision_status='none_visual_only')
        write_json(p,previous)
    write_json(catalog_path,catalog)
    write_json(OUT/'ledger_registration.json',dict(asset_id='ENV-OPENWORLD-SKYLINE08',ledger=domain.path.relative_to(ROOT).as_posix(),root_row=known['ENV-OPENWORLD-SKYLINE08'],new_component_rows=252,prefab_rows_added=253,unrelated_content_unchanged=True,unrelated_fingerprints_unchanged=True))
    print('SKYLINE_RUNTIME_LEDGER_OK',len(records),'prefabs')

if __name__=='__main__':
    main()
