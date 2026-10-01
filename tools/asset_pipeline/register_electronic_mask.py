"""Register a static Bunny face component, preserving historical ledger facts.

Separate transactions: asset main row first, then component/Prefab tables.
"""
from pathlib import Path
from copy import copy,deepcopy
from collections import Counter
from datetime import datetime
import sys,json,hashlib,shutil,re,subprocess,zipfile
import openpyxl
from openpyxl.worksheet.cell_range import MultiCellRange
ROOT=Path(__file__).resolve().parents[2]
sys.path.insert(0,str(ROOT/'scripts')); sys.path.insert(0,str(ROOT/'tools/asset_pipeline'))
from ledger_registry import LedgerIndex
from split_asset_ledger import HEADER_ROW,FIRST_DATA_ROW,read_source_rows,dedupe_key_formula,dedupe_result_formula,_row_digest,col_digest,CONTENT_COLUMNS,sheet_digest
ASSET=ROOT/'assets/art/characters/player/chr_player_capsule01_3d/variants/bunny01/components/face/electronic_mask'
ID='CHR-PLY-BUNNY01-FACE-ELECTRONIC-MASK'
PREFAB=ASSET/'runtime/chr_bunny01_electronic_mask_root.tscn'
GLB=ASSET/'components/chr_bunny01_electronic_mask_visual.glb'
BLEND=ASSET/'source/chr_bunny01_electronic_mask_model_v001.blend'
BASE=ROOT/'assets/registry/ledger_split_baseline.json'
OUT=ROOT/'outputs/electronic_mask/ledger'
def rel(p): return p.relative_to(ROOT).as_posix()
def sha(p): return hashlib.sha256(p.read_bytes()).hexdigest()
def cells(wb): return {(w.title,c.coordinate):c.value for w in wb for row in w for c in row if c.value is not None}
def gates(phase):
    for path,args in [('scripts/check_asset_registry.py',['--scope','structure','--ledger','characters']),('tools/asset_pipeline/verify_ledger_split.py',[])]:
        p=subprocess.run([sys.executable,'-X','utf8',str(ROOT/path),*args],cwd=ROOT,text=True,encoding='utf-8',capture_output=True)
        (OUT/(phase+'_'+Path(path).stem+'.log')).write_text(p.stdout+p.stderr,encoding='utf-8')
        assert p.returncode==0,p.stdout+p.stderr
def append(w,values,template):
    last=max(c.row for row in w for c in row if c.column==1 and c.value is not None)
    r=last+1
    for c,v in enumerate(values,1):
        w.cell(r,c)._style=copy(w.cell(template,c)._style); w.cell(r,c).value=v
    w.row_dimensions[r].height=w.row_dimensions[template].height
    return r
def main():
    phase=sys.argv[1]; assert phase in ['main','prefab']
    OUT.mkdir(parents=True,exist_ok=True)
    index=LedgerIndex.load(ROOT); path=index.path_for_category('角色')
    shutil.copy2(path,OUT/f'before_{phase}.xlsx'); shutil.copy2(BASE,OUT/f'before_{phase}_baseline.json')
    wb=openpyxl.load_workbook(path); baseline=json.loads(BASE.read_text('utf-8')); old=deepcopy(baseline)
    before=cells(wb); allowed=set(); report={}
    if phase=='main':
        w=wb['资产主表']; rows=read_source_rows(w)
        assert not any(v[0]==ID for _,v in rows),'Asset already exists'
        assert not any(sha(PREFAB)==str(v[19]) for _,v in rows),'Same-hash Prefab already exists'
        last=HEADER_ROW+len(rows); r=last+1
        values=[ID,'兔子主角黑色蓝光电子面具','角色','player_accessory','player_capsule01','glasses','CHR-PLY-CAPSULE01-3D-BUNNY01','3D/正面+俯视','default','Bunny01面饰槽','正式美术已接入','P1','v001','源脸部曲面复制+Y前移9mm；黑色亮面壳4mm；2×4×14蓝光方块；HeadJoint局部输出；根缩放1；无骨架/碰撞',rel(PREFAB),'src/player3d/PlayerAvatar3D.gd；docs/v0.1/16_技术施工_主页面与角色换装.md','电子面具;黑色面罩;蓝光;方块眼睛;面饰;electronic_mask',dedupe_key_formula(r),dedupe_result_formula(r,r),sha(PREFAB),'Codex',datetime(2026,10,1),'用户参考图；项目既有头部曲面派生；代码制作蓝光颗粒','—','按耳朵的角色配件分类：01_部件/眼镜/眼镜__electronic_mask。原11个模型网格、UV、材质、权重、骨架与源哈希保留；新增默认配置，可在衣柜面饰中替换/卸下，合法旧存档选择保留。源、GLB、挂点、实际保存/读档、动画跟随和Forward+渲染通过。']
        for c,v in enumerate(values,1):
            w.cell(r,c)._style=copy(w.cell(16,c)._style); w.cell(r,c).value=v; allowed.add((w.title,w.cell(r,c).coordinate))
        w.row_dimensions[r].height=w.row_dimensions[16].height
        for row in range(FIRST_DATA_ROW,r+1):
            w.cell(row,19).value=dedupe_result_formula(row,r); allowed.add((w.title,w.cell(row,19).coordinate))
        for ref in ['A6','C6','E6','G6','B10','C10','B11','C11']:
            cell=wb['总览'][ref]; assert f'${last}' in cell.value
            cell.value=cell.value.replace(f'${last}',f'${r}'); allowed.add(('总览',ref))
        for dv in w.data_validations.dataValidation:
            new=[]
            for rng in dv.sqref.ranges:
                if rng.max_row<r and rng.min_row==FIRST_DATA_ROW: rng.max_row=r
                new.append(str(rng))
            dv.sqref=MultiCellRange(' '.join(new))
        if w.auto_filter.ref: w.auto_filter.ref=f'A{HEADER_ROW}:Y{r}'
        log=wb['域变更日志']; lr=append(log,['v0.1.5','2026-10-01','新增默认电子面具','角色 / Bunny01 / glasses',f'新增{ID}；独立面饰默认佩戴、可换装和保存；原角色模型保留。','只新增一条资产，历史内容与其他资产指纹保持；查重与统计公式扩至新行。','Codex'],log.max_row)
        allowed.update((log.title,c.coordinate) for c in log[lr] if c.value is not None)
        row_values=next(v for _,v in read_source_rows(w) if v[0]==ID)
        baseline['assets'][ID]={'v':_row_digest(row_values),'c':'角色','d':'characters'}
        baseline['asset_count']=len(baseline['assets'])
        union=[]
        for domain in index.domains:
            union.extend(read_source_rows((wb if domain.key=='characters' else openpyxl.load_workbook(domain.path))['资产主表']))
        baseline['column_digests']={str(c):col_digest(union,c) for c in CONTENT_COLUMNS}
        baseline['category_counts']=dict(sorted(Counter(v[2] for _,v in union).items()))
        report['asset_main_row']=r
    else:
        w=wb['角色组件']
        r=append(w,['player_capsule01_bunny01_3d','glasses/electronic_mask','VisualRoot/BunnyRig/HeadJoint/FaceAccessorySocket/ElectronicMask','head',rel(GLB),'源0.630×0.479×0.257m；运行时继承玩家0.8表现倍率','HeadJoint局部原点；源曲面+Y前移0.009m','Blender +Y / Godot -Z；独立根缩放1',5,'继承现有head骨剪辑，无新动作','面饰槽，可独立替换/卸下','否',f'{ID}；原头部不变；与兔耳同为角色配件；默认款；112发光方块；layer2、GI禁用，无碰撞；v001。'],42)
        allowed.update((w.title,c.coordinate) for c in w[r] if c.value is not None)
        baseline['sheet_digests'][w.title]=sheet_digest(w); report['component_row']=r
        w=wb['3D-角色']; r=append(w,[ID,'兔子主角黑色蓝光电子面具',rel(PREFAB),rel(GLB),rel(BLEND),'可替换电子面饰组件；新角色默认款；实际换装/读档沿用glasses槽','src/player3d/PlayerAvatar3D.gd','无','Player3D玩法根','纯视觉配件不定义碰撞','源0.630×0.479×0.257m；独立根缩放1','HeadJoint局部原点；Blender +Y / Godot -Z','正式Bunny01玩家、衣柜面饰与3D预览','正式美术已接入','v001','复制原头部前侧曲面并前移9mm；黑色亮面+112蓝色方块；旧母版、骨架、动作不变；character_transfer_ledger.json与source_validation.json记录依赖及验收。'],5)
        allowed.update((w.title,c.coordinate) for c in w[r] if c.value is not None)
        w['A2']='玩家、NPC与角色表现/玩法根及制作母版；当前 5 个登记项。'; allowed.add((w.title,'A2'))
        baseline['sheet_digests'][w.title]=sheet_digest(w); report['prefab_row']=r
    after=cells(wb)
    changed={k for k in set(before)|set(after) if before.get(k)!=after.get(k)}
    assert changed<=allowed,changed-allowed
    assert all(baseline['assets'][k]==v for k,v in old['assets'].items()),'Historical asset fingerprint drift'
    temp=path.with_suffix('.electronic_mask_tmp.xlsx'); wb.save(temp)
    assert cells(openpyxl.load_workbook(temp))==after,'Save/reload changed workbook values'
    temp.replace(path); BASE.write_text(json.dumps(baseline,ensure_ascii=False,indent=1)+'\n',encoding='utf-8')
    report.update({'phase':phase,'changed_cells':sorted(changed),'historical_asset_fingerprints_preserved':len(old['assets']),'hash':sha(path)})
    (OUT/f'{phase}_transaction.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    with zipfile.ZipFile(path) as z:
        assert any('C6:C36' in z.read(n).decode('utf-8') for n in z.namelist() if n.startswith('xl/worksheets/') and n.endswith('.xml')),'Validation range lost'
    gates(phase); print('ELECTRONIC_MASK_LEDGER_OK',phase,report)
if __name__=='__main__': main()
