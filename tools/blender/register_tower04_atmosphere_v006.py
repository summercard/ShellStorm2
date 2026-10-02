"""Advance existing Tower04 row from v004 to accepted v006; preserve all other rows."""
import hashlib
import json
import re
import shutil
import sys
from pathlib import Path
from collections import Counter
from openpyxl import load_workbook
R=Path(__file__).resolve().parents[2]
sys.path.insert(0,str(R/'tools/asset_pipeline'))
from split_asset_ledger import FIRST_DATA_ROW,CONTENT_COLUMNS,read_source_rows,_row_digest,col_digest
F=R/'assets/art/environments/open_world/source/tower_04/v006'
cat=json.loads((F/'catalog.json').read_text(encoding='utf8'))
for name in ('source_audit','palette_validation','apocalypse_audit','visual_review','atmosphere_revision'):
    assert json.loads((F/'qa'/f'{name}.json').read_text(encoding='utf8'))['passed'],name
index_path=R/'assets/registry/ledger_index.json'; idx=json.loads(index_path.read_text(encoding='utf8'))
domain=next(d for d in idx['domains'] if d['key']=='scenes')
path=R/idx['ledger_dir']/domain['file']; base=R/'assets/registry/ledger_split_baseline.json'
digest=lambda p:hashlib.sha256(p.read_bytes()).hexdigest()
start={p:digest(p) for p in (path,base,index_path)}
backup=R.parent/'_scratch/tower04_v006_ledger_backup'; backup.mkdir(parents=True,exist_ok=True)
for p in (path,base):
    assert not (backup/p.name).exists(),'Do not replay a completed transaction'
    shutil.copy2(p,backup/p.name)
w=load_workbook(path); s=w['资产主表']; asset=cat['asset_id']
rows=[r for r in range(FIRST_DATA_ROW,s.max_row+1) if s.cell(r,1).value==asset]
assert len(rows)==1; row=rows[0]
assert s.cell(row,13).value=='v004'
assert digest(R/s.cell(row,15).value)==s.cell(row,20).value,'Prior source drift'
source_sha=digest(R/cat['source_blend'])
allowed={13:'v006',14:'150×50m；五层×5m；天台25m；1109独立包/五区域；平台末级XY≤10m；54营地设施/12藤蔓；原4材质',15:cat['source_blend'],16:'用户六张末世平台参考图；docs/v0.1/design/tower04_mall_source.md r6',17:'塔4;末世平台;五区模块;帆布营地;笼式水箱;光伏板;局部藤蔓;近景磨损',20:source_sha,22:'2026-10-02',25:'源v006：继承v005完整几何，补12度夕阳、Nishita天空与展示合成薄霭；保留平台/跑道/草坪轮廓和下部建筑。平台长件按10m网格拆分，694地板保留附着磨损；6营地、54独立设施、12藤蔓挂点。原四材质/公共贴图/import/MipMap均未修改。保存重开、逐面UV、固定镜头及独立暖光展示验收；v001–v005保留。仅Blender源，未导出GLB/PackedScene，碰撞/LOD/导航未制作。'}
assert cat['package_count']==1109
before={(r,c):(s.cell(r,c).value,s.cell(r,c).style_id) for r in range(1,s.max_row+1) for c in range(1,26)}
other={sheet.title:[tuple(c.value for c in cells) for cells in sheet] for sheet in w if sheet.title not in ('资产主表','域变更日志')}
for c,value in allowed.items(): s.cell(row,c,value)
log=w['域变更日志']; m=re.fullmatch(r'v(\d+)\.(\d+)\.(\d+)',str(log.cell(log.max_row,1).value)); assert m
log.append([f'v{m[1]}.{m[2]}.{int(m[3])+1}','2026-10-02','塔4末世平台与黄昏展示收尾','关卡场景 / 开放世界',asset+'：平台模块化与54件营地设施、12处局部藤蔓；v006仅补独立黄昏天空、低角度暖光和距离薄霭。','源v006；原四材质与MipMap不变；未导入运行时。','Codex'])
assert all(digest(p)==h for p,h in start.items()),'Concurrent ledger write; no overwrite'
w.save(path)
again=load_workbook(path); ss=again['资产主表']; diffs=[]
for (r,c),(v,style) in before.items():
    assert ss.cell(r,c).style_id==style
    if ss.cell(r,c).value!=v:
        assert r==row and c in allowed,('Unrelated cell changed',r,c)
        diffs.append((r,c))
assert set(c for _,c in diffs)==set(allowed)
assert all([tuple(c.value for c in cells) for cells in again[name]]==values for name,values in other.items())
bl=json.loads(base.read_text(encoding='utf8')); prior=dict(bl['assets'])
bl['assets'][asset]={'v':_row_digest(dict(read_source_rows(ss))[row]),'c':'场景','d':'scenes'}
all_rows=[]
for d in idx['domains']: all_rows.extend(read_source_rows(load_workbook(R/idx['ledger_dir']/d['file'])['资产主表']))
bl['column_digests']={str(c):col_digest(all_rows,c) for c in CONTENT_COLUMNS}
bl['category_counts']=dict(Counter(v[2] for _,v in all_rows))
assert all(bl['assets'][k]==v for k,v in prior.items() if k!=asset)
assert digest(base)==start[base] and digest(index_path)==start[index_path],'Concurrent baseline change'
base.write_text(json.dumps(bl,ensure_ascii=False,indent=1)+'\n',encoding='utf8')
report=dict(asset_id=asset,row=row,source_version='v006',source_sha256=source_sha,changed_cells=[ss.cell(r,c).coordinate for r,c in diffs],other_rows_and_formulas_unchanged=True,styles_unchanged=True,unrelated_sheets_unchanged=True,other_baseline_assets_unchanged=True,asset_count_unchanged=bl['asset_count']==len(prior),prefab_pages_not_modified=True,backup=str(backup))
(F/'qa/ledger_registration.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf8')
print(json.dumps(report,ensure_ascii=False,indent=2))
