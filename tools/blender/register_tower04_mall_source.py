"""Append Tower04 source only, preserving all existing ledger asset contents.

Project's ledger-specific authoring contract requires openpyxl, not re-export
through a spreadsheet editor which can rewrite its frozen formula contracts.
"""
from collections import Counter
from copy import copy
import hashlib
import json
from pathlib import Path
import re
import shutil
import sys
from xml.etree import ElementTree as ET
from zipfile import ZipFile
from openpyxl import load_workbook
from openpyxl.worksheet.datavalidation import DataValidation

R=Path(__file__).resolve().parents[2]
sys.path.insert(0,str(R/'tools/asset_pipeline'))
from split_asset_ledger import FIRST_DATA_ROW, CONTENT_COLUMNS, read_source_rows, _row_digest, col_digest, dedupe_key_formula, dedupe_result_formula

asset='ENV-OPENWORLD-TOWER04'
folder=R/'assets/art/environments/open_world/source/tower_04/v001'
catalog=json.loads((folder/'catalog.json').read_text(encoding='utf8'))
blend=R/catalog['source_blend']
audit=json.loads((folder/'qa/source_audit.json').read_text(encoding='utf8'))
uv=json.loads((folder/'qa/palette_validation.json').read_text(encoding='utf8'))
assert audit['passed'] and uv['passed'] and audit['original_shader_signatures_match']
idx=json.loads((R/'assets/registry/ledger_index.json').read_text(encoding='utf8'))
domain=next(d for d in idx['domains'] if d['key']=='scenes')
path=R/idx['ledger_dir']/domain['file']
baseline=R/'assets/registry/ledger_split_baseline.json'
backup=R.parent/'_scratch/tower04_ledger_backup'
backup.mkdir(parents=True,exist_ok=True)
w=load_workbook(path); sheet=w[idx['asset_sheet']]; old=sheet.max_row
if '--refresh-draft' in sys.argv:
    prior=json.loads((folder/'qa/ledger_registration.json').read_text(encoding='utf8'))
    row=prior['row']; assert sheet.cell(row,1).value==asset
    assert sheet.cell(row,20).value==prior['source_sha256'], 'Concurrent ledger update'
    content={(r,c):sheet.cell(r,c).value for r in range(1,old+1) for c in CONTENT_COLUMNS if (r,c)!=(row,20)}
    sha=hashlib.sha256(blend.read_bytes()).hexdigest(); sheet.cell(row,20,sha); w.save(path)
    ss=load_workbook(path)['资产主表']
    assert all(ss.cell(r,c).value==v for (r,c),v in content.items()), 'Unexpected ledger mutation'
    bl=json.loads(baseline.read_text(encoding='utf8'))
    bl['assets'][asset]={'v':_row_digest(dict(read_source_rows(ss))[row]),'c':'场景','d':'scenes'}
    rows=[]
    for d in idx['domains']:
        rows.extend(read_source_rows(load_workbook(R/idx['ledger_dir']/d['file'])['资产主表']))
    bl['column_digests']={str(c):col_digest(rows,c) for c in CONTENT_COLUMNS}
    baseline.write_text(json.dumps(bl,ensure_ascii=False,indent=1)+'\n',encoding='utf8')
    prior['source_sha256']=sha
    (folder/'qa/ledger_registration.json').write_text(json.dumps(prior,ensure_ascii=False,indent=2),encoding='utf8')
    print('TOWER04_DRAFT_SHA_REFRESHED',sha)
    raise SystemExit(0)
assert not any(sheet.cell(r,1).value==asset for r in range(FIRST_DATA_ROW,old+1)), 'Already registered; do not append'
for p in (path,baseline):
    if (backup/p.name).exists():
        assert hashlib.sha256((backup/p.name).read_bytes()).digest()==hashlib.sha256(p.read_bytes()).digest(), 'Previous write or concurrent edit; inspect transaction'
    else: shutil.copy2(p,backup/p.name)
original={(r,c):sheet.cell(r,c).value for r in range(1,old+1) for c in CONTENT_COLUMNS}
other_sheets={s.title:[[c.value for c in row] for row in s] for s in w if s.title not in ('资产主表','总览','域变更日志')}
templates=[r for r in range(FIRST_DATA_ROW,old+1) if sheet.cell(r,1).value=='ENV-OPENWORLD-TOWER03']
template=templates[-1]
row=old+1
sha=hashlib.sha256(blend.read_bytes()).hexdigest()
values=[asset,'塔4 曲线天台大型商场','场景','environment_kit_3d','open_world_tower_04','building_source',None,'Top3D / Blender Z-up','default','独立开放世界商场美术源；不替换现有建筑','Blender源已完成','P1','v001','150×50m包络；最高主体5层×5m=25m；错层露台15m；112独立包，122输出网格；4原材质；6张真实渲染预览',catalog['source_blend'],'用户参考图及150×50m约五层要求；docs/v0.1/design/tower04_mall_source.md','塔4;商场;圆弧左翼;波浪悬挑;格构遮阳;曲线花园;回旋楼梯;玻璃亭',None,None,sha,'Codex','2026-09-30','用户提供参考图，据图建模；未下载第三方模型','TOWER04','仅Blender源已完成，未导出GLB、未建立PackedScene、未接主场景，未制作碰撞/导航/LOD。五层为左翼最高主体；长翼三层，露台15m；按参考保持错层体块。复用塔楼03原四材质，原节点签名一致，公共贴图/导入设置未改。']
for c,value in enumerate(values,1):
    sheet.cell(row,c,value); sheet.cell(row,c)._style=copy(sheet.cell(template,c)._style)
sheet.row_dimensions[row].height=sheet.row_dimensions[template].height
sheet.cell(row,18,dedupe_key_formula(row))
for r in range(FIRST_DATA_ROW,row+1): sheet.cell(r,19,dedupe_result_formula(r,row))
for cells in w['总览']:
    for cell in cells:
        if cell.data_type=='f': cell.value=re.sub(r'(\$[A-Z]+\$)'+str(old)+r'\b',lambda m:m[1]+str(row),cell.value)
for validation in list(sheet.data_validations.dataValidation):
    ranges=[]
    for rng in validation.sqref.ranges:
        expanded=copy(rng)
        if expanded.max_row==old: expanded.max_row=row
        ranges.append(str(expanded))
    validation.sqref=' '.join(ranges)
sheet.auto_filter.ref=f'A{FIRST_DATA_ROW-1}:Y{row}'
log=w['域变更日志']; match=re.fullmatch(r'v(\d+)\.(\d+)\.(\d+)',str(log.cell(log.max_row,1).value))
assert match, 'Unexpected domain version'
version=f'v{match[1]}.{match[2]}.{int(match[3])+1}'
log.append([version,'2026-09-30','新增Blender建筑源','关卡场景 / 开放世界',asset+'：150×50m曲线商场，最高主体五层，112独立包，原四材质复用。','源v001完成；未导入运行时。','Codex'])
w.save(path)
again=load_workbook(path); ss=again['资产主表']
assert all(ss.cell(r,c).value==v for (r,c),v in original.items()), 'Existing asset content changed'
assert all([[c.value for c in cells] for cells in again[s]]==v for s,v in other_sheets.items()), 'Unrelated sheet changed'
assert all(ss.cell(r,18).value==dedupe_key_formula(r) and ss.cell(r,19).value==dedupe_result_formula(r,row) for r in range(FIRST_DATA_ROW,row+1))
with ZipFile(path) as z:
    ns={'s':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
    validations=[v.attrib['sqref'] for name in z.namelist() if re.match(r'xl/worksheets/sheet\d+.xml',name) for v in ET.fromstring(z.read(name)).findall('.//s:dataValidation',ns)]
assert any(f'C{row}' in v for v in validations) and any(f'K{row}' in v for v in validations) and any(f'L{row}' in v for v in validations)
bl=json.loads(baseline.read_text(encoding='utf8')); old_assets=dict(bl['assets'])
assert asset not in old_assets
bl['assets'][asset]={'v':_row_digest(dict(read_source_rows(ss))[row]),'c':'场景','d':'scenes'}
bl['asset_count']=len(bl['assets'])
all_rows=[]
for d in idx['domains']:
    wb=load_workbook(R/idx['ledger_dir']/d['file']); all_rows.extend(read_source_rows(wb['资产主表']))
bl['column_digests']={str(c):col_digest(all_rows,c) for c in CONTENT_COLUMNS}
bl['category_counts']=dict(Counter(values[2] for _,values in all_rows))
assert all(bl['assets'][k]==v for k,v in old_assets.items()), 'Existing baseline asset changed'
baseline.write_text(json.dumps(bl,ensure_ascii=False,indent=1)+'\n',encoding='utf8')
report=dict(asset_id=asset,row=row,ledger=path.relative_to(R).as_posix(),source_sha256=sha,source_version='v001',status='Blender源已完成',prior_asset_content_unchanged=True,prior_baseline_fingerprints_unchanged=True,prior_unrelated_sheets_unchanged=True,prefab_page_not_modified=True,baseline_count=bl['asset_count'],data_validations_verified=True)
(folder/'qa/ledger_registration.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf8')
print(json.dumps(report,ensure_ascii=False,indent=2))
