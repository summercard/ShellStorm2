"""Promote the validated v005 Blender source in the existing scenes ledger row."""

import hashlib
import json
import re
import shutil
import sys
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "tools/asset_pipeline"))
from split_asset_ledger import (  # noqa: E402
    CONTENT_COLUMNS,
    FIRST_DATA_ROW,
    _row_digest,
    col_digest,
    read_source_rows,
)

ASSET = "ENV-OPENWORLD-LOGISTICS07"
FOLDER = ROOT / "assets/art/environments/open_world/source/logistics_07/v005"
BLEND = FOLDER / "env_logistics_07_source_v005.blend"
AUDIT = json.loads((FOLDER / "qa/optimization_v005.json").read_text(encoding="utf8"))
SOURCE_AUDIT = json.loads((FOLDER / "qa/source_audit.json").read_text(encoding="utf8"))
PALETTE_AUDIT = json.loads((FOLDER / "qa/palette_validation.json").read_text(encoding="utf8"))
assert AUDIT["passed"] and SOURCE_AUDIT["passed"] and PALETTE_AUDIT["passed"]
assert AUDIT["optimized_hvac_count"] == 11
assert all(item["triangle_reduction"] >= 0.5 for item in AUDIT["optimized_hvac"])
new_sha = hashlib.sha256(BLEND.read_bytes()).hexdigest()
assert new_sha == AUDIT["source_after_sha256"]

index = json.loads((ROOT / "assets/registry/ledger_index.json").read_text(encoding="utf8"))
domain = next(item for item in index["domains"] if item["key"] == "scenes")
ledger = ROOT / index["ledger_dir"] / domain["file"]
baseline = ROOT / "assets/registry/ledger_split_baseline.json"
backup = ROOT / "_scratch/logistics07_v005_ledger_backup"
backup.mkdir(parents=True, exist_ok=True)
for path in (ledger, baseline):
    target = backup / path.name
    if not target.exists():
        shutil.copy2(path, target)

workbook = load_workbook(ledger)
sheet = workbook["资产主表"]
prior_workbook = load_workbook(backup / ledger.name)
rows_before = dict(read_source_rows(prior_workbook["资产主表"]))
matches = [row for row, values in rows_before.items() if values[0] == ASSET]
assert matches == [824], matches
row = matches[0]
assert rows_before[row][12] == "v004"
assert rows_before[row][14].endswith("/v004/env_logistics_07_source_v004.blend")
assert rows_before[row][19] == AUDIT["source_before_sha256"]

needs_update = sheet.cell(row, 13).value == "v004"
if needs_update:
    sheet.cell(row, 13, "v005")
    sheet.cell(row, 14, "主体28×20×18m；334独立包；4角色公共色盘；6张Cycles实图；11台空调减面59%–61%，整栋87,939三角面")
    sheet.cell(row, 15, BLEND.relative_to(ROOT).as_posix())
    sheet.cell(row, 20, new_sha)
    sheet.cell(row, 25, "Blender源v005已完成；11台空调三角面至少减半，其余323包几何不变。未导出GLB/接入Godot，未制作碰撞/LOD。独立外观美术尺度，不作为主塔12m战斗层或可进入物流关卡契约。")
else:
    assert sheet.cell(row, 13).value == "v005"
    assert sheet.cell(row, 20).value == new_sha

log = workbook["域变更日志"]
previous_version = str(log.cell(log.max_row, 1).value)
match = re.fullmatch(r"v(\d+)\.(\d+)\.(\d+)", previous_version)
assert match, previous_version
if needs_update:
    new_log_version = f"v{match[1]}.{match[2]}.{int(match[3]) + 1}"
    log.append([
        new_log_version,
        "2026-10-03",
        "优化07号物流建筑空调源",
        "关卡场景 / 开放世界",
        f"{ASSET}：源v005；11台空调减面59%–61%，整栋116,011→87,939三角面。",
        "Blender源完成；未导入运行时。",
        "Codex",
    ])
    workbook.save(ledger)
else:
    new_log_version = previous_version
    assert "v005" in str(log.cell(log.max_row, 5).value)

saved = load_workbook(ledger)
rows_after = dict(read_source_rows(saved["资产主表"]))
assert set(rows_after) == set(rows_before)
assert all(rows_after[number] == values for number, values in rows_before.items() if number != row)
assert rows_after[row][12] == "v005"
assert rows_after[row][19] == new_sha

data = json.loads(baseline.read_text(encoding="utf8"))
old_assets = dict(data["assets"])
assert ASSET in old_assets
data["assets"][ASSET]["v"] = _row_digest(rows_after[row])
assert all(data["assets"][asset] == value for asset, value in old_assets.items() if asset != ASSET)
all_rows = []
for item in index["domains"]:
    book = load_workbook(ROOT / index["ledger_dir"] / item["file"])
    all_rows.extend(read_source_rows(book["资产主表"]))
    book.close()
data["column_digests"] = {str(column): col_digest(all_rows, column) for column in CONTENT_COLUMNS}
data["category_counts"] = dict(Counter(values[2] for _, values in all_rows))
assert data["asset_count"] == len(data["assets"])
baseline.write_text(json.dumps(data, ensure_ascii=False, indent=1) + "\n", encoding="utf8")

report = {
    "asset_id": ASSET,
    "row": row,
    "version": "v005",
    "source_sha256": new_sha,
    "previous_version": "v004",
    "previous_sha256": AUDIT["source_before_sha256"],
    "other_asset_rows_unchanged": True,
    "other_asset_fingerprints_unchanged": True,
    "backup_directory": str(backup),
    "log_version": new_log_version,
}
(FOLDER / "qa/ledger_registration.json").write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf8")
print("LOGISTICS07_V005_LEDGER=" + json.dumps(report, ensure_ascii=False))
