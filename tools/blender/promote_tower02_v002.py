"""Update exactly the existing Tower02 ledger row after final source checks."""
from pathlib import Path
import json,sys,hashlib,shutil
from collections import Counter
from openpyxl import load_workbook
R=Path(__file__).resolve().parents[2]; sys.path.insert(0,str(R/'tools/asset_pipeline'))
from split_asset_ledger import read_source_rows,_row_digest,col_digest,CONTENT_COLUMNS
idx=json.loads((R/'assets/registry/ledger_index.json').read_text(encoding='utf8')); domain=next(d for d in idx['domains'] if d['key']=='scenes')
path=R/idx['ledger_dir']/domain['file']; base=R/'assets/registry/ledger_split_baseline.json'
folder=R/'assets/art/environments/open_world/source/tower_02/v002'; blend=folder/'塔2_施工高楼_70x50m_v002.blend'
for name in ['scope_lock.json','source_audit.json','palette_validation.json']: assert json.loads((folder/'qa'/name).read_text(encoding='utf8'))['passed'],name
backup=R/'_scratch/tower02_v002_ledger_backup'; backup.mkdir(exist_ok=True)
for p in [path,base]:
 if not (backup/p.name).exists(): shutil.copy2(p,backup/p.name)
w=load_workbook(path); s=w['资产主表']; entries=read_source_rows(s); asset='ENV-OPENWORLD-TOWER02'; r=next(r for r,v in entries if v[0]==asset)
assert s.cell(r,13).value=='v001'
updates={13:'v002',14:'70×50m；92独立包；立柱3.5–12.2m高低错落；多层模板平台/梁笼顶撑；三台塔吊机械深化；4材质公共色盘',15:blend.relative_to(R).as_posix(),16:'用户两张深化参考图；docs/v0.1/development/2026-09-29_openworld_tower02_v002.md',20:hashlib.sha256(blend.read_bytes()).hexdigest(),22:'2026-09-29',25:'v002仅深化屋顶及三台塔吊；下部楼体、灰色顶板、参考镜头/灯光锁定。v001保留。源已完成；未导出或接入Godot。'}
for c,v in updates.items(): s.cell(r,c,v)
before=dict(entries); after=dict(read_source_rows(s)); differences=[(row,c+1) for row,vals in before.items() for c,v in enumerate(vals) if v!=after[row][c]]
assert set(differences)=={(r,c) for c in updates},differences
log=w['域变更日志']; bits=str(log.cell(log.max_row,1).value).split('.'); bits[-1]=str(int(bits[-1])+1)
log.append(['.'.join(bits),'2026-09-29','塔2屋顶与塔吊深化','关卡场景 / 开放世界','ENV-OPENWORLD-TOWER02 升v002，92包；错层钢筋笼、多层作业平台及塔吊机械深化。','原v001保留；下部楼体锁区检查通过；仅Blender源。','Codex']); w.save(path)
bl=json.loads(base.read_text(encoding='utf8')); bl['assets'][asset]['v']=_row_digest(after[r]); rows=[]
for d in idx['domains']: rows.extend(read_source_rows(load_workbook(R/idx['ledger_dir']/d['file'])['资产主表']))
bl['column_digests']={str(c):col_digest(rows,c) for c in CONTENT_COLUMNS}; bl['category_counts']=dict(Counter(v[2] for _,v in rows)); base.write_text(json.dumps(bl,ensure_ascii=False,indent=1)+'\n',encoding='utf8')
(folder/'qa/ledger_changes.json').write_text(json.dumps(dict(asset_id=asset,row=r,changed_columns=sorted(updates),other_asset_rows_unchanged=True),indent=2),encoding='utf8'); print('TOWER02_V002_REGISTERED',r)
