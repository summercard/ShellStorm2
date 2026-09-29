"""Narrow scene-ledger append using the project's lossless ledger convention."""
from pathlib import Path
import sys,json,hashlib,shutil,re
from copy import copy
from collections import Counter
from openpyxl import load_workbook
R=Path(__file__).resolve().parents[2]; sys.path.insert(0,str(R/'tools/asset_pipeline'))
from split_asset_ledger import FIRST_DATA_ROW,read_source_rows,_row_digest,col_digest,CONTENT_COLUMNS,dedupe_key_formula,dedupe_result_formula
idx=json.loads((R/'assets/registry/ledger_index.json').read_text(encoding='utf8')); domain=next(d for d in idx['domains'] if d['key']=='scenes')
path=R/idx['ledger_dir']/domain['file']; baseline=R/'assets/registry/ledger_split_baseline.json'
asset='ENV-OPENWORLD-TOWER02'; folder=R/'assets/art/environments/open_world/source/tower_02/v001'; blend=folder/'塔2_施工高楼_70x50m_v001.blend'
backup=R/'_scratch/tower02_ledger_backup'; backup.mkdir(exist_ok=True)
for p in [path,baseline]:
 if not (backup/p.name).exists(): shutil.copy2(p,backup/p.name)
w=load_workbook(path); s=w['资产主表']; old=s.max_row
assert not any(s.cell(r,1).value==asset for r in range(FIRST_DATA_ROW,old+1))
new=old+1
values=[asset,'开放世界 塔2 施工高楼 70×50m','场景','environment_kit_3d','open_world_tower_02','building_source',None,'Top3D / Blender Z-up','construction','开放世界独立建筑，与主塔平级','Blender源已完成','P1','v001','主体70×50m；20层，层高4.2m为参考图推定；62独立包；4角色公共色盘；灰色屋顶',blend.relative_to(R).as_posix(),'用户参考图；docs/v0.1/development/2026-09-29_openworld_tower02_source.md','塔2;施工高楼;塔吊;脚手架;灰色屋顶',None,None,hashlib.sha256(blend.read_bytes()).hexdigest(),'Codex','2026-09-29','用户提供参考图，仅据图建模','塔2','仅Blender原始文件；未导出GLB、未接入Godot、无碰撞/LOD。高度为美术推定，不继承主塔玩法层高。']
for c,value in enumerate(values,1): s.cell(new,c,value); s.cell(new,c)._style=copy(s.cell(old,c)._style)
s.row_dimensions[new].height=s.row_dimensions[old].height
s.cell(new,18,dedupe_key_formula(new))
for r in range(FIRST_DATA_ROW,new+1): s.cell(r,19,dedupe_result_formula(r,new))
for row in w['总览']:
 for cell in row:
  if cell.data_type=='f': cell.value=re.sub(r'(\$[A-Z]+\$)'+str(old)+r'\b',lambda m:m[1]+str(new),cell.value)
for dv in s.data_validations.dataValidation:
 ranges=[]
 for rng in dv.sqref.ranges:
  q=copy(rng)
  if q.max_row==old: q.max_row=new
  ranges.append(str(q))
 dv.sqref=' '.join(ranges)
s.auto_filter.ref=f'A{FIRST_DATA_ROW-1}:Y{new}'
log=w['域变更日志']; version=log.cell(log.max_row,1).value; bits=version.split('.'); bits[-1]=str(int(bits[-1])+1)
log.append(['.'.join(bits),'2026-09-29','新增Blender建筑源','关卡场景 / 开放世界',asset+' 塔2：70×50m，62包，灰色顶面。','仅源资产完成；运行时未接入。','Codex'])
w.save(path)
bl=json.loads(baseline.read_text(encoding='utf8')); current=dict(read_source_rows(s))[new]; bl['assets'][asset]={'v':_row_digest(current),'c':'场景','d':'scenes'}; bl['asset_count']=len(bl['assets'])
rows=[]
for d in idx['domains']:
 book=load_workbook(R/idx['ledger_dir']/d['file']); rows.extend(read_source_rows(book['资产主表']))
bl['column_digests']={str(c):col_digest(rows,c) for c in CONTENT_COLUMNS}; bl['category_counts']=dict(Counter(v[2] for _,v in rows)); baseline.write_text(json.dumps(bl,ensure_ascii=False,indent=1)+'\n',encoding='utf8')
catalog=json.loads((folder/'catalog.json').read_text(encoding='utf8')); catalog['ledger_status']='Blender源已完成'; catalog['asset_id']=asset; catalog['asset_ledger']=path.relative_to(R).as_posix(); (folder/'catalog.json').write_text(json.dumps(catalog,ensure_ascii=False,indent=2),encoding='utf8')
print('TOWER02_LEDGER_REGISTERED',new,asset)
