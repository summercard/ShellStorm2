"""Promote only the existing Tower04 source row to v002, with frozen-row checks."""
from collections import Counter
from pathlib import Path
import hashlib
import json
import re
import shutil
import sys
from openpyxl import load_workbook

R=Path(__file__).resolve().parents[2]
sys.path.insert(0,str(R/'tools/asset_pipeline'))
from split_asset_ledger import FIRST_DATA_ROW, CONTENT_COLUMNS, read_source_rows, _row_digest, col_digest

asset='ENV-OPENWORLD-TOWER04'
folder=R/'assets/art/environments/open_world/source/tower_04/v002'
cat=json.loads((folder/'catalog.json').read_text(encoding='utf8'))
for name in ('source_audit','palette_validation','reference_plan_audit'):
    assert json.loads((folder/'qa'/f'{name}.json').read_text(encoding='utf8'))['passed'], name
idx=json.loads((R/'assets/registry/ledger_index.json').read_text(encoding='utf8'))
d=next(d for d in idx['domains'] if d['key']=='scenes')
path=R/idx['ledger_dir']/d['file']; base=R/'assets/registry/ledger_split_baseline.json'
refresh='--refresh-final-source' in sys.argv
backup=R.parent/('_scratch/tower04_v002_ledger_sha_backup' if refresh else '_scratch/tower04_v002_ledger_backup'); backup.mkdir(parents=True,exist_ok=True)
for p in (path,base):
    assert not (backup/p.name).exists(), 'Transaction already begun; inspect before replay'
    shutil.copy2(p,backup/p.name)
start_hash=hashlib.sha256(path.read_bytes()).hexdigest()
start_base=hashlib.sha256(base.read_bytes()).hexdigest()
w=load_workbook(path); s=w['资产主表']
rows=[r for r in range(FIRST_DATA_ROW,s.max_row+1) if s.cell(r,1).value==asset]
assert len(rows)==1; row=rows[0]
if refresh:
    prior_report=json.loads((folder/'qa/ledger_registration.json').read_text(encoding='utf8'))
    assert s.cell(row,13).value=='v002' and s.cell(row,15).value==cat['source_blend']
    assert s.cell(row,20).value==prior_report['source_sha256'], 'Concurrent source hash edit'
else:
    assert s.cell(row,13).value=='v001', 'Do not overwrite concurrent source revision'
    old_source=R/s.cell(row,15).value
    assert hashlib.sha256(old_source.read_bytes()).hexdigest()==s.cell(row,20).value
source=R/cat['source_blend']; sha=hashlib.sha256(source.read_bytes()).hexdigest()
allowed={13:'v002',14:f"150×50m；五层×5m；连续天台25m；叶形棚净高4.8m；栏杆1.2m、花池0.6m；{cat['package_count']}独立包；4原材质；图纸校准与近景深化",
15:cat['source_blend'],16:'用户最新详细平面图及功能分区图；docs/v0.1/design/tower04_mall_source.md r2',
17:'塔4;商场;回旋屋顶;月牙凹口;弯曲叶棚;Y形柱;异形花池;斜置玻璃亭;草坪庭院',20:sha,22:'2026-10-01',
25:'Blender源v002；统一25m连续天台；斜置玻璃亭、四组异形绿岛、草坪雕塑，深化柱脚、采光格构、栏杆夹件及风扇设备。未导出GLB/PackedScene、碰撞/LOD/导航未制作。原四材质及公共贴图不变，MipMap未启用；v001保留。图纸带透视且东侧裁切，不宣称CAD测绘精度。'}
if refresh: allowed={20:sha}
before={(r,c):(s.cell(r,c).value,s.cell(r,c).style_id) for r in range(1,s.max_row+1) for c in range(1,26)}
other={sheet.title:[tuple(c.value for c in cells) for cells in sheet] for sheet in w if sheet.title not in ('资产主表','域变更日志')}
for c,value in allowed.items(): s.cell(row,c,value)
log=w['域变更日志']; m=re.fullmatch(r'v(\d+)\.(\d+)\.(\d+)',str(log.cell(log.max_row,1).value)); assert m
if not refresh: log.append([f'v{m[1]}.{m[2]}.{int(m[3])+1}','2026-10-01','塔4源版本升级','关卡场景 / 开放世界',asset+'：按详细平面重建月牙凹口、曲线顶棚、异形花池、斜置玻璃亭与庭院，深化近景设备。','v002源完成；原四材质不变；未导入运行时。','Codex'])
assert hashlib.sha256(path.read_bytes()).hexdigest()==start_hash and hashlib.sha256(base.read_bytes()).hexdigest()==start_base, 'Concurrent ledger edit'
w.save(path)
again=load_workbook(path); ss=again['资产主表']; diffs=[]
for (r,c),(value,style) in before.items():
    assert ss.cell(r,c).style_id==style, ('style drift',r,c)
    if ss.cell(r,c).value!=value: diffs.append((r,c)); assert r==row and c in allowed, ('unexpected edit',r,c)
assert set(c for r,c in diffs)==set(allowed)
assert all([tuple(c.value for c in cells) for cells in again[name]]==values for name,values in other.items())
bl=json.loads(base.read_text(encoding='utf8')); prior=dict(bl['assets'])
bl['assets'][asset]={'v':_row_digest(dict(read_source_rows(ss))[row]),'c':'场景','d':'scenes'}
all_rows=[]
for domain in idx['domains']:
    all_rows.extend(read_source_rows(load_workbook(R/idx['ledger_dir']/domain['file'])['资产主表']))
bl['column_digests']={str(c):col_digest(all_rows,c) for c in CONTENT_COLUMNS}
bl['category_counts']=dict(Counter(v[2] for _,v in all_rows))
assert all(bl['assets'][k]==v for k,v in prior.items() if k!=asset)
assert hashlib.sha256(base.read_bytes()).hexdigest()==start_base
base.write_text(json.dumps(bl,ensure_ascii=False,indent=1)+'\n',encoding='utf8')
report=dict(asset_id=asset,row=row,source_version='v002',source_sha256=sha,changed_cells=[f'{ss.cell(r,c).coordinate}' for r,c in diffs],other_rows_and_formulas_unchanged=True,styles_unchanged=True,unrelated_sheets_unchanged=True,other_baseline_assets_unchanged=True,asset_count_unchanged=bl['asset_count']==len(prior),prefab_pages_not_modified=True,backup=str(backup))
if refresh:
    report['initial_upgrade_changed_cells']=prior_report['changed_cells']
    report['initial_backup']=prior_report['backup']
    report['final_walkway_fix_sha_only']=True
(folder/'qa/ledger_registration.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf8')
print(json.dumps(report,ensure_ascii=False,indent=2))
