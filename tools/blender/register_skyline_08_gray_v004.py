"""Update only the existing SKYLINE row for the user-directed palette revision."""
import sys,json,hashlib,shutil
from pathlib import Path
from collections import Counter
from openpyxl import load_workbook
R=Path(__file__).resolve().parents[2]; sys.path.insert(0,str(R/'tools/asset_pipeline'))
from split_asset_ledger import FIRST_DATA_ROW,read_source_rows,_row_digest,col_digest,CONTENT_COLUMNS
idx=json.loads((R/'assets/registry/ledger_index.json').read_text(encoding='utf8')); domain=next(d for d in idx['domains'] if d['key']=='scenes')
path=R/idx['ledger_dir']/domain['file']; baseline=R/'assets/registry/ledger_split_baseline.json'
asset='ENV-OPENWORLD-SKYLINE08'; folder=R/'assets/art/environments/open_world/source/skyline_08/v004'; blend=folder/'SKYLINE大楼_8层_精细天台_v004.blend'
for filename in ['source_audit.json','palette_validation.json','white_darkening.json']:
 assert json.loads((folder/'qa'/filename).read_text(encoding='utf8'))['passed']
backup=R/'_scratch/skyline08_gray_v004_ledger_backup'; backup.mkdir(exist_ok=True)
for p in [path,baseline]:
 if not (backup/p.name).exists(): shutil.copy2(p,backup/p.name)
w=load_workbook(path); s=w['资产主表']; row=next(r for r in range(FIRST_DATA_ROW,s.max_row+1) if s.cell(r,1).value==asset)
assert s.cell(row,13).value=='v003'
before={(r,c):s.cell(r,c).value for r in range(1,s.max_row+1) for c in CONTENT_COLUMNS if r!=row}
s.cell(row,13,'v004'); s.cell(row,15,blend.relative_to(R).as_posix()); s.cell(row,20,hashlib.sha256(blend.read_bytes()).hexdigest())
s.cell(row,14,str(s.cell(row,14).value)+'；最浅白灰色按用户要求下调两档，几何与灯光不变')
s.cell(row,25,str(s.cell(row,25).value)+' v004：R10C10→R8C10、R9C10→R7C10；0面继续使用最浅两档。')
log=w['域变更日志']; bits=str(log.cell(log.max_row,1).value).split('.'); bits[-1]=str(int(bits[-1])+1)
log.append(['.'.join(bits),'2026-09-30','SKYLINE白色灰阶压暗','关卡场景 / 开放世界',asset+' v004，最浅两档下调两档；灯光、相机、几何保持。','仅源完成；未导入运行时。','Codex']); w.save(path)
check=load_workbook(path)['资产主表']; assert all(check.cell(r,c).value==v for (r,c),v in before.items())
bl=json.loads(baseline.read_text(encoding='utf8')); previous=dict(bl['assets']); rv=dict(read_source_rows(check))[row]; bl['assets'][asset]={'v':_row_digest(rv),'c':'场景','d':'scenes'}
rows=[]
for d in idx['domains']: rows.extend(read_source_rows(load_workbook(R/idx['ledger_dir']/d['file'])['资产主表']))
bl['column_digests']={str(c):col_digest(rows,c) for c in CONTENT_COLUMNS}; bl['category_counts']=dict(Counter(v[2] for _,v in rows))
assert all(bl['assets'][k]==v for k,v in previous.items() if k!=asset)
baseline.write_text(json.dumps(bl,ensure_ascii=False,indent=1)+'\n',encoding='utf8')
report=dict(asset_id=asset,row=row,ledger=path.relative_to(R).as_posix(),source_version='v004',source_sha256=s.cell(row,20).value,other_asset_content_unchanged=True,other_asset_fingerprints_unchanged=True,status='Blender源已完成')
(folder/'qa/ledger_registration.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf8')
print('SKYLINE08_GRAY_V004_LEDGER_UPDATED',row)
