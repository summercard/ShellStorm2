"""Advance only Tower04 source metadata; preserve concurrent unrelated assets."""
import hashlib
import json
import re
import shutil
import sys
from pathlib import Path
from collections import Counter
from openpyxl import load_workbook
R=Path(__file__).resolve().parents[2]
sys.path.insert(0,str(R/'tools/asset_pipeline'))
from split_asset_ledger import FIRST_DATA_ROW,CONTENT_COLUMNS,read_source_rows,_row_digest,col_digest
F=R/'assets/art/environments/open_world/source/tower_04/v003'
cat=json.loads((F/'catalog.json').read_text(encoding='utf8'))
for name in ('source_audit','palette_validation','reference_plan_audit'):
    assert json.loads((F/'qa'/f'{name}.json').read_text(encoding='utf8'))['passed'], name
index_path=R/'assets/registry/ledger_index.json'
idx=json.loads(index_path.read_text(encoding='utf8'))
domain=next(d for d in idx['domains'] if d['key']=='scenes')
path=R/idx['ledger_dir']/domain['file']; base=R/'assets/registry/ledger_split_baseline.json'
digest=lambda p:hashlib.sha256(p.read_bytes()).hexdigest()
start={p:digest(p) for p in (path,base,index_path)}
refresh='--refresh-final-source' in sys.argv
prior_report=json.loads((F/'qa/ledger_registration.json').read_text(encoding='utf8')) if refresh else None
backup=R.parent/('_scratch/tower04_v003_sha_backup' if refresh else '_scratch/tower04_v003_ledger_backup'); backup.mkdir(parents=True,exist_ok=True)
for p in (path,base):
    assert not (backup/p.name).exists(),'Transaction already exists; do not replay'
    shutil.copy2(p,backup/p.name)
w=load_workbook(path); s=w['资产主表']; asset=cat['asset_id']
rows=[r for r in range(FIRST_DATA_ROW,s.max_row+1) if s.cell(r,1).value==asset]
assert len(rows)==1; row=rows[0]
assert s.cell(row,13).value==('v003' if refresh else 'v002')
if refresh:
    assert s.cell(row,15).value==cat['source_blend'] and s.cell(row,20).value==prior_report['source_sha256'],'Concurrent source SHA edit'
else:
    assert digest(R/s.cell(row,15).value)==s.cell(row,20).value,'Previous source drift'
source_sha=digest(R/cat['source_blend'])
allowed={13:'v003',14:'150×50m；五层×5m；天台25m；90独立包；4原材质；第一步仅校正2m跑道回环、分支及绿化布局',15:cat['source_blend'],16:'用户最新天台功能流线图；docs/v0.1/design/tower04_mall_source.md r3',17:'塔4;商场;天台路线;广场回环;绿岛绕行;玻璃亭支路;后侧花园;草坪庭院',20:source_sha,25:'源v003：第一步只调整天台流线和绿化，含平台局部悬挑及附着边缘护栏；其余118个对象及原四材质/公共贴图/import签名不变，未启用MipMap。两个回环和亭入口支路连通，绕开原有设施。源和逐面UV通过，未导出GLB/PackedScene、碰撞/LOD/导航未制作。v001/v002保留；图纸非CAD、裁切外沿用原端头，其它设施暂不优化。'}
if refresh: allowed={20:source_sha}
before={(r,c):(s.cell(r,c).value,s.cell(r,c).style_id) for r in range(1,s.max_row+1) for c in range(1,26)}
other={sheet.title:[tuple(c.value for c in cells) for cells in sheet] for sheet in w if sheet.title not in ('资产主表','域变更日志')}
for c,value in allowed.items(): s.cell(row,c,value)
log=w['域变更日志']; m=re.fullmatch(r'v(\d+)\.(\d+)\.(\d+)',str(log.cell(log.max_row,1).value)); assert m
if not refresh: log.append([f'v{m[1]}.{m[2]}.{int(m[3])+1}','2026-10-01','塔4第一步天台流线','关卡场景 / 开放世界',asset+'：仅天台跑道回环、分支、绿化位置与附着边界调整，其它设施保持v002。','源v003；原四材质与MipMap不变；未导入运行时。','Codex'])
assert all(digest(p)==h for p,h in start.items()),'Concurrent ledger transaction; no overwrite'
w.save(path)
again=load_workbook(path); ss=again['资产主表']; diffs=[]
for (r,c),(v,style) in before.items():
    assert ss.cell(r,c).style_id==style
    if ss.cell(r,c).value!=v:
        assert r==row and c in allowed,('Unrelated cell changed',r,c)
        diffs.append((r,c))
assert set(c for _,c in diffs)==set(allowed)
assert all([tuple(c.value for c in cells) for cells in again[name]]==values for name,values in other.items())
bl=json.loads(base.read_text(encoding='utf8')); prior=dict(bl['assets'])
bl['assets'][asset]={'v':_row_digest(dict(read_source_rows(ss))[row]),'c':'场景','d':'scenes'}
all_rows=[]
for d in idx['domains']:
    all_rows.extend(read_source_rows(load_workbook(R/idx['ledger_dir']/d['file'])['资产主表']))
bl['column_digests']={str(c):col_digest(all_rows,c) for c in CONTENT_COLUMNS}
bl['category_counts']=dict(Counter(v[2] for _,v in all_rows))
assert all(bl['assets'][k]==v for k,v in prior.items() if k!=asset)
assert digest(base)==start[base] and digest(index_path)==start[index_path],'Concurrent baseline edit; preserve it'
base.write_text(json.dumps(bl,ensure_ascii=False,indent=1)+'\n',encoding='utf8')
report=dict(asset_id=asset,row=row,source_version='v003',source_sha256=source_sha,changed_cells=[ss.cell(r,c).coordinate for r,c in diffs],other_rows_and_formulas_unchanged=True,styles_unchanged=True,unrelated_sheets_unchanged=True,other_baseline_assets_unchanged=True,asset_count_unchanged=bl['asset_count']==len(prior),prefab_pages_not_modified=True,backup=str(backup))
if refresh:
    report.update(initial_changed_cells=prior_report['changed_cells'],initial_backup=prior_report['backup'],final_curve_tangent_sha_only=True)
(F/'qa/ledger_registration.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf8')
print(json.dumps(report,ensure_ascii=False,indent=2))
