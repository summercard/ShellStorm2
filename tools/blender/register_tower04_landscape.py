"""Register five skyline-level source assets; preserve unrelated cells and sheet digests."""
from pathlib import Path
import sys,json,hashlib,shutil,re,os
from copy import copy
from collections import Counter
from openpyxl import load_workbook
from zipfile import ZipFile
from xml.etree import ElementTree as ET
R=Path(__file__).resolve().parents[2];sys.path.insert(0,str(R/'tools/asset_pipeline'))
from split_asset_ledger import FIRST_DATA_ROW,CONTENT_COLUMNS,read_source_rows,_row_digest,col_digest,dedupe_key_formula,dedupe_result_formula
O=R/'assets/art/environments/open_world/source/landscape_tower04/v001'
cat=json.loads((O/'catalog.json').read_text(encoding='utf8'))
for report in ['source_audit','palette_validation']:assert json.loads((O/'qa'/f'{report}.json').read_text(encoding='utf8'))['passed']
sha=lambda p:hashlib.sha256(p.read_bytes()).hexdigest()
assert sha(R/cat['source_blend'])==cat['source_sha256']
idx=json.loads((R/'assets/registry/ledger_index.json').read_text(encoding='utf8'));domain=next(d for d in idx['domains'] if d['key']=='scenes')
path=R/idx['ledger_dir']/domain['file'];baseline=R/'assets/registry/ledger_split_baseline.json'
backup=R.parent/'_scratch/landscape04_ledger_backup';backup.mkdir(exist_ok=True)
start={p:sha(p) for p in (path,baseline)}
for p in start:
 assert not (backup/p.name).exists(),'Do not replay an already registered transaction'
 shutil.copy2(p,backup/p.name)
w=load_workbook(path);s=w['资产主表'];old=s.max_row
before={(r,c):s.cell(r,c).value for r in range(1,old+1) for c in range(1,26)}
other={sn:tuple(tuple(row) for row in w[sn].values) for sn in w.sheetnames if sn not in ['资产主表','总览','域变更日志']}
registered=[]
for i,g in enumerate(cat['groups'],1):
 aid=g['asset_id'];assert not any(s.cell(r,1).value==aid for r in range(FIRST_DATA_ROW,old+1))
 r=s.max_row+1
 vals=[aid,'景观建筑 塔4周边 '+g['name'],'场景','environment_kit_3d',f'open_world_landscape04_{i:02d}','building_source',None,'Top3D / Blender Z-up','default','大地图场景景观；与SKYLINE同级；塔4周边远景布景','Blender源已完成','P1','v001',f"{g['buildings']}栋；{g['instances']}模块实例；{g['triangles']}三角面（全部楼层、残顶和藤蔓）；每组≤3000；4层楼身模块＋3种大片藤蔓；共享色盘标准材质",cat['source_blend'],'docs/v0.1/development/2026-10-02_tower04_landscape_buildings.md','景观建筑;塔4;SKYLINE同级;低面数;废墟;破败楼房;藤蔓',None,None,cat['source_sha256'],'Codex','2026-10-02','用户参考图；项目内原创程序建模；无第三方模型下载',f'LANDSCAPE04-{i:02d}',f"独立母版中的第{i:02d}组，Collection={g['collection']}；独立场景原点居中、底部Z=0。五个资产共享同一母版文件，哈希相同属明确共源，非重复资产。制作方自检通过；用户视觉待复核。未导出GLB、未生成PackedScene、未摆入Godot大地图。"]
 for c,v in enumerate(vals,1):s.cell(r,c,v);s.cell(r,c)._style=copy(s.cell(old,c)._style)
 s.row_dimensions[r].height=s.row_dimensions[old].height
 registered.append({'asset_id':aid,'row':r,'triangles':g['triangles']})
last=s.max_row
for r in range(FIRST_DATA_ROW,last+1):
 if r>old:s.cell(r,18,dedupe_key_formula(r))
 s.cell(r,19,dedupe_result_formula(r,last))
for row in w['总览']:
 for cell in row:
  if cell.data_type=='f':cell.value=re.sub(r'(\$[A-Z]+\$)'+str(old)+r'\b',lambda m:m[1]+str(last),cell.value)
for dv in s.data_validations.dataValidation:
 ranges=[]
 for q in dv.sqref.ranges:
  q=copy(q)
  if q.max_row==old:q.max_row=last
  ranges.append(str(q))
 dv.sqref=' '.join(ranges)
s.auto_filter.ref=f'A{FIRST_DATA_ROW-1}:Y{last}'
log=w['域变更日志'];m=re.fullmatch(r'v(\d+)\.(\d+)\.(\d+)',str(log.cell(log.max_row,1).value));assert m
log.append([f'v{m[1]}.{m[2]}.{int(m[3])+1}','2026-10-02','新增五组远景景观建筑','关卡场景 / 开放世界','LANDSCAPE04-01至05；2098/2334/2450/2098/2991三角；独立母版，SKYLINE同级','Blender源已完成；用户视觉待复核；未导入Godot','Codex'])
temp=backup/'candidate_scene_ledger.xlsx';w.save(temp);a=load_workbook(temp);ss=a['资产主表']
assert all(ss.cell(r,c).value==v for (r,c),v in before.items() if not(c==19 and r>=FIRST_DATA_ROW))
assert all(tuple(tuple(row) for row in a[sn].values)==values for sn,values in other.items())
with ZipFile(temp) as z:
 root=ET.fromstring(z.read('xl/worksheets/sheet3.xml'));ns={'s':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
 dv_xml=[e.attrib['sqref'] for e in root.findall('.//s:dataValidation',ns)]
 assert all(any(f'{col}{last}' in v for v in dv_xml) for col in ['C','K','L']),dv_xml
bl=json.loads(baseline.read_text(encoding='utf8'));previous=dict(bl['assets']);source_rows=dict(read_source_rows(ss))
for entry in registered:bl['assets'][entry['asset_id']]={'v':_row_digest(source_rows[entry['row']]),'c':'场景','d':'scenes'}
bl['asset_count']=len(bl['assets']);rows=[]
for d in idx['domains']:
 book=a if d['key']=='scenes' else load_workbook(R/idx['ledger_dir']/d['file'])
 rows.extend(read_source_rows(book['资产主表']))
bl['column_digests']={str(c):col_digest(rows,c) for c in CONTENT_COLUMNS};bl['category_counts']=dict(Counter(v[2] for _,v in rows))
assert all(bl['assets'][k]==v for k,v in previous.items())
btemp=backup/'candidate_baseline.json';btemp.write_text(json.dumps(bl,ensure_ascii=False,indent=1)+'\n',encoding='utf8')
assert all(sha(p)==h for p,h in start.items()),'Concurrent ledger change'
os.replace(temp,path);os.replace(btemp,baseline)
report={'registered':registered,'prior_cells_unchanged_except_dedupe_range':True,'other_sheets_unchanged':True,'prior_fingerprints_unchanged':True,'validation_ranges':dv_xml,'status':'Blender源已完成','source_sha256':cat['source_sha256']}
(O/'qa/ledger_registration.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf8')
print('LANDSCAPE_LEDGER_REGISTERED',registered)
