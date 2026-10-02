"""Record the revised dense Blender source; distinguish self-check from user approval."""
import hashlib
import json
import re
import shutil
import sys
from pathlib import Path
from collections import Counter
from openpyxl import load_workbook
R=Path(__file__).resolve().parents[2]
sys.path.insert(0,str(R/'tools/asset_pipeline'))
from split_asset_ledger import FIRST_DATA_ROW,CONTENT_COLUMNS,read_source_rows,_row_digest,col_digest
F=R/'assets/art/environments/open_world/source/tower_04/v009'
cat=json.loads((F/'catalog.json').read_text(encoding='utf8'))
for name in ('lower_structure_audit','palette_validation','visual_review','metadata_finalization'):
    assert json.loads((F/'qa'/f'{name}.json').read_text(encoding='utf8'))['passed'],name
index_path=R/'assets/registry/ledger_index.json'; idx=json.loads(index_path.read_text(encoding='utf8'))
domain=next(d for d in idx['domains'] if d['key']=='scenes')
path=R/idx['ledger_dir']/domain['file']; base=R/'assets/registry/ledger_split_baseline.json'
digest=lambda p:hashlib.sha256(p.read_bytes()).hexdigest()
start={p:digest(p) for p in (path,base,index_path)}
backup=R.parent/'_scratch/tower04_v009_ledger_backup'; backup.mkdir(parents=True,exist_ok=True)
for p in (path,base):
    assert not (backup/p.name).exists(),'Do not replay a completed transaction'
    shutil.copy2(p,backup/p.name)
w=load_workbook(path); s=w['资产主表']; asset=cat['asset_id']
rows=[r for r in range(FIRST_DATA_ROW,s.max_row+1) if s.cell(r,1).value==asset]
assert len(rows)==1; row=rows[0]
assert s.cell(row,13).value=='v008'
assert digest(R/s.cell(row,15).value)==s.cell(row,20).value,'Prior source drift'
source_sha=digest(R/cat['source_blend'])
allowed={13:'v009',14:'顶层150×50m/屋面25m锁定；下部错层架空；1911独立包/41组件族；375新结构包XY≤10m；原4材质',15:cat['source_blend'],16:'用户红线范围图与商城结构参考；docs/v0.1/design/tower04_mall_source.md r8',17:'塔4;末世商城;架空大柱;弧形商场;环形退台;内退窗带;下层联系走廊;顶层锁定',20:source_sha,22:'2026-10-02',25:'源v009：仅重建顶层以下。左楼9m架空底板、14.3m退台、跨层圆角肋；中央大柱和下层联系走廊；右翼厚混凝土实墙与内退破损窗带，移除旧五层实心幕墙。2942锁定对象完整签名一致，原694地板/顶层植被/设施/流线/四材质/公共色盘/MipMap不变；88下部植被包随宿主调整，119旧悬浮生长包退役。参考尺寸按既有25m屋面推定，遮挡背面为必要补全；制作方源验收与看图复核完成，用户视觉待确认。约499万输出面；未导出GLB/PackedScene或制作LOD/碰撞/导航，未验收运行性能；v001–v008保留。'}
assert cat['package_count']==1911
for name in ('lower_structure_audit','metadata_finalization','visual_review'):
    assert json.loads((F/'qa'/f'{name}.json').read_text(encoding='utf8'))['source_sha256']==source_sha
before={(r,c):(s.cell(r,c).value,s.cell(r,c).style_id) for r in range(1,s.max_row+1) for c in range(1,26)}
other={sheet.title:[tuple(c.value for c in cells) for cells in sheet] for sheet in w if sheet.title not in ('资产主表','域变更日志')}
for c,value in allowed.items(): s.cell(row,c,value)
log=w['域变更日志']; m=re.fullmatch(r'v(\d+)\.(\d+)\.(\d+)',str(log.cell(log.max_row,1).value)); assert m
log.append([f'v{m[1]}.{m[2]}.{int(m[3])+1}','2026-10-02','塔4参考图下部架空退台结构重建','关卡场景 / 开放世界',asset+'：v009顶层完整锁定，重建左楼架空退台、中央下层联系走廊及右翼混凝土窗墙。','制作方自检完成；用户视觉待复核；原四材质/MipMap不变；未导入运行时。','Codex'])
assert all(digest(p)==h for p,h in start.items()),'Concurrent ledger write; no overwrite'
w.save(path)
again=load_workbook(path); ss=again['资产主表']; diffs=[]
for (r,c),(v,style) in before.items():
    assert ss.cell(r,c).style_id==style
    if ss.cell(r,c).value!=v:
        assert r==row and c in allowed,('Unrelated cell changed',r,c)
        diffs.append((r,c))
assert set(c for _,c in diffs)=={c for c,v in allowed.items() if before[(row,c)][0]!=v}
assert all([tuple(c.value for c in cells) for cells in again[name]]==values for name,values in other.items())
bl=json.loads(base.read_text(encoding='utf8')); prior=dict(bl['assets'])
bl['assets'][asset]={'v':_row_digest(dict(read_source_rows(ss))[row]),'c':'场景','d':'scenes'}
all_rows=[]
for d in idx['domains']: all_rows.extend(read_source_rows(load_workbook(R/idx['ledger_dir']/d['file'])['资产主表']))
bl['column_digests']={str(c):col_digest(all_rows,c) for c in CONTENT_COLUMNS}
bl['category_counts']=dict(Counter(v[2] for _,v in all_rows))
assert all(bl['assets'][k]==v for k,v in prior.items() if k!=asset)
assert digest(base)==start[base] and digest(index_path)==start[index_path],'Concurrent baseline change'
base.write_text(json.dumps(bl,ensure_ascii=False,indent=1)+'\n',encoding='utf8')
report=dict(asset_id=asset,row=row,source_version='v009',source_sha256=source_sha,changed_cells=[ss.cell(r,c).coordinate for r,c in diffs],other_rows_and_formulas_unchanged=True,styles_unchanged=True,unrelated_sheets_unchanged=True,other_baseline_assets_unchanged=True,asset_count_unchanged=bl['asset_count']==len(prior),prefab_pages_not_modified=True,user_visual_confirmation='pending',backup=str(backup))
(F/'qa/ledger_registration.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf8')
print(json.dumps(report,ensure_ascii=False,indent=2))
