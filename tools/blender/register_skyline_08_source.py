"""Append only the new source asset row and extend formulas, keeping prior edits."""
from pathlib import Path
import sys,json,hashlib,shutil,re
from copy import copy
from collections import Counter
from openpyxl import load_workbook
R=Path(__file__).resolve().parents[2]; sys.path.insert(0,str(R/'tools/asset_pipeline'))
from split_asset_ledger import FIRST_DATA_ROW,read_source_rows,_row_digest,col_digest,CONTENT_COLUMNS,dedupe_key_formula,dedupe_result_formula
idx=json.loads((R/'assets/registry/ledger_index.json').read_text(encoding='utf8')); domain=next(d for d in idx['domains'] if d['key']=='scenes')
path=R/idx['ledger_dir']/domain['file']; baseline=R/'assets/registry/ledger_split_baseline.json'
asset='ENV-OPENWORLD-SKYLINE08'; folder=R/'assets/art/environments/open_world/source/skyline_08/v003'; blend=folder/'SKYLINE大楼_8层_精细天台_v003.blend'
assert json.loads((folder/'qa/source_audit.json').read_text(encoding='utf8'))['passed']
assert json.loads((folder/'qa/palette_validation.json').read_text(encoding='utf8'))['passed']
backup=R/'_scratch/skyline08_ledger_backup'; backup.mkdir(exist_ok=True)
for p in [path,baseline]:
 if not (backup/p.name).exists(): shutil.copy2(p,backup/p.name)
w=load_workbook(path); s=w['资产主表']; old=s.max_row
assert not any(s.cell(r,1).value==asset for r in range(FIRST_DATA_ROW,old+1)), 'Already registered; do not append duplicate'
before={(r,c):s.cell(r,c).value for r in range(1,old+1) for c in CONTENT_COLUMNS}
new=old+1
values=[asset,'SKYLINE大楼 8层 精细天台','场景','environment_kit_3d','open_world_skyline_08','building_source',None,'Top3D / Blender Z-up','default','独立开放世界建筑美术源；不替换现有塔楼','Blender源已完成','P1','v003','主体32×24m；8层×4m，屋顶Z=32m；广告牌最高45.68m；252包；4角色公共色盘；6张真实渲染预览',blend.relative_to(R).as_posix(),'用户参考图和8层要求；docs/v0.1/development/2026-09-30_skyline08_building_source.md','SKYLINE;8层;精细天台;破损广告牌;灯泡招牌;风格化建筑',None,None,hashlib.sha256(blend.read_bytes()).hexdigest(),'Codex','2026-09-30','用户提供参考图，据图建模；无第三方模型下载','SKYLINE08','仅Blender源完成；未导出GLB/接入Godot，未制作碰撞/LOD。下方7层重复窗墙结构；层高为美术推定，不套用主塔12m战斗层契约。']
for c,value in enumerate(values,1): s.cell(new,c,value); s.cell(new,c)._style=copy(s.cell(old,c)._style)
s.row_dimensions[new].height=s.row_dimensions[old].height
s.cell(new,18,dedupe_key_formula(new))
for r in range(FIRST_DATA_ROW,new+1): s.cell(r,19,dedupe_result_formula(r,new))
for row in w['总览']:
 for cell in row:
  if cell.data_type=='f': cell.value=re.sub(r'(\$[A-Z]+\$)'+str(old)+r'\b',lambda m:m[1]+str(new),cell.value)
for dv in s.data_validations.dataValidation:
 ranges=[]
 for rng in dv.sqref.ranges:
  q=copy(rng)
  if q.max_row==old: q.max_row=new
  ranges.append(str(q))
 dv.sqref=' '.join(ranges)
s.auto_filter.ref=f'A{FIRST_DATA_ROW-1}:Y{new}'
log=w['域变更日志']; bits=str(log.cell(log.max_row,1).value).split('.'); bits[-1]=str(int(bits[-1])+1)
log.append(['.'.join(bits),'2026-09-30','新增Blender建筑源','关卡场景 / 开放世界',asset+'：8层重复窗墙与精细天台，252包，源v003。','源完成；未导入运行时。','Codex'])
w.save(path)
again=load_workbook(path); ss=again['资产主表']
assert all(ss.cell(r,c).value==v for (r,c),v in before.items()), 'Prior asset content changed'
bl=json.loads(baseline.read_text(encoding='utf8')); old_assets=dict(bl['assets'])
rowvalues=dict(read_source_rows(ss))[new]
bl['assets'][asset]={'v':_row_digest(rowvalues),'c':'场景','d':'scenes'}; bl['asset_count']=len(bl['assets'])
rows=[]
for d in idx['domains']:
 book=load_workbook(R/idx['ledger_dir']/d['file']); rows.extend(read_source_rows(book['资产主表']))
bl['column_digests']={str(c):col_digest(rows,c) for c in CONTENT_COLUMNS}; bl['category_counts']=dict(Counter(v[2] for _,v in rows))
assert all(bl['assets'][k]==v for k,v in old_assets.items())
baseline.write_text(json.dumps(bl,ensure_ascii=False,indent=1)+'\n',encoding='utf8')
report=dict(asset_id=asset,row=new,ledger=path.relative_to(R).as_posix(),source_version='v003',source_sha256=values[19],prior_asset_content_unchanged=True,prior_asset_fingerprints_unchanged=True,baseline_count=bl['asset_count'],status='Blender源已完成')
(folder/'qa/ledger_registration.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf8')
print('SKYLINE08_LEDGER_REGISTERED',new,asset)
